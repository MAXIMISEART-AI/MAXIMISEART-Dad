"""Step 3 - Excel extract + safe buffer append.

Rule #1 Floor Never Drops: this module never writes into tata's existing sheets.
It appends normalized rows only to MAXIMISEART_DAILY_APPEND.
"""
from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from lib.excel_safe_writer import snapshot, wait_for_unlock

BUFFER_SHEET_NAME = "MAXIMISEART_DAILY_APPEND"
HEADERS = ["Data", "Pracownik", "Kod", "Ilość", "Lokacja", "Uwagi", "Status", "Email ID", "Timestamp"]


class ExcelLockedError(RuntimeError):
    """Raised when Excel/OneDrive lock files prevent a safe write."""


def append_rows(
    excel_path,
    routed_emails: dict,
    code_mapping: dict,
    schema: dict,
    *,
    snapshots_dir: Path | None = None,
    employees_config: dict | None = None,
    date: str | None = None,
    timestamp: str | None = None,
) -> list[dict]:
    """Append extracted rows to MAXIMISEART_DAILY_APPEND and return added rows."""
    path = Path(excel_path)
    if not path.exists():
        raise FileNotFoundError(path)

    retry_intervals = schema.get("retry_on_lock", [120, 300, 600]) if schema else [120, 300, 600]
    if not wait_for_unlock(path, retry_intervals):
        raise ExcelLockedError(
            "Excel jest otwarty albo OneDrive blokuje plik. Zamknij Excel i uruchom ponownie."
        )

    rows = _extract_rows(routed_emails, code_mapping, employees_config or {}, date, timestamp)
    if not rows:
        return []

    snapshot(path, snapshots_dir or path.parent / "state" / "excel-snapshots")

    wb = load_workbook(path)
    ws = _ensure_buffer_sheet(wb)
    existing_keys = _existing_row_keys(ws)

    appended: list[dict] = []
    for row in rows[: _max_rows_per_run(schema)]:
        key = _row_key(row)
        if key in existing_keys:
            continue
        ws.append([row.get(header) for header in HEADERS])
        existing_keys.add(key)
        appended.append(row)

    wb.save(path)
    return appended


def _ensure_buffer_sheet(wb):
    if BUFFER_SHEET_NAME in wb.sheetnames:
        ws = wb[BUFFER_SHEET_NAME]
    else:
        ws = wb.create_sheet(BUFFER_SHEET_NAME)

    first_row = [cell.value for cell in ws[1]]
    if first_row != HEADERS:
        if ws.max_row == 1 and all(value is None for value in first_row):
            for idx, header in enumerate(HEADERS, start=1):
                ws.cell(row=1, column=idx, value=header)
        else:
            raise ValueError(f"Arkusz {BUFFER_SHEET_NAME} ma niezgodne nagłówki.")
    return ws


def _extract_rows(
    routed_emails: dict,
    code_mapping: dict,
    employees_config: dict,
    date: str | None,
    timestamp: str | None,
) -> list[dict]:
    rows: list[dict] = []
    employees = {emp.get("slug"): emp for emp in employees_config.get("employees", [])}
    code_index = _active_codes(code_mapping)
    now = timestamp or datetime.now().isoformat(timespec="seconds")

    for slug, emails in routed_emails.items():
        employee = employees.get(slug, {})
        employee_name = employee.get("name") or _name_from_slug(slug)
        for email in emails:
            email_date = date or str(email.get("received_at_iso", ""))[:10] or datetime.now().date().isoformat()
            for item in _extract_work_items(email.get("body_text", ""), code_index):
                rows.append({
                    "Data": email_date,
                    "Pracownik": employee_name,
                    "Kod": item["code"],
                    "Ilość": item["quantity"],
                    "Lokacja": item["location"],
                    "Uwagi": item["notes"],
                    "Status": item["status"],
                    "Email ID": email.get("id", ""),
                    "Timestamp": now,
                })
    return rows


def _extract_work_items(body: str, code_index: list[dict]) -> list[dict[str, Any]]:
    lines = [_clean_line(line) for line in body.splitlines()]
    candidates = [line for line in lines if line and _looks_like_work_line(line)]
    if not candidates and body.strip():
        candidates = [_clean_line(body)]

    items = []
    for line in candidates:
        code = _match_code(line, code_index)
        items.append({
            "code": code,
            "quantity": _extract_quantity(line),
            "location": _extract_location(line),
            "notes": line,
            "status": "ANSWER" if code is not None else "ASK",
        })
    return items


def _clean_line(line: str) -> str:
    return re.sub(r"^\s*[-*•]\s*", "", line.strip())


def _looks_like_work_line(line: str) -> bool:
    text = line.lower()
    return bool(re.search(r"\d", text)) and not text.startswith(("cześć", "czesc", "pozdrawiam"))


def _active_codes(code_mapping: dict) -> list[dict]:
    return [
        code for code in code_mapping.get("codes", [])
        if code.get("active", True) and code.get("code") is not None
    ]


def _match_code(line: str, code_index: list[dict]) -> int | None:
    text = line.lower()
    for code in code_index:
        keywords = [str(keyword).lower() for keyword in code.get("keywords", []) or []]
        name = str(code.get("name", "")).lower()
        if any(keyword and keyword in text for keyword in keywords) or (name and name in text):
            return int(code["code"])
    return None


def _extract_quantity(line: str) -> int | float | None:
    match = re.search(r"(\d+(?:[,.]\d+)?)\s*x\b", line, flags=re.IGNORECASE)
    if match is None:
        match = re.search(r"\b(\d+(?:[,.]\d+)?)\b", line)
    if match is None:
        return None
    raw = match.group(1).replace(",", ".")
    value = float(raw)
    return int(value) if value.is_integer() else value


def _extract_location(line: str) -> str:
    match = re.search(r"\b(?:na ulicy|ulicy|ul\.|przy|na)\s+(.+)$", line, flags=re.IGNORECASE)
    if match is None:
        return ""
    location = match.group(1).strip()
    location = re.sub(r"\s*\([^)]*\)\s*$", "", location).strip()
    return location


def _existing_row_keys(ws) -> set[tuple]:
    keys = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        row_dict = dict(zip(HEADERS, row))
        if any(value is not None for value in row):
            keys.add(_row_key(row_dict))
    return keys


def _row_key(row: dict) -> tuple:
    return (
        row.get("Data"),
        row.get("Pracownik"),
        row.get("Kod"),
        row.get("Ilość"),
        row.get("Lokacja"),
        row.get("Uwagi"),
        row.get("Email ID"),
    )


def _max_rows_per_run(schema: dict) -> int:
    try:
        return int(schema.get("max_rows_per_run", 200))
    except (TypeError, ValueError):
        return 200


def _name_from_slug(slug: str) -> str:
    return " ".join(part.capitalize() for part in slug.split("-"))
