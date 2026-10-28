"""INTERMAX (Rejestr zleceń) w Uporządkowaniu INTERMAX: fakty arkusza zleceń i zapis dnia podmianą wierszy XML.

``Register.read`` czyta bajty INTERMAX przez jedyne miejsce odczytu bajtów skoroszytu (``workbook_files``) i
znajduje arkusze zleceń po nagłówkach z wiersza 1, tak jak Przygotowanie pliku zbiorczego (``COLUMNS``), z
kolumnami ``AKCJE`` i ``UWAGI``. Arkusz zleceń daje swoje wiersze od dołu jako fakty (``raw_orders.RowFacts``),
bez openpyxl: ~42 tys. wierszy taty to dziesiątki MB XML, a rozpoznanie potrzebuje tylko dołu arkusza.

``save`` zapisuje plan dnia (ADR 0019, wzorem kopii Wzorca z ADR 0008). Zmieniają się tylko:

- w XML arkusza zleceń wiersze pod wierszem, pod którym staje dzień (``RawBlock.anchor``): wiersz zlecenia
  dostaje komórki Surowego zlecenia (w przesuniętym pierwszym wierszu o dwie kolumny w prawo) w kolumnach od A
  do ``UWAGI`` z wyglądem reszty INTERMAX i Kolorem operatora z planu dnia, a kolumny z aplikacji za ``UWAGI``
  znikają; zlecenie z Wpisem urządzenia i czynności INTERMAX ma w ``URZĄDZENIA`` i ``CZYNNOŚCI`` tekst Wpisu
  zamiast wartości z aplikacji; pusty wiersz między segmentami ma puste komórki od A do ``UWAGI`` z tym
  samym wyglądem; inny pusty wiersz, który był pusty, zostaje bajt w bajt, a wiersz, z którego zlecenie odeszło, i
  wiersz, którego nie było, są pustym wierszem taty; wygląd ``<row>`` (styl wiersza, wysokość) daje ``_RowLooks``;
- ``ref`` Hiperłączy z tych wierszy, które idą razem ze swoją komórką, i ``<dimension>``, gdy dzień sięga dalej;
- widok arkusza zleceń (``<sheetViews>``): karta zaznaczona, a górny wiersz widoku pod zamrożonym nagłówkiem i
  zaznaczenie na wierszu widoku z planu dnia (``DayPlan.view_row``);
- aktywna karta skoroszytu (``activeTab`` w ``xl/workbook.xml``) na arkuszu zleceń i znacznik zaznaczenia karty
  (``tabSelected``) zdjęty z pozostałych arkuszy;
- w ``xl/styles.xml`` czcionki, wypełnienia, obramowanie i style komórek dopisane na końcu swoich list.

Każda inna część pliku i każdy bajt nad tymi wierszami zostają. Arkusz, w którego zmienianych wierszach jest
formuła albo zakres (scalenie, formatowanie warunkowe, sprawdzanie poprawności, autofiltr, Hiperłącze na kilka
komórek albo za ``UWAGI``), to ``UnsupportedShape``. Wynik trafia do pliku tymczasowego obok INTERMAX, zostaje
odczytany jeszcze raz i porównany z oryginałem, a komórki Wpisu z Wpisem z planu dnia; różnica to
``WriteMismatch``, a INTERMAX zostaje nietknięty. Dopiero po zgodnym
sprawdzeniu ``before_replace`` (Kopia INTERMAX) i atomowa podmiana. Podmianę, której system odmawia, gdy Excel nie
trzyma INTERMAX (OneDrive go synchronizuje), ``save`` ponawia przez ``_RETRY_S`` sekund zegara procesu (``Waiting``),
a potem zgłasza ``HeldByOneDrive``. Przed każdą ponowną próbą czyta INTERMAX jeszcze raz: gdy OneDrive wgrał w
przerwie inną wersję, to też ``HeldByOneDrive``, bez podmiany.

``restore`` podmienia tą samą drogą cały INTERMAX bajtami Kopii INTERMAX (Przywracanie Kopii INTERMAX): plik
tymczasowy obok musi mieć po odczycie te bajty, a ponowna próba podmienia tylko INTERMAX o bajtach sprzed przywrócenia.
"""

from __future__ import annotations

import html
import os
import re
import time
import zipfile
from functools import cached_property
from collections.abc import Callable, Iterator, Mapping, Sequence
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Any, Final, cast
from xml.etree import ElementTree

from openpyxl.styles.borders import Border
from openpyxl.styles.colors import Color
from openpyxl.styles.fills import Fill, PatternFill
from openpyxl.styles.fonts import Font
from openpyxl.styles.numbers import BUILTIN_FORMATS, is_date_format
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string, get_column_letter, range_boundaries
from openpyxl.utils.datetime import MAC_EPOCH, WINDOWS_EPOCH, from_excel
from openpyxl.workbook.properties import WorkbookProperties

from . import workbook_files
from .preparation_register import ACTIVITIES, COLUMNS, DATE, DEVICES, WORKER
from .raw_orders import SHIFT, CellFacts, DayPlan, IntermaxEntry, OperatorColor, OrderColumns, RawBlock, RowFacts
from .row_look import BORDER
from .workbook_files import WorkbookError, header_columns, is_blank
from .workbook_parts import MAIN, OFFICE_REL, IrregularSheetData, SheetData, SheetRow, office_document, relationships
from .workbook_save import WorkbookSaveError, write_atomically
from .worker_config import normalize_text

__all__ = ["OrderSheet", "Register", "UnsupportedShape", "WriteMismatch", "save"]

_FIRST: Final = "AKCJE"
_LAST: Final = "UWAGI"
_DATE_FORMAT: Final = 22  # m/d/yy h:mm, format daty uporządkowanych dni
_FACE: Final = "Calibri"
_SIZE: Final = 8
_WORKER_SIZE: Final = 9
_BLACK: Final = "FF000000"
_YELLOW: Final = "FFFF00"

_CELL = re.compile(rb'<c r="([A-Z]{1,3})(\d+)"([^>]*?)(?:/>|>(.*?)</c>)', re.S)
_STYLE = re.compile(rb'\ss="(\d+)"')
_TYPE = re.compile(rb'\st="(\w+)"')
_VALUE = re.compile(rb"<v(?:\s[^>]*)?>([^<]*)</v>")
_TEXT = re.compile(rb"<t(?:\s[^>]*)?>([^<]*)</t>")
_FORMULA = re.compile(rb"<f[\s>/]")
_DIMENSION = re.compile(rb'<dimension ref="([^"]*)"\s*/>')
_HYPERLINK_REF = re.compile(rb'(<hyperlink\b[^>]*?\sref=")([^"]*)(")')
_HYPERLINKS = re.compile(rb"<hyperlinks>.*?</hyperlinks>", re.S)
_RANGES = re.compile(rb'\s(?:ref|sqref)="([^"]*)"')
_HEADER_ROW = re.compile(rb'<row r="1"[^>]*?(?:/>|>.*?</row>)', re.S)
_ATTRIBUTE = re.compile(rb'\s([\w:]+)="([^"]*)"')
_ROW_LOOK: Final = (b"s", b"customFormat", b"ht", b"customHeight", b"x14ac:dyDescent")
"""Atrybuty ``<row>``, które dają wiersz wygląd: styl wiersza i wysokość; ``x14ac`` deklaruje arkusz taty."""
_SHEET_VIEWS: Final = re.compile(rb"<sheetViews\b.*?</sheetViews>", re.S)
_SHEET_VIEW: Final = re.compile(rb"<sheetView\b([^>]*?)(?:/>|>.*?</sheetView>)", re.S)
_PANE: Final = re.compile(rb"<pane\b([^>]*?)\s*/>")
_CONTRACT_NUMBER: Final = "NUMER UMOWY/ZLECENIA ORANGE"
"""Nagłówek docelowy kolumny ``NUMER UMOWY`` w ``COLUMNS`` Rejestru."""
_TAB_SELECTED: Final = re.compile(rb'(<sheetView\b[^>]*?)\stabSelected="(?:1|true)"')
_WORKBOOK_VIEW: Final = re.compile(rb"<workbookView\b([^>]*?)(/?>)")
_STYLE_LISTS: Final = (b"fonts", b"fills", b"borders", b"cellXfs")
"""Listy ``xl/styles.xml``, na których końcu zapis dopisuje wpisy."""
_RETRY_S: Final = 5.0
"""Tyle sekund ``save`` ponawia podmianę, której system odmawia, choć Excel nie trzyma INTERMAX."""
_PAUSE_S: Final = 0.5


class UnsupportedShape(Exception):
    """Arkusz zleceń ma w zmienianych wierszach kształt, którego zapis nie przenosi (formuła, zakres)."""


class WriteMismatch(Exception):
    """Plik zapisany przed podmianą różni się od oryginału czymś więcej niż uporządkowanym dniem."""


class HeldByOneDrive(Exception):
    """System odmawia podmiany INTERMAX dłużej niż krótka ponowna próba, a Excel go nie trzyma, albo OneDrive wgrał
    w przerwie między próbami inną wersję INTERMAX."""


@dataclass(frozen=True)
class Waiting:
    """Zegar procesu w sekundach i przerwa między próbami podmiany INTERMAX."""

    clock: Callable[[], float] = time.monotonic
    pause: Callable[[float], None] = time.sleep


@dataclass(frozen=True)
class _Styles:
    """Style komórek z ``xl/styles.xml``: format liczby, czcionka i wypełnienie każdego stylu."""

    number_formats: tuple[int, ...]
    faces: tuple[tuple[str | None, float | None], ...]
    yellow: tuple[bool, ...]
    date_formats: frozenset[int]
    fonts: tuple[Font, ...]
    fills: tuple[Fill, ...]
    borders: tuple[Border, ...]
    xfs: tuple[bytes, ...]

    @classmethod
    def of(cls, xml: bytes) -> _Styles:
        root = ElementTree.fromstring(xml)
        codes = {int(item.get("numFmtId", "0")): item.get("formatCode", "") for item in root.iter(f"{MAIN}numFmt")}
        fonts = [Font.from_tree(cast(Any, font)) for font in root.iterfind(f"{MAIN}fonts/{MAIN}font")]
        borders = [cast(Border, Border.from_tree(cast(Any, border))) for border in root.iterfind(f"{MAIN}borders/{MAIN}border")]
        fill_trees = list(root.iterfind(f"{MAIN}fills/{MAIN}fill"))
        fills = [_is_yellow(fill) for fill in fill_trees]
        xfs = list(root.iterfind(f"{MAIN}cellXfs/{MAIN}xf"))
        raw_xfs = re.findall(rb"<xf\b[^>]*?(?:/>|>.*?</xf>)", _style_list(xml, b"cellXfs").group(2), re.S)
        number_formats = tuple(int(xf.get("numFmtId", "0")) for xf in xfs)
        return cls(
            number_formats=number_formats,
            faces=tuple(_face(fonts[int(xf.get("fontId", "0"))]) for xf in xfs),
            yellow=tuple(fills[int(xf.get("fillId", "0"))] for xf in xfs),
            date_formats=frozenset(
                number for number in number_formats if is_date_format(codes.get(number) or BUILTIN_FORMATS.get(number, ""))
            ),
            fonts=tuple(fonts),
            fills=tuple(cast(Fill, Fill.from_tree(cast(Any, fill))) for fill in fill_trees),
            borders=tuple(borders),
            xfs=tuple(raw_xfs),
        )


def _face(font: Font) -> tuple[str | None, float | None]:
    return font.name, font.sz


def _is_yellow(fill: ElementTree.Element) -> bool:
    color = fill.find(f"{MAIN}patternFill[@patternType='solid']/{MAIN}fgColor")
    return color is not None and color.get("rgb", "").upper().endswith(_YELLOW)


@dataclass(frozen=True)
class _Cell:
    column: int
    attributes: bytes
    content: bytes | None

    @property
    def style(self) -> int:
        found = _STYLE.search(self.attributes)
        return int(found.group(1)) if found else 0

    @property
    def kind(self) -> bytes | None:
        found = _TYPE.search(self.attributes)
        return found.group(1) if found else None

    def xml(self, column: int, row: int, attributes: bytes | None = None) -> bytes:
        reference = get_column_letter(column).encode("ascii") + b"%d" % row
        attributes = self.attributes if attributes is None else attributes
        if self.content is None:
            return b'<c r="' + reference + b'"' + attributes + b"/>"
        return b'<c r="' + reference + b'"' + attributes + b">" + self.content + b"</c>"


def _cells(row_xml: bytes) -> list[_Cell]:
    return [
        _Cell(column_index_from_string(found.group(1).decode("ascii")), found.group(3), found.group(4))
        for found in _CELL.finditer(row_xml)
    ]


@dataclass(frozen=True)
class OrderSheet:
    """Arkusz zleceń INTERMAX: nazwa, część zip i kolumny po nagłówkach z wiersza 1."""

    name: str
    part: str
    columns: OrderColumns
    _register: Register

    @cached_property
    def sheet_xml(self) -> _SheetXml:
        """XML arkusza z indeksem wierszy, liczonym raz; ``UnsupportedShape``, gdy wiersze nie dają się rozdzielić."""

        xml = self._register._parts[self.part]
        try:
            return _SheetXml(xml, SheetData.of(xml))
        except IrregularSheetData as error:
            raise UnsupportedShape(str(error)) from error

    def rows_from_bottom(self) -> Iterator[RowFacts]:
        """Wiersze arkusza od ostatniego do drugiego jako fakty dla rozpoznania Surowych zleceń."""

        sheet = self.sheet_xml
        for row in reversed(sheet.data.rows):
            if row.number < 2:
                return
            yield self._register._facts(_row_tag(sheet.xml, row), row.number, _cells(sheet.xml[row.start : row.end]))


def _row_tag(xml: bytes, row: SheetRow) -> bytes:
    return xml[row.start : xml.index(b">", row.start) + 1]


def _row_attributes(tag: bytes) -> dict[bytes, bytes]:
    return dict(_ATTRIBUTE.findall(tag))


@dataclass(frozen=True)
class _SheetXml:
    xml: bytes
    data: SheetData


class Register:
    """INTERMAX odczytany raz: bajty, części zip, style, teksty wspólne i arkusze zleceń."""

    def __init__(self, path: Path, content: bytes) -> None:
        self.path = path
        self.content = content
        with workbook_files.unreadable_as_error(path):
            archive = zipfile.ZipFile(BytesIO(content))
            self._infos = archive.infolist()
            self._parts = {info.filename: archive.read(info) for info in self._infos}
            workbook = cast(str, office_document(self._parts))  # bez dokumentu: KeyError, czyli nieczytelny
            rels = relationships(self._parts, workbook)
            root = ElementTree.fromstring(self._parts[workbook])
            styles_part = next(rel.target for rel in rels.values() if rel.kind.endswith("/styles"))
            self._styles_part = styles_part
            self._styles = _Styles.of(self._parts[styles_part])
            strings = next((rel.target for rel in rels.values() if rel.kind.endswith("/sharedStrings")), None)
            self._shared = _shared_strings(self._parts[strings]) if strings else []
            properties = root.find(f"{MAIN}workbookPr")
            date1904 = properties is not None and cast(WorkbookProperties, WorkbookProperties.from_tree(cast(Any, properties))).date1904
            self._epoch = MAC_EPOCH if date1904 else WINDOWS_EPOCH
            sheets = [
                (sheet.get("name", ""), rels[sheet.get(f"{{{OFFICE_REL}}}id", "")].target)
                for sheet in root.iter(f"{MAIN}sheet")
            ]
            self._workbook_part = workbook
            self._sheet_parts = [part for _, part in sheets]
            self.order_sheets = [
                OrderSheet(name, part, columns, self)
                for name, part in sheets
                if (columns := self._order_columns(part)) is not None
            ]

    @classmethod
    def read(cls, path: Path) -> Register:
        """INTERMAX z dysku; ``WorkbookError``, gdy go nie ma, jest otwarty w Excelu, niedostępny albo nieczytelny."""

        return cls(path, workbook_files.workbook_content(path))

    def _order_columns(self, part: str) -> OrderColumns | None:
        header_row = _HEADER_ROW.search(self._parts[part])
        if header_row is None:
            return None
        cells = _cells(header_row.group(0))
        header: list[object] = [None] * max((cell.column for cell in cells), default=0)
        for cell in cells:
            header[cell.column - 1] = self._value(cell)
        found = header_columns(header)
        named = {
            column.target_header: next((found[name] for name in map(normalize_text, column.register_headers) if name in found), None)
            for column in COLUMNS
        }
        first, last = found.get(normalize_text(_FIRST)), found.get(normalize_text(_LAST))
        if first is None or last is None or None in named.values():
            return None
        return OrderColumns(
            first,
            cast(int, named[DATE.target_header]),
            cast(int, named[WORKER.target_header]),
            last,
            devices=cast(int, named[DEVICES.target_header]),
            activities=cast(int, named[ACTIVITIES.target_header]),
            contract_number=cast(int, named[_CONTRACT_NUMBER]),
        )

    def _value(self, cell: _Cell) -> object:
        """Wartość komórki: tekst (także prawda/fałsz jako 1 i 0), liczba, data w formacie daty; ``None`` dla pustej."""

        content = cell.content or b""
        kind = cell.kind or b"n"
        if kind == b"inlineStr":
            text: object = html.unescape(b"".join(_TEXT.findall(content)).decode("utf-8"))
        else:
            found = _VALUE.search(content)
            if found is None:
                return None
            raw = found.group(1).decode("utf-8")
            if kind == b"s":
                text = self._shared[int(raw)]
            elif kind == b"n":
                number = float(raw)
                if self._styles.number_formats[cell.style] in self._styles.date_formats:
                    return from_excel(number, self._epoch)
                return number
            else:
                text = html.unescape(raw)
        return None if is_blank(text) else text

    def _facts(self, row_tag: bytes, number: int, cells: Sequence[_Cell]) -> RowFacts:
        styles = self._styles
        attributes = _row_attributes(row_tag)
        row_style = attributes.get(b"s") if attributes.get(b"customFormat") == b"1" else None
        return RowFacts(
            number,
            {
                cell.column: CellFacts(self._value(cell), styles.faces[cell.style], styles.yellow[cell.style])
                for cell in cells
            },
            row_style is not None and styles.yellow[int(row_style)],
        )

    def ordered(self, sheet: OrderSheet, block: RawBlock, plan: DayPlan) -> bytes:
        """Bajty INTERMAX z uporządkowanym dniem; ``UnsupportedShape``, gdy arkusza nie da się tak zapisać."""

        xml = sheet.sheet_xml
        looks = _Looks(self._styles, sheet.columns)
        sheet_xml = _ordered_sheet(xml, sheet.columns, block, plan, looks)
        parts = {
            **self._tabs(sheet),
            sheet.part: sheet_xml,
            self._styles_part: looks.styles(self._parts[self._styles_part]),
        }
        out = BytesIO()
        with zipfile.ZipFile(out, "w") as target:
            for info in self._infos:
                copied = zipfile.ZipInfo(info.filename, info.date_time)
                copied.compress_type = info.compress_type
                copied.external_attr = info.external_attr
                target.writestr(copied, parts.get(info.filename, self._parts[info.filename]))
        return out.getvalue()

    def check(self, written: bytes, sheet: OrderSheet, block: RawBlock, plan: DayPlan) -> None:
        """``WriteMismatch``, gdy ``written`` różni się od INTERMAX czymś więcej niż dniem z ``plan``.

        Uszkodzenie ``written`` zgłasza jak odczyt skoroszytu (``workbook_files.unreadable_as_error``).
        """

        with zipfile.ZipFile(BytesIO(written)) as archive:
            names = [info.filename for info in archive.infolist()]
            parts = {name: archive.read(name) for name in names}
        if names != [info.filename for info in self._infos]:
            raise WriteMismatch("części")
        tabs = self._tabs(sheet)
        for name in names:
            if name not in (sheet.part, self._styles_part) and parts[name] != tabs.get(name, self._parts[name]):
                raise WriteMismatch(name)
        _check_styles(self._parts[self._styles_part], parts[self._styles_part])
        _check_sheet(sheet.sheet_xml, parts[sheet.part], sheet.columns, block, plan)

    def _tabs(self, sheet: OrderSheet) -> dict[str, bytes]:
        """Części, które zapis zmienia poza arkuszem zleceń: ``sheet`` jako aktywna karta i inne karty bez zaznaczenia.

        Znacznik zaznaczenia karty (``<sheetView tabSelected>``) może stać tylko w arkuszu.
        """

        return {
            **{part: _TAB_SELECTED.sub(rb"\1", self._parts[part]) for part in self._sheet_parts if part != sheet.part},
            self._workbook_part: _with_active_tab(self._parts[self._workbook_part], self._sheet_parts.index(sheet.part)),
        }


def _with_active_tab(xml: bytes, index: int) -> bytes:
    """``xl/workbook.xml`` z aktywną kartą ``index`` (od 0) w każdym widoku skoroszytu."""

    tab = b' activeTab="%d"' % index

    def active(found: re.Match[bytes]) -> bytes:
        attributes, replaced = re.subn(rb'\sactiveTab="[^"]*"', tab, found.group(1))
        return b"<workbookView" + attributes + (b"" if replaced else tab) + found.group(2)

    return _WORKBOOK_VIEW.sub(active, xml)


def _with_view(head: bytes, row: int, column: int) -> bytes:
    """Nagłówek arkusza z kartą zaznaczoną, górnym wierszem widoku ``row`` i zaznaczeniem komórki ``column`` w nim.

    Zamrożony nagłówek (``<pane>``) zostaje, a zmienia się jego ``topLeftCell``; bez niego ``topLeftCell`` widoku.
    Widok zaczyna się od kolumny A (u taty nie ma zamrożonych kolumn) i ma potem tylko ``<pane>`` i jedno
    zaznaczenie, w aktywnym panelu.
    """

    cell = f"{get_column_letter(column)}{row}".encode("ascii")

    def top(attributes: bytes) -> bytes:
        return re.sub(rb'\stopLeftCell="[^"]*"', b"", attributes) + b' topLeftCell="A%d"' % row

    def view(found: re.Match[bytes]) -> bytes:
        attributes = re.sub(rb'\stabSelected="[^"]*"', b"", found.group(1))
        pane = _PANE.search(found.group(0))
        if pane is None:
            attributes, panes, name = top(attributes), b"", b""
        else:
            active = re.search(rb'\sactivePane="(\w+)"', pane.group(1))
            panes = b"<pane" + top(pane.group(1)) + b"/>"
            name = b' pane="%s"' % (active.group(1) if active else b"topLeft")
        selection = b'<selection%s activeCell="%s" sqref="%s"/>' % (name, cell, cell)
        return b'<sheetView tabSelected="1"' + attributes + b">" + panes + selection + b"</sheetView>"

    return _SHEET_VIEW.sub(view, head)


def _shared_strings(xml: bytes) -> list[str]:
    """Teksty wspólne bez podpowiedzi wymowy (``rPh``): zwykły tekst albo tekst każdego fragmentu."""

    return [
        "".join(text.text or "" for text in [*item.iterfind(f"{MAIN}t"), *item.iterfind(f"{MAIN}r/{MAIN}t")])
        for item in ElementTree.fromstring(xml).iterfind(f"{MAIN}si")
    ]


class _Looks:
    """Style wyglądu reszty INTERMAX dla kolumn zlecenia, dopisane do ``xl/styles.xml`` tylko raz każdy."""

    def __init__(self, styles: _Styles, columns: OrderColumns) -> None:
        self._styles = styles
        self._columns = columns
        self._fonts: list[Font] = []
        self._fills: list[Fill] = []
        self._borders: list[Border] = []
        self._xfs: list[bytes] = []
        self._style_ids: dict[tuple[bool, int, str | None], int] = {}

    def style(self, column: int, number_format: int, fill: str | None) -> int:
        """Styl komórki kolumny zlecenia ``column``: Calibri 8 (WYKONAWCA pogrubiony Calibri 9), czarna, cienka ramka,
        do lewej i środka.

        ``fill``: wypełnienie RGB z Koloru operatora, jak to, które nadaje Excel; ``None``: bez wypełnienia.
        """

        if column == self._columns.date:
            number_format = _DATE_FORMAT
        worker = column == self._columns.worker
        key = (worker, number_format, fill)
        if key not in self._style_ids:
            size, bold = (_WORKER_SIZE, True) if worker else (_SIZE, None)
            font = _index(self._styles.fonts, self._fonts, Font(name=_FACE, sz=size, b=bold, color=_BLACK, family=2))
            fill_id = 0 if fill is None else _index(self._styles.fills, self._fills, _solid(fill))
            border = _index(self._styles.borders, self._borders, BORDER)
            xf = (
                b'<xf numFmtId="%d" fontId="%d" fillId="%d" borderId="%d" xfId="0" applyNumberFormat="1" applyFont="1" '
                b'applyFill="1" applyBorder="1" applyAlignment="1"><alignment horizontal="left" vertical="center"/></xf>'
                % (number_format, font, fill_id, border)
            )
            if xf in self._styles.xfs:
                self._style_ids[key] = self._styles.xfs.index(xf)
            else:
                self._xfs.append(xf)
                self._style_ids[key] = len(self._styles.xfs) + len(self._xfs) - 1
        return self._style_ids[key]

    def number_format(self, style: int) -> int:
        return self._styles.number_formats[style]

    def styles(self, xml: bytes) -> bytes:
        styles = self._styles
        xml = _appended(xml, b"fonts", [ElementTree.tostring(font.to_tree()) for font in self._fonts], len(styles.fonts))
        xml = _appended(xml, b"fills", [ElementTree.tostring(fill.to_tree()) for fill in self._fills], len(styles.fills))
        borders = [ElementTree.tostring(border.to_tree()) for border in self._borders]
        xml = _appended(xml, b"borders", borders, len(styles.borders))
        return _appended(xml, b"cellXfs", self._xfs, len(styles.xfs))


def _solid(rgb: str) -> PatternFill:
    """Pełne wypełnienie ``rgb`` tak, jak zapisuje je Excel taty (``bgColor`` 64)."""

    return PatternFill(patternType="solid", fgColor=Color(rgb=f"FF{rgb}"), bgColor=Color(indexed=64))


def _index(existing: Sequence[object], added: list[Any], item: object) -> int:
    if item in existing:
        return list(existing).index(item)
    if item not in added:
        added.append(item)
    return len(existing) + added.index(item)


def _style_list(xml: bytes, tag: bytes) -> re.Match[bytes]:
    """Lista ``tag`` w ``xl/styles.xml``: atrybuty (grupa 1) i wpisy (grupa 2); ``WriteMismatch``, gdy jej nie ma."""

    found = re.search(rb"<" + tag + rb"\b([^>]*)>(.*?)</" + tag + rb">", xml, re.S)
    if found is None:
        raise WriteMismatch(f"style: {tag.decode('ascii')}")
    return found


def _appended(xml: bytes, tag: bytes, entries: Sequence[bytes], existing: int) -> bytes:
    """``entries`` na końcu listy ``tag`` (``existing`` pozycji) w ``xl/styles.xml`` z licznikiem ``count``."""

    block = _style_list(xml, tag)
    count = existing + len(entries)
    attributes = re.sub(rb'\scount="\d+"', b"", block.group(1))
    opening = b"<" + tag + b' count="%d"' % count + attributes + b">"
    return xml[: block.start()] + opening + block.group(2) + b"".join(entries) + b"</" + tag + b">" + xml[block.end() :]


@dataclass(frozen=True)
class _RowLooks:
    """Wygląd ``<row>`` wierszy nowego dnia z ostatniego uporządkowanego dnia taty; jedyne miejsce tej decyzji.

    Wiersz zlecenia i pusty wiersz między segmentami mają atrybuty ostatniego uporządkowanego wiersza
    (``RawBlock.last_ordered``), a inny pusty wiersz atrybuty pustego wiersza taty pod nim (``RawBlock.blank_row``),
    bez niego te same co wiersz zlecenia. Bez uporządkowanego wiersza nowe wiersze nie mają atrybutów wyglądu.
    """

    order: bytes
    blank: bytes

    @classmethod
    def of(cls, xml: bytes, rows: Mapping[int, SheetRow], block: RawBlock) -> _RowLooks:
        def look(number: int) -> bytes:
            attributes = _row_attributes(_row_tag(xml, rows[number]))
            return b"".join(b' %s="%s"' % (name, attributes[name]) for name in _ROW_LOOK if name in attributes)

        order = look(block.last_ordered) if block.last_ordered > 1 else b""
        return cls(order, order if block.blank_row is None else look(block.blank_row))


def _ordered_sheet(sheet: _SheetXml, columns: OrderColumns, block: RawBlock, plan: DayPlan, looks: _Looks) -> bytes:
    xml, data = sheet.xml, sheet.data
    region = [row for row in data.rows if row.number > block.anchor]
    by_number = {row.number: row for row in data.rows}
    row_looks = _RowLooks.of(xml, by_number, block)
    raw = {order.row for order in block.orders}
    tail_at = data.end + len(b"</sheetData>")
    _require_writable(xml[tail_at:], block.anchor)
    for number in raw:
        if _FORMULA.search(xml, by_number[number].start, by_number[number].end):
            raise UnsupportedShape("formuła")
    placed = {placement.row: placement for placement in plan.placements}
    separators = set(range(plan.placements[0].row, plan.placements[-1].row)) - set(placed)
    last = max([*placed, *(row.number for row in region)])
    rows: list[bytes] = []
    for number in range(block.anchor + 1, last + 1):
        placement = placed.get(number)
        if placement is not None:
            source = by_number[placement.order.row]
            offset = SHIFT if placement.order.shifted else 0
            rows.append(
                _order_row(
                    xml[source.start : source.end],
                    number,
                    offset,
                    columns,
                    looks,
                    row_looks.order,
                    placement.color,
                    _entry_cells(placement.entry, columns),
                )
            )
        elif number in separators:
            rows.append(_order_row(b"", number, 0, columns, looks, row_looks.order, OperatorColor(), {}))
        elif number in by_number and number not in raw:
            rows.append(xml[by_number[number].start : by_number[number].end])
        else:
            rows.append(b'<row r="%d"%s/>' % (number, row_looks.blank))
    head = _with_view(_with_dimension(xml[: region[0].start], last, columns.last), plan.view_row, columns.first)
    return head + b"".join(rows) + xml[data.rows[-1].end : tail_at] + _moved_hyperlinks(xml[tail_at:], columns, block, plan)


def _order_row(
    row_xml: bytes,
    number: int,
    offset: int,
    columns: OrderColumns,
    looks: _Looks,
    row_look: bytes,
    color: OperatorColor,
    entry: Mapping[int, _Cell],
) -> bytes:
    """Wiersz zlecenia ``number`` z komórek Surowego zlecenia od A do ``UWAGI``, w kolumnach ``entry`` z komórkami
    Wpisu, z Kolorem operatora ``color`` i wyglądem ``row_look``."""

    fills = {columns.activities: color.activities, columns.contract_number: color.contract_number, columns.devices: color.devices}
    written: dict[int, bytes] = {}
    for source in _cells(row_xml):
        column = source.column + offset
        if column > columns.last:
            continue
        cell = entry.get(column, source)
        style = looks.style(column, looks.number_format(source.style), fills.get(column))
        written[column] = cell.xml(column, number, _STYLE.sub(b"", cell.attributes) + b' s="%d"' % style)
    for column in range(1, columns.last + 1):
        if column not in written:
            style = looks.style(column, 0, fills.get(column))
            written[column] = _Cell(column, b"", None).xml(column, number, b' s="%d"' % style)
    cells = b"".join(written[column] for column in sorted(written))
    return b'<row r="%d"%s>' % (number, row_look) + cells + b"</row>"


def _entry_cells(entry: IntermaxEntry | None, columns: OrderColumns) -> dict[int, _Cell]:
    """Komórki Wpisu urządzenia i czynności INTERMAX po kolumnie: tekst w ``URZĄDZENIA`` i ``CZYNNOŚCI``; bez Wpisu
    żadnych."""

    if entry is None:
        return {}
    return {
        column: _Cell(column, b' t="inlineStr"', b"<is><t>%s</t></is>" % html.escape(text, quote=False).encode("utf-8"))
        for column, text in ((columns.devices, entry.devices), (columns.activities, entry.activities))
    }


def _with_dimension(head: bytes, last_row: int, last_column: int) -> bytes:
    """Nagłówek arkusza z ``<dimension>`` sięgającym co najmniej ``last_row`` i ``last_column``."""

    found = _DIMENSION.search(head)
    if found is None:
        return head
    corners = found.group(1).decode("ascii").split(":")
    column, row = coordinate_from_string(corners[-1])
    end = f"{get_column_letter(max(column_index_from_string(column), last_column))}{max(row, last_row)}"
    return head[: found.start(1)] + f"{corners[0]}:{end}".encode("ascii") + head[found.end(1) :]


def _require_writable(tail: bytes, anchor: int) -> None:
    """Zakresy za ``sheetData`` (poza Hiperłączami) nie sięgają wierszy pod ``anchor``."""

    for found in _RANGES.finditer(_HYPERLINKS.sub(b"", tail)):
        for reference in found.group(1).decode("ascii").split():
            _, _, _, max_row = range_boundaries(reference)
            if max_row is None or max_row > anchor:
                raise UnsupportedShape("zakres")


def _moves(block: RawBlock, plan: DayPlan) -> dict[int, tuple[int, int]]:
    """Wiersz Surowego zlecenia → (przesunięcie kolumn, docelowy wiersz)."""

    return {placement.order.row: (SHIFT if placement.order.shifted else 0, placement.row) for placement in plan.placements}


def _moved_hyperlinks(tail: bytes, columns: OrderColumns, block: RawBlock, plan: DayPlan) -> bytes:
    moves = _moves(block, plan)

    def moved(found: re.Match[bytes]) -> bytes:
        reference = found.group(2).decode("ascii")
        first = reference.split(":")[0]
        column, number = coordinate_from_string(first)
        if number <= block.anchor:
            return found.group(0)
        if reference != first or number not in moves:
            raise UnsupportedShape("Hiperłącze")
        offset, row = moves[number]
        moved_column = column_index_from_string(column) + offset
        if moved_column > columns.last:
            raise UnsupportedShape("Hiperłącze")
        return found.group(1) + f"{get_column_letter(moved_column)}{row}".encode("ascii") + found.group(3)

    return _HYPERLINK_REF.sub(moved, tail)


def _check_styles(original: bytes, written: bytes) -> None:
    """``WriteMismatch``, gdy arkusz stylów ``written`` to coś więcej niż ``original`` z wpisami na końcu list."""

    for tag in _STYLE_LISTS:
        before, after = _style_list(original, tag), _style_list(written, tag)
        if not after.group(2).startswith(before.group(2)):
            raise WriteMismatch(f"style: {tag.decode('ascii')}")
        written = written[: after.start()] + before.group(0) + written[after.end() :]
    if written != original:
        raise WriteMismatch("style")


def _check_sheet(original: _SheetXml, written: bytes, columns: OrderColumns, block: RawBlock, plan: DayPlan) -> None:
    try:
        data = SheetData.of(written)
    except IrregularSheetData as error:
        raise WriteMismatch("arkusz") from error
    xml, before = original.xml, original.data
    start = next((row.start for row in before.rows if row.number > block.anchor), before.end)
    new_start = next((row.start for row in data.rows if row.number > block.anchor), data.end)
    if _without_view(xml[:start]) != _without_view(written[:new_start]):
        raise WriteMismatch("wiersze nad dniem")
    tail, new_tail = xml[before.end :], written[data.end :]
    if _moved_hyperlinks(tail, columns, block, plan) != new_tail:
        raise WriteMismatch("ogon arkusza")
    rows = {row.number: row for row in data.rows[data.row_index(new_start) :]}
    sources = {row.number: row for row in before.rows}
    raw = {order.row for order in block.orders}
    kept = {number: xml[row.start : row.end] for number, row in sources.items() if number not in raw}
    placed = {placement.row: placement for placement in plan.placements}
    for number, row in rows.items():
        cells = _cells(written[row.start : row.end])
        valued = {cell.column: cell for cell in cells if _has_value(cell)}
        placement = placed.get(number)
        if placement is None:
            if valued and written[row.start : row.end] != kept.get(number):
                raise WriteMismatch(f"wiersz {number}")
            continue
        source = sources[placement.order.row]
        offset = SHIFT if placement.order.shifted else 0
        moved = ((cell.column + offset, cell) for cell in _cells(xml[source.start : source.end]))
        expected = {column: cell for column, cell in moved if column <= columns.last}
        expected.update(_entry_cells(placement.entry, columns))
        if {column for column, cell in expected.items() if _has_value(cell)} != set(valued):
            raise WriteMismatch(f"wiersz {number}")
        found = {cell.column: cell for cell in cells}
        if any(column > columns.last for column in found):
            raise WriteMismatch(f"wiersz {number}")
        for column, cell in expected.items():
            new = found.get(column)
            if new is None or new.content != cell.content or new.kind != cell.kind:
                raise WriteMismatch(f"wiersz {number}")
    if set(placed) - set(rows):
        raise WriteMismatch("brak zlecenia")


def _without_view(head: bytes) -> bytes:
    """Nagłówek arkusza zleceń bez tego, co zmienia zapis: ``<dimension>`` i widoku arkusza."""

    return _SHEET_VIEWS.sub(b"", _DIMENSION.sub(b"", head))


def _has_value(cell: _Cell) -> bool:
    content = cell.content or b""
    return _VALUE.search(content) is not None or b"<is>" in content


def save(
    register: Register,
    sheet: OrderSheet,
    block: RawBlock,
    plan: DayPlan,
    *,
    before_replace: Callable[[], None],
    waiting: Waiting,
) -> None:
    """Zapisuje dzień w INTERMAX: plik tymczasowy obok, sprawdzenie zapisu, ``before_replace`` i atomowa podmiana.

    ``WriteMismatch``, ``UnsupportedShape`` albo wyjątek z ``before_replace`` zostawiają INTERMAX nietknięty i bez
    pliku tymczasowego. Nieudana podmiana to ``WorkbookSaveError``, a odmowa podmiany bez pliku blokady Excela po
    ponownych próbach przez ``_RETRY_S`` sekund zegara ``waiting`` to ``HeldByOneDrive``. Ponowna próba podmienia
    tylko INTERMAX, który ma wciąż bajty ``register``: inne bajty albo plik, którego nie da się przeczytać, to
    ``HeldByOneDrive``, a plik blokady Excela to ``WorkbookSaveError``.
    """

    def check(written: bytes) -> None:
        register.check(written, sheet, block, plan)

    _replace(register.path, register.content, register.ordered(sheet, block, plan), check, before_replace, waiting)


def restore(
    path: Path, current: bytes, content: bytes, *, before_replace: Callable[[], None], waiting: Waiting
) -> None:
    """Podmienia cały INTERMAX ``path`` o bajtach ``current`` bajtami Kopii INTERMAX ``content``, jak ``save``.

    Plik tymczasowy obok musi mieć po odczycie bajty ``content`` (inaczej ``WriteMismatch``); dalej ``before_replace``
    i ta sama atomowa podmiana z ponowną próbą, która podmienia tylko INTERMAX o bajtach ``current``.
    """

    def check(written: bytes) -> None:
        if written != content:
            raise WriteMismatch("kopia")

    _replace(path, current, content, check, before_replace, waiting)


def _replace(
    path: Path,
    current: bytes,
    content: bytes,
    check: Callable[[bytes], None],
    before_replace: Callable[[], None],
    waiting: Waiting,
) -> None:
    """``content`` w pliku tymczasowym obok ``path``, ``check`` jego odczytu, ``before_replace`` i podmiana ``path``.

    Podmianę, której system odmawia bez pliku blokady Excela, ponawia przez ``_RETRY_S`` sekund zegara ``waiting``,
    a przed każdą ponowną próbą czyta ``path`` jeszcze raz: bajty inne niż ``current`` to ``HeldByOneDrive``.
    """

    def write(temporary: Path) -> None:
        temporary.write_bytes(content)
        try:
            written = workbook_files.workbook_content(temporary)
            with workbook_files.unreadable_as_error(temporary):
                check(written)
        except WorkbookError as error:
            raise WriteMismatch("odczyt") from error
        before_replace()

    def replace(temporary: Path, target: Path) -> None:
        deadline = waiting.clock() + _RETRY_S
        while True:
            try:
                os.replace(temporary, target)
                return
            except PermissionError as error:
                if workbook_files.is_open_in_excel(target):
                    raise
                if waiting.clock() >= deadline:
                    raise HeldByOneDrive(target.name) from error
            waiting.pause(_PAUSE_S)
            try:
                unchanged = workbook_files.workbook_content(target) == current
            except WorkbookError as error:
                if error.problem is workbook_files.Problem.OPEN_IN_EXCEL:
                    raise WorkbookSaveError(target.name) from error
                raise HeldByOneDrive(target.name) from error
            if not unchanged:
                raise HeldByOneDrive(target.name)

    write_atomically(path, write, replace=replace)
