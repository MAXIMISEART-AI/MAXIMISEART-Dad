"""Kopia skoroszytu Wzorca z podmienionymi Wierszami danych, bez wczytywania przez openpyxl (ADR 0008).

Moduł jest szczegółem ``template_file.TemplateFile`` i nikt poza nim go nie używa.
``TemplateCopy.of`` czyta bajty Wzorca raz: rozpakowuje części zip,
sprawdza, że każda część XML jest poprawnym XML-em, indeksuje wiersze XML
aktywnego arkusza i liczy kontrole Szablonu pracownika (nagłówek, dane w
Wierszach danych); na tym opiera się Sprawdzanie Wzorca. Arkusz jest kompresowany
raz, w kawałkach od granic wierszy, przy pierwszej kopii. ``TemplateCopy.with_rows``
składa z tego jedną kopię:

- każda część zip poza arkuszem, ``xl/workbook.xml`` i ``xl/styles.xml`` jest
  przepisana bajt w bajt razem ze skompresowaną treścią, w tym
  ``xl/externalLinks/*``, ``xl/calcChain.xml`` i ``xl/sharedStrings.xml``;
- ``xl/styles.xml`` zmienia się tylko wtedy, gdy wpisana komórka potrzebuje stylu,
  którego Wzorzec nie ma: Wygląd Wiersza danych (``row_look``) w kolumnach wyglądu, data
  w komórce bez formatu daty albo Wyróżnienia inne niż w komórce Wzorca; nowe czcionki,
  wypełnienia, obramowania, formaty liczb i style komórek są dopisywane na końcu swoich
  list, a istniejące zostają bajt w bajt;
- w arkuszu zmieniają się tylko wiersze od pierwszego Wiersza danych do ostatniego
  wpisanego: komórki kolumn Wiersza danych dostają wartości, a kolumn wyglądu Wygląd
  Wiersza danych, za nimi Wyróżnienia na stylu Wzorca, tak jak przy wpisaniu przez openpyxl
  (``template_file._write_rows``); wpisany wiersz traci zapisaną wysokość (``ht``,
  ``customHeight``), a pozostałe komórki, w tym formuły, wspólne formuły i style, zostają
  bajt w bajt; kompresowany jest tylko początek arkusza do granicy kawałka za ostatnim
  wpisanym wierszem, reszta to gotowe kawałki z pierwszej kopii;
- ``calcPr`` w ``xl/workbook.xml`` dostaje ``calcMode="auto"``, ``fullCalcOnLoad="1"``
  i ``forceFullCalc="1"``, jak przy zapisie przez openpyxl;
- Plik okresu (``SheetLayout.period_file``), czyli Plik zbiorczy, dostaje też element ``autoFilter`` na
  nagłówek i wpisane wiersze oraz ten sam zakres w nazwie ``_xlnm._FilterDatabase`` w
  ``xl/workbook.xml``;
- Hiperłącza wpisanych komórek trafiają do jednego bloku ``<hyperlinks>`` w miejscu,
  które wyznacza mu schemat arkusza, a ich relacje są dopisywane na końcu relacji
  arkusza pod numerami ``rId`` za najwyższym we Wzorcu; Wzorzec bez relacji arkusza
  dostaje nową część na końcu zip. Bez Hiperłączy arkusz i relacje są takie jak bez
  tej zmiany.

Wzorzec, którego kształtu ten moduł nie obsługuje (np. formuła w kolumnach Wiersza
danych, scalenie w Wierszach danych, własne hiperłącza arkusza, XML z prefiksami albo
bez atrybutów ``r``, uszkodzona część XML, a w Pliku zbiorczym brak autofiltra), zgłasza
``UnsupportedShape``;
tak samo wartość, której typu moduł nie zapisuje. Wtedy ``template_file`` sprawdza Wzorzec
i tworzy plik dawną drogą przez openpyxl, z jej klasyfikacją uszkodzeń.
Wartość, której nie da się zapisać w ogóle (znak niedozwolony w XML, data ze
strefą czasową), to ``InvalidValue``.
"""

from __future__ import annotations

import html
import math
import re
import struct
import zipfile
import zlib
from bisect import bisect_left, bisect_right
from collections.abc import Iterator, Mapping, Sequence
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from io import BytesIO
from typing import IO, Any, Final, Protocol, TypeVar, cast
from xml.etree import ElementTree
from xml.parsers import expat

from openpyxl.cell.cell import ERROR_CODES, ILLEGAL_CHARACTERS_RE
from openpyxl.styles.numbers import (
    BUILTIN_FORMATS,
    BUILTIN_FORMATS_REVERSE,
    FORMAT_DATE_DATETIME,
    FORMAT_DATE_TIME6,
    FORMAT_DATE_TIMEDELTA,
    FORMAT_DATE_YYYYMMDD2,
    is_date_format,
)
from openpyxl.styles.fills import Fill, PatternFill
from openpyxl.styles.fonts import Font
from openpyxl.utils.cell import absolute_coordinate, column_index_from_string, get_column_letter, range_boundaries
from openpyxl.utils.datetime import MAC_EPOCH, WINDOWS_EPOCH, to_excel

from openpyxl.styles.alignment import Alignment
from openpyxl.styles.borders import Border

from .highlights import PLAIN, Highlights
from .hyperlinks import Hyperlink
from .quantity_rules import QuantityColumn
from .row_look import GENERAL, CellLook, RowLook, wraps
from .workbook_files import DataCell
from .workbook_parts import (
    MAIN,
    OFFICE_REL,
    PACKAGE_REL,
    IrregularSheetData,
    Relationship,
    SheetData,
    SheetRow,
    office_document,
    relationships,
    relationships_part,
)


_T = TypeVar("_T")


class UnsupportedShape(Exception):
    """Wzorca albo wartości nie da się zapisać podmianą wierszy; zapis idzie dawną drogą."""


class InvalidValue(Exception):
    """Wartości Wiersza danych nie da się zapisać w skoroszycie."""


_CONTENT_DEFAULT: Final = "{http://schemas.openxmlformats.org/package/2006/content-types}Default"
# Błędy odczytu bajtów Wzorca, które znaczą kształt spoza kopii albo uszkodzenie: Wzorzec idzie wtedy dawną
# drogą przez openpyxl, z jej klasyfikacją. Kopia czyta tylko bajty w pamięci, więc błąd wejścia-wyjścia
# (``OSError``) ani inny nieoczekiwany wyjątek nie jest tu kształtem i nie jest maskowany.
_READ_ERRORS: Final = (
    zipfile.BadZipFile,
    zipfile.LargeZipFile,
    KeyError,
    IndexError,
    ValueError,
    ElementTree.ParseError,
    expat.ExpatError,
    zlib.error,
    struct.error,
    EOFError,
)

# Wiersze arkusza w jednym kawałku kompresji. Mniejszy kawałek: mniej kompresji przy każdej kopii,
# nieco gorszy stopień kompresji całości.
_ROWS_PER_CHUNK: Final = 64
_COMPRESSION_LEVEL: Final = 6
_MAX_TEXT: Final = 32767  # openpyxl przycina tekst komórki do tej długości

_CELL = re.compile(rb'<c r="([A-Z]{1,3})(\d+)"([^>]*?)(?:/>|>(.*?)</c>)', re.S)
_CELL_WITH_CONTENT = re.compile(rb'<c r="([A-Z]{1,3})\d+"[^>]*(?<!/)>(.*?)</c>', re.S)  # komórka z treścią
_STYLE = re.compile(rb'\ss="(\d+)"')
_TYPE = re.compile(rb'\st="(\w+)"')
_FORMULA = re.compile(rb"<f[\s>/]")
_NON_EMPTY_VALUE = re.compile(rb"<v(?:\s[^>]*)?>[^<]")
_VALUE = re.compile(rb"<v(?:\s[^>]*)?>([^<]*)</v>")
_ATTRIBUTE = re.compile(rb'([\w:.-]+)\s*=\s*(?:"([^"]*)"|\'([^\']*)\')')
_CALC_PR = re.compile(rb"<((?:[\w.-]+:)?)calcPr\b([^>]*?)(\s*/)?>")
_DIMENSION = re.compile(rb'<dimension ref="([A-Z]+\d+(?::[A-Z]+\d+)?)"')
_MERGE = re.compile(rb'<mergeCell ref="([^"]+)"')
_CELL_XFS = re.compile(rb"<cellXfs\b[^>]*>(.*?)</cellXfs>", re.S)
_FONTS = re.compile(rb"<fonts\b[^>]*>(.*?)</fonts>", re.S)
_FILLS = re.compile(rb"<fills\b[^>]*>(.*?)</fills>", re.S)
_BORDERS = re.compile(rb"<borders\b[^>]*>(.*?)</borders>", re.S)
_ALIGNMENT = re.compile(rb"<alignment\b[^>]*?(?:/>|>.*?</alignment>)", re.S)
_HEIGHT = re.compile(rb'\s(?:ht|customHeight)="[^"]*"')
_XF = re.compile(rb"<xf\b[^>]*?(?:/>|>.*?</xf>)", re.S)
_NUM_FMT = re.compile(rb"<numFmt\b([^>]*?)/?>")
_NUM_FMT_ID = re.compile(rb'numFmtId="(\d+)"')
_CALC_FLAGS: Final = ((b"calcMode", b"auto"), (b"fullCalcOnLoad", b"1"), (b"forceFullCalc", b"1"))
_WORKSHEET_START = re.compile(rb"<((?:[\w.-]+:)?)worksheet\b[^>]*>")
_RELATIONSHIPS_PREFIX = re.compile(rb"""xmlns:([\w.-]+)\s*=\s*["']""" + OFFICE_REL.encode() + rb"""["']""")
_RELATIONSHIP_NUMBER = re.compile(r"rId(\d+)")
# Elementy arkusza, które schemat (CT_Worksheet) stawia między ``sheetData`` a ``hyperlinks``; każdy inny element
# za ``sheetData`` (``printOptions``, ``pageMargins``, …, ``extLst``, także z innej przestrzeni nazw) idzie za nimi.
_BEFORE_HYPERLINKS: Final = frozenset(
    {
        "sheetCalcPr",
        "sheetProtection",
        "protectedRanges",
        "scenarios",
        "autoFilter",
        "sortState",
        "dataConsolidate",
        "customSheetViews",
        "mergeCells",
        "phoneticPr",
        "conditionalFormatting",
        "dataValidations",
    }
)


@dataclass(frozen=True)
class _Part:
    """Część zip gotowa do przepisania: nagłówek i skompresowana treść."""

    name: bytes
    flags: int
    method: int
    dos_time: int
    dos_date: int
    crc: int
    size: int
    data: Sequence[bytes]

    @property
    def compressed_size(self) -> int:
        return sum(len(chunk) for chunk in self.data)


class Package:
    """Gotowe Rozliczenie pracownika: części zip do zapisania jednym przebiegiem."""

    def __init__(self, parts: Sequence[_Part]) -> None:
        self._parts = parts

    def write(self, out: IO[bytes]) -> None:
        central: list[bytes] = []
        offset = 0
        for part in self._parts:
            version = 20 if part.method == zipfile.ZIP_DEFLATED else 10
            fields = (part.flags, part.method, part.dos_time, part.dos_date, part.crc, part.compressed_size, part.size)
            local = struct.pack("<IHHHHHIIIHH", 0x04034B50, version, *fields, len(part.name), 0) + part.name
            out.write(local)
            for chunk in part.data:
                out.write(chunk)
            central.append(
                struct.pack("<IHHHHHHIIIHHHHHII", 0x02014B50, 20, version, *fields, len(part.name), 0, 0, 0, 0, 0, offset)
                + part.name
            )
            offset += len(local) + part.compressed_size
        directory = b"".join(central)
        out.write(directory)
        out.write(struct.pack("<IHHHHIIH", 0x06054B50, 0, 0, len(central), len(central), len(directory), offset, 0))


class Row(Protocol):
    """Wiersz danych do wpisania: komórki kolumn Wiersza danych od A i Ilości usług."""

    @property
    def cells(self) -> Mapping[int, DataCell]:
        """Kolumna od 1 → komórka z wartością, Wyróżnieniami, formatem liczby i Hiperłączem; każda kolumna Wiersza danych."""
        ...

    @property
    def quantities(self) -> Mapping[QuantityColumn, int]:
        """Litera kolumny → liczba wpisana w komórkę Wzorca z jej stylem; tylko w ``SheetLayout.quantity_columns``."""
        ...


@dataclass(frozen=True)
class SheetLayout:
    """Układ arkusza Wzorca, jeden dla kopii i drogi przez openpyxl (``template_file``): gdzie są nagłówek, Wiersze
    danych, dane wejściowe i kolumny Wiersza danych (od 1)."""

    header_row: int
    first_data_row: int
    input_columns: int
    data_columns: int
    look_columns: int
    """Kolumny wyglądu (``row_look``); te za kolumnami Wiersza danych dostają sam wygląd, a treść zostaje z Wzorca."""
    period_file: bool
    """Plik okresu z Przygotowania: autofiltr obejmuje nagłówek i wpisane wiersze, a wygląd ma każda wpisana kolumna."""
    quantity_columns: range = range(0)
    """Kolumny Ilości usług; formuła w nich w Wierszach danych to ``UnsupportedShape``, jak w kolumnach Wiersza danych."""


class TemplateCopy:
    """Wzorzec wczytany raz; z niego powstaje dowolnie wiele kopii z Wierszami danych."""

    def __init__(
        self,
        parts: list[_Part | None],
        sheet: _Sheet,
        styles: _Styles,
        sheet_rels: _SheetRels,
        filter_name: _FilterName | None,
        title: str,
        header: list[object],
        has_input_data: bool,
        look: RowLook,
    ) -> None:
        self._parts = parts  # ``None``: miejsce arkusza, stylów albo skoroszytu, składanych dla każdej kopii
        self._sheet = sheet
        self._look = look
        self._filter_name = filter_name
        self._styles = styles
        self._sheet_rels = sheet_rels
        self.title = title
        """Nazwa aktywnego arkusza."""
        self.header = header
        """Wartości wiersza nagłówka w kolumnach danych wejściowych (od A), jak do sprawdzenia WYKONAWCA."""
        self.has_input_data = has_input_data
        """Wiersze danych Wzorca mają już wartość w kolumnach danych wejściowych; formuły nie są danymi."""

    @classmethod
    def of(cls, content: bytes, layout: SheetLayout) -> TemplateCopy:
        try:
            return _read_template(content, layout)
        except _READ_ERRORS as exc:
            raise UnsupportedShape(type(exc).__name__) from exc

    def with_rows(self, rows: Sequence[Row]) -> Package:
        links = self._sheet.links(rows)
        rels_part, ids = self._sheet_rels.with_links([link for _, link in links])
        styles = _StyleAdditions(self._styles)
        filter_ref = self._sheet.filter_ref(len(rows))
        sheet_part = self._sheet.with_rows(rows, styles, self._look, self._sheet.hyperlinks(links, ids), filter_ref)
        assembled = {self._sheet.index: sheet_part, self._styles.index: self._styles.part(styles)}
        if self._filter_name is not None and filter_ref is not None:
            assembled[self._filter_name.index] = self._filter_name.part(filter_ref)
        parts = [part if part is not None else assembled[index] for index, part in enumerate(self._parts)]
        if rels_part is not None:
            if self._sheet_rels.index is None:
                parts.append(rels_part)
            else:
                parts[self._sheet_rels.index] = rels_part
        return Package(parts)


# --- wczytanie ---------------------------------------------------------------------------------


def _read_template(content: bytes, layout: SheetLayout) -> TemplateCopy:
    archive = zipfile.ZipFile(BytesIO(content))
    infos = archive.infolist()
    for info in infos:
        if info.flag_bits & 0x1:
            raise UnsupportedShape("encrypted")
        if info.compress_type not in (zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED):
            raise UnsupportedShape("compression")
    plain = {info.filename: archive.read(info) for info in infos}  # sprawdza CRC każdej części
    _require_well_formed_xml(plain)
    workbook_name = office_document(plain)
    if workbook_name is None:
        raise UnsupportedShape("officeDocument")
    workbook_rels = relationships(plain, workbook_name)
    workbook_xml = plain[workbook_name]
    sheet_name, sheet_title, sheet_position, date1904 = _active_sheet(workbook_xml, workbook_rels)
    styles_name = _target_of_type(workbook_rels, "/styles")
    strings_name = _target_of_type(workbook_rels, "/sharedStrings")
    if styles_name is None or styles_name not in plain:
        raise UnsupportedShape("styles")
    shared_strings = _shared_strings(plain[strings_name]) if strings_name in plain else []

    rels_name = relationships_part(sheet_name)
    sheet_index = styles_index = -1
    rels_index: int | None = None
    filter_name: _FilterName | None = None
    parts: list[_Part | None] = []
    for index, info in enumerate(infos):
        if info.filename == rels_name:
            rels_index = index
        if info.filename == sheet_name:
            sheet_index = index
            parts.append(None)
        elif info.filename == styles_name:
            styles_index = index
            parts.append(None)
        elif info.filename == workbook_name and layout.period_file:
            filter_name = _FilterName(info, index, _with_full_recalculation(workbook_xml), sheet_position)
            parts.append(None)
        elif info.filename == workbook_name:
            parts.append(_compressed_part(info, _with_full_recalculation(workbook_xml)))
        else:
            parts.append(_raw_part(content, info))
    if sheet_index < 0:
        raise UnsupportedShape("sheet")
    sheet = _Sheet(infos[sheet_index], sheet_index, plain[sheet_name], layout, epoch=MAC_EPOCH if date1904 else WINDOWS_EPOCH)
    if layout.period_file:
        sheet.read_filter_columns(shared_strings)
    styles = _Styles(infos[styles_index], styles_index, plain[styles_name], _raw_part(content, infos[styles_index]))
    if rels_index is None:
        sheet_rels = _SheetRels(
            zipfile.ZipInfo(rels_name, infos[sheet_index].date_time), None, None, _rels_content_type(plain)
        )
    else:
        sheet_rels = _SheetRels(infos[rels_index], rels_index, plain[rels_name], True)
    look = RowLook([styles.font(style) for style in sheet.header_styles()])
    return TemplateCopy(
        parts, sheet, styles, sheet_rels, filter_name, sheet_title, sheet.header(shared_strings), sheet.has_input_data, look
    )


def _rels_content_type(plain: dict[str, bytes]) -> bool:
    """``[Content_Types].xml`` ma typ dla rozszerzenia ``rels``, więc nowa część relacji go nie zmienia."""

    types = ElementTree.fromstring(plain["[Content_Types].xml"])
    return any(default.get("Extension", "").lower() == "rels" for default in types.iter(_CONTENT_DEFAULT))


def _require_well_formed_xml(plain: dict[str, bytes]) -> None:
    """Każda część XML jest poprawnym XML-em; skoroszyt, którego nie przeczyta openpyxl, idzie dawną drogą.

    Kopia czyta arkusz tylko w miejscach, które zmienia, więc np. arkusz urwany za
    Wierszami danych przeszedłby niezauważony; dawna droga klasyfikuje go jak przy
    każdym otwarciu skoroszytu.
    """

    for name, data in plain.items():
        if name.endswith((".xml", ".rels")):
            expat.ParserCreate().Parse(data, True)


def _target_of_type(workbook_rels: Mapping[str, Relationship], suffix: str) -> str | None:
    return next((rel.target for rel in workbook_rels.values() if rel.kind.endswith(suffix)), None)


def _active_sheet(workbook_xml: bytes, workbook_rels: Mapping[str, Relationship]) -> tuple[str, str, int, bool]:
    """Część zip aktywnego arkusza (tak jak ``workbook.active`` w openpyxl), jego nazwa, pozycja i czy daty liczą się od 1904."""

    root = ElementTree.fromstring(workbook_xml)
    view = root.find(f"{MAIN}bookViews/{MAIN}workbookView")
    active = int(view.get("activeTab", "0")) if view is not None else 0
    sheets = root.findall(f"{MAIN}sheets/{MAIN}sheet")
    sheet = sheets[active]
    if sheet.get("state", "visible") != "visible":
        raise UnsupportedShape("hidden active sheet")
    relationship = workbook_rels[sheet.get(f"{{{OFFICE_REL}}}id", "")]
    if not relationship.kind.endswith("/worksheet"):
        raise UnsupportedShape("chartsheet")
    properties = root.find(f"{MAIN}workbookPr")
    date1904 = properties is not None and properties.get("date1904", "0").lower() in ("1", "true")
    return relationship.target, sheet.get("name", ""), active, date1904


def _shared_strings(xml: bytes) -> list[str]:
    return [_plain_text(item) for item in ElementTree.fromstring(xml).iter(f"{MAIN}si")]


def _plain_text(element: ElementTree.Element) -> str:
    """Tekst ``<si>`` albo ``<is>``: zwykły ``<t>`` albo przebiegi ``<r>``, bez transkrypcji fonetycznej."""

    texts: list[str] = []
    for child in element:
        local = child.tag.rsplit("}", 1)[-1]
        if local == "t":
            texts.append(child.text or "")
        elif local == "r":
            texts.extend(run.text or "" for run in child if run.tag.rsplit("}", 1)[-1] == "t")
    return "".join(texts)


def _with_full_recalculation(workbook_xml: bytes) -> bytes:
    """``calcPr`` z flagami przeliczania takimi jak przy zapisie przez openpyxl; pozostałe atrybuty bez zmian."""

    found = _CALC_PR.search(workbook_xml)
    if found is None:
        raise UnsupportedShape("calcPr")
    attributes = [(match.group(1), match.group(0)) for match in _ATTRIBUTE.finditer(found.group(2))]
    names = [name for name, _ in attributes]
    rendered = [
        next((name + b'="' + value + b'"' for flag, value in _CALC_FLAGS if flag == name), text)
        for name, text in attributes
    ]
    rendered += [flag + b'="' + value + b'"' for flag, value in _CALC_FLAGS if flag not in names]
    tag = b"<" + found.group(1) + b"calcPr " + b" ".join(rendered) + (b"/>" if found.group(3) else b">")
    return workbook_xml[: found.start()] + tag + workbook_xml[found.end() :]


_FILTER_DATABASE = re.compile(
    rb'<(?:[\w.-]+:)?definedName\b[^>]*\bname="_xlnm\._FilterDatabase"[^>]*>([^<]*)</(?:[\w.-]+:)?definedName>'
)
_LOCAL_SHEET = re.compile(rb'\blocalSheetId="(\d+)"')
_ABSOLUTE_RANGE = re.compile(rb"\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?")


class _FilterName:
    """``xl/workbook.xml`` Pliku zbiorczego: nazwa ``_xlnm._FilterDatabase`` aktywnego arkusza idzie za autofiltrem."""

    def __init__(self, info: zipfile.ZipInfo, index: int, xml: bytes, sheet_position: int) -> None:
        self._info = info
        self.index = index
        self._xml = xml
        for found in _FILTER_DATABASE.finditer(xml):
            local = _LOCAL_SHEET.search(found.group(0))
            if local is not None and int(local.group(1)) == sheet_position and b"!" in found.group(1):
                if not _ABSOLUTE_RANGE.fullmatch(found.group(1), found.group(1).rindex(b"!") + 1):
                    break  # np. ``#REF!``: nazwę naprawi dawna droga przez openpyxl
                self._range = (found.start(1) + found.group(1).rindex(b"!") + 1, found.end(1))
                return
        raise UnsupportedShape("filter name")

    def part(self, reference: str) -> _Part:
        start, end = self._range
        absolute = absolute_coordinate(reference).encode("ascii")
        return _compressed_part(self._info, self._xml[:start] + absolute + self._xml[end:])


def _dos_stamp(info: zipfile.ZipInfo) -> tuple[int, int]:
    year, month, day, hour, minute, second = info.date_time
    return hour << 11 | minute << 5 | second // 2, (year - 1980) << 9 | month << 5 | day


def _encoded_name(name: str) -> tuple[bytes, int]:
    try:
        return name.encode("ascii"), 0
    except UnicodeEncodeError:
        return name.encode("utf-8"), 0x800


def _raw_part(content: bytes, info: zipfile.ZipInfo) -> _Part:
    """Część przepisana razem ze skompresowaną treścią z zip Wzorca."""

    header = struct.unpack("<IHHHHHIIIHH", content[info.header_offset : info.header_offset + 30])
    if header[0] != 0x04034B50:
        raise UnsupportedShape("local header")
    start = info.header_offset + 30 + header[9] + header[10]
    name, flags = _encoded_name(info.filename)
    dos_time, dos_date = _dos_stamp(info)
    data = content[start : start + info.compress_size]
    return _Part(name, flags, info.compress_type, dos_time, dos_date, info.CRC, info.file_size, [data])


def _compressed_part(info: zipfile.ZipInfo, data: bytes) -> _Part:
    name, flags = _encoded_name(info.filename)
    dos_time, dos_date = _dos_stamp(info)
    compressor = zlib.compressobj(_COMPRESSION_LEVEL, zlib.DEFLATED, -15)
    compressed = compressor.compress(data) + compressor.flush()
    return _Part(name, flags, zipfile.ZIP_DEFLATED, dos_time, dos_date, zlib.crc32(data), len(data), [compressed])


# --- arkusz ------------------------------------------------------------------------------------


_COLUMN_INDEX: dict[bytes, int] = {}


def _column(letters: bytes) -> int:
    index = _COLUMN_INDEX.get(letters)
    if index is None:
        index = _COLUMN_INDEX[letters] = column_index_from_string(letters.decode("ascii"))
    return index


class _Sheet:
    """XML aktywnego arkusza z indeksem wierszy i kawałkami skompresowanymi od granic wierszy."""

    def __init__(self, info: zipfile.ZipInfo, index: int, xml: bytes, layout: SheetLayout, *, epoch: datetime) -> None:
        self.index = index
        self._info = info
        self._xml = xml
        self._layout = layout
        self._epoch = epoch
        self._letters = [
            get_column_letter(column).encode("ascii")
            for column in range(1, max(layout.data_columns, layout.look_columns, layout.quantity_columns.stop - 1) + 1)
        ]
        try:
            sheet_data = SheetData.of(xml)
        except IrregularSheetData as exc:
            raise UnsupportedShape(str(exc)) from exc
        self._rows = sheet_data.rows
        data_end = sheet_data.end
        self._numbers = [row.number for row in self._rows]
        first = bisect_left(self._numbers, layout.first_data_row)
        self._data_start = self._rows[first].start if first < len(self._rows) else data_end
        self._first_data_index = first
        self.has_input_data = _writable_data_rows_have_input(xml, self._data_start, data_end, layout)
        """Jak sprawdzenie Wzorca rozliczenia pracownika przez openpyxl: niepusta wartość bez formuły w kolumnach danych wejściowych."""
        self._links_at, self._links_open, self._links_prefix = _hyperlinks_place(xml, data_end)
        self._filter: tuple[int, int, bytes] | None = None
        self._filter_columns = (1, 1)
        if layout.period_file:
            self._filter = _autofilter(xml, data_end, self._links_at)
        # Ostatni kawałek zaczyna się w miejscu bloku Hiperłączy, więc kopia z nimi kompresuje tylko ogon arkusza;
        # przedostatni w Pliku zbiorczym od autofiltra, który zmienia się w każdej kopii.
        self._boundaries = [
            self._data_start,
            *(row.start for row in self._rows[first + _ROWS_PER_CHUNK :: _ROWS_PER_CHUNK]),
            *([self._filter[0]] if self._filter is not None else []),
            self._links_at,
        ]
        self._chunks: list[bytes] | None = None  # dopiero przy pierwszej kopii: Sprawdzanie ich nie potrzebuje
        dimension = _DIMENSION.search(xml, 0, sheet_data.start)
        self._dimension = None if dimension is None else (dimension.start(1), dimension.end(1), range_boundaries(dimension.group(1).decode("ascii")))

    def header(self, shared_strings: list[str]) -> list[object]:
        values: list[object] = [None] * self._layout.input_columns
        position = bisect_left(self._numbers, self._layout.header_row)
        if position == len(self._rows) or self._rows[position].number != self._layout.header_row:
            return values
        row = self._rows[position]
        for cell in _CELL.finditer(self._xml, row.start, row.end):
            column = _column(cell.group(1))
            if column <= self._layout.input_columns:
                values[column - 1] = _header_value(cell.group(3), cell.group(4), shared_strings)
        return values

    def header_styles(self) -> list[int | None]:
        """Style komórek nagłówka w kolumnach wyglądu (od A); ``None``: komórki nie ma w XML."""

        styles: list[int | None] = [None] * self._layout.look_columns
        position = bisect_left(self._numbers, self._layout.header_row)
        if position == len(self._rows) or self._rows[position].number != self._layout.header_row:
            return styles
        row = self._rows[position]
        for cell in _CELL.finditer(self._xml, row.start, row.end):
            column = _column(cell.group(1))
            style = _STYLE.search(cell.group(3))
            if column <= self._layout.look_columns and style is not None:
                styles[column - 1] = int(style.group(1))
        return styles

    def read_filter_columns(self, shared_strings: list[str]) -> None:
        """Kolumny autofiltra Pliku zbiorczego: od pierwszej do ostatniej niepustej komórki wiersza nagłówka."""

        position = bisect_left(self._numbers, self._layout.header_row)
        if position == len(self._rows) or self._rows[position].number != self._layout.header_row:
            raise UnsupportedShape("header row")
        row = self._rows[position]
        columns = []
        for cell in _CELL.finditer(self._xml, row.start, row.end):
            value = _header_value(cell.group(3), cell.group(4), shared_strings)
            if value is not None and not (isinstance(value, str) and not value.strip()):
                columns.append(_column(cell.group(1)))
        if not columns:
            raise UnsupportedShape("header row")
        self._filter_columns = (min(columns), max(columns))

    def filter_ref(self, row_count: int) -> str | None:
        """Zakres autofiltra Pliku zbiorczego: nagłówek i wpisane wiersze; ``None``: autofiltr zostaje z Wzorca."""

        if self._filter is None:
            return None
        first, last = self._filter_columns
        header_row = self._layout.header_row
        return f"{get_column_letter(first)}{header_row}:{get_column_letter(last)}{header_row + row_count}"

    def links(self, rows: Sequence[Row]) -> list[tuple[bytes, Hyperlink]]:
        """Hiperłącza wpisywanych komórek kolumn Wiersza danych: adres komórki → Hiperłącze, wierszami."""

        first_row = self._layout.first_data_row
        return [
            (self._letters[column - 1] + b"%d" % (first_row + index), link)
            for index, row in enumerate(rows)
            for column, link in sorted(_row_links(row))
            if column <= self._layout.data_columns
        ]

    def hyperlinks(self, links: Sequence[tuple[bytes, Hyperlink]], ids: Sequence[bytes]) -> bytes:
        """Blok ``<hyperlinks>`` z odnośnikiem każdej komórki do relacji o tym samym miejscu w ``ids``; pusty bez nich."""

        if not links:
            return b""
        entries = b"".join(
            b'<hyperlink ref="' + reference + b'" ' + self._links_prefix + b':id="' + rid + b'"/>'
            for (reference, _), rid in zip(links, ids, strict=True)
        )
        return self._links_open + entries + b"</hyperlinks>"

    def with_rows(
        self, rows: Sequence[Row], styles: _StyleAdditions, look: RowLook, hyperlinks: bytes, filter_ref: str | None
    ) -> _Part:
        xml = self._xml
        first_row = self._layout.first_data_row
        last_row = first_row + len(rows) - 1
        existing = self._rows[self._first_data_index : bisect_right(self._numbers, last_row)]
        by_number = {row.number: row for row in existing}
        region_end = existing[-1].end if existing else self._data_start
        patched = [
            self._patched_row(by_number.get(first_row + index), first_row + index, row, styles, look)
            for index, row in enumerate(rows)
        ]
        chunk = bisect_left(self._boundaries, region_end)
        tail_start = self._boundaries[chunk]
        links_at = self._links_at
        last_column = max(
            [self._layout.data_columns, *(column_index_from_string(letter) for row in rows for letter in row.quantities)]
        )
        head = b"".join([self._prefix(last_row if rows else None, last_column), *patched, xml[region_end:tail_start]])
        if self._chunks is None:
            self._chunks = _compressed_chunks(xml, self._boundaries)
        compressor = zlib.compressobj(_COMPRESSION_LEVEL, zlib.DEFLATED, -15)
        tail = self._chunks[-1]
        if hyperlinks:
            tail_compressor = zlib.compressobj(_COMPRESSION_LEVEL, zlib.DEFLATED, -15)
            tail = tail_compressor.compress(hyperlinks + xml[links_at:]) + tail_compressor.flush()
        compressed = [compressor.compress(head) + compressor.flush(zlib.Z_SYNC_FLUSH), *self._chunks[chunk:-1], tail]
        middle: tuple[bytes | memoryview, ...] = (memoryview(xml)[tail_start:links_at],)
        if self._filter is not None and filter_ref is not None:
            # Przedostatni kawałek zaczyna się od autofiltra (granica z ``__init__``) i sięga do bloku Hiperłączy.
            start, end, element = self._filter
            filtered = element.replace(b"{ref}", filter_ref.encode("ascii")) + xml[end:links_at]
            filter_compressor = zlib.compressobj(_COMPRESSION_LEVEL, zlib.DEFLATED, -15)
            compressed[-2] = filter_compressor.compress(filtered) + filter_compressor.flush(zlib.Z_FULL_FLUSH)
            middle = (memoryview(xml)[tail_start:start], filtered)
        crc = zlib.crc32(head)
        for piece in (*middle, hyperlinks, memoryview(xml)[links_at:]):
            crc = zlib.crc32(piece, crc)
        name, flags = _encoded_name(self._info.filename)
        dos_time, dos_date = _dos_stamp(self._info)
        size = len(head) + sum(len(piece) for piece in middle) + len(hyperlinks) + len(xml) - links_at
        return _Part(name, flags, zipfile.ZIP_DEFLATED, dos_time, dos_date, crc, size, compressed)

    def _prefix(self, last_row: int | None, last_column: int) -> bytes:
        """Arkusz do pierwszego Wiersza danych; ``dimension`` obejmuje wpisane komórki, jeśli wcześniej nie obejmował."""

        prefix = self._xml[: self._data_start]
        if self._dimension is None or last_row is None:
            return prefix
        start, end, (min_column, min_row, max_column, max_row) = self._dimension
        covering = (
            min(min_column or 1, 1),
            min(min_row or 1, self._layout.first_data_row),
            max(max_column or 1, last_column),
            max(max_row or 1, last_row),
        )
        if covering == (min_column, min_row, max_column, max_row):
            return prefix
        reference = f"{get_column_letter(covering[0])}{covering[1]}:{get_column_letter(covering[2])}{covering[3]}"
        return prefix[:start] + reference.encode("ascii") + prefix[end:]

    def _patched_row(
        self, row: SheetRow | None, number: int, data: Row, styles: _StyleAdditions, look: RowLook
    ) -> bytes:
        cells: dict[int, bytes] = {}
        if row is None:
            start_tag = b'<row r="%d">' % number
        else:
            start_tag, cells = _split_row(self._xml[row.start : row.end], number)
            start_tag = _HEIGHT.sub(b"", start_tag)
        suffix = b"%d" % number
        columns = self._layout.data_columns
        written = dict(cells)
        for column in range(1, columns + 1):
            value, highlights, number_format, _ = data.cells[column]
            cell_look = None
            if column <= look.columns:
                cell_look = look.cell(column, highlights, number_format, wrapped=wraps(value))
            reference = self._letters[column - 1] + suffix
            written[column] = self._cell(reference, cells.get(column), value, highlights, cell_look, styles)
        for column in range(columns + 1, look.columns + 1):
            reference = self._letters[column - 1] + suffix
            existing = cells.get(column)
            style = styles.look_style(_cell_style(existing), look.cell(column, PLAIN, GENERAL, wrapped=False))
            written[column] = _restyled(reference, existing, style)
        for letter, quantity in data.quantities.items():
            column = column_index_from_string(letter)
            reference = self._letters[column - 1] + suffix
            written[column] = _cell_xml(reference, _cell_style(cells.get(column)), *_value_xml(quantity))
        return start_tag + b"".join(cell for _, cell in sorted(written.items())) + b"</row>"

    def _cell(
        self,
        reference: bytes,
        existing: bytes | None,
        value: object,
        highlights: Highlights,
        look: CellLook | None,
        styles: _StyleAdditions,
    ) -> bytes:
        template_style = _cell_style(existing)

        def style(date_format: str | None) -> int | None:
            if look is not None:
                return styles.look_style(template_style, look)
            return styles.style(template_style, date_format, highlights)

        if isinstance(value, str) and value[:_MAX_TEXT] == "":
            value = None
        if value is None:
            empty_style = style(None)
            if empty_style == template_style and existing is not None and existing.endswith(b"/>"):
                return existing
            if empty_style is None and existing is None:
                return b""  # openpyxl nie zapisuje pustej komórki bez stylu; Excel daje jej styl kolumny
            return _cell_xml(reference, empty_style, None, b"")
        if isinstance(value, (datetime, date, time, timedelta)):
            if getattr(value, "tzinfo", None) is not None:
                raise InvalidValue("strefa czasowa")
            date_format = None if is_date_format(styles.number_format(template_style)) else _date_format(value)
            serial: float = to_excel(value, self._epoch)  # type: ignore[no-untyped-call]
            return _cell_xml(reference, style(date_format), b"n", b"<v>" + _number(serial) + b"</v>")
        kind, content = _value_xml(value)
        return _cell_xml(reference, style(None), kind, content)


def _row_links(row: Row) -> Iterator[tuple[int, Hyperlink]]:
    return ((column, cell.hyperlink) for column, cell in row.cells.items() if cell.hyperlink is not None)


def _writable_data_rows_have_input(xml: bytes, data_start: int, data_end: int, layout: SheetLayout) -> bool:
    """Czy Wiersze danych mają już wartość w kolumnach danych wejściowych; jedno przejście po komórkach.

    Kolumny Wiersza danych od pierwszego Wiersza danych nie mogą mieć formuł ani
    scaleń, a kolumny Ilości usług formuł: wpisanie wartości w taką komórkę zmienia
    ``calcChain`` albo wspólną formułę (formuła) lub nie jest możliwe (scalenie); to
    robi dawna droga przez openpyxl (``UnsupportedShape``). Formuły nie są danymi.
    """

    has_input = False
    for cell in _CELL_WITH_CONTENT.finditer(xml, data_start, data_end):
        column = _column(cell.group(1))
        if column > layout.data_columns and column not in layout.quantity_columns:
            continue
        content = cell.group(2)
        if _FORMULA.search(content):
            raise UnsupportedShape("formula in data row")
        if not has_input and column <= layout.input_columns:
            has_input = b"<is" in content or _NON_EMPTY_VALUE.search(content) is not None
    for merged in _MERGE.finditer(xml, data_end):
        min_column, min_row, _max_column, max_row = range_boundaries(merged.group(1).decode("ascii"))
        if (min_column or 1) <= layout.data_columns and (max_row or min_row or 1) >= layout.first_data_row:
            raise UnsupportedShape("merged data rows")
    return has_input


def _hyperlinks_place(xml: bytes, data_end: int) -> tuple[int, bytes, bytes]:
    """Miejsce bloku ``<hyperlinks>`` w arkuszu, jego znacznik otwierający i przedrostek przestrzeni nazw relacji.

    Miejsce to początek pierwszego elementu za ``sheetData``, który schemat stawia za
    ``hyperlinks``, a bez takiego koniec arkusza. Ogon arkusza jest mały, więc jest
    parsowany raz, z deklaracjami przestrzeni nazw z korzenia. Wzorzec z własnymi
    hiperłączami to ``UnsupportedShape``.
    """

    root = _WORKSHEET_START.search(xml)
    if root is None:
        raise UnsupportedShape("worksheet")
    after_data = data_end + len(b"</sheetData>")
    document = root.group(0) + xml[after_data:]
    places: list[tuple[str, int]] = []
    depth = 0
    parser = expat.ParserCreate()

    def start(name: str, _attributes: dict[str, str]) -> None:
        nonlocal depth
        depth += 1
        if depth == 2:
            places.append((name, parser.CurrentByteIndex - len(root.group(0)) + after_data))

    def end(_name: str) -> None:
        nonlocal depth
        depth -= 1

    parser.StartElementHandler = start
    parser.EndElementHandler = end
    parser.Parse(document, True)
    if any(name == "hyperlinks" for name, _ in places):
        raise UnsupportedShape("hyperlinks in template")
    after = [position for name, position in places if name not in _BEFORE_HYPERLINKS]
    position = after[0] if after else xml.rindex(b"</" + root.group(1) + b"worksheet>")
    bound = _RELATIONSHIPS_PREFIX.search(root.group(0))
    if bound is not None:
        return position, b"<hyperlinks>", bound.group(1)
    return position, b'<hyperlinks xmlns:r="' + OFFICE_REL.encode() + b'">', b"r"


def _cell_style(cell: bytes | None) -> int | None:
    """Styl komórki z XML arkusza; ``None``: komórki nie ma albo nie ma atrybutu ``s``."""

    if cell is None:
        return None
    style = _STYLE.search(cell, 0, cell.index(b">"))
    return None if style is None else int(style.group(1))


def _restyled(reference: bytes, cell: bytes | None, style: int | None) -> bytes:
    """Komórka ``cell`` ze stylem ``style`` i treścią bez zmian; brakująca komórka powstaje pusta."""

    if cell is None:
        return _cell_xml(reference, style, None, b"")
    tag_end = cell.index(b">")
    tag = _STYLE.sub(b"", cell[:tag_end])
    reference_end = tag.index(b'"', len(b'<c r="')) + 1
    attribute = b"" if style is None else b' s="%d"' % style
    return tag[:reference_end] + attribute + tag[reference_end:] + cell[tag_end:]


_AUTOFILTER = re.compile(rb"<autoFilter\b([^>]*?)(/?)>")
_REF = re.compile(rb'\sref="[^"]*"')


def _autofilter(xml: bytes, data_end: int, links_at: int) -> tuple[int, int, bytes]:
    """Miejsce elementu ``autoFilter`` za ``sheetData`` i jego nowa postać z ``{ref}`` w miejscu zakresu.

    Jak ``AutoFilter(ref=…)`` w openpyxl: warunki filtrów Wzorca nie zostają. Wzorzec bez autofiltra
    to ``UnsupportedShape``.
    """

    found = _AUTOFILTER.search(xml, data_end, links_at)
    if found is None:
        raise UnsupportedShape("autofilter")
    end = found.end() if found.group(2) else xml.index(b"</autoFilter>", found.end()) + len(b"</autoFilter>")
    attributes = _REF.sub(b"", found.group(1))
    return found.start(), end, b'<autoFilter ref="{ref}"' + attributes + b"/>"


def _compressed_chunks(xml: bytes, boundaries: list[int]) -> list[bytes]:
    """Kawałki od każdej granicy do następnej, każdy z własnego kompresora i z pełnym opróżnieniem, więc nie sięga wstecz.

    Ostatni kończy strumień. Kawałki nie zależą od siebie, a zlib kompresuje bez blokady GIL, więc powstają równolegle.
    """

    view = memoryview(xml)
    ends = [*boundaries[1:], len(xml)]

    def compressed(index: int) -> bytes:
        compressor = zlib.compressobj(_COMPRESSION_LEVEL, zlib.DEFLATED, -15)
        last = index == len(boundaries) - 1
        return compressor.compress(view[boundaries[index] : ends[index]]) + compressor.flush(
            zlib.Z_FINISH if last else zlib.Z_FULL_FLUSH
        )

    with ThreadPoolExecutor() as pool:
        return list(pool.map(compressed, range(len(boundaries))))


def _split_row(row_xml: bytes, number: int) -> tuple[bytes, dict[int, bytes]]:
    """Znacznik otwierający wiersza i jego komórki według kolumny; inna treść wiersza to ``UnsupportedShape``."""

    tag_end = row_xml.index(b">")
    if row_xml[tag_end - 1 : tag_end] == b"/":
        return row_xml[: tag_end - 1].rstrip() + b">", {}
    body_start, body_end = tag_end + 1, len(row_xml) - len(b"</row>")
    cells: dict[int, bytes] = {}
    position = body_start
    previous = 0
    for cell in _CELL.finditer(row_xml, body_start, body_end):
        column = _column(cell.group(1))
        if row_xml[position : cell.start()].strip() or int(cell.group(2)) != number or column <= previous:
            raise UnsupportedShape("row content")
        cells[column] = cell.group(0)
        position, previous = cell.end(), column
    if row_xml[position:body_end].strip():
        raise UnsupportedShape("row content")
    return row_xml[: tag_end + 1], cells


def _header_value(attributes: bytes, content: bytes | None, shared_strings: list[str]) -> object:
    """Wartość komórki nagłówka na tyle, na ile potrzeba do sprawdzenia WYKONAWCA."""

    if not content:
        return None
    if _FORMULA.search(content):
        return "="
    kind = _TYPE.search(attributes)
    if kind is not None and kind.group(1) == b"inlineStr":
        return _plain_text(ElementTree.fromstring(content[content.index(b"<is") : content.rindex(b"</is>") + 5]))
    value = _VALUE.search(content)
    if value is None:
        return None
    text = html.unescape(value.group(1).decode("utf-8"))
    if kind is not None and kind.group(1) == b"s":
        return shared_strings[int(text)]
    return text


class _SheetRels:
    """Relacje arkusza Wzorca: relacje Hiperłączy kopii są dopisywane na końcu, istniejące zostają bajt w bajt."""

    def __init__(self, info: zipfile.ZipInfo, index: int | None, xml: bytes | None, addable: bool) -> None:
        self._info = info
        self.index = index
        """Miejsce części w zip Wzorca; ``None``: Wzorzec jej nie ma, a kopia z Hiperłączami dopisuje ją na końcu."""
        self._xml = xml
        numbers = [0]
        if xml is not None:
            for relationship in ElementTree.fromstring(xml).iter(PACKAGE_REL):
                number = _RELATIONSHIP_NUMBER.fullmatch(relationship.get("Id", ""))
                if number is not None:
                    numbers.append(int(number.group(1)))
        self._first = max(numbers) + 1
        self._addable = addable and (xml is None or b"</Relationships>" in xml)

    def with_links(self, links: Sequence[Hyperlink]) -> tuple[_Part | None, list[bytes]]:
        """Część relacji z relacją każdego Hiperłącza i ich identyfikatory; bez Hiperłączy ``None``, część bez zmian."""

        if not links:
            return None, []
        if not self._addable:
            raise UnsupportedShape("sheet rels")
        ids = [b"rId%d" % number for number in range(self._first, self._first + len(links))]
        entries = b"".join(
            b'<Relationship Id="' + rid + b'" Type="' + OFFICE_REL.encode() + b'/hyperlink" Target="'
            + _escaped_attribute(link.address) + b'" TargetMode="External"/>'
            for rid, link in zip(ids, links, strict=True)
        )
        if self._xml is None:
            xml = (
                b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
                b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                + entries
                + b"</Relationships>"
            )
        else:
            end = self._xml.rindex(b"</Relationships>")
            xml = self._xml[:end] + entries + self._xml[end:]
        return _compressed_part(self._info, xml), ids


# --- wartości ----------------------------------------------------------------------------------


def _cell_xml(reference: bytes, style: int | None, kind: bytes | None, content: bytes) -> bytes:
    attributes = b'<c r="' + reference + b'"'
    if style is not None:
        attributes += b' s="%d"' % style
    if kind is not None:
        attributes += b' t="' + kind + b'"'
    return attributes + (b">" + content + b"</c>" if content else b"/>")


def _date_format(value: datetime | date | time | timedelta) -> str:
    """Format daty, który openpyxl nadaje komórce z datą bez formatu daty."""

    if isinstance(value, datetime):
        return FORMAT_DATE_DATETIME
    if isinstance(value, date):
        return FORMAT_DATE_YYYYMMDD2
    if isinstance(value, time):
        return FORMAT_DATE_TIME6
    return FORMAT_DATE_TIMEDELTA


def _number(value: int | float | Decimal) -> bytes:
    """Liczba jak w zapisie openpyxl (``safe_string``): 16 cyfr znaczących; NaN i nieskończoność są puste."""

    number = float(value)
    if math.isnan(number) or math.isinf(number):
        return b""
    return ("%.16g" % value).encode("ascii")


def _escaped_attribute(text: str) -> bytes:
    """Tekst wartości atrybutu XML; tabulacja i końce linii jako odwołania, bo parser zamieniłby je na spacje."""

    escaped = html.escape(text, quote=True).replace("\t", "&#9;").replace("\n", "&#10;").replace("\r", "&#13;")
    try:
        return escaped.encode("utf-8")
    except UnicodeEncodeError as exc:
        raise InvalidValue("kodowanie") from exc


def _escaped(text: str) -> bytes:
    try:
        return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").encode("utf-8")
    except UnicodeEncodeError as exc:
        raise InvalidValue("kodowanie") from exc


def _value_xml(value: object) -> tuple[bytes | None, bytes]:
    """Typ i treść komórki dla wartości innej niż pusta i data, jak przy wpisaniu przez openpyxl."""

    if isinstance(value, bool):
        return b"b", b"<v>1</v>" if value else b"<v>0</v>"
    if isinstance(value, (int, float, Decimal)):
        return b"n", b"<v>" + _number(value) + b"</v>"
    if isinstance(value, str):
        text = value[:_MAX_TEXT]
        if ILLEGAL_CHARACTERS_RE.search(text):
            raise InvalidValue("znak niedozwolony w XML")
        if len(text) > 1 and text.startswith("="):
            return None, b"<f>" + _escaped(text[1:]) + b"</f>"
        if text in ERROR_CODES:
            return b"e", b"<v>" + _escaped(text) + b"</v>"
        space = b' xml:space="preserve"' if text != text.strip() else b""
        return b"inlineStr", b"<is><t" + space + b">" + _escaped(text) + b"</t></is>"
    raise UnsupportedShape(type(value).__name__)


# --- style -------------------------------------------------------------------------------------


# Styl komórki bez atrybutu ``s``: taki, jaki openpyxl nadaje nowej komórce.
_DEFAULT_XF: Final = b'<xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>'
_FONT_ID: Final = re.compile(rb'\sfontId="(\d+)"')
_FILL_ID: Final = re.compile(rb'\sfillId="(\d+)"')
_BORDER_ID: Final = re.compile(rb'\sborderId="(\d+)"')


class _Styles:
    """``xl/styles.xml`` Wzorca: style komórek z ich formatami liczb, czcionkami, wypełnieniami i obramowaniami.

    Czcionki i wypełnienia z nałożonymi Wyróżnieniami są liczone raz na Wzorzec, wspólnie dla wszystkich kopii.
    """

    def __init__(self, info: zipfile.ZipInfo, index: int, xml: bytes, unchanged: _Part) -> None:
        self._info = info
        self.index = index
        self.xml = xml
        cell_xfs = _CELL_XFS.search(xml)
        if cell_xfs is None:
            raise UnsupportedShape("cellXfs")
        self.xfs: list[bytes] = _XF.findall(cell_xfs.group(1))
        self.custom_formats: dict[int, str] = {}
        for number_format in _NUM_FMT.finditer(xml):
            attributes = {match.group(1): match.group(2) or match.group(3) or b"" for match in _ATTRIBUTE.finditer(number_format.group(1))}
            self.custom_formats[int(attributes[b"numFmtId"])] = html.unescape(attributes[b"formatCode"].decode("utf-8"))
        root = ElementTree.fromstring(xml)
        fonts = root.find(f"{MAIN}fonts")
        fills = root.find(f"{MAIN}fills")
        borders = root.find(f"{MAIN}borders")
        if fonts is None or fills is None or borders is None or None in (_FONTS.search(xml), _FILLS.search(xml), _BORDERS.search(xml)):
            raise UnsupportedShape("fonts")
        self.fonts = [Font.from_tree(cast(Any, font)) for font in fonts.iterfind(f"{MAIN}font")]
        self.fills = [Fill.from_tree(cast(Any, fill)) or PatternFill() for fill in fills.iterfind(f"{MAIN}fill")]
        self.borders = [Border.from_tree(cast(Any, border)) or Border() for border in borders.iterfind(f"{MAIN}border")]
        self.font_ids = _first_ids(self.fonts)
        self.fill_ids = _first_ids(self.fills)
        self.border_ids = _first_ids(self.borders)
        self.xf_ids = {xf: index for index, xf in reversed(list(enumerate(self.xfs)))}
        self._fonts_over: dict[tuple[int, Highlights], Font | None] = {}
        self._fills_over: dict[tuple[int, Highlights], Fill | None] = {}
        self._unchanged = unchanged

    def xf(self, style: int | None) -> bytes:
        if style is None:
            return _DEFAULT_XF
        if style >= len(self.xfs):
            raise UnsupportedShape("style")
        return self.xfs[style]

    def font(self, style: int | None) -> Font:
        return _item(self.fonts, _xf_id(self.xf(style), _FONT_ID))

    def alignment(self, style: int | None) -> Alignment:
        found = _ALIGNMENT.search(self.xf(style))
        if found is None:
            return Alignment()
        return Alignment.from_tree(cast(Any, ElementTree.fromstring(found.group(0)))) or Alignment()

    def format_id(self, style: int) -> int:
        xf = self.xf(style)
        found = _NUM_FMT_ID.search(xf, 0, xf.index(b">"))
        return int(found.group(1)) if found else 0

    def font_over(self, font_id: int, highlights: Highlights) -> Font | None:
        """Czcionka ``font_id`` z Wyróżnieniami; ``None``, gdy zostaje bez zmian."""

        key = (font_id, highlights)
        if key not in self._fonts_over:
            base = _item(self.fonts, font_id)
            font = highlights.font_over(base)
            self._fonts_over[key] = None if font == base else font
        return self._fonts_over[key]

    def fill_over(self, fill_id: int, highlights: Highlights) -> Fill | None:
        """Wypełnienie ``fill_id`` zastąpione Wyróżnieniami; ``None``, gdy zostaje bez zmian."""

        key = (fill_id, highlights)
        if key not in self._fills_over:
            base = _item(self.fills, fill_id)
            fill = highlights.fill_over(base)
            self._fills_over[key] = None if fill == base else fill
        return self._fills_over[key]

    def part(self, additions: _StyleAdditions) -> _Part:
        if not additions.xfs:
            return self._unchanged
        xml = self.xml
        fonts = [ElementTree.tostring(font.to_tree()) for font in additions.fonts]
        fills = [ElementTree.tostring(fill.to_tree()) for fill in additions.fills]
        borders = [ElementTree.tostring(border.to_tree()) for border in additions.borders]
        xml = _appended(xml, _FONTS, b"fonts", fonts, len(self.fonts))
        xml = _appended(xml, _FILLS, b"fills", fills, len(self.fills))
        xml = _appended(xml, _BORDERS, b"borders", borders, len(self.borders))
        xml = _appended(xml, _CELL_XFS, b"cellXfs", additions.xfs, len(self.xfs))
        if additions.formats:
            entries = b"".join(
                b'<numFmt numFmtId="%d" formatCode="' % number + html.escape(code).encode("utf-8") + b'"/>'
                for number, code in additions.formats
            )
            count = len(self.custom_formats) + len(additions.formats)
            block = re.search(rb"<numFmts\b[^>]*?(?:/>|>(.*?)</numFmts>)", xml, re.S)
            if block is None:
                root = re.search(rb"<styleSheet\b[^>]*>", xml)
                if root is None:
                    raise UnsupportedShape("styleSheet")
                xml = xml[: root.end()] + b'<numFmts count="%d">' % count + entries + b"</numFmts>" + xml[root.end() :]
            elif block.group(1) is None:
                xml = xml[: block.start()] + b'<numFmts count="%d">' % count + entries + b"</numFmts>" + xml[block.end() :]
            else:
                xml = xml[: block.end(1)] + entries + xml[block.end(1) :]
                xml = _with_count(xml, b"numFmts", count)
        return _compressed_part(self._info, xml)


def _first_ids(items: Sequence[Font] | Sequence[Fill] | Sequence[Border]) -> dict[object, int]:
    ids: dict[object, int] = {}
    for index, item in enumerate(items):
        ids.setdefault(item, index)
    return ids


def _item(items: Sequence[_T], index: int) -> _T:
    if index >= len(items):
        raise UnsupportedShape("style")
    return items[index]


def _appended(xml: bytes, block: re.Pattern[bytes], tag: bytes, entries: Sequence[bytes], existing: int) -> bytes:
    """``entries`` dopisane na końcu listy ``tag`` (pierwszej w pliku) z licznikiem ``existing`` + ich liczba."""

    if not entries:
        return xml
    found = block.search(xml)
    assert found is not None
    xml = xml[: found.end(1)] + b"".join(entries) + xml[found.end(1) :]
    return _with_count(xml, tag, existing + len(entries))


def _with_count(xml: bytes, tag: bytes, count: int) -> bytes:
    start = re.search(rb"<" + tag + rb"\b[^>]*>", xml)
    assert start is not None
    opening = re.sub(rb'\scount="\d+"', b"", start.group(0))
    opening = opening[: len(tag) + 1] + b' count="%d"' % count + opening[len(tag) + 1 :]
    return xml[: start.start()] + opening + xml[start.end() :]


class _StyleAdditions:
    """Style dopisane dla jednej kopii: Wygląd Wiersza danych, format daty jak w openpyxl i Wyróżnienia inne niż w
    komórce Wzorca."""

    def __init__(self, styles: _Styles) -> None:
        self._styles = styles
        self.xfs: list[bytes] = []
        self.formats: list[tuple[int, str]] = []
        self.fonts: list[Font] = []
        self.fills: list[Fill] = []
        self.borders: list[Border] = []
        self._added: dict[tuple[int | None, str | None, Highlights], int | None] = {}
        self._looks: dict[tuple[int | None, CellLook], int | None] = {}
        self._xf_ids: dict[bytes, int] = {}

    def number_format(self, style: int | None) -> str:
        if style is None:
            return "General"
        number = self._styles.format_id(style)
        return self._styles.custom_formats.get(number) or BUILTIN_FORMATS.get(number, "General")

    def style(self, style: int | None, number_format: str | None, highlights: Highlights) -> int | None:
        """Styl wpisanej komórki: styl Wzorca ``style`` z formatem ``number_format`` (``None``: bez zmiany) i Wyróżnieniami."""

        key = (style, number_format, highlights)
        if key not in self._added:
            self._added[key] = self._resolved(style, number_format, highlights)
        return self._added[key]

    def _resolved(self, style: int | None, number_format: str | None, highlights: Highlights) -> int | None:
        base = self._styles.xf(style)
        xf = base
        if number_format is not None and self.number_format(style) != number_format:
            xf = _with_id(xf, b"numFmtId", self._format_id(number_format), b"applyNumberFormat")
        font = self._styles.font_over(_xf_id(xf, _FONT_ID), highlights)
        if font is not None:
            font_id = _id(self._styles.font_ids, len(self._styles.fonts), self.fonts, font)
            xf = _with_id(xf, b"fontId", font_id, b"applyFont")
        fill = self._styles.fill_over(_xf_id(xf, _FILL_ID), highlights)
        if fill is not None:
            fill_id = _id(self._styles.fill_ids, len(self._styles.fills), self.fills, fill)
            xf = _with_id(xf, b"fillId", fill_id, b"applyFill")
        if xf is base:
            return style
        self.xfs.append(xf)
        return len(self._styles.xfs) + len(self.xfs) - 1

    def look_style(self, style: int | None, look: CellLook) -> int | None:
        """Styl komórki Wzorca ``style`` z wyglądem ``look``; ochrona, styl nazwany i reszta stylu zostają z Wzorca."""

        key = (style, look)
        if key not in self._looks:
            self._looks[key] = self._look_resolved(style, look)
        return self._looks[key]

    def _look_resolved(self, style: int | None, look: CellLook) -> int | None:
        styles = self._styles
        base = styles.xf(style)
        xf = base
        if self.number_format(style) != look.number_format:
            xf = _with_id(xf, b"numFmtId", self._format_id(look.number_format), b"applyNumberFormat")
        if _item(styles.fonts, _xf_id(xf, _FONT_ID)) != look.font:
            xf = _with_id(xf, b"fontId", _id(styles.font_ids, len(styles.fonts), self.fonts, look.font), b"applyFont")
        if _item(styles.fills, _xf_id(xf, _FILL_ID)) != look.fill:
            xf = _with_id(xf, b"fillId", _id(styles.fill_ids, len(styles.fills), self.fills, look.fill), b"applyFill")
        if _item(styles.borders, _xf_id(xf, _BORDER_ID)) != look.border:
            border_id = _id(styles.border_ids, len(styles.borders), self.borders, look.border)
            xf = _with_id(xf, b"borderId", border_id, b"applyBorder")
        if styles.alignment(style) != look.alignment:
            xf = _with_alignment(xf, ElementTree.tostring(look.alignment.to_tree()))
        if xf is base:
            return style
        if xf in styles.xf_ids:
            return styles.xf_ids[xf]
        if xf not in self._xf_ids:
            self.xfs.append(xf)
            self._xf_ids[xf] = len(styles.xfs) + len(self.xfs) - 1
        return self._xf_ids[xf]

    def _format_id(self, code: str) -> int:
        if code in BUILTIN_FORMATS_REVERSE:
            return int(BUILTIN_FORMATS_REVERSE[code])
        known = {existing: number for number, existing in [*self._styles.custom_formats.items(), *self.formats]}
        if code in known:
            return known[code]
        number = max([163, *self._styles.custom_formats, *(added for added, _ in self.formats)]) + 1
        self.formats.append((number, code))
        return number


def _id(existing: dict[object, int], count: int, added: list[_T], item: _T) -> int:
    """Numer ``item`` na liście Wzorca o ``count`` pozycjach, a gdy go tam nie ma, wśród dopisanych za nią."""

    if item in existing:
        return existing[item]
    if item not in added:
        added.append(item)
    return count + added.index(item)


def _xf_id(xf: bytes, attribute: re.Pattern[bytes]) -> int:
    found = attribute.search(xf, 0, xf.index(b">"))
    return int(found.group(1)) if found else 0


def _with_alignment(xf: bytes, alignment: bytes) -> bytes:
    """Styl ``xf`` z elementem ``alignment`` jako pierwszym dzieckiem (CT_Xf) i flagą ``applyAlignment``."""

    tag_end = xf.index(b">")
    if xf[tag_end - 1 : tag_end] == b"/":
        tag, children = xf[: tag_end - 1].rstrip(), b""
    else:
        tag, children = xf[:tag_end], _ALIGNMENT.sub(b"", xf[tag_end + 1 : xf.rindex(b"</xf>")])
    tag = re.sub(rb'\sapplyAlignment="\w+"', b"", tag)
    return tag + b' applyAlignment="1">' + alignment + children + b"</xf>"


def _with_id(xf: bytes, name: bytes, value: int, apply: bytes) -> bytes:
    """Styl ``xf`` z atrybutem ``name`` równym ``value`` i flagą ``apply``; reszta stylu bez zmian."""

    tag_end = xf.index(b">")
    tag = xf[:tag_end]
    tag = re.sub(rb"\s" + name + rb'="\d+"', b"", tag)
    tag = re.sub(rb"\s" + apply + rb'="\w+"', b"", tag)
    tag = tag.rstrip(b"/ ")
    closing = xf[tag_end:] if not xf[tag_end - 1 : tag_end] == b"/" else b"/>"
    return tag[:3] + b" " + name + b'="%d"' % value + tag[3:] + b" " + apply + b'="1"' + closing
