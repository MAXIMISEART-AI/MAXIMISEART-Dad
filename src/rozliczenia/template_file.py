"""Plik z Wzorca: Sprawdzanie, Planowanie i zapis Wierszy danych w kopii Wzorca jednego z trzech plików (ADR 0008).

Rozliczenie pracownika, Plik zbiorczy i Rozliczenie wewnętrzne powstają z Wzorca (``Layout``) przez
``TemplateFile``: ``check`` to Sprawdzanie, ``plan`` to Planowanie, a ``create`` zapisuje plik atomowo
(``workbook_save``). Plik powstaje kopią bajtów Wzorca (``template_copy``), gdy kopia przyjmuje kształt Wzorca i
wierszy, a inaczej przez openpyxl; obie drogi dają ten sam plik. Wzorzec czyta moduł skoroszytu
(``workbook_files``) z jego klasyfikacją uszkodzeń, a zły układ to ``WorkbookError`` z miejscem.
"""

from __future__ import annotations

import re
from collections.abc import Callable, Sequence
from contextlib import AbstractContextManager, suppress
from copy import copy
from dataclasses import dataclass
from enum import Enum
from functools import lru_cache
from pathlib import Path
from typing import cast
from xml.parsers import expat

from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.styles.fills import Fill
from openpyxl.styles.fonts import Font
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import IllegalCharacterError
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.cell_range import CellRange, MultiCellRange
from openpyxl.worksheet.filters import AutoFilter
from openpyxl.worksheet.worksheet import Worksheet

from .highlights import PLAIN
from .row_look import GENERAL, RowLook, wraps
from .template_copy import InvalidValue, Package, SheetLayout, TemplateCopy, UnsupportedShape
from .workbook_files import (
    DATA_ROW_COLUMNS,
    FIRST_DATA_ROW,
    HEADER_ROW,
    INPUT_COLUMNS,
    LOOK_COLUMNS,
    QUANTITY_COLUMNS,
    DataRow,
    MergedCells,
    MissingHeader,
    Problem,
    WorkbookError,
    check_worker_header,
    classified,
    ensure_available,
    header_columns,
    is_blank,
    open_workbook,
    unreadable_as_error,
    workbook_content,
)
from .workbook_parts import package_parts, sheet_part
from .workbook_save import WorkbookSaveError, save_from_template, write_atomically
from .worker_config import normalize_text


# Rozszerzenia skoroszytów, które otwiera openpyxl; inny plik jest nieczytelny jak przy otwieraniu przez openpyxl.
_WORKBOOK_SUFFIXES = (".xlsx", ".xlsm", ".xltx", ".xltm")
_INTERNAL_LOOK_COLUMNS = range(1, 14)  # A:M Rozliczenia wewnętrznego, z ``zwrot`` i ``UWAGI``


@dataclass(frozen=True)
class _Rules:
    """Układ Wzorca jednego pliku: jego arkusz (``SheetLayout``, ten sam dla kopii i drogi przez openpyxl) i dwie
    reguły, którymi zapis tego pliku różni się od pozostałych."""

    sheet: SheetLayout
    worker_header: bool
    """Nagłówek w kolumnach danych wejściowych z WYKONAWCA w H; bez niego nagłówek to cały wiersz."""
    copy: bool
    """Plik może powstać kopią Wzorca (ADR 0008). Kopia zawsze każe Excelowi przeliczyć formuły przy otwarciu, więc
    droga przez openpyxl robi wtedy to samo; Rozliczenie wewnętrzne powstaje tylko przez openpyxl, bez tego."""

    @property
    def last_header_column(self) -> int | None:
        """Ostatnia kolumna nagłówka; ``None``: cały wiersz nagłówka."""

        return self.sheet.input_columns if self.worker_header else None


_WORKER_SETTLEMENT = _Rules(
    SheetLayout(
        header_row=HEADER_ROW,
        first_data_row=FIRST_DATA_ROW,
        input_columns=INPUT_COLUMNS[-1],
        data_columns=DATA_ROW_COLUMNS[-1],
        look_columns=LOOK_COLUMNS[-1],
        period_file=False,
    ),
    worker_header=True,
    copy=True,
)
_COLLECTIVE_FILE = _Rules(
    SheetLayout(
        header_row=HEADER_ROW,
        first_data_row=FIRST_DATA_ROW,
        input_columns=INPUT_COLUMNS[-1],
        data_columns=INPUT_COLUMNS[-1],
        look_columns=LOOK_COLUMNS[-1],
        period_file=True,
        quantity_columns=QUANTITY_COLUMNS,
    ),
    worker_header=True,
    copy=True,
)
_INTERNAL_SETTLEMENT = _Rules(
    SheetLayout(
        header_row=1,
        first_data_row=2,
        input_columns=INPUT_COLUMNS[-1],
        data_columns=INPUT_COLUMNS[-1],
        look_columns=_INTERNAL_LOOK_COLUMNS[-1],
        period_file=True,
    ),
    worker_header=False,
    copy=False,
)


def _is_value(cell: Cell | MergedCell) -> bool:
    return cell.value is not None and cell.data_type != "f"


def _is_zero(value: object) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool) and value == 0


@dataclass(frozen=True)
class _Outline:
    """Wzorzec pliku okresu w tym, czego potrzebują Sprawdzanie i Planowanie."""

    title: str
    header: tuple[object, ...]
    """Wartości wiersza nagłówka od kolumny A."""
    merged_ranges: tuple[CellRange, ...]
    """Scalenia aktywnego arkusza w kolejności, w jakiej daje je openpyxl."""
    zero_row_count: int
    """Liczba kolejnych Wierszy danych z zerami w kolumnach ilości N:AP."""


def _read_outline(path: Path, rules: _Rules) -> _Outline:
    """Wzorzec do sprawdzenia bez wczytania każdej komórki; niczego nie zapisuje.

    Pełne wczytanie Wzorca pliku zbiorczego przez openpyxl trwa ~2 s, a Sprawdzanie i Planowanie
    potrzebują tylko nagłówka, scaleń i wierszy z zerami. openpyxl czyta więc tylko te wiersze, a reszta
    arkusza jest sprawdzana tak, żeby Wzorzec, którego pełne wczytanie by się nie udało, był nieczytelny.
    """

    return _outline(path, workbook_content(path), rules)


# Sprawdzanie i Planowanie Przygotowania czytają zarys tych samych Wzorców; zarys zależy tylko od bajtów,
# więc niezmieniony Wzorzec nie jest parsowany drugi raz, a zmieniony po Sprawdzaniu jest czytany od nowa.
@lru_cache(maxsize=4)
def _outline(path: Path, content: bytes, rules: _Rules) -> _Outline:
    with open_workbook(path, content, read_only=True) as workbook:
        sheet = cast(Worksheet, workbook.active)  # tylko do odczytu: ReadOnlyWorksheet
        header = next(
            classified(
                path,
                sheet.iter_rows(
                    min_row=rules.sheet.header_row,
                    max_row=rules.sheet.header_row,
                    max_col=rules.last_header_column,
                    values_only=True,
                ),
            ),
            (),
        )
        zero_row_count = 0
        for values in classified(
            path,
            sheet.iter_rows(
                min_row=rules.sheet.first_data_row,
                min_col=QUANTITY_COLUMNS[0],
                max_col=QUANTITY_COLUMNS[-1],
                values_only=True,
            ),
        ):
            if not any(_is_zero(value) for value in values):
                break
            zero_row_count += 1
        with unreadable_as_error(path):
            # openpyxl nie daje liczby stylów komórek inaczej niż przez atrybut prywatny.
            cell_style_count = len(workbook._cell_styles)  # type: ignore[attr-defined]
            merged_ranges = _merged_ranges(content, sheet.title, cell_style_count)
        return _Outline(sheet.title, header, merged_ranges, zero_row_count)


_MERGE_CELL = "http://schemas.openxmlformats.org/spreadsheetml/2006/main mergeCell"
_CELL_OR_ROW_STYLE = re.compile(rb'<(?:[\w.-]+:)?(?:c|row)\s[^>]*?\bs="([0-9]+)"')
_COLUMN_STYLE = re.compile(rb'<(?:[\w.-]+:)?col\s[^>]*?\bstyle="([0-9]+)"')
_ANY_STYLE = re.compile(rb's="([0-9]+)"')


def _merged_ranges(content: bytes, sheet_title: str, cell_style_count: int) -> tuple[CellRange, ...]:
    """Scalenia arkusza z bloku ``<mergeCells>`` za ``</sheetData>``.

    Arkusz musi dać się wczytać w całości jak przy pełnym wczytaniu przez openpyxl: zepsuty XML to
    ``ExpatError``, a styl komórki, wiersza albo kolumny spoza stylów skoroszytu ``IndexError``.
    """

    with package_parts(content) as parts:
        xml = parts[sheet_part(parts, sheet_title)]
    # Styl każdej komórki i wiersza jest też w zbiorze wszystkich ``s="N"``, który znajduje się ~3× szybciej; dokładne
    # wyrażenie rozstrzyga tylko wtedy, gdy ten szerszy zbiór ma styl spoza stylów skoroszytu.
    cell_or_row_styles = set(_ANY_STYLE.findall(xml))
    if max(map(int, cell_or_row_styles), default=0) >= cell_style_count:
        cell_or_row_styles = set(_CELL_OR_ROW_STYLE.findall(xml))
    styles = [int(index) for index in (*cell_or_row_styles, *set(_COLUMN_STYLE.findall(xml)))]
    if max(styles, default=0) >= cell_style_count:
        raise IndexError(f"styl {max(styles)} spoza {cell_style_count} stylów skoroszytu")
    references: list[str] = []
    parser = expat.ParserCreate(namespace_separator=" ")
    # Obsługa elementów dopiero za danymi arkusza: wywołanie Pythona dla każdej z ~500 tys. komórek kosztuje.
    tail = max(xml.rfind(b"sheetData"), 0)
    parser.Parse(xml[:tail], False)
    parser.StartElementHandler = lambda name, attributes: (
        references.append(attributes["ref"]) if name == _MERGE_CELL else None
    )
    parser.Parse(xml[tail:], True)
    return tuple(MultiCellRange([CellRange(reference) for reference in references]).ranges)


class Route(Enum):
    """Droga, którą powstał plik z Wzorca (ADR 0008); oba dają ten sam plik, a kopia jest szybsza."""

    COPY = "kopia"
    OPENPYXL = "openpyxl"


@dataclass(frozen=True)
class TemplatePlan:
    """Planowanie Wierszy danych we Wzorcu: wiersze, do których trafią, ile kolejnych Wierszy danych ma zera w N:AP
    i który wiersz jest za nimi."""

    data_rows: range
    zero_row_count: int
    first_row_without_zeros: int


class TemplateFile:
    """Wzorzec jednego układu: jego Sprawdzanie, Planowanie i pliki z Wierszami danych (ADR 0008).

    Powstaje dla jednego z trzech plików: ``worker_settlement``, ``collective_file`` albo ``internal_settlement``.
    ``headers`` pliku okresu to nagłówek docelowy każdej pozycji Wiersza danych, szukany w wierszu nagłówka Wzorca
    bez względu na wielkość liter i spacje; w Rozliczeniu pracownika pozycja ``i`` to kolumna ``i`` Wiersza danych
    A:AT. Utworzenie nie czyta dysku. Plik powstaje kopią bajtów Wzorca, gdy kopia przyjmuje kształt Wzorca i
    wierszy, a inaczej przez openpyxl; obie drogi dają ten sam plik. Pierwszy plik czyta bajty Wzorca
    ponownie: Wzorzec zmieniony albo otwarty w Excelu po Sprawdzaniu działa tak, jakby był czytany dopiero
    wtedy, a kolejne pliki powstają z tych samych bajtów. Wzorzec na dysku pozostaje niezmieniony.
    """

    def __init__(self, path: Path, rules: _Rules, headers: Sequence[str]) -> None:
        self.path = path
        self._rules = rules
        self._headers = tuple(headers)
        self._content: bytes | None = None
        self._copy: TemplateCopy | None = None
        self._copy_read = False
        self._read_for_files = False

    @classmethod
    def worker_settlement(cls, path: Path) -> TemplateFile:
        """Wzorzec rozliczenia pracownika: bez danych w Wierszach danych, sprawdzany też przed każdym plikiem."""

        return cls(path, _WORKER_SETTLEMENT, ())

    @classmethod
    def collective_file(cls, path: Path, headers: Sequence[str]) -> TemplateFile:
        """Wzorzec pliku zbiorczego z nagłówkiem docelowym każdej pozycji Wiersza danych i Ilościami usług w N:AP."""

        return cls(path, _COLLECTIVE_FILE, headers)

    @classmethod
    def internal_settlement(cls, path: Path, headers: Sequence[str]) -> TemplateFile:
        """Wzorzec rozliczenia wewnętrznego z nagłówkiem docelowym każdej pozycji Wiersza danych w wierszu 1."""

        return cls(path, _INTERNAL_SETTLEMENT, headers)

    def check(self) -> None:
        """Sprawdzanie bez zapisu: ``WorkbookError``, gdy Wzorzec nie może posłużyć jako wzór.

        Powód to brak pliku, otwarcie w Excelu, nieczytelny skoroszyt, brak nagłówka z ``headers`` albo
        WYKONAWCA poza H17 (zły układ z ``MissingHeader``), a we Wzorcu rozliczenia pracownika także
        wartości w Wierszach danych (``HAS_DATA``).
        """

        if self._rules.sheet.period_file:
            outline = _read_outline(self.path, self._rules)
            _columns(self.path, outline.title, self._rules, outline.header, self._headers)
            if self._rules.worker_header:
                check_worker_header(self.path, outline.title, outline.header)
            return
        self._read()
        copy = self._template_copy()
        if copy is not None:
            self._check_settlement_template(copy.title, copy.header, copy.has_input_data)
            return
        with self._openpyxl_workbook() as workbook:
            sheet = _active_worksheet(self.path, workbook)
            self._check_settlement_template(sheet.title, self._header(sheet), self._has_input_data(sheet))

    def plan(self, row_count: int) -> TemplatePlan:
        """Planowanie ``row_count`` Wierszy danych: nagłówki z ``headers`` jak przy Sprawdzaniu, a scalenie w
        komórkach, do których trafią, to zły układ z ``MergedCells``."""

        outline = _read_outline(self.path, self._rules)
        columns = _columns(self.path, outline.title, self._rules, outline.header, self._headers)
        _ensure_unmerged(self.path, outline.title, outline.merged_ranges, self._rules.sheet, columns, row_count)
        first_row = self._rules.sheet.first_data_row
        return TemplatePlan(
            range(first_row, first_row + row_count), outline.zero_row_count, first_row + outline.zero_row_count
        )

    def create(self, target: Path, rows: Sequence[DataRow]) -> Route:
        """Zapisuje atomowo pod ``target`` plik z Wzorca z ``rows`` od pierwszego Wiersza danych i podaje drogę zapisu.

        Wpisana komórka dostaje Hiperłącze wiersza albo żadne i wartość; w kolumnach wyglądu Wygląd Wiersza
        danych (``row_look``), a dalej Wyróżnienia na czcionce i wypełnieniu komórki Wzorca. Formuły,
        ``calcChain`` kopii, reszta formatowania i ``xl/externalLinks/*`` zostają z Wzorca. ``WorkbookError``,
        gdy Wzorca nie da się już użyć, brak nagłówka albo komórki zapisu pliku okresu są scalone;
        ``WorkbookSaveError``, gdy zapis się nie uda, także przez znak niedozwolony w XML. W obu przypadkach ``target`` nie powstaje.
        """

        if not self._read_for_files:
            self._read()
            self._read_for_files = True
        copy = self._template_copy()
        if copy is not None:
            columns = _columns(self.path, copy.title, self._rules, copy.header, self._headers)
            if not self._rules.sheet.period_file:
                self._check_settlement_template(copy.title, copy.header, copy.has_input_data)
            placed = _placed(rows, columns)
            if placed is not None:
                try:
                    package = copy.with_rows(placed)
                except InvalidValue as exc:
                    raise WorkbookSaveError(f"Nie można zapisać skoroszytu: {target.name}") from exc
                except UnsupportedShape:
                    pass
                else:
                    write_atomically(target, _package_writer(package, target))
                    return Route.COPY
        self._create_with_openpyxl(target, rows)
        return Route.OPENPYXL

    def _read(self) -> None:
        """Czyta bajty Wzorca; wczytany Wzorzec zostaje, jeśli plik się nie zmienił."""

        ensure_available(self.path)
        if self.path.suffix.lower() not in _WORKBOOK_SUFFIXES:
            raise WorkbookError(self.path, Problem.UNREADABLE)
        content = workbook_content(self.path)
        if content != self._content:
            self._content = content
            self._copy = None
            self._copy_read = False

    def _template_copy(self) -> TemplateCopy | None:
        """Wzorzec do kopii z podmienionymi Wierszami danych; ``None``, gdy jego kształt wymaga openpyxl (ADR 0008)."""

        if not self._copy_read:
            self._copy_read = True
            self._copy = None
            if self._rules.copy:
                with suppress(UnsupportedShape):
                    self._copy = TemplateCopy.of(self._loaded(), self._rules.sheet)
        return self._copy

    def _openpyxl_workbook(self) -> AbstractContextManager[Workbook]:
        """Pełne wczytanie Wzorca przez openpyxl z jego klasyfikacją uszkodzeń: jedyna droga Rozliczenia wewnętrznego,
        a dla pozostałych Wzorców droga, gdy kopia nie przyjmuje ich kształtu (ADR 0008)."""

        return open_workbook(self.path, self._loaded(), read_only=False)

    def _create_with_openpyxl(self, target: Path, rows: Sequence[DataRow]) -> None:
        with self._openpyxl_workbook() as workbook:
            sheet = _active_worksheet(self.path, workbook)
            header = self._header(sheet)
            columns = _columns(self.path, sheet.title, self._rules, header, self._headers)
            if self._rules.sheet.period_file:
                _ensure_unmerged(
                    self.path, sheet.title, tuple(sheet.merged_cells.ranges), self._rules.sheet, columns, len(rows)
                )
            else:
                self._check_settlement_template(sheet.title, header, self._has_input_data(sheet))
            try:
                _write_rows(sheet, self._rules.sheet, rows, columns)
            except IllegalCharacterError as exc:
                raise WorkbookSaveError(f"Nie można zapisać skoroszytu: {target.name}") from exc
            if self._rules.copy:
                _request_full_recalculation(workbook)
            save_from_template(workbook, target, self._loaded(), self.path.name)

    def _header(self, sheet: Worksheet) -> list[object]:
        """Wartości wiersza nagłówka od kolumny A."""

        last_column = self._rules.last_header_column or sheet.max_column
        return [sheet.cell(self._rules.sheet.header_row, column).value for column in range(1, last_column + 1)]

    def _has_input_data(self, sheet: Worksheet) -> bool:
        """Wiersze danych mają wartość w A:K; formuły Wzorca nie są danymi."""

        for row in range(self._rules.sheet.first_data_row, sheet.max_row + 1):
            if any(_is_value(sheet.cell(row, column)) for column in INPUT_COLUMNS):
                return True
        return False

    def _check_settlement_template(self, sheet: str, header: Sequence[object], has_input_data: bool) -> None:
        """Kontrole Szablonu pracownika na Wzorcu: WYKONAWCA w H17, potem brak wartości w Wierszach danych."""

        check_worker_header(self.path, sheet, header)
        if has_input_data:
            raise WorkbookError(self.path, Problem.HAS_DATA)

    def _loaded(self) -> bytes:
        assert self._content is not None, "Wzorzec jest czytany przed użyciem"
        return self._content


def _columns(path: Path, sheet: str, rules: _Rules, header: Sequence[object], headers: Sequence[str]) -> Sequence[int]:
    """Kolumna Wzorca każdej pozycji Wiersza danych: pod nagłówkiem z ``headers``, a bez nich kolumny A:AT."""

    if not headers:
        return DATA_ROW_COLUMNS
    found = header_columns(header)
    columns: list[int] = []
    for target in headers:
        column = found.get(normalize_text(target))
        if column is None:
            raise WorkbookError(path, Problem.LAYOUT, MissingHeader(sheet, rules.sheet.header_row, target))
        columns.append(column)
    return columns


def _ensure_unmerged(
    path: Path, sheet: str, merged_ranges: Sequence[CellRange], layout: SheetLayout, columns: Sequence[int], row_count: int
) -> None:
    """Zły układ z ``MergedCells``, gdy komórki ``row_count`` zapisanych Wierszy danych są scalone."""

    if row_count == 0:
        return
    first_row, last_row = layout.first_data_row, layout.first_data_row + row_count - 1
    last_column = max(_last_written_column(layout, columns), layout.quantity_columns.stop - 1)
    for merged in merged_ranges:
        if merged.min_row <= last_row and merged.max_row >= first_row and merged.min_col <= last_column:
            raise WorkbookError(path, Problem.LAYOUT, MergedCells(sheet, merged.coord))


def _last_written_column(layout: SheetLayout, columns: Sequence[int]) -> int:
    return max(layout.look_columns, *columns)


def _placed(rows: Sequence[DataRow], columns: Sequence[int]) -> Sequence[DataRow] | None:
    """Wiersze z pozycjami w kolumnach od A, jak je wpisuje kopia; ``None``, gdy ``columns`` nie zajmują kolumn od A."""

    in_order = list(range(1, len(columns) + 1))
    if list(columns) == in_order:
        return rows
    if sorted(columns) != in_order:
        return None
    order = sorted(range(len(columns)), key=columns.__getitem__)
    return [
        DataRow(
            row.number,
            tuple(row.values[index] for index in order),
            tuple(row.highlights[index] for index in order),
            tuple(row.number_formats[index] for index in order),
            {columns[position - 1]: link for position, link in row.hyperlinks.items()},
            row.quantities,
        )
        for row in rows
    ]


def _write_rows(sheet: Worksheet, layout: SheetLayout, rows: Sequence[DataRow], columns: Sequence[int]) -> None:
    """Jedyny zapis Wierszy danych przez openpyxl: pozycja ``i`` wiersza trafia do kolumny ``columns[i-1]``.

    Komórka kolumny wyglądu dostaje Wygląd Wiersza danych, a bez pozycji w wierszu sam wygląd i treść Wzorca; w
    pliku okresu wygląd ma każda zapisana kolumna. Dalsze komórki dostają Wyróżnienia na stylu Wzorca. Zapisany
    wiersz traci wysokość z Wzorca.
    """

    last_column = _last_written_column(layout, columns)
    look_columns = last_column if layout.period_file else layout.look_columns
    look = RowLook([cast(Font, sheet.cell(layout.header_row, column).font) for column in range(1, look_columns + 1)])
    positions = {column: position for position, column in enumerate(columns, start=1)}
    for index, row in enumerate(rows):
        number = layout.first_data_row + index
        dimension = sheet.row_dimensions.get(number)  # bez tworzenia wymiaru, którego nie ma we Wzorcu
        if dimension is not None:
            dimension.ht = None
        for column in range(1, last_column + 1):
            cell = sheet.cell(number, column)
            position = positions.get(column)
            if position is None:
                look.cell(column, PLAIN, GENERAL, wrapped=False).apply(cell)
                continue
            value, highlights = row.values[position - 1], row.highlights[position - 1]
            link = row.hyperlinks.get(position)
            # Przed wartością: openpyxl wpisuje adres jako wartość pustej komórki.
            cell.hyperlink = None if link is None else link.address
            cell.value = value
            if column <= look.columns:
                look.cell(column, highlights, row.number_formats[position - 1], wrapped=wraps(value)).apply(cell)
            else:
                # font_over zawsze zwraca kopię, a fill_over może zwrócić wypełnienie, które dostał.
                cell.font = highlights.font_over(cast(Font, cell.font))
                cell.fill = highlights.fill_over(cast(Fill, copy(cell.fill)))
        for column, quantity in row.quantities.items():
            sheet.cell(number, column).value = quantity
    if layout.period_file:
        _set_autofilter(sheet, layout.header_row, len(rows))


def _set_autofilter(sheet: Worksheet, header_row: int, row_count: int) -> None:
    """Autofiltr na nagłówek i faktyczne Wiersze danych; zakres zapisany we Wzorcu jest ignorowany."""

    columns = [column for column, cell in enumerate(sheet[header_row], start=1) if not is_blank(cell.value)]
    first_column, last_column = get_column_letter(min(columns)), get_column_letter(max(columns))
    sheet.auto_filter = AutoFilter(ref=f"{first_column}{header_row}:{last_column}{header_row + row_count}")


def _package_writer(package: Package, target: Path) -> Callable[[Path], None]:
    def write(path: Path) -> None:
        try:
            with path.open("wb") as out:
                package.write(out)
        except OSError as exc:
            raise WorkbookSaveError(f"Nie można zapisać skoroszytu: {target.name}") from exc

    return write


def _active_worksheet(path: Path, workbook: Workbook) -> Worksheet:
    worksheet = workbook.active
    if not isinstance(worksheet, Worksheet):
        raise WorkbookError(path, Problem.UNREADABLE)
    return worksheet


def _request_full_recalculation(workbook: Workbook) -> None:
    """Excel przeliczy formuły przy otwarciu; openpyxl nie zapisuje wyników formuł."""

    calculation = getattr(workbook, "calculation", None)
    if calculation is not None:
        calculation.calcMode = "auto"
        calculation.fullCalcOnLoad = True
        calculation.forceFullCalc = True
