"""Deterministyczne rozdzielanie pliku zbiorczego do szablonów pracowników."""

from __future__ import annotations

import re
import shutil
import tempfile
import time
import unicodedata
from collections import defaultdict
from pathlib import Path
from typing import Callable, TypeVar

import yaml
from openpyxl import load_workbook

from . import template_settlement
from .domain import (
    Issue,
    ProgressEvent,
    ProgressEventFactory,
    ProgressObserver,
    ProgressPhase,
    SettlementSummary,
    WorkerResult,
)
from .template_settlement import (
    ExcelRow,
    PlaceholderValidationError,
    validate_placeholder,
)


ResultT = TypeVar("ResultT")

PERIOD_PATTERN = r"\d{2}_\d{2}_\d{2}_\d{4}"
SOURCE_PATTERN = re.compile(
    rf"^Rozliczenie (?P<period>{PERIOD_PATTERN}) - zbiorcze\.xlsx$"
)
HEADER_ROW = 17
DATA_START_ROW = 18
WORKER_COLUMN = 8  # H
_INVALID_FILENAME_CHARACTERS = frozenset('<>:"/\\|?*')
_MAX_WINDOWS_FILENAME_CODE_UNITS = 255


class SettlementError(Exception):
    """Błąd uniemożliwiający bezpieczne rozpoczęcie procesu."""


def _worker_filename_fits_windows(name: str) -> bool:
    filename = f"Rozliczenie 00_00_00_0000 - {name}.xlsx"
    try:
        filename_length = len(filename.encode("utf-16-le")) // 2
    except UnicodeEncodeError:
        return False
    return filename_length <= _MAX_WINDOWS_FILENAME_CODE_UNITS


def normalize_text(value: object) -> str:
    """Normalizuje identyfikator bez usuwania polskich znaków."""

    text = unicodedata.normalize("NFC", str(value or ""))
    return " ".join(text.replace("\u00a0", " ").split()).casefold()


def load_worker_mapping(config_path: Path) -> dict[str, str]:
    """Wczytuje jawne mapowanie identyfikatorów na nazwy plików."""

    try:
        config = yaml.safe_load(config_path.read_text(encoding="utf-8")) or {}
    except OSError as exc:
        raise SettlementError(f"Nie można odczytać konfiguracji: {config_path}") from exc
    except yaml.YAMLError as exc:
        raise SettlementError(f"Niepoprawny YAML konfiguracji: {config_path}") from exc

    if not isinstance(config, dict):
        raise SettlementError("Konfiguracja musi być mapą z sekcją workers.")
    if type(config.get("schema_version")) is not int or config["schema_version"] != 1:
        raise SettlementError("Nieobsługiwana wersja konfiguracji.")

    workers = config.get("workers")
    if not isinstance(workers, dict) or not workers:
        raise SettlementError("Konfiguracja workers musi zawierać co najmniej jeden wpis.")

    mapping: dict[str, str] = {}
    seen_identifiers: set[str] = set()
    seen_names: dict[str, str] = {}
    for source_id, target_name in workers.items():
        if not isinstance(source_id, str) or not isinstance(target_name, str):
            raise SettlementError("Konfiguracja zawiera niepoprawny identyfikator lub nazwę pracownika.")
        normalized_id = normalize_text(source_id)
        display_name = target_name.strip()
        normalized_name = normalize_text(display_name)
        if not normalized_id or not normalized_name:
            raise SettlementError("Konfiguracja zawiera pusty identyfikator lub nazwę pracownika.")
        if normalized_id in seen_identifiers:
            raise SettlementError("Konfiguracja zawiera powtórzony identyfikator pracownika.")
        if (
            any(character in _INVALID_FILENAME_CHARACTERS for character in display_name)
            or any(ord(character) < 32 for character in display_name)
            or display_name.endswith((".", " "))
        ):
            raise SettlementError("Konfiguracja zawiera nazwę pracownika nieprawidłową dla nazwy pliku.")
        if not _worker_filename_fits_windows(display_name):
            raise SettlementError("Konfiguracja zawiera nazwę pracownika zbyt długą dla nazwy pliku.")
        if normalized_name in seen_names:
            raise SettlementError(
                "Dwóch identyfikatorów wskazuje ten sam plik pracownika: "
                f"{seen_names[normalized_name]} / {display_name}"
            )
        mapping[normalized_id] = display_name
        seen_identifiers.add(normalized_id)
        seen_names[normalized_name] = display_name
    return mapping


def period_from_source(source_path: Path) -> str:
    """Waliduje nazwę i lokalizację pliku zbiorczego."""

    if "Rozliczenie wew." in source_path.name:
        raise SettlementError("Plik Rozliczenie wew. jest poza zakresem tego procesu.")
    match = SOURCE_PATTERN.fullmatch(source_path.name)
    if not match:
        raise SettlementError(
            "Wybierz plik nazwany dokładnie: "
            "Rozliczenie {dd}_{dd}_{mm}_{rrrr} - zbiorcze.xlsx."
        )
    period = match.group("period")
    if source_path.parent.name != period:
        raise SettlementError(
            f"Folder pliku musi nazywać się tak samo jak okres: {period}."
        )
    return period


def default_placeholder_path() -> Path:
    """Wskazuje wspólny Placeholder wersjonowany w katalogu konfiguracji."""

    return Path(__file__).resolve().parents[2] / "config" / "placeholder.xlsx"


def excel_lock_path(path: Path) -> Path:
    """Zwraca standardowy plik blokady tworzony przez desktopowy Excel."""

    return path.with_name(f"~${path.name}")


def ensure_unlocked(path: Path) -> None:
    if excel_lock_path(path).exists():
        raise SettlementError(f"Plik jest otwarty lub zablokowany: {path.name}")


def _read_source_rows(source_path: Path) -> tuple[dict[str, list[ExcelRow]], list[Issue]]:
    """Czyta wyłącznie dane wejściowe i grupuje je po kolumnie H."""

    workbook = load_workbook(source_path, read_only=True, data_only=False)
    try:
        sheet = workbook.active
        if sheet is None:
            raise SettlementError("Skoroszyt źródłowy nie zawiera arkusza.")
        header = next(
            sheet.iter_rows(min_row=HEADER_ROW, max_row=HEADER_ROW, max_col=WORKER_COLUMN, values_only=True),
            (),
        )
        if len(header) < WORKER_COLUMN or normalize_text(header[WORKER_COLUMN - 1]) != "wykonawca":
            raise SettlementError("W wierszu 17 nie znaleziono nagłówka WYKONAWCA w kolumnie H.")

        rows_by_worker: defaultdict[str, list[ExcelRow]] = defaultdict(list)
        issues: list[Issue] = []
        for row_number, values in enumerate(
            sheet.iter_rows(min_row=DATA_START_ROW, max_col=max(46, sheet.max_column), values_only=True),
            start=DATA_START_ROW,
        ):
            input_values = values[:11]
            if not any(value is not None for value in input_values):
                continue
            worker_value = values[WORKER_COLUMN - 1] if len(values) >= WORKER_COLUMN else None
            worker_id = normalize_text(worker_value)
            if not worker_id:
                issues.append(
                    Issue("BRAK_WYKONAWCY", f"Wiersz {row_number} ma dane, ale nie ma wykonawcy w kolumnie H.")
                )
                continue
            rows_by_worker[worker_id].append(tuple(values[:46]))
        return dict(rows_by_worker), issues
    finally:
        workbook.close()


def _elapsed_ms(started: float) -> int:
    return max(0, round((time.perf_counter() - started) * 1000))


class _ObserverDispatcher:
    """Detaches a broken observer without affecting the settlement engine."""

    def __init__(self, observer: ProgressObserver | None):
        self._observer = observer

    def notify(self, event: ProgressEvent) -> None:
        if self._observer is None:
            return
        try:
            self._observer(event)
        except Exception:
            self._observer = None


def _run_phase(
    phase: ProgressPhase,
    notify: Callable[[ProgressEvent], None],
    run_started: float,
    phase_durations_ms: dict[ProgressPhase, int],
    operation: Callable[[], ResultT],
) -> ResultT:
    phase_started = time.perf_counter()
    notify(ProgressEventFactory.phase_started(phase, elapsed_ms=_elapsed_ms(run_started)))
    try:
        result = operation()
    except Exception:
        phase_elapsed_ms = _elapsed_ms(phase_started)
        phase_durations_ms[phase] = phase_elapsed_ms
        notify(
            ProgressEventFactory.phase_failed(
                phase,
                elapsed_ms=_elapsed_ms(run_started),
                phase_elapsed_ms=phase_elapsed_ms,
            )
        )
        raise
    phase_elapsed_ms = _elapsed_ms(phase_started)
    phase_durations_ms[phase] = phase_elapsed_ms
    notify(
        ProgressEventFactory.phase_ended(
            phase,
            elapsed_ms=_elapsed_ms(run_started),
            phase_elapsed_ms=phase_elapsed_ms,
        )
    )
    return result


def run_settlements(
    source_path: Path,
    config_path: Path,
    *,
    dry_run: bool = False,
    placeholder_path: Path | None = None,
    observer: ProgressObserver | None = None,
) -> SettlementSummary:
    """Uruchamia pełny proces z ochroną przed błędną lokalizacją i duplikacją."""

    run_started = time.perf_counter()
    source_path = source_path.expanduser().resolve()
    config_path = config_path.expanduser().resolve()
    placeholder_path = (placeholder_path or default_placeholder_path()).expanduser().resolve()
    dispatcher = _ObserverDispatcher(observer)
    notify = dispatcher.notify
    phase_durations_ms: dict[ProgressPhase, int] = {}

    def check_inputs() -> tuple[str, Path, dict[str, str]]:
        if not source_path.is_file():
            raise SettlementError(f"Nie znaleziono pliku źródłowego: {source_path}")
        ensure_unlocked(source_path)
        period = period_from_source(source_path)
        target_directory = source_path.parent / f"Rozliczenie pracowników {period}"
        if target_directory.exists() or target_directory.is_symlink():
            raise SettlementError(f"Folder docelowy już istnieje: {target_directory.name}.")
        mapping = load_worker_mapping(config_path)
        try:
            validate_placeholder(placeholder_path)
        except PlaceholderValidationError as exc:
            raise SettlementError(str(exc)) from exc
        return period, target_directory, mapping

    period, target_directory, mapping = _run_phase(
        ProgressPhase.CHECKING,
        notify,
        run_started,
        phase_durations_ms,
        check_inputs,
    )
    rows_by_worker, source_issues = _run_phase(
        ProgressPhase.READING,
        notify,
        run_started,
        phase_durations_ms,
        lambda: _read_source_rows(source_path),
    )
    notify(
        ProgressEventFactory.issue_count(
            elapsed_ms=_elapsed_ms(run_started),
            issue_count=len(source_issues),
        )
    )
    summary = SettlementSummary(
        period,
        source_path,
        target_directory,
        issues=source_issues,
        phase_durations_ms=phase_durations_ms,
    )

    def plan_rows() -> tuple[list[tuple[str, Path]], dict[str, list[ExcelRow]], list[Issue]]:
        template_files = [
            (name, target_directory / f"Rozliczenie {period} - {name}.xlsx")
            for name in mapping.values()
        ]
        rows_by_target: defaultdict[str, list[ExcelRow]] = defaultdict(list)
        issues: list[Issue] = []

        for source_worker, rows in rows_by_worker.items():
            target_name = mapping.get(source_worker)
            if target_name is None:
                issues.append(
                    Issue(
                        "BRAK_MAPOWANIA",
                        "Nieznany identyfikator WYKONAWCA; sprawdź konfigurację pracowników.",
                    )
                )
                continue
            rows_by_target[normalize_text(target_name)].extend(rows)
        if not dry_run and (source_issues or issues):
            if any(issue.code == "BRAK_MAPOWANIA" for issue in issues):
                raise SettlementError(
                    "WYKONAWCA bez mapowania; folder wynikowy nie został opublikowany."
                )
            raise SettlementError(
                "Nie wszystkie wiersze danych mają przypisanego WYKONAWCĘ; "
                "folder wynikowy nie został opublikowany."
            )
        return template_files, rows_by_target, issues

    template_files, rows_by_target, plan_issues = _run_phase(
        ProgressPhase.PLANNING,
        notify,
        run_started,
        phase_durations_ms,
        plan_rows,
    )
    summary.issues.extend(plan_issues)
    notify(
        ProgressEventFactory.plan_ready(
            template_total=len(template_files),
            elapsed_ms=_elapsed_ms(run_started),
            issue_count=len(summary.issues),
        )
    )

    worker_failure_notified = False

    def save_targets(files: list[tuple[str, Path]]) -> None:
        nonlocal worker_failure_notified
        for template_index, (worker_name, target_path) in enumerate(files, start=1):
            target_key = normalize_text(worker_name)
            worker_started = time.perf_counter()
            notify(
                ProgressEventFactory.worker_started(
                    worker_name,
                    template_index=template_index,
                    template_total=len(template_files),
                    elapsed_ms=_elapsed_ms(run_started),
                )
            )
            rows = rows_by_target.get(target_key, [])
            try:
                if dry_run:
                    result = WorkerResult(
                        worker_name,
                        target_path,
                        len(rows),
                        "PLAN" if rows else "PUSTY_SZABLON",
                    )
                else:
                    result = template_settlement.process_template(
                        worker_name,
                        target_path,
                        rows,
                        dry_run=False,
                    )
                    if result.status not in {"ZAPISANO", "PUSTY_SZABLON"}:
                        summary.issues.extend(result.issues)
                        raise SettlementError(
                            f"Nie udało się przygotować skoroszytu pracownika {worker_name}; "
                            "folder wynikowy nie został opublikowany."
                        )
            except Exception as exc:
                worker_failure_notified = True
                notify(
                    ProgressEventFactory.worker_failed(
                        worker_name,
                        template_index=template_index,
                        template_total=len(template_files),
                        elapsed_ms=_elapsed_ms(run_started),
                        phase_elapsed_ms=_elapsed_ms(save_started),
                        worker_elapsed_ms=_elapsed_ms(worker_started),
                    )
                )
                if isinstance(exc, SettlementError):
                    raise
                raise SettlementError(
                    f"Nie udało się przygotować skoroszytu pracownika {worker_name}; "
                    "folder wynikowy nie został opublikowany."
                ) from exc
            summary.issues.extend(result.issues)
            summary.results.append(result)
            notify(
                ProgressEventFactory.worker_ended(
                    worker_name,
                    template_index=template_index,
                    template_total=len(template_files),
                    status=result.status,
                    rows=result.rows,
                    elapsed_ms=_elapsed_ms(run_started),
                    worker_elapsed_ms=_elapsed_ms(worker_started),
                    issue_count=len(summary.issues),
                )
            )
    save_started = time.perf_counter()
    notify(ProgressEventFactory.phase_started("Zapisywanie", elapsed_ms=_elapsed_ms(run_started)))
    try:
        if dry_run:
            save_targets(template_files)
        else:
            try:
                with tempfile.TemporaryDirectory(
                    prefix=f".{target_directory.name}.staging-",
                    dir=source_path.parent,
                ) as staging_name:
                    staging_directory = Path(staging_name)
                    staged_files = [
                        (worker_name, staging_directory / target_path.name)
                        for worker_name, target_path in template_files
                    ]
                    try:
                        for _, staged_path in staged_files:
                            shutil.copy2(placeholder_path, staged_path)
                    except OSError as exc:
                        raise SettlementError(
                            "Nie można utworzyć skoroszytów z Placeholdera; "
                            "folder wynikowy nie został opublikowany."
                        ) from exc

                    save_targets(staged_files)
                    if target_directory.exists() or target_directory.is_symlink():
                        raise SettlementError(f"Folder docelowy już istnieje: {target_directory.name}.")
                    try:
                        staging_directory.rename(target_directory)
                    except OSError as exc:
                        raise SettlementError(
                            "Nie można opublikować gotowego folderu rozliczeń pracowników."
                        ) from exc
                    summary.results = [
                        WorkerResult(
                            result.worker_name,
                            target_directory / result.output_file.name,
                            result.rows,
                            result.status,
                            result.issues,
                        )
                        for result in summary.results
                    ]
            except OSError as exc:
                raise SettlementError(
                    "Nie można przygotować folderu rozliczeń pracowników; "
                    "folder wynikowy nie został opublikowany."
                ) from exc
    except Exception:
        phase_elapsed_ms = _elapsed_ms(save_started)
        phase_durations_ms[ProgressPhase.SAVING] = phase_elapsed_ms
        if not worker_failure_notified:
            notify(
                ProgressEventFactory.phase_failed(
                    "Zapisywanie",
                    elapsed_ms=_elapsed_ms(run_started),
                    phase_elapsed_ms=phase_elapsed_ms,
                )
            )
        raise
    else:
        phase_elapsed_ms = _elapsed_ms(save_started)
        phase_durations_ms[ProgressPhase.SAVING] = phase_elapsed_ms
        notify(
            ProgressEventFactory.phase_ended(
                "Zapisywanie",
                template_index=len(template_files),
                template_total=len(template_files),
                elapsed_ms=_elapsed_ms(run_started),
                phase_elapsed_ms=phase_elapsed_ms,
            )
        )

    return summary
