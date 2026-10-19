"""Przetwarzanie pojedynczego Szablonu pracownika."""

from __future__ import annotations

import zipfile
from pathlib import Path
from typing import Iterable, Literal

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from . import collective_file, workbook_files
from .collective_file import CollectiveFileError, ExcelRow
from .domain import Issue, WorkerResult
from .workbook_files import WORKER_HEADER, Layout, Problem, WorkbookError


class TemplateWriteError(Exception):
    """Błąd właściwego zapisu Szablonu pracownika."""


# Nazwa pliku sprzed zmiany nazwy Placeholdera; służy tylko podpowiedzi w komunikacie.
_LEGACY_WORKER_TEMPLATE_NAME = "placeholder.xlsx"


WorkerTemplateProblem = Literal["missing", "open_in_excel", "unreadable"]


class WorkerTemplateValidationError(Exception):
    """Wzorzec rozliczenia pracownika nie może bezpiecznie posłużyć jako wzór nowych skoroszytów.

    ``message`` to bezpieczny komunikat dla taty; ``problem`` mówi, czy Wzorca
    nie ma, czy jest otwarty w Excelu, czy nie przechodzi Sprawdzania.
    """

    def __init__(self, message: str, problem: WorkerTemplateProblem = "unreadable") -> None:
        super().__init__(message)
        self.message = message
        self.problem: WorkerTemplateProblem = problem


def _is_real_value(cell: Cell | MergedCell) -> bool:
    """Pomija formuły szablonu przy wykrywaniu istniejących danych."""

    return cell.value is not None and cell.data_type != "f"


def _target_has_input_data(sheet: Worksheet) -> bool:
    """Sprawdza dane wejściowe A:K, ignorując formuły w szablonie."""

    for row in range(collective_file.FIRST_DATA_ROW, sheet.max_row + 1):
        if any(_is_real_value(sheet.cell(row=row, column=column)) for column in collective_file.INPUT_COLUMNS):
            return True
    return False


def validate_worker_template(path: Path) -> None:
    """Sprawdza Wzorzec rozliczenia pracownika bez modyfikowania jego zawartości."""

    try:
        with workbook_files.open_template(path, Layout.COLLECTIVE) as sheet:
            sheet.check_worker_header()
            if sheet.has_input_data():
                raise WorkerTemplateValidationError(
                    "Wzorzec rozliczenia pracownika zawiera dane i nie może być użyty jako wzór."
                )
    except WorkbookError as error:
        raise _validation_error(path, error) from error


def _validation_error(path: Path, error: WorkbookError) -> WorkerTemplateValidationError:
    if error.problem is Problem.MISSING:
        message = f"Nie znaleziono Wzorca rozliczenia pracownika: {path.name}."
        if path.with_name(_LEGACY_WORKER_TEMPLATE_NAME).is_file():
            message += f" Zmień nazwę pliku {_LEGACY_WORKER_TEMPLATE_NAME} w tym samym folderze na „{path.name}”."
        return WorkerTemplateValidationError(message, "missing")
    if error.problem is Problem.OPEN_IN_EXCEL:
        return WorkerTemplateValidationError("Wzorzec rozliczenia pracownika jest otwarty lub zablokowany.", "open_in_excel")
    if error.problem is Problem.LAYOUT:
        return WorkerTemplateValidationError(
            f"Wzorzec rozliczenia pracownika nie ma nagłówka WYKONAWCA w {WORKER_HEADER.column}:{WORKER_HEADER.row}."
        )
    return WorkerTemplateValidationError("Nie można odczytać Wzorca rozliczenia pracownika.")


def _active_worksheet(workbook: Workbook) -> Worksheet | None:
    sheet = workbook.active
    return sheet if isinstance(sheet, Worksheet) else None


def _write_workbook(path: Path, workbook: Workbook, rows: list[ExcelRow]) -> None:
    sheet = _active_worksheet(workbook)
    if sheet is None:
        raise ValueError("Szablon nie zawiera arkusza roboczego.")
    for destination_row, values in enumerate(rows, start=collective_file.FIRST_DATA_ROW):
        for column, value in zip(collective_file.COPIED_COLUMNS, values):
            sheet.cell(row=destination_row, column=column).value = value

    workbook_files.request_full_recalculation(workbook)
    workbook_files.save_workbook(workbook, path)


def process_template(
    worker_name: str,
    target_path: Path,
    rows: Iterable[ExcelRow],
) -> WorkerResult:
    """Przetwarza jeden Szablon pracownika i zwraca jego bezpieczny wynik."""

    rows = list(rows)
    if workbook_files.is_open_in_excel(target_path):
        issue = Issue("PLIK_ZABLOKOWANY", f"Plik jest otwarty lub zablokowany: {target_path.name}")
        return WorkerResult(worker_name, target_path, 0, "ZABLOKOWANY", (issue,))
    try:
        workbook = load_workbook(target_path, read_only=False, data_only=False, keep_links=True)
        try:
            sheet = _active_worksheet(workbook)
            if sheet is None:
                raise ValueError("Szablon nie zawiera arkusza.")
            try:
                collective_file.check_worker_header(sheet)
            except CollectiveFileError:
                issue = Issue("ZLY_SZABLON", f"Szablon nie ma WYKONAWCA w H:17: {target_path.name}.")
                return WorkerResult(worker_name, target_path, 0, "ZLY_SZABLON", (issue,))
            if _target_has_input_data(sheet):
                issue = Issue("NADPISANIE_ZABLOKOWANE", f"Szablon zawiera już dane: {target_path.name}.")
                return WorkerResult(worker_name, target_path, 0, "ZABLOKOWANY", (issue,))
            if not rows:
                return WorkerResult(worker_name, target_path, 0, "PUSTY_SZABLON")

            try:
                _write_workbook(target_path, workbook, rows)
            except (
                workbook_files.WorkbookSaveError,
                OSError,
                InvalidFileException,
                ValueError,
                zipfile.BadZipFile,
            ) as exc:
                raise TemplateWriteError(f"Nie można zapisać szablonu: {target_path.name}") from exc
            return WorkerResult(worker_name, target_path, len(rows), "ZAPISANO")
        finally:
            workbook.close()
    except (OSError, InvalidFileException, ValueError, zipfile.BadZipFile):
        issue = Issue("BLAD_SZABLONU", f"Nie można obsłużyć szablonu: {target_path.name}.")
        return WorkerResult(worker_name, target_path, 0, "POMINIĘTO", (issue,))
