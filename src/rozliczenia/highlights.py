"""Wyróżnienia komórki: wypełnienie, kolor czcionki, pogrubienie i kursywa.

Przebieg rozliczeń przenosi je z każdej komórki A:AT Wiersza danych Pliku
zbiorczego do tej samej komórki Rozliczenia pracownika. Krój i rozmiar czcionki,
format liczby, obramowanie, wyrównanie i ochrona zostają z Wzorca rozliczenia
pracownika. Wyróżnienia zastępują wygląd komórki Wzorca w całości: komórka bez
Wyróżnień, także brakująca w Pliku zbiorczym, daje komórkę bez wypełnienia, z
kolorem automatycznym, bez pogrubienia i bez kursywy.

Typ odczytuje ``workbook_files.read_data_rows``; obie drogi zapisu Rozliczenia
pracownika (``template_copy`` i openpyxl w ``workbook_files``) nakładają go na
czcionkę i wypełnienie komórki Wzorca tymi samymi metodami, więc dają ten sam
wynik.
"""

from __future__ import annotations

from copy import copy
from dataclasses import dataclass, field

from openpyxl.styles.colors import Color
from openpyxl.styles.fills import Fill, PatternFill
from openpyxl.styles.fonts import Font


@dataclass(frozen=True)
class Highlights:
    """Wyróżnienia jednej komórki; porównywane po wartości."""

    fill: Fill | None = None
    """``None``: bez wypełnienia."""
    color: Color | None = None
    """``None``: kolor automatyczny."""
    bold: bool = False
    italic: bool = False
    # Zapis szuka stylu po Wyróżnieniach każdej wpisanej komórki, a openpyxl liczy hash
    # stylu od nowa przy każdym wywołaniu.
    _hash: int = field(init=False, repr=False, compare=False)

    def __post_init__(self) -> None:
        object.__setattr__(self, "_hash", hash((self.fill, self.color, self.bold, self.italic)))

    def __hash__(self) -> int:
        return self._hash

    @classmethod
    def of(cls, font: Font | None, fill: Fill | None) -> Highlights:
        """Wyróżnienia komórki o czcionce ``font`` i wypełnieniu ``fill``; brakująca komórka ma oba ``None``."""

        return cls(
            fill=None if fill is None or _is_empty(fill) else fill,
            color=None if font is None else font.color,
            bold=font is not None and bool(font.b),
            italic=font is not None and bool(font.i),
        )

    def font_over(self, font: Font) -> Font:
        """Czcionka komórki Wzorca z tymi Wyróżnieniami; atrybut, który już znaczy to samo, zostaje bez zmian."""

        result = copy(font)
        if result.color != self.color:
            result.color = copy(self.color)
        if bool(result.b) != self.bold:
            result.b = self.bold
        if bool(result.i) != self.italic:
            result.i = self.italic
        return result

    def fill_over(self, fill: Fill) -> Fill:
        """Wypełnienie komórki Wzorca zastąpione tym; równe albo puste po obu stronach zostaje bez zmian."""

        if self.fill is None:
            return fill if _is_empty(fill) else PatternFill()
        return fill if fill == self.fill else copy(self.fill)


PLAIN = Highlights()
"""Komórka bez Wyróżnień."""


def _is_empty(fill: Fill) -> bool:
    return isinstance(fill, PatternFill) and fill.patternType is None
