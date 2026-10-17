"""Przygotowanie pliku zbiorczego: z Rejestru zleceń do Rozliczenia wewnętrznego i Pliku zbiorczego.

Jedyną operacją publiczną jest ``prepare_collective_file()``. Silnik nie pyta o
nic i nie proponuje okresu; zatrzymania zgłasza jako ``PreparationError`` z
bezpiecznym komunikatem dla taty, a wynik nie zawiera danych klientów.
"""

from __future__ import annotations

import os
import re
import tempfile
import zipfile
from collections.abc import Sequence
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet._read_only import ReadOnlyWorksheet
from openpyxl.worksheet.worksheet import Worksheet

from . import preparation_messages as messages
from . import workbook_files
from .domain import ProgressEventFactory, ProgressObserver, ProgressPhase
from .engine import SettlementError, load_worker_mapping, normalize_text
from .phase_runner import PhaseRunner
from .preparation_messages import PreparationError, PreparationMessage

__all__ = [
    "PreparationError",
    "PreparationMessage",
    "PreparationSummary",
    "default_collective_template_path",
    "default_internal_template_path",
    "prepare_collective_file",
]


@dataclass(frozen=True)
class _Column:
    """Kolumna przenoszona z Rejestru zleceń; nagłówek docelowy jest szukany we wzorcach."""

    register_header: str
    target_header: str


# Kolejność tej krotki jest kolejnością wartości w ``_RegisterOrder.values``.
_COLUMNS: tuple[_Column, ...] = (
    _Column("DATA GODZ. POCZ.", "DATA"),
    _Column("MIASTO", "MIASTO"),
    _Column("adres", "ADRES"),
    _Column("UMOWA", "UMOWA"),
    _Column("URZĄDZENIA", "URZĄDZENIA"),
    _Column("CZYNNOŚCI", "CZYNNOŚCI"),
    _Column("nr zlecenia", "NUMER ZL."),
    _Column("WYKONAWCA", "WYKONAWCA"),
    _Column("NUMER UMOWY/ZLECENIA ORANGE", "NUMER UMOWY/ZLECENIA ORANGE"),
    _Column("SKUTECZNE", "SKUTECZNE"),
    _Column("CZYNNOŚCI ZREALIZOWANE", "CZYNNOŚCI ZREALIZOWANE"),
)
_DATE_INDEX = 0
_SUCCESSFUL_INDEX = 9
_SUCCESSFUL_VALUE = "tak"

_PERIOD_TOKEN = re.compile(r"^(\d{2})_(\d{2})_(\d{2})_(\d{4})$")
_INTERNAL_HEADER_ROW = 1
_COLLECTIVE_HEADER_ROW = 17
_COLLECTIVE_DATA_START_ROW = 18
_COLLECTIVE_COLUMNS = 11  # A:K
# Błędy openpyxl przy pliku, który nie jest poprawnym skoroszytem .xlsx.
_UNREADABLE_WORKBOOK_ERRORS = (OSError, InvalidFileException, zipfile.BadZipFile, KeyError, ValueError)


@dataclass
class PreparationSummary:
    """Podsumowanie Przygotowania pliku zbiorczego bez danych klientów."""

    period: str
    period_directory: Path
    internal_settlement_file: Path
    collective_file: Path
    period_order_count: int
    successful_order_count: int
    skipped_empty_rows: int
    warnings: list[PreparationMessage] = field(default_factory=list)
    phase_durations_ms: dict[ProgressPhase, int] = field(default_factory=dict)


@dataclass(frozen=True)
class _Period:
    token: str
    start: datetime  # wtorek 00:00
    end: datetime  # następny wtorek 00:00, wyłącznie


@dataclass(frozen=True)
class _RegisterOrder:
    row_number: int
    values: tuple[object, ...]  # w kolejności ``_COLUMNS``

    @property
    def is_successful(self) -> bool:
        return normalize_text(self.values[_SUCCESSFUL_INDEX]) == _SUCCESSFUL_VALUE


@dataclass(frozen=True)
class _RegisterSheet:
    name: str
    rows: list[tuple[int, tuple[object, ...] | None]]  # ``None``: całkowicie pusty wiersz


@dataclass(frozen=True)
class _Plan:
    orders: list[_RegisterOrder]
    skipped_empty_rows: int

    @property
    def successful_orders(self) -> list[_RegisterOrder]:
        return [order for order in self.orders if order.is_successful]


def default_internal_template_path() -> Path:
    """Wzorzec rozliczenia wewnętrznego w katalogu konfiguracji obok Placeholdera."""

    return Path(__file__).resolve().parents[2] / "config" / "Wzorzec rozliczenia wewnętrznego.xlsx"


def default_collective_template_path() -> Path:
    """Wzorzec pliku zbiorczego w katalogu konfiguracji obok Placeholdera."""

    return Path(__file__).resolve().parents[2] / "config" / "Wzorzec pliku zbiorczego.xlsx"


def _internal_settlement_filename(period: str) -> str:
    return f"Rozliczenie wew. {period}.xlsx"


def _collective_filename(period: str) -> str:
    return f"Rozliczenie {period} - zbiorcze.xlsx"


def _parse_period(token: str) -> _Period:
    """Token ``dd_dd_mm_rrrr``: wtorek, poniedziałek, miesiąc i rok jednego z tych dni."""

    match = _PERIOD_TOKEN.fullmatch(token)
    if match is None:
        raise messages.invalid_period_token(token)
    tuesday_day, monday_day, month, year = (int(part) for part in match.groups())
    candidates: list[date] = []
    try:
        candidates.append(date(year, month, tuesday_day))
    except ValueError:
        pass
    try:
        candidates.append(date(year, month, monday_day) - timedelta(days=6))
    except ValueError:
        pass
    for tuesday in candidates:
        monday = tuesday + timedelta(days=6)
        if tuesday.weekday() == 1 and tuesday.day == tuesday_day and monday.day == monday_day:
            start = datetime(tuesday.year, tuesday.month, tuesday.day)
            return _Period(token, start, start + timedelta(days=7))
    raise messages.invalid_period_token(token)


def _is_blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _header_columns(values: Sequence[object]) -> dict[str, int]:
    """Mapuje znormalizowane nagłówki na numery kolumn (od 1)."""

    return {
        normalize_text(value): column
        for column, value in enumerate(values, start=1)
        if not _is_blank(value)
    }


def _missing_register_headers(header: Sequence[object]) -> list[str]:
    present = set(_header_columns(header))
    return [column.register_header for column in _COLUMNS if normalize_text(column.register_header) not in present]


def _duplicated_register_headers(header: Sequence[object]) -> list[str]:
    normalized = [normalize_text(value) for value in header if not _is_blank(value)]
    return [
        column.register_header
        for column in _COLUMNS
        if normalized.count(normalize_text(column.register_header)) > 1
    ]


def _read_register(register_path: Path) -> list[_RegisterSheet]:
    try:
        return _read_register_sheets(register_path)
    except _UNREADABLE_WORKBOOK_ERRORS as exc:
        raise messages.register_unreadable(register_path) from exc


def _read_register_sheets(register_path: Path) -> list[_RegisterSheet]:
    """Czyta arkusze Rejestru rozpoznane po komplecie wymaganych nagłówków w pierwszym wierszu.

    Arkusze bez większości wymaganych nagłówków (listy, komentarze, ukryte) są pomijane.
    Arkusz z większością, ale nie kompletem nagłówków, wskazuje brakujący nagłówek.
    """

    workbook = load_workbook(register_path, read_only=True, data_only=False)
    sheets: list[_RegisterSheet] = []
    closest: tuple[str, list[str]] | None = None  # arkusz z najmniejszą liczbą braków
    try:
        for sheet in workbook.worksheets:
            header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
            missing = _missing_register_headers(header)
            if missing:
                if len(missing) < len(_COLUMNS) / 2 and (closest is None or len(missing) < len(closest[1])):
                    closest = (sheet.title, missing)
                continue
            duplicated = _duplicated_register_headers(header)
            if duplicated:
                raise messages.register_with_duplicated_headers(register_path, sheet.title, duplicated)
            columns = _header_columns(header)
            positions = [columns[normalize_text(column.register_header)] for column in _COLUMNS]
            rows: list[tuple[int, tuple[object, ...] | None]] = []
            for row_number, cells in enumerate(sheet.iter_rows(min_row=2), start=2):
                values = [getattr(cell, "value", None) for cell in cells]
                if all(_is_blank(value) for value in values):
                    rows.append((row_number, None))
                    continue
                rows.append(
                    (
                        row_number,
                        tuple(values[position - 1] if position <= len(values) else None for position in positions),
                    )
                )
            sheets.append(_RegisterSheet(sheet.title, rows))
    finally:
        workbook.close()
    if sheets:
        return sheets
    if closest is not None:
        raise messages.register_without_headers(register_path, *closest)
    raise messages.register_without_data_sheet(register_path, [column.register_header for column in _COLUMNS])


def _order_date(value: object) -> datetime | None:
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    return None


def _plan(register: _RegisterSheet, period: _Period) -> _Plan:
    """Wybiera zlecenia okresu po dacie, w kolejności Rejestru.

    Pominięte puste wiersze to te, które leżą między dwoma kolejnymi zleceniami okresu.
    """

    orders: list[_RegisterOrder] = []
    skipped_empty_rows = 0
    pending_empty_rows = 0
    previous_in_period = False
    for row_number, values in register.rows:
        if values is None:
            pending_empty_rows += 1
            continue
        started = _order_date(values[_DATE_INDEX])
        in_period = started is not None and period.start <= started < period.end
        if in_period:
            if previous_in_period:
                skipped_empty_rows += pending_empty_rows
            orders.append(_RegisterOrder(row_number, values))
        previous_in_period = in_period
        pending_empty_rows = 0
    return _Plan(orders, skipped_empty_rows)


def _plan_register(register_path: Path, sheets: Sequence[_RegisterSheet], period: _Period) -> _Plan:
    """Planuje okres z jedynego arkusza Rejestru, który ma zlecenia tego okresu."""

    plans = [(sheet.name, _plan(sheet, period)) for sheet in sheets]
    with_orders = [(name, plan) for name, plan in plans if plan.orders]
    if len(with_orders) > 1:
        raise messages.period_orders_in_many_sheets(register_path, [name for name, _ in with_orders], period.token)
    return (with_orders or plans)[0][1]


def _worksheet(workbook: Workbook) -> Worksheet:
    sheet = workbook.active
    if not isinstance(sheet, Worksheet):
        raise ValueError("Wzorzec nie zawiera arkusza roboczego.")
    return sheet


def _target_columns(
    template_path: Path,
    sheet_title: str,
    header_row: int,
    header_values: Sequence[object],
) -> list[int]:
    """Kolumny wzorca dla ``_COLUMNS``, rozpoznane po nagłówku docelowym albo nagłówku Rejestru."""

    columns = _header_columns(header_values)
    positions: list[int] = []
    for column in _COLUMNS:
        position = columns.get(normalize_text(column.target_header)) or columns.get(
            normalize_text(column.register_header)
        )
        if position is None:
            raise messages.template_without_header(template_path, sheet_title, header_row, column.target_header)
        positions.append(position)
    return positions


def _write_orders(sheet: Worksheet, positions: list[int], orders: list[_RegisterOrder], first_row: int) -> None:
    for row, order in enumerate(orders, start=first_row):
        for position, value in zip(positions, order.values):
            sheet.cell(row, position).value = value  # type: ignore[assignment]


def _write_internal_settlement(template_path: Path, target: Path, orders: list[_RegisterOrder]) -> None:
    workbook = load_workbook(template_path, read_only=False, data_only=False, keep_links=True)
    try:
        sheet = _worksheet(workbook)
        header = [cell.value for cell in sheet[_INTERNAL_HEADER_ROW]]
        positions = _target_columns(template_path, sheet.title, _INTERNAL_HEADER_ROW, header)
        _write_orders(sheet, positions, orders, _INTERNAL_HEADER_ROW + 1)
        workbook_files.save_workbook(workbook, target, external_links_from=template_path)
    finally:
        workbook.close()


def _write_collective_file(template_path: Path, target: Path, orders: list[_RegisterOrder]) -> None:
    workbook = load_workbook(template_path, read_only=False, data_only=False, keep_links=True)
    try:
        sheet = _worksheet(workbook)
        header = [
            sheet.cell(_COLLECTIVE_HEADER_ROW, column).value
            for column in range(1, _COLLECTIVE_COLUMNS + 1)
        ]
        positions = _target_columns(template_path, sheet.title, _COLLECTIVE_HEADER_ROW, header)
        _write_orders(sheet, positions, orders, _COLLECTIVE_DATA_START_ROW)
        calculation = getattr(workbook, "calculation", None)
        if calculation is not None:
            calculation.calcMode = "auto"
            calculation.fullCalcOnLoad = True
            calculation.forceFullCalc = True
        workbook_files.save_workbook(workbook, target, external_links_from=template_path)
    finally:
        workbook.close()


def _ensure_register_available(register_path: Path) -> None:
    if not register_path.is_file():
        raise messages.register_missing(register_path)
    if workbook_files.is_open_in_excel(register_path):
        raise messages.register_open_in_excel(register_path)


def _ensure_template_available(template_path: Path) -> None:
    if not template_path.is_file():
        raise messages.template_missing(template_path)
    if workbook_files.is_open_in_excel(template_path):
        raise messages.template_open_in_excel(template_path)


def _check_template(template_path: Path, header_row: int, last_column: int | None) -> None:
    """Wzorzec istnieje, nie jest otwarty, daje się odczytać i ma oczekiwane nagłówki."""

    _ensure_template_available(template_path)
    try:
        workbook = load_workbook(template_path, read_only=True, data_only=False)
    except _UNREADABLE_WORKBOOK_ERRORS as exc:
        raise messages.template_unreadable(template_path) from exc
    try:
        sheet = workbook.active
        if not isinstance(sheet, ReadOnlyWorksheet):
            raise messages.template_unreadable(template_path)
        header = next(
            sheet.iter_rows(min_row=header_row, max_row=header_row, max_col=last_column, values_only=True),
            (),
        )
    except _UNREADABLE_WORKBOOK_ERRORS as exc:
        raise messages.template_unreadable(template_path) from exc
    finally:
        workbook.close()
    _target_columns(template_path, sheet.title, header_row, header)


def _ensure_period_files_absent(period_directory: Path, filenames: Sequence[str]) -> None:
    for filename in filenames:
        target = period_directory / filename
        if target.exists() or target.is_symlink():
            raise messages.period_file_exists(period_directory, filename)


def _publish(staged_directory: Path, period_directory: Path, filenames: Sequence[str]) -> None:
    """Przenosi oba pliki do Folderu okresu razem; przy błędzie nie zostawia żadnego."""

    _ensure_period_files_absent(period_directory, filenames)
    if not period_directory.exists():
        staged_directory.rename(period_directory)
        return
    published: list[Path] = []
    try:
        for filename in filenames:
            target = period_directory / filename
            os.rename(staged_directory / filename, target)
            published.append(target)
    except OSError:
        for target in published:
            target.unlink(missing_ok=True)
        raise


def prepare_collective_file(
    register_path: Path,
    periods_directory: Path,
    period: str,
    config_path: Path,
    *,
    internal_template_path: Path | None = None,
    collective_template_path: Path | None = None,
    observer: ProgressObserver | None = None,
) -> PreparationSummary:
    """Tworzy w Folderze okresu razem Rozliczenie wewnętrzne i Plik zbiorczy.

    Zgłasza ``PreparationError``, gdy nie można działać bezpiecznie; wtedy nie
    powstaje żaden plik ani nowy Folder okresu. Rejestr zleceń i wzorce są
    tylko czytane.
    """

    phases = PhaseRunner(observer)
    register_path = register_path.expanduser().resolve()
    periods_directory = periods_directory.expanduser().resolve()
    config_path = config_path.expanduser().resolve()
    internal_template_path = (internal_template_path or default_internal_template_path()).expanduser().resolve()
    collective_template_path = (
        (collective_template_path or default_collective_template_path()).expanduser().resolve()
    )
    period_directory = periods_directory / period
    filenames = (_internal_settlement_filename(period), _collective_filename(period))

    def check_inputs() -> tuple[_Period, dict[str, str]]:
        parsed_period = _parse_period(period)
        _ensure_register_available(register_path)
        _check_template(internal_template_path, _INTERNAL_HEADER_ROW, None)
        _check_template(collective_template_path, _COLLECTIVE_HEADER_ROW, _COLLECTIVE_COLUMNS)
        if not periods_directory.is_dir():
            raise messages.periods_directory_missing(periods_directory)
        _ensure_period_files_absent(period_directory, filenames)
        try:
            worker_mapping = load_worker_mapping(config_path)
        except SettlementError as exc:
            raise messages.unreadable_worker_config(config_path) from exc
        return parsed_period, worker_mapping

    parsed_period, _worker_mapping = phases.run(ProgressPhase.CHECKING, check_inputs)
    register = phases.run(ProgressPhase.READING, lambda: _read_register(register_path))
    phases.notify(ProgressEventFactory.issue_count(elapsed_ms=phases.elapsed_ms(), issue_count=0))
    plan = phases.run(ProgressPhase.PLANNING, lambda: _plan_register(register_path, register, parsed_period))
    summary = PreparationSummary(
        period=period,
        period_directory=period_directory,
        internal_settlement_file=period_directory / filenames[0],
        collective_file=period_directory / filenames[1],
        period_order_count=len(plan.orders),
        successful_order_count=len(plan.successful_orders),
        skipped_empty_rows=plan.skipped_empty_rows,
        phase_durations_ms=phases.phase_durations_ms,
    )
    phases.notify(
        ProgressEventFactory.plan_ready(
            template_total=0,
            elapsed_ms=phases.elapsed_ms(),
            issue_count=len(summary.warnings),
        )
    )

    def save_both() -> None:
        try:
            with tempfile.TemporaryDirectory(prefix=f".{period}.staging-", dir=periods_directory) as staging:
                staged_directory = Path(staging) / period
                staged_directory.mkdir()
                _write_internal_settlement(
                    internal_template_path, staged_directory / filenames[0], plan.orders
                )
                _write_collective_file(
                    collective_template_path, staged_directory / filenames[1], plan.successful_orders
                )
                _publish(staged_directory, period_directory, filenames)
        except (OSError, workbook_files.WorkbookSaveError) as exc:
            raise messages.write_failed(period_directory) from exc

    phases.run(ProgressPhase.SAVING, save_both)
    return summary
