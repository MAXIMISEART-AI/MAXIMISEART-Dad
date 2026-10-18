"""Plik zbiorczy: układ wspólny dla obu procesów i nazwa ``Rozliczenie {okres} - zbiorcze.xlsx``.

Przygotowanie pliku zbiorczego zapisuje Plik zbiorczy w tym układzie, a Przebieg
rozliczeń go czyta i kopiuje Wiersze danych po pozycji kolumn do Szablonów
pracowników. Dlatego ten sam układ mają Wzorzec pliku zbiorczego, Placeholder i
Szablony pracowników. Moduł nie zna mapowania Wykonawców. ``CollectiveFileError``
ma bezpieczny komunikat, a każdy proces tłumaczy go na swój błąd.
"""

from __future__ import annotations

import re
import zipfile
import zlib
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from pathlib import Path
from xml.etree.ElementTree import ParseError

from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula
from openpyxl.worksheet.worksheet import Worksheet

from .domain import Issue
from .settlement_period import PERIOD_TOKEN_PATTERN

HEADER_ROW = 17
FIRST_DATA_ROW = 18
WORKER_COLUMN = 8  # H
INPUT_COLUMNS = range(1, 12)  # A:K: nagłówki w wierszu 17 i dane wejściowe Wiersza danych
QUANTITY_COLUMNS = range(14, 43)  # N:AP
COPIED_COLUMNS = range(1, 47)  # A:AT: kopiowane z Pliku zbiorczego do Szablonu pracownika

_ExcelCellValue = (
    bool
    | int
    | float
    | Decimal
    | str
    | CellRichText
    | date
    | datetime
    | time
    | timedelta
    | DataTableFormula
    | ArrayFormula
    | bytes
    | None
)
ExcelRow = tuple[_ExcelCellValue, ...]

# Błędy openpyxl przy pliku, który nie jest poprawnym skoroszytem .xlsx: przy otwieraniu
# i, w trybie tylko do odczytu, także przy czytaniu wierszy uszkodzonego arkusza.
_UNREADABLE_WORKBOOK_ERRORS = (
    OSError,
    InvalidFileException,
    zipfile.BadZipFile,
    KeyError,
    ValueError,
    zlib.error,
    ParseError,
)

_NAME_PREFIX = "Rozliczenie "
_NAME_SUFFIX = " - zbiorcze.xlsx"
_NAME_PATTERN = re.compile(
    rf"{re.escape(_NAME_PREFIX)}(?P<period>{PERIOD_TOKEN_PATTERN}){re.escape(_NAME_SUFFIX)}"
)


class CollectiveFileError(Exception):
    """Plik zbiorczy albo skoroszyt w jego układzie nie nadaje się do użycia; komunikat jest bezpieczny."""


class LayoutError(CollectiveFileError):
    """Arkusz nie ma nagłówka WYKONAWCA w H17.

    ``worker_column`` to litera kolumny A:K, w której wiersz 17 ma WYKONAWCA,
    albo ``None``, gdy w A17:K17 go nie ma.
    """

    def __init__(self, worker_column: str | None) -> None:
        super().__init__("W wierszu 17 nie znaleziono nagłówka WYKONAWCA w kolumnie H.")
        self.worker_column = worker_column


@dataclass(frozen=True)
class DataRow:
    """Wiersz danych Pliku zbiorczego z surową wartością WYKONAWCA i wartościami A:AT."""

    worker: _ExcelCellValue
    values: ExcelRow


def _filename(period: str) -> str:
    return f"{_NAME_PREFIX}{period}{_NAME_SUFFIX}"


def path_for_period(period_directory: Path, period: str) -> Path:
    """Plik zbiorczy okresu w Folderze okresu."""

    return period_directory / _filename(period)


def period_from_path(path: Path) -> str:
    """Token okresu z nazwy Pliku zbiorczego leżącego w Folderze okresu o tej samej nazwie."""

    if "Rozliczenie wew." in path.name:
        raise CollectiveFileError("Plik Rozliczenie wew. jest poza zakresem tego procesu.")
    match = _NAME_PATTERN.fullmatch(path.name)
    if match is None:
        raise CollectiveFileError(f"Wybierz plik nazwany dokładnie: {_filename('{dd}_{dd}_{mm}_{rrrr}')}.")
    period = match.group("period")
    if path.parent.name != period:
        raise CollectiveFileError(f"Folder pliku musi nazywać się tak samo jak okres: {period}.")
    return period


def _is_worker_header(value: object) -> bool:
    return str(value or "").casefold().split() == ["wykonawca"]


def check_worker_header(sheet: Worksheet) -> None:
    """Arkusz ma nagłówek WYKONAWCA w H17, bez względu na wielkość liter i spacje.

    Inaczej ``LayoutError`` podaje kolumnę A:K z WYKONAWCA albo jej brak.
    Działa także dla arkusza otwartego tylko do odczytu.
    """

    if _is_worker_header(sheet.cell(row=HEADER_ROW, column=WORKER_COLUMN).value):
        return
    header = next(
        sheet.iter_rows(min_row=HEADER_ROW, max_row=HEADER_ROW, max_col=INPUT_COLUMNS[-1], values_only=True), ()
    )
    found = next((column for column, value in enumerate(header, start=1) if _is_worker_header(value)), None)
    raise LayoutError(None if found is None else get_column_letter(found))


def read_data_rows(path: Path) -> tuple[list[DataRow], list[Issue]]:
    """Wiersze danych w kolejności z pliku oraz problemy dla Wierszy danych bez Wykonawcy.

    Wiersz bez wartości w A:K nie jest Wierszem danych i jest pomijany. Plik jest tylko czytany.
    """

    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
    except _UNREADABLE_WORKBOOK_ERRORS as exc:
        raise _unreadable(path) from exc
    try:
        sheet = workbook.active
        if sheet is None:
            raise CollectiveFileError("Skoroszyt źródłowy nie zawiera arkusza.")
        check_worker_header(sheet)

        rows: list[DataRow] = []
        issues: list[Issue] = []
        last_copied_column = COPIED_COLUMNS[-1]
        for number, values in enumerate(
            sheet.iter_rows(
                min_row=FIRST_DATA_ROW, max_col=max(last_copied_column, sheet.max_column), values_only=True
            ),
            start=FIRST_DATA_ROW,
        ):
            if all(value is None for value in values[: INPUT_COLUMNS[-1]]):
                continue
            worker = values[WORKER_COLUMN - 1] if len(values) >= WORKER_COLUMN else None
            if not str(worker or "").split():
                issues.append(Issue("BRAK_WYKONAWCY", f"Wiersz {number} ma dane, ale nie ma wykonawcy w kolumnie H."))
                continue
            rows.append(DataRow(worker, tuple(values[:last_copied_column])))
        return rows, issues
    except _UNREADABLE_WORKBOOK_ERRORS as exc:
        # Tylko do odczytu arkusz jest parsowany leniwie, więc uszkodzony XML wychodzi dopiero przy czytaniu.
        raise _unreadable(path) from exc
    finally:
        workbook.close()


def _unreadable(path: Path) -> CollectiveFileError:
    return CollectiveFileError(
        f"Nie można odczytać Pliku zbiorczego {path.name}. "
        "Sprawdź, czy plik otwiera się w Excelu i nie jest uszkodzony."
    )
