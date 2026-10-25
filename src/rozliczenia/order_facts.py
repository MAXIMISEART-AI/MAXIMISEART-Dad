"""Treść Wiersza danych, która może opuścić komputer taty: reguła prywatności Kontroli ilości (ADR 0013, ADR 0015).

To jedyny typ z treścią komórek Pliku zbiorczego, który przeżywa silnik Przebiegu rozliczeń, i jedyne miejsce,
które wylicza dozwolone kolumny: UMOWA, URZĄDZENIA i CZYNNOŚCI (po nagłówku z wiersza 17, pod którym zapisuje je
Przygotowanie), dzień tygodnia i godzina z DATA oraz N:AP (po literze kolumny). Treść zlecenia
robią z komórek Reguły ilości (``OrderContent.of``), tak samo jak z Rejestru zleceń w Przygotowaniu. Adres, numery
zlecenia i umowy, Wykonawca, data dzienna, CZYNNOŚCI ZREALIZOWANE i reszta A:AT nie mają pola, więc nie da się ich
wysłać przez ten typ. Tekst wpisany przez tatę w komórce ilości też nie przechodzi: zostaje z niego „nieliczbowa”.
"""

from __future__ import annotations

from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from decimal import Decimal
from typing import Final, Literal

from openpyxl.utils import get_column_letter

from .preparation_register import ACTIVITIES, CONTRACT, DATE, DEVICES, Column
from .quantity_rules import OrderContent
from .workbook_files import QUANTITY_COLUMNS, DataRow, cell_datetime, is_blank
from .worker_config import normalize_text

NON_NUMERIC: Final = "nieliczbowa"
TataValue = int | Literal["nieliczbowa"]
"""Ilość taty: liczba (pusta komórka, 0 i ``"0"`` to 0; ``"2"`` i ``2.0`` to 2) albo „nieliczbowa” (formuła, słowo, 1.5)."""

N_AP: Final[tuple[str, ...]] = tuple(get_column_letter(column) for column in QUANTITY_COLUMNS)
_INPUTS: Final[tuple[Column, ...]] = (CONTRACT, DEVICES, ACTIVITIES, DATE)


@dataclass(frozen=True)
class OrderFacts:
    """Wiersz danych po filtrze prywatności."""

    content: OrderContent | None
    """Treść zlecenia dla Reguł ilości; ``None``, gdy ``DATA`` nie jest datą."""
    quantities: Mapping[str, TataValue]
    """Litera kolumny N:AP → ilość taty, tylko różna od 0."""

    def quantity(self, column: str) -> TataValue:
        return self.quantities.get(column, 0)


@dataclass(frozen=True)
class MissingHeaders:
    """W wierszu 17 Pliku zbiorczego nie ma nagłówka kolumny wejścia, więc nie wiadomo, gdzie ją czytać."""

    headers: tuple[str, ...]


CollectiveOrders = tuple[OrderFacts, ...] | MissingHeaders


def collective_orders(header: Mapping[str, int], rows: Sequence[DataRow]) -> CollectiveOrders:
    """Dozwolona treść każdego Wiersza danych.

    ``header``: znormalizowany nagłówek wiersza 17 → kolumna od 1.
    """

    names = [normalize_text(column.target_header) for column in _INPUTS]
    missing = tuple(column.target_header for column, name in zip(_INPUTS, names, strict=True) if name not in header)
    if missing:
        return MissingHeaders(missing)
    contract, devices, activities, started = (get_column_letter(header[name]) for name in names)
    return tuple(
        OrderFacts(
            _content(row.value(contract), row.value(devices), row.value(activities), row.value(started)), _quantities(row)
        )
        for row in rows
    )


def tata_value(cell: object) -> TataValue:
    """Decyzja 11: puste i 0 to 0, liczba całkowita (także tekstem albo jako float) to ta liczba, reszta „nieliczbowa”."""

    if is_blank(cell):
        return 0
    if isinstance(cell, bool):
        return NON_NUMERIC
    if isinstance(cell, int):
        return cell
    if isinstance(cell, float):
        return int(cell) if cell.is_integer() else NON_NUMERIC
    if isinstance(cell, Decimal):
        return int(cell) if cell.is_finite() and cell == cell.to_integral_value() else NON_NUMERIC
    if isinstance(cell, str):
        return _number_text(cell.strip())
    return NON_NUMERIC


def _number_text(text: str) -> TataValue:
    try:
        return int(text)
    except ValueError:
        pass
    try:
        number = float(text.replace(",", "."))
    except ValueError:
        return NON_NUMERIC
    return int(number) if number.is_integer() else NON_NUMERIC


def _content(contract: object, devices: object, activities: object, date_cell: object) -> OrderContent | None:
    started = cell_datetime(date_cell)
    return None if started is None else OrderContent.of(contract, devices, activities, started)


def _quantities(row: DataRow) -> dict[str, TataValue]:
    values = {letter: tata_value(row.value(letter)) for letter in N_AP}
    return {letter: value for letter, value in values.items() if value != 0}
