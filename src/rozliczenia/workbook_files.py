"""Skoroszyty na dysku: otwarcie z jedną klasyfikacją uszkodzeń, układ arkusza i odczyt Wierszy danych.

Każde otwarcie skoroszytu w obu procesach przechodzi przez ten moduł: Plik
zbiorczy (``read_data_rows``) i Rejestr zleceń (``read_workbook``) dają Wiersze
danych z Wyróżnieniami, formatami liczb i Hiperłączami komórek (``DataRow``), a
Wzorce czyta przez ``open_workbook`` plik z Wzorca (``template_file``). Otwarcie
sprawdza dostępność i klasyfikuje każde uszkodzenie, a problem zgłasza jako
``WorkbookError`` z powodem, przy złym układzie z jego miejscem. Moduł nie zna
Komunikatów: proces tłumaczy powód na Komunikat ze swojego katalogu.

Moduł jest właścicielem układu arkusza Pliku zbiorczego, wspólnego dla Wzorca
pliku zbiorczego, Wzorca rozliczenia pracownika i Szablonów pracowników:
nagłówek w wierszu 17 z WYKONAWCA w H, Wiersze danych od 18, dane wejściowe
A:K, kolumny wyglądu A:L (``row_look``), ilości N:AP i kolumny Wiersza danych A:AT. Publiczne są etykieta komórki
nagłówka ``WORKER_HEADER`` do treści Komunikatów oraz wiersze i kolumny układu dla pliku z Wzorca i
narzędzi benchmarku w ``tools/``, żeby nie powtarzały układu.
"""

from __future__ import annotations

import os
import zipfile
import zlib
from collections.abc import Iterator, Mapping, Sequence
from contextlib import contextmanager
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from enum import Enum
from io import BytesIO
from pathlib import Path
import stat
from types import MappingProxyType
from typing import IO, Never, TypeVar, cast
from xml.etree.ElementTree import ParseError
from xml.parsers.expat import ExpatError

from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet._read_only import ReadOnlyWorksheet
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula
from openpyxl.worksheet.worksheet import Worksheet

from .highlights import Highlights
from .hyperlinks import Hyperlink, read_sheet_hyperlinks
from .quantity_rules import QuantityColumn
from .row_look import GENERAL
from .workbook_parts import rows_dated_between, rows_with_input
from .worker_config import normalize_text


# Błędy openpyxl przy pliku, który nie jest poprawnym skoroszytem .xlsx: przy otwieraniu, a w trybie
# tylko do odczytu także przy czytaniu wierszy, bo arkusz jest wtedy parsowany leniwie. ``IndexError``:
# komórka ze stylem w skoroszycie bez ``xl/styles.xml``. ``ExpatError``: zepsuty XML arkusza przy odczycie
# Wierszy danych (``workbook_parts.rows_with_input``).
_UNREADABLE_WORKBOOK_ERRORS = (
    OSError,
    InvalidFileException,
    zipfile.BadZipFile,
    KeyError,
    IndexError,
    ValueError,
    zlib.error,
    ParseError,
    ExpatError,
)

# Układ arkusza Pliku zbiorczego i Wzorców tej rodziny. Publiczne są tylko liczby, których potrzebują narzędzia
# spoza pakietu (benchmark wpisuje ilości jak tata, generator składa syntetyczne Wzorce kształtu Wzorców taty).
HEADER_ROW = 17
FIRST_DATA_ROW = 18
_WORKER_COLUMN = 8  # H
INPUT_COLUMNS = range(1, 12)  # A:K: nagłówki w wierszu 17 i dane wejściowe Wiersza danych
LOOK_COLUMNS = range(1, 13)  # A:L: kolumny wyglądu Wiersza danych
QUANTITY_COLUMNS = range(14, 43)  # N:AP
DATA_ROW_COLUMNS = range(1, 47)  # A:AT: kopiowane z Pliku zbiorczego do Szablonu pracownika


_ExcelCellValue = (
    bool
    | int
    | float
    | Decimal
    | str
    | CellRichText
    | date
    | datetime
    | time
    | timedelta
    | DataTableFormula
    | ArrayFormula
    | bytes
    | None
)
ExcelRow = tuple[_ExcelCellValue, ...]


class Problem(Enum):
    """Powód, dla którego skoroszytu nie da się użyć."""

    MISSING = "brak pliku"
    OPEN_IN_EXCEL = "otwarty w Excelu"
    INACCESSIBLE = "niedostępny"
    """Plik jest cały, ale system nie oddał jego bajtów, np. zajęty przez inny program albo OneDrive bez sieci."""
    UNREADABLE = "nieczytelny"
    LAYOUT = "zły układ"
    HAS_DATA = "zawiera dane"
    """Wzorzec rozliczenia pracownika ma już wartości w Wierszach danych."""


@dataclass(frozen=True)
class MissingHeader:
    """Nagłówek ``header`` nie stoi na swoim miejscu w wierszu ``row`` arkusza ``sheet``.

    ``found_in``: litera kolumny A:K, w której stoi WYKONAWCA zamiast w H; ``None``: nagłówka w wierszu nie ma.
    """

    sheet: str
    row: int
    header: str
    found_in: str | None = None


@dataclass(frozen=True)
class MergedCells:
    """Komórki ``cells`` arkusza ``sheet``, do których trafiłyby Wiersze danych, są scalone."""

    sheet: str
    cells: str


LayoutFault = MissingHeader | MergedCells


class WorkbookError(Exception):
    """Skoroszytu nie da się użyć; ``problem`` mówi dlaczego, a przy złym układzie (i tylko wtedy) ``fault`` mówi gdzie."""

    def __init__(self, path: Path, problem: Problem, fault: LayoutFault | None = None) -> None:
        super().__init__(f"{problem.value}: {path.name}")
        self.path = path
        self.problem = problem
        self.fault = fault


@dataclass(frozen=True)
class HeaderCell:
    """Etykieta komórki nagłówka do treści Komunikatów: ``str()`` daje np. ``H17``."""

    column: str
    row: int

    def __str__(self) -> str:
        return f"{self.column}{self.row}"


WORKER_HEADER = HeaderCell(get_column_letter(_WORKER_COLUMN), HEADER_ROW)
_WORKER_LABEL = "WYKONAWCA"


_NO_ENTRIES: MappingProxyType[int, Never] = MappingProxyType({})
_NO_QUANTITIES: MappingProxyType[QuantityColumn, Never] = MappingProxyType({})


@dataclass(frozen=True)
class DataRow:
    """Wiersz danych: wartości komórek po kolei, ich Wyróżnienia, formaty liczb i Hiperłącza.

    Daje go odczyt Pliku zbiorczego (pozycje to kolumny A:AT, formaty liczb kolumn wyglądu A:L) i odczyt
    Rejestru zleceń (pozycje w kolejności ``preparation_register.COLUMNS``, czyli kolumn A:K Pliku
    zbiorczego, z formatem liczby każdej). ``number`` to numer wiersza w pliku, z którego go odczytano.
    Klucze ``hyperlinks`` to pozycje od 1; Hiperłącze ma tylko komórka z wartością.
    """

    number: int
    values: ExcelRow
    highlights: tuple[Highlights, ...]
    """Wyróżnienia komórek po pozycji jak ``values``."""
    number_formats: tuple[str, ...]
    """Formaty liczb komórek po pozycji; trafiają do kolumn wyglądu pliku z Wzorca."""
    hyperlinks: Mapping[int, Hyperlink] = field(default_factory=lambda: _NO_ENTRIES)
    quantities: Mapping[QuantityColumn, int] = field(default_factory=lambda: _NO_QUANTITIES)
    """Ilości usług Wiersza danych Pliku zbiorczego z Przygotowania: litera kolumny N:AP → liczba różna od zera.

    Trafiają do komórki Wzorca z jej wyglądem, bez Wyróżnień; reszta kolumn ilości zostaje jak we Wzorcu.
    Odczyt Pliku zbiorczego ma ilości w ``values``.
    """

    @property
    def worker(self) -> _ExcelCellValue:
        """Surowa wartość WYKONAWCA z kolumny H Pliku zbiorczego."""

        return self.values[_WORKER_COLUMN - 1]


def _status(path: Path) -> os.stat_result | None:
    """``stat`` pliku albo ``None``, gdy go nie ma; folder bez dostępu to ``OSError``.

    ``Path.exists`` i ``Path.is_file`` od Pythona 3.14 dają wtedy ``False``, więc plik bez dostępu byłby brakiem pliku.
    """

    try:
        return path.stat()
    except (FileNotFoundError, NotADirectoryError):
        return None


def is_open_in_excel(path: Path) -> bool:
    """Rozpoznaje skoroszyt otwarty w desktopowym Excelu po pliku blokady ``~$…``.

    W folderze bez dostępu blokady nie widać; niedostępny skoroszyt zgłasza jego odczyt (``ensure_available``).
    """

    return os.path.exists(path.with_name(f"~${path.name}"))


def ensure_available(path: Path, *, despite_excel_lock: bool = False) -> None:
    """Skoroszyt jest na dysku i nie jest otwarty w Excelu; inaczej ``WorkbookError``.

    Folder, którego nie da się odczytać (np. brak dostępu), czyni skoroszyt niedostępnym, a nie uszkodzonym.
    ``despite_excel_lock`` pozwala czytać skoroszyt, który trzyma Excel: jego ostatnią zapisaną wersję (ADR 0016).
    """

    try:
        status, open_in_excel = _status(path), is_open_in_excel(path)
    except OSError as exc:
        raise WorkbookError(path, Problem.INACCESSIBLE) from exc
    if status is None or not stat.S_ISREG(status.st_mode):
        raise WorkbookError(path, Problem.MISSING)
    if open_in_excel and not despite_excel_lock:
        raise WorkbookError(path, Problem.OPEN_IN_EXCEL)


_T = TypeVar("_T")


@contextmanager
def unreadable_as_error(path: Path) -> Iterator[None]:
    """Uszkodzenie skoroszytu zgłoszone przez openpyxl staje się ``WorkbookError``; inne błędy przechodzą."""

    try:
        yield
    except _UNREADABLE_WORKBOOK_ERRORS as exc:
        raise WorkbookError(path, Problem.UNREADABLE) from exc


def _file_bytes(path: Path) -> bytes:
    """Bajty skoroszytu z dysku; błąd systemu plików to ``Problem.INACCESSIBLE``, nie uszkodzenie."""

    try:
        return path.read_bytes()
    except OSError as exc:
        raise WorkbookError(path, Problem.INACCESSIBLE) from exc


def classified(path: Path, items: Iterator[_T]) -> Iterator[_T]:
    """Elementy leniwego odczytu openpyxl; klasyfikowane jest tylko pobranie elementu."""

    while True:
        with unreadable_as_error(path):
            try:
                item = next(items)
            except StopIteration:
                return
        yield item


def _open(path: Path, source: IO[bytes], *, read_only: bool) -> Workbook:
    """Skoroszyt ``path`` otwarty przez openpyxl z jego bajtów (``_file_bytes``) z klasyfikacją uszkodzeń."""

    with unreadable_as_error(path):
        if read_only:
            workbook = load_workbook(source, read_only=True, data_only=False)
        else:
            workbook = load_workbook(source, read_only=False, data_only=False, keep_links=True)
    if not isinstance(workbook.active, (Worksheet, ReadOnlyWorksheet)):
        workbook.close()
        raise WorkbookError(path, Problem.UNREADABLE)
    return workbook


def _is_worker_header(value: object) -> bool:
    return str(value or "").casefold().split() == ["wykonawca"]


def check_worker_header(path: Path, sheet: str, header: Sequence[object]) -> None:
    """Jedno sprawdzenie nagłówka WYKONAWCA w H17 (bez względu na wielkość liter i spacje) dla wartości A17:K17."""

    columns = [column for column, value in enumerate(header[: INPUT_COLUMNS[-1]], start=1) if _is_worker_header(value)]
    if _WORKER_COLUMN not in columns:
        found_in = get_column_letter(columns[0]) if columns else None
        raise WorkbookError(path, Problem.LAYOUT, MissingHeader(sheet, HEADER_ROW, _WORKER_LABEL, found_in))


def is_blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def cell_datetime(value: object) -> datetime | None:
    """Data i godzina z komórki ``DATA``; data bez godziny to północ, inna treść to ``None``."""

    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    return None


def header_columns(values: Sequence[object]) -> dict[str, int]:
    """Mapuje znormalizowane nagłówki na numery kolumn (od 1)."""

    return {normalize_text(value): column for column, value in enumerate(values, start=1) if not is_blank(value)}


@dataclass(frozen=True)
class DataRows:
    """Wiersze danych Pliku zbiorczego i nagłówki A17:K17 z tego samego odczytu.

    ``header``: znormalizowany nagłówek (``header_columns``) → kolumna od 1. Przygotowanie zapisuje kolumny A:K pod
    nagłówkami Wzorca, więc odczyt po nagłówku jest lustrem zapisu.
    """

    header: Mapping[str, int]
    rows: list[DataRow]


def read_data_rows(path: Path, *, despite_excel_lock: bool = False) -> DataRows:
    """Wiersze danych Pliku zbiorczego z Wyróżnieniami i Hiperłączami komórek w kolejności z pliku; plik jest tylko czytany.

    Wiersz bez wartości w A:K nie jest Wierszem danych i jest pomijany. Brak
    nagłówka WYKONAWCA w H17 to ``WorkbookError`` ze złym układem. Plik otwarty w Excelu to
    ``WorkbookError``, chyba że wywołujący jawnie pozwala go czytać (``despite_excel_lock``, ADR 0016).

    openpyxl czyta tylko wiersze, które mogą mieć wartości w A:K
    (``workbook_parts.rows_with_input``), a nie każdy z ~10 tys. sformatowanych wierszy Pliku
    zbiorczego; wartości i klasyfikacja uszkodzeń są takie jak przy odczycie
    całego arkusza.
    """

    ensure_available(path, despite_excel_lock=despite_excel_lock)
    content = _file_bytes(path)
    with unreadable_as_error(path):
        content = rows_with_input(content, INPUT_COLUMNS[-1])
    workbook = _open(path, BytesIO(content), read_only=True)
    try:
        sheet = cast(Worksheet, workbook.active)  # tylko do odczytu: ReadOnlyWorksheet
        header = next(
            classified(
                path,
                sheet.iter_rows(min_row=HEADER_ROW, max_row=HEADER_ROW, max_col=INPUT_COLUMNS[-1], values_only=True),
            ),
            (),
        )
        check_worker_header(path, sheet.title, header)
        with unreadable_as_error(path):
            max_column = max(DATA_ROW_COLUMNS[-1], sheet.max_column)
            links = read_sheet_hyperlinks(content, sheet.title)
        rows: list[DataRow] = []
        last_column = DATA_ROW_COLUMNS[-1]
        highlights = _HighlightsReader()
        for number, cells in enumerate(
            classified(path, sheet.iter_rows(min_row=FIRST_DATA_ROW, max_col=max_column)), start=FIRST_DATA_ROW
        ):
            values = tuple(cell.value for cell in cells[:last_column])
            if all(value is None for value in values[: INPUT_COLUMNS[-1]]):
                continue
            with unreadable_as_error(path):
                row_highlights = tuple(highlights.of(cell) for cell in cells[:last_column])
                number_formats = tuple(getattr(cell, "number_format", None) or GENERAL for cell in cells[: LOOK_COLUMNS[-1]])
            rows.append(DataRow(number, values, row_highlights, number_formats, links.row(number, values)))
    finally:
        workbook.close()
    return DataRows(header_columns(header), rows)


class _HighlightsReader:
    """Wyróżnienia komórek jednego odczytu; komórki o tej samej czcionce i wypełnieniu dzielą jeden obiekt."""

    def __init__(self) -> None:
        self._known: dict[tuple[int, int], Highlights] = {}

    def of(self, cell: object) -> Highlights:
        font, fill = getattr(cell, "font", None), getattr(cell, "fill", None)
        key = (id(font), id(fill))
        found = self._known.get(key)
        if found is None:
            found = self._known[key] = Highlights.of(font, fill)
        return found


class ReadSheet:
    """Arkusz skoroszytu otwartego tylko do odczytu; każdy odczyt komórek jest klasyfikowany.

    ``rows`` to zakres wierszy danych arkusza zawężonego przez ``ReadWorkbook.narrowed``; ``None``: wszystkie wiersze.
    """

    def __init__(self, path: Path, content: bytes, sheet: Worksheet, rows: range | None = None) -> None:
        self._path = path
        self._content = content
        self._sheet = sheet
        self._rows = rows

    @property
    def title(self) -> str:
        return self._sheet.title

    def values(self, min_row: int, max_row: int | None = None) -> Iterator[tuple[object, ...]]:
        return classified(self._path, self._sheet.iter_rows(min_row=min_row, max_row=max_row, values_only=True))

    def data_rows(self, min_row: int, columns: Sequence[int]) -> Iterator[tuple[tuple[object, ...], DataRow]]:
        """Wartości całego wiersza i Wiersz danych z komórek ``columns`` (od 1) po kolei, od ``min_row`` w dół.

        Wyróżnienia czyta ta sama reguła co w Pliku zbiorczym (``Highlights.of``), a format liczby każdej
        komórki to jej format, ogólny dla brakującej. W arkuszu zawężonym tylko wiersze z jego zakresu.
        """

        return classified(self._path, self._data_rows(min_row, columns))

    def _data_rows(self, min_row: int, columns: Sequence[int]) -> Iterator[tuple[tuple[object, ...], DataRow]]:
        max_row = None
        if self._rows is not None:
            min_row, max_row = max(min_row, self._rows.start), self._rows.stop - 1
            if min_row > max_row:
                return
        links = read_sheet_hyperlinks(self._content, self.title)
        highlights = _HighlightsReader()
        for number, cells in enumerate(self._sheet.iter_rows(min_row=min_row, max_row=max_row), start=min_row):
            values = tuple(getattr(cell, "value", None) for cell in cells)
            row_links = links.row(number, values)
            picked = [cells[column - 1] if column <= len(cells) else None for column in columns]
            yield values, DataRow(
                number,
                tuple(getattr(cell, "value", None) for cell in picked),
                tuple(highlights.of(cell) for cell in picked),
                tuple(getattr(cell, "number_format", None) or GENERAL for cell in picked),
                MappingProxyType(
                    {position: row_links[column] for position, column in enumerate(columns, start=1) if column in row_links}
                ),
            )


class ReadWorkbook:
    """Skoroszyt otwarty tylko do odczytu, np. Rejestr zleceń."""

    def __init__(self, path: Path, content: bytes, workbook: Workbook, ranges: Mapping[str, range] | None = None) -> None:
        self._path = path
        self._content = content
        self._workbook = workbook
        self._ranges = ranges or {}

    def sheets(self) -> Iterator[ReadSheet]:
        with unreadable_as_error(self._path):
            worksheets = list(self._workbook.worksheets)
        for sheet in worksheets:
            yield ReadSheet(self._path, self._content, cast(Worksheet, sheet), self._ranges.get(sheet.title))

    @contextmanager
    def narrowed(self, date_columns: Mapping[str, int], start: datetime, end: datetime) -> Iterator[ReadWorkbook]:
        """Ten skoroszyt z arkuszami ``date_columns`` (tytuł → kolumna daty od 1) zawężonymi do dat z [``start``, ``end``).

        Arkusz zawężony czyta tylko wiersze od pierwszego do ostatniego, którego
        data może leżeć w przedziale (``workbook_parts.rows_dated_between``); wartości, formaty i
        Hiperłącza tych wierszy są takie jak przy odczycie całego arkusza.
        """

        # Liczba dni od początku systemu dat skoroszytu (1900 albo 1904), z dniem zapasu z obu stron: openpyxl
        # zaokrągla część dnia do milisekund. Wiersz zakresu z datą spoza przedziału czytający i tak pomija.
        day = timedelta(days=1)
        epoch = self._workbook.epoch
        low, high = (start - day - epoch) / day, (end + day - epoch) / day
        with unreadable_as_error(self._path):
            content, ranges = rows_dated_between(self._content, date_columns, low, high)
        workbook = _open(self._path, BytesIO(content), read_only=True)
        try:
            yield ReadWorkbook(self._path, content, workbook, ranges)
        finally:
            workbook.close()


def workbook_content(path: Path) -> bytes:
    """Bajty skoroszytu, który jest na dysku i nie jest otwarty w Excelu; inaczej ``WorkbookError``."""

    ensure_available(path)
    return _file_bytes(path)


@contextmanager
def open_workbook(path: Path, content: bytes, *, read_only: bool) -> Iterator[Workbook]:
    """Skoroszyt ``path`` z jego bajtów (``workbook_content``) otwarty przez openpyxl z klasyfikacją uszkodzeń.

    Bez ``read_only`` z odwołaniami zewnętrznymi, do zapisu. Zamykany przy wyjściu z bloku.
    """

    workbook = _open(path, BytesIO(content), read_only=read_only)
    try:
        yield workbook
    finally:
        workbook.close()


@contextmanager
def read_workbook(path: Path, content: bytes) -> Iterator[ReadWorkbook]:
    """Skoroszyt ``path`` z jego bajtów (``workbook_content``) tylko do odczytu z formatem i Hiperłączami komórek.

    Zamykany przy wyjściu z bloku. openpyxl i Hiperłącza czytają te same bajty, odczytane z dysku raz.
    """

    workbook = _open(path, BytesIO(content), read_only=True)
    try:
        yield ReadWorkbook(path, content, workbook)
    finally:
        workbook.close()
