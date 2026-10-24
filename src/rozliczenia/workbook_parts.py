"""Części zip skoroszytu .xlsx: nawigacja po relacjach i wierszach arkusza oraz skoroszyt zawężony do potrzebnych wierszy.

Moduł jest szczegółem ``workbook_files``, ``template_copy`` i ``hyperlinks`` i
nikt poza nimi go nie używa. Wszyscy czytają skoroszyt po częściach zip, bez
openpyxl: ``relationships`` i ``sheet_part`` znajdują część po relacjach, a
``SheetData`` dzieli XML arkusza na wiersze bez parsera.

``SheetData`` przyjmuje tylko kształt, który zapisują Excel i openpyxl: jeden
``<sheetData>``, każdy wiersz i każda komórka z ``r`` jako pierwszym atrybutem,
wiersze rosnąco po numerze i bez komentarzy ani CDATA w danych arkusza. Inny
kształt to ``IrregularSheetData``: zawężanie zostawia wtedy arkusz w całości, a
``template_copy`` oddaje zapis openpyxl.

Plik zbiorczy ma ~10 tys. sformatowanych wierszy, a dane tylko w kilkuset;
Rejestr zleceń dziesiątki tysięcy wierszy od 2021 r., a okres kilkaset. openpyxl
czytałby każdą komórkę każdego wiersza. ``rows_with_input`` i
``rows_dated_between`` składają więc z bajtów skoroszytu taki sam skoroszyt
(części przepisane bez kompresji), w którego arkuszach zostają tylko potrzebne
wiersze i wiersze z formułą główną wspólnej formuły, bo openpyxl rozwija z niej
formuły zależne w dalszych wierszach. Wartości komórek dalej odczytuje openpyxl,
więc są takie same jak przy odczycie całego arkusza (typy, daty, teksty wspólne
i wbudowane, formuły). Pominięte wiersze nie są czytane przez openpyxl, więc XML
zawężanego arkusza jest najpierw sprawdzany parserem expat, żeby zepsuty arkusz
dalej był nieczytelny (``xml.parsers.expat.ExpatError``).

``rows_with_input`` (Wiersze danych Pliku zbiorczego) zostawia wiersze, które
mogą mieć dane w kolumnach danych wejściowych, i nagłówek. Pominięty wiersz nie
ma tam wartości: każda jego komórka w tych kolumnach jest pustym elementem
``<c …/>``, a openpyxl czyta taki wiersz tak samo jak brakujący.

``rows_dated_between`` (zlecenia okresu z Rejestru zleceń) zostawia wiersze od
pierwszego do ostatniego z liczbą z przedziału w kolumnie daty; nagłówki
czytający bierze wcześniej z całego skoroszytu. openpyxl daje datę tylko z liczby
w komórce bez typu albo z ``t="n"`` (ze stylem daty), więc wiersz, którego data
leży w przedziale, ma w kolumnie daty liczbę z przedziału. Zakres od pierwszej do
ostatniej takiej liczby zawiera więc każdy taki wiersz i wszystkie wiersze między
nimi. Liczba bez stylu daty albo wartość zapamiętana przy formule może zakres
tylko poszerzyć. Arkusz z komórką daty typu ``d`` (data tekstem ISO, której
Excel nie zapisuje) albo z wartością liczbową, której nie da się przeczytać bez
parsera XML, zostaje w całości.
"""

from __future__ import annotations

import posixpath
import re
import zipfile
from bisect import bisect_right
from collections.abc import Callable, Iterable, Iterator, Mapping
from contextlib import contextmanager, suppress
from dataclasses import dataclass
from functools import cached_property
from io import BytesIO
from typing import Final, NoReturn
from xml.etree import ElementTree
from xml.parsers import expat

from openpyxl.utils import get_column_letter

MAIN: Final = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
OFFICE_REL: Final = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PACKAGE_REL: Final = "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship"

_WITHOUT_LEADING_R: Final = re.compile(rb'<(?:c|row)(?=[\s>/])(?! r=")')
_ROW_START_OR_IRREGULAR: Final = re.compile(rb'<(?:row r="(\d+)"|!|(?:c|row)(?=[\s>/])(?! r="))')
_WORKSHEET: Final = re.compile(rb"\s*(?:<\?xml[^>]*\?>\s*)?<worksheet[\s>]")
_SHARED_FORMULA_MASTER: Final = re.compile(rb'<f\b[^>]*\bt="shared"[^>]*(?<!/)>')
_TYPE: Final = re.compile(rb"""\st=["']([^"']*)["']""")
_VALUE: Final = re.compile(rb"<v\b[^>]*>([^<]*)</v>")
_NUMBER_TYPES: Final = (None, b"n")


# --- części i relacje --------------------------------------------------------------------------


@contextmanager
def package_parts(content: bytes) -> Iterator[Mapping[str, bytes]]:
    """Części skoroszytu ``content`` (.xlsx) po nazwie w zip; zepsuty zip zgłasza ``zipfile`` (np. ``BadZipFile``)."""

    with zipfile.ZipFile(BytesIO(content)) as archive:
        yield _ZipParts(archive)


class _ZipParts(Mapping[str, bytes]):
    """Części zip po nazwie, rozpakowane dopiero przy odczycie: arkusz Rejestru zleceń ma dziesiątki MB."""

    def __init__(self, archive: zipfile.ZipFile) -> None:
        self._archive = archive

    def __getitem__(self, name: str) -> bytes:
        return self._archive.read(name)

    def __contains__(self, name: object) -> bool:
        try:
            self._archive.getinfo(str(name))
        except KeyError:
            return False
        return True

    def __iter__(self) -> Iterator[str]:
        return iter(self._archive.namelist())

    def __len__(self) -> int:
        return len(self._archive.infolist())


@dataclass(frozen=True)
class Relationship:
    """Relacja części: typ (``Type``) i cel, czyli ścieżka części w zip albo adres zewnętrzny."""

    kind: str
    target: str
    external: bool


def relationships_part(part: str) -> str:
    """Nazwa części z relacjami części ``part``."""

    folder, name = posixpath.split(part)
    return posixpath.join(folder, "_rels", f"{name}.rels")


def relationships(parts: Mapping[str, bytes], part: str) -> dict[str, Relationship]:
    """Relacje części ``part`` po identyfikatorze; bez części relacji żadnych."""

    rels = relationships_part(part)
    if rels not in parts:
        return {}
    folder = posixpath.dirname(part)
    found: dict[str, Relationship] = {}
    for relationship in ElementTree.fromstring(parts[rels]).iter(PACKAGE_REL):
        target = relationship.get("Target", "")
        external = relationship.get("TargetMode") == "External"
        if not external:
            target = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(folder, target))
        found[relationship.get("Id", "")] = Relationship(relationship.get("Type", ""), target, external)
    return found


def office_document(parts: Mapping[str, bytes]) -> str | None:
    """Część skoroszytu (``xl/workbook.xml``) według relacji paczki; ``None``, gdy żadna relacja jej nie wskazuje."""

    for relationship in ElementTree.fromstring(parts["_rels/.rels"]).iter(PACKAGE_REL):
        if relationship.get("Type", "").endswith("/officeDocument"):
            return relationship.get("Target", "").lstrip("/")
    return None


def sheet_part(parts: Mapping[str, bytes], title: str) -> str:
    """Część arkusza o nazwie ``title`` według ``workbook.xml`` i jego relacji; brak którejś to ``KeyError``."""

    workbook = office_document(parts) or ""
    sheets = ElementTree.fromstring(parts[workbook]).iter(f"{MAIN}sheet")
    identifier = next((sheet.get(f"{{{OFFICE_REL}}}id", "") for sheet in sheets if sheet.get("name") == title), "")
    return relationships(parts, workbook)[identifier].target


# --- wiersze arkusza ---------------------------------------------------------------------------


class IrregularSheetData(Exception):
    """``<sheetData>`` spoza kształtu, który zapisują Excel i openpyxl."""


@dataclass(frozen=True)
class SheetRow:
    """Wiersz arkusza: numer i granice jego elementu ``<row>`` w XML arkusza."""

    number: int
    start: int
    end: int


@dataclass(frozen=True)
class SheetData:
    """Jedyny ``<sheetData>`` arkusza: treść od ``start`` (za ``<sheetData>``) do ``end`` (``</sheetData>``) i jej wiersze."""

    start: int
    end: int
    rows: list[SheetRow]
    """Wiersze ``<row>`` w kolejności z pliku."""

    @classmethod
    def of(cls, xml: bytes) -> SheetData:
        """``<sheetData>`` XML arkusza; kształt spoza tego, który zapisują Excel i openpyxl, to ``IrregularSheetData``."""

        start = xml.find(b"<sheetData>")
        end = xml.find(b"</sheetData>")
        if start < 0 or end < start or xml.count(b"<sheetData") != 1:
            raise IrregularSheetData("sheetData")
        start += len(b"<sheetData>")
        return cls(start, end, _sheet_rows(xml, start, end))

    @cached_property
    def _starts(self) -> list[int]:
        return [row.start for row in self.rows]

    def row_index(self, position: int) -> int:
        """Indeks w ``rows`` ostatniego wiersza, który zaczyna się przed miejscem ``position``; -1 przed pierwszym."""

        return bisect_right(self._starts, position) - 1


def _sheet_rows(xml: bytes, start: int, end: int) -> list[SheetRow]:
    # Jedno przejście po arkuszu (dziesiątki MB w Rejestrze zleceń) zamiast osobnego dla każdego warunku kształtu.
    rows: list[SheetRow] = []
    for match in _ROW_START_OR_IRREGULAR.finditer(xml, start, end):
        if match.group(1) is None:
            _raise_irregular(xml, start, end)
        number = int(match.group(1))
        if rows and number <= rows[-1].number:
            _raise_irregular(xml, start, end, "kolejność wierszy")
        tag_end = xml.index(b">", match.end())
        row_end = tag_end + 1 if xml[tag_end - 1 : tag_end] == b"/" else xml.index(b"</row>", tag_end) + len(b"</row>")
        rows.append(SheetRow(number, match.start(), row_end))
    return rows


def _raise_irregular(xml: bytes, start: int, end: int, otherwise: str = "") -> NoReturn:
    """Pierwszy niespełniony warunek kształtu w kolejności: komentarz albo CDATA, ``r`` jako pierwszy atrybut, ``otherwise``."""

    if xml.find(b"<!", start, end) >= 0:
        raise IrregularSheetData("komentarz albo CDATA")
    if _WITHOUT_LEADING_R.search(xml, start, end):
        raise IrregularSheetData("r")
    raise IrregularSheetData(otherwise)


# --- zawężanie ---------------------------------------------------------------------------------

_Narrow = Callable[[bytes, SheetData], bytes]
"""Zawężenie XML arkusza o danym ``SheetData``; arkusz spoza jego kształtu to ``IrregularSheetData``."""
_Choose = Callable[[str, bytes], _Narrow | None]
"""Zawężenie części zip o tej nazwie i treści; ``None``: część zostaje bez zmian."""


def rows_with_input(content: bytes, input_columns: int) -> bytes:
    """Skoroszyt ``content`` (.xlsx) z arkuszami zawężonymi do wierszy, które mogą mieć dane w kolumnach 1..``input_columns``.

    Pozostałe części są przepisane bez zmian. Uszkodzenie skoroszytu zgłaszają
    ``zipfile`` (np. ``BadZipFile``), ``zlib`` albo ``expat`` (``ExpatError``).
    """

    letters = b"|".join(get_column_letter(column).encode("ascii") for column in range(1, input_columns + 1))
    with_input = re.compile(rb'<c r="(?:' + letters + rb')[0-9]+"[^>]*(?<!/)>')

    def narrow(xml: bytes, sheet: SheetData) -> bytes:
        cells = with_input.finditer(xml, sheet.start, sheet.end)
        return _with_rows(xml, sheet, {sheet.row_index(cell.start()) for cell in cells})

    def every_worksheet(_: Mapping[str, bytes]) -> _Choose:
        return lambda _, xml: narrow if _WORKSHEET.match(xml) else None

    return _rewritten(content, every_worksheet)


def rows_dated_between(
    content: bytes, date_columns: Mapping[str, int], low: float, high: float
) -> tuple[bytes, dict[str, range]]:
    """Skoroszyt ``content`` (.xlsx) z arkuszami ``date_columns`` (tytuł → kolumna daty od 1) zawężonymi do zakresu dat.

    Zakres arkusza to wiersze od pierwszego do ostatniego z liczbą z [``low``,
    ``high``] w kolumnie daty; zwracany jest dla każdego zawężonego arkusza (pusty,
    gdy żadna komórka daty nie ma takiej liczby). Arkusz bez zwróconego zakresu
    zostaje w całości, a pozostałe części są przepisane bez zmian. Uszkodzenie
    skoroszytu zgłaszają ``zipfile`` (np. ``BadZipFile``), ``zlib``, brak części
    ``KeyError`` albo ``expat`` (``ExpatError``).
    """

    ranges: dict[str, range] = {}

    def dated(title: str, column: int) -> _Narrow:
        letter = get_column_letter(column).encode("ascii")
        date_cell = re.compile(rb'<c r="' + letter + rb'(\d+)"([^>]*)(?<!/)>(.*?)</c>', re.DOTALL)

        def narrow(xml: bytes, sheet: SheetData) -> bytes:
            numbers = [
                int(found.group(1))
                for found in date_cell.finditer(xml, sheet.start, sheet.end)
                if _is_number_between(found.group(2), found.group(3), low, high)
            ]
            kept = range(min(numbers), max(numbers) + 1) if numbers else range(0)
            narrowed = _with_rows(xml, sheet, {index for index, row in enumerate(sheet.rows) if row.number in kept})
            ranges[title] = kept
            return narrowed

        return narrow

    def by_title(parts: Mapping[str, bytes]) -> _Choose:
        narrowing = {sheet_part(parts, title): dated(title, column) for title, column in date_columns.items()}
        return lambda name, _: narrowing.get(name)

    return _rewritten(content, by_title), ranges


def _rewritten(content: bytes, choose: Callable[[Mapping[str, bytes]], _Choose]) -> bytes:
    """Skoroszyt ``content`` przepisany bez kompresji, z częściami zawężonymi tak, jak wybiera ``choose`` z jego części.

    Arkusz spoza kształtu (``IrregularSheetData``) zostaje w całości.
    """

    reduced = BytesIO()
    with zipfile.ZipFile(BytesIO(content)) as archive, zipfile.ZipFile(reduced, "w", zipfile.ZIP_STORED) as target:
        narrowing = choose(_ZipParts(archive))
        for info in archive.infolist():
            data = archive.read(info)
            narrow = narrowing(info.filename, data)
            if narrow is not None:
                with suppress(IrregularSheetData):
                    data = narrow(data, SheetData.of(data))
            target.writestr(info.filename, data)
    return reduced.getvalue()


def _with_rows(xml: bytes, sheet: SheetData, kept: Iterable[int]) -> bytes:
    """XML arkusza tylko z wierszami o indeksach ``kept`` i z wierszami formuł głównych wspólnych formuł.

    Najpierw sprawdza cały XML: openpyxl nie przeczyta pominiętych wierszy, a zepsuty arkusz ma być nieczytelny.
    """

    masters = (sheet.row_index(found.start()) for found in _SHARED_FORMULA_MASTER.finditer(xml, sheet.start, sheet.end))
    indices = sorted({*kept, *masters} - {-1})
    expat.ParserCreate().Parse(xml, True)
    rows = sheet.rows
    return b"".join([xml[: sheet.start], *(xml[rows[index].start : rows[index].end] for index in indices), xml[sheet.end :]])


def _is_number_between(attributes: bytes, content: bytes, low: float, high: float) -> bool:
    """Komórka z liczbą z [``low``, ``high``]: tylko taką openpyxl może odczytać jako datę z przedziału."""

    found_type = _TYPE.search(attributes)
    cell_type = found_type.group(1) if found_type else None
    if cell_type == b"d":
        raise IrregularSheetData("data tekstem ISO")
    if cell_type not in _NUMBER_TYPES:
        return False
    value = _VALUE.search(content)
    if value is None:
        return False
    try:
        number = float(value.group(1))
    except ValueError:
        raise IrregularSheetData("wartość liczbowa") from None
    return low <= number <= high
