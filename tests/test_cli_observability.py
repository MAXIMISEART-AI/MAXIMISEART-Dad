from __future__ import annotations

from datetime import date
from io import StringIO
import json
import os
from pathlib import Path
import subprocess
import sys
import zipfile
from collections.abc import Iterable
from typing import Any, cast

from openpyxl import load_workbook
import pytest

from rozliczenia import settings as settings_file
from rozliczenia.cli import main
from rozliczenia.domain import WorkerResult
import rozliczenia.template_settlement as template_settlement
from rozliczenia.telemetry import MetricsStore
from rozliczenia.workbook_files import ExcelRow, WorkerSettlementTemplate

from tests.fixtures.terminal import TtyOutput, install_fake_rich
from tests.fixtures.settlement import PERIOD, make_fixture
from tests.fixtures.workbooks import active_worksheet
from tests.fixtures.settlement import append_completed_records, install_source_dialog, make_preview_fixture, run_cli


def assert_fixture_workbooks(target_directory: Path, *, dry_run: bool) -> None:
    if dry_run:
        assert not target_directory.exists()
        return

    expected_filenames = {
        f"Rozliczenie {PERIOD} - {worker_name}.xlsx"
        for worker_name in ("Adrian Testowy", "Darek Fikcyjny", "Kamil Syntetyczny")
    }
    assert {path.name for path in target_directory.glob("*.xlsx")} == expected_filenames
    for worker_name in ("Adrian Testowy", "Darek Fikcyjny"):
        workbook = load_workbook(target_directory / f"Rozliczenie {PERIOD} - {worker_name}.xlsx")
        try:
            sheet = active_worksheet(workbook)
            assert sheet["A18"].value == "POZNAŃ"
            assert sheet["AU18"].value == "=N18"
        finally:
            workbook.close()

    workbook = load_workbook(target_directory / f"Rozliczenie {PERIOD} - Kamil Syntetyczny.xlsx")
    try:
        sheet = active_worksheet(workbook)
        assert sheet["A18"].value is None
        assert sheet["AU18"].value == "=N18"
    finally:
        workbook.close()


def assert_completed_metrics(
    metrics_path: Path,
    *,
    dry_run: bool,
    issues: int = 0,
) -> dict[str, Any]:
    record = json.loads(metrics_path.read_text(encoding="utf-8"))
    assert record["schema_version"] == 1
    assert record["mode"] == ("DRY-RUN" if dry_run else "RUN")
    assert record["completed"] is True
    assert record["counters"] == {
        "templates_total": 3,
        "templates_completed": 3,
        "rows": 3,
        "written": 0 if dry_run else 2,
        "empty": 1,
        "planned": 2 if dry_run else 0,
        "skipped": 0,
        "issues": issues,
    }
    return cast(dict[str, Any], record)


def test_cli_runs_without_rich(tmp_path: Path) -> None:
    source_path, target_directory, config_path = make_fixture(tmp_path)
    environment = os.environ.copy()
    environment["PYTHONPATH"] = str(Path(__file__).parents[1] / "src")
    arguments = [
        "--source",
        str(source_path),
        "--config",
        str(config_path),
        "--worker-template",
        str(config_path.parent / "Wzorzec rozliczenia pracownika.xlsx"),
        "--metrics",
        str(tmp_path / "metrics.jsonl"),
    ]
    result = subprocess.run(
        [
            sys.executable,
            "-c",
            "import sys; sys.modules['rich'] = None; from rozliczenia.cli import main; raise SystemExit(main())",
            *arguments,
        ],
        capture_output=True,
        text=True,
        env=environment,
    )

    assert result.returncode == 0, result.stderr
    assert "Etap: Sprawdzanie" in result.stdout
    assert "Postęp szablonów: 3/3 (100%)" in result.stdout
    assert "WYKONAWCA: Adrian Testowy" in result.stdout
    assert "Status końcowy: OK" in result.stdout
    assert_fixture_workbooks(target_directory, dry_run=False)


@pytest.mark.parametrize(
    ("dry_run", "mode"),
    [(False, "RUN"), (True, "DRY-RUN")],
)
def test_interactive_rich_cli_shows_szablony_pracownika(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    dry_run: bool,
    mode: str,
) -> None:
    install_fake_rich(monkeypatch)
    fixture = make_preview_fixture(tmp_path) if dry_run else make_fixture(tmp_path)
    source_path, _, config_path = fixture
    output = TtyOutput()
    arguments = [
        "--source",
        str(source_path),
        "--config",
        str(config_path),
        "--worker-template",
        str(config_path.parent / "Wzorzec rozliczenia pracownika.xlsx"),
        "--metrics",
        str(tmp_path / "metrics.jsonl"),
    ]
    if dry_run:
        arguments.append("--dry-run")

    exit_code = main(arguments, output=output)

    assert exit_code == (2 if dry_run else 0)
    rendered = output.getvalue()
    assert f"Rozliczenia | {PERIOD} | {mode}" in rendered
    assert "Postęp szablonów: 3/3 (100%)" in rendered
    assert "Bieżący WYKONAWCA: Adrian Testowy" in rendered
    assert ("PLAN" if dry_run else "ZAPISANO") in rendered
    assert "syntetyczny adres" not in rendered
    assert "#1" not in rendered


def test_cli_dry_run_changes_no_file_and_records_dry_run_metrics(tmp_path: Path) -> None:
    source_path, target_directory, config_path = make_preview_fixture(tmp_path)
    worker_template_path = tmp_path / "Wzorzec rozliczenia pracownika.xlsx"
    source_before = source_path.read_bytes()
    worker_template_before = worker_template_path.read_bytes()

    exit_code, _ = run_cli(source_path, config_path, tmp_path / "metrics.jsonl", "--dry-run")

    assert exit_code == 2
    assert not target_directory.exists()
    assert source_path.read_bytes() == source_before
    assert worker_template_path.read_bytes() == worker_template_before
    assert_completed_metrics(tmp_path / "metrics.jsonl", dry_run=True, issues=1)


def test_critical_failure_is_recorded_as_incomplete_run(tmp_path: Path) -> None:
    period_directory = tmp_path / PERIOD
    period_directory.mkdir()
    source_path = period_directory / f"Rozliczenie {PERIOD} - zbiorcze.xlsx"
    metrics_path = tmp_path / "metrics.jsonl"

    exit_code, output = run_cli(source_path, tmp_path / "missing.yaml", metrics_path)

    assert exit_code == 1
    assert "Nie wykonano" in output
    assert "Status semantyczny: BŁĄD" in output
    record = json.loads(metrics_path.read_text(encoding="utf-8"))
    assert record["completed"] is False
    assert record["result"] == "BLAD_KRYTYCZNY"


def run_cli_from_ustawienia_taty(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, *, today: date
) -> tuple[int, str, Path, list[str]]:
    """Launcher bez ``--source``; Ustawienia taty wskazują ``Załączniki`` w układzie folderów taty."""

    periods_directory = tmp_path / "Załączniki"
    month = periods_directory / "Rozliczenia 2026" / "09_Wrzesień_2026" / "INEA"
    month.mkdir(parents=True)
    _, target_directory, config_path = make_fixture(month)
    shown = install_source_dialog(monkeypatch, selected="")
    monkeypatch.setattr(settings_file, "load_settings", lambda: settings_file.Settings(periods_directory=periods_directory))
    output = StringIO()
    exit_code = main(
        [
            "--config",
            str(config_path),
            "--worker-template",
            str(month / "Wzorzec rozliczenia pracownika.xlsx"),
            "--metrics",
            str(tmp_path / "metrics.jsonl"),
        ],
        output=output,
        input_stream=StringIO("\n"),
        today=today,
    )
    return exit_code, output.getvalue(), target_directory, shown


def test_with_ustawienia_taty_launcher_asks_for_period_instead_of_plik_zbiorczy_dialog(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    exit_code, output, target_directory, shown = run_cli_from_ustawienia_taty(
        tmp_path, monkeypatch, today=date(2026, 9, 16)
    )

    assert exit_code == 0, output
    assert shown == []
    assert "Proponowany okres: 08_14_09_2026 (wtorek 08.09.2026 – poniedziałek 14.09.2026)" in output
    assert f"Rozliczenia | okres: {PERIOD} | tryb: RUN" in output
    assert len(list(target_directory.glob("*.xlsx"))) == 3


def test_with_ustawienia_taty_period_without_folder_okresu_stops_with_one_line(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    exit_code, output, _, shown = run_cli_from_ustawienia_taty(tmp_path, monkeypatch, today=date(2026, 9, 23))

    assert exit_code == 1
    assert shown == []
    lines = output.splitlines()
    status = lines.index("Status semantyczny: BŁĄD")
    assert lines[status + 1] == "Nie wykonano: Nie znaleziono folderu 15_21_09_2026 w żadnym folderze miesiąca."


def test_cancelled_plik_zbiorczy_selection_stops_with_statistics_and_without_dialog_time(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    metrics_path = tmp_path / "metrics.jsonl"
    append_completed_records(metrics_path, "RUN")
    shown = install_source_dialog(monkeypatch, selected="", answer_after_s=1.0)

    output = StringIO()
    exit_code = main(["--metrics", str(metrics_path)], output=output)

    assert exit_code == 1
    assert shown == ["Wybierz plik zbiorczy rozliczenia"]
    lines = output.getvalue().splitlines()
    assert lines[0] == "Rozliczenia | okres: nieznany | tryb: RUN"
    status = lines.index("Status semantyczny: BŁĄD")
    assert lines[status + 1] == "Nie wykonano: Nie wybrano pliku zbiorczego."
    assert lines[status + 2].startswith("Czas uruchomienia: ")
    assert int(lines[status + 2].removeprefix("Czas uruchomienia: ").removesuffix(" ms")) < 1000
    assert lines[status + 3] == "Statystyki RUN | próbek: 5"
    record = MetricsStore(metrics_path).read()[-1]
    assert (record["period"], record["completed"], record["result"]) == ("nieznany", False, "BLAD_KRYTYCZNY")
    assert record["total_duration_ms"] < 1000


def test_broken_wzorzec_rozliczenia_pracownika_shows_unreadable_wzorzec_message(tmp_path: Path) -> None:
    source_path, target_directory, config_path = make_fixture(tmp_path)
    worker_template_path = config_path.parent / "Wzorzec rozliczenia pracownika.xlsx"
    with zipfile.ZipFile(source_path) as source:
        entries = [(entry.filename, source.read(entry)) for entry in source.infolist()]
    with zipfile.ZipFile(worker_template_path, "w") as broken:
        for name, data in entries:
            broken.writestr(name, b"<x" if name.startswith("xl/worksheets/") else data)

    exit_code, output = run_cli(source_path, config_path, tmp_path / "metrics.jsonl")

    assert exit_code == 1
    lines = output.splitlines()
    status = lines.index("Status semantyczny: BŁĄD")
    assert lines[status + 1] == "Nie wykonano: Nie można odczytać Wzorca rozliczenia pracownika."
    assert lines[status + 2].startswith("Czas uruchomienia: ")
    assert "unclosed token" not in output
    assert "Traceback" not in output
    assert not target_directory.exists()


def test_critical_write_failure_shows_interrupted_szablon_pracownika_and_records_incomplete_run(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source_path, target_directory, config_path = make_fixture(tmp_path)
    metrics_path = tmp_path / "metrics.jsonl"
    original_process_template = template_settlement.process_template
    save_calls = 0

    def fail_on_second_save(
        worker_name: str,
        template: WorkerSettlementTemplate,
        path: Path,
        rows: Iterable[ExcelRow],
    ) -> WorkerResult:
        nonlocal save_calls
        save_calls += 1
        if save_calls == 2:
            raise template_settlement.TemplateWriteError("synthetic write failure")
        return original_process_template(worker_name, template, path, rows)

    monkeypatch.setattr(template_settlement, "process_template", fail_on_second_save)

    exit_code, output = run_cli(source_path, config_path, metrics_path)

    assert exit_code == 1
    assert not target_directory.exists()
    assert not list(target_directory.parent.glob(f".{target_directory.name}.staging-*"))
    assert "folder wynikowy nie został opublikowany" in output
    assert "Postęp szablonów: 1/3 (33%)" in output
    assert (
        "Etap przerwany: Zapisywanie | WYKONAWCA: Darek Fikcyjny | "
        "Szablon pracownika: 2/3"
    ) in output
    record = json.loads(metrics_path.read_text(encoding="utf-8"))
    assert record["completed"] is False
    assert record["result"] == "BLAD_KRYTYCZNY"
    assert set(record["counters"].values()) == {0}
