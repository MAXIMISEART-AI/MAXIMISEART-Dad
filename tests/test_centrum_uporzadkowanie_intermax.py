"""Uporządkowanie INTERMAX z Centrum rozliczeń przez HTTP API i strumień zdarzeń, jak widzi je przeglądarka (#470).

Karta „Uporządkuj INTERMAX” stoi poza okresami: INTERMAX to Rejestr zleceń z Ustawień taty, a okres Uruchomienia to
okres dzisiejszej daty z zegara procesu, ustalony przy starcie (ADR 0019).
"""

from __future__ import annotations

from collections.abc import Callable
from datetime import datetime, timedelta
import json
from pathlib import Path
from typing import Any
from uuid import uuid4

from openpyxl import Workbook, load_workbook
import pytest

from rozliczenia import intermax_copies
from rozliczenia.centrum.run_logs import RunLog

from tests.fixtures.centrum import phase_states, views_until
from tests.fixtures.centrum_intermax import (
    PERIOD,
    REGISTER_NAME,
    Clock,
    OrderingCentrum,
    copies,
    open_in_excel,
    ordering_centrum,
)
from tests.fixtures.intermax import ACTIVITIES, DEVICES, ORDER_SHEET, Intermax, Order
from tests.fixtures.intermax_disk import denied, failing_replace

COPY = "INTERMAX przed uporządkowaniem 2026-10-26 14-05-07.xlsx"
SUMMARY: dict[str, Any] = {
    "figures": [
        {"label": "Zlecenia", "value": 6, "unit": "count"},
        {"label": "Wykonawcy", "value": 3, "unit": "count"},
    ],
    "counts": {"title": "Uporządkowany dzień", "counts": [{"label": "pon 26.10.2026", "count": 6}]},
    "files": [],
    "worker_settlements": None,
    "intermax_copy": COPY,
}
OPEN = (
    "INTERMAX jest otwarty w Excelu. Zapisz go i zamknij. Jeśli Excel jest zamknięty, usuń z tego samego folderu "
    f"pozostawiony plik ~${REGISTER_NAME}."
)


def _stopped(centrum: OrderingCentrum) -> dict[str, Any]:
    view = centrum.run_to_end()
    assert view["result"]["ending"] == "stop", view["result"]
    return view


def test_bez_wskazanego_rejestru_zlecen_karta_mowi_co_ustawic_a_start_jest_odrzucony(tmp_path: Path) -> None:
    centrum = ordering_centrum(tmp_path, chosen=False)

    setup = centrum.setup()
    started = centrum.start()

    assert setup == {
        "title": "Uporządkowanie INTERMAX",
        "reminder": "Zapisz i zamknij INTERMAX w Excelu",
        "blocker": "Najpierw wskaż INTERMAX.",
        "open_in_excel": [],
        "last": None,
    }
    assert (started.status_code, started.json()) == (409, {"message": "Najpierw wskaż INTERMAX."})
    assert centrum.client.get("/api/runs/current").json()["state"] == "idle"


def test_karta_mowi_ze_intermax_jest_otwarty_w_excelu_az_do_jego_zamkniecia(tmp_path: Path) -> None:
    centrum = ordering_centrum(tmp_path)

    lock = open_in_excel(centrum.register)
    while_open = centrum.setup()
    lock.unlink()
    closed = centrum.setup()

    assert (while_open["blocker"], while_open["open_in_excel"]) == (None, [OPEN])
    assert closed["open_in_excel"] == []


def test_uporzadkuj_porzadkuje_intermax_z_ustawien_taty_i_pokazuje_liczby_dzien_i_kopie(tmp_path: Path) -> None:
    centrum = ordering_centrum(tmp_path)
    original = centrum.register.read_bytes()

    started = centrum.start()
    views = views_until(centrum.client, "finished")
    finished = views[-1]

    assert started.status_code == 200
    assert started.json()["state"] == "running"
    run = finished["run"]
    assert (run["process"], run["title"], run["period"], run["dry_run"], run["confirms_readiness"]) == (
        "intermax",
        "Uporządkowanie INTERMAX",
        PERIOD,
        False,
        False,
    )
    result = finished["result"]
    assert (result["result"], result["ending"], result["headline"], result["next_action"]) == (
        "OK",
        "finished",
        "INTERMAX uporządkowany",
        "open_order_register",
    )
    assert result["summary"] == SUMMARY
    assert (result["stops"], result["warnings"], result["files_unchanged"]) == ([], [], False)
    assert phase_states(finished) == dict.fromkeys(("Sprawdzanie", "Odczyt danych", "Planowanie", "Zapisywanie"), "completed")
    assert copies(centrum) == [COPY]
    assert (centrum.copies / COPY).read_bytes() == original
    assert centrum.register.read_bytes() != original


def test_centrum_bez_wskazanego_folderu_kopii_kladzie_kopie_intermax_w_folderze_narzedzia(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    tool = tmp_path / "narzędzie"
    tool.mkdir()
    monkeypatch.setattr(intermax_copies, "TOOL_ROOT", tool)
    centrum = ordering_centrum(tmp_path, tool_copies=True)
    original = centrum.register.read_bytes()

    view = centrum.run_to_end()

    assert view["result"]["headline"] == "INTERMAX uporządkowany"
    assert (tool / "Kopie INTERMAX" / COPY).read_bytes() == original


def test_po_chwilowej_odmowie_onedrive_uporzadkuj_z_centrum_ponawia_podmiane_na_zegarze_centrum(tmp_path: Path) -> None:
    centrum = ordering_centrum(tmp_path)
    centrum.disk.replace, attempts = failing_replace(denied, times=2)
    started_at = centrum.surroundings.clock.seconds

    view = centrum.run_to_end()

    assert (view["result"]["ending"], view["result"]["headline"]) == ("finished", "INTERMAX uporządkowany")
    assert (len(attempts), centrum.surroundings.clock.seconds - started_at) == (3, 1.0)
    assert copies(centrum) == [COPY]


def test_gdy_onedrive_dalej_trzyma_intermax_centrum_mowi_tacie_o_synchronizacji_po_pieciu_sekundach_zegara(
    tmp_path: Path,
) -> None:
    centrum = ordering_centrum(tmp_path)
    original = centrum.register.read_bytes()
    centrum.disk.replace, attempts = failing_replace(denied)
    started_at = centrum.surroundings.clock.seconds

    view = _stopped(centrum)

    result = view["result"]
    assert (result["stop_kind"], result["stops"][0]["what"]) == (
        "for_tata",
        "OneDrive jeszcze synchronizuje INTERMAX. INTERMAX nie został zmieniony.",
    )
    assert (len(attempts), centrum.surroundings.clock.seconds - started_at) == (11, 5.0)
    assert centrum.register.read_bytes() == original
    assert copies(centrum) == []


def test_okres_uruchomienia_to_okres_dnia_startu_takze_gdy_skonczy_sie_po_polnocy(tmp_path: Path) -> None:
    # Poniedziałek tuż przed północą: następne pytania zegara są już we wtorek, w następnym okresie.
    clock = Clock([datetime(2026, 10, 26, 23, 59, 59), datetime(2026, 10, 27, 0, 0, 3)])
    centrum = ordering_centrum(tmp_path, clock=clock)

    view = centrum.run_to_end()

    assert view["run"]["period"] == PERIOD
    assert [record["period"] for record in centrum.metric_records()] == [PERIOD]
    assert [record["period"] for record in centrum.run_logs()] == [PERIOD]
    assert centrum.last_started()["period"] == PERIOD
    assert copies(centrum) == ["INTERMAX przed uporządkowaniem 2026-10-27 00-00-03.xlsx"]


def test_intermax_otwarty_to_zatrzymanie_dla_taty_przed_odczytem_w_okresie_dzisiejszej_daty(tmp_path: Path) -> None:
    centrum = ordering_centrum(tmp_path)
    original = centrum.register.read_bytes()
    open_in_excel(centrum.register)

    view = _stopped(centrum)

    result = view["result"]
    assert (result["result"], result["stop_kind"], result["headline"], result["next_action"]) == (
        "BLAD_KRYTYCZNY",
        "for_tata",
        "Do poprawy",
        "rerun",
    )
    assert result["stops"][0]["what"] == "INTERMAX jest otwarty w Excelu. INTERMAX nie został zmieniony."
    assert result["stops"][0]["where"] == f"plik {REGISTER_NAME}, folder OneDrive"
    assert (result["opens"], result["diagnosis"], result["files_unchanged"]) == ("order_register", None, True)
    # Karta proponuje otworzyć Rejestr zleceń pod nazwą z Komunikatu: „Otwórz INTERMAX”.
    assert view["run"]["register_name"] == "INTERMAX"
    assert phase_states(view)["Sprawdzanie"] == "failed"
    assert view["run"]["period"] == PERIOD
    assert centrum.last_started()["period"] == PERIOD
    assert [(record["period"], record["stop"]["code"]) for record in centrum.metric_records()] == [
        (PERIOD, "INTERMAX_OTWARTY")
    ]
    assert centrum.register.read_bytes() == original
    assert copies(centrum) == []


def test_logi_zatrzymania_przed_odczytem_ida_do_skrzynki_przy_okresie_dzisiejszej_daty(tmp_path: Path) -> None:
    centrum = ordering_centrum(tmp_path)
    (tmp_path / "config" / "zgloszenia-klucz.txt").write_text("github_pat_SYNTETYCZNY123\n", encoding="utf-8")
    open_in_excel(centrum.register)

    run_id = _stopped(centrum)["run"]["id"]

    [logs] = centrum.run_logs()
    assert (logs["run_id"], logs["process"], logs["mode"], logs["period"]) == (run_id, "intermax", "zapis", PERIOD)
    assert [message["code"] for message in logs["messages"]] == ["INTERMAX_OTWARTY"]
    assert [file["name"] for file in logs["environment"]["files"]] == [REGISTER_NAME]
    box = centrum.surroundings.report_box
    line = f"- Logi Uruchomienia okresu: {PERIOD}"
    assert box.wait_until(lambda box: any(line in issue.body and issue.comments for issue in box.background_issues))
    [issue] = box.background_issues
    assert issue.title == f"Logi Uruchomienia {PERIOD}"
    assert run_id in issue.comments[0]


def test_zatrzymanie_do_maksa_prowadzi_do_diagnozy_intermax_i_zgloszenia(tmp_path: Path) -> None:
    centrum = ordering_centrum(tmp_path)
    workbook = load_workbook(centrum.register)
    workbook[ORDER_SHEET].cell(centrum.layout.first_raw + 1, 4).value = "=1+1"
    workbook.save(centrum.register)

    view = _stopped(centrum)
    diagnosis = centrum.client.post(f"/api/runs/{view['run']['id']}/diagnosis")
    report = centrum.client.post(f"/api/runs/{view['run']['id']}/report")

    result = view["result"]
    assert (result["stop_kind"], result["headline"], result["diagnosis"]) == ("for_maks", "Błąd", "on_request")
    assert result["stops"][0]["where"] == f"plik {REGISTER_NAME}, folder OneDrive, arkusz {ORDER_SHEET}"
    assert diagnosis.status_code == 200
    checks = {check["label"]: (check["state"], check["detail"]) for check in diagnosis.json()["checks"]}
    assert list(checks) == [
        "Ustawienia taty",
        "Pliki w OneDrive",
        "INTERMAX",
        "Aktualność narzędzia",
        "Token do Zgłoszeń",
        "Dziennik Centrum rozliczeń",
    ]
    assert checks["Ustawienia taty"] == ("passed", "INTERMAX jest na miejscu.")
    assert checks["Pliki w OneDrive"] == ("passed", "Sprawdzone pliki są na tym komputerze: 1.")
    assert checks["INTERMAX"] == ("passed", "Nie jest otwarty w Excelu, otwiera się i ma arkusz zleceń.")
    assert diagnosis.json()["next_action"] == "report"
    assert "- Proces: Uporządkowanie INTERMAX" in diagnosis.json()["report"]["body"]
    assert report.status_code == 200
    assert report.json()["state"] in ("sent", "queued")


def test_powtorzone_zatrzymanie_intermax_wyjasnia_diagnoza_intermax(tmp_path: Path) -> None:
    centrum = ordering_centrum(tmp_path)
    open_in_excel(centrum.register)

    _stopped(centrum)
    repeated = _stopped(centrum)
    diagnosis = centrum.client.post(f"/api/runs/{repeated['run']['id']}/diagnosis").json()

    assert (repeated["result"]["stop_kind"], repeated["result"]["diagnosis"]) == ("repeated", "on_request")
    assert diagnosis["confirms_stop"] is True
    assert diagnosis["explanation"] == repeated["result"]["stops"][0]
    assert diagnosis["next_action"] == "rerun"
    assert {check["label"]: check["state"] for check in diagnosis["checks"]}["INTERMAX"] == "failed"


def test_diagnoza_mowi_ktorego_intermax_z_ustawien_taty_nie_ma(tmp_path: Path) -> None:
    centrum = ordering_centrum(tmp_path, Intermax())
    centrum.engines.fail_summary = True
    view = centrum.run_to_end()
    centrum.register.unlink()

    diagnosis = centrum.client.post(f"/api/runs/{view['run']['id']}/diagnosis").json()

    assert view["result"]["stop_kind"] == "for_maks"
    assert diagnosis["explanation"]["what"] == "Nie znaleziono INTERMAX. INTERMAX nie został zmieniony."
    assert diagnosis["next_action"] == "rerun"
    states = {check["label"]: check["state"] for check in diagnosis["checks"]}
    assert (states["Ustawienia taty"], states["INTERMAX"]) == ("failed", "not_checked")


@pytest.mark.parametrize(
    ("broken", "what"),
    [
        (lambda register: register.write_bytes(b"to nie jest skoroszyt"), "Nie można odczytać INTERMAX jako skoroszytu Excela. INTERMAX nie został zmieniony."),
        (
            lambda register: Workbook().save(register),
            "W INTERMAX nie ma arkusza zleceń z nagłówkami od AKCJE do UWAGI w wierszu 1. INTERMAX nie został zmieniony.",
        ),
    ],
    ids=["uszkodzony", "bez-arkusza-zlecen"],
)
def test_diagnoza_odczytuje_intermax_i_wskazuje_ten_ktorego_nie_da_sie_uporzadkowac(
    tmp_path: Path, broken: Callable[[Path], None], what: str
) -> None:
    centrum = ordering_centrum(tmp_path)
    centrum.engines.fail_summary = True
    view = centrum.run_to_end()
    broken(centrum.register)

    diagnosis = centrum.client.post(f"/api/runs/{view['run']['id']}/diagnosis").json()

    [intermax] = [check for check in diagnosis["checks"] if check["label"] == "INTERMAX"]
    assert (intermax["state"], intermax["detail"]) == ("failed", what)
    assert diagnosis["explanation"]["what"] == what


def test_diagnoza_intermax_sprawdza_w_onedrive_tylko_intermax_a_nie_pliki_okresu(tmp_path: Path) -> None:
    periods = tmp_path / "INEA"
    collective = periods / PERIOD / f"Rozliczenie {PERIOD} - zbiorcze.xlsx"
    collective.parent.mkdir(parents=True)
    collective.write_bytes(b"tylko w chmurze")
    centrum = ordering_centrum(tmp_path, periods_directory=periods)
    assert centrum.surroundings.onedrive is not None
    centrum.surroundings.onedrive.online_only.add(collective)
    centrum.engines.fail_summary = True
    view = centrum.run_to_end()

    diagnosis = centrum.client.post(f"/api/runs/{view['run']['id']}/diagnosis").json()

    checks = {check["label"]: (check["state"], check["detail"]) for check in diagnosis["checks"]}
    assert checks["Pliki w OneDrive"] == ("passed", "Sprawdzone pliki są na tym komputerze: 1.")


def test_historia_metryki_i_karta_maja_uporzadkowanie_intermax(tmp_path: Path) -> None:
    centrum = ordering_centrum(tmp_path)
    run_id = centrum.run_to_end()["run"]["id"]

    metrics = centrum.client.get("/api/metrics").json()
    card = centrum.client.get(f"/api/runs/{run_id}/card").json()

    [run] = metrics["runs"]
    assert (run["process"], run["process_label"], run["period"], run["mode"]) == (
        "intermax",
        "Uporządkowanie INTERMAX",
        PERIOD,
        "RUN",
    )
    assert (run["result"], run["origin"], run["run_id"]) == ("OK", "centrum", run_id)
    # Uporządkowanie INTERMAX nie ma Podglądu bez zapisu, więc Metryki nie mają jego trybu DRY-RUN.
    assert [(entry["process"], entry["mode"]) for entry in metrics["statistics"]] == [
        ("preparation", "RUN"),
        ("preparation", "DRY-RUN"),
        ("settlement", "RUN"),
        ("settlement", "DRY-RUN"),
        ("intermax", "RUN"),
    ]
    assert card["state"] == "restored"
    assert (card["run"]["title"], card["result"]["headline"]) == ("Uporządkowanie INTERMAX", "INTERMAX uporządkowany")
    assert card["result"]["summary"] == SUMMARY
    pinnable = centrum.client.get("/api/dad-notes/runs").json()
    assert [(entry["run_id"], entry["process_label"]) for entry in pinnable] == [(run_id, "Uporządkowanie INTERMAX")]


def test_karta_pokazuje_ostatnie_uporzadkowanie_z_logow_takze_po_ponownym_starcie_centrum(tmp_path: Path) -> None:
    centrum = ordering_centrum(tmp_path)
    ordered = centrum.run_to_end()["run"]["id"]
    after_ordering = centrum.setup()["last"]
    open_in_excel(centrum.register)
    stopped = _stopped(centrum)["run"]["id"]

    centrum.restart()
    last = centrum.setup()["last"]
    previous = centrum.client.get("/api/runs/current").json()["previous"]

    assert (after_ordering["run"]["id"], after_ordering["result"]["summary"]) == (ordered, SUMMARY)
    assert (last["run"]["id"], last["result"]["headline"], last["result"]["stops"][0]["what"]) == (
        stopped,
        "Do poprawy",
        "INTERMAX jest otwarty w Excelu. INTERMAX nie został zmieniony.",
    )
    assert previous == last
    assert last["run"]["register_name"] == "INTERMAX"


def _logs_like(centrum: OrderingCentrum, like: str, later: timedelta, **changes: Any) -> Path:
    """Logi drugiego Uruchomienia jak ``like``, rozpoczętego ``later`` po pełnej sekundzie jego startu.

    Start trafia do nazwy pliku i do Logów; ``changes`` zmienia inne pola Logów.
    """

    [path] = (centrum.tmp_path / "logs" / "uruchomienia").glob(f"*-{like}.json")
    record = RunLog.model_validate_json(path.read_text(encoding="utf-8"))
    started_at = record.started_at.replace(microsecond=0) + later
    times = {"started_at": started_at, "ended_at": started_at + (record.ended_at - record.started_at)}
    other = record.model_copy(update={"run_id": uuid4().hex, **times, **changes})
    other_path = path.with_name(other.file_name)
    other_path.write_text(other.model_dump_json(indent=2), encoding="utf-8")
    return other_path


def test_karta_pokazuje_najnowsze_uporzadkowanie_a_nie_starsze_z_logow(tmp_path: Path) -> None:
    centrum = ordering_centrum(tmp_path)
    ordered = centrum.run_to_end()["run"]["id"]
    _logs_like(centrum, ordered, -timedelta(days=7))

    assert centrum.setup()["last"]["run"]["id"] == ordered


def test_karta_bierze_ostatnie_uporzadkowanie_a_nie_pozniejsze_uruchomienie_innego_procesu(tmp_path: Path) -> None:
    centrum = ordering_centrum(tmp_path)
    ordered = centrum.run_to_end()["run"]["id"]
    _logs_like(centrum, ordered, timedelta(hours=1), process="preparation")

    assert centrum.setup()["last"]["run"]["id"] == ordered


def test_bez_logow_uporzadkowania_karta_nie_ma_ostatniego_choc_sa_logi_innego_procesu(tmp_path: Path) -> None:
    centrum = ordering_centrum(tmp_path)
    ordered = centrum.run_to_end()["run"]["id"]
    _logs_like(centrum, ordered, timedelta(hours=1), process="preparation")
    [intermax_logs] = (centrum.tmp_path / "logs" / "uruchomienia").glob(f"*-{ordered}.json")
    intermax_logs.unlink()
    centrum.restart()

    setup = centrum.client.get("/api/intermax")

    assert (setup.status_code, setup.json()["last"]) == (200, None)


@pytest.mark.parametrize(
    ("earlier", "later"), [("0" * 32, "f" * 32), ("f" * 32, "0" * 32)], ids=["pozniejszy-dalej", "pozniejszy-blizej"]
)
def test_z_uporzadkowan_w_tej_samej_sekundzie_karta_pokazuje_pozniejsze(tmp_path: Path, earlier: str, later: str) -> None:
    centrum = ordering_centrum(tmp_path)
    ordered = centrum.run_to_end()["run"]["id"]
    _logs_like(centrum, ordered, timedelta(hours=1, milliseconds=100), run_id=earlier)
    _logs_like(centrum, ordered, timedelta(hours=1, milliseconds=200), run_id=later)

    assert centrum.setup()["last"]["run"]["id"] == later


def test_karta_pomija_nieczytelne_logi_uporzadkowania_z_tej_samej_sekundy(tmp_path: Path) -> None:
    centrum = ordering_centrum(tmp_path)
    ordered = centrum.run_to_end()["run"]["id"]
    _logs_like(centrum, ordered, timedelta(0)).write_text("{", encoding="utf-8")

    assert centrum.setup()["last"]["run"]["id"] == ordered


def test_nieczytelne_najnowsze_logi_uporzadkowania_to_karta_bez_ostatniego_a_nie_starsze(tmp_path: Path) -> None:
    centrum = ordering_centrum(tmp_path)
    ordered = centrum.run_to_end()["run"]["id"]
    _logs_like(centrum, ordered, timedelta(hours=1)).write_text("{", encoding="utf-8")

    assert centrum.setup()["last"] is None


def test_nieczytelne_najnowsze_logi_to_karta_bez_ostatniego_takze_gdy_starsze_sa_z_tej_samej_minuty(tmp_path: Path) -> None:
    centrum = ordering_centrum(tmp_path)
    ordered = centrum.run_to_end()["run"]["id"]
    [path] = (centrum.tmp_path / "logs" / "uruchomienia").glob(f"*-{ordered}.json")
    second = RunLog.model_validate_json(path.read_text(encoding="utf-8")).started_at.second
    # Godzinę później: czytelne Logi w pełnej minucie, a nieczytelne najnowsze pół minuty po nich.
    _logs_like(centrum, ordered, timedelta(hours=1, seconds=-second))
    _logs_like(centrum, ordered, timedelta(hours=1, seconds=30 - second)).write_text("{", encoding="utf-8")

    assert centrum.setup()["last"] is None


def test_logi_uporzadkowania_maja_wiersze_wpisow_na_operatora_i_instalacje_hurt_bez_wpisu_a_karta_sie_nie_zmienia(
    tmp_path: Path,
) -> None:
    raw = (
        Order(datetime(2026, 10, 26, 8), "jan.nowak", 201, "P4", "Instalacja Hurt"),
        Order(datetime(2026, 10, 26, 9), "jan.nowak", 202, "T-Mobile", "Instalacja Hurt"),
        Order(datetime(2026, 10, 26, 10), "jan.nowak", 203, "P4", "Instalacja Hurt"),
        Order(datetime(2026, 10, 26, 11), "jan.nowak", 204, "INEA S.A.", "Instalacja Hurt"),
        Order(datetime(2026, 10, 26, 12), "jan.nowak", 205, None, "Instalacja Hurt"),
        Order(datetime(2026, 10, 26, 13), "jan.nowak", 206, "P4", "Usterka"),
    )
    centrum = ordering_centrum(tmp_path, Intermax(raw=raw))
    day = centrum.layout.yellow
    assert day is not None

    view = centrum.run_to_end()

    result = view["result"]
    assert (result["headline"], result["stops"], result["warnings"]) == ("INTERMAX uporządkowany", [], [])
    assert result["summary"] == SUMMARY | {
        "figures": [
            {"label": "Zlecenia", "value": 6, "unit": "count"},
            {"label": "Wykonawcy", "value": 1, "unit": "count"},
        ],
    }
    [logs] = centrum.run_logs()
    assert logs["intermax_entries"] == {
        "rows": {"P4": [day + 2, day + 4], "TMOBILE": [day + 3]},
        "without_entry": [{"row": day + 5, "operator": "INEA S.A."}, {"row": day + 6, "operator": None}],
    }
    [path] = (tmp_path / "logs" / "uruchomienia").glob("*.json")
    text = path.read_text(encoding="utf-8")
    for order in raw:
        for column, value in order.values().items():
            if isinstance(value, str) and column not in (DEVICES, ACTIVITIES) and value not in ("Miasto", "Umowa"):
                assert value not in text, value
    assert "akt pętli" not in text and "nokia" not in text, "bez wartości Wpisu"


def test_karta_wraca_z_logow_uporzadkowania_sprzed_wpisow_intermax(tmp_path: Path) -> None:
    centrum = ordering_centrum(tmp_path)
    ordered = centrum.run_to_end()["run"]["id"]
    [path] = (tmp_path / "logs" / "uruchomienia").glob(f"*-{ordered}.json")
    record = json.loads(path.read_text(encoding="utf-8"))
    del record["intermax_entries"]
    path.write_text(json.dumps(record, ensure_ascii=False), encoding="utf-8")

    centrum.restart()

    last = centrum.setup()["last"]
    assert (last["run"]["id"], last["result"]["summary"]) == (ordered, SUMMARY)


def test_start_z_okresem_nie_przyjmuje_uporzadkowania_intermax(tmp_path: Path) -> None:
    centrum = ordering_centrum(tmp_path)

    refused = centrum.client.post("/api/runs", json={"process": "intermax", "period": PERIOD, "dry_run": False})

    assert refused.status_code == 422
    assert copies(centrum) == []
