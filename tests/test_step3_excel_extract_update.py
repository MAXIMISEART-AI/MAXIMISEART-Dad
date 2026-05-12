"""Step 3 Excel buffer append tests."""
from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from scripts.step3_excel_extract_update import ExcelLockedError, append_rows


BUFFER_SHEET = "MAXIMISEART_DAILY_APPEND"
HEADERS = ["Data", "Pracownik", "Kod", "Ilość", "Lokacja", "Uwagi", "Status", "Email ID", "Timestamp"]


@pytest.fixture
def workbook_path(tmp_path: Path) -> Path:
    path = tmp_path / "tata.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Istniejacy"
    ws["A1"] = "Nie dotykac"
    ws["B2"] = "=SUM(1,2)"
    wb.save(path)
    return path


@pytest.fixture
def code_mapping() -> dict:
    return {
        "codes": [
            {
                "code": 1,
                "name": "Instalacja FTTH",
                "keywords": ["instalacja", "ftth"],
                "active": True,
            },
            {
                "code": 2,
                "name": "Spaw",
                "keywords": ["spaw", "mufa"],
                "active": True,
            },
        ]
    }


@pytest.fixture
def employees_config() -> dict:
    return {
        "employees": [
            {
                "slug": "przyklad-jan-kowalski",
                "name": "Jan Kowalski",
                "email": "jan.kowalski@example.com",
                "active": True,
            }
        ]
    }


def test_append_rows_writes_only_safe_buffer_sheet(
    workbook_path: Path,
    sample_emails: list[dict],
    code_mapping: dict,
    employees_config: dict,
    tmp_path: Path,
) -> None:
    routed = {"przyklad-jan-kowalski": sample_emails[:2]}

    rows = append_rows(
        workbook_path,
        routed,
        code_mapping,
        {"retry_on_lock": [0]},
        snapshots_dir=tmp_path / "snapshots",
        employees_config=employees_config,
        date="2026-04-24",
    )

    assert len(rows) == 3
    assert [row["Kod"] for row in rows] == [1, 2, None]
    assert [row["Status"] for row in rows] == ["ANSWER", "ANSWER", "ASK"]

    wb = load_workbook(workbook_path, data_only=False)
    assert "Istniejacy" in wb.sheetnames
    assert BUFFER_SHEET in wb.sheetnames
    assert wb["Istniejacy"]["A1"].value == "Nie dotykac"
    assert wb["Istniejacy"]["B2"].value == "=SUM(1,2)"

    ws = wb[BUFFER_SHEET]
    assert [cell.value for cell in ws[1]] == HEADERS
    assert ws.max_row == 4
    assert ws["A2"].value == "2026-04-24"
    assert ws["B2"].value == "Jan Kowalski"
    assert ws["C2"].value == 1
    assert ws["D2"].value == 3
    assert "Krańcowej 12" in ws["E2"].value
    assert ws["H2"].value == "18f2a3b4c5d6e7f8"


def test_append_rows_takes_snapshot_before_write(
    workbook_path: Path,
    sample_emails: list[dict],
    code_mapping: dict,
    employees_config: dict,
    tmp_path: Path,
) -> None:
    snapshots_dir = tmp_path / "snapshots"

    append_rows(
        workbook_path,
        {"przyklad-jan-kowalski": sample_emails[:1]},
        code_mapping,
        {"retry_on_lock": [0]},
        snapshots_dir=snapshots_dir,
        employees_config=employees_config,
        date="2026-04-24",
    )

    snapshots = list(snapshots_dir.glob("tata-*.xlsx"))
    assert len(snapshots) == 1
    snapshot_wb = load_workbook(snapshots[0], data_only=False)
    assert BUFFER_SHEET not in snapshot_wb.sheetnames
    assert snapshot_wb["Istniejacy"]["A1"].value == "Nie dotykac"


def test_append_rows_aborts_when_excel_lock_file_exists(
    workbook_path: Path,
    sample_emails: list[dict],
    code_mapping: dict,
    employees_config: dict,
    tmp_path: Path,
) -> None:
    lock_file = workbook_path.parent / f"~${workbook_path.name}"
    lock_file.write_text("locked", encoding="utf-8")

    with pytest.raises(ExcelLockedError):
        append_rows(
            workbook_path,
            {"przyklad-jan-kowalski": sample_emails[:1]},
            code_mapping,
            {"retry_on_lock": [0]},
            snapshots_dir=tmp_path / "snapshots",
            employees_config=employees_config,
            date="2026-04-24",
        )

    wb = load_workbook(workbook_path, data_only=False)
    assert BUFFER_SHEET not in wb.sheetnames
    assert not (tmp_path / "snapshots").exists()
