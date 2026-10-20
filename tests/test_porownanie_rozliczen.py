"""Macierz reguł porównania znaczeniowego dwóch Folderów rozliczeń pracowników (bramka poprawności, #144)."""

from __future__ import annotations

import json
import shutil
from collections.abc import Callable
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.worksheet import Worksheet

from tests.fixtures.porownanie import (
    CATEGORIES,
    SECRET_VALUES,
    SETTLEMENT_NAME,
    add_calc_chain,
    compare_worker_settlements,
    formulas_in_row_13,
    renumber_cell_styles,
    settlement_folders,
    share_formulas_in_row_13,
    write_settlement,
)
from tests.fixtures.workbooks import rewrite_parts

NO_DIFFERENCES = dict.fromkeys(CATEGORIES, 0)


def _only(category: str, count: int = 1) -> dict[str, int]:
    return NO_DIFFERENCES | {category: count}


def _set(coordinate: str, attribute: str, value: object) -> Callable[[Workbook, Worksheet], None]:
    def change(_workbook: Workbook, sheet: Worksheet) -> None:
        setattr(sheet[coordinate], attribute, value)

    return change


def _change_workbook(change: Callable[[Workbook, Worksheet], None]) -> Callable[[Workbook, Worksheet], None]:
    return change


def test_identyczne_foldery_rozliczen_pracownikow_nie_maja_roznic(tmp_path: Path) -> None:
    reference, candidate = settlement_folders(tmp_path)

    assert compare_worker_settlements(reference, candidate) == NO_DIFFERENCES


@pytest.mark.parametrize(
    ("change", "expected"),
    [
        pytest.param(_set("N18", "value", 4), _only("wartosci"), id="wartosc"),
        pytest.param(_set("B19", "value", "ul. Inna 3"), _only("wartosci"), id="tekst"),
        pytest.param(_set("N18", "value", True), _only("wartosci"), id="prawda-to-nie-liczba-3"),
        pytest.param(_set("AU18", "value", "=SUM(N18:AO18)"), _only("formuly"), id="formula"),
        pytest.param(_set("AU18", "value", 3), _only("formuly"), id="formula-zastapiona-wartoscia"),
        pytest.param(_set("N18", "font", Font(bold=True)), _only("style"), id="czcionka"),
        pytest.param(
            _set("N18", "fill", PatternFill(fill_type="solid", fgColor="FF0000")), _only("style"), id="wypelnienie"
        ),
        pytest.param(_set("N18", "border", Border(left=Side(style="thick"))), _only("style"), id="obramowanie"),
        pytest.param(_set("N18", "alignment", Alignment(horizontal="left")), _only("style"), id="wyrownanie"),
        pytest.param(_set("N18", "number_format", "0.00"), _only("style"), id="format-liczb"),
        pytest.param(_set("N18", "protection", Protection(locked=True)), _only("style"), id="ochrona"),
        pytest.param(
            _change_workbook(lambda _workbook, sheet: sheet.unmerge_cells("A16:M16")), _only("scalenia"), id="scalenie"
        ),
        pytest.param(
            _change_workbook(lambda _workbook, sheet: setattr(sheet.row_dimensions[17], "height", 20)),
            _only("wymiary"),
            id="wysokosc-wiersza",
        ),
        pytest.param(
            _change_workbook(lambda _workbook, sheet: setattr(sheet.row_dimensions[5], "hidden", True)),
            _only("wymiary"),
            id="ukryty-wiersz",
        ),
        pytest.param(
            _change_workbook(lambda _workbook, sheet: setattr(sheet.column_dimensions["B"], "width", 12)),
            _only("wymiary"),
            id="szerokosc-kolumny",
        ),
        pytest.param(
            _change_workbook(
                lambda workbook, _sheet: workbook.defined_names.__setitem__(
                    "Stawki", DefinedName("Stawki", attr_text="Sheet1!$N$2:$AP$2")
                )
            ),
            _only("nazwy_zdefiniowane"),
            id="nazwa-zdefiniowana",
        ),
        pytest.param(
            _change_workbook(lambda workbook, _sheet: setattr(workbook.calculation, "forceFullCalc", False)),
            _only("flagi_przeliczania"),
            id="flaga-przeliczania",
        ),
    ],
)
def test_bramka_wykrywa_kazda_kategorie_roznicy_w_rozliczeniu_pracownika(
    tmp_path: Path, change: Callable[[Workbook, Worksheet], None], expected: dict[str, int]
) -> None:
    reference, candidate = settlement_folders(tmp_path)
    write_settlement(candidate / SETTLEMENT_NAME, change)

    assert compare_worker_settlements(reference, candidate) == expected


def test_bramka_wykrywa_zmienione_bajty_odwolania_zewnetrznego(tmp_path: Path) -> None:
    reference, candidate = settlement_folders(tmp_path)
    rewrite_parts(
        candidate / SETTLEMENT_NAME,
        lambda name, data: data.replace(b"S\xc5\x82ownik", b"Slownik") if name.startswith("xl/externalLinks/") else data,
    )

    assert compare_worker_settlements(reference, candidate) == _only("external_links")


def test_bramka_wykrywa_brakujacy_nadmiarowy_i_nieczytelny_plik(tmp_path: Path) -> None:
    reference, candidate = settlement_folders(tmp_path)
    write_settlement(reference / "Rozliczenie 08_14_09_2026 - Darek Fikcyjny.xlsx")
    shutil.copy2(candidate / SETTLEMENT_NAME, candidate / "Rozliczenie 08_14_09_2026 - Kamil Syntetyczny.xlsx")
    write_settlement(reference / "Rozliczenie 08_14_09_2026 - Tomasz Lis.xlsx")
    (candidate / "Rozliczenie 08_14_09_2026 - Tomasz Lis.xlsx").write_bytes(b"to nie jest skoroszyt")

    assert compare_worker_settlements(reference, candidate) == _only("pliki", 3)


def test_wspolna_i_rozpisana_formula_o_tej_samej_tresci_sa_zgodne(tmp_path: Path) -> None:
    reference, candidate = settlement_folders(tmp_path)
    write_settlement(reference / SETTLEMENT_NAME, formulas_in_row_13)
    share_formulas_in_row_13(reference / SETTLEMENT_NAME)
    write_settlement(candidate / SETTLEMENT_NAME, formulas_in_row_13)

    assert compare_worker_settlements(reference, candidate) == NO_DIFFERENCES

    def other_formula(workbook: Workbook, sheet: Worksheet) -> None:
        formulas_in_row_13(workbook, sheet)
        sheet["P13"] = "=P14+P15"

    write_settlement(candidate / SETTLEMENT_NAME, other_formula)
    assert compare_worker_settlements(reference, candidate) == _only("formuly")


def test_bramka_nie_porownuje_calc_chain(tmp_path: Path) -> None:
    reference, candidate = settlement_folders(tmp_path)
    add_calc_chain(candidate / SETTLEMENT_NAME)

    assert compare_worker_settlements(reference, candidate) == NO_DIFFERENCES


def test_ten_sam_styl_pod_innym_numerem_stylu_jest_zgodny(tmp_path: Path) -> None:
    reference, candidate = settlement_folders(tmp_path)
    renumber_cell_styles(candidate / SETTLEMENT_NAME)

    assert compare_worker_settlements(reference, candidate) == NO_DIFFERENCES


def test_komorka_nieobecna_w_xml_ma_styl_kolumny_jak_w_excelu(tmp_path: Path) -> None:
    def column_style(_workbook: Workbook, sheet: Worksheet) -> None:
        sheet.column_dimensions["C"].font = Font(italic=True)

    def explicit_cell(style: Font) -> Callable[[Workbook, Worksheet], None]:
        def change(workbook: Workbook, sheet: Worksheet) -> None:
            column_style(workbook, sheet)
            sheet["C20"].font = style

        return change

    reference, candidate = settlement_folders(tmp_path)
    write_settlement(reference / SETTLEMENT_NAME, column_style)
    write_settlement(candidate / SETTLEMENT_NAME, explicit_cell(Font(italic=True)))
    assert compare_worker_settlements(reference, candidate) == NO_DIFFERENCES

    write_settlement(candidate / SETTLEMENT_NAME, explicit_cell(Font()))
    assert compare_worker_settlements(reference, candidate) == _only("style")


def test_wynik_porownania_to_same_liczby_bez_tresci_komorek_adresow_i_numerow_zlecen(tmp_path: Path) -> None:
    def changed_rows(_workbook: Workbook, sheet: Worksheet) -> None:
        sheet["B18"] = "ul. Inna 9"
        sheet["F18"] = "ZL-INNE-9"

    reference, candidate = settlement_folders(tmp_path)
    write_settlement(candidate / SETTLEMENT_NAME, changed_rows)

    differences = compare_worker_settlements(reference, candidate)
    published = json.dumps(differences, ensure_ascii=False)

    assert differences == _only("wartosci", 2)
    assert all(type(count) is int for count in differences.values())
    for secret in (*SECRET_VALUES, "ul. Inna 9", "ZL-INNE-9", "B18", "F18", "Adrian"):
        assert secret not in published
