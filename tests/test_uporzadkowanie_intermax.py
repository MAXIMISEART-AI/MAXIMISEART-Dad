"""Uporządkowanie INTERMAX przez rdzeń Uruchomienia procesu: wynikowy INTERMAX z dysku, Kopie INTERMAX i Zatrzymania."""

from __future__ import annotations

import errno
import os
import re
import zipfile
from collections.abc import Callable, Iterator, Sequence
from contextlib import contextmanager
from dataclasses import dataclass, replace
from datetime import datetime, timedelta
from io import BytesIO
from pathlib import Path
from typing import Literal

import pytest
from openpyxl import load_workbook
from openpyxl.styles import Color, Font, PatternFill
from openpyxl.utils.cell import coordinate_from_string, get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from rozliczenia import TOOL_ROOT
from rozliczenia.domain import ProgressPhase
from rozliczenia.intermax_file import Waiting
from rozliczenia.intermax_ordering import OrderingSummary
from rozliczenia.intermax_process import PROCESS, OrderingInputs
from rozliczenia.message import Message
from rozliczenia.process_run import Finished, Outcome, Stop, run_process
from rozliczenia.progress import PhaseState

from tests.fixtures.clock import ManualClock
from tests.fixtures.intermax import (
    ACTIVITIES,
    APP_LAST,
    BLANK_HEIGHT,
    CONTRACT_NUMBER,
    DATE,
    DEVICES,
    DATE_FORMAT,
    FIRST,
    HEADERS,
    LAST,
    NUMBER,
    ORDER_HEIGHT,
    ORDER_SHEET,
    WORKER,
    YELLOW,
    Intermax,
    Layout,
    Order,
    order_rows,
    parts,
    rewrite_parts,
)
from tests.fixtures.interlocutor import Tata

NOW = datetime(2026, 10, 25, 21, 5, 7)
COPY_NAME = "INTERMAX przed uporządkowaniem 2026-10-25 21-05-07.xlsx"
REGISTER_NAME = "ZLECENIA INTERMAX OD 08_11_2021.xlsx"


def _run(tmp_path: Path, *, now: datetime = NOW, clock: ManualClock | None = None) -> Outcome[OrderingSummary]:
    """Uruchomienie Uporządkowania; ``clock`` to zegar procesu krótkiej ponownej próby podmiany."""

    waiting = clock or ManualClock()
    inputs = OrderingInputs(
        tmp_path / "OneDrive" / REGISTER_NAME, tmp_path / "Kopie INTERMAX", lambda: now, Waiting(waiting, waiting.pause)
    )
    return run_process(
        PROCESS, inputs, Tata(), view=lambda _inputs, summary: summary, metrics_path=tmp_path / "m.jsonl", clock=ManualClock()
    )


def _register(tmp_path: Path, intermax: Intermax | None = None) -> tuple[Path, Layout]:
    folder = tmp_path / "OneDrive"
    folder.mkdir(exist_ok=True)
    path = folder / REGISTER_NAME
    return path, (intermax or Intermax()).write(path)


def _expected(order: Order) -> tuple[object, ...]:
    values = order.values()
    return tuple(values.get(column) for column in range(1, APP_LAST + 1))


def _stop(outcome: Outcome[OrderingSummary]) -> Stop:
    assert isinstance(outcome.ending, Stop), outcome.ending
    return outcome.ending


def _without_order_data(message: Message, orders: Sequence[Order]) -> None:
    """Komunikat bez adresów, numerów zleceń, numerów umów, loginów Wykonawców i innych wartości zleceń."""

    for order in orders:
        for value in order.values().values():
            if isinstance(value, str) and value not in ("Miasto", "Umowa", "Instalacja", "INEA"):
                assert value not in message.text, value


def _raw_row(order: Order) -> Callable[[Worksheet, int], None]:
    """Wiersz ``order`` z wyglądem z aplikacji (Roboto 4) wpisany w wiersz arkusza."""

    def write(sheet: Worksheet, number: int) -> None:
        for column, value in order.values().items():
            sheet.cell(number, column, value).font = Font(name="Roboto", sz=4)

    return write


def _in_other_font(number: int) -> Callable[[Worksheet], None]:
    """Uporządkowany wiersz ``number`` przepisany czcionką inną niż Calibri 8, jak wiersz z wyglądem z aplikacji."""

    def edit(sheet: Worksheet) -> None:
        for column in range(FIRST, LAST + 1):
            sheet.cell(number, column).font = Font(name="Arial", sz=10)

    return edit


def _ordered_row(order: Order) -> Callable[[Worksheet, int], None]:
    """Wiersz ``order`` z wyglądem uporządkowanego zlecenia (Calibri 8) wpisany w wiersz arkusza."""

    def write(sheet: Worksheet, number: int) -> None:
        for column, value in order.values().items():
            cell = sheet.cell(number, column, value)
            cell.font = Font(name="Calibri", sz=8)
            if column == DATE:
                cell.number_format = DATE_FORMAT

    return write


def _ordered(outcome: Outcome[OrderingSummary]) -> OrderingSummary:
    assert isinstance(outcome.ending, Finished) and outcome.ending.result == "OK", outcome.ending
    assert outcome.view is not None
    return outcome.view


def _copies(tmp_path: Path) -> set[str]:
    folder = tmp_path / "Kopie INTERMAX"
    return {entry.name for entry in folder.iterdir()} if folder.is_dir() else set()


@contextmanager
def _nothing_written(tmp_path: Path, path: Path) -> Iterator[None]:
    """INTERMAX bajt w bajt, bez pliku tymczasowego obok i bez nowej Kopii INTERMAX po bloku."""

    original, neighbours, copies = path.read_bytes(), set(path.parent.iterdir()), _copies(tmp_path)
    yield
    assert path.read_bytes() == original
    assert set(path.parent.iterdir()) <= neighbours | {path.with_name(f"~${REGISTER_NAME}")}, "plik tymczasowy nie zostaje"
    assert _copies(tmp_path) == copies


def _edited(path: Path, edit: Callable[[Worksheet], None]) -> None:
    """INTERMAX zmieniony w openpyxl tak, jak tata zmienia go w Excelu: ``edit(arkusz zleceń)``."""

    workbook = load_workbook(path)
    edit(workbook[ORDER_SHEET])
    workbook.save(path)


def _row_xml(path: Path, number: int) -> bytes:
    found = re.search(rb'<row r="%d"[^>]*?(?:/>|>.*?</row>)' % number, parts(path)["xl/worksheets/sheet2.xml"], re.S)
    assert found is not None, number
    return found.group(0)


def test_dzien_staje_pod_zoltym_wierszem_w_segmentach_wykonawcow_alfabetycznie_i_po_godzinie(tmp_path: Path) -> None:
    path, layout = _register(tmp_path)
    anna_10, anna_10_later, anna_16 = layout.raw[1], layout.raw[4], layout.raw[2]
    bartek, jan_8, jan_14_shifted = layout.raw[5], layout.raw[3], layout.raw[0]

    outcome = _run(tmp_path)

    assert isinstance(outcome.ending, Finished) and outcome.ending.result == "OK"
    assert outcome.view is not None
    assert (outcome.view.day.isoformat(), outcome.view.order_count, outcome.view.worker_count) == ("2026-10-26", 6, 3)
    assert layout.yellow is not None
    assert order_rows(path, layout.yellow + 1, layout.yellow + 12) == [
        None,
        _expected(anna_10),
        _expected(anna_10_later),
        _expected(anna_16),
        None,
        _expected(bartek),
        None,
        _expected(jan_8),
        _expected(jan_14_shifted),
        None,
        None,
        None,
    ]


def test_bez_zoltego_wiersza_dzien_staje_10_pustych_wierszy_pod_ostatnim_uporzadkowanym_zleceniem(tmp_path: Path) -> None:
    path, layout = _register(tmp_path, Intermax(yellow=None, gap=1))

    _run(tmp_path)

    rows = order_rows(path, layout.last_ordered + 1, layout.last_ordered + 11)
    assert rows[:10] == [None] * 10
    assert rows[10] == _expected(layout.raw[1])
    sheet = load_workbook(path, read_only=True).worksheets[1]
    assert sheet.max_row == layout.last_ordered + 10 + 6 + 2, "<dimension> obejmuje dłuższy dzień"


def test_pierwszy_wiersz_na_miejscu_nie_jest_przesuwany(tmp_path: Path) -> None:
    path, layout = _register(tmp_path, Intermax(shifted=False))

    _run(tmp_path)

    assert layout.yellow is not None
    assert order_rows(path, layout.yellow + 9, layout.yellow + 9) == [_expected(layout.raw[0])]


def test_zlecenia_maja_wyglad_reszty_intermax_a_kolumny_za_uwagi_zostaja_jak_przyszly(tmp_path: Path) -> None:
    path, layout = _register(tmp_path)

    _run(tmp_path)

    assert layout.yellow is not None
    sheet = load_workbook(path).worksheets[1]
    for row in (layout.yellow + 2, layout.yellow + 9):
        for column in range(FIRST, LAST + 1):
            cell = sheet.cell(row, column)
            font = cell.font
            assert (font.name, font.sz, font.b, font.i, font.color.rgb) == (
                "Calibri",
                9 if column == WORKER else 8,
                False,
                False,
                "FF000000",
            ), cell.coordinate
            assert cell.fill.fill_type is None, cell.coordinate
            assert {side.style for side in (cell.border.left, cell.border.right, cell.border.top, cell.border.bottom)} == {
                "thin"
            }
            assert (cell.alignment.horizontal, cell.alignment.vertical) == ("left", "center")
        assert sheet.cell(row, DATE).number_format == "m/d/yy h:mm"
        assert sheet.cell(row, DATE + 1).number_format == "General"
        for column in range(LAST + 1, APP_LAST + 1):
            cell = sheet.cell(row, column)
            assert (cell.font.name, cell.font.sz, cell.font.b, cell.fill.fill_type) == ("Roboto", 4, True, "solid")


_GREEN, _ORANGE, _DARK_BLUE, _LIGHT_BLUE, _PINK = "FF99FF66", "FFFFC000", "FF69D8FF", "FF00FFFF", "FFE89CA7"


def _fills(sheet: Worksheet, row: int) -> dict[int, str | None]:
    """Wypełnienie RGB każdej komórki kolumn zlecenia wiersza ``row``; ``None``: bez wypełnienia."""

    return {
        column: cell.fill.fgColor.rgb if (cell := sheet.cell(row, column)).fill.fill_type == "solid" else None
        for column in range(FIRST, LAST + 1)
    }


def _colored(activities: str | None, contract_number: str | None, devices: str | None) -> dict[int, str | None]:
    return {column: None for column in range(FIRST, LAST + 1)} | {
        ACTIVITIES: activities,
        CONTRACT_NUMBER: contract_number,
        DEVICES: devices,
    }


@pytest.mark.parametrize(
    ("operator", "kind", "expected"),
    [
        ("TMOBILE", "Instalacja Hurt", _colored(_GREEN, _GREEN, None)),
        ("T-Mobile", "Instalacja", _colored(_GREEN, _GREEN, None)),
        (" t mobile ", "Instalacja", _colored(_GREEN, _GREEN, None)),
        ("ORANGE", "Instalacja Hurt", _colored(_ORANGE, _ORANGE, None)),
        ("Netia", "Instalacja Hurt", _colored(_DARK_BLUE, _LIGHT_BLUE, None)),
        ("P4", "Instalacja Hurt", _colored(_LIGHT_BLUE, _LIGHT_BLUE, None)),
        ("VECTRA", "Instalacja Hurt", _colored(_LIGHT_BLUE, _LIGHT_BLUE, None)),
        ("INEA", "Instalacja Hurt", _colored(None, None, None)),
        (None, "Instalacja", _colored(None, None, None)),
        ("TMOBILE", "Usterka", _colored(None, None, _PINK)),
        ("TMOBILE", "Usterka Hurt", _colored(None, None, _PINK)),
        ("ORANGE", "USTERKA  HURT", _colored(None, None, _PINK)),
        (None, "Usterka", _colored(None, None, _PINK)),
    ],
    ids=[
        "tmobile",
        "t-mobile",
        "t mobile ze spacjami",
        "orange",
        "netia",
        "p4",
        "vectra",
        "inea",
        "puste czynnosci",
        "usterka przy tmobile",
        "usterka hurt przy tmobile",
        "usterka hurt wielkimi literami",
        "usterka bez operatora",
    ],
)
def test_kolor_operatora_w_czynnosciach_numerze_umowy_i_urzadzeniach_a_inne_komorki_bez_wypelnienia(
    tmp_path: Path, operator: str | None, kind: str, expected: dict[int, str | None]
) -> None:
    raw = (
        Order(datetime(2026, 10, 26, 14), "jan.nowak", 101, operator, kind),
        Order(datetime(2026, 10, 26, 8), "jan.nowak", 102, operator, kind),
    )
    path, layout = _register(tmp_path, Intermax(raw=raw))

    _ordered(_run(tmp_path))

    assert layout.yellow is not None
    sheet = load_workbook(path).worksheets[1]
    assert order_rows(path, layout.yellow + 2, layout.yellow + 3) == [_expected(raw[1]), _expected(raw[0])]
    assert _fills(sheet, layout.yellow + 2) == expected, "zlecenie na miejscu"
    assert _fills(sheet, layout.yellow + 3) == expected, "przesunięty pierwszy wiersz"


def test_pusta_komorka_numeru_umowy_dostaje_kolor_operatora(tmp_path: Path) -> None:
    raw = (Order(datetime(2026, 10, 26, 8), "jan.nowak", 101, "TMOBILE"),)
    path, layout = _register(tmp_path, Intermax(raw=raw, shifted=False))
    contract = rb'<c r="%s%d"[^>]*?(?:/>|>.*?</c>)' % (get_column_letter(CONTRACT_NUMBER).encode(), layout.first_raw)
    rewrite_parts(path, lambda name, data: re.sub(contract, b"", data, flags=re.S) if name == _SHEET else data)

    _ordered(_run(tmp_path))

    assert layout.yellow is not None
    sheet = load_workbook(path).worksheets[1]
    assert sheet.cell(layout.yellow + 2, CONTRACT_NUMBER).value is None
    assert _fills(sheet, layout.yellow + 2) == _colored(_GREEN, _GREEN, None)


def _links(path: Path, first: int, last: int) -> dict[str, str | None]:
    sheet = load_workbook(path)[ORDER_SHEET]
    return {
        cell.coordinate: link.target
        for row in range(first, last + 1)
        if (link := (cell := sheet.cell(row, NUMBER)).hyperlink) is not None
    }


def test_inne_czesci_pliku_i_wiersze_nad_dniem_zostaja_bajt_w_bajt_a_hiperlacza_ida_ze_zleceniem(tmp_path: Path) -> None:
    """Arkusz zleceń był aktywny, więc poza nim i stylami nic się nie zmienia; nad dniem zmienia się tylko widok."""

    path, layout = _register(tmp_path)
    before = parts(path)
    ordered_links = _links(path, 2, layout.last_ordered)
    assert len(ordered_links) == 3

    _run(tmp_path)

    assert _links(path, 2, layout.last_ordered) == ordered_links

    after = parts(path)
    sheet_part, styles_part = "xl/worksheets/sheet2.xml", "xl/styles.xml"
    assert list(after) == list(before)
    assert {name for name in before if before[name] != after[name]} <= {sheet_part, styles_part}
    assert layout.yellow is not None
    region = next(row for row in re.finditer(rb'<row r="(\d+)"', before[sheet_part]) if int(row.group(1)) > layout.yellow)
    without_dimension_and_view = re.compile(rb"<dimension [^>]*>|<sheetViews>.*?</sheetViews>")
    above = without_dimension_and_view.sub(b"", before[sheet_part][: region.start()])
    assert b"<sheetData>" in above and without_dimension_and_view.sub(b"", after[sheet_part]).startswith(above)
    sheet = load_workbook(path).worksheets[1]
    links = {
        cell.value: link.target
        for row in range(layout.yellow + 2, layout.yellow + 10)
        if (link := (cell := sheet.cell(row, NUMBER)).hyperlink) is not None
    }
    assert links == {f"ZL-{order.number}": order.link for order in layout.raw}


@dataclass(frozen=True)
class _View:
    """Jak INTERMAX otwiera się w Excelu: aktywny arkusz, zaznaczone karty, górny wiersz widoku i zaznaczenie."""

    active: str
    selected: tuple[str, ...]
    top_row: int
    selection: str
    frozen_header: bool


def _view(path: Path) -> _View:
    workbook = load_workbook(path)
    sheet = workbook[ORDER_SHEET]
    view = sheet.sheet_view
    pane = view.pane
    top = pane.topLeftCell if pane is not None else view.topLeftCell
    [selection] = view.selection
    assert top is not None and selection.activeCell == selection.sqref
    assert selection.pane == (pane.activePane if pane is not None else None)
    assert workbook.active is not None
    return _View(
        workbook.active.title,
        tuple(sheet.title for sheet in workbook.worksheets if sheet.sheet_view.tabSelected),
        coordinate_from_string(top)[1],
        selection.activeCell,
        pane is not None and (pane.ySplit, pane.state, pane.activePane) == (1, "frozen", "bottomLeft"),
    )


# Uporządkowane dni fixture'u ``_VIEW_DAYS``: 19.10 w wierszu 2, 20.10 od wiersza 3 (dwóch Wykonawców, pusty wiersz 5).
_VIEW_DAYS = (
    Order(datetime(2026, 10, 19, 9), "anna.kowalska", 1),
    Order(datetime(2026, 10, 20, 9), "anna.kowalska", 2),
    Order(datetime(2026, 10, 20, 12), "anna.kowalska", 3),
    Order(datetime(2026, 10, 20, 8), "jan.nowak", 4),
)


@pytest.mark.parametrize(
    ("today", "top_row"),
    [(datetime(2026, 10, 20, 7), 3), (datetime(2026, 10, 19, 7), 2), (datetime(2026, 10, 23, 7), 3)],
    ids=["dzis-ze-zleceniami", "dzis-wczesniejszy-dzien", "dzis-bez-zlecen-ostatni-wczesniejszy-dzien"],
)
def test_intermax_otwiera_sie_na_pierwszym_zleceniu_dzisiejszego_dnia_choc_uporzadkowany_dzien_jest_inny(
    tmp_path: Path, today: datetime, top_row: int
) -> None:
    path, _ = _register(tmp_path, Intermax(ordered=_VIEW_DAYS))

    _ordered(_run(tmp_path, now=today))

    assert _view(path) == _View(ORDER_SHEET, (ORDER_SHEET,), top_row, f"B{top_row}", frozen_header=True)


@pytest.mark.parametrize(
    "today", [datetime(2026, 10, 26, 7), datetime(2026, 10, 28, 7), datetime(2026, 10, 1, 7)], ids=["dzis", "pozniej", "przed"]
)
def test_intermax_otwiera_sie_na_uporzadkowanym_dniu_gdy_to_dzis_ostatni_dzien_przed_dzis_albo_nie_ma_wczesniejszych(
    tmp_path: Path, today: datetime
) -> None:
    path, layout = _register(tmp_path, Intermax(ordered=_VIEW_DAYS))

    _ordered(_run(tmp_path, now=today))

    assert layout.yellow is not None
    assert _view(path).top_row == layout.yellow + 2


def test_usterka_poprzedniego_dnia_dopisana_pod_dzisiejszym_dniem_nie_przesuwa_widoku(tmp_path: Path) -> None:
    ordered = (*_VIEW_DAYS[1:3], Order(datetime(2026, 10, 19, 15), "anna.kowalska", 5))
    path, _ = _register(tmp_path, Intermax(ordered=ordered))

    _ordered(_run(tmp_path, now=datetime(2026, 10, 20, 7)))

    assert _view(path).top_row == 2


# Zlecenie na dziś w wierszu 2, pod nim ``below`` zleceń z 19.10, pusty wiersz i dzisiejszy dzień Wykonawcy jan.nowak.
@pytest.mark.parametrize(("below", "top_row"), [(9, 2), (10, 14)])
def test_zlecenie_na_dzis_wpisane_we_wczesniejszym_dniu_jest_poczatkiem_widoku_tylko_tuz_nad_dzisiejszym_dniem(
    tmp_path: Path, below: int, top_row: int
) -> None:
    earlier = (Order(datetime(2026, 10, 19, 8 + hour % 10), "anna.kowalska", 10 + hour) for hour in range(below))
    ordered = (Order(datetime(2026, 10, 20, 7), "anna.kowalska", 1), *earlier, Order(datetime(2026, 10, 20, 9), "jan.nowak", 2))
    path, _ = _register(tmp_path, Intermax(ordered=ordered))

    _ordered(_run(tmp_path, now=datetime(2026, 10, 20, 7)))

    assert _view(path).top_row == top_row


def test_zlecenie_na_dzis_wsrod_wczesniejszych_zaczyna_widok_gdy_10_wczesniejszych_nie_stoi_pod_rzad(tmp_path: Path) -> None:
    """Granica 10 wierszy liczy wcześniejsze dni pod rząd, jak przegląd nad wklejeniem: zlecenie na dziś ją zeruje."""

    earlier = tuple(Order(datetime(2026, 10, 19, 8 + hour), "anna.kowalska", 10 + hour) for hour in range(10))
    ordered = (
        Order(datetime(2026, 10, 20, 7), "anna.kowalska", 1),
        *earlier[:9],
        Order(datetime(2026, 10, 20, 8), "anna.kowalska", 2),
        earlier[9],
        Order(datetime(2026, 10, 20, 9), "jan.nowak", 3),
    )
    path, _ = _register(tmp_path, Intermax(ordered=ordered))

    _ordered(_run(tmp_path, now=datetime(2026, 10, 20, 7)))

    assert _view(path).top_row == 2


def test_dzien_bez_zoltego_wiersza_z_jedynym_zleceniem_dzisiaj_w_ostatnim_uporzadkowanym_wierszu_jest_widokiem(
    tmp_path: Path,
) -> None:
    ordered = (_VIEW_DAYS[0], _VIEW_DAYS[3])
    path, layout = _register(tmp_path, Intermax(ordered=ordered, yellow=None))

    _ordered(_run(tmp_path, now=datetime(2026, 10, 20, 7)))

    assert _view(path).top_row == layout.last_ordered == 4


def test_intermax_zapisany_z_innym_aktywnym_arkuszem_otwiera_sie_na_arkuszu_zlecen(tmp_path: Path) -> None:
    path, _ = _register(tmp_path, Intermax(ordered=_VIEW_DAYS, active_sheet="LISTA OD 08.11.2021."))
    before = parts(path)

    _ordered(_run(tmp_path, now=datetime(2026, 10, 20, 7)))

    assert _view(path) == _View(ORDER_SHEET, (ORDER_SHEET,), 3, "B3", frozen_header=True)
    after = parts(path)
    assert {name for name in before if before[name] != after[name]} == {
        _SHEET,
        _STYLES,
        "xl/workbook.xml",
        "xl/worksheets/sheet3.xml",
    }


def test_intermax_zapisany_na_ostatnim_arkuszu_otwiera_sie_na_samym_arkuszu_zlecen(tmp_path: Path) -> None:
    path, _ = _register(tmp_path, Intermax(ordered=_VIEW_DAYS, active_sheet="komentarze"))

    _ordered(_run(tmp_path, now=datetime(2026, 10, 20, 7)))

    assert _view(path).selected == (ORDER_SHEET,)


def test_karta_zaznaczona_jako_true_w_innym_arkuszu_przestaje_byc_zaznaczona(tmp_path: Path) -> None:
    path, _ = _register(tmp_path, Intermax(ordered=_VIEW_DAYS, active_sheet="LISTA OD 08.11.2021."))
    rewrite_parts(
        path, lambda name, data: data.replace(b'tabSelected="1"', b'tabSelected="true"') if name == _LIST_SHEET else data
    )

    _ordered(_run(tmp_path, now=datetime(2026, 10, 20, 7)))

    assert _view(path).selected == (ORDER_SHEET,)
    assert b"tabSelected" not in parts(path)[_LIST_SHEET]


def test_w_drugim_oknie_excela_intermax_tez_otwiera_sie_na_arkuszu_zlecen_i_dzisiejszym_dniu(tmp_path: Path) -> None:
    """Każde okno skoroszytu ma własny ``<workbookView>`` i własny ``<sheetView>`` w każdym arkuszu."""

    path, _ = _register(tmp_path, Intermax(ordered=_VIEW_DAYS, active_sheet="LISTA OD 08.11.2021."))

    def second_window(name: str, data: bytes) -> bytes:
        if name == "xl/workbook.xml":
            return re.sub(rb"(<workbookView\b[^>]*/>)", rb"\1\1", data)
        if name == _SHEET:
            view = re.search(rb"<sheetView\b.*?</sheetView>", data, re.S)
            assert view is not None
            second = view.group(0).replace(b'workbookViewId="0"', b'workbookViewId="1"')
            return data.replace(view.group(0), view.group(0) + second)
        return data

    rewrite_parts(path, second_window)

    _ordered(_run(tmp_path, now=datetime(2026, 10, 20, 7)))

    after = parts(path)
    assert re.findall(rb'<workbookView\b[^>]*\sactiveTab="(\d+)"', after["xl/workbook.xml"]) == [b"1", b"1"]
    views = re.findall(rb"<sheetView\b.*?</sheetView>", after[_SHEET], re.S)
    assert len(views) == 2
    for view in views:
        assert b'tabSelected="1"' in view and b'topLeftCell="A3"' in view and b'activeCell="B3"' in view


def test_intermax_zapisany_na_pierwszym_arkuszu_bez_activetab_otwiera_sie_na_arkuszu_zlecen(tmp_path: Path) -> None:
    """Excel nie zapisuje ``activeTab``, gdy aktywny jest pierwszy arkusz."""

    path, _ = _register(tmp_path, Intermax(ordered=_VIEW_DAYS))
    rewrite_parts(path, lambda name, data: data.replace(b' activeTab="1"', b"") if name == "xl/workbook.xml" else data)
    assert load_workbook(path).active.title == "Arkusz1"  # type: ignore[union-attr]

    _ordered(_run(tmp_path, now=datetime(2026, 10, 20, 7)))

    assert _view(path).active == ORDER_SHEET


def test_zamrozony_naglowek_bez_activepane_ma_zaznaczenie_w_gornym_panelu(tmp_path: Path) -> None:
    """Bez ``activePane`` aktywny jest górny lewy panel, więc zaznaczenie staje w nim."""

    path, _ = _register(tmp_path, Intermax(ordered=_VIEW_DAYS))
    rewrite_parts(path, lambda name, data: data.replace(b' activePane="bottomLeft"', b"") if name == _SHEET else data)

    _ordered(_run(tmp_path, now=datetime(2026, 10, 20, 7)))

    view = load_workbook(path)[ORDER_SHEET].sheet_view
    assert view.pane is not None and (view.pane.ySplit, view.pane.state, view.pane.topLeftCell) == (1, "frozen", "A3")
    assert [(selection.pane, selection.activeCell) for selection in view.selection] == [("topLeft", "B3")]


def test_bez_zamrozonego_naglowka_widok_bez_panelu_otwiera_sie_na_dzisiejszym_dniu(tmp_path: Path) -> None:
    path, _ = _register(tmp_path, Intermax(ordered=_VIEW_DAYS, frozen=False))

    _ordered(_run(tmp_path, now=datetime(2026, 10, 20, 7)))

    assert _view(path) == _View(ORDER_SHEET, (ORDER_SHEET,), 3, "B3", frozen_header=False)
    assert load_workbook(path)[ORDER_SHEET].sheet_view.pane is None


def test_kopia_intermax_ma_nazwe_z_zegara_procesu_i_bajty_sprzed_uporzadkowania(tmp_path: Path) -> None:
    path, _ = _register(tmp_path)
    original = path.read_bytes()

    _run(tmp_path)

    copy = tmp_path / "Kopie INTERMAX" / COPY_NAME
    assert copy.read_bytes() == original
    assert path.read_bytes() != original


def test_po_jedenastej_kopii_zostaje_10_najnowszych_a_inne_pliki_folderu_zostaja(tmp_path: Path) -> None:
    _register(tmp_path)
    copies = tmp_path / "Kopie INTERMAX"
    copies.mkdir()
    older = [copies / f"INTERMAX przed uporządkowaniem 2026-10-{day:02d} 07-00-00.xlsx" for day in range(10, 20)]
    for copy in older:
        copy.write_bytes(b"kopia")
    (copies / "notatka.txt").write_text("tata", encoding="utf-8")

    _run(tmp_path)

    assert sorted(path.name for path in copies.iterdir()) == sorted(
        [*(copy.name for copy in older[1:]), COPY_NAME, "notatka.txt"]
    )


def _part_changed(content: bytes, name: str, change: Callable[[bytes], bytes]) -> bytes:
    out = BytesIO()
    with zipfile.ZipFile(BytesIO(content)) as source, zipfile.ZipFile(out, "w") as target:
        for info in source.infolist():
            data = source.read(info)
            target.writestr(info, change(data) if info.filename == name else data)
    return out.getvalue()


def _parts_reordered(content: bytes) -> bytes:
    out = BytesIO()
    with zipfile.ZipFile(BytesIO(content)) as source, zipfile.ZipFile(out, "w") as target:
        infos = source.infolist()
        for info in [infos[1], infos[0], *infos[2:]]:
            target.writestr(info, source.read(info))
    return out.getvalue()


_SHEET, _STYLES = "xl/worksheets/sheet2.xml", "xl/styles.xml"
_LIST_SHEET = "xl/worksheets/sheet3.xml"
"""Arkusz ``LISTA OD 08.11.2021.`` fixture'u, za arkuszem zleceń."""


def _in_sheet(change: Callable[[bytes], bytes]) -> Callable[[bytes, Layout], bytes]:
    return lambda content, _layout: _part_changed(content, _SHEET, change)


def _cell_value_changed(content: bytes, layout: Layout) -> bytes:
    assert layout.yellow is not None
    row = rb'(<c r="D%d"[^>]*><v>)\d+' % (layout.yellow + 2)
    return _part_changed(content, _SHEET, lambda xml: re.sub(row, rb"\g<1>0", xml, count=1))


def _cell_kind_changed(content: bytes, layout: Layout) -> bytes:
    assert layout.yellow is not None
    row = rb'(<c r="D%d"[^>]*?) t="s"' % (layout.yellow + 2)
    return _part_changed(content, _SHEET, lambda xml: re.sub(row, rb'\g<1> t="str"', xml, count=1))


def _blank_row_valued(cell: bytes) -> Callable[[bytes, Layout], bytes]:
    def change(content: bytes, layout: Layout) -> bytes:
        assert layout.yellow is not None
        number = layout.yellow + 5
        separator = rb'(<row r="%d"[^>]*>)' % number
        return _part_changed(content, _SHEET, lambda xml: re.sub(separator, rb"\g<1>" + cell % number, xml))

    return change


def _in_first_order(change: Callable[[bytes, int], bytes]) -> Callable[[bytes, Layout], bytes]:
    def changed(content: bytes, layout: Layout) -> bytes:
        assert layout.yellow is not None
        return _part_changed(content, _SHEET, lambda xml: change(xml, layout.yellow + 2))  # type: ignore[operator]

    return changed


def _order_row_removed(content: bytes, layout: Layout) -> bytes:
    assert layout.yellow is not None
    row = rb'<row r="%d"[^>]*>.*?</row>' % (layout.yellow + 9)
    return _part_changed(content, _SHEET, lambda xml: re.sub(row, b"", xml, count=1, flags=re.S))


# Każde zepsucie zmienia w pliku tymczasowym jedną rzecz, której pilnuje jedno sprawdzenie zapisu.
CORRUPTIONS: dict[str, Callable[[bytes, Layout], bytes]] = {
    "nie zip": lambda _content, _layout: b"to nie jest skoroszyt",
    "kolejność części": lambda content, _layout: _parts_reordered(content),
    "inna część": lambda content, _layout: _part_changed(content, "docProps/app.xml", lambda xml: xml + b" "),
    "skoroszyt poza aktywną kartą": lambda content, _layout: _part_changed(
        content, "xl/workbook.xml", lambda xml: xml.replace(b'tabRatio="600"', b'tabRatio="601"', 1)
    ),
    "inny arkusz poza zaznaczeniem karty": lambda content, _layout: _part_changed(
        content, _LIST_SHEET, lambda xml: xml.replace(b'activeCell="A1" sqref="A1"', b'activeCell="A2" sqref="A2"', 1)
    ),
    "inna aktywna karta": lambda content, _layout: _part_changed(
        content, "xl/workbook.xml", lambda xml: xml.replace(b'activeTab="1"', b'activeTab="2"', 1)
    ),
    "inny arkusz zaznaczony": lambda content, _layout: _part_changed(
        content, _LIST_SHEET, lambda xml: xml.replace(b"<sheetView ", b'<sheetView tabSelected="true" ', 1)
    ),
    "istniejący styl": lambda content, _layout: _part_changed(
        content, _STYLES, lambda xml: xml.replace(b'<sz val="11"', b'<sz val="12"', 1)
    ),
    "istniejące wypełnienie": lambda content, _layout: _part_changed(
        content, _STYLES, lambda xml: xml.replace(b'patternType="gray125"', b'patternType="gray0625"', 1)
    ),
    "lista stylów": lambda content, _layout: _part_changed(
        content, _STYLES, lambda xml: re.sub(rb"<borders\b.*?</borders>", b"", xml, flags=re.S)
    ),
    "wiersze nad dniem": _in_sheet(lambda xml: xml.replace(b'<row r="2" ', b'<row r="2" hidden="1" ', 1)),
    "ogon arkusza": _in_sheet(lambda xml: xml.replace(b'<pageMargins left="0.75"', b'<pageMargins left="0.5"', 1)),
    "wartość zlecenia": _cell_value_changed,
    "rodzaj wartości zlecenia": _cell_kind_changed,
    "liczba w pustym wierszu": _blank_row_valued(b'<c r="B%d"><v>1</v></c>'),
    "tekst w pustym wierszu": _blank_row_valued(b'<c r="B%d" t="inlineStr"><is><t>x</t></is></c>'),
    "brak zlecenia": _order_row_removed,
    "dodatkowa wartość w zleceniu": _in_first_order(
        lambda xml, row: re.sub(rb'(<row r="%d"[^>]*>.*?)</row>' % row, rb'\g<1><c r="AZ%d"><v>1</v></c></row>' % row, xml, count=1, flags=re.S)
    ),
    "brak pustej komórki zlecenia": _in_first_order(lambda xml, row: re.sub(rb'<c r="B%d"[^>]*/>' % row, b"", xml, count=1)),
    "wygląd kolumny za UWAGI": _in_first_order(
        lambda xml, row: re.sub(rb'(<c r="P%d" s=")\d+' % row, rb"\g<1>0", xml, count=1)
    ),
}


@pytest.mark.parametrize("corruption", CORRUPTIONS)
def test_niezgodnosc_w_sprawdzeniu_zapisu_to_zatrzymanie_do_maksa_bez_zmiany_intermax_i_kopii(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, corruption: str
) -> None:
    path, layout = _register(tmp_path)
    write_bytes = Path.write_bytes

    def corrupted(target: Path, data: bytes) -> int:
        if target.parent == path.parent and target.name.startswith(f".{path.stem}."):
            data = CORRUPTIONS[corruption](bytes(data), layout)
        return write_bytes(target, data)

    monkeypatch.setattr(Path, "write_bytes", corrupted)

    with _nothing_written(tmp_path, path):
        stop = _stop(_run(tmp_path))

    assert (stop.message.code, stop.stop_kind) == ("ZAPIS_INTERMAX_NIEZGODNY", "for_maks")

def test_bez_surowych_zlecen_zatrzymanie_dla_taty_i_nic_nie_zapisano(tmp_path: Path) -> None:
    path, _ = _register(tmp_path, Intermax(raw=()))
    original = path.read_bytes()

    stop = _stop(_run(tmp_path))

    assert (stop.message.code, stop.stop_kind) == ("BRAK_SUROWYCH_ZLECEN", "for_tata")
    assert ORDER_SHEET in stop.message.where
    assert path.read_bytes() == original
    assert not (tmp_path / "Kopie INTERMAX").exists()


def test_intermax_otwarty_w_excelu_to_zatrzymanie_dla_taty_w_sprawdzaniu_bez_zapisu(tmp_path: Path) -> None:
    path, _ = _register(tmp_path)
    original = path.read_bytes()
    lock = path.with_name(f"~${path.name}")
    lock.write_bytes(b"")

    outcome = _run(tmp_path)

    stop = _stop(outcome)
    assert outcome.progress.phase_states[ProgressPhase.CHECKING] == PhaseState.FAILED

    assert (stop.message.code, stop.stop_kind) == ("INTERMAX_OTWARTY", "for_tata")
    assert path.read_bytes() == original
    assert not (tmp_path / "Kopie INTERMAX").exists()


def test_scalone_komorki_w_surowych_zleceniach_to_zatrzymanie_do_maksa_bez_zapisu(tmp_path: Path) -> None:
    path, layout = _register(tmp_path)
    workbook = load_workbook(path)
    workbook.worksheets[1].merge_cells(start_row=layout.first_raw + 1, start_column=4, end_row=layout.first_raw + 2, end_column=4)
    workbook.save(path)
    original = path.read_bytes()

    stop = _stop(_run(tmp_path))

    assert (stop.message.code, stop.stop_kind) == ("NIEOBSLUGIWANY_UKLAD_INTERMAX", "for_maks")
    assert path.read_bytes() == original


@pytest.mark.parametrize("now", [datetime(2026, 10, 25, 21, 5, 7), datetime(2026, 10, 27, 8, 0, 0)])
def test_uruchomienie_ma_okres_dzisiejszej_daty_z_zegara_procesu(tmp_path: Path, now: datetime) -> None:
    inputs = replace(OrderingInputs(tmp_path / REGISTER_NAME), now=lambda: now)

    assert inputs.period == {25: "20_26_10_2026", 27: "27_02_11_2026"}[now.day]
    assert inputs.dry_run is False


def _first_day_row(path: Path, layout: Layout, order: Order) -> int:
    """Wiersz arkusza zleceń, w którym stoi zlecenie ``order`` (po ``NUMER ZL.``)."""

    sheet = load_workbook(path, read_only=True)[ORDER_SHEET]
    numbers = [row[NUMBER - 1].value for row in sheet.iter_rows(min_row=1, max_col=NUMBER)]
    return numbers.index(f"ZL-{order.number}") + 1


@pytest.mark.parametrize("yellow", ["row", "cells"])
@pytest.mark.parametrize("above_yellow", [1, 3])
def test_dzien_staje_jeden_pusty_wiersz_pod_zoltym_wierszem_taty(
    tmp_path: Path, yellow: Literal["row", "cells"], above_yellow: int
) -> None:
    path, layout = _register(tmp_path, Intermax(yellow=yellow, above_yellow=above_yellow))

    _ordered(_run(tmp_path))

    assert layout.yellow == layout.last_ordered + above_yellow + 1
    assert _first_day_row(path, layout, layout.raw[1]) == layout.yellow + 2
    assert order_rows(path, layout.last_ordered + 1, layout.yellow + 1) == [None] * (above_yellow + 2)


@pytest.mark.parametrize(("cells", "under_yellow"), [(7, True), (6, False)])
def test_zolty_wiersz_to_co_najmniej_polowa_kolumn_zlecenia_w_zoltym(tmp_path: Path, cells: int, under_yellow: bool) -> None:
    path, layout = _register(tmp_path, Intermax(yellow="cells", yellow_cells=cells))

    _ordered(_run(tmp_path))

    assert layout.yellow is not None
    expected = layout.yellow + 2 if under_yellow else layout.last_ordered + 11
    assert _first_day_row(path, layout, layout.raw[1]) == expected


def test_z_dwoch_zoltych_wierszy_dzien_staje_pod_nizszym(tmp_path: Path) -> None:
    path, layout = _register(tmp_path, Intermax(above_yellow=3))
    _edited(path, lambda sheet: setattr(sheet.row_dimensions[layout.last_ordered + 1], "fill", YELLOW))

    _ordered(_run(tmp_path))

    assert layout.yellow is not None
    assert _first_day_row(path, layout, layout.raw[1]) == layout.yellow + 2


def test_zolty_wiersz_wewnatrz_wklejenia_nie_wyznacza_miejsca_dnia(tmp_path: Path) -> None:
    path, layout = _register(tmp_path, Intermax(yellow=None, gap=12))
    inside = layout.first_raw + 2

    def emptied(sheet: Worksheet) -> None:
        for column in range(1, APP_LAST + 1):
            cell = sheet.cell(inside, column)
            cell.value, cell.fill, cell.hyperlink = None, YELLOW, None

    _edited(path, emptied)

    summary = _ordered(_run(tmp_path))

    assert summary.order_count == len(layout.raw) - 1
    assert _first_day_row(path, layout, layout.raw[1]) == layout.last_ordered + 11


def test_wiersz_wklejenia_z_wykonawca_w_calibri_8_dalej_jest_surowym_zleceniem(tmp_path: Path) -> None:
    path, layout = _register(tmp_path)
    last_raw = layout.first_raw + len(layout.raw) - 1
    _edited(path, lambda sheet: setattr(sheet.cell(last_raw, WORKER), "font", Font(name="Calibri", sz=8)))

    assert _ordered(_run(tmp_path)).order_count == len(layout.raw)


def test_komorka_w_calibri_8_za_uwagi_nie_czyni_wiersza_wklejenia_uporzadkowanym(tmp_path: Path) -> None:
    path, layout = _register(tmp_path)
    last_raw = layout.first_raw + len(layout.raw) - 1
    _edited(path, lambda sheet: setattr(sheet.cell(last_raw, APP_LAST), "font", Font(name="Calibri", sz=8)))

    assert _ordered(_run(tmp_path)).order_count == len(layout.raw)


def test_uporzadkowany_wiersz_z_jedna_komorka_w_innej_czcionce_dalej_jest_uporzadkowany(tmp_path: Path) -> None:
    path, layout = _register(tmp_path)
    _edited(path, lambda sheet: setattr(sheet.cell(layout.last_ordered, 4), "font", Font(name="Arial", sz=10)))

    assert _ordered(_run(tmp_path)).order_count == len(layout.raw)
    assert layout.yellow is not None
    assert _first_day_row(path, layout, layout.raw[1]) == layout.yellow + 2


def test_pierwszy_wiersz_z_data_na_miejscu_nie_jest_przesuwany_nawet_z_data_w_kolumnie_a(tmp_path: Path) -> None:
    path, layout = _register(tmp_path, Intermax(shifted=False))
    stray = datetime(2026, 10, 26, 7)
    _edited(path, lambda sheet: setattr(sheet.cell(layout.first_raw, DATE - 2), "value", stray))

    _ordered(_run(tmp_path))

    assert layout.yellow is not None
    assert order_rows(path, layout.yellow + 9, layout.yellow + 9) == [(stray, *_expected(layout.raw[0])[1:])]


def test_segmenty_wykonawcow_ida_alfabetycznie_bez_wzgledu_na_wielkosc_liter(tmp_path: Path) -> None:
    raw = (
        Order(datetime(2026, 10, 26, 9), "Cezary", 201),
        Order(datetime(2026, 10, 26, 9), "anna", 202),
        Order(datetime(2026, 10, 26, 9), "Bartek", 203),
        Order(datetime(2026, 10, 26, 10), "ANNA", 204),
    )
    path, layout = _register(tmp_path, Intermax(raw=raw))

    summary = _ordered(_run(tmp_path))

    assert layout.yellow is not None
    rows = order_rows(path, layout.yellow + 2, layout.yellow + 8)
    assert [None if row is None else row[NUMBER - 1] for row in rows] == [
        "ZL-202",
        "ZL-204",
        None,
        "ZL-203",
        None,
        "ZL-201",
        None,
    ]
    assert summary.worker_count == 3


def test_bez_uporzadkowanych_dni_dzien_staje_10_pustych_wierszy_pod_naglowkiem_bez_wygladu_naglowka(
    tmp_path: Path,
) -> None:
    path, layout = _register(tmp_path, Intermax(ordered=(), yellow=None, grid=False, excel=False, gap=0))
    _edited(path, lambda sheet: setattr(sheet.row_dimensions[1], "height", 20))
    assert layout.first_raw == 2

    assert _ordered(_run(tmp_path)).order_count == len(layout.raw)

    assert _first_day_row(path, layout, layout.raw[1]) == 12
    assert _row_look(path, 12) == {}


def test_porzadkowany_jest_arkusz_zlecen_z_wklejeniem_a_starszy_zostaje(tmp_path: Path) -> None:
    path, layout = _register(tmp_path, Intermax(older_sheet=True))
    older = load_workbook(path)["ZLECENIA OD 08_11_2021"]
    before = [[cell.value for cell in row] for row in older.iter_rows()]

    assert _ordered(_run(tmp_path)).order_count == len(layout.raw)

    assert [[cell.value for cell in row] for row in load_workbook(path)["ZLECENIA OD 08_11_2021"].iter_rows()] == before
    assert layout.yellow is not None
    assert _first_day_row(path, layout, layout.raw[1]) == layout.yellow + 2


def _without_workbook_properties(path: Path) -> None:
    rewrite_parts(path, lambda name, data: re.sub(rb"<workbookPr\b[^>]*/>", b"", data) if name == "xl/workbook.xml" else data)


def test_data_wklejenia_we_wbudowanym_formacie_daty_excela_jest_data(tmp_path: Path) -> None:
    path, layout = _register(tmp_path, Intermax(raw_date_format=DATE_FORMAT))

    summary = _ordered(_run(tmp_path))

    assert (summary.day.isoformat(), summary.order_count) == ("2026-10-26", len(layout.raw))


@pytest.mark.parametrize("dates", ["1904", "bez workbookPr"])
def test_dzien_ma_date_z_systemu_dat_intermax(tmp_path: Path, dates: str) -> None:
    path, layout = _register(tmp_path, Intermax(epoch_1904=dates == "1904"))
    if dates != "1904":
        _without_workbook_properties(path)

    summary = _ordered(_run(tmp_path))

    assert summary.day.isoformat() == "2026-10-26"
    assert layout.yellow is not None
    assert order_rows(path, layout.yellow + 2, layout.yellow + 2) == [_expected(layout.raw[1])]


@pytest.mark.parametrize("value", [True, 7])
def test_wiersz_z_sama_wartoscia_pod_wklejeniem_to_zlecenie_bez_daty_z_wierszem_i_kolumna(
    tmp_path: Path, value: object
) -> None:
    path, layout = _register(tmp_path)
    stray = layout.first_raw + len(layout.raw) + 1
    _edited(path, lambda sheet: setattr(sheet.cell(stray, 12), "value", value))

    with _nothing_written(tmp_path, path):
        stop = _stop(_run(tmp_path))

    assert (stop.message.code, stop.stop_kind) == ("SUROWE_ZLECENIE_BEZ_DATY", "for_tata")
    assert stop.message.where.endswith(f"komórka C{stray}")


def test_wiersz_z_samymi_spacjami_pod_wklejeniem_jest_pusty(tmp_path: Path) -> None:
    path, layout = _register(tmp_path)
    _edited(path, lambda sheet: setattr(sheet.cell(layout.first_raw + len(layout.raw) + 1, 12), "value", "   "))

    assert _ordered(_run(tmp_path)).order_count == len(layout.raw)


def test_pusta_komorka_w_calibri_8_nie_czyni_wiersza_wklejenia_uporzadkowanym(tmp_path: Path) -> None:
    path, layout = _register(tmp_path)
    last_raw = layout.first_raw + len(layout.raw) - 1
    _edited(path, lambda sheet: setattr(sheet.cell(last_raw, FIRST), "font", Font(name="Calibri", sz=8)))

    assert _ordered(_run(tmp_path)).order_count == len(layout.raw)


def test_blad_i_prawda_w_komorkach_wklejenia_przechodza_do_dnia_bez_zmian(tmp_path: Path) -> None:
    path, layout = _register(tmp_path)

    def marked(sheet: Worksheet) -> None:
        sheet.cell(layout.first_raw + 1, 12).value = True
        sheet.cell(layout.first_raw + 1, 13).value = "#N/A"

    _edited(path, marked)

    _ordered(_run(tmp_path))

    assert layout.yellow is not None
    sheet = load_workbook(path)[ORDER_SHEET]
    assert (sheet.cell(layout.yellow + 2, 12).value, sheet.cell(layout.yellow + 2, 13).value) == (True, "#N/A")


@pytest.mark.parametrize("value", ["jutro", 46321.5, None], ids=["tekst", "liczba bez formatu daty", "pusta"])
def test_surowe_zlecenie_bez_daty_to_zatrzymanie_dla_taty_z_wierszem_i_kolumna(
    tmp_path: Path, value: str | float | None
) -> None:
    path, layout = _register(tmp_path)
    row = layout.first_raw + 1

    def undated(sheet: Worksheet) -> None:
        cell = sheet.cell(row, DATE)
        cell.value, cell.number_format = value, "General"

    _edited(path, undated)

    with _nothing_written(tmp_path, path):
        stop = _stop(_run(tmp_path))

    assert (stop.message.code, stop.stop_kind) == ("SUROWE_ZLECENIE_BEZ_DATY", "for_tata")
    assert f"w wierszu {row} nie ma daty z godziną w kolumnie C." in stop.message.what
    assert stop.message.where == f"plik {REGISTER_NAME}, folder OneDrive, arkusz {ORDER_SHEET}, komórka C{row}"
    assert stop.message.action == f"Wpisz datę z godziną w komórce C{row}, zapisz i zamknij INTERMAX i uruchom ponownie."
    _without_order_data(stop.message, layout.raw)


@pytest.mark.parametrize(("shifted", "column"), [(True, "H"), (False, "J")])
def test_surowe_zlecenie_bez_wykonawcy_to_zatrzymanie_dla_taty_z_wierszem_i_kolumna(
    tmp_path: Path, shifted: bool, column: str
) -> None:
    path, layout = _register(tmp_path, Intermax(shifted=shifted))
    _edited(path, lambda sheet: setattr(sheet[f"{column}{layout.first_raw}"], "value", None))

    with _nothing_written(tmp_path, path):
        stop = _stop(_run(tmp_path))

    assert (stop.message.code, stop.stop_kind) == ("SUROWE_ZLECENIE_BEZ_WYKONAWCY", "for_tata")
    assert f"w wierszu {layout.first_raw} nie ma Wykonawcy w kolumnie {column}." in stop.message.what
    assert stop.message.where.endswith(f"komórka {column}{layout.first_raw}")
    assert stop.message.action.startswith(f"Wpisz Wykonawcę w komórce {column}{layout.first_raw}")
    _without_order_data(stop.message, layout.raw)


def test_zlecenie_bez_daty_we_wklejeniu_z_kilku_dni_to_najpierw_brak_daty(tmp_path: Path) -> None:
    """Bez daty nie wiadomo, z którego dnia jest zlecenie, więc tata najpierw wpisuje datę."""

    raw = (
        Order(datetime(2026, 10, 27, 9), "jan.nowak", 101),
        Order(datetime(2026, 10, 26, 10), "anna.kowalska", 102),
        Order(datetime(2026, 10, 26, 11), "jan.nowak", 103),
    )
    path, layout = _register(tmp_path, Intermax(raw=raw))
    _edited(path, lambda sheet: setattr(sheet.cell(layout.first_raw + 2, DATE), "value", "jutro"))

    with _nothing_written(tmp_path, path):
        stop = _stop(_run(tmp_path))

    assert stop.message.code == "SUROWE_ZLECENIE_BEZ_DATY"
    assert stop.message.where.endswith(f"komórka C{layout.first_raw + 2}")


def test_przesuniety_pierwszy_wiersz_bez_daty_wskazuje_kolumne_a(tmp_path: Path) -> None:
    path, layout = _register(tmp_path)
    _edited(path, lambda sheet: setattr(sheet.cell(layout.first_raw, DATE - 2), "value", "jutro"))

    with _nothing_written(tmp_path, path):
        stop = _stop(_run(tmp_path))

    assert stop.message.code == "SUROWE_ZLECENIE_BEZ_DATY"
    assert stop.message.where.endswith(f"komórka A{layout.first_raw}")


def test_wklejenie_z_kilku_dni_to_zatrzymanie_dla_taty_z_dniami(tmp_path: Path) -> None:
    raw = (
        Order(datetime(2026, 10, 27, 9), "jan.nowak", 101),
        Order(datetime(2026, 10, 26, 10), "anna.kowalska", 102),
        Order(datetime(2026, 10, 26, 11), "jan.nowak", 103),
    )
    path, layout = _register(tmp_path, Intermax(raw=raw))

    with _nothing_written(tmp_path, path):
        stop = _stop(_run(tmp_path))

    assert (stop.message.code, stop.stop_kind) == ("WKLEJENIE_Z_KILKU_DNI", "for_tata")
    assert "26.10.2026, 27.10.2026" in stop.message.what
    assert stop.message.where.endswith(f"wiersze {layout.first_raw}–{layout.first_raw + 2}")
    assert stop.message.action == (
        "Zostaw na dole arkusza zleceń zlecenia jednego dnia, a pozostałe usuń. Zapisz i zamknij INTERMAX i uruchom "
        "ponownie, a następny dzień wklej i uporządkuj osobno."
    )
    _without_order_data(stop.message, layout.raw)


def test_surowe_zlecenia_w_dwoch_arkuszach_to_zatrzymanie_dla_taty_z_nazwami_arkuszy(tmp_path: Path) -> None:
    path, layout = _register(tmp_path, Intermax(older_sheet=True))
    workbook = load_workbook(path)
    older = workbook["ZLECENIA OD 08_11_2021"]
    _raw_row(layout.raw[1])(older, older.max_row + 3)
    workbook.save(path)

    with _nothing_written(tmp_path, path):
        stop = _stop(_run(tmp_path))

    assert (stop.message.code, stop.stop_kind) == ("SUROWE_ZLECENIA_W_KILKU_ARKUSZACH", "for_tata")
    assert stop.message.where.endswith(f"arkusze ZLECENIA OD 08_11_2021, {ORDER_SHEET}")
    _without_order_data(stop.message, layout.raw)


@pytest.mark.parametrize("raw_below", [True, False])
def test_surowe_zlecenia_nad_uporzadkowanym_wierszem_to_zatrzymanie_dla_taty(tmp_path: Path, raw_below: bool) -> None:
    path, layout = _register(tmp_path)
    last_raw = layout.first_raw + len(layout.raw) - 1
    typed = Order(datetime(2026, 10, 25, 12), "jan.nowak", 301)

    def typed_below(sheet: Worksheet) -> None:
        _ordered_row(typed)(sheet, last_raw + 2)
        _ordered_row(replace(typed, number=303))(sheet, last_raw + 3)
        if raw_below:
            _raw_row(Order(datetime(2026, 10, 26, 7), "jan.nowak", 302))(sheet, last_raw + 5)

    _edited(path, typed_below)

    with _nothing_written(tmp_path, path):
        stop = _stop(_run(tmp_path))

    assert (stop.message.code, stop.stop_kind) == ("SUROWE_ZLECENIA_NAD_UPORZADKOWANYMI", "for_tata")
    assert f"Wiersz {last_raw} wygląda jak wklejone zlecenie, ale stoi nad uporządkowanym wierszem {last_raw + 2}." in (
        stop.message.what
    )
    assert stop.message.where.endswith(f"arkusz {ORDER_SHEET}, wiersz {last_raw}")
    assert stop.message.action == (
        "Przenieś wklejone zlecenia na dół arkusza zleceń, pod ostatni uporządkowany wiersz. Jeśli wiersz "
        f"{last_raw} jest już uporządkowany, nadaj mu czcionkę Calibri 8. Zapisz i zamknij INTERMAX i uruchom ponownie."
    )
    _without_order_data(stop.message, [*layout.raw, typed, replace(typed, number=303)])


def test_bez_wklejenia_wiersz_z_wygladem_z_aplikacji_w_starszym_dniu_nie_zatrzymuje(tmp_path: Path) -> None:
    ordered = tuple(
        Order(datetime(2026, 10, 19 if index < 12 else 20, 7 + index % 12), "anna.kowalska", index + 1) for index in range(14)
    )
    path, _ = _register(tmp_path, Intermax(ordered=ordered, raw=()))
    _edited(path, _in_other_font(2))

    with _nothing_written(tmp_path, path):
        stop = _stop(_run(tmp_path))

    assert stop.message.code == "BRAK_SUROWYCH_ZLECEN"


def test_bez_wklejenia_wiersz_z_wygladem_z_aplikacji_wysoko_w_ostatnim_dniu_zatrzymuje(tmp_path: Path) -> None:
    """Wiersze ostatniego dnia nie są wcześniejszym dniem, więc nie kończą przeglądu."""

    ordered = tuple(Order(datetime(2026, 10, 19, 7 + index % 12), "anna.kowalska", index + 1) for index in range(14))
    path, _ = _register(tmp_path, Intermax(ordered=ordered, raw=()))
    _edited(path, _in_other_font(2))

    with _nothing_written(tmp_path, path):
        stop = _stop(_run(tmp_path))

    assert stop.message.code == "SUROWE_ZLECENIA_NAD_UPORZADKOWANYMI"
    assert stop.message.where.endswith("wiersz 2")


def test_pojedyncze_wiersze_wczesniejszych_dni_miedzy_dniami_nie_koncza_przegladu(tmp_path: Path) -> None:
    def day(number: int, count: int, first: int) -> tuple[Order, ...]:
        return tuple(Order(datetime(2026, 10, number, 7 + index), "anna.kowalska", first + index) for index in range(count))

    path, _ = _register(tmp_path, Intermax(ordered=(*day(26, 2, 1), *day(25, 6, 3), *day(27, 1, 9), *day(25, 6, 10))))

    with _nothing_written(tmp_path, path):
        stop = _stop(_run(tmp_path))

    assert stop.message.code == "DZIEN_JUZ_UPORZADKOWANY"


@pytest.mark.parametrize(("above", "stops"), [(9, True), (10, False)])
def test_wiersz_z_wygladem_z_aplikacji_w_poprzednim_dniu_zatrzymuje_tylko_tuz_nad_wklejeniem(
    tmp_path: Path, above: int, stops: bool
) -> None:
    ordered = tuple(Order(datetime(2026, 10, 19, 7 + index), "anna.kowalska", index + 1) for index in range(12))
    path, layout = _register(tmp_path, Intermax(ordered=ordered))
    strange = layout.last_ordered - above
    _edited(path, _in_other_font(strange))

    outcome = _run(tmp_path)

    if stops:
        assert _stop(outcome).message.code == "SUROWE_ZLECENIA_NAD_UPORZADKOWANYMI"
        assert _stop(outcome).message.where.endswith(f"wiersz {strange}")
    else:
        assert _ordered(outcome).order_count == len(layout.raw)


def test_drugie_wklejenie_tego_samego_dnia_mowi_ze_dzien_jest_uporzadkowany_i_nic_nie_zmienia(tmp_path: Path) -> None:
    path, layout = _register(tmp_path)
    original = path.read_bytes()
    _ordered(_run(tmp_path))
    first_raw = _pasted_again(path, original, layout)

    with _nothing_written(tmp_path, path):
        stop = _stop(_run(tmp_path, now=NOW + timedelta(minutes=5)))

    assert (stop.message.code, stop.stop_kind) == ("DZIEN_JUZ_UPORZADKOWANY", "for_tata")
    assert stop.message.what == "Dzień 26.10.2026 jest już uporządkowany. INTERMAX nie został zmieniony."
    assert stop.message.action == "Nowe zlecenia na ten dzień dopisz ręcznie, a wklejone usuń."
    assert stop.message.where.endswith(f"wklejone wiersze {first_raw}–{first_raw + len(layout.raw) - 1}")
    _without_order_data(stop.message, layout.raw)


def test_wklejony_dzien_uporzadkowany_nad_nowszymi_dniami_jest_juz_uporzadkowany(tmp_path: Path) -> None:
    ordered = (
        Order(datetime(2026, 10, 26, 8), "anna.kowalska", 1),
        *(Order(datetime(2026, 10, 27, 7 + index), "jan.nowak", index + 2) for index in range(12)),
    )
    path, _ = _register(tmp_path, Intermax(ordered=ordered))

    with _nothing_written(tmp_path, path):
        stop = _stop(_run(tmp_path))

    assert (stop.message.code, stop.message.what) == (
        "DZIEN_JUZ_UPORZADKOWANY",
        "Dzień 26.10.2026 jest już uporządkowany. INTERMAX nie został zmieniony.",
    )


def test_zlecenie_na_jutro_w_srodku_dzisiejszego_dnia_z_jego_wierszami_pod_nim_nie_czyni_jutra_uporzadkowanym(tmp_path: Path) -> None:
    def day(number: int, count: int, first: int) -> tuple[Order, ...]:
        return tuple(Order(datetime(2026, 10, number, 7 + index), "anna.kowalska", first + index) for index in range(count))

    path, layout = _register(tmp_path, Intermax(ordered=(*day(25, 6, 1), *day(26, 1, 7), *day(25, 4, 8))))

    assert _ordered(_run(tmp_path)).order_count == len(layout.raw)


def test_uporzadkowany_wiersz_bez_daty_nad_wklejeniem_nie_zatrzymuje(tmp_path: Path) -> None:
    path, layout = _register(tmp_path)
    _edited(path, lambda sheet: setattr(sheet.cell(layout.last_ordered, DATE), "value", None))

    assert _ordered(_run(tmp_path)).order_count == len(layout.raw)


def test_dzien_jest_uporzadkowany_takze_z_usterka_poprzedniego_dnia_dopisana_pod_nim(tmp_path: Path) -> None:
    path, layout = _register(tmp_path, Intermax(grid=False))
    original = path.read_bytes()
    _ordered(_run(tmp_path))
    first_raw = _pasted_again(path, original, layout)
    _edited(path, lambda sheet: _ordered_row(Order(datetime(2026, 10, 25, 15), "jan.nowak", 401))(sheet, first_raw - 2))

    with _nothing_written(tmp_path, path):
        stop = _stop(_run(tmp_path, now=NOW + timedelta(minutes=5)))

    assert stop.message.code == "DZIEN_JUZ_UPORZADKOWANY"


def test_recznie_przesuniety_pierwszy_wiersz_jest_surowym_zleceniem_i_nie_przesuwa_sie_drugi_raz(tmp_path: Path) -> None:
    path, layout = _register(tmp_path)
    _edited(path, lambda sheet: sheet.move_range(f"A{layout.first_raw}:{get_column_letter(APP_LAST - 2)}{layout.first_raw}", cols=2))

    summary = _ordered(_run(tmp_path))

    assert summary.order_count == len(layout.raw)
    assert layout.yellow is not None
    assert order_rows(path, layout.yellow + 9, layout.yellow + 9) == [_expected(layout.raw[0])]


def test_formula_we_wklejeniu_to_zatrzymanie_do_maksa_bez_zapisu(tmp_path: Path) -> None:
    path, layout = _register(tmp_path)
    _edited(path, lambda sheet: setattr(sheet.cell(layout.first_raw + 1, 4), "value", "=1+1"))

    with _nothing_written(tmp_path, path):
        stop = _stop(_run(tmp_path))

    assert (stop.message.code, stop.stop_kind) == ("NIEOBSLUGIWANY_UKLAD_INTERMAX", "for_maks")


def test_hiperlacze_w_pustym_wierszu_pod_dniem_to_zatrzymanie_do_maksa_bez_zapisu(tmp_path: Path) -> None:
    path, layout = _register(tmp_path)
    assert layout.yellow is not None
    first_link = rb'(<hyperlink [^>]*?ref=")I2("[^>]*>)'
    blank_link = rb"\g<0>\g<1>I%d\g<2>" % (layout.yellow + 1)
    rewrite_parts(path, lambda name, data: re.sub(first_link, blank_link, data, count=1) if name == _SHEET else data)

    with _nothing_written(tmp_path, path):
        stop = _stop(_run(tmp_path))

    assert (stop.message.code, stop.stop_kind) == ("NIEOBSLUGIWANY_UKLAD_INTERMAX", "for_maks")


def test_arkusz_bez_dimension_dostaje_dzien_bez_dimension(tmp_path: Path) -> None:
    path, layout = _register(tmp_path)
    rewrite_parts(path, lambda name, data: re.sub(rb"<dimension [^>]*>", b"", data) if name == _SHEET else data)

    _ordered(_run(tmp_path))

    assert b"<dimension" not in parts(path)[_SHEET]
    assert layout.yellow is not None
    assert order_rows(path, layout.yellow + 2, layout.yellow + 2) == [_expected(layout.raw[1])]


def _style_entries(xml: bytes, tag: bytes, item: bytes) -> tuple[int, list[bytes]]:
    found = re.search(rb"<" + tag + rb'\b[^>]*?count="(\d+)"[^>]*>(.*?)</' + tag + rb">", xml, re.S)
    assert found is not None, tag
    entries = re.findall(rb"<" + item + rb"\b(?:[^>]*/>|[^>]*>.*?</" + item + rb">)", found.group(2), re.S)
    return int(found.group(1)), entries


_LISTS = ((b"fonts", b"font"), (b"fills", b"fill"), (b"borders", b"border"), (b"cellXfs", b"xf"))
_COLORED = Intermax(
    raw=(
        Order(datetime(2026, 10, 26, 14), "jan.nowak", 101, "TMOBILE"),
        Order(datetime(2026, 10, 26, 10), "anna.kowalska", 102, "NETIA"),
        Order(datetime(2026, 10, 26, 16), "anna.kowalska", 103, "T-Mobile"),
        Order(datetime(2026, 10, 26, 8), "jan.nowak", 104, None, "Usterka"),
        Order(datetime(2026, 10, 26, 9), "bartek.zielinski", 105, "P4"),
    )
)
"""Surowe zlecenia z Kolorami operatora, które powtarzają wypełnienia."""


def test_dopisane_style_sa_raz_kazdy_a_licznik_count_zgadza_sie_z_lista(tmp_path: Path) -> None:
    path, _ = _register(tmp_path, _COLORED)
    before = parts(path)[_STYLES]

    _ordered(_run(tmp_path))

    after = parts(path)[_STYLES]
    for tag, item in _LISTS:
        count, entries = _style_entries(after, tag, item)
        added = entries[len(_style_entries(before, tag, item)[1]) :]
        assert count == len(entries), tag
        assert len(set(added)) == len(added), tag
    assert len(_style_entries(after, b"cellXfs", b"xf")[1]) > len(_style_entries(before, b"cellXfs", b"xf")[1])
    # Zielony, ciemniejszy i jasny niebieski oraz różowy, każdy raz.
    assert len(_style_entries(after, b"fills", b"fill")[1]) == len(_style_entries(before, b"fills", b"fill")[1]) + 4


def _pasted_again(path: Path, original: bytes, layout: Layout, *, days: int = 0) -> int:
    """Te same wiersze wklejenia co w ``original`` doklejone w XML pod ostatni wiersz ``path``, jak drugie wklejenie.

    Daty zleceń są ``days`` dni później. Zwraca pierwszy wiersz nowego wklejenia.
    """

    numbers = range(layout.first_raw, layout.first_raw + len(layout.raw))
    source = zipfile.ZipFile(BytesIO(original)).read(_SHEET)
    rows = [re.search(rb'<row r="%d"[^>]*>.*?</row>' % number, source, re.S) for number in numbers]
    last = max(int(number) for number in re.findall(rb'<row r="(\d+)"', parts(path)[_SHEET]))

    def renumbered(index: int, row: re.Match[bytes] | None) -> bytes:
        assert row is not None
        number = b"%d" % (last + 3 + index)
        cells = re.sub(rb'(<c r="[A-Z]+)\d+"', rb'\g<1>' + number + b'"', row.group(0))
        cells = re.sub(
            rb'(<c r="[AC]\d+"(?:(?! t=")[^>])*><v>)([\d.]+)(</v>)',
            lambda date: date.group(1) + b"%r" % (float(date.group(2)) + days) + date.group(3),
            cells,
        )
        return re.sub(rb'<row r="\d+"', b'<row r="' + number + b'"', cells, count=1)

    pasted = b"".join(renumbered(index, row) for index, row in enumerate(rows))
    rewrite_parts(path, lambda name, data: data.replace(b"</sheetData>", pasted + b"</sheetData>") if name == _SHEET else data)
    return last + 3


def test_kolor_operatora_bierze_wypelnienie_ktore_tata_juz_ma_w_intermax(tmp_path: Path) -> None:
    path, layout = _register(tmp_path, Intermax(raw=(Order(datetime(2026, 10, 26, 8), "jan.nowak", 101, "TMOBILE"),)))
    green = PatternFill(patternType="solid", fgColor=Color(rgb=_GREEN), bgColor=Color(indexed=64))
    _edited(path, lambda sheet: setattr(sheet.cell(2, ACTIVITIES), "fill", green))
    before = _style_entries(parts(path)[_STYLES], b"fills", b"fill")

    _ordered(_run(tmp_path))

    assert layout.yellow is not None
    assert _fills(load_workbook(path).worksheets[1], layout.yellow + 2)[ACTIVITIES] == _GREEN
    assert _style_entries(parts(path)[_STYLES], b"fills", b"fill") == before


def test_kolejne_uporzadkowanie_nie_dopisuje_stylow_ktore_juz_sa(tmp_path: Path) -> None:
    path, layout = _register(tmp_path, replace(_COLORED, grid=False))
    original = path.read_bytes()
    _ordered(_run(tmp_path))
    first = parts(path)[_STYLES]
    _pasted_again(path, original, layout, days=1)

    assert _ordered(_run(tmp_path, now=NOW + timedelta(days=1))).order_count == len(layout.raw)

    second = parts(path)[_STYLES]
    for tag, item in _LISTS:
        assert _style_entries(second, tag, item) == _style_entries(first, tag, item), tag


def _thin_row(sheet: Worksheet, number: int) -> tuple[float | None, str | None]:
    look = sheet.row_dimensions[number]
    return look.height, look.border.left.style


def _row_look(path: Path, number: int) -> dict[bytes, bytes]:
    """Atrybuty ``<row>`` wiersza ``number`` arkusza zleceń bez numeru i ``spans``."""

    tag = re.match(rb"<row\b[^>]*", _row_xml(path, number))
    assert tag is not None
    attributes = dict(re.findall(rb'\s([\w:]+)="([^"]*)"', tag.group(0)))
    return {name: value for name, value in attributes.items() if name not in (b"r", b"spans")}


@pytest.mark.parametrize("grid", [True, False])
def test_wiersze_nowego_dnia_maja_wysokosc_i_styl_wierszy_poprzednich_dni(tmp_path: Path, grid: bool) -> None:
    path, layout = _register(tmp_path, Intermax(gap=10, grid=grid))
    assert layout.yellow is not None
    yellow = layout.yellow
    order_look = _row_look(path, layout.last_ordered)
    blank_look = _row_look(path, layout.last_ordered + 1) if grid else order_look
    kept = {number: _row_xml(path, number) for number in (yellow + 1, yellow + 10)} if grid else {}

    _ordered(_run(tmp_path))

    sheet = load_workbook(path)[ORDER_SHEET]
    for number in (yellow + 2, yellow + 3, yellow + 4, yellow + 6, yellow + 8, yellow + 9):
        assert _thin_row(sheet, number) == (ORDER_HEIGHT, "thin"), number
        assert _row_look(path, number) == order_look, number
    for number in (yellow + 5, yellow + 7):
        assert _thin_row(sheet, number) == (ORDER_HEIGHT, "thin"), number
        assert _row_look(path, number) == order_look, number
        assert {sheet.cell(number, column).border.left.style for column in range(FIRST, LAST + 1)} == {"thin"}, number
    blank_height = BLANK_HEIGHT if grid else ORDER_HEIGHT
    vacated = range(layout.first_raw, layout.first_raw + len(layout.raw))
    for number in (yellow + 1, yellow + 10, *vacated):
        assert _thin_row(sheet, number) == (blank_height, "thin"), number
        assert _row_look(path, number) == blank_look, number
    assert order_rows(path, yellow + 10, vacated[-1]) == [None] * (vacated[-1] - yellow - 9)
    assert {number: _row_xml(path, number) for number in kept} == kept, "pusty wiersz taty zostaje bajt w bajt"


def test_pusty_wiersz_pod_dniem_ma_wyglad_pustego_wiersza_taty_a_nie_zoltego(tmp_path: Path) -> None:
    path, layout = _register(tmp_path, Intermax(above_yellow=0, gap=10))
    blank_look = _row_look(path, layout.first_raw - 1)

    _ordered(_run(tmp_path))

    for number in range(layout.first_raw, layout.first_raw + len(layout.raw)):
        assert _row_look(path, number) == blank_look, number


def test_zolty_wiersz_pod_wklejeniem_nie_wyznacza_miejsca_dnia(tmp_path: Path) -> None:
    path, layout = _register(tmp_path)
    below = layout.first_raw + len(layout.raw) + 1
    _edited(path, lambda sheet: setattr(sheet.row_dimensions[below], "fill", YELLOW))

    _ordered(_run(tmp_path))

    assert layout.yellow is not None
    assert _first_day_row(path, layout, layout.raw[1]) == layout.yellow + 2


def test_styl_wiersza_bez_customformat_nie_czyni_wiersza_zoltym(tmp_path: Path) -> None:
    path, layout = _register(tmp_path)
    assert layout.yellow is not None
    yellow_row = rb'(<row r="%d"[^>]*?) customFormat="1"' % layout.yellow
    rewrite_parts(path, lambda name, data: re.sub(yellow_row, rb"\g<1>", data) if name == _SHEET else data)

    _ordered(_run(tmp_path))

    assert _first_day_row(path, layout, layout.raw[1]) == layout.last_ordered + 11


def test_kopia_intermax_powstaje_przed_podmiana_intermax(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    path, _ = _register(tmp_path)
    original = path.read_bytes()
    copy = tmp_path / "Kopie INTERMAX" / COPY_NAME
    replace_file = os.replace
    copied_before: list[bool] = []

    def replace(source: str | os.PathLike[str], target: str | os.PathLike[str]) -> None:
        if Path(target).name == REGISTER_NAME:
            copied_before.append(copy.is_file() and copy.read_bytes() == original)
        replace_file(source, target)

    monkeypatch.setattr(os, "replace", replace)

    _ordered(_run(tmp_path))

    assert copied_before == [True]


def test_kopia_o_tej_samej_nazwie_zostaje_a_intermax_nie_jest_zmieniony(tmp_path: Path) -> None:
    path, _ = _register(tmp_path)
    copies = tmp_path / "Kopie INTERMAX"
    copies.mkdir()
    (copies / COPY_NAME).write_bytes(b"kopia sprzed minuty")

    with _nothing_written(tmp_path, path):
        stop = _stop(_run(tmp_path))

    assert (stop.message.code, stop.stop_kind) == ("KOPIA_INTERMAX_NIEUDANA", "for_tata")
    assert (copies / COPY_NAME).read_bytes() == b"kopia sprzed minuty"


class _FullDisk:
    """Plik Kopii, którego zapis kończy się brakiem miejsca na dysku po pierwszych bajtach."""

    def __init__(self, descriptor: int) -> None:
        self._file = open(descriptor, "wb")

    def __enter__(self) -> _FullDisk:
        return self

    def __exit__(self, *_exc: object) -> None:
        self._file.close()

    def write(self, data: bytes) -> int:
        self._file.write(data[:100])
        raise OSError(errno.ENOSPC, "No space left on device")


def test_niepelna_kopia_intermax_nie_zostaje_a_intermax_nie_jest_zmieniony(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path, _ = _register(tmp_path)
    monkeypatch.setattr(os, "fdopen", lambda descriptor, _mode="r": _FullDisk(descriptor))

    with _nothing_written(tmp_path, path):
        stop = _stop(_run(tmp_path))

    assert (stop.message.code, stop.stop_kind) == ("KOPIA_INTERMAX_NIEUDANA", "for_tata")


def test_folder_kopii_ktorego_nie_da_sie_utworzyc_to_zatrzymanie_dla_taty_bez_zapisu(tmp_path: Path) -> None:
    path, _ = _register(tmp_path)
    (tmp_path / "Kopie INTERMAX").write_bytes(b"plik zamiast folderu")
    original = path.read_bytes()

    stop = _stop(_run(tmp_path))

    assert (stop.message.code, stop.stop_kind) == ("KOPIA_INTERMAX_NIEUDANA", "for_tata")
    assert path.read_bytes() == original
    assert [entry.name for entry in path.parent.iterdir()] == [REGISTER_NAME]


def _refused(
    monkeypatch: pytest.MonkeyPatch, times: int | None, *, excel: bool = False, code: int = errno.EACCES
) -> list[float]:
    """Podmiana INTERMAX odmawiana ``times`` razy (``None``: zawsze), z plikiem blokady Excela (``excel``) albo bez.

    ``code``: błąd systemu przy podmianie; domyślnie odmowa dostępu.

    Zwraca listę, do której trafia liczba prób.
    """

    replace_file = os.replace
    attempts: list[float] = []

    def refused(source: str | os.PathLike[str], target: str | os.PathLike[str]) -> None:
        if Path(target).name == REGISTER_NAME:
            attempts.append(1)
            assert len(attempts) < 100, "ponowna próba bez przerwy na zegarze procesu"
            if excel:
                Path(target).with_name(f"~${REGISTER_NAME}").write_bytes(b"")
            if times is None or len(attempts) <= times:
                raise OSError(code, os.strerror(code))
        replace_file(source, target)

    monkeypatch.setattr(os, "replace", refused)
    return attempts


def test_odmowa_podmiany_bez_excela_po_krotkiej_ponownej_probie_to_onedrive_dla_taty(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path, layout = _register(tmp_path)
    attempts = _refused(monkeypatch, None)
    clock = ManualClock(seconds=4000.0)

    with _nothing_written(tmp_path, path):
        stop = _stop(_run(tmp_path, clock=clock))

    assert (stop.message.code, stop.stop_kind) == ("INTERMAX_SYNCHRONIZOWANY", "for_tata")
    assert stop.message.what == "OneDrive jeszcze synchronizuje INTERMAX. INTERMAX nie został zmieniony."
    assert stop.message.action == "Spróbuj za chwilę: poczekaj, aż OneDrive skończy synchronizację, i uruchom ponownie."
    assert (clock.seconds, len(attempts)) == (4005.0, 11)
    _without_order_data(stop.message, layout.raw)


def test_podmiana_po_chwilowej_odmowie_onedrive_porzadkuje_intermax(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    path, layout = _register(tmp_path)
    original = path.read_bytes()
    attempts = _refused(monkeypatch, 2)
    clock = ManualClock(seconds=4000.0)

    assert _ordered(_run(tmp_path, clock=clock)).order_count == len(layout.raw)

    assert (clock.seconds, len(attempts)) == (4001.0, 3)
    assert path.read_bytes() != original
    assert (tmp_path / "Kopie INTERMAX" / COPY_NAME).read_bytes() == original


class _PauseThen(ManualClock):
    """Zegar procesu, w którego przerwie przed ponowną próbą podmiany dzieje się ``meanwhile``."""

    def __init__(self, meanwhile: Callable[[], object]) -> None:
        super().__init__()
        self._meanwhile = meanwhile

    def pause(self, seconds: float) -> None:
        super().pause(seconds)
        self._meanwhile()


def test_wersja_pracownika_wgrana_przez_onedrive_miedzy_probami_podmiany_zostaje_a_tata_slyszy_o_synchronizacji(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path, _ = _register(tmp_path)
    worker = Intermax()
    worker.ordered = (*worker.ordered, Order(datetime(2026, 10, 20, 11), "jan.nowak", 4))
    attempts = _refused(monkeypatch, 1)
    uploaded: list[bytes] = []

    def upload() -> None:
        worker.write(path)
        uploaded.append(path.read_bytes())

    stop = _stop(_run(tmp_path, clock=_PauseThen(upload)))

    assert (stop.message.code, stop.stop_kind) == ("INTERMAX_SYNCHRONIZOWANY", "for_tata")
    assert len(attempts) == 1
    assert uploaded == [path.read_bytes()]
    assert [entry.name for entry in path.parent.iterdir()] == [REGISTER_NAME]
    assert _copies(tmp_path) == set()


def _unreadable(path: Path, patch: pytest.MonkeyPatch) -> None:
    """System nie oddaje bajtów ``path`` (OneDrive trzyma plik w trakcie synchronizacji)."""

    read_bytes = Path.read_bytes

    def refused(target: Path) -> bytes:
        if target == path:
            raise PermissionError(errno.EACCES, os.strerror(errno.EACCES))
        return read_bytes(target)

    patch.setattr(Path, "read_bytes", refused)


@pytest.mark.parametrize(
    ("meanwhile", "code"),
    [
        (lambda path, _patch: path.with_name(f"~${REGISTER_NAME}").write_bytes(b""), "INTERMAX_OTWARTY"),
        (_unreadable, "INTERMAX_SYNCHRONIZOWANY"),
    ],
    ids=["tata otwiera INTERMAX w Excelu", "INTERMAX niedostępny"],
)
def test_intermax_nieczytelny_w_przerwie_miedzy_probami_podmiany_to_zatrzymanie_dla_taty_bez_zapisu(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    meanwhile: Callable[[Path, pytest.MonkeyPatch], object],
    code: str,
) -> None:
    path, _ = _register(tmp_path)
    attempts = _refused(monkeypatch, 1)

    with _nothing_written(tmp_path, path), monkeypatch.context() as during:
        stop = _stop(_run(tmp_path, clock=_PauseThen(lambda: meanwhile(path, during))))

    assert (stop.message.code, stop.stop_kind) == (code, "for_tata")
    assert len(attempts) == 1


def test_odmowa_podmiany_przy_otwartym_excelu_to_zatrzymanie_dla_taty_bez_czekania(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path, _ = _register(tmp_path)
    attempts = _refused(monkeypatch, None, excel=True)
    clock = ManualClock()

    with _nothing_written(tmp_path, path):
        stop = _stop(_run(tmp_path, clock=clock))

    assert (stop.message.code, stop.stop_kind) == ("INTERMAX_OTWARTY", "for_tata")
    assert (clock.seconds, len(attempts)) == (0.0, 1)


def test_inny_blad_podmiany_niz_odmowa_to_zatrzymanie_dla_taty_bez_czekania(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path, _ = _register(tmp_path)
    attempts = _refused(monkeypatch, None, code=errno.ENOSPC)
    clock = ManualClock()

    with _nothing_written(tmp_path, path):
        stop = _stop(_run(tmp_path, clock=clock))

    assert (stop.message.code, stop.stop_kind) == ("INTERMAX_OTWARTY", "for_tata")
    assert (clock.seconds, len(attempts)) == (0.0, 1)


def test_intermax_otwarty_w_excelu_w_trakcie_zapisu_to_zatrzymanie_dla_taty_bez_zapisu(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path, _ = _register(tmp_path)
    write_bytes = Path.write_bytes

    def opened_meanwhile(target: Path, data: bytes) -> int:
        if target.parent == path.parent and target.name.startswith(f".{path.stem}."):
            write_bytes(path.with_name(f"~${path.name}"), b"")
        return write_bytes(target, data)

    monkeypatch.setattr(Path, "write_bytes", opened_meanwhile)

    with _nothing_written(tmp_path, path):
        stop = _stop(_run(tmp_path))

    assert (stop.message.code, stop.stop_kind) == ("INTERMAX_OTWARTY", "for_tata")


@pytest.mark.parametrize(("content", "code"), [(None, "BRAK_INTERMAX"), (b"to nie jest skoroszyt", "NIECZYTELNY_INTERMAX")])
def test_brak_albo_nieczytelny_intermax_to_zatrzymanie_dla_taty(tmp_path: Path, content: bytes | None, code: str) -> None:
    folder = tmp_path / "OneDrive"
    folder.mkdir()
    if content is not None:
        (folder / REGISTER_NAME).write_bytes(content)

    outcome = _run(tmp_path)

    stop = _stop(outcome)
    assert (stop.message.code, stop.stop_kind) == (code, "for_tata")
    stopped = ProgressPhase.CHECKING if content is None else ProgressPhase.READING
    assert outcome.progress.phase_states[stopped] == PhaseState.FAILED
    assert sorted(entry.name for entry in folder.iterdir()) == ([] if content is None else [REGISTER_NAME])
    assert not _copies(tmp_path)


@pytest.mark.parametrize("header", ["AKCJE", "UWAGI", "WYKONAWCA"])
def test_bez_naglowka_zlecen_w_wierszu_1_zatrzymanie_dla_taty_bez_zapisu(tmp_path: Path, header: str) -> None:
    path, _ = _register(tmp_path)
    _edited(path, lambda sheet: setattr(sheet.cell(1, FIRST + HEADERS.index(header)), "value", f"{header} X"))

    with _nothing_written(tmp_path, path):
        stop = _stop(_run(tmp_path))

    assert (stop.message.code, stop.stop_kind) == ("BRAK_ARKUSZA_ZLECEN", "for_tata")


def test_nieudane_usuniecie_starszej_kopii_to_ostrzezenie_przy_uporzadkowanym_intermax(tmp_path: Path) -> None:
    path, _ = _register(tmp_path)
    original = path.read_bytes()
    copies = tmp_path / "Kopie INTERMAX"
    copies.mkdir()
    for day in range(10, 20):
        (copies / f"INTERMAX przed uporządkowaniem 2026-10-{day:02d} 07-00-00.xlsx").write_bytes(b"kopia")
    oldest = copies / "INTERMAX przed uporządkowaniem 2026-10-09 07-00-00.xlsx"
    oldest.mkdir()  # usunięcie dostaje błąd systemu plików, jak Kopia otwarta w Excelu na Windows

    summary = _ordered(_run(tmp_path))

    assert [warning.code for warning in summary.warnings] == ["STARE_KOPIE_INTERMAX"]
    assert "uporządkowany" in summary.warnings[0].what
    assert path.read_bytes() != original
    assert (copies / COPY_NAME).read_bytes() == original


def test_podsumowanie_ktorego_nie_da_sie_pokazac_mowi_ze_intermax_jest_uporzadkowany(tmp_path: Path) -> None:
    _register(tmp_path)
    inputs = OrderingInputs(tmp_path / "OneDrive" / REGISTER_NAME, tmp_path / "Kopie INTERMAX", lambda: NOW)

    def broken(_inputs: OrderingInputs, _summary: OrderingSummary) -> OrderingSummary:
        raise RuntimeError("widok")

    outcome = run_process(PROCESS, inputs, Tata(), view=broken, metrics_path=tmp_path / "m.jsonl", clock=ManualClock())

    stop = _stop(outcome)
    assert (stop.kind, stop.message.code, stop.stop_kind) == ("after_engine", "BLAD_PODSUMOWANIA", "for_maks")


def test_kopie_intermax_trafiaja_do_folderu_narzedzia_poza_onedrive(tmp_path: Path) -> None:
    assert OrderingInputs(tmp_path / REGISTER_NAME).copies_directory == TOOL_ROOT / "Kopie INTERMAX"


def test_niedostepny_intermax_to_zatrzymanie_dla_taty_bez_zapisu(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    path, _ = _register(tmp_path)
    read_bytes = Path.read_bytes

    def blocked(target: Path) -> bytes:
        if target.name == REGISTER_NAME:
            raise PermissionError(errno.EACCES, "Permission denied")
        return read_bytes(target)

    with _nothing_written(tmp_path, path):
        monkeypatch.setattr(Path, "read_bytes", blocked)
        stop = _stop(_run(tmp_path))
        monkeypatch.undo()

    assert (stop.message.code, stop.stop_kind) == ("NIEDOSTEPNY_INTERMAX", "for_tata")


def test_nieoczekiwany_blad_przed_podmiana_mowi_ze_intermax_nie_zostal_zmieniony(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path, _ = _register(tmp_path)
    write_bytes = Path.write_bytes

    def broken(target: Path, data: bytes) -> int:
        if target.parent == path.parent:
            raise RuntimeError("zapis")
        return write_bytes(target, data)

    monkeypatch.setattr(Path, "write_bytes", broken)

    with _nothing_written(tmp_path, path):
        stop = _stop(_run(tmp_path))

    assert (stop.kind, stop.message.code, stop.stop_kind) == ("unexpected", "BLAD_NIEOCZEKIWANY", "for_maks")
    assert "INTERMAX nie został zmieniony." in stop.message.what


@pytest.mark.parametrize("ending", ["unexpected_failure", "plan_not_shown", "summary_not_shown"])
def test_komunikaty_zakonczen_uporzadkowania_maja_trzy_czesci_i_nazywaja_proces(ending: str) -> None:
    message = getattr(PROCESS, ending)

    assert isinstance(message, Message)
    assert message.code.startswith("BLAD_") and message.what and message.action
    assert message.where == "Uporządkowanie INTERMAX"


def test_metryki_uporzadkowania_maja_wlasny_plik_obok_innych_procesow() -> None:
    assert PROCESS.metrics_path() == TOOL_ROOT / ".uporzadkowanie-intermax-metrics.jsonl"


@pytest.mark.parametrize(("rows", "stops"), [("yellow", False), ("above", False), ("whole column", True)])
def test_zakres_za_danymi_arkusza_zatrzymuje_tylko_gdy_siega_pod_zolty_wiersz(
    tmp_path: Path, rows: str, stops: bool
) -> None:
    path, layout = _register(tmp_path)
    assert layout.yellow is not None
    yellow = layout.yellow

    def ranged(sheet: Worksheet) -> None:
        if rows == "yellow":
            sheet.merge_cells(start_row=yellow - 1, start_column=FIRST, end_row=yellow, end_column=FIRST)
        elif rows == "above":
            sheet.merge_cells(start_row=2, start_column=LAST, end_row=3, end_column=LAST)

    if rows == "whole column":
        rule = b'<conditionalFormatting sqref="L:L"><cfRule type="expression" priority="1"><formula>1</formula></cfRule></conditionalFormatting>'
        rewrite_parts(path, lambda name, data: data.replace(b"<hyperlinks>", rule + b"<hyperlinks>") if name == _SHEET else data)
    else:
        _edited(path, ranged)

    outcome = _run(tmp_path)

    if stops:
        assert (_stop(outcome).message.code, _stop(outcome).stop_kind) == ("NIEOBSLUGIWANY_UKLAD_INTERMAX", "for_maks")
    else:
        assert _ordered(outcome).order_count == len(layout.raw)


def test_hiperlacze_na_kilka_komorek_we_wklejeniu_to_zatrzymanie_do_maksa_bez_zapisu(tmp_path: Path) -> None:
    path, layout = _register(tmp_path)
    first = rb'(<hyperlink [^>]*?ref=")I%d(")' % (layout.first_raw + 1)
    wide = rb"\g<1>I%d:I%d\g<2>" % (layout.first_raw + 1, layout.first_raw + 2)
    rewrite_parts(path, lambda name, data: re.sub(first, wide, data, count=1) if name == _SHEET else data)

    with _nothing_written(tmp_path, path):
        stop = _stop(_run(tmp_path))

    assert (stop.message.code, stop.stop_kind) == ("NIEOBSLUGIWANY_UKLAD_INTERMAX", "for_maks")


def test_wiersz_wklejenia_pozostawiony_na_miejscu_to_niezgodnosc_zapisu(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path, layout = _register(tmp_path, Intermax(gap=10))
    raw_row = _row_xml(path, layout.first_raw + 1)
    write_bytes = Path.write_bytes

    def left_in_place(target: Path, data: bytes) -> int:
        if target.parent == path.parent and target.name.startswith(f".{path.stem}."):
            row = rb'<row r="%d"[^>]*?(?:/>|>.*?</row>)' % (layout.first_raw + 1)
            data = _part_changed(bytes(data), _SHEET, lambda xml: re.sub(row, lambda _found: raw_row, xml, count=1, flags=re.S))
        return write_bytes(target, data)

    monkeypatch.setattr(Path, "write_bytes", left_in_place)

    with _nothing_written(tmp_path, path):
        stop = _stop(_run(tmp_path))

    assert (stop.message.code, stop.stop_kind) == ("ZAPIS_INTERMAX_NIEZGODNY", "for_maks")
