from __future__ import annotations

from pathlib import Path
import zipfile

import pytest
from openpyxl import Workbook, load_workbook

import rozliczenia.template_settlement as template_settlement
from rozliczenia.template_settlement import ExcelRow, TemplateWriteError, process_template

from tests.test_settlement_engine import PERIOD, active_worksheet, write_template


def make_template_fixture(tmp_path: Path) -> Path:
    target_directory = tmp_path / PERIOD / f"Rozliczenie pracowników {PERIOD}"
    target_directory.mkdir(parents=True)
    for worker_name in ("Adrian Maciejewski", "Darek Nowak", "Kamil Frontczak"):
        write_template(target_directory / f"Rozliczenie {PERIOD} - {worker_name}.xlsx")
    return target_directory


def synthetic_rows() -> list[ExcelRow]:
    return [("syntetyczne miasto",) + (None,) * 45]


def test_przetworzenie_szablonu_pracownika_zglasza_bledny_naglowek(tmp_path: Path) -> None:
    target_directory = make_template_fixture(tmp_path)
    target_path = target_directory / f"Rozliczenie {PERIOD} - Kamil Frontczak.xlsx"
    workbook = load_workbook(target_path)
    try:
        active_worksheet(workbook)["H17"] = "NIE WYKONAWCA"
        workbook.save(target_path)
    finally:
        workbook.close()

    result = process_template("Kamil Frontczak", target_path, [])

    assert result.status == "ZLY_SZABLON"
    assert [issue.code for issue in result.issues] == ["ZLY_SZABLON"]


def test_przetworzenie_szablonu_pracownika_przerywa_po_bledzie_zapisu(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    target_directory = make_template_fixture(tmp_path)
    target_path = target_directory / f"Rozliczenie {PERIOD} - Adrian Maciejewski.xlsx"
    before = target_path.read_bytes()

    def fail_write(
        path: Path,
        workbook: Workbook,
        rows: list[ExcelRow],
        external_link_parts: dict[str, tuple[zipfile.ZipInfo, bytes]],
    ) -> None:
        raise OSError("synthetic write failure")

    monkeypatch.setattr(template_settlement, "_write_workbook", fail_write)

    with pytest.raises(TemplateWriteError):
        process_template("Adrian Maciejewski", target_path, [("row",)])

    assert target_path.read_bytes() == before


def test_przetworzenie_szablonu_pracownika_przerywa_po_bledzie_snapshotu_linkow(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    target_directory = make_template_fixture(tmp_path)
    target_path = target_directory / f"Rozliczenie {PERIOD} - Adrian Maciejewski.xlsx"
    before = target_path.read_bytes()

    def fail_snapshot(path: Path) -> dict[str, tuple[zipfile.ZipInfo, bytes]]:
        raise OSError("synthetic external link snapshot failure")

    monkeypatch.setattr(template_settlement, "_external_link_parts", fail_snapshot)

    with pytest.raises(TemplateWriteError):
        process_template("Adrian Maciejewski", target_path, synthetic_rows())

    assert target_path.read_bytes() == before


def test_przetworzenie_szablonu_pracownika_zglasza_blokade_pliku(tmp_path: Path) -> None:
    target_directory = make_template_fixture(tmp_path)
    target_path = target_directory / f"Rozliczenie {PERIOD} - Adrian Maciejewski.xlsx"
    (target_directory / f"~${target_path.name}").touch()

    result = process_template("Adrian Maciejewski", target_path, synthetic_rows())

    assert result.status == "ZABLOKOWANY"
    assert [issue.code for issue in result.issues] == ["PLIK_ZABLOKOWANY"]


def test_przetworzenie_szablonu_pracownika_nie_nadpisuje_danych(tmp_path: Path) -> None:
    target_directory = make_template_fixture(tmp_path)
    target_path = target_directory / f"Rozliczenie {PERIOD} - Adrian Maciejewski.xlsx"
    workbook = load_workbook(target_path)
    try:
        active_worksheet(workbook)["A18"] = "wczesniejsza wartosc"
        workbook.save(target_path)
    finally:
        workbook.close()

    result = process_template("Adrian Maciejewski", target_path, synthetic_rows())

    assert result.status == "ZABLOKOWANY"
    assert [issue.code for issue in result.issues] == ["NADPISANIE_ZABLOKOWANE"]


def test_przetworzenie_pustego_szablonu_pracownika_nic_nie_zapisuje(tmp_path: Path) -> None:
    target_directory = make_template_fixture(tmp_path)
    target_path = target_directory / f"Rozliczenie {PERIOD} - Adrian Maciejewski.xlsx"

    result = process_template("Adrian Maciejewski", target_path, [])

    assert result.status == "PUSTY_SZABLON"
    assert result.issues == ()


def test_przetworzenie_szablonu_pracownika_w_trybie_dry_run_nie_zapisuje(
    tmp_path: Path,
) -> None:
    target_directory = make_template_fixture(tmp_path)
    target_path = target_directory / f"Rozliczenie {PERIOD} - Adrian Maciejewski.xlsx"

    result = process_template("Adrian Maciejewski", target_path, synthetic_rows(), dry_run=True)

    assert result.status == "PLAN"
    assert result.rows == 1
    workbook = load_workbook(target_path, data_only=False)
    try:
        assert active_worksheet(workbook)["A18"].value is None
    finally:
        workbook.close()


def test_przetworzenie_szablonu_pracownika_zapisuje_wiersze_i_zachowuje_formule(
    tmp_path: Path,
) -> None:
    target_directory = make_template_fixture(tmp_path)
    target_path = target_directory / f"Rozliczenie {PERIOD} - Adrian Maciejewski.xlsx"

    result = process_template("Adrian Maciejewski", target_path, synthetic_rows())

    workbook = load_workbook(target_path, data_only=False)
    try:
        sheet = active_worksheet(workbook)
        assert sheet["A18"].value == "syntetyczne miasto"
        assert sheet["AU18"].value == "=N18"
    finally:
        workbook.close()
    assert result.status == "ZAPISANO"
    assert result.rows == 1
