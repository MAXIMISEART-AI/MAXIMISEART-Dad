"""Ilości usług w Pliku zbiorczym z Przygotowania pliku zbiorczego (ADR 0012).

Każdy przypadek to jedno Zlecenie skuteczne syntetycznego Rejestru, a oczekiwane ilości pochodzą z poprawionej
instrukcji taty (spec #268), nie z kodu. Wzorce mają kształt Wzorców taty: zera tylko w N:V i X.
"""

from __future__ import annotations

from copy import copy
from dataclasses import dataclass, replace
from datetime import datetime, timedelta
from pathlib import Path
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from rozliczenia.workbook_files import QUANTITY_COLUMNS

from tests.fixtures.benchmark import COLLECTIVE_TEMPLATE, INTERNAL_TEMPLATE, write_synthetic_templates
from tests.fixtures.preparation import (
    REGISTER_HEADERS,
    Highlight,
    PreparationFixture,
    cell_values,
    make_preparation_fixture,
    register_order,
    run_preparation,
    write_register_sheet,
)
from tests.fixtures.workbooks import active_worksheet, refused_by_copy, replace_cells, rewrite_parts

QUANTITY_LETTERS = [get_column_letter(column) for column in QUANTITY_COLUMNS]

# Linie instrukcji taty o składnikach URZĄDZENIA, dokładnie z jej zapisem (także podwójne spacje), z UMOWA bez druku,
# CZYNNOŚCI bez liczby sztuk i DATA we wtorek rano.
INSTRUCTION_LINES: list[tuple[str, dict[str, int]]] = [
    ("usterka", {"AI": 1}),
    ("usterka hurt", {"AI": 1}),
    ("usterka wew", {"AI": 1}),
    ("zmiana instalacji w lokalu abonenta", {"AM": 1}),
    ("technicolor", {"O": 1, "S": 1, "Z": 1, "AK": 1}),
    ("technicolor globo", {"O": 1, "P": 1, "Z": 1, "AK": 1, "S": 2}),
    ("gateway wifi6extra", {"O": 1, "AG": 1, "AO": 1, "S": 2}),
    ("fut", {"Q": 1, "S": 1}),
    ("globo", {"Q": 1, "S": 1}),
    *((f"{count}*{device}", {"Q": count, "S": count}) for count in range(2, 7) for device in ("fut", "globo")),
    ("wifi6extra", {"AG": 1, "S": 1}),
    ("2*wifi6extra", {"AG": 2, "S": 2}),
    ("3*wifi6extra", {"AG": 3, "S": 3}),
    ("zyxel, wifi6extra", {"O": 1, "AG": 1, "AO": 1, "S": 2}),
    ("gniazdo", {"AE": 1, "AP": 1}),
    ("instalacja ręczna", {"AN": 1}),
    ("nokia", {"Z": 1}),
    ("nokia getewayHomeBox  ", {"O": 1, "S": 1, "Z": 1}),
    ("nokia getewayHomeBox  stb2", {"O": 1, "P": 1, "Z": 1, "S": 2}),
    ("nokia zte", {"O": 1, "S": 1, "Z": 1}),
    ("nokia zte XG", {"O": 1, "S": 1, "Z": 1}),
    ("nokia zte  fut", {"O": 1, "P": 1, "Z": 1, "S": 2}),
    ("nokia zte  fut XG", {"O": 1, "P": 1, "Z": 1, "S": 2}),
    ("nokia zte  2*fut", {"O": 1, "P": 1, "Q": 1, "Z": 1, "S": 3}),
    ("nokia zte  2*fut XG", {"O": 1, "P": 1, "Q": 1, "Z": 1, "S": 3}),
    ("nokia zte  3*fut", {"O": 1, "P": 1, "Z": 1, "Q": 2, "S": 4}),
    ("nokia zte  3*fut XG", {"O": 1, "P": 1, "Z": 1, "Q": 2, "S": 4}),
    ("nokia zte  4*fut", {"O": 1, "P": 1, "Z": 1, "Q": 3, "S": 5}),
    ("nokia zte  4*fut XG", {"O": 1, "P": 1, "Z": 1, "Q": 3, "S": 5}),
    ("nokia zte  5*fut", {"O": 1, "P": 1, "Z": 1, "Q": 4, "S": 6}),
    ("nokia zte  5*fut XG", {"O": 1, "P": 1, "Z": 1, "Q": 4, "S": 6}),
    ("nokia zte  wifi6extra", {"O": 1, "Z": 1, "AG": 1, "S": 2}),
    ("nokia zte  2*wifi6extra", {"O": 1, "Z": 1, "AG": 2, "S": 3}),
    ("nokia zte  3*wifi6extra", {"O": 1, "Z": 1, "AG": 3, "S": 4}),
    ("nokia zte  fut wifi6extra", {"O": 1, "P": 1, "Z": 1, "AG": 1, "S": 3}),
    ("nokia zte  fut 2*wifi6extra", {"O": 1, "P": 1, "Z": 1, "AG": 2, "S": 4}),
    ("nokia zte  2*fut wifi6extra", {"O": 1, "P": 1, "Q": 1, "Z": 1, "AG": 1, "S": 4}),
    ("nokia zte  2*fut 2*wifi6extra", {"O": 1, "P": 1, "Q": 1, "Z": 1, "AG": 2, "S": 5}),
    ("nokia zte  3*fut wifi6extra", {"O": 1, "P": 1, "Z": 1, "AG": 1, "Q": 2, "S": 5}),
    ("nokia zte  3*fut 2*wifi6extra", {"O": 1, "P": 1, "Z": 1, "Q": 2, "AG": 2, "S": 6}),
    ("nokia zte  4*fut wifi6extra", {"O": 1, "P": 1, "Z": 1, "AG": 1, "Q": 3, "S": 6}),
    ("nokia zte  4*fut 2*wifi6extra", {"O": 1, "P": 1, "Z": 1, "Q": 3, "AG": 2, "S": 7}),
]
# Nowe kombinacje znanych składników i zapis różniący się wielkością liter, spacjami, przecinkami i ``XG``.
NEW_COMBINATIONS: list[tuple[str, dict[str, int]]] = [
    ("nokia zte 6*fut", {"O": 1, "P": 1, "Z": 1, "Q": 5, "S": 7}),
    ("nokia zte 3*wifi6extra 2*fut", {"O": 1, "P": 1, "Q": 1, "Z": 1, "AG": 3, "S": 6}),
    ("technicolor 2*globo", {"O": 1, "P": 1, "Q": 1, "Z": 1, "AK": 1, "S": 3}),
    ("fut wifi6extra", {"Q": 1, "AG": 1, "S": 2}),
    ("gateway", {"O": 1, "S": 1}),
    ("gateway 2*wifi6extra", {"O": 1, "AG": 2, "AO": 1, "S": 3}),
    ("NOKIA  ZTE", {"O": 1, "S": 1, "Z": 1}),
    ("Nokia Zte Fut", {"O": 1, "P": 1, "Z": 1, "S": 2}),
    ("nokia zteXG", {"O": 1, "S": 1, "Z": 1}),
    ("nokiaXG  zte  2*fut", {"O": 1, "P": 1, "Q": 1, "Z": 1, "S": 3}),
    ("nokia zte 2* fut", {"O": 1, "P": 1, "Q": 1, "Z": 1, "S": 3}),
    ("nokia zte  4*fut, 2*wifi6extra", {"O": 1, "P": 1, "Z": 1, "Q": 3, "AG": 2, "S": 7}),
    ("TECHNICOLOR", {"O": 1, "S": 1, "Z": 1, "AK": 1}),
]
# Instalacja ręczna i wizja: liczba sztuk z CZYNNOŚCI w AN i V, bez liczby 1.
PIECES: list[tuple[str, str, int]] = [
    ("instalacja ręczna", "6 szt", 6),
    ("instalacja ręczna", "instalacja 7szt", 7),
    ("instalacja ręczna", "instalacja", 1),
    ("wizja", "6 szt", 6),
    ("wizja", "instalacja", 1),
    ("Instalacja  Ręczna", "2 SZT.", 2),
]
# Dodatek 45 zł (N = 45) za usterkę pon–pt od 17:00 albo w sobotę, z godziny w DATA. Okres 08–14.09.2026 zaczyna się
# we wtorek; w niedzielę firma nie pracuje.
FAULT = {"AI": 1}
LATE: list[tuple[str, datetime, dict[str, int]]] = [
    ("usterka", datetime(2026, 9, 8, 16, 0), FAULT),
    ("usterka", datetime(2026, 9, 8, 16, 59), FAULT),
    ("usterka", datetime(2026, 9, 8, 17, 0), FAULT | {"N": 45}),
    ("usterka hurt", datetime(2026, 9, 11, 23, 30), FAULT | {"N": 45}),
    ("usterka wew", datetime(2026, 9, 14, 18, 0), FAULT | {"N": 45}),
    ("usterka", datetime(2026, 9, 12, 8, 0), FAULT | {"N": 45}),
    ("usterka hurt", datetime(2026, 9, 12, 0, 0), FAULT | {"N": 45}),
    ("usterka", datetime(2026, 9, 13, 10, 0), FAULT),
    ("usterka, gniazdo", datetime(2026, 9, 12, 10, 0), FAULT | {"AE": 1, "AP": 1, "N": 45}),
    ("gniazdo", datetime(2026, 9, 12, 8, 0), {"AE": 1, "AP": 1}),
    ("nokia zte", datetime(2026, 9, 9, 19, 0), {"O": 1, "S": 1, "Z": 1}),
]
# Druk dokumentów (T) z pierwszego słowa UMOWA.
CONTRACTS: list[tuple[str | None, int]] = [
    ("u", 1),
    ("U", 1),
    ("u z pro", 1),
    (" U ", 1),
    ("UM-SYN-001", 0),
    ("pr", 0),
    ("prot", 0),
    (None, 0),
]
# Zlecenia rozliczane z opisu: bez żadnych ilości, także bez V i T.
FROM_DESCRIPTION: list[str | None] = ["prewencja", "Prewencja", "Niedostępność Usług", None, "   "]
# Składnik, którego Reguły ilości nie znają: tylko V, T i N, reszta dla taty.
WITHOUT_RULE: list[str] = ["dekoder", "przeniesienie", "nokia zte przeniesienie", "nokia zte 2x fut", "zte", "prewencja fut"]


@dataclass(frozen=True)
class Case:
    """Zlecenie skuteczne syntetycznego Rejestru i Ilości usług różne od zera, których jego wiersz oczekuje."""

    devices: str | None
    contract: str | None
    quantities: dict[str, int]
    activities: str = "instalacja"
    started: datetime | None = None
    """``DATA``; bez niej wtorek rano."""


CASES: list[Case] = [
    *(Case(devices, "pr", quantities | {"V": 1}) for devices, quantities in INSTRUCTION_LINES + NEW_COMBINATIONS),
    *(
        Case("nokia zte", contract, {"O": 1, "S": 1, "Z": 1, "V": 1} | ({"T": 1} if t else {}))
        for contract, t in CONTRACTS
    ),
    *(Case(devices, "pr", {"AN": pieces, "V": pieces}, activities) for devices, activities, pieces in PIECES),
    *(Case(devices, "pr", quantities | {"V": 1}, started=started) for devices, started, quantities in LATE),
    *(Case(devices, "u", {}) for devices in FROM_DESCRIPTION),
    Case("prewencja", "u", {}, started=datetime(2026, 9, 12, 9, 0)),
    *(Case(devices, "u", {"V": 1, "T": 1}) for devices in WITHOUT_RULE),
    Case("dekoder", "pr", {"V": 1}),
    Case("usterka przeniesienie", "u", {"N": 45, "V": 1, "T": 1}, started=datetime(2026, 9, 8, 18, 0)),
    Case("2*przeniesienie usterka wew", "pr", {"N": 45, "V": 1}, started=datetime(2026, 9, 12, 8, 0)),
    Case("przeniesienie", "u", {"V": 1, "T": 1}, started=datetime(2026, 9, 8, 18, 0)),
    Case("usterka przeniesienie", "u", {"V": 1, "T": 1}, started=datetime(2026, 9, 8, 16, 0)),
]


def _fixture(tmp_path: Path, cases: list[Case], highlight: Highlight | None = None) -> PreparationFixture:
    """Rejestr z jednym Zleceniem skutecznym na przypadek; z ``highlight`` każda komórka zlecenia jest wyróżniona."""

    start = datetime(2026, 9, 8, 8, 0)
    rows: list[dict[str, object] | None] = [
        register_order(number, case.started or start + timedelta(minutes=number))
        | {"URZĄDZENIA": case.devices, "UMOWA": case.contract, "CZYNNOŚCI": case.activities}
        for number, case in enumerate(cases, start=1)
    ]
    templates = write_synthetic_templates(tmp_path / "wzorce")
    (tmp_path / "okres").mkdir()
    fixture = replace(
        make_preparation_fixture(tmp_path / "okres", rows),
        internal_template_path=templates / INTERNAL_TEMPLATE,
        collective_template_path=templates / COLLECTIVE_TEMPLATE,
    )
    if highlight is not None:
        workbook = Workbook()
        sheet = active_worksheet(workbook)
        sheet.title = "ZLECENIA"
        marked = {(row, header): highlight for row in range(2, 2 + len(rows)) for header in REGISTER_HEADERS if header}
        write_register_sheet(sheet, rows, highlights=marked)
        workbook.save(fixture.register_path)
    return fixture


def _quantity_rows(path: Path, count: int) -> list[dict[str, object]]:
    workbook = load_workbook(path)
    try:
        sheet = active_worksheet(workbook)
        return [
            {letter: sheet[f"{letter}{row}"].value for letter in QUANTITY_LETTERS} for row in range(18, 18 + count)
        ]
    finally:
        workbook.close()


def _looks(path: Path, count: int) -> dict[str, object]:
    workbook = load_workbook(path)
    try:
        sheet = active_worksheet(workbook)
        return {
            cell.coordinate: (
                cell.number_format,
                (cell.font.name, cell.font.sz, bool(cell.font.b), bool(cell.font.i), copy(cell.font.color)),
                copy(cell.fill),
                copy(cell.border),
                copy(cell.alignment),
                copy(cell.protection),
            )
            for row in sheet.iter_rows(min_row=18, max_row=17 + count, min_col=14, max_col=42)
            for cell in row
        }
    finally:
        workbook.close()


def test_plik_zbiorczy_ma_ilosci_uslug_z_kazdej_linii_instrukcji_taty(tmp_path: Path) -> None:
    fixture = _fixture(tmp_path, CASES)

    preview = run_preparation(fixture, dry_run=True)
    summary = run_preparation(fixture)

    assert replace(preview, phase_durations_ms={}) == replace(summary, phase_durations_ms={})
    written = _quantity_rows(fixture.collective_file_path, len(CASES))
    template = _quantity_rows(fixture.collective_template_path, len(CASES))
    assert {value for row in template for value in row.values()} == {0, None}
    assert [row["O"] for row in (template[0], template[-1])] == [0, None]
    expected = [cells | case.quantities for cells, case in zip(template, CASES, strict=True)]
    mismatched = [
        (case, {letter: (row[letter], want[letter]) for letter in row if row[letter] != want[letter]})
        for case, row, want in zip(CASES, written, expected, strict=True)
        if row != want
    ]
    assert mismatched == []
    internal = cell_values(
        fixture.internal_settlement_path, min_row=1, max_row=len(CASES) + 1, min_col=14, max_col=47
    )
    assert {value for row in internal for value in row} == {None}


def test_komorki_ilosci_maja_wyglad_komorek_wzorca_bez_wyroznien_rejestru(tmp_path: Path) -> None:
    cases: list[Case] = [
        Case("nokia zte  2*fut wifi6extra", "u", {"O": 1, "P": 1, "Q": 1, "S": 4, "T": 1, "V": 1, "Z": 1, "AG": 1}),
        Case("gateway 2*wifi6extra", "u", {"O": 1, "S": 3, "T": 1, "V": 1, "AG": 2, "AO": 1}),
        Case("technicolor", "u", {"O": 1, "S": 1, "T": 1, "V": 1, "Z": 1, "AK": 1}),
    ]
    fixture = _fixture(tmp_path, cases, Highlight(fill="FFFF0000", color="FF0000FF", bold=True, italic=True))

    run_preparation(fixture)

    created, template = fixture.collective_file_path, fixture.collective_template_path
    written = _quantity_rows(created, len(cases))
    assert [{letter: value for letter, value in row.items() if value} for row in written] == [
        case.quantities for case in cases
    ]
    assert _looks(created, len(cases)) == _looks(template, len(cases))


def test_ilosci_uslug_z_kopii_wzorca_sa_takie_jak_z_openpyxl(tmp_path: Path) -> None:
    (tmp_path / "kopia").mkdir()
    (tmp_path / "openpyxl").mkdir()
    by_copy = _fixture(tmp_path / "kopia", CASES)
    by_openpyxl = _fixture(tmp_path / "openpyxl", CASES)
    refused_by_copy(by_openpyxl.collective_template_path)

    run_preparation(by_copy)
    run_preparation(by_openpyxl)

    created, reference = by_copy.collective_file_path, by_openpyxl.collective_file_path
    with zipfile.ZipFile(created) as archive:
        assert "xl/calcChain.xml" in archive.namelist()
    assert _quantity_rows(created, len(CASES)) == _quantity_rows(reference, len(CASES))
    assert _looks(created, len(CASES)) == _looks(reference, len(CASES))


def test_ilosc_w_komorce_wzorca_z_formula_zastepuje_formule_bez_sladu_w_calcchain(tmp_path: Path) -> None:
    cases: list[Case] = [Case("nokia zte", "pr", {})]
    fixture = _fixture(tmp_path, cases)
    replace_cells(fixture.collective_template_path, {"V18": '<c r="V18"><f>1-1</f><v>0</v></c>'})
    rewrite_parts(
        fixture.collective_template_path,
        lambda name, data: data.replace(b"</calcChain>", b'<c r="V18"/></calcChain>') if name == "xl/calcChain.xml" else data,
    )

    run_preparation(fixture)

    created = fixture.collective_file_path
    assert _quantity_rows(created, 1)[0]["V"] == 1
    with zipfile.ZipFile(created) as archive:
        names = archive.namelist()
        chain = archive.read("xl/calcChain.xml").decode("utf-8") if "xl/calcChain.xml" in names else ""
    assert 'r="V18"' not in chain
