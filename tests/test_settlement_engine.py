from __future__ import annotations

import errno
import os
import shutil
import tempfile
from pathlib import Path
import time
from collections.abc import Callable

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Border, Color, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet
import pytest

from rozliczenia.domain import (
    PhaseEnded,
    PhaseFailed,
    PhaseStarted,
    ProgressEvent,
    ProgressPhase,
    WorkerEnded,
    WorkerFailed,
    WorkerStarted,
)
from rozliczenia.engine import run_settlements
from rozliczenia.message import StopError
from rozliczenia.process_run import Stop
from rozliczenia.progress import PhaseState, ProgressProjection
from rozliczenia.workbook_files import Layout, TemplateFile, read_data_rows
from tests.fixtures.benchmark import WORKER_TEMPLATE, write_synthetic_templates
from tests.fixtures.porownanie import compare_worker_settlements
from tests.fixtures.workbooks import (
    active_worksheet,
    add_external_link_fixture,
    column_widths,
    cut_sheet_after_data_rows,
    hyperlinks,
    refused_by_copy,
    without_sheet,
)
from tests.fixtures.settlement import (
    NOT_A_PERIOD,
    PERIOD,
    SOURCE_NAME,
    WORKER_TEMPLATE_COLUMN_WIDTHS,
    link_order_numbers,
    make_fixture,
    projected,
    write_source,
)


UNREADABLE_SOURCE_MESSAGE = (
    f"Nie można odczytać Pliku zbiorczego {SOURCE_NAME}. "
    "Sprawdź, czy plik otwiera się w Excelu i nie jest uszkodzony."
)


def overwrite_with_non_workbook(path: Path) -> None:
    path.write_bytes(b"uszkodzony")


UNFIT_COLLECTIVE_FILE_PREFIX = "Plik zbiorczy nie nadaje się do Przebiegu rozliczeń."
UNUSABLE_WORKER_TEMPLATE_PREFIX = "Nie można użyć Wzorca rozliczenia pracownika."


def make_complete_run_fixture(tmp_path: Path) -> tuple[Path, Path, Path, Path]:
    source_path, target_directory, config_path = make_fixture(tmp_path)
    return source_path, target_directory, config_path, tmp_path / "Wzorzec rozliczenia pracownika.xlsx"


@pytest.fixture(scope="session")
def synthetic_worker_template(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """Mała odmiana syntetycznego Wzorca rozliczenia pracownika kształtu Wzorca taty (``tools/syntetyczne_wzorce.py``)."""

    return write_synthetic_templates(tmp_path_factory.mktemp("syntetyczne-wzorce") / "wzorce") / WORKER_TEMPLATE


@pytest.fixture(params=["maly-wzorzec", "syntetyczny-wzorzec"])
def complete_run(
    request: pytest.FixtureRequest, tmp_path: Path, synthetic_worker_template: Path
) -> tuple[Path, Path, Path, Path]:
    """Pełne wejście Przebiegu rozliczeń z małym Wzorcem z fixture albo z małym syntetycznym Wzorcem (#145)."""

    source_path, target_directory, config_path, worker_template_path = make_complete_run_fixture(tmp_path)
    if request.param == "syntetyczny-wzorzec":
        shutil.copyfile(synthetic_worker_template, worker_template_path)
    return source_path, target_directory, config_path, worker_template_path


def expected_source_row(address: str, order_number: str, worker_id: str, quantity: int) -> tuple[object, ...]:
    return (
        "POZNAŃ",
        address,
        None,
        None,
        None,
        order_number,
        None,
        worker_id,
        None,
        None,
        None,
        None,
        None,
        quantity,
    ) + (None,) * 32


def worksheet_input_row(sheet: Worksheet, row_number: int) -> tuple[object, ...]:
    return next(sheet.iter_rows(min_row=row_number, max_row=row_number, max_col=46, values_only=True))


def test_przygotowanie_skoroszytow_pracownikow_publikuje_folder_rozliczen_pracownikow(
    complete_run: tuple[Path, Path, Path, Path],
) -> None:
    source_path, target_directory, config_path, worker_template_path = complete_run
    worker_template_before = worker_template_path.read_bytes()
    template_workbook = load_workbook(worker_template_path)
    try:
        template_sheetnames = template_workbook.sheetnames
        template_rows = [worksheet_input_row(active_worksheet(template_workbook), row) for row in (18, 19)]
    finally:
        template_workbook.close()

    summary = run_settlements(source_path, config_path, worker_template_path=worker_template_path)

    expected_files = {
        f"Rozliczenie {PERIOD} - {name}.xlsx"
        for name in ("Adrian Testowy", "Darek Fikcyjny", "Kamil Syntetyczny")
    }
    assert {path.name for path in target_directory.glob("*.xlsx")} == expected_files
    assert summary.written_count == 2
    assert summary.empty_count == 1
    assert summary.warnings == []
    assert all(result.output_file.parent == target_directory for result in summary.results)
    assert worker_template_path.read_bytes() == worker_template_before

    adrian_workbook = load_workbook(
        target_directory / f"Rozliczenie {PERIOD} - Adrian Testowy.xlsx",
        data_only=False,
    )
    try:
        adrian = active_worksheet(adrian_workbook)
        assert adrian_workbook.sheetnames == template_sheetnames
        assert worksheet_input_row(adrian, 18) == expected_source_row(
            "syntetyczny adres 1", "#1", "adrian.testowy", 1
        )
        assert worksheet_input_row(adrian, 19) == expected_source_row(
            "syntetyczny adres 3", "#3", "adrian.testowy", 3
        )
    finally:
        adrian_workbook.close()

    darek_workbook = load_workbook(
        target_directory / f"Rozliczenie {PERIOD} - Darek Fikcyjny.xlsx",
        data_only=False,
    )
    try:
        darek = active_worksheet(darek_workbook)
        assert worksheet_input_row(darek, 18) == expected_source_row(
            "syntetyczny adres 2", "#2", "dariusz.fikcyjny2", 2
        )
    finally:
        darek_workbook.close()

    empty_workbook = load_workbook(
        target_directory / f"Rozliczenie {PERIOD} - Kamil Syntetyczny.xlsx",
        data_only=False,
    )
    try:
        empty = active_worksheet(empty_workbook)
        assert [worksheet_input_row(empty, row) for row in (18, 19)] == template_rows  # pusty Szablon to Wzorzec
    finally:
        empty_workbook.close()


def style_cells(cells: dict[str, tuple[Font, PatternFill]]) -> Callable[[Path], None]:
    def change(path: Path) -> None:
        workbook = load_workbook(path)
        try:
            sheet = active_worksheet(workbook)
            for coordinate, (font, fill) in cells.items():
                sheet[coordinate].font = font
                sheet[coordinate].fill = fill
            workbook.save(path)
        finally:
            workbook.close()

    return change


def highlights(cell: Cell) -> tuple[object, ...]:
    """Wyróżnienia komórki tak, jak widzi je tata: wypełnienie, kolor czcionki, pogrubienie, kursywa."""

    fill = cell.fill
    return (fill.fill_type, fill.fgColor if fill.fill_type else None, cell.font.color, bool(cell.font.b), bool(cell.font.i))


def template_look(cell: Cell) -> tuple[object, ...]:
    """Krój i rozmiar czcionki, obramowanie i format liczby komórki."""

    return (cell.font.name, cell.font.sz, cell.border.left.style, cell.number_format)


def test_wiersze_danych_rozliczenia_pracownika_maja_wyroznienia_pliku_zbiorczego_i_wyglad_wiersza_danych(
    complete_run: tuple[Path, Path, Path, Path],
) -> None:
    source_path, target_directory, config_path, worker_template_path = complete_run
    style_cells(
        {
            "B18": (
                Font(name="Arial", sz=8, color="FFFF0000", b=True, i=True),
                PatternFill(fill_type="solid", fgColor="FFFFFF00"),
            ),
            "N18": (Font(name="Arial", sz=8), PatternFill(fill_type="solid", fgColor=Color(theme=5, tint=0.4))),
        }
    )(source_path)
    template_cell = (
        Font(name="Times New Roman", sz=9, color="FF00AA55", b=True, i=True),
        PatternFill(fill_type="solid", fgColor="FF0000FF"),
    )
    style_cells({"B18": template_cell, "A19": template_cell, "C18": template_cell, "N18": template_cell})(
        worker_template_path
    )
    workbook = load_workbook(worker_template_path)
    try:
        for coordinate in ("B18", "N18"):
            active_worksheet(workbook)[coordinate].border = Border(left=Side(style="dashed"))
            active_worksheet(workbook)[coordinate].number_format = "@"
        workbook.save(worker_template_path)
    finally:
        workbook.close()
    template_workbook = load_workbook(worker_template_path)
    source_workbook = load_workbook(source_path)
    try:
        template = active_worksheet(template_workbook)
        source = active_worksheet(source_workbook)
        header_font = (template["B17"].font.name, template["B17"].font.sz)
        expected_template_look = template_look(template["N18"])
        # Szablon pracownika Adriana dostaje pierwsze i trzecie zlecenie Pliku zbiorczego: wiersze 18 i 20.
        expected_highlights = {
            written: highlights(source[read]) for written, read in (("A19", "A20"), ("B18", "B18"), ("N18", "N18"))
        }
        source_formats = {written: source[read].number_format for written, read in (("A19", "A20"), ("B18", "B18"))}
    finally:
        template_workbook.close()
        source_workbook.close()
    assert expected_highlights["B18"] == ("solid", Color("FFFFFF00"), Color("FFFF0000"), True, True)
    assert expected_highlights["A19"][:2] == (None, None)
    assert header_font != ("Times New Roman", 9)

    summary = run_settlements(source_path, config_path, worker_template_path=worker_template_path)

    assert summary.written_count == 2
    written_workbook = load_workbook(target_directory / f"Rozliczenie {PERIOD} - Adrian Testowy.xlsx")
    try:
        written = active_worksheet(written_workbook)
        assert {cell: highlights(written[cell]) for cell in expected_highlights} == expected_highlights
        # Komórki C18 nie ma w Pliku zbiorczym: bez Wyróżnień.
        assert highlights(written["C18"]) == (None, None, None, False, False)
        # A:L: Wygląd Wiersza danych z czcionką nagłówka Wzorca i formatem liczby z Pliku zbiorczego; N poza nim.
        assert {cell: template_look(written[cell]) for cell in ("A19", "B18", "C18")} == {
            "A19": (*header_font, "thin", source_formats["A19"]),
            "B18": (*header_font, "thin", source_formats["B18"]),
            "C18": (*header_font, "thin", "General"),
        }
        assert template_look(written["N18"]) == expected_template_look
    finally:
        written_workbook.close()


def test_numer_zlecenia_w_rozliczeniu_pracownika_ma_hiperlacze_z_pliku_zbiorczego(
    complete_run: tuple[Path, Path, Path, Path],
) -> None:
    source_path, target_directory, config_path, worker_template_path = complete_run
    links = link_order_numbers(source_path, "https://przyklad.invalid/zlecenie?tw=1&caseId=SYN-")

    summary = run_settlements(source_path, config_path, worker_template_path=worker_template_path)

    assert (summary.written_count, summary.empty_count, summary.warnings) == (2, 1, [])
    # Szablon pracownika Adriana dostaje pierwsze i trzecie zlecenie Pliku zbiorczego, Darka drugie.
    assert hyperlinks(target_directory / f"Rozliczenie {PERIOD} - Adrian Testowy.xlsx") == {"F18": links[18], "F19": links[20]}
    assert hyperlinks(target_directory / f"Rozliczenie {PERIOD} - Darek Fikcyjny.xlsx") == {"F18": links[19]}
    assert hyperlinks(target_directory / f"Rozliczenie {PERIOD} - Kamil Syntetyczny.xlsx") == {}


def test_hiperlacze_wzorca_w_wierszu_danych_znika_a_poza_wierszami_danych_zostaje(tmp_path: Path) -> None:
    source_path, target_directory, config_path, worker_template_path = make_complete_run_fixture(tmp_path)
    workbook = load_workbook(worker_template_path)
    try:
        sheet = active_worksheet(workbook)
        sheet["A5"] = "instrukcja"
        sheet["A5"].hyperlink = "https://przyklad.invalid/instrukcja"
        sheet["AT18"] = "uwaga Wzorca"
        sheet["AT18"].hyperlink = "https://przyklad.invalid/wzorzec"
        workbook.save(worker_template_path)
    finally:
        workbook.close()

    summary = run_settlements(source_path, config_path, worker_template_path=worker_template_path)

    assert summary.written_count == 2
    written = target_directory / f"Rozliczenie {PERIOD} - Adrian Testowy.xlsx"
    assert hyperlinks(written) == {"A5": "https://przyklad.invalid/instrukcja"}


def test_przebieg_rozliczen_na_wzorcu_z_formula_w_wierszu_danych_zapisuje_jak_openpyxl(tmp_path: Path) -> None:
    source_path, target_directory, config_path, worker_template_path = make_complete_run_fixture(tmp_path)
    set_cell("K18", "=N18+1")(worker_template_path)
    add_external_link_fixture(worker_template_path)
    worker_template_before = worker_template_path.read_bytes()
    expected_directory = tmp_path / "przez-openpyxl"
    expected_directory.mkdir()
    adrian = f"Rozliczenie {PERIOD} - Adrian Testowy.xlsx"
    # Szablon pracownika Adriana dostaje pierwsze i trzecie zlecenie Pliku zbiorczego.
    adrian_rows = [row for row in read_data_rows(source_path).rows if row.number in (18, 20)]
    twin = refused_by_copy(Path(shutil.copy(worker_template_path, tmp_path / "bliźniak Wzorca.xlsx")))
    TemplateFile(twin, Layout.WORKER_SETTLEMENT).create(expected_directory / adrian, adrian_rows)

    summary = run_settlements(source_path, config_path, worker_template_path=worker_template_path)

    assert (summary.written_count, summary.empty_count, summary.warnings) == (2, 1, [])
    assert worker_template_path.read_bytes() == worker_template_before
    written_directory = tmp_path / "przebieg"
    written_directory.mkdir()
    shutil.copy2(target_directory / adrian, written_directory / adrian)
    assert set(compare_worker_settlements(expected_directory, written_directory).values()) == {0}
    empty_workbook = load_workbook(target_directory / f"Rozliczenie {PERIOD} - Kamil Syntetyczny.xlsx")
    try:
        assert active_worksheet(empty_workbook)["K18"].value == "=N18+1"  # formuła Wzorca nie jest danymi
    finally:
        empty_workbook.close()


def test_kazde_rozliczenie_pracownika_takze_puste_ma_szerokosci_kolumn_wzorca_rozliczenia_pracownika(
    tmp_path: Path,
) -> None:
    source_path, target_directory, config_path, worker_template_path = make_complete_run_fixture(tmp_path)

    summary = run_settlements(source_path, config_path, worker_template_path=worker_template_path)

    assert summary.empty_count == 1
    expected = [WORKER_TEMPLATE_COLUMN_WIDTHS.get(column) for column in range(1, 48)]
    published = sorted(target_directory.glob("*.xlsx"))
    assert len(published) == 3
    for path in published:
        workbook = load_workbook(path)
        try:
            assert column_widths(active_worksheet(workbook), 47) == expected, path.name
        finally:
            workbook.close()


def with_unmapped_worker(source_path: Path) -> None:
    write_source(source_path, include_unmapped=True)


def with_workers(source_path: Path, workers: dict[int, str | None]) -> None:
    """Zmienia WYKONAWCĘ (kolumna H) w podanych Wierszach danych Pliku zbiorczego."""

    workbook = load_workbook(source_path)
    try:
        sheet = active_worksheet(workbook)
        for row, worker in workers.items():
            sheet.cell(row, 8).value = worker
        workbook.save(source_path)
    finally:
        workbook.close()


def with_data_row_without_worker(source_path: Path) -> None:
    with_workers(source_path, {20: None})


def with_unmapped_worker_and_data_row_without_worker(source_path: Path) -> None:
    with_unmapped_worker(source_path)
    with_data_row_without_worker(source_path)


@pytest.mark.parametrize(
    ("damage", "code"),
    [
        (with_unmapped_worker, "NIEZNANY_WYKONAWCA"),
        (with_data_row_without_worker, "WIERSZE_BEZ_WYKONAWCY"),
        (with_unmapped_worker_and_data_row_without_worker, "NIEZNANY_WYKONAWCA"),
    ],
    ids=["wykonawca-bez-mapowania", "wiersz-danych-bez-wykonawcy", "oba"],
)
def test_wykonawca_bez_mapowania_albo_wiersz_danych_bez_wykonawcy_wstrzymuje_publikacje_folderu_rozliczen_pracownikow(
    tmp_path: Path, damage: Callable[[Path], None], code: str
) -> None:
    source_path, target_directory, config_path = make_fixture(tmp_path)
    damage(source_path)

    with pytest.raises(StopError) as raised:
        run_settlements(source_path, config_path, worker_template_path=tmp_path / "Wzorzec rozliczenia pracownika.xlsx")

    assert (raised.value.message.code, raised.value.message.where) == (code, f"plik {SOURCE_NAME}, folder {PERIOD}, kolumna H")
    assert not target_directory.exists()
    assert not list(target_directory.parent.glob(f".{target_directory.name}.staging-*"))


def test_podglad_bez_zapisu_pokazuje_problemy_wierszy_danych_i_plan_kazdego_wykonawcy_z_konfiguracji(tmp_path: Path) -> None:
    source_path, target_directory, config_path = make_fixture(tmp_path)
    with_workers(source_path, {18: "andrzej.nieznany2", 20: None})  # nieznany Wykonawca w pierwszym wierszu, przed znanymi

    summary = run_settlements(
        source_path,
        config_path,
        dry_run=True,
        worker_template_path=tmp_path / "Wzorzec rozliczenia pracownika.xlsx",
    )

    assert [warning.code for warning in summary.warnings] == ["BRAK_WYKONAWCY", "BRAK_MAPOWANIA"]
    assert [(result.worker_name, result.status, result.rows) for result in summary.results] == [
        ("Adrian Testowy", "PUSTY_SZABLON", 0),
        ("Darek Fikcyjny", "PLAN", 1),
        ("Kamil Syntetyczny", "PUSTY_SZABLON", 0),
    ]
    assert not target_directory.exists()


@pytest.mark.parametrize(
    "damage",
    [overwrite_with_non_workbook, cut_sheet_after_data_rows, without_sheet],
    ids=["nie-skoroszyt", "urwany-arkusz", "bez-arkusza"],
)
def test_nieczytelny_plik_zbiorczy_zatrzymuje_przebieg_rozliczen_bezpiecznym_komunikatem(
    tmp_path: Path,
    damage: Callable[[Path], None],
) -> None:
    source_path, target_directory, config_path = make_fixture(tmp_path)
    damage(source_path)

    with pytest.raises(StopError) as stop:
        run_settlements(source_path, config_path, worker_template_path=tmp_path / "Wzorzec rozliczenia pracownika.xlsx")

    assert (stop.value.message.code, stop.value.message.what) == (
        "ZLY_PLIK_ZBIORCZY",
        f"{UNFIT_COLLECTIVE_FILE_PREFIX} {UNREADABLE_SOURCE_MESSAGE}",
    )
    assert not target_directory.exists()
    assert not list(target_directory.parent.glob(f".{target_directory.name}.staging-*"))


@pytest.mark.parametrize(
    ("failing", "error"),
    [
        ("read_bytes", PermissionError(errno.EACCES, "The process cannot access the file because it is being used by another process")),
        ("read_bytes", OSError(errno.EIO, "The cloud file provider is not running")),
        ("is_file", PermissionError(errno.EACCES, "Permission denied")),
    ],
    ids=["zablokowany-przez-inny-program", "onedrive-bez-sieci", "folder-bez-dostepu"],
)
def test_plik_zbiorczy_niedostepny_do_odczytu_nie_jest_nazwany_uszkodzonym(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, failing: str, error: OSError
) -> None:
    source_path, target_directory, config_path = make_fixture(tmp_path)
    original = getattr(Path, failing)

    def blocked(path: Path) -> object:
        if path == source_path:
            raise error
        return original(path)

    monkeypatch.setattr(Path, failing, blocked)

    with pytest.raises(StopError) as stop:
        run_settlements(source_path, config_path, worker_template_path=tmp_path / "Wzorzec rozliczenia pracownika.xlsx")

    assert stop.value.message.code == "NIEDOSTEPNY_PLIK_ZBIORCZY"
    assert "uszkodzony" not in stop.value.message.text
    assert not target_directory.exists()


def test_plik_zbiorczy_otwarty_w_excelu_zatrzymuje_przebieg_rozliczen_i_zostawia_plik_blokady(tmp_path: Path) -> None:
    source_path, target_directory, config_path = make_fixture(tmp_path)
    lock_file = source_path.with_name(f"~${source_path.name}")
    lock_file.write_bytes(b"blokada Excela")
    source_before = source_path.read_bytes()

    with pytest.raises(StopError) as stop:
        run_settlements(source_path, config_path, worker_template_path=tmp_path / "Wzorzec rozliczenia pracownika.xlsx")

    assert stop.value.message.code == "PLIK_ZBIORCZY_OTWARTY"
    assert stop.value.message.where == f"plik {source_path.name}, folder {PERIOD}"
    assert lock_file.read_bytes() == b"blokada Excela"
    assert source_path.read_bytes() == source_before
    assert not target_directory.exists()


def test_plik_zbiorczy_bez_wykonawca_w_h17_zatrzymuje_przebieg_rozliczen(tmp_path: Path) -> None:
    source_path, target_directory, config_path = make_fixture(tmp_path)
    workbook = load_workbook(source_path)
    try:
        sheet = active_worksheet(workbook)
        sheet["H17"], sheet["C17"] = None, "WYKONAWCA"
        workbook.save(source_path)
    finally:
        workbook.close()

    with pytest.raises(StopError) as stop:
        run_settlements(source_path, config_path, worker_template_path=tmp_path / "Wzorzec rozliczenia pracownika.xlsx")

    assert stop.value.message.code == "ZLY_PLIK_ZBIORCZY"
    assert stop.value.message.what == f"{UNFIT_COLLECTIVE_FILE_PREFIX} W wierszu 17 nie znaleziono nagłówka WYKONAWCA w kolumnie H."
    assert not target_directory.exists()


def test_istniejacy_folder_rozliczen_pracownikow_nie_jest_nadpisywany(complete_run: tuple[Path, Path, Path, Path]) -> None:
    source_path, target_directory, config_path, worker_template_path = complete_run
    target_directory.mkdir()
    sentinel = target_directory / "existing.txt"
    sentinel.write_text("zachowaj", encoding="utf-8")

    with pytest.raises(StopError) as raised:
        run_settlements(source_path, config_path, worker_template_path=worker_template_path)

    assert raised.value.message.code == "FOLDER_ROZLICZEN_ISTNIEJE"
    assert sentinel.read_text(encoding="utf-8") == "zachowaj"


def lock_in_excel(path: Path) -> None:
    path.with_name(f"~${path.name}").touch()


def set_cell(cell: str, value: str) -> Callable[[Path], None]:
    def change(path: Path) -> None:
        workbook = load_workbook(path)
        try:
            active_worksheet(workbook)[cell] = value
            workbook.save(path)
        finally:
            workbook.close()

    return change


def changes(*steps: Callable[[Path], None]) -> Callable[[Path], None]:
    def change(path: Path) -> None:
        for step in steps:
            step(path)

    return change


@pytest.mark.parametrize("dry_run", [False, True], ids=["RUN", "DRY-RUN"])
@pytest.mark.parametrize(
    ("damage", "detail"),
    [
        (lock_in_excel, "Wzorzec rozliczenia pracownika jest otwarty lub zablokowany."),
        (overwrite_with_non_workbook, "Nie można odczytać Wzorca rozliczenia pracownika."),
        (set_cell("H17", "NIE WYKONAWCA"), "Wzorzec rozliczenia pracownika nie ma nagłówka WYKONAWCA w H:17."),
        (
            set_cell("A18", "SENSITIVE-SYNTHETIC"),
            "Wzorzec rozliczenia pracownika zawiera dane i nie może być użyty jako wzór.",
        ),
        (
            set_cell("K25", "SENSITIVE-SYNTHETIC"),
            "Wzorzec rozliczenia pracownika zawiera dane i nie może być użyty jako wzór.",
        ),
        (
            changes(set_cell("K18", "=N18+1"), set_cell("K25", "SENSITIVE-SYNTHETIC")),
            "Wzorzec rozliczenia pracownika zawiera dane i nie może być użyty jako wzór.",
        ),
    ],
    ids=[
        "otwarty-w-excelu",
        "nie-skoroszyt",
        "bez-wykonawca-w-h17",
        "z-danymi",
        "z-danymi-dalej",
        "z-formula-w-wierszu-danych-i-danymi",
    ],
)
def test_wzorzec_rozliczenia_pracownika_ktorego_nie_mozna_uzyc_zatrzymuje_sprawdzanie(
    complete_run: tuple[Path, Path, Path, Path], damage: Callable[[Path], None], detail: str, dry_run: bool
) -> None:
    source_path, target_directory, config_path, worker_template_path = complete_run
    damage(worker_template_path)

    with pytest.raises(StopError) as raised:
        run_settlements(source_path, config_path, dry_run=dry_run, worker_template_path=worker_template_path)

    assert raised.value.message.code == "ZLY_WZORZEC_PRACOWNIKA"
    assert raised.value.message.what == f"{UNUSABLE_WORKER_TEMPLATE_PREFIX} {detail}"
    assert "SENSITIVE-SYNTHETIC" not in raised.value.message.text
    assert not target_directory.exists()


def test_wzorzec_rozliczenia_pracownika_spoza_formatow_skoroszytu_jest_nieczytelny_w_sprawdzaniu(tmp_path: Path) -> None:
    source_path, target_directory, config_path = make_fixture(tmp_path)
    worker_template_path = tmp_path / "Wzorzec rozliczenia pracownika.xls"
    (tmp_path / "Wzorzec rozliczenia pracownika.xlsx").rename(worker_template_path)

    with pytest.raises(StopError) as raised:
        run_settlements(source_path, config_path, worker_template_path=worker_template_path)

    assert raised.value.message.code == "ZLY_WZORZEC_PRACOWNIKA"
    assert raised.value.message.what == f"{UNUSABLE_WORKER_TEMPLATE_PREFIX} Nie można odczytać Wzorca rozliczenia pracownika."
    assert not target_directory.exists()


def test_niepoprawna_konfiguracja_wykonawcow_zatrzymuje_sprawdzanie(tmp_path: Path) -> None:
    source_path, target_directory, config_path = make_fixture(tmp_path)
    config_path.write_text("workers: [", encoding="utf-8")

    with pytest.raises(StopError) as raised:
        run_settlements(source_path, config_path, worker_template_path=tmp_path / "Wzorzec rozliczenia pracownika.xlsx")

    assert raised.value.message.code == "ZLA_KONFIGURACJA_WYKONAWCOW"
    assert not target_directory.exists()


def test_wzorzec_rozliczenia_pracownika_otwarty_w_zapisywaniu_wstrzymuje_publikacje(complete_run: tuple[Path, Path, Path, Path]) -> None:
    source_path, target_directory, config_path, worker_template_path = complete_run
    lock_file = worker_template_path.with_name(f"~${worker_template_path.name}")

    def lock_in_saving(event: ProgressEvent) -> None:
        if isinstance(event, WorkerStarted):
            lock_file.touch()

    with pytest.raises(StopError) as raised:
        run_settlements(source_path, config_path, worker_template_path=worker_template_path, observer=lock_in_saving)

    assert raised.value.message.code == "BLAD_KOPII_WZORCA"
    assert not target_directory.exists()
    assert not list(target_directory.parent.glob(f".{target_directory.name}.staging-*"))


TEMPLATE_NOT_COPIED = (
    "Nie można utworzyć skoroszytów z Wzorca rozliczenia pracownika. Nie utworzono żadnego pliku pracownika."
)
FIRST_WORKER_FILE_FAILED = (
    "Nie udało się przygotować skoroszytu pracownika Adrian Testowy. Nie utworzono żadnego pliku pracownika."
)


@pytest.mark.parametrize(
    ("damage", "code", "komunikat"),
    [
        (overwrite_with_non_workbook, "BLAD_KOPII_WZORCA", TEMPLATE_NOT_COPIED),
        (set_cell("H17", "NIE WYKONAWCA"), "BLAD_SZABLONU_PRACOWNIKA", FIRST_WORKER_FILE_FAILED),
        (set_cell("A18", "SENSITIVE-SYNTHETIC"), "BLAD_SZABLONU_PRACOWNIKA", FIRST_WORKER_FILE_FAILED),
    ],
    ids=["nie-skoroszyt", "bez-wykonawca-w-h17", "z-danymi"],
)
def test_wzorzec_rozliczenia_pracownika_zepsuty_po_sprawdzaniu_wstrzymuje_publikacje(
    complete_run: tuple[Path, Path, Path, Path], damage: Callable[[Path], None], code: str, komunikat: str
) -> None:
    source_path, target_directory, config_path, worker_template_path = complete_run
    recorded: list[ProgressEvent] = []

    def damage_in_saving(event: ProgressEvent) -> None:
        recorded.append(event)
        if isinstance(event, WorkerStarted) and event.template_index == 1:
            damage(worker_template_path)

    with pytest.raises(StopError) as raised:
        run_settlements(source_path, config_path, worker_template_path=worker_template_path, observer=damage_in_saving)

    assert raised.value.message.code == code
    assert raised.value.message.what == komunikat
    assert "SENSITIVE-SYNTHETIC" not in raised.value.message.text
    assert not target_directory.exists()
    assert not list(target_directory.parent.glob(f".{target_directory.name}.staging-*"))
    failed = [event for event in recorded if isinstance(event, (WorkerEnded, WorkerFailed))]
    assert [(type(event), event.template_index) for event in failed] == [(WorkerFailed, 1)]


def test_wzorzec_rozliczenia_pracownika_ze_scaleniem_w_wierszach_danych_zatrzymuje_zapisywanie(
    complete_run: tuple[Path, Path, Path, Path],
) -> None:
    source_path, target_directory, config_path, worker_template_path = complete_run
    workbook = load_workbook(worker_template_path)
    try:
        active_worksheet(workbook).merge_cells("A18:B18")
        workbook.save(worker_template_path)
    finally:
        workbook.close()
    worker_template_before = worker_template_path.read_bytes()

    with pytest.raises(StopError) as raised:
        run_settlements(source_path, config_path, worker_template_path=worker_template_path)

    assert raised.value.message.code == "BLAD_SZABLONU_PRACOWNIKA"
    assert raised.value.message.what == FIRST_WORKER_FILE_FAILED
    assert not target_directory.exists()
    assert not list(target_directory.parent.glob(f".{target_directory.name}.staging-*"))
    assert worker_template_path.read_bytes() == worker_template_before


@pytest.mark.parametrize(
    ("failing_phase", "code"),
    [(ProgressPhase.CHECKING, "ZLY_WZORZEC_PRACOWNIKA"), (ProgressPhase.SAVING, "BLAD_KOPII_WZORCA")],
    ids=["w-sprawdzaniu", "w-zapisywaniu"],
)
def test_blad_odczytu_wzorca_rozliczenia_pracownika_z_dysku_wstrzymuje_publikacje(
    complete_run: tuple[Path, Path, Path, Path], monkeypatch: pytest.MonkeyPatch, failing_phase: ProgressPhase, code: str
) -> None:
    source_path, target_directory, config_path, worker_template_path = complete_run
    read_bytes = Path.read_bytes
    phase: list[ProgressPhase] = []

    def fail_reading_template(path: Path) -> bytes:
        if path == worker_template_path and phase == [failing_phase]:
            raise PermissionError("synthetic read failure")
        return read_bytes(path)

    def follow_phase(event: ProgressEvent) -> None:
        if isinstance(event, PhaseStarted):
            phase[:] = [event.phase]

    monkeypatch.setattr(Path, "read_bytes", fail_reading_template)

    with pytest.raises(StopError) as raised:
        run_settlements(source_path, config_path, worker_template_path=worker_template_path, observer=follow_phase)

    assert raised.value.message.code == code
    if failing_phase is ProgressPhase.CHECKING:
        assert raised.value.message.what == f"{UNUSABLE_WORKER_TEMPLATE_PREFIX} Nie można odczytać Wzorca rozliczenia pracownika."
    assert not target_directory.exists()
    assert not list(target_directory.parent.glob(f".{target_directory.name}.staging-*"))


def test_nieudany_zapis_rozliczenia_pracownika_wstrzymuje_publikacje(
    complete_run: tuple[Path, Path, Path, Path], monkeypatch: pytest.MonkeyPatch
) -> None:
    source_path, target_directory, config_path, worker_template_path = complete_run
    worker_template_before = worker_template_path.read_bytes()
    replace = os.replace

    def fail_replacing_workbook(source: str | os.PathLike[str], target: str | os.PathLike[str]) -> None:
        if Path(target).suffix == ".xlsx":
            raise OSError("synthetic write failure")
        replace(source, target)

    monkeypatch.setattr(os, "replace", fail_replacing_workbook)

    with pytest.raises(StopError) as raised:
        run_settlements(source_path, config_path, worker_template_path=worker_template_path)

    assert raised.value.message.code == "BLAD_SZABLONU_PRACOWNIKA"
    assert "Adrian Testowy" in raised.value.message.what
    assert not target_directory.exists()
    assert not list(target_directory.parent.glob(f".{target_directory.name}.staging-*"))
    assert sorted(path.name for path in target_directory.parent.iterdir()) == [SOURCE_NAME]
    assert worker_template_path.read_bytes() == worker_template_before


def fail_renaming_staging_folder(monkeypatch: pytest.MonkeyPatch, target_directory: Path) -> None:
    real_rename = os.rename

    def rename(source: str | os.PathLike[str], target: str | os.PathLike[str]) -> None:
        if Path(source).name.startswith(f".{target_directory.name}.staging-"):
            raise PermissionError("wymuszony: Windows nie zmienia nazwy folderu z otwartym plikiem")
        real_rename(source, target)

    monkeypatch.setattr(os, "rename", rename)


def fail_saving_second_worker_file(monkeypatch: pytest.MonkeyPatch, target_directory: Path) -> None:
    real_replace = os.replace
    saved: list[str] = []

    def replace(source: str | os.PathLike[str], target: str | os.PathLike[str]) -> None:
        if Path(target).suffix == ".xlsx":
            saved.append(Path(target).name)
            if len(saved) == 2:
                raise OSError("wymuszony: zapis drugiego pliku pracownika")
        real_replace(source, target)

    monkeypatch.setattr(os, "replace", replace)


def create_worker_settlements_folder_while_saving(monkeypatch: pytest.MonkeyPatch, target_directory: Path) -> None:
    real_mkstemp = tempfile.mkstemp

    def mkstemp(*args: object, **kwargs: object) -> tuple[int, str]:
        if not target_directory.exists():
            target_directory.mkdir()
            (target_directory / "existing.txt").write_text("zachowaj", encoding="utf-8")
        return real_mkstemp(*args, **kwargs)  # type: ignore[call-overload,no-any-return]

    monkeypatch.setattr(tempfile, "mkstemp", mkstemp)


def test_nieudana_publikacja_folderu_rozliczen_pracownikow_zatrzymuje_kodem_publikacji(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source_path, target_directory, config_path = make_fixture(tmp_path)
    fail_renaming_staging_folder(monkeypatch, target_directory)

    with pytest.raises(StopError) as raised:
        run_settlements(source_path, config_path, worker_template_path=tmp_path / "Wzorzec rozliczenia pracownika.xlsx")

    assert raised.value.message.code == "BLAD_PUBLIKACJI"
    assert not target_directory.exists()
    assert not list(target_directory.parent.glob(f".{target_directory.name}.staging-*"))


def hold_worker_files_like_onedrive(monkeypatch: pytest.MonkeyPatch, refuses_permission_change: bool) -> None:
    worker_file_prefix = f"Rozliczenie {PERIOD} - "
    real_unlink, real_chmod = os.unlink, os.chmod

    def unlink(path: str | os.PathLike[str], **kwargs: object) -> None:
        if Path(path).name.startswith(worker_file_prefix):
            raise PermissionError("wymuszony: plik pracownika trzyma OneDrive")
        real_unlink(path, **kwargs)  # type: ignore[arg-type]

    def chmod(path: str | os.PathLike[str], mode: int, **kwargs: object) -> None:
        if refuses_permission_change:
            raise PermissionError("wymuszony: OneDrive nie pozwala zmienić uprawnień")
        real_chmod(path, mode, **kwargs)  # type: ignore[arg-type]

    monkeypatch.setattr(os, "unlink", unlink)
    monkeypatch.setattr(os, "chmod", chmod)


@pytest.mark.parametrize(
    ("failure", "code", "original_action"),
    [
        (fail_renaming_staging_folder, "BLAD_PUBLIKACJI", "Sprawdź, czy Folder okresu jest dostępny."),
        (fail_saving_second_worker_file, "BLAD_SZABLONU_PRACOWNIKA", "Uruchom ponownie."),
        (
            create_worker_settlements_folder_while_saving,
            "FOLDER_ROZLICZEN_ISTNIEJE",
            "przenieś Folder rozliczeń pracowników poza Folder okresu",
        ),
    ],
    ids=["publikacja", "zapis-drugiego-pracownika", "folder-rozliczen-powstal-w-trakcie"],
)
@pytest.mark.parametrize("refuses_permission_change", [False, True], ids=["usuwanie", "usuwanie-i-zmiana-uprawnien"])
def test_zatrzymanie_wskazuje_folder_tymczasowy_trzymany_przez_onedrive(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    failure: Callable[[pytest.MonkeyPatch, Path], None],
    code: str,
    original_action: str,
    refuses_permission_change: bool,
) -> None:
    source_path, target_directory, config_path = make_fixture(tmp_path)
    failure(monkeypatch, target_directory)
    hold_worker_files_like_onedrive(monkeypatch, refuses_permission_change)

    with pytest.raises(StopError) as raised:
        run_settlements(source_path, config_path, worker_template_path=tmp_path / "Wzorzec rozliczenia pracownika.xlsx")

    message = raised.value.message
    staging_prefix = f".{target_directory.name}.staging-"
    staging_left = list(target_directory.parent.glob(f"{staging_prefix}*"))
    assert message.code == code
    assert len(staging_left) == 1
    assert any(staging_left[0].iterdir())
    assert f"zaczyna się od „{staging_prefix}”" in message.action
    assert original_action in message.action
    assert "folder tymczasowy" in message.what
    assert "Nie utworzono żadnego pliku pracownika." not in message.what
    assert not target_directory.exists() or [path.name for path in target_directory.iterdir()] == ["existing.txt"]


def test_ponowna_nieudana_publikacja_z_folderem_tymczasowym_jest_powtorka(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source_path, target_directory, config_path = make_fixture(tmp_path)
    fail_renaming_staging_folder(monkeypatch, target_directory)
    hold_worker_files_like_onedrive(monkeypatch, refuses_permission_change=False)
    stops: list[Stop] = []
    for _ in range(2):
        with pytest.raises(StopError) as raised:
            run_settlements(source_path, config_path, worker_template_path=tmp_path / "Wzorzec rozliczenia pracownika.xlsx")
        stops.append(Stop("stop", raised.value.message))

    assert len(list(target_directory.parent.glob(f".{target_directory.name}.staging-*"))) == 2
    assert stops[1].message == stops[0].message
    assert stops[1].stop_kind == "for_tata"


def test_folder_rozliczen_pracownikow_utworzony_w_trakcie_zapisywania_nie_jest_nadpisywany(complete_run: tuple[Path, Path, Path, Path]) -> None:
    source_path, target_directory, config_path, worker_template_path = complete_run
    sentinel = target_directory / "existing.txt"

    def create_folder_in_saving(event: ProgressEvent) -> None:
        if isinstance(event, WorkerStarted) and not target_directory.exists():
            target_directory.mkdir()
            sentinel.write_text("zachowaj", encoding="utf-8")

    with pytest.raises(StopError) as raised:
        run_settlements(
            source_path, config_path, worker_template_path=worker_template_path, observer=create_folder_in_saving
        )

    assert raised.value.message.code == "FOLDER_ROZLICZEN_ISTNIEJE"
    assert [path.name for path in target_directory.iterdir()] == [sentinel.name]
    assert sentinel.read_text(encoding="utf-8") == "zachowaj"
    assert not list(target_directory.parent.glob(f".{target_directory.name}.staging-*"))


def test_brak_wzorca_rozliczenia_pracownika_podpowiada_zmiane_nazwy_starego_pliku(
    tmp_path: Path,
) -> None:
    source_path, target_directory, config_path = make_fixture(tmp_path)
    worker_template_path = tmp_path / "Wzorzec rozliczenia pracownika.xlsx"
    worker_template_path.rename(tmp_path / "placeholder.xlsx")

    with pytest.raises(StopError) as raised:
        run_settlements(source_path, config_path, worker_template_path=worker_template_path)

    assert raised.value.message.what == (
        f"{UNUSABLE_WORKER_TEMPLATE_PREFIX} "
        "Nie znaleziono Wzorca rozliczenia pracownika: Wzorzec rozliczenia pracownika.xlsx. "
        "Zmień nazwę pliku placeholder.xlsx w tym samym folderze na „Wzorzec rozliczenia pracownika.xlsx”."
    )
    assert not target_directory.exists()


@pytest.mark.parametrize(
    ("source", "detail"),
    [
        (f"{PERIOD}/Rozliczenie wew. {PERIOD}.xlsx", "Plik Rozliczenie wew. jest poza zakresem tego procesu."),
        (f"INEA/{SOURCE_NAME}", f"Folder pliku musi nazywać się tak samo jak okres: {PERIOD}."),
        (
            f"{PERIOD}/Rozliczenie {PERIOD} zbiorcze.xlsx",
            "Wybierz plik nazwany dokładnie: Rozliczenie {dd}_{dd}_{mm}_{rrrr} - zbiorcze.xlsx.",
        ),
    ],
    ids=["rozliczenie_wewnetrzne", "folder_nazwany_inaczej_niz_okres", "nazwa_inna_niz_plik_zbiorczy"],
)
def test_rozliczenie_wewnetrzne_albo_plik_zbiorczy_poza_folderem_okresu_zatrzymuje_przebieg_rozliczen(
    tmp_path: Path, source: str, detail: str
) -> None:
    source_path = tmp_path / source
    source_path.parent.mkdir()
    source_path.touch()

    with pytest.raises(StopError) as raised:
        run_settlements(source_path, tmp_path / "missing.yaml")

    assert raised.value.message.code == "ZLY_PLIK_ZBIORCZY"
    assert raised.value.message.what == f"{UNFIT_COLLECTIVE_FILE_PREFIX} {detail}"


def test_plik_zbiorczy_z_tokenem_ktory_nie_jest_okresem_zatrzymuje_przebieg_rozliczen(tmp_path: Path) -> None:
    source_path, target_directory, config_path = make_fixture(tmp_path, period=NOT_A_PERIOD)

    with pytest.raises(StopError) as raised:
        run_settlements(source_path, config_path, worker_template_path=tmp_path / "Wzorzec rozliczenia pracownika.xlsx")

    assert (raised.value.message.code, raised.value.message.what) == (
        "ZLY_PLIK_ZBIORCZY",
        f"{UNFIT_COLLECTIVE_FILE_PREFIX} Okres {NOT_A_PERIOD} w nazwie pliku nie jest okresem od wtorku do poniedziałku "
        "(przykład okresu: 08_14_09_2026).",
    )
    assert not target_directory.exists()
    assert [path.name for path in source_path.parent.iterdir()] == [source_path.name]


@pytest.mark.parametrize(
    "failure",
    [OSError(errno.EIO, "synthetic failure"), RuntimeError("synthetic failure")],
    ids=["blad-zapisu", "nieoczekiwany-blad"],
)
def test_failed_przetworzenie_szablonu_pracownika_emits_failure_without_false_completion(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, failure: Exception
) -> None:
    source_path, target_directory, config_path = make_fixture(tmp_path)
    real_replace = os.replace
    saved = 0
    recorded: list[ProgressEvent] = []

    def fail_on_second_save(source: str | os.PathLike[str], target: str | os.PathLike[str]) -> None:
        nonlocal saved
        if Path(target).suffix == ".xlsx":
            saved += 1
            if saved == 2:
                raise failure
        real_replace(source, target)

    monkeypatch.setattr(os, "replace", fail_on_second_save)

    with pytest.raises(StopError) as raised:
        run_settlements(
            source_path,
            config_path,
            worker_template_path=tmp_path / "Wzorzec rozliczenia pracownika.xlsx",
            observer=recorded.append,
        )

    assert raised.value.message.code == "BLAD_SZABLONU_PRACOWNIKA"
    assert "Darek Fikcyjny" in raised.value.message.what
    assert not target_directory.exists()
    assert not list(target_directory.parent.glob(f".{target_directory.name}.staging-*"))

    failed = [event for event in recorded if isinstance(event, (PhaseFailed, WorkerFailed))]
    assert len(failed) == 1
    assert isinstance(failed[0], WorkerFailed)
    assert failed[0].worker_name == "Darek Fikcyjny"
    assert failed[0].template_index == 2
    assert all("synthetic failure" not in repr(event) and "#1" not in repr(event) for event in failed)
    assert not any(
        isinstance(event, WorkerEnded) and event.worker_name == "Darek Fikcyjny" for event in recorded
    )
    assert not any(isinstance(event, PhaseEnded) and event.phase is ProgressPhase.SAVING for event in recorded)

    monkeypatch.setattr(os, "replace", real_replace)
    summary = run_settlements(source_path, config_path, worker_template_path=tmp_path / "Wzorzec rozliczenia pracownika.xlsx")
    assert summary.written_count == 2
    assert target_directory.is_dir()


@pytest.mark.parametrize(
    ("dry_run", "include_unmapped", "issue_count"),
    [(False, False, 0), (True, True, 1)],
    ids=["zapis", "dry-run-z-problemem"],
)
def test_progress_projection_accepts_whole_przebieg_rozliczen(
    tmp_path: Path, dry_run: bool, include_unmapped: bool, issue_count: int
) -> None:
    source_path, _, config_path = make_fixture(tmp_path)
    write_source(source_path, include_unmapped=include_unmapped)
    projection = ProgressProjection()

    summary = run_settlements(
        source_path,
        config_path,
        dry_run=dry_run,
        worker_template_path=tmp_path / "Wzorzec rozliczenia pracownika.xlsx",
        observer=projected(projection),
    )

    snapshot = projection.snapshot
    assert set(snapshot.phase_states.values()) == {PhaseState.COMPLETED}
    assert snapshot.template_total == 3
    assert snapshot.templates_completed == 3
    assert snapshot.issue_count == issue_count
    assert set(summary.phase_durations_ms) == set(ProgressPhase)
    with pytest.raises(TypeError):
        summary.phase_durations_ms[ProgressPhase.SAVING] = 0  # type: ignore[index]


def test_progress_projection_accepts_przebieg_rozliczen_interrupted_in_zapisywanie_with_its_time(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source_path, _, config_path = make_fixture(tmp_path)
    projection = ProgressProjection()
    save_s = 0.05

    real_replace = os.replace

    def fail_save(source: str | os.PathLike[str], target: str | os.PathLike[str]) -> None:
        if Path(target).suffix == ".xlsx":
            time.sleep(save_s)
            raise OSError(errno.EIO, "synthetic failure")
        real_replace(source, target)

    monkeypatch.setattr(os, "replace", fail_save)

    with pytest.raises(StopError):
        run_settlements(
            source_path,
            config_path,
            worker_template_path=tmp_path / "Wzorzec rozliczenia pracownika.xlsx",
            observer=projected(projection),
        )

    snapshot = projection.snapshot
    assert snapshot.phase_states == {
        ProgressPhase.CHECKING: PhaseState.COMPLETED,
        ProgressPhase.READING: PhaseState.COMPLETED,
        ProgressPhase.PLANNING: PhaseState.COMPLETED,
        ProgressPhase.SAVING: PhaseState.FAILED,
    }
    assert snapshot.template_total == 3
    assert snapshot.current_worker == "Adrian Testowy"
    assert snapshot.phase_durations_ms[ProgressPhase.SAVING] >= save_s * 1000
