"""Filtr prywatności Wiersza danych (``order_facts``): jedyna treść komórek Pliku zbiorczego, która opuszcza komputer taty.

Seam z dwoma wejściami (Kontrola ilości teraz, Logi Uruchomienia później), więc macierz reguły jest tu raz
(ADR 0013, ADR 0015). Dozwolone: UMOWA, URZĄDZENIA, CZYNNOŚCI, dzień tygodnia i godzina z DATA oraz N:AP.
"""

from __future__ import annotations

from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path

from openpyxl.utils import get_column_letter
import pytest

from rozliczenia.order_facts import NON_NUMERIC, MissingHeaders, OrderFacts, collective_orders, tata_value
from rozliczenia.quantity_rules import OrderContent, WeekTime
from rozliczenia.collective_file import read_data_rows
from tests.fixtures.data_rows import data_row
from tests.fixtures.settlement import write_quantity_source

# Nagłówki A17:K17 Wzorca pliku zbiorczego po normalizacji, jak je daje odczyt Pliku zbiorczego.
HEADER = {
    "miasto": "A",
    "adres": "B",
    "umowa": "C",
    "urządzenia": "D",
    "czynności": "E",
    "numer zl.": "F",
    "data": "G",
    "wykonawca": "H",
    "numer umowy/zlecenia orange": "I",
    "skuteczne": "J",
    "czynności zrealizowane": "K",
}
STARTED = datetime(2026, 9, 17, 13, 45, 21)


def _marked_row() -> list[object]:
    """Wiersz A:AT ze znacznikiem w każdej komórce, z prawdziwą datą w G."""

    row: list[object] = [f"ZNACZNIK-{get_column_letter(column)}" for column in range(1, 47)]
    row[6] = STARTED
    return row


def _filtered(header: dict[str, str], row: list[object]) -> OrderFacts:
    orders = collective_orders(header, [data_row(tuple(row))])  # type: ignore[arg-type]
    assert isinstance(orders, tuple)
    [facts] = orders
    return facts


def test_wiersz_po_filtrze_ma_tylko_umowe_urzadzenia_czynnosci_pore_i_ilosci() -> None:
    facts = _filtered(HEADER, _marked_row())

    assert facts == OrderFacts(
        content=OrderContent("ZNACZNIK-C", "ZNACZNIK-D", "ZNACZNIK-E", WeekTime(3, time(13, 45))),
        quantities={get_column_letter(column): NON_NUMERIC for column in range(14, 43)},
    )
    forbidden = [f"ZNACZNIK-{letter}" for letter in ("A", "B", "F", "G", "H", "I", "J", "K", "AQ", "AR", "AS", "AT")]
    text = repr(facts)
    assert [marker for marker in forbidden if marker in text] == []
    assert "2026" not in text and "17" not in text.replace("ZNACZNIK-", "")


def test_kolumny_wejscia_czytane_po_naglowku_a_nie_po_pozycji() -> None:
    reordered = {**HEADER, "umowa": "E", "czynności": "C"}
    row: list[object] = [None] * 46
    row[2], row[3], row[4], row[6] = "instalacja", "nokia zte", "u", STARTED

    content = _filtered(reordered, row).content

    assert content is not None
    assert (content.contract, content.devices, content.activities) == ("u", "nokia zte", "instalacja")


def test_brak_naglowka_kolumny_wejscia_to_brak_porownania_a_nie_zgadywanie_pozycji() -> None:
    header = {name: column for name, column in HEADER.items() if name not in ("umowa", "data")}

    assert collective_orders(header, [data_row(tuple(_marked_row()))]) == MissingHeaders(("UMOWA", "DATA"))  # type: ignore[arg-type]


def test_data_bez_daty_w_komorce_daje_wiersz_bez_tresci_zlecenia() -> None:
    row = _marked_row()
    row[6] = "wtorek rano"

    facts = _filtered(HEADER, row)

    assert facts.content is None


def test_data_bez_godziny_to_polnoc_tego_dnia_tygodnia() -> None:
    row = _marked_row()
    row[6] = date(2026, 9, 12)

    content = _filtered(HEADER, row).content

    assert content is not None
    assert content.started == WeekTime(5, time(0, 0))


def test_tresc_zlecenia_dla_regul_ilosci_ma_pore_bez_daty_dziennej() -> None:
    row: list[object] = [None] * 46
    row[2], row[3], row[4], row[6] = "u", "usterka", None, STARTED

    content = _filtered(HEADER, row).content

    assert content is not None
    assert (content.contract, content.devices, content.activities, content.started) == (
        "u",
        "usterka",
        "",
        WeekTime(3, time(13, 45)),
    )


@pytest.mark.parametrize(
    ("cell", "value"),
    [
        (None, 0),
        ("", 0),
        ("  ", 0),
        (0, 0),
        (0.0, 0),
        ("0", 0),
        (2, 2),
        (2.0, 2),
        ("2", 2),
        (" 2 ", 2),
        ("2,0", 2),
        (Decimal("3"), 3),
        (45, 45),
        (1.5, NON_NUMERIC),
        ("=O18+P18", NON_NUMERIC),
        ("dwa", NON_NUMERIC),
        (True, NON_NUMERIC),
        (STARTED, NON_NUMERIC),
    ],
    ids=repr,
)
def test_ilosc_taty_puste_i_zero_to_samo_liczba_tekstem_to_liczba_reszta_nieliczbowa(cell: object, value: object) -> None:
    assert tata_value(cell) == value


def test_ilosci_rowne_zero_nie_sa_zapamietane() -> None:
    row: list[object] = [None] * 46
    row[13], row[14], row[15], row[6] = 0, "0", 2, STARTED

    facts = _filtered(HEADER, row)

    assert facts.quantities == {"P": 2}


def test_wiersz_danych_bez_wykonawcy_tez_przechodzi_przez_filtr(tmp_path: Path) -> None:
    path = tmp_path / "Rozliczenie 08_14_09_2026 - zbiorcze.xlsx"
    write_quantity_source(path, [{"V": 1}, {"H": None, "V": 2}])

    read = read_data_rows(path)

    assert read.without_worker == [19]
    assert isinstance(read.orders, tuple)
    assert [facts.quantity("V") for facts in read.orders] == [1, 2]
