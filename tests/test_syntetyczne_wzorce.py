"""Syntetyczne Wzorce do benchmarku i CI mają kształt Wzorców taty, bez ich danych.

Generator leży w ``tools/``; testy uruchamiają jego CLI w małej odmianie (kilkaset wierszy
zamiast ~10 tys.), żeby pytest był szybki. Kształt części zip jest ten sam w każdym rozmiarze.
"""

from __future__ import annotations

import re
from pathlib import Path

import pytest
from openpyxl import load_workbook

from rozliczenia.worker_config import load_worker_mapping
from tests.fixtures.benchmark import (
    COLLECTIVE_TEMPLATE,
    INTERNAL_TEMPLATE,
    SMALL_DATA_ROWS,
    WORKER_TEMPLATE,
    write_synthetic_templates,
    zip_part,
    zip_part_names,
)
from tests.fixtures.workbooks import active_worksheet

LAST_ROW = 17 + SMALL_DATA_ROWS
# Wzorzec taty: AU:AY w każdym Wierszu danych oraz wspólne formuły N:AP i SUBTOTAL w wierszach 13 i 15.
FORMULAS_ABOVE_DATA_ROWS = 62
FORMULAS_PER_DATA_ROW = 5
# Wiersze z zerami w N:V i X: w małej odmianie połowa Wierszy danych (u taty 483 z ~10 tys.).
ZERO_ROWS = SMALL_DATA_ROWS // 2
INTERNAL_HEADERS = (
    "MIASTO",
    "ADRES",
    "UMOWA",
    "URZĄDZENIA",
    "CZYNNOŚCI",
    "NUMER ZL.",
    "DATA",
    "WYKONAWCA",
    "NUMER UMOWY/ZLECENIA ORANGE",
    "SKUTECZNE",
    "CZYNNOŚCI ZREALIZOWANE",
    "zwrot",
    "UWAGI",
)
SHEET = "xl/worksheets/sheet1.xml"


def formula_cells(sheet_xml: str) -> list[str]:
    return re.findall(r'<c r="([A-Z]+\d+)"[^>]*><f[ >/]', sheet_xml)


def rows_without_cell(sheet_xml: str, column: str) -> list[int]:
    rows = re.findall(r'<row r="(\d+)"[^>]*>(.*?)</row>', sheet_xml, re.S)
    return [int(number) for number, cells in rows if int(number) >= 18 and f'<c r="{column}{number}"' not in cells]


@pytest.mark.parametrize(
    ("template", "first_header"),
    [(WORKER_TEMPLATE, "GODZINA"), (COLLECTIVE_TEMPLATE, "MIASTO")],
)
def test_syntetyczny_wzorzec_ma_wspolne_formuly_calcchain_i_odwolanie_zewnetrzne_jak_wzorzec_taty(
    tmp_path: Path, template: str, first_header: str
) -> None:
    path = write_synthetic_templates(tmp_path / "wzorce") / template

    assert {
        "xl/calcChain.xml",
        "xl/sharedStrings.xml",
        "xl/externalLinks/externalLink1.xml",
        "xl/externalLinks/_rels/externalLink1.xml.rels",
    } <= zip_part_names(path)
    assert '<externalReference r:id="' in zip_part(path, "xl/workbook.xml")
    sheet_xml = zip_part(path, SHEET)
    formulas = formula_cells(sheet_xml)
    assert len(formulas) == FORMULAS_ABOVE_DATA_ROWS + FORMULAS_PER_DATA_ROW * SMALL_DATA_ROWS
    assert sorted(re.findall(r'<c r="(\w+)"', zip_part(path, "xl/calcChain.xml"))) == sorted(formulas)
    # Każda formuła ma zapisaną wartość, jak po zapisie w Excelu.
    assert len(re.findall(r"<f[^>]*(?:/>|>[^<]*</f>)<v>", sheet_xml)) == len(formulas)
    # Wspólne formuły AU: wzór co 64 wiersze, pozostałe komórki tylko odwołują się do niego.
    assert re.findall(r'<c r="(AU\d+)"[^>]*><f t="shared" ref="', sheet_xml) == ["AU19", "AU83", "AU147"]
    assert sheet_xml.count('<f t="shared" si="') > 3 * SMALL_DATA_ROWS
    assert '<f t="shared" ref="N13:AP13"' in sheet_xml
    assert rows_without_cell(sheet_xml, "A")
    assert rows_without_cell(sheet_xml, "B")

    workbook = load_workbook(path)
    try:
        sheet = active_worksheet(workbook)
        assert (sheet.title, sheet.max_row) == ("Sheet1", LAST_ROW)
        assert (sheet["A17"].value, sheet["H17"].value) == (first_header, "WYKONAWCA")
        assert [sheet.cell(row, 14).value for row in range(18, LAST_ROW + 1)] == [0] * ZERO_ROWS + [None] * (
            SMALL_DATA_ROWS - ZERO_ROWS
        )
        assert [sheet.cell(row, 24).value for row in (18, 17 + ZERO_ROWS, 18 + ZERO_ROWS)] == [0, 0, None]
        assert {sheet.cell(row, column).value for row in range(18, LAST_ROW + 1) for column in range(1, 13)} == {None}
        assert str(sheet["AU20"].value).startswith("=N20")
        assert "[1]" in str(sheet["AY18"].value)
        assert {"A16:M16", "AU13:AU15"} <= {str(merged) for merged in sheet.merged_cells.ranges}
        assert len({sheet.cell(row, 2).style_id for row in range(18, LAST_ROW + 1)}) >= 3
    finally:
        workbook.close()


def test_pelny_syntetyczny_wzorzec_ma_rozmiar_wzorca_taty(tmp_path: Path) -> None:
    sheet_xml = zip_part(write_synthetic_templates(tmp_path / "wzorce", data_rows=None) / WORKER_TEMPLATE, SHEET)

    # Wzorzec taty: Wiersze danych 18..10093, 496 710 komórek, 50 442 formuły, z tego 40 360 wspólnych.
    data_rows = [int(row) for row in re.findall(r'<row r="(\d+)"[^>]*>(?=<c )', sheet_xml) if int(row) >= 18]
    assert (min(data_rows), max(data_rows), len(data_rows)) == (18, 10_093, 10_076)
    assert 470_000 < sheet_xml.count("<c ") < 520_000
    assert len(formula_cells(sheet_xml)) == 50_442
    assert sheet_xml.count('<f t="shared"') > 40_000


def test_syntetyczny_wzorzec_rozliczenia_wewnetrznego_ma_naglowki_w_wierszu_1(tmp_path: Path) -> None:
    path = write_synthetic_templates(tmp_path / "wzorce") / INTERNAL_TEMPLATE

    workbook = load_workbook(path)
    try:
        sheet = active_worksheet(workbook)
        assert sheet.title == "Arkusz1"
        assert next(sheet.iter_rows(min_row=1, max_row=1, values_only=True)) == INTERNAL_HEADERS
        assert sheet.freeze_panes == "A2"
        assert {value for row in sheet.iter_rows(min_row=2, values_only=True) for value in row} == {None}
    finally:
        workbook.close()


def test_syntetyczne_wzorce_maja_konfiguracje_dwudziestu_wykonawcow(tmp_path: Path) -> None:
    directory = write_synthetic_templates(tmp_path / "wzorce")

    assert len(load_worker_mapping(directory / "worker_mapping.yaml")) == 20


def test_generator_tworzy_te_same_bajty_za_kazdym_razem(tmp_path: Path) -> None:
    first = write_synthetic_templates(tmp_path / "pierwsze")
    second = write_synthetic_templates(tmp_path / "drugie")

    names = sorted(path.name for path in first.iterdir())
    assert names == sorted(path.name for path in second.iterdir())
    assert len(names) == 4
    assert all((first / name).read_bytes() == (second / name).read_bytes() for name in names)
