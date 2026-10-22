"""Zatrzymania Przygotowania pliku zbiorczego przy problemach z Rejestrem zleceń i wzorcami."""

from __future__ import annotations

from collections.abc import Callable
from dataclasses import replace
from datetime import datetime
import os
from pathlib import Path

from openpyxl import Workbook, load_workbook
import pytest

from rozliczenia.message import StopError
from rozliczenia.workbook_files import ExternalLinksError, WorkbookSaveError
import rozliczenia.workbook_files as workbook_files

from tests.fixtures.preparation import (
    COLLECTIVE_FILE_NAME,
    INTERNAL_SETTLEMENT_NAME,
    PERIOD,
    REGISTER_HEADERS,
    PreparationFixture,
    RegisterRow,
    make_dad_layout_fixture,
    make_preparation_fixture,
    month_periods_folder,
    register_order,
    run_preparation,
    write_register,
    write_register_sheet,
)
from tests.fixtures.workbooks import (
    break_row_far_below_data_rows,
    cut_sheet_after_data_rows,
    without_sheet,
    without_styles,
)


def input_snapshot(fixture: PreparationFixture) -> dict[Path, bytes]:
    paths = (fixture.register_path, fixture.internal_template_path, fixture.collective_template_path)
    return {path: path.read_bytes() for path in paths if path.exists()}


def stop_preparation(fixture: PreparationFixture, *, dry_run: bool = False) -> StopError:
    """Uruchamia proces, oczekuje zatrzymania i sprawdza, że nic nie powstało ani nie zmieniło się."""

    inputs_before = input_snapshot(fixture)
    periods_before = sorted(path.name for path in fixture.periods_directory.iterdir())

    with pytest.raises(StopError) as raised:
        run_preparation(fixture, dry_run=dry_run)

    assert sorted(path.name for path in fixture.periods_directory.iterdir()) == periods_before
    assert input_snapshot(fixture) == inputs_before
    assert "SYN-" not in raised.value.message.text
    return raised.value


def test_missing_rejestr_zlecen_stops_without_files(tmp_path: Path) -> None:
    fixture = make_preparation_fixture(tmp_path)
    fixture.register_path.unlink()

    error = stop_preparation(fixture)

    assert error.message.code == "BRAK_REJESTRU"
    assert error.message.text == (
        "Co się stało: Nie znaleziono Rejestru zleceń. Nie utworzono żadnego pliku.\n"
        f"Gdzie: plik ZLECENIA INTERMAX OD 08_11_2021.xlsx, folder {tmp_path.name}\n"
        "Co zrobić: Wybierz właściwy plik Rejestru zleceń i uruchom ponownie."
    )


def test_rejestr_open_in_excel_stops_and_leaves_lock_file(tmp_path: Path) -> None:
    fixture = make_preparation_fixture(tmp_path)
    lock_file = fixture.register_path.with_name(f"~${fixture.register_path.name}")
    lock_file.write_bytes(b"")

    error = stop_preparation(fixture)

    assert error.message.code == "REJESTR_OTWARTY"
    assert error.message.text == (
        "Co się stało: Rejestr zleceń jest otwarty w Excelu. Nie utworzono żadnego pliku.\n"
        f"Gdzie: plik ZLECENIA INTERMAX OD 08_11_2021.xlsx, folder {tmp_path.name}\n"
        "Co zrobić: Zapisz i zamknij Rejestr zleceń w Excelu, potem uruchom ponownie. "
        "Jeśli Excel jest zamknięty, a komunikat wraca, usuń z tego samego folderu "
        "pozostawiony plik ~$ZLECENIA INTERMAX OD 08_11_2021.xlsx i uruchom ponownie."
    )
    assert lock_file.exists()


@pytest.mark.parametrize(
    ("removed", "label"),
    [
        ("SKUTECZNE", "SKUTECZNE"),
        ("ADRES", "ADRES"),
        ("NUMER UMOWY", "NUMER UMOWY albo NUMER UMOWY/ZLECENIA ORANGE"),
    ],
)
def test_rejestr_without_required_header_stops_with_header_name(tmp_path: Path, removed: str, label: str) -> None:
    fixture = make_preparation_fixture(tmp_path)
    write_register(
        fixture.register_path,
        headers=tuple(header for header in REGISTER_HEADERS if header != removed),
    )

    error = stop_preparation(fixture)

    assert error.message.code == "BRAK_NAGLOWKA_REJESTRU"
    assert error.message.text == (
        f"Co się stało: W Rejestrze zleceń brakuje wymaganego nagłówka {label}. "
        "Nie utworzono żadnego pliku.\n"
        "Gdzie: plik ZLECENIA INTERMAX OD 08_11_2021.xlsx, arkusz ZLECENIA, wiersz 1\n"
        f"Co zrobić: Dopisz albo popraw nagłówek {label} w pierwszym wierszu arkusza ze zleceniami, "
        "zapisz i zamknij Rejestr zleceń i uruchom ponownie."
    )


def test_rejestr_with_duplicated_required_header_stops_with_header_name(tmp_path: Path) -> None:
    fixture = make_preparation_fixture(tmp_path)
    write_register(fixture.register_path, headers=(*REGISTER_HEADERS, "WYKONAWCA "))

    error = stop_preparation(fixture)

    assert error.message.code == "POWTORZONY_NAGLOWEK_REJESTRU"
    assert error.message.text == (
        "Co się stało: W Rejestrze zleceń nagłówek WYKONAWCA występuje więcej niż raz, "
        "więc nie wiadomo, z której kolumny brać dane. Nie utworzono żadnego pliku.\n"
        "Gdzie: plik ZLECENIA INTERMAX OD 08_11_2021.xlsx, arkusz ZLECENIA, wiersz 1\n"
        "Co zrobić: Zostaw nagłówek WYKONAWCA tylko w jednej kolumnie, "
        "zapisz i zamknij Rejestr zleceń i uruchom ponownie."
    )


def test_rejestr_with_both_spellings_of_one_header_stops_as_duplicated(tmp_path: Path) -> None:
    fixture = make_preparation_fixture(tmp_path)
    write_register(fixture.register_path, headers=(*REGISTER_HEADERS, "nr zlecenia"))

    error = stop_preparation(fixture)

    assert error.message.code == "POWTORZONY_NAGLOWEK_REJESTRU"
    assert error.message.what.startswith(
        "W Rejestrze zleceń nagłówek NUMER ZL. albo nr zlecenia występuje więcej niż raz"
    )


def test_rejestr_with_header_repeated_in_other_letter_case_names_it_once(tmp_path: Path) -> None:
    fixture = make_preparation_fixture(tmp_path)
    write_register(fixture.register_path, headers=(*REGISTER_HEADERS, "adres"))

    error = stop_preparation(fixture)

    assert error.message.code == "POWTORZONY_NAGLOWEK_REJESTRU"
    assert error.message.what.startswith("W Rejestrze zleceń nagłówek ADRES występuje więcej niż raz")


def test_rejestr_without_sheet_with_required_headers_stops(tmp_path: Path) -> None:
    fixture = make_preparation_fixture(tmp_path)
    write_register(fixture.register_path, headers=("KLIENT", "TELEFON", "WYKONAWCA"))

    error = stop_preparation(fixture)

    assert error.message.code == "BRAK_ARKUSZA_ZLECEN"
    assert error.message.text == (
        "Co się stało: W Rejestrze zleceń nie znaleziono arkusza ze zleceniami, czyli arkusza "
        "z wymaganymi nagłówkami w pierwszym wierszu. Nie utworzono żadnego pliku.\n"
        "Gdzie: plik ZLECENIA INTERMAX OD 08_11_2021.xlsx\n"
        "Co zrobić: Sprawdź, czy wybrano właściwy plik Rejestru zleceń. Pierwszy wiersz arkusza "
        "ze zleceniami musi zawierać nagłówki: DATA GODZ. POCZ., MIASTO, ADRES, UMOWA, URZĄDZENIA, "
        "CZYNNOŚCI, NUMER ZL. albo nr zlecenia, WYKONAWCA, NUMER UMOWY albo NUMER UMOWY/ZLECENIA ORANGE, "
        "SKUTECZNE, "
        "CZYNNOŚCI ZREALIZOWANE."
    )


def add_register_sheet(
    register_path: Path,
    title: str,
    rows: list[RegisterRow],
    *,
    headers: tuple[str, ...] = REGISTER_HEADERS,
    hidden: bool = False,
    first: bool = False,
) -> None:
    """Dopisuje do syntetycznego Rejestru arkusz z nagłówkami w wierszu 1 i wierszami od 2."""

    workbook = load_workbook(register_path)
    sheet = workbook.create_sheet(title, 0 if first else None)
    write_register_sheet(sheet, rows, headers=headers)
    if hidden:
        sheet.sheet_state = "hidden"
    workbook.save(register_path)
    workbook.close()


def test_decoy_sheets_do_not_hide_sheet_with_orders(tmp_path: Path) -> None:
    fixture = make_preparation_fixture(tmp_path)
    add_register_sheet(
        fixture.register_path,
        "komentarze",
        [{"UWAGI": "sprawdzić w poniedziałek", "WYKONAWCA": "adrian.testowy"}],
        headers=("DATA GODZ. POCZ.", "WYKONAWCA", "UWAGI"),
        first=True,
    )
    add_register_sheet(
        fixture.register_path,
        "słowniki",
        [{"MIASTO": "POZNAŃ", "UMOWA": "abonament", "SKUTECZNE": "tak"}],
        headers=("MIASTO", "UMOWA", "SKUTECZNE", "URZĄDZENIA"),
        hidden=True,
    )

    summary = run_preparation(fixture)

    assert (summary.period_order_count, summary.successful_order_count) == (4, 3)
    assert fixture.internal_settlement_path.is_file()
    assert fixture.collective_file_path.is_file()


def test_period_orders_in_more_than_one_sheet_stop_with_sheet_names(tmp_path: Path) -> None:
    fixture = make_preparation_fixture(tmp_path)
    add_register_sheet(fixture.register_path, "ZLECENIA 2", [register_order(7, datetime(2026, 9, 10, 9, 0))])

    error = stop_preparation(fixture)

    assert error.message.code == "ZLECENIA_W_WIELU_ARKUSZACH"
    assert error.message.text == (
        "Co się stało: Zlecenia tego okresu są w więcej niż jednym arkuszu Rejestru zleceń, "
        "więc nie wiadomo, który arkusz jest właściwy. Nie utworzono żadnego pliku.\n"
        f"Gdzie: plik ZLECENIA INTERMAX OD 08_11_2021.xlsx, arkusze ZLECENIA, ZLECENIA 2, okres {PERIOD}\n"
        "Co zrobić: Zostaw zlecenia tego okresu tylko w jednym arkuszu, "
        "zapisz i zamknij Rejestr zleceń i uruchom ponownie."
    )


def test_sheet_with_orders_of_other_periods_does_not_block_period(tmp_path: Path) -> None:
    fixture = make_preparation_fixture(tmp_path)
    add_register_sheet(
        fixture.register_path,
        "archiwum 2025",
        [register_order(8, datetime(2025, 9, 9, 9, 0)), None, register_order(9, datetime(2025, 9, 10, 9, 0))],
        hidden=True,
        first=True,
    )

    summary = run_preparation(fixture)

    assert (summary.period_order_count, summary.successful_order_count) == (4, 3)


@pytest.mark.parametrize(
    "damage",
    [
        lambda path: path.write_bytes(b"to nie jest skoroszyt Excela"),
        lambda path: cut_sheet_after_data_rows(path, "xl/worksheets/sheet2.xml"),  # ZLECENIA za arkuszem lista
        without_styles,
        without_sheet,
    ],
    ids=["nie-skoroszyt", "urwany-arkusz-zlecen", "bez-styles-xml", "bez-arkusza"],
)
def test_damaged_rejestr_zlecen_stops_without_files(tmp_path: Path, damage: Callable[[Path], object]) -> None:
    fixture = make_preparation_fixture(tmp_path)
    damage(fixture.register_path)

    error = stop_preparation(fixture)

    assert error.message.code == "USZKODZONY_REJESTR"
    assert error.message.text == (
        "Co się stało: Nie można odczytać Rejestru zleceń. Plik jest uszkodzony "
        "albo nie jest skoroszytem .xlsx. Nie utworzono żadnego pliku.\n"
        f"Gdzie: plik ZLECENIA INTERMAX OD 08_11_2021.xlsx, folder {tmp_path.name}\n"
        "Co zrobić: Sprawdź, czy Rejestr zleceń otwiera się w Excelu i jest zapisany jako .xlsx, "
        "zamknij go i uruchom ponownie. Jeśli błąd wraca, przekaż to osobie utrzymującej narzędzie."
    )


WZORCE = pytest.mark.parametrize(
    "template_name",
    ["Wzorzec rozliczenia wewnętrznego.xlsx", "Wzorzec pliku zbiorczego.xlsx"],
)


def template_path(fixture: PreparationFixture, template_name: str) -> Path:
    path = fixture.internal_template_path.with_name(template_name)
    assert path in (fixture.internal_template_path, fixture.collective_template_path)
    return path


@WZORCE
def test_missing_wzorzec_stops_without_files(tmp_path: Path, template_name: str) -> None:
    fixture = make_preparation_fixture(tmp_path)
    template_path(fixture, template_name).unlink()

    error = stop_preparation(fixture)

    assert error.message.code == "BRAK_WZORCA"
    assert error.message.text == (
        "Co się stało: Nie znaleziono wzorca, z którego powstają pliki okresu. Nie utworzono żadnego pliku.\n"
        f"Gdzie: plik {template_name}, folder config\n"
        "Co zrobić: Przekaż ten komunikat osobie utrzymującej narzędzie."
    )


@WZORCE
def test_wzorzec_open_in_excel_stops_and_leaves_lock_file(tmp_path: Path, template_name: str) -> None:
    fixture = make_preparation_fixture(tmp_path)
    lock_file = template_path(fixture, template_name).with_name(f"~${template_name}")
    lock_file.write_bytes(b"")

    error = stop_preparation(fixture)

    assert error.message.code == "WZORZEC_OTWARTY"
    assert error.message.text == (
        "Co się stało: Wzorzec jest otwarty w Excelu. Nie utworzono żadnego pliku.\n"
        f"Gdzie: plik {template_name}, folder config\n"
        "Co zrobić: Zamknij wzorzec w Excelu bez zapisywania zmian i uruchom ponownie. "
        "Jeśli Excel jest zamknięty, a komunikat wraca, usuń z tego samego folderu "
        f"pozostawiony plik ~${template_name} i uruchom ponownie."
    )
    assert lock_file.exists()


@WZORCE
@pytest.mark.parametrize(
    "damage",
    [lambda path: path.write_bytes(b"uszkodzony wzorzec"), without_styles, break_row_far_below_data_rows],
    ids=["nie-skoroszyt", "bez-styles-xml", "zepsuty-wiersz-daleko-pod-danymi"],
)
@pytest.mark.parametrize("dry_run", [False, True])
def test_damaged_wzorzec_stops_before_reading_rejestr(
    tmp_path: Path, template_name: str, damage: Callable[[Path], object], dry_run: bool
) -> None:
    fixture = make_preparation_fixture(tmp_path)
    damage(template_path(fixture, template_name))
    fixture.register_path.write_bytes(b"uszkodzony Rejestr")

    error = stop_preparation(fixture, dry_run=dry_run)

    assert error.message.code == "USZKODZONY_WZORZEC"
    assert error.message.text == (
        "Co się stało: Nie można odczytać wzorca. Plik jest uszkodzony "
        "albo nie jest skoroszytem .xlsx. Nie utworzono żadnego pliku.\n"
        f"Gdzie: plik {template_name}, folder config\n"
        "Co zrobić: Przekaż ten komunikat osobie utrzymującej narzędzie."
    )


@pytest.mark.parametrize(
    ("template_name", "header_cell", "header", "sheet", "row"),
    [
        ("Wzorzec rozliczenia wewnętrznego.xlsx", "G1", "NUMER ZL.", "Arkusz1", 1),
        ("Wzorzec pliku zbiorczego.xlsx", "H17", "WYKONAWCA", "Sheet1", 17),
    ],
)
@pytest.mark.parametrize("dry_run", [False, True])
def test_wzorzec_without_expected_header_stops_before_reading_rejestr(
    tmp_path: Path, template_name: str, header_cell: str, header: str, sheet: str, row: int, dry_run: bool
) -> None:
    fixture = make_preparation_fixture(tmp_path)
    path = template_path(fixture, template_name)
    workbook = load_workbook(path, keep_links=True)
    workbook[sheet][header_cell].value = "inna kolumna"
    workbook.save(path)
    workbook.close()
    fixture.register_path.write_bytes(b"uszkodzony Rejestr")

    error = stop_preparation(fixture, dry_run=dry_run)

    assert error.message.code == "ZLY_WZORZEC"
    assert error.message.text == (
        f"Co się stało: We wzorcu brakuje nagłówka {header}, więc nie wiadomo, gdzie wpisać zlecenia. "
        "Nie utworzono żadnego pliku.\n"
        f"Gdzie: plik {template_name}, arkusz {sheet}, wiersz {row}\n"
        "Co zrobić: Przekaż ten komunikat osobie utrzymującej narzędzie."
    )


@pytest.mark.parametrize("dry_run", [False, True])
def test_wzorzec_pliku_zbiorczego_with_wykonawca_outside_h_stops_before_reading_rejestr(
    tmp_path: Path, dry_run: bool
) -> None:
    fixture = make_preparation_fixture(tmp_path)
    workbook = load_workbook(fixture.collective_template_path, keep_links=True)
    sheet = workbook["Sheet1"]
    sheet["G17"].value, sheet["H17"].value = sheet["H17"].value, sheet["G17"].value
    workbook.save(fixture.collective_template_path)
    workbook.close()
    fixture.register_path.write_bytes(b"uszkodzony Rejestr")

    error = stop_preparation(fixture, dry_run=dry_run)

    assert error.message.code == "WYKONAWCA_POZA_H"
    assert error.message.text == (
        "Co się stało: Wzorzec pliku zbiorczego ma nagłówek WYKONAWCA w kolumnie G zamiast w H (wiersz 17). "
        "Nie utworzono żadnego pliku.\n"
        "Gdzie: plik Wzorzec pliku zbiorczego.xlsx, arkusz Sheet1, wiersz 17\n"
        "Co zrobić: Przenieś WYKONAWCA do kolumny H we wzorcu i uruchom ponownie."
    )


@pytest.mark.parametrize(
    ("template_name", "header_cell", "register_header", "header", "sheet", "row"),
    [
        ("Wzorzec rozliczenia wewnętrznego.xlsx", "G1", "nr zlecenia", "NUMER ZL.", "Arkusz1", 1),
        ("Wzorzec rozliczenia wewnętrznego.xlsx", "A1", "DATA GODZ. POCZ.", "DATA", "Arkusz1", 1),
        ("Wzorzec pliku zbiorczego.xlsx", "F17", "nr zlecenia", "NUMER ZL.", "Sheet1", 17),
    ],
)
@pytest.mark.parametrize("dry_run", [False, True])
def test_wzorzec_with_rejestr_header_instead_of_its_own_stops(
    tmp_path: Path,
    template_name: str,
    header_cell: str,
    register_header: str,
    header: str,
    sheet: str,
    row: int,
    dry_run: bool,
) -> None:
    fixture = make_preparation_fixture(tmp_path)
    path = template_path(fixture, template_name)
    workbook = load_workbook(path, keep_links=True)
    workbook[sheet][header_cell].value = register_header
    workbook.save(path)
    workbook.close()

    error = stop_preparation(fixture, dry_run=dry_run)

    assert error.message.code == "ZLY_WZORZEC"
    assert error.message.where == f"plik {template_name}, arkusz {sheet}, wiersz {row}"
    assert f"brakuje nagłówka {header}," in error.message.what


@pytest.mark.parametrize(
    ("template_name", "merged_range", "sheet"),
    [
        ("Wzorzec rozliczenia wewnętrznego.xlsx", "C5:D5", "Arkusz1"),  # 4. zlecenie okresu
        ("Wzorzec pliku zbiorczego.xlsx", "B19:C20", "Sheet1"),  # 2. i 3. Zlecenie skuteczne
    ],
)
@pytest.mark.parametrize("dry_run", [False, True])
def test_wzorzec_with_merged_cells_where_orders_go_stops_without_files(
    tmp_path: Path, template_name: str, merged_range: str, sheet: str, dry_run: bool
) -> None:
    fixture = make_preparation_fixture(tmp_path)
    path = template_path(fixture, template_name)
    workbook = load_workbook(path, keep_links=True)
    workbook[sheet].merge_cells(merged_range)
    workbook.save(path)
    workbook.close()

    error = stop_preparation(fixture, dry_run=dry_run)

    assert error.message.code == "SCALONE_KOMORKI_WZORCA"
    assert error.message.text == (
        "Co się stało: We wzorcu są scalone komórki tam, gdzie trafiają zlecenia, więc nie można ich wpisać. "
        "Nie utworzono żadnego pliku.\n"
        f"Gdzie: plik {template_name}, arkusz {sheet}, komórki {merged_range}\n"
        "Co zrobić: Przekaż ten komunikat osobie utrzymującej narzędzie."
    )


def test_merged_cells_of_wzorzec_below_written_orders_do_not_stop(tmp_path: Path) -> None:
    fixture = make_preparation_fixture(tmp_path)
    workbook = load_workbook(fixture.collective_template_path, keep_links=True)
    workbook["Sheet1"].merge_cells("A21:K21")  # pod trzema Zleceniami skutecznymi z wiersza 18
    workbook.save(fixture.collective_template_path)
    workbook.close()

    summary = run_preparation(fixture)

    assert summary.successful_order_count == 3
    assert fixture.collective_file_path.is_file()


@pytest.mark.parametrize("dry_run", [False, True])
def test_in_dad_layout_missing_folder_okresu_stops_without_creating_it(tmp_path: Path, dry_run: bool) -> None:
    fixture = make_dad_layout_fixture(tmp_path)
    fixture.period_directory.rmdir()
    tree_before = sorted(fixture.periods_directory.rglob("*"))

    with pytest.raises(StopError) as raised:
        run_preparation(fixture, dry_run=dry_run)

    assert raised.value.message.code == "BRAK_FOLDERU_OKRESU"
    assert raised.value.message.text == (
        "Co się stało: Nie znaleziono folderu 08_14_09_2026 w żadnym folderze miesiąca. "
        "Nie utworzono żadnego pliku.\n"
        "Gdzie: foldery INEA w folderze Załączniki\n"
        "Co zrobić: Załóż folder 08_14_09_2026 w folderze INEA właściwego miesiąca i uruchom ponownie."
    )
    assert sorted(fixture.periods_directory.rglob("*")) == tree_before


def test_in_dad_layout_folder_inea_of_another_month_does_not_get_missing_folder_okresu(tmp_path: Path) -> None:
    fixture = make_dad_layout_fixture(tmp_path)
    fixture.period_directory.rmdir()
    august = fixture.periods_directory / "Rozliczenia 2026" / "08_Sierpień_2026" / "INEA"
    tree_before = sorted(fixture.periods_directory.rglob("*"))

    with pytest.raises(StopError) as raised:
        run_preparation(replace(fixture, periods_directory=august))

    assert raised.value.message.code == "BRAK_FOLDERU_OKRESU"
    assert raised.value.message.where == "foldery INEA w folderze Załączniki"
    assert sorted(fixture.periods_directory.rglob("*")) == tree_before


def test_in_dad_layout_folder_okresu_in_two_months_stops_with_both_months(tmp_path: Path) -> None:
    fixture = make_dad_layout_fixture(tmp_path)
    (month_periods_folder(fixture.periods_directory, "10_Październik_2026") / PERIOD).mkdir()

    with pytest.raises(StopError) as raised:
        run_preparation(fixture)

    assert raised.value.message.code == "POWTORZONY_FOLDER_OKRESU"
    assert raised.value.message.where == (
        "Rozliczenia 2026\\09_Wrzesień_2026\\INEA, Rozliczenia 2026\\10_Październik_2026\\INEA"
    )
    assert [list(folder.iterdir()) for folder in fixture.periods_directory.glob(f"*/*/INEA/{PERIOD}")] == [[], []]


def test_missing_folder_okresow_stops_without_creating_it(tmp_path: Path) -> None:
    fixture = make_preparation_fixture(tmp_path)
    fixture.periods_directory.rmdir()

    with pytest.raises(StopError) as raised:
        run_preparation(fixture)

    assert raised.value.message.code == "BRAK_FOLDERU_OKRESOW"
    assert raised.value.message.text == (
        "Co się stało: Nie znaleziono folderu okresów. Nie utworzono żadnego pliku.\n"
        f"Gdzie: folder {tmp_path.name}\\INEA\n"
        "Co zrobić: Wybierz istniejący folder okresów, w którym są foldery poprzednich tygodni, "
        "i uruchom ponownie."
    )
    assert not fixture.periods_directory.exists()


def test_zlamana_regula_konfiguracji_wykonawcow_zatrzymuje_przygotowanie_ze_szczegolem(tmp_path: Path) -> None:
    fixture = make_preparation_fixture(tmp_path)
    fixture.config_path.write_text("schema_version: 1\nworkers:\n  a: Ten Sam\n  b: ten sam\n", encoding="utf-8")

    error = stop_preparation(fixture)

    assert error.message.code == "ZLA_KONFIGURACJA_WYKONAWCOW"
    assert error.message.text == (
        "Co się stało: Nie można odczytać konfiguracji wykonawców. "
        "Dwóch identyfikatorów wskazuje ten sam plik pracownika: Ten Sam / ten sam\n"
        "Gdzie: plik worker_mapping.yaml\n"
        "Co zrobić: Przekaż to osobie utrzymującej narzędzie."
    )


def test_nieczytelna_konfiguracja_wykonawcow_zatrzymuje_przygotowanie_bez_powtorzenia(tmp_path: Path) -> None:
    fixture = make_preparation_fixture(tmp_path)
    fixture.config_path.unlink()
    fixture.config_path.mkdir()

    error = stop_preparation(fixture)

    assert error.message.code == "ZLA_KONFIGURACJA_WYKONAWCOW"
    assert error.message.text == (
        "Co się stało: Nie można odczytać konfiguracji wykonawców.\n"
        "Gdzie: plik worker_mapping.yaml\n"
        "Co zrobić: Przekaż to osobie utrzymującej narzędzie."
    )


@pytest.mark.parametrize(
    ("failure", "period_directory_exists"),
    [
        (OSError("wymuszony błąd dysku"), False),
        (ExternalLinksError("wymuszony błąd linków"), True),
        (WorkbookSaveError("wymuszony błąd zapisu"), True),
    ],
)
def test_failed_save_of_plik_zbiorczy_leaves_no_period_file(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, period_directory_exists: bool, failure: Exception
) -> None:
    fixture = make_preparation_fixture(tmp_path)
    if period_directory_exists:
        fixture.period_directory.mkdir()
    real_save = workbook_files.save_workbook
    saved: list[str] = []

    def save_first_then_fail(workbook: Workbook, target: Path, *, external_links_from: Path | None = None) -> None:
        if target.name == COLLECTIVE_FILE_NAME:
            raise failure
        saved.append(target.name)
        real_save(workbook, target, external_links_from=external_links_from)

    monkeypatch.setattr(workbook_files, "save_workbook", save_first_then_fail)

    error = stop_preparation(fixture)

    assert saved == [INTERNAL_SETTLEMENT_NAME]
    assert error.message.code == "BLAD_ZAPISU"
    assert "wymuszony" not in error.message.text
    assert [path.name for path in fixture.periods_directory.iterdir()] == ([PERIOD] if period_directory_exists else [])
    if period_directory_exists:
        assert list(fixture.period_directory.iterdir()) == []


def test_failed_publication_of_plik_zbiorczy_removes_published_rozliczenie_wewnetrzne(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    fixture = make_preparation_fixture(tmp_path)
    fixture.period_directory.mkdir()
    real_rename = os.rename

    def rename_fails_for_plik_zbiorczy(source: str | os.PathLike[str], target: str | os.PathLike[str]) -> None:
        if Path(target) == fixture.collective_file_path:
            raise OSError("wymuszony błąd publikacji")
        real_rename(source, target)

    monkeypatch.setattr(os, "rename", rename_fails_for_plik_zbiorczy)

    error = stop_preparation(fixture)

    assert error.message.code == "BLAD_ZAPISU"
    assert list(fixture.period_directory.iterdir()) == []


def test_failed_removal_of_published_rozliczenie_wewnetrzne_names_the_file_left(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    fixture = make_preparation_fixture(tmp_path)
    fixture.period_directory.mkdir()
    real_rename, real_unlink = os.rename, os.unlink

    def rename_fails_for_plik_zbiorczy(source: str | os.PathLike[str], target: str | os.PathLike[str]) -> None:
        if Path(target) == fixture.collective_file_path:
            raise OSError("wymuszony błąd publikacji")
        real_rename(source, target)

    def unlink_fails_for_rozliczenie_wewnetrzne(path: str | os.PathLike[str], **kwargs: object) -> None:
        if Path(path) == fixture.internal_settlement_path:
            raise PermissionError("wymuszony: plik używany przez inny proces")
        real_unlink(path, **kwargs)  # type: ignore[arg-type]

    monkeypatch.setattr(os, "rename", rename_fails_for_plik_zbiorczy)
    monkeypatch.setattr(os, "unlink", unlink_fails_for_rozliczenie_wewnetrzne)

    error = stop_preparation(fixture)

    assert error.message.code == "BLAD_ZAPISU_POZOSTAL_PLIK"
    assert error.message.text == (
        "Co się stało: Zapisano tylko jeden z dwóch plików okresu i nie udało się go usunąć.\n"
        f"Gdzie: folder INEA\\{PERIOD}, plik {INTERNAL_SETTLEMENT_NAME}\n"
        "Co zrobić: Usuń wskazany plik, bo bez drugiego pliku okresu jest niepełny, i uruchom ponownie. "
        "Jeśli nie daje się go usunąć, odczekaj chwilę. Jeśli błąd wraca, przekaż komunikat osobie "
        "utrzymującej narzędzie."
    )
    assert [path.name for path in fixture.period_directory.iterdir()] == [INTERNAL_SETTLEMENT_NAME]


@pytest.mark.parametrize("period_directory_exists", [False, True])
def test_staging_folder_that_cannot_be_removed_after_publication_does_not_stop_the_run(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, period_directory_exists: bool
) -> None:
    fixture = make_preparation_fixture(tmp_path)
    if period_directory_exists:
        fixture.period_directory.mkdir()
    real_rmdir = os.rmdir

    def rmdir_fails_for_staging_folder(path: str | os.PathLike[str], **kwargs: object) -> None:
        if Path(path).name.startswith(f".{PERIOD}.staging-"):
            raise PermissionError("wymuszony: folder trzyma OneDrive")
        real_rmdir(path, **kwargs)  # type: ignore[arg-type]

    monkeypatch.setattr(os, "rmdir", rmdir_fails_for_staging_folder)

    run_preparation(fixture)

    assert sorted(path.name for path in fixture.period_directory.iterdir()) == sorted(
        [INTERNAL_SETTLEMENT_NAME, COLLECTIVE_FILE_NAME]
    )
