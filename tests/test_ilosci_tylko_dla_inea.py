"""Ilości tylko dla INEA (ADR 0018, #457): Rozliczenie pracownika bez AP przy ``gniazdo`` i bez AK przy ``technicolor``.

Plik zbiorczy jest syntetyczny; wszystkie Wiersze danych należą do Wykonawcy ``adrian.testowy`` („Adrian Testowy”).
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
import pytest

from rozliczenia.engine import run_settlements
from rozliczenia.quantity_rules import QuantityColumn
from tests.fixtures.settlement import PERIOD, CellValue, make_fixture, write_quantity_source
from tests.fixtures.workbooks import active_worksheet, hyperlinks, marked_for_copy, refused_by_copy, written_by_copy

AP, AK = QuantityColumn.AP, QuantityColumn.AK
HIGHLIGHT = "FFFFFF00"
LINK = "https://przyklad.invalid/ilosc"

ROWS: list[dict[str, CellValue]] = [
    {"D": "nokia zte, gniazdo", "AP": 1, "AK": 1},
    {"D": "2*globo Technicolor", "AK": 2, "AP": 1},
    {"D": "GNIAZDO", "AK": 3},
    {"D": "globo", "AP": 1, "AK": 1},
]
EMPTIED = {("AP", 18), ("AP", 20), ("AK", 19)}


def a_at(path: Path, rows: range) -> dict[tuple[str, int], object]:
    workbook = load_workbook(path)
    try:
        sheet = active_worksheet(workbook)
        return {
            (get_column_letter(column), number): value
            for number, values in enumerate(
                sheet.iter_rows(min_row=rows.start, max_row=rows.stop - 1, max_col=46, values_only=True),
                start=rows.start,
            )
            for column, value in enumerate(values, start=1)
        }
    finally:
        workbook.close()


def highlight_and_link_ap18(path: Path) -> None:
    workbook = load_workbook(path)
    try:
        cell = active_worksheet(workbook)["AP18"]
        cell.fill = PatternFill(fill_type="solid", fgColor=HIGHLIGHT)
        cell.hyperlink = LINK
        workbook.save(path)
    finally:
        workbook.close()


@pytest.mark.parametrize("route", ["kopia", "openpyxl"])
def test_rozliczenie_pracownika_nie_ma_ilosci_tylko_dla_inea_a_plik_zbiorczy_i_kontrola_ilosci_je_maja(
    tmp_path: Path, route: str
) -> None:
    source_path, target_directory, config_path = make_fixture(tmp_path)
    template_path = tmp_path / "Wzorzec rozliczenia pracownika.xlsx"
    if route == "kopia":
        marked_for_copy(template_path)
    else:
        refused_by_copy(template_path)
    write_quantity_source(source_path, ROWS)
    highlight_and_link_ap18(source_path)
    source_before = source_path.read_bytes()

    preview = run_settlements(source_path, config_path, worker_template_path=template_path, dry_run=True)
    summary = run_settlements(source_path, config_path, worker_template_path=template_path)

    assert dict(summary.inea_only_cleared) == {AP: 1, AK: 1}, "bez pustego AP w wierszu GNIAZDO"
    assert dict(preview.inea_only_cleared) == dict(summary.inea_only_cleared)
    assert source_path.read_bytes() == source_before
    assert isinstance(summary.orders, tuple)
    assert [(order.quantity("AP"), order.quantity("AK")) for order in summary.orders] == [(1, 1), (1, 2), (0, 3), (1, 1)]
    written = target_directory / f"Rozliczenie {PERIOD} - Adrian Testowy.xlsx"
    assert written_by_copy(written) is (route == "kopia")
    source, settlement = a_at(source_path, range(18, 22)), a_at(written, range(18, 22))
    assert {key: settlement[key] for key in EMPTIED} == dict.fromkeys(EMPTIED)
    assert {key: value for key, value in settlement.items() if key not in EMPTIED} == {
        key: value for key, value in source.items() if key not in EMPTIED
    }
    workbook = load_workbook(written)
    try:
        fill = active_worksheet(workbook)["AP18"].fill
        assert (fill.fill_type, fill.fgColor.rgb) == ("solid", HIGHLIGHT), "wyczyszczona komórka z Wyróżnieniem"
    finally:
        workbook.close()
    assert "AP18" not in hyperlinks(written)


def test_bez_naglowka_urzadzenia_rozliczenie_pracownika_ma_ilosci_jak_plik_zbiorczy(tmp_path: Path) -> None:
    source_path, target_directory, config_path = make_fixture(tmp_path)
    write_quantity_source(source_path, ROWS)
    workbook = load_workbook(source_path)
    try:
        active_worksheet(workbook)["D17"] = None
        workbook.save(source_path)
    finally:
        workbook.close()

    summary = run_settlements(
        source_path, config_path, worker_template_path=tmp_path / "Wzorzec rozliczenia pracownika.xlsx"
    )

    assert dict(summary.inea_only_cleared) == {AP: 0, AK: 0}
    written = target_directory / f"Rozliczenie {PERIOD} - Adrian Testowy.xlsx"
    assert a_at(written, range(18, 22)) == a_at(source_path, range(18, 22))
