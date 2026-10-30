"""Moduł skoroszytu przez jego interfejs: otwarcie z jedną klasyfikacją, układ arkusza i plik z Wzorca."""

from __future__ import annotations

import re
import shutil
import sys
from copy import copy
from dataclasses import replace
import xml.etree.ElementTree as ET
from collections.abc import Callable
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import IO, Any, cast

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Color, Font, PatternFill
from openpyxl.styles.fills import Fill
from openpyxl.styles.numbers import is_date_format
from openpyxl.utils import get_column_letter

from rozliczenia.highlights import Highlights
from rozliczenia.preparation_register import TARGET_HEADERS
from rozliczenia.quantity_rules import QuantityColumn
from rozliczenia.workbook_files import (
    DataRow,
    ExcelRow,
    MissingHeader,
    Problem,
    RegisterRow,
    WorkbookError,
    read_data_rows,
    read_workbook,
    workbook_content,
)
from rozliczenia.template_copy import Package
from rozliczenia.template_file import Route, TemplateFile
from rozliczenia.workbook_save import (
    ExternalLinksError,
    SaveProblem,
    WorkbookSaveError,
    save_from_template,
    save_workbook,
)

from tests.fixtures.benchmark import (
    COLLECTIVE_TEMPLATE,
    SMALL_DATA_ROWS,
    WORKER_TEMPLATE,
    write_synthetic_templates,
    zip_part,
)
from tests.fixtures.preparation import INTERNAL_TEMPLATE_HEADERS, write_collective_template, write_internal_template
from tests.fixtures.porownanie import compare_worker_settlements, renumber_cell_styles
from tests.fixtures.workbooks import (
    active_worksheet,
    add_external_link_fixture,
    break_empty_data_row,
    break_row_far_below_data_rows,
    cut_sheet_after_data_rows,
    cut_sheet_after_sheet_data,
    external_link_parts,
    has_xml_comment,
    hyperlinks,
    marked_for_copy,
    refused_by_copy,
    written_by_copy,
    replace_cells,
    rewrite_parts,
    without_sheet,
    without_styles,
)
from tests.fixtures.data_rows import data_row, register_row
from tests.fixtures.locked_file import lock_file
from tests.fixtures.settlement import COLLECTIVE_HEADERS, write_source, write_template


def not_a_zip(path: Path) -> None:
    path.write_bytes(b"to nie jest skoroszyt")


def with_target_headers(path: Path) -> None:
    """Nagłówki docelowe pliku zbiorczego w A17:K17, z WYKONAWCA w H17."""

    workbook = load_workbook(path)
    try:
        sheet = active_worksheet(workbook)
        for column, header in enumerate(COLLECTIVE_HEADERS, start=1):
            sheet.cell(17, column).value = header
        workbook.save(path)
    finally:
        workbook.close()


def styled_template(path: Path) -> Path:
    """Skoroszyt w układzie Pliku zbiorczego z nagłówkami docelowymi i komórką ze stylem w Wierszu danych."""

    write_template(path)
    with_target_headers(path)
    workbook = load_workbook(path)
    try:
        sheet = active_worksheet(workbook)
        sheet["N18"].font = Font(bold=True)
        workbook.save(path)
    finally:
        workbook.close()
    return path


def _read_with_format(path: Path) -> None:
    with read_workbook(path, workbook_content(path)) as workbook:
        for sheet in workbook.sheets():
            list(sheet.data_rows(min_row=1, columns={get_column_letter(column): column for column in range(1, 15)}))


def _check_template(path: Path) -> None:
    TemplateFile.collective_file(path).check()


def _plan_template(path: Path) -> None:
    TemplateFile.collective_file(path).plan(1)


def _period_file(path: Path) -> None:
    target = path.with_name("Plik z Wzorca.xlsx")
    try:
        TemplateFile.collective_file(path).create(target, [])
    finally:
        assert not target.exists()


def _check_worker_settlement_template(path: Path) -> None:
    TemplateFile.worker_settlement(path).check()


def _worker_settlement(path: Path) -> None:
    target = path.with_name("Rozliczenie - Wykonawca.xlsx")
    try:
        TemplateFile.worker_settlement(path).create(target, [_data_row(_row(A="syntetyczne miasto"))])
    finally:
        assert not target.exists()


OPENINGS: dict[str, Callable[[Path], object]] = {
    "odczyt": read_data_rows,
    "odczyt-z-formatem": _read_with_format,
    "wzorzec": _check_template,
    "planowanie-wzorca": _plan_template,
    "plik-okresu-z-wzorca": _period_file,
    "sprawdzanie-wzorca-rozliczenia-pracownika": _check_worker_settlement_template,
    "rozliczenie-pracownika": _worker_settlement,
}


DAMAGES: dict[str, Callable[[Path], None]] = {
    "nie-zip": not_a_zip,
    "uciety-xml-arkusza": cut_sheet_after_data_rows,
    "xml-arkusza-uciety-za-wierszami-danych": cut_sheet_after_sheet_data,
    "zepsuty-pusty-wiersz-danych": break_empty_data_row,
    "zepsuty-wiersz-daleko-pod-danymi": break_row_far_below_data_rows,
    "bez-styles-xml": without_styles,
    "bez-arkusza": without_sheet,
}
# Odczyt sięga po style tylko w Wierszach danych, a ten skoroszyt ich nie ma, więc bez styles.xml daje się przeczytać.
UNREADABLE_CASES = [
    pytest.param(damage, opening, id=f"{damage}-{opening}")
    for damage in DAMAGES
    for opening in OPENINGS
    if (damage, opening) != ("bez-styles-xml", "odczyt")
]


@pytest.mark.parametrize(("damage", "opening"), UNREADABLE_CASES)
def test_uszkodzony_skoroszyt_jest_nieczytelny_przy_kazdym_otwarciu(tmp_path: Path, damage: str, opening: str) -> None:
    path = styled_template(tmp_path / "Skoroszyt.xlsx")
    DAMAGES[damage](path)

    with pytest.raises(WorkbookError) as raised:
        OPENINGS[opening](path)

    assert raised.value.problem is Problem.UNREADABLE
    assert raised.value.path == path


def test_skoroszyt_bez_stylow_daje_sie_przeczytac_bez_formatu(tmp_path: Path) -> None:
    path = tmp_path / "Rozliczenie - zbiorcze.xlsx"
    write_source(path, include_unmapped=False)
    without_styles(path)

    assert [row.number for row in read_data_rows(path).rows] == [18, 19, 20]


def test_numer_stylu_poza_komorka_i_wierszem_nie_czyni_wzorca_nieczytelnym(tmp_path: Path) -> None:
    path = styled_template(tmp_path / "Skoroszyt.xlsx")
    rewrite_parts(
        path,
        lambda name, data: (
            data.replace(b"</worksheet>", b'<!-- s="999" --></worksheet>', 1) if name.startswith("xl/worksheets/") else data
        ),
    )

    _check_template(path)
    _plan_template(path)


@pytest.mark.parametrize("opening", [_check_template, _plan_template])
def test_styl_wiersza_spoza_stylow_skoroszytu_czyni_wzorzec_nieczytelnym(tmp_path: Path, opening: Callable[[Path], None]) -> None:
    path = styled_template(tmp_path / "Skoroszyt.xlsx")
    rewrite_parts(
        path,
        lambda name, data: (
            re.sub(rb'<row r="([0-9]+)"', rb'<row r="\1" s="99999" customFormat="1"', data, count=1)
            if name.startswith("xl/worksheets/")
            else data
        ),
    )

    with pytest.raises(WorkbookError) as raised:
        opening(path)

    assert raised.value.problem is Problem.UNREADABLE


@pytest.mark.parametrize("opening", OPENINGS)
def test_brak_pliku_przy_kazdym_otwarciu(tmp_path: Path, opening: str) -> None:
    with pytest.raises(WorkbookError) as raised:
        OPENINGS[opening](tmp_path / "Brak.xlsx")

    assert raised.value.problem is Problem.MISSING


@pytest.mark.parametrize(("exists", "problem"), [(True, Problem.UNREADABLE), (False, Problem.MISSING)], ids=["jest", "brak"])
def test_wzorzec_rozliczenia_pracownika_bez_rozszerzenia_skoroszytu(tmp_path: Path, exists: bool, problem: Problem) -> None:
    path = tmp_path / "Wzorzec rozliczenia pracownika.txt"
    if exists:
        styled_template(tmp_path / "Wzorzec.xlsx").rename(path)

    with pytest.raises(WorkbookError) as raised:
        TemplateFile.worker_settlement(path).check()

    assert raised.value.problem is problem


@pytest.mark.parametrize("opening", OPENINGS)
def test_folder_albo_plik_zamiast_skoroszytu_to_brak_pliku_przy_kazdym_otwarciu(tmp_path: Path, opening: str) -> None:
    (tmp_path / "Folder.xlsx").mkdir()
    (tmp_path / "plik").write_bytes(b"")

    for path in (tmp_path / "Folder.xlsx", tmp_path / "plik" / "Skoroszyt.xlsx"):
        with pytest.raises(WorkbookError) as raised:
            OPENINGS[opening](path)
        assert raised.value.problem is Problem.MISSING, path.name


# Te otwarcia sprawdzają po sobie, że w folderze skoroszytu nie został plik, a w folderze bez dostępu nie da się tego sprawdzić.
CHECK_FOLDER_AFTER = frozenset({"plik-okresu-z-wzorca", "rozliczenie-pracownika"})


@pytest.mark.skipif(sys.platform == "win32", reason="chmod w Windows nie odbiera prawa odczytu")
@pytest.mark.parametrize("opening", sorted(OPENINGS.keys() - CHECK_FOLDER_AFTER))
def test_skoroszyt_w_folderze_bez_dostepu_jest_niedostepny_przy_kazdym_otwarciu(tmp_path: Path, opening: str) -> None:
    folder = tmp_path / "bez dostępu"
    folder.mkdir()
    path = styled_template(folder / "Skoroszyt.xlsx")
    folder.chmod(0)
    try:
        with pytest.raises(WorkbookError) as raised:
            OPENINGS[opening](path)
    finally:
        folder.chmod(0o755)

    assert raised.value.problem is Problem.INACCESSIBLE


@pytest.mark.parametrize("opening", OPENINGS)
def test_plik_otwarty_w_excelu_przy_kazdym_otwarciu(tmp_path: Path, opening: str) -> None:
    path = styled_template(tmp_path / "Skoroszyt.xlsx")
    lock_file = tmp_path / "~$Skoroszyt.xlsx"
    lock_file.touch()

    with pytest.raises(WorkbookError) as raised:
        OPENINGS[opening](path)

    assert raised.value.problem is Problem.OPEN_IN_EXCEL
    assert lock_file.exists()


def move_worker_header(path: Path, cell: str | None) -> None:
    """WYKONAWCA z H17 w komórce ``cell``, a jej nagłówek w H17; bez ``cell`` H17 jest pusta."""

    workbook = load_workbook(path)
    try:
        sheet = active_worksheet(workbook)
        sheet["H17"] = None
        if cell is not None:
            sheet["H17"] = sheet[cell].value
            sheet[cell] = " wykonawca "
        workbook.save(path)
    finally:
        workbook.close()


HEADER_CHECKS: dict[str, Callable[[Path], object]] = {
    "odczyt-wierszy-danych": read_data_rows,
    "wzorzec": _check_template,
    "sprawdzanie-wzorca-rozliczenia-pracownika": _check_worker_settlement_template,
}


@pytest.mark.parametrize("check", HEADER_CHECKS)
@pytest.mark.parametrize(("cell", "worker_column"), [("C17", "C"), (None, None)], ids=["WYKONAWCA-w-C", "bez-WYKONAWCA"])
def test_wykonawca_poza_h17_to_zly_uklad(tmp_path: Path, check: str, cell: str | None, worker_column: str | None) -> None:
    path = tmp_path / "Rozliczenie - zbiorcze.xlsx"
    write_source(path, include_unmapped=False)
    with_target_headers(path)
    move_worker_header(path, cell)

    with pytest.raises(WorkbookError) as raised:
        HEADER_CHECKS[check](path)

    assert raised.value.problem is Problem.LAYOUT
    assert raised.value.fault == MissingHeader("Sheet1", 17, "WYKONAWCA", worker_column)


def test_naglowek_wykonawca_nie_zalezy_od_wielkosci_liter_i_spacji(tmp_path: Path) -> None:
    path = tmp_path / "Rozliczenie - zbiorcze.xlsx"
    write_source(path, include_unmapped=False)
    with_target_headers(path)
    move_worker_header(path, "H17")

    _check_template(path)
    assert len(read_data_rows(path).rows) == 3


def test_wiersze_danych_pliku_zbiorczego_w_kolejnosci_z_wykonawca_i_wartosciami_a_at(tmp_path: Path) -> None:
    path = tmp_path / "Rozliczenie - zbiorcze.xlsx"
    write_source(path, include_unmapped=False)
    workbook = load_workbook(path)
    try:
        sheet = active_worksheet(workbook)
        sheet.insert_rows(19)  # pusty wiersz między Wierszami danych nie jest Wierszem danych
        sheet["L22"] = "poza A:K"  # wiersz bez wartości w A:K nie jest Wierszem danych
        sheet["AT18"] = "uwagi"
        sheet["AU18"] = "poza A:AT"
        workbook.save(path)
    finally:
        workbook.close()
    before = path.read_bytes()

    rows = read_data_rows(path).rows

    assert [(row.number, row.worker) for row in rows] == [
        (18, "adrian.testowy"),
        (20, "dariusz.fikcyjny2"),
        (21, "adrian.testowy"),
    ]
    assert all(len(row.cells) == 46 for row in rows)
    assert (rows[0].value("A"), rows[0].value("B")) == ("POZNAŃ", "syntetyczny adres 1")
    assert rows[0].value("AT") == "uwagi"
    assert path.read_bytes() == before


def _inline(reference: str, text: str) -> str:
    return f'<c r="{reference}" s="13" t="inlineStr"><is><t>{text}</t></is></c>'


def excel_collective_file(directory: Path) -> Path:
    """Plik zbiorczy w kształcie z Excela: tekst wspólny i wbudowany, wspólne formuły, sformatowane puste wiersze.

    Wiersze danych 18, 100 i 150 przedzielają dziesiątki sformatowanych pustych
    wierszy; wspólna formuła w Q100 ma formułę główną w wierszu 30 bez danych w A:K.
    """

    path = write_synthetic_templates(directory) / COLLECTIVE_TEMPLATE
    replace_cells(
        path,
        {
            "A18": '<c r="A18" s="13" t="s"><v>4</v></c>',  # sharedStrings[4] == "Instalacje"
            "B18": _inline("B18", "syntetyczny adres 1"),
            "F18": '<c r="F18" s="15"><v>1001</v></c>',
            "G18": '<c r="G18" s="16"><v>46301.375</v></c>',  # format daty 14
            "H18": _inline("H18", "adrian.testowy"),
            "J18": '<c r="J18" s="15" t="b"><v>1</v></c>',
            "K18": '<c r="K18" s="15" t="str"><v>wynik formuły</v></c>',
            "N18": '<c r="N18" s="26"><v>2</v></c>',
            "O18": '<c r="O18" s="26"><v>2.5</v></c>',
            "P18": '<c r="P18" s="26"><f>N18*2</f><v>4</v></c>',
            "AT18": '<c r="AT18" s="31" t="s"><v>0</v></c>',  # sharedStrings[0] == "Suma"
            "Q30": '<c r="Q30" s="26"><f t="shared" ref="Q30:Q140" si="90">N30+1</f><v>1</v></c>',
            "L60": _inline("L60", "poza A:K"),
            "C100": _inline("C100", "tylko C"),
            "H100": _inline("H100", "beata.fikcyjna"),
            "I100": '<c r="I100" s="13" t="e"><v>#N/A</v></c>',
            "Q100": '<c r="Q100" s="28"><f t="shared" si="90"/><v>1</v></c>',
            "E150": '<c r="E150" s="24"><v>3</v></c>',
        },
    )
    return path


def _values_by_column(row: DataRow) -> dict[str, object]:
    return {get_column_letter(column): cell.value for column, cell in enumerate(row.cells, start=1) if cell.value is not None}


def without_cell_references_in_row_100(path: Path) -> None:
    """Komórki A:P wiersza 100 bez adresu ``r``; openpyxl liczy wtedy kolumny po kolei, jak pozwala format."""

    def change(name: str, data: bytes) -> bytes:
        if name != "xl/worksheets/sheet1.xml":
            return data
        start = data.index(b'<row r="100"')
        end = data.index(b"</row>", start)
        return data[:start] + re.sub(rb'<c r="[A-P]100"', b"<c", data[start:end]) + data[end:]

    rewrite_parts(path, change)


SHEET_SHAPES: dict[str, Callable[[Path], None]] = {
    "jak-z-excela": lambda path: None,
    "komorki-bez-adresu": without_cell_references_in_row_100,
}


@pytest.mark.parametrize("shape", SHEET_SHAPES)
def test_wiersze_danych_po_pustych_wierszach_maja_wartosci_jak_w_excelu(tmp_path: Path, shape: str) -> None:
    path = excel_collective_file(tmp_path / "wzorce")
    SHEET_SHAPES[shape](path)
    zeros = {letter: 0 for letter in ("Q", "R", "S", "T", "U", "V", "X")}

    rows = read_data_rows(path).rows

    assert [(row.number, row.worker) for row in rows] == [(18, "adrian.testowy"), (100, "beata.fikcyjna"), (150, None)]
    assert all(len(row.cells) == 46 for row in rows)
    assert _values_by_column(rows[0]) == {
        "A": "Instalacje",
        "B": "syntetyczny adres 1",
        "F": 1001,
        "G": datetime(2026, 10, 6, 9, 0),
        "H": "adrian.testowy",
        "J": True,
        "K": "wynik formuły",
        "N": 2,
        "O": 2.5,
        "P": "=N18*2",
        **zeros,
        "AT": "Suma",
    }
    assert [type(rows[0].value(column)) for column in ("F", "N", "O")] == [int, int, float]
    assert _values_by_column(rows[1]) == {"C": "tylko C", "H": "beata.fikcyjna", "I": "#N/A", "Q": "=N100+1"}
    assert _values_by_column(rows[2]) == {"E": 3}


def linked_template(path: Path) -> Path:
    write_template(path)
    add_external_link_fixture(path)
    return path


def test_plik_z_wzorca_przez_openpyxl_ma_wiersze_danych_formuly_formatowanie_i_kazdy_link_zewnetrzny(tmp_path: Path) -> None:
    template_path = refused_by_copy(linked_template(tmp_path / "Wzorzec.xlsx"))
    template_before = template_path.read_bytes()
    output_directory = tmp_path / "okres"
    output_directory.mkdir()
    target_path = output_directory / "Rozliczenie - Wykonawca.xlsx"
    rows = [
        ("syntetyczne miasto 1",) + (None,) * 12 + (2,) + (None,) * 31 + ("uwagi 1",),
        ("syntetyczne miasto 2",) + (None,) * 45,
    ]

    TemplateFile.worker_settlement(template_path).create(target_path, [_data_row(values) for values in rows])

    assert not has_xml_comment(target_path)
    assert external_link_parts(target_path) == external_link_parts(template_path)
    assert len(external_link_parts(target_path)) == 2
    saved = load_workbook(target_path)
    try:
        written = active_worksheet(saved)
        assert [
            next(written.iter_rows(min_row=row, max_row=row, max_col=46, values_only=True)) for row in (18, 19)
        ] == rows
        assert written["AU18"].value == "=N18"
        assert written["AU19"].value == "=N19"
        assert written["H17"].value == "WYKONAWCA"
        assert written["A17"].fill.fgColor.rgb == "0000AA55"
        assert saved["Rates"]["B1"].value == 17.5
        assert saved.calculation.fullCalcOnLoad is True
    finally:
        saved.close()
    assert sorted(path.name for path in output_directory.iterdir()) == [target_path.name]
    assert template_path.read_bytes() == template_before


def shared_formula_masters(path: Path) -> list[str]:
    """Wzory wspólnych formuł arkusza ``Sheet1`` (zakres, numer i treść) w kolejności z pliku."""

    return re.findall(r'<f t="shared" ref="[^"]+" si="\d+">[^<]*</f>', zip_part(path, "xl/worksheets/sheet1.xml"))


def calculation_flags(path: Path) -> dict[str, str | None]:
    calculation = ET.fromstring(zip_part(path, "xl/workbook.xml")).find(
        "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}calcPr"
    )
    assert calculation is not None
    return {name: calculation.get(name) for name in ("calcMode", "fullCalcOnLoad", "forceFullCalc")}


def test_rozliczenie_pracownika_z_wzorca_zachowuje_wspolne_formuly_calcchain_i_flagi_przeliczania(tmp_path: Path) -> None:
    template_path = write_synthetic_templates(tmp_path / "wzorce") / WORKER_TEMPLATE
    template_before = template_path.read_bytes()
    output_directory = tmp_path / "okres"
    output_directory.mkdir()
    target_path = output_directory / "Rozliczenie - Wykonawca.xlsx"
    rows: list[ExcelRow] = [
        ("syntetyczne miasto 1",) + (None,) * 12 + (2,) + (None,) * 31 + ("uwagi 1",),
        ("syntetyczne miasto 2",) + (None,) * 45,
    ]

    TemplateFile.worker_settlement(template_path).create(target_path, [_data_row(values) for values in rows])

    assert shared_formula_masters(target_path) == shared_formula_masters(template_path)
    assert len(shared_formula_masters(template_path)) > 4
    assert zip_part(target_path, "xl/calcChain.xml") == zip_part(template_path, "xl/calcChain.xml")
    assert calculation_flags(target_path) == {"calcMode": "auto", "fullCalcOnLoad": "1", "forceFullCalc": "1"}
    assert external_link_parts(target_path) == external_link_parts(template_path)
    saved = load_workbook(target_path)
    try:
        written = active_worksheet(saved)
        assert [
            next(written.iter_rows(min_row=row, max_row=row, max_col=46, values_only=True)) for row in (18, 19)
        ] == rows
        assert written["N20"].value == 0
        assert str(written["AU19"].value).startswith("=N19*")
        assert written["H17"].value == "WYKONAWCA"
    finally:
        saved.close()
    assert sorted(path.name for path in output_directory.iterdir()) == [target_path.name]
    assert template_path.read_bytes() == template_before


def _settlement_folders(
    tmp_path: Path, template_path: Path, rows: list[DataRow], *, copied: bool = True
) -> tuple[Path, Path]:
    """Rozliczenie pracownika z Wzorca (kopią, gdy ``copied``) i to samo z jego bliźniaka drogą openpyxl."""

    created = tmp_path / "rozliczenie"
    expected = tmp_path / "przez-openpyxl"
    created.mkdir()
    expected.mkdir()
    marked_for_copy(template_path)
    twin = refused_by_copy(Path(shutil.copy(template_path, tmp_path / f"bliźniak {template_path.name}")))
    name = "Rozliczenie - Wykonawca.xlsx"
    route = TemplateFile.worker_settlement(template_path).create(created / name, rows)
    assert TemplateFile.worker_settlement(twin).create(expected / name, rows) is Route.OPENPYXL
    assert written_by_copy(created / name) == copied
    assert route is (Route.COPY if copied else Route.OPENPYXL)
    assert not has_xml_comment(expected / name)
    return created, expected


def _row(**columns: Any) -> ExcelRow:
    values: list[Any] = [None] * 46
    for letter, value in columns.items():
        values[_column_index(letter) - 1] = value
    return tuple(values)


def _data_row(values: ExcelRow, **highlights: Highlights) -> DataRow:
    """Wiersz danych z wartościami ``values`` i Wyróżnieniami w podanych kolumnach; pozostałe komórki bez Wyróżnień."""

    return data_row(values, {_column_index(letter): cell for letter, cell in highlights.items()})


def _column_index(letter: str) -> int:
    index = 0
    for character in letter:
        index = index * 26 + ord(character) - 64
    return index


VALUE_ROWS: list[ExcelRow] = [
    _row(
        A='<b>Miasto & "Wieś"</b> \'1\' ]]>',
        B=" spacje na brzegach ",
        C="=N18*2",
        D="=",
        E="#N/A",
        F="",
        G=date(2026, 9, 8),
        H="adrian.testowy",
        I=12345678901234567,
        J=True,
        K=datetime(2026, 9, 8, 7, 30, 15),
        N=None,
        O=1.25,
        P=Decimal("2.5"),
        Q=time(13, 45),
        R=False,
        AT="zażółć gęślą jaźń\nnowa linia",
    ),
    _row(A="tylko miasto"),
    _row(A="wiersz poza Wzorcem", N=3),
]


@pytest.mark.parametrize("template", ["maly", "syntetyczny"])
def test_rozliczenie_pracownika_ma_wartosci_kazdego_typu_jak_plik_z_wzorca(tmp_path: Path, template: str) -> None:
    if template == "maly":
        template_path = linked_template(tmp_path / "Wzorzec.xlsx")
    else:
        template_path = write_synthetic_templates(tmp_path / "wzorce") / WORKER_TEMPLATE

    created, expected = _settlement_folders(tmp_path, template_path, [_data_row(values) for values in VALUE_ROWS])

    assert set(compare_worker_settlements(expected, created).values()) == {0}
    saved = load_workbook(created / "Rozliczenie - Wykonawca.xlsx")
    try:
        written = active_worksheet(saved)
        assert written["A18"].value == '<b>Miasto & "Wieś"</b> \'1\' ]]>'
        assert written["B18"].value == " spacje na brzegach "
        assert written["C18"].value == "=N18*2"
        assert (written["D18"].value, written["E18"].value, written["F18"].value) == ("=", "#N/A", None)
        assert written["G18"].value == datetime(2026, 9, 8)
        assert is_date_format(written["G18"].number_format)
        assert written["K18"].value == datetime(2026, 9, 8, 7, 30, 15)
        assert is_date_format(written["K18"].number_format)
        assert written["Q18"].value == time(13, 45)
        assert (written["I18"].value, written["J18"].value, written["R18"].value) == (1.234567890123457e16, True, False)
        assert (written["O18"].value, written["P18"].value) == (1.25, 2.5)
        assert written["AT18"].value == "zażółć gęślą jaźń\nnowa linia"
        assert (written["N18"].value, written["N19"].value, written["N20"].value) == (None, None, 3)
        assert written["A20"].value == "wiersz poza Wzorcem"
    finally:
        saved.close()


RICH_TEXT_ROWS: list[ExcelRow] = [_row(A=CellRichText("zwykły ", TextBlock(InlineFont(b=True), "pogrubiony")), N=1)]


def test_rozliczenie_pracownika_z_tekstem_sformatowanym_jest_jak_plik_z_wzorca(tmp_path: Path) -> None:
    """Tekst sformatowany nie przychodzi z Pliku zbiorczego (odczyt daje zwykły tekst), więc ``run_settlements()`` go nie sięga."""

    template_path = linked_template(tmp_path / "Wzorzec.xlsx")

    created, expected = _settlement_folders(
        tmp_path, template_path, [_data_row(values) for values in RICH_TEXT_ROWS], copied=False
    )

    assert set(compare_worker_settlements(expected, created).values()) == {0}


YELLOW = Highlights(PatternFill(fill_type="solid", fgColor="FFFFFF00"), Color("FFFF0000"), bold=True, italic=True)
THEME_FILL = Highlights(PatternFill(fill_type="solid", fgColor=Color(theme=5, tint=0.4)), Color(theme=1))
ITALIC = Highlights(color=Color(theme=1), italic=True)
HIGHLIGHTED_ROWS: list[DataRow] = [
    _data_row(_row(A="syntetyczne miasto", G=date(2026, 9, 8), N=2, P=1.5), A=YELLOW, G=THEME_FILL, N=ITALIC, AT=YELLOW),
    _data_row(_row(A="tylko miasto")),
    _data_row(_row(A="wiersz poza Wzorcem", N=3), A=THEME_FILL, N=YELLOW),
]


def _highlights(cell: Any) -> tuple[object, ...]:
    """Wyróżnienia komórki tak, jak widzi je tata: wypełnienie, kolor czcionki, pogrubienie, kursywa."""

    fill = cell.fill
    return (fill.fill_type, fill.fgColor if fill.fill_type else None, cell.font.color, bool(cell.font.b), bool(cell.font.i))


@pytest.mark.parametrize("template", ["maly", "z-powtorzona-czcionka", "syntetyczny"])
def test_rozliczenie_pracownika_ma_wyroznienia_wierszy_danych_i_czcionke_naglowka_wzorca_jak_plik_z_wzorca(
    tmp_path: Path, template: str
) -> None:
    if template == "syntetyczny":
        template_path = write_synthetic_templates(tmp_path / "wzorce") / WORKER_TEMPLATE
    else:
        template_path = linked_template(tmp_path / "Wzorzec.xlsx")
    if template == "z-powtorzona-czcionka":
        renumber_cell_styles(template_path)

    created, expected = _settlement_folders(tmp_path, template_path, HIGHLIGHTED_ROWS)

    assert set(compare_worker_settlements(expected, created).values()) == {0}
    saved = load_workbook(created / "Rozliczenie - Wykonawca.xlsx")
    original = load_workbook(template_path)
    try:
        written, wzorzec = active_worksheet(saved), active_worksheet(original)
        yellow = ("solid", Color("FFFFFF00"), Color("FFFF0000"), True, True)
        assert [_highlights(written[cell]) for cell in ("A18", "AT18", "N20")] == [yellow] * 3
        theme_fill = ("solid", Color(theme=5, tint=0.4), Color(theme=1), False, False)
        assert [_highlights(written[cell]) for cell in ("G18", "A20")] == [theme_fill] * 2
        assert _highlights(written["N18"]) == (None, None, Color(theme=1), False, True)
        assert _highlights(written["A19"]) == (None, None, None, False, False)
        assert (written["G18"].value, is_date_format(written["G18"].number_format)) == (datetime(2026, 9, 8), True)
        # A w kolumnach wyglądu: krój nagłówka A17; N poza nimi: krój komórki Wzorca.
        assert [(written[cell].font.name, written[cell].font.sz) for cell in ("A18", "N18", "A19")] == [
            (wzorzec[cell].font.name, wzorzec[cell].font.sz) for cell in ("A17", "N18", "A17")
        ]
    finally:
        saved.close()
        original.close()


def test_pusta_komorka_bez_wlasnego_stylu_za_wierszami_wzorca_ma_styl_kolumny_jak_plik_z_wzorca(tmp_path: Path) -> None:
    """Komórki C nie ma we Wzorcu za jego Wierszami danych; czarna czcionka to czcionka domyślna Wzorca."""

    template_path = write_synthetic_templates(tmp_path / "wzorce") / WORKER_TEMPLATE
    black = Highlights.of(Font(color=Color("FF000000")), None)
    rows = [_data_row(_row(A="syntetyczne miasto"), C=black)] * (SMALL_DATA_ROWS + 5)

    created, expected = _settlement_folders(tmp_path, template_path, rows)

    assert set(compare_worker_settlements(expected, created).values()) == {0}


def _row_styles(path: Path, row: int, columns: range) -> dict[str, str | None]:
    """Numer stylu (atrybut ``s``) każdej komórki wiersza ``row`` w kolumnach ``columns`` w XML arkusza."""

    found = re.search(rf'<row r="{row}"[^>]*?(?:/>|>(.*?)</row>)', zip_part(path, "xl/worksheets/sheet1.xml"), re.S)
    assert found is not None
    cells = re.finditer(r'<c r="([A-Z]+)\d+"([^>]*?)/?>', found.group(1) or "")
    styles = {cell.group(1): re.search(r'\ss="(\d+)"', cell.group(2)) for cell in cells}
    return {
        letter: None if style is None else style.group(1)
        for letter, style in styles.items()
        if _column_index(letter) in columns
    }


@pytest.mark.parametrize("template", ["maly", "syntetyczny"])
def test_wyroznienia_takie_jak_we_wzorcu_zostawiaja_style_wzorca_poza_kolumnami_wygladu(tmp_path: Path, template: str) -> None:
    if template == "maly":
        template_path = linked_template(tmp_path / "Wzorzec.xlsx")
    else:
        template_path = write_synthetic_templates(tmp_path / "wzorce") / WORKER_TEMPLATE
    workbook = load_workbook(template_path)
    try:
        sheet = active_worksheet(workbook)
        same = {
            get_column_letter(column): Highlights.of(
                cast(Font, copy(sheet.cell(18, column).font)), cast(Fill, copy(sheet.cell(18, column).fill))
            )
            for column in range(1, 47)
        }
    finally:
        workbook.close()
    target = tmp_path / "Rozliczenie - Wykonawca.xlsx"

    TemplateFile.worker_settlement(template_path).create(target, [_data_row(_row(A="syntetyczne miasto", N=2), **same)])

    beyond_look = range(13, 47)  # M:AT
    assert _row_styles(target, 18, beyond_look) == _row_styles(template_path, 18, beyond_look)
    assert _row_styles(template_path, 18, beyond_look)


# Adresy, które aplikacja firmowa może dać w Rejestrze zleceń: znaki specjalne XML, fragment, spacje na brzegach, polskie znaki.
LINKED_CELLS = {
    "F18": 'https://przyklad.invalid/zlecenie?caseId=SYN-1&tw=1&q="a"<b>#sekcja',
    "B19": " https://przyklad.invalid/ścieżka/zażółć gęślą?x='1' ",
    "F20": "https://przyklad.invalid/zlecenie?caseId=SYN-3",
}


def linked_collective_file(directory: Path) -> Path:
    """Plik zbiorczy z Wierszami danych 18–20 i Hiperłączami ``LINKED_CELLS``."""

    path = directory / "Rozliczenie 08_14_09_2026 - zbiorcze.xlsx"
    write_source(path, include_unmapped=False)
    workbook = load_workbook(path)
    try:
        sheet = active_worksheet(workbook)
        for coordinate, address in LINKED_CELLS.items():
            sheet[coordinate].hyperlink = address
        workbook.save(path)
    finally:
        workbook.close()
    return path


@pytest.mark.parametrize("template", ["maly", "syntetyczny"])
def test_rozliczenie_pracownika_ma_hiperlacza_wierszy_danych_jak_plik_z_wzorca(tmp_path: Path, template: str) -> None:
    if template == "maly":
        template_path = linked_template(tmp_path / "Wzorzec.xlsx")
    else:
        template_path = write_synthetic_templates(tmp_path / "wzorce") / WORKER_TEMPLATE
    rows = read_data_rows(linked_collective_file(tmp_path)).rows

    created, expected = _settlement_folders(tmp_path, template_path, rows)

    assert set(compare_worker_settlements(expected, created).values()) == {0}
    assert hyperlinks(created / "Rozliczenie - Wykonawca.xlsx") == LINKED_CELLS


@pytest.mark.parametrize("by_copy", [True, False], ids=["kopia", "openpyxl"])
def test_znak_niedozwolony_w_xml_przerywa_zapis_rozliczenia_pracownika_bez_pliku(tmp_path: Path, by_copy: bool) -> None:
    template_path = linked_template(tmp_path / "Wzorzec.xlsx")
    if not by_copy:
        refused_by_copy(template_path)
    template_before = template_path.read_bytes()
    output_directory = tmp_path / "okres"
    output_directory.mkdir()

    with pytest.raises(WorkbookSaveError) as raised:
        TemplateFile.worker_settlement(template_path).create(
            output_directory / "Rozliczenie.xlsx", [_data_row(_row(A="syntetyczny\x01znak"))]
        )

    assert (raised.value.target, raised.value.problem) == (output_directory / "Rozliczenie.xlsx", SaveProblem.INVALID_VALUE)
    assert list(output_directory.iterdir()) == []
    assert template_path.read_bytes() == template_before


def test_znak_niedozwolony_w_xml_przerywa_zapis_rozliczenia_wewnetrznego_bez_pliku(tmp_path: Path) -> None:
    template_path = write_internal_template(tmp_path / "Wzorzec rozliczenia wewnętrznego.xlsx")
    output_directory = tmp_path / "okres"
    output_directory.mkdir()
    row = register_row({"DATA": "syntetyczny\x01znak"})

    with pytest.raises(WorkbookSaveError) as raised:
        TemplateFile.internal_settlement(template_path).create(
            output_directory / "Rozliczenie wew.xlsx", [row]
        )

    assert (raised.value.target, raised.value.problem) == (output_directory / "Rozliczenie wew.xlsx", SaveProblem.INVALID_VALUE)
    assert list(output_directory.iterdir()) == []


@pytest.mark.parametrize("by_copy", [True, False], ids=["kopia", "openpyxl"])
def test_nieudany_zapis_pliku_z_wzorca_nie_zostawia_pliku(tmp_path: Path, by_copy: bool) -> None:
    template_path = linked_template(tmp_path / "Wzorzec.xlsx")
    if not by_copy:
        refused_by_copy(template_path)
    missing_directory = tmp_path / "brak"

    with pytest.raises(WorkbookSaveError) as raised:
        TemplateFile.worker_settlement(template_path).create(missing_directory / "Plik.xlsx", [])

    assert (raised.value.target, raised.value.problem) == (missing_directory / "Plik.xlsx", SaveProblem.TEMPORARY_FILE)
    assert not missing_directory.exists()


def test_pelny_dysk_przy_zapisie_kopii_wzorca_przerywa_zapis_bez_pliku(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    template_path = linked_template(tmp_path / "Wzorzec.xlsx")
    output_directory = tmp_path / "okres"
    output_directory.mkdir()

    def full_disk(package: Package, out: IO[bytes]) -> None:
        out.write(b"PK")
        raise OSError(28, "wymuszony: brak miejsca na dysku")

    monkeypatch.setattr(Package, "write", full_disk)

    with pytest.raises(WorkbookSaveError) as raised:
        TemplateFile.worker_settlement(template_path).create(output_directory / "Plik.xlsx", [_data_row(_row())])

    assert (raised.value.target, raised.value.problem) == (output_directory / "Plik.xlsx", SaveProblem.CONTENT)
    assert list(output_directory.iterdir()) == []


def test_nieczytelne_linki_zewnetrzne_przerywaja_zapis_bez_zmiany_skoroszytu(tmp_path: Path) -> None:
    source_path = linked_template(tmp_path / "Plik zbiorczy.xlsx")
    target_path = tmp_path / "Uszkodzony.xlsx"
    target_path.write_bytes(b"to nie jest skoroszyt")
    workbook = load_workbook(source_path, keep_links=True)
    try:
        active_worksheet(workbook)["A18"] = "syntetyczne miasto"

        with pytest.raises(ExternalLinksError) as raised:
            save_workbook(workbook, target_path)
    finally:
        workbook.close()

    assert (raised.value.target, raised.value.problem) == (target_path, SaveProblem.EXTERNAL_LINKS_UNREADABLE)
    assert target_path.read_bytes() == b"to nie jest skoroszyt"
    assert sorted(path.name for path in tmp_path.iterdir()) == ["Plik zbiorczy.xlsx", "Uszkodzony.xlsx"]


def test_nieczytelne_odwolania_wzorca_przerywaja_zapis_i_nazywaja_skoroszyt_docelowy(tmp_path: Path) -> None:
    target_path = linked_template(tmp_path / "Plik.xlsx")
    before = target_path.read_bytes()

    with pytest.raises(ExternalLinksError) as raised:
        save_from_template(Workbook(), target_path, b"to nie jest Wzorzec")

    assert (raised.value.target, raised.value.problem) == (target_path, SaveProblem.EXTERNAL_LINKS_UNREADABLE)
    assert target_path.read_bytes() == before
    assert [path.name for path in tmp_path.iterdir()] == ["Plik.xlsx"]


class _FullDisk(Workbook):
    """Skoroszyt, którego zapis pada jak na pełnym dysku."""

    def save(self, filename: object) -> None:
        raise OSError(28, "wymuszony: brak miejsca na dysku")


class _NotAZip(Workbook):
    """Skoroszyt, którego zapis daje plik bez archiwum ZIP, więc odwołań zewnętrznych nie da się w nim przywrócić."""

    def save(self, filename: object) -> None:
        assert isinstance(filename, Path), "write_atomically podaje ścieżkę pliku tymczasowego"
        filename.write_bytes(b"to nie jest skoroszyt")


@pytest.mark.parametrize(
    ("workbook", "replace_refused", "problem", "kind"),
    [
        (Workbook(), True, SaveProblem.REPLACE, WorkbookSaveError),
        (_FullDisk(), False, SaveProblem.CONTENT, WorkbookSaveError),
        (_NotAZip(), False, SaveProblem.EXTERNAL_LINKS_NOT_RESTORED, ExternalLinksError),
    ],
    ids=["podmiana", "tresc", "odwolania"],
)
def test_nieudany_krok_zapisu_nazywa_skoroszyt_i_krok_a_skoroszyt_zostaje_bez_zmian(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    workbook: Workbook,
    replace_refused: bool,
    problem: SaveProblem,
    kind: type[WorkbookSaveError],
) -> None:
    target_path = linked_template(tmp_path / "Plik zbiorczy.xlsx")
    before = target_path.read_bytes()
    if replace_refused:
        lock_file(monkeypatch, target_path.name)

    with pytest.raises(WorkbookSaveError) as raised:
        save_workbook(workbook, target_path)

    assert (type(raised.value), raised.value.target, raised.value.problem) == (kind, target_path, problem)
    assert target_path.read_bytes() == before
    assert [path.name for path in tmp_path.iterdir()] == ["Plik zbiorczy.xlsx"]


def test_tekst_bledu_zapisu_to_krok_i_nazwa_pliku_bez_folderu(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """Zapis wyjątku do logu nie przenosi pełnej ścieżki, a z nią nazwy konta taty (ADR 0004)."""

    target_path = linked_template(tmp_path / "Plik zbiorczy.xlsx")
    lock_file(monkeypatch, target_path.name)

    with pytest.raises(WorkbookSaveError) as raised:
        save_workbook(Workbook(), target_path)

    assert str(raised.value) == "podmiana: Plik zbiorczy.xlsx"


def _register_row() -> RegisterRow:
    """Wiersz Rejestru zleceń: komórka pod każdym nagłówkiem docelowym go nazywa."""

    return register_row({header: f"syntetyczne {header}" for header in TARGET_HEADERS})


@pytest.mark.parametrize("by_copy", [True, False], ids=["kopia", "openpyxl"])
def test_naglowek_docelowy_pliku_zbiorczego_jest_szukany_tylko_w_a_k(tmp_path: Path, by_copy: bool) -> None:
    template_path = write_collective_template(tmp_path / "Wzorzec pliku zbiorczego.xlsx")
    workbook = load_workbook(template_path, keep_links=True)
    active_worksheet(workbook)["AV17"] = "ADRES"
    workbook.save(template_path)
    workbook.close()
    if not by_copy:
        refused_by_copy(template_path)
    target = tmp_path / "Plik zbiorczy.xlsx"

    TemplateFile.collective_file(template_path).create(target, [_register_row()])

    sheet = active_worksheet(load_workbook(target))
    assert (sheet["B18"].value, sheet["AV18"].value) == ("syntetyczne ADRES", None)


def test_naglowek_docelowy_rozliczenia_wewnetrznego_jest_szukany_w_calym_wierszu_1(tmp_path: Path) -> None:
    template_path = tmp_path / "Wzorzec rozliczenia wewnętrznego.xlsx"
    workbook = Workbook()
    for column, header in enumerate(("LP", *INTERNAL_TEMPLATE_HEADERS), start=1):
        active_worksheet(workbook).cell(1, column).value = header
    workbook.save(template_path)
    target = tmp_path / "Rozliczenie wewnętrzne.xlsx"

    TemplateFile.internal_settlement(template_path).create(target, [_register_row()])

    sheet = active_worksheet(load_workbook(target))
    assert [sheet.cell(2, column).value for column in (1, 2, 12)] == [
        None,
        "syntetyczne DATA",
        "syntetyczne CZYNNOŚCI ZREALIZOWANE",
    ]


def test_rozliczenie_wewnetrzne_powstaje_przez_openpyxl_bez_wymuszonego_przeliczania(tmp_path: Path) -> None:
    template_path = write_internal_template(tmp_path / "Wzorzec rozliczenia wewnętrznego.xlsx")
    target = tmp_path / "Rozliczenie wewnętrzne.xlsx"

    route = TemplateFile.internal_settlement(template_path).create(target, [_register_row()])

    assert route is Route.OPENPYXL
    assert not load_workbook(target).calculation.forceFullCalc


def test_plik_zbiorczy_z_kopii_ma_wymiar_arkusza_z_wpisanymi_iloscami_uslug(tmp_path: Path) -> None:
    template_path = write_collective_template(tmp_path / "Wzorzec pliku zbiorczego.xlsx")
    rewrite_parts(
        template_path,
        lambda name, data: re.sub(rb'<dimension ref="[^"]*"', b'<dimension ref="A1:K17"', data)
        if name == "xl/worksheets/sheet1.xml"
        else data,
    )
    target = tmp_path / "Plik zbiorczy.xlsx"
    row = replace(_register_row(), quantities={QuantityColumn.AP: 1})

    route = TemplateFile.collective_file(template_path).create(target, [row])

    assert route is Route.COPY
    assert re.search(r'<dimension ref="A1:AP18"', zip_part(target, "xl/worksheets/sheet1.xml"))


@pytest.mark.parametrize("copied", [True, False], ids=["kopia", "openpyxl"])
def test_wzorzec_rozliczenia_pracownika_z_wartoscia_w_wierszu_danych_nie_daje_pliku_na_zadnej_drodze(
    tmp_path: Path, copied: bool
) -> None:
    template_path = tmp_path / "Wzorzec rozliczenia pracownika.xlsx"
    write_template(template_path)
    workbook = load_workbook(template_path)
    try:
        active_worksheet(workbook)["A19"] = "wpis taty"
        workbook.save(template_path)
    finally:
        workbook.close()
    (marked_for_copy if copied else refused_by_copy)(template_path)
    target = tmp_path / "Rozliczenie - Wykonawca.xlsx"

    with pytest.raises(WorkbookError) as raised:
        TemplateFile.worker_settlement(template_path).create(target, [])

    assert raised.value.problem is Problem.HAS_DATA
    assert not target.exists()


def _worker_template_with_merge(tmp_path: Path, merged: str, *, copied: bool) -> Path:
    template_path = tmp_path / "Wzorzec rozliczenia pracownika.xlsx"
    write_template(template_path)
    workbook = load_workbook(template_path)
    try:
        active_worksheet(workbook).merge_cells(merged)
        workbook.save(template_path)
    finally:
        workbook.close()
    return (marked_for_copy if copied else refused_by_copy)(template_path)


# Zapis Wzorca rozliczenia pracownika obejmuje kolumny A:AT Wierszy danych zapisanych wierszy (18 i 19).
@pytest.mark.parametrize("merged", ["A20:B20", "AU19:AV19"], ids=["wiersz-pod-zapisanymi", "kolumna-za-zapisem"])
@pytest.mark.parametrize("copied", [True, False], ids=["kopia", "openpyxl"])
def test_scalenie_wzorca_rozliczenia_pracownika_poza_komorkami_zapisu_nie_przeszkadza(
    tmp_path: Path, copied: bool, merged: str
) -> None:
    template_path = _worker_template_with_merge(tmp_path, merged, copied=copied)
    target = tmp_path / "Rozliczenie - Wykonawca.xlsx"

    TemplateFile.worker_settlement(template_path).create(target, [_data_row(_row(A="syntetyczne miasto"))] * 2)

    assert target.exists()


@pytest.mark.parametrize("copied", [True, False], ids=["kopia", "openpyxl"])
def test_wykonawca_bez_wierszy_danych_dostaje_plik_ze_scalonego_wzorca_rozliczenia_pracownika(
    tmp_path: Path, copied: bool
) -> None:
    template_path = _worker_template_with_merge(tmp_path, "A18:B18", copied=copied)
    target = tmp_path / "Rozliczenie - Wykonawca.xlsx"

    TemplateFile.worker_settlement(template_path).create(target, [])

    assert target.exists()


@pytest.mark.parametrize(
    "merged",
    ["A19:B19", "N19:O19", "AT19:AU19", "A18:A19", "N18:N19", "A17:A18"],
    ids=["kolumny-opisu", "kolumny-ilosci", "ostatnia-kolumna-zapisu", "dwa-wiersze-danych", "dwa-wiersze-ilosci", "z-wiersza-naglowka"],
)
@pytest.mark.parametrize("copied", [True, False], ids=["kopia", "openpyxl"])
def test_wzorzec_rozliczenia_pracownika_ze_scaleniem_w_wierszu_danych_nie_daje_pliku_na_zadnej_drodze(
    tmp_path: Path, copied: bool, merged: str
) -> None:
    template_path = _worker_template_with_merge(tmp_path, merged, copied=copied)
    target = tmp_path / "Rozliczenie - Wykonawca.xlsx"

    with pytest.raises(WorkbookError) as raised:
        TemplateFile.worker_settlement(template_path).create(target, [_data_row(_row(A="syntetyczne miasto"))] * 2)

    assert raised.value.problem is Problem.LAYOUT
    assert not target.exists()
