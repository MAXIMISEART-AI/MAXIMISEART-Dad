"""Syntetyczne skoroszyty dla Przygotowania pliku zbiorczego.

Rejestr zleceń, Wzorzec rozliczenia wewnętrznego i Wzorzec pliku zbiorczego są
budowane w katalogu tymczasowym. Nie zawierają prawdziwych danych klientów.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import PatternFill

from tests.test_settlement_engine import active_worksheet, add_external_link_fixture, write_mapping


PERIOD = "08_14_09_2026"
INTERNAL_SETTLEMENT_NAME = f"Rozliczenie wew. {PERIOD}.xlsx"
COLLECTIVE_FILE_NAME = f"Rozliczenie {PERIOD} - zbiorcze.xlsx"

# Kolejność kolumn Rejestru celowo różni się od wzorców; "STATUS APLIKACJI" nie jest przenoszony.
REGISTER_HEADERS: tuple[str, ...] = (
    "nr zlecenia",
    "DATA GODZ. POCZ.",
    "WYKONAWCA",
    "MIASTO",
    "adres",
    "STATUS APLIKACJI",
    "UMOWA",
    "URZĄDZENIA",
    "CZYNNOŚCI",
    "SKUTECZNE",
    "CZYNNOŚCI ZREALIZOWANE",
    "NUMER UMOWY/ZLECENIA ORANGE",
)
INTERNAL_TEMPLATE_HEADERS: tuple[str, ...] = (
    "DATA",
    "MIASTO",
    "ADRES",
    "UMOWA",
    "URZĄDZENIA",
    "CZYNNOŚCI",
    "NUMER ZL.",
    "WYKONAWCA",
    "NUMER UMOWY/ZLECENIA ORANGE",
    "SKUTECZNE",
    "CZYNNOŚCI ZREALIZOWANE",
    "zwrot",
    "UWAGI",
)
# Nagłówki A:K w wierszu 17 Wzorca pliku zbiorczego; WYKONAWCA w H jak w Przebiegu rozliczeń.
COLLECTIVE_TEMPLATE_HEADERS: tuple[str, ...] = (
    "MIASTO",
    "ADRES",
    "UMOWA",
    "URZĄDZENIA",
    "CZYNNOŚCI",
    "NUMER ZL.",
    "DATA",
    "WYKONAWCA",
    "NUMER UMOWY/ZLECENIA ORANGE",
    "SKUTECZNE",
    "CZYNNOŚCI ZREALIZOWANE",
)
# Wiersze Wzorca pliku zbiorczego przygotowane zerami w N:AP.
COLLECTIVE_TEMPLATE_ZERO_ROWS = 6

RegisterRow = dict[str, object] | None
"""Wiersz Rejestru: wartości po nagłówkach Rejestru albo ``None`` dla całkowicie pustego wiersza."""


def register_order(
    number: int,
    started: datetime,
    *,
    successful: str = "tak",
    worker: str = "adrian.maciejewski",
) -> dict[str, object]:
    """Syntetyczne zlecenie z unikalnymi znacznikami treści ``SYN-…``."""

    return {
        "nr zlecenia": f"ZL-SYN-{number:03d}",
        "DATA GODZ. POCZ.": started,
        "WYKONAWCA": worker,
        "MIASTO": "POZNAŃ",
        "adres": f"syntetyczny adres SYN-{number:03d}",
        "STATUS APLIKACJI": "zamknięte",
        "UMOWA": f"UM-SYN-{number:03d}",
        "URZĄDZENIA": "dekoder",
        "CZYNNOŚCI": "instalacja",
        "SKUTECZNE": successful,
        "CZYNNOŚCI ZREALIZOWANE": "montaż",
        "NUMER UMOWY/ZLECENIA ORANGE": f"OR-SYN-{number:03d}",
    }


def default_register_rows() -> list[RegisterRow]:
    """Rejestr z okresem 08_14_09_2026 pośród innych okresów, pustymi wierszami i spóźnionym zleceniem."""

    return [
        register_order(1, datetime(2026, 9, 7, 15, 0)),  # poprzedni okres (poniedziałek)
        None,
        register_order(2, datetime(2026, 9, 8, 0, 0), worker="adrian.maciejewski"),
        None,
        register_order(3, datetime(2026, 9, 9, 10, 30), successful="NIE", worker="dariusz.nowak2"),
        None,
        None,
        register_order(4, datetime(2026, 9, 12, 12, 0), successful=" Tak ", worker="kamil.frontczak"),
        None,
        register_order(5, datetime(2026, 9, 15, 0, 0)),  # następny okres (wtorek 00:00)
        None,
        register_order(6, datetime(2026, 9, 14, 18, 30), successful="TAK", worker="dariusz.nowak2"),
    ]


def write_register(
    path: Path,
    rows: list[RegisterRow] | None = None,
    *,
    headers: tuple[str, ...] = REGISTER_HEADERS,
) -> Path:
    """Zapisuje Rejestr zleceń: arkusz-wabik ``lista`` i arkusz danych z nagłówkami w wierszu 1."""

    workbook = Workbook()
    decoy = active_worksheet(workbook)
    decoy.title = "lista"
    decoy["A1"] = "WYKONAWCY"
    decoy["A2"] = "adrian.maciejewski"
    sheet = workbook.create_sheet("ZLECENIA")
    for column, header in enumerate(headers, start=1):
        sheet.cell(1, column).value = header
    for row_number, row in enumerate(rows if rows is not None else default_register_rows(), start=2):
        if row is None:
            continue
        for column, header in enumerate(headers, start=1):
            if header in row:
                sheet.cell(row_number, column).value = row[header]
    workbook.save(path)
    return path


def write_internal_template(path: Path) -> Path:
    """Wzorzec rozliczenia wewnętrznego: nagłówki w wierszu 1, bez zleceń."""

    workbook = Workbook()
    sheet = active_worksheet(workbook)
    sheet.title = "Arkusz1"
    for column, header in enumerate(INTERNAL_TEMPLATE_HEADERS, start=1):
        sheet.cell(1, column).value = header
    workbook.save(path)
    return path


def write_collective_template(path: Path) -> Path:
    """Wzorzec pliku zbiorczego: stawki w wierszach 1–16, nagłówki w 17, zera i formuły, link zewnętrzny."""

    workbook = Workbook()
    sheet = active_worksheet(workbook)
    sheet.title = "Sheet1"
    sheet["A1"] = "STAWKI"
    sheet["N1"] = 17.5
    sheet["O1"] = 30
    sheet["N16"] = "=SUM(N18:N40)"
    sheet["A17"].fill = PatternFill(fill_type="solid", fgColor="00AA55")
    for column, header in enumerate(COLLECTIVE_TEMPLATE_HEADERS, start=1):
        sheet.cell(17, column).value = header
    sheet["N17"] = "instalacja"
    sheet["AP17"] = "demontaż"
    for row in range(18, 18 + COLLECTIVE_TEMPLATE_ZERO_ROWS):
        for column in range(14, 43):  # N:AP
            sheet.cell(row, column).value = 0
        sheet.cell(row, 47).value = f"=N{row}*$N$1"  # AU
    rates = workbook.create_sheet("Rates")
    rates["A1"] = "synthetic rate"
    rates["B1"] = 17.5
    workbook.save(path)
    add_external_link_fixture(path)
    return path


@dataclass(frozen=True)
class PreparationFixture:
    """Ścieżki jednego syntetycznego uruchomienia Przygotowania pliku zbiorczego."""

    register_path: Path
    periods_directory: Path
    config_path: Path
    internal_template_path: Path
    collective_template_path: Path

    @property
    def period_directory(self) -> Path:
        return self.periods_directory / PERIOD

    @property
    def internal_settlement_path(self) -> Path:
        return self.period_directory / INTERNAL_SETTLEMENT_NAME

    @property
    def collective_file_path(self) -> Path:
        return self.period_directory / COLLECTIVE_FILE_NAME


def make_preparation_fixture(tmp_path: Path, rows: list[RegisterRow] | None = None) -> PreparationFixture:
    """Rejestr, folder okresów ``INEA``, konfiguracja wykonawców i oba wzorce w ``tmp_path``."""

    config_directory = tmp_path / "config"
    config_directory.mkdir()
    periods_directory = tmp_path / "INEA"
    periods_directory.mkdir()
    return PreparationFixture(
        register_path=write_register(tmp_path / "ZLECENIA INTERMAX OD 08_11_2021.xlsx", rows),
        periods_directory=periods_directory,
        config_path=write_mapping(config_directory / "worker_mapping.yaml"),
        internal_template_path=write_internal_template(
            config_directory / "Wzorzec rozliczenia wewnętrznego.xlsx"
        ),
        collective_template_path=write_collective_template(
            config_directory / "Wzorzec pliku zbiorczego.xlsx"
        ),
    )
