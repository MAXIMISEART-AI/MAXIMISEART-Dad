"""Plik zbiorczy z kopii Wzorca (ADR 0008) jest taki sam jak zapisany przez openpyxl.

Test właściciel drogi kopii w Przygotowaniu pliku zbiorczego: obie drogi zapisu dają te same
wartości, Wyróżnienia, formaty liczb, Hiperłącza, scalenia, autofiltr i odwołania zewnętrzne.
"""

from __future__ import annotations

from collections.abc import Callable
from copy import copy
from dataclasses import replace
from pathlib import Path
import re
import zipfile

from openpyxl import load_workbook
import pytest

from rozliczenia import preparation_files

from tests.fixtures.benchmark import COLLECTIVE_TEMPLATE, INTERNAL_TEMPLATE, write_synthetic_templates
from tests.fixtures.preparation import (
    PreparationFixture,
    link_register_order_numbers,
    make_dad_layout_fixture,
    make_preparation_fixture,
    run_preparation,
    write_collective_template,
)
from tests.fixtures.workbooks import active_worksheet, external_link_parts


def _synthetic(tmp_path: Path) -> PreparationFixture:
    templates = write_synthetic_templates(tmp_path / "wzorce")
    (tmp_path / "syntetyczne").mkdir()
    return replace(
        make_preparation_fixture(tmp_path / "syntetyczne"),
        internal_template_path=templates / INTERNAL_TEMPLATE,
        collective_template_path=templates / COLLECTIVE_TEMPLATE,
    )


def _linked(tmp_path: Path) -> PreparationFixture:
    fixture = make_preparation_fixture(tmp_path)
    link_register_order_numbers(fixture.register_path, "https://przyklad.invalid/zlecenie?caseId=SYN-1&tw=1")
    return fixture


def _row_style(tmp_path: Path) -> PreparationFixture:
    fixture = make_preparation_fixture(tmp_path)
    write_collective_template(fixture.collective_template_path, row_style=True)
    return fixture


def _filter_database(path: Path) -> list[str]:
    """Zakres nazwy ``_xlnm._FilterDatabase``: openpyxl jej nie wczytuje, Excel filtruje według niej."""

    with zipfile.ZipFile(path) as archive:
        workbook_xml = archive.read("xl/workbook.xml").decode("utf-8")
    found = re.findall(r'<(?:\w+:)?definedName\b[^>]*name="_xlnm\._FilterDatabase"[^>]*>([^<]*)<', workbook_xml)
    return [name.replace("'", "") for name in found]


def _filter_name_ref_error(tmp_path: Path) -> PreparationFixture:
    """Wzorzec, w którym nazwa ``_xlnm._FilterDatabase`` wskazuje ``#REF!``, jak po usunięciu zakresu w Excelu."""

    fixture = make_preparation_fixture(tmp_path)
    path = fixture.collective_template_path
    with zipfile.ZipFile(path) as source:
        parts = {info.filename: source.read(info) for info in source.infolist()}
    parts["xl/workbook.xml"] = re.sub(
        rb"('Sheet1'!)\$B\$17:\$C\$57", rb"\1#REF!", parts["xl/workbook.xml"]
    )
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
        for name, data in parts.items():
            target.writestr(name, data)
    return fixture


def _cells(path: Path) -> dict[str, object]:
    workbook = load_workbook(path)
    try:
        sheet = active_worksheet(workbook)
        cells = {
            cell.coordinate: (
                cell.value,
                cell.number_format,
                (cell.font.name, cell.font.sz, bool(cell.font.b), bool(cell.font.i), copy(cell.font.color)),
                copy(cell.fill),
                copy(cell.border),
                copy(cell.alignment),
                copy(cell.protection),
                cell.hyperlink.target if cell.hyperlink is not None else None,
            )
            for row in sheet.iter_rows()
            for cell in row
        }
        names = {name: str(value.attr_text) for name, value in sheet.defined_names.items()}
        return {
            **cells,
            "autofiltr": sheet.auto_filter.ref,
            "scalenia": sorted(str(merged) for merged in sheet.merged_cells.ranges),
            "nazwy arkusza": names,
            "_FilterDatabase": _filter_database(path),
            "przeliczenie": workbook.calculation.fullCalcOnLoad,
        }
    finally:
        workbook.close()


@pytest.mark.parametrize(
    ("make_fixture", "copied"),
    [
        pytest.param(make_preparation_fixture, True, id="syntetyczny"),
        pytest.param(_linked, True, id="hiperlacza"),
        pytest.param(make_dad_layout_fixture, True, id="uklad-taty"),
        pytest.param(_synthetic, True, id="wzorce-ksztaltu-taty"),
        pytest.param(_row_style, False, id="styl-wiersza-przez-openpyxl"),
        pytest.param(_filter_name_ref_error, False, id="nazwa-filtra-ref-przez-openpyxl"),
    ],
)
def test_plik_zbiorczy_z_kopii_wzorca_jest_taki_jak_z_openpyxl(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, make_fixture: Callable[[Path], PreparationFixture], copied: bool
) -> None:
    (tmp_path / "kopia").mkdir()
    (tmp_path / "openpyxl").mkdir()
    by_copy = make_fixture(tmp_path / "kopia")
    by_openpyxl = make_fixture(tmp_path / "openpyxl")
    run_preparation(by_copy)
    with monkeypatch.context() as patch:
        patch.setattr(preparation_files, "_written_as_copy", lambda *_: False)
        run_preparation(by_openpyxl)

    created, reference = by_copy.collective_file_path, by_openpyxl.collective_file_path
    with zipfile.ZipFile(created) as archive:
        keeps_template_calc_chain = "xl/calcChain.xml" in archive.namelist()
    with zipfile.ZipFile(by_copy.collective_template_path) as archive:
        template_has_calc_chain = "xl/calcChain.xml" in archive.namelist()

    assert _cells(created) == _cells(reference)
    assert external_link_parts(created) == external_link_parts(reference)
    assert keeps_template_calc_chain == (copied and template_has_calc_chain)
