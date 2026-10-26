"""Przebieg rozliczeń w testach: syntetyczne pliki i obserwator postępu.

Plik zbiorczy, Wzorzec rozliczenia pracownika i konfiguracja wykonawców są
budowane w katalogu tymczasowym. Nie zawierają prawdziwych danych klientów.
"""

from __future__ import annotations

from collections.abc import Mapping, Sequence
from datetime import date, datetime
from pathlib import Path

import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

from rozliczenia.domain import ProgressEvent, ProgressObserver
from rozliczenia.progress import ProgressProjection
from tests.fixtures.workbooks import TEMPLATE_VIEW, View, active_worksheet, set_view

PERIOD = "08_14_09_2026"


NOT_A_PERIOD = "09_15_09_2026"  # wygląda jak token okresu, ale 9 września 2026 to środa


SOURCE_NAME = f"Rozliczenie {PERIOD} - zbiorcze.xlsx"


def write_mapping(path: Path) -> Path:
    path.write_text(
        yaml.safe_dump(
            {
                "schema_version": 1,
                "workers": {
                    "adrian.testowy": "Adrian Testowy",
                    "dariusz.fikcyjny2": "Darek Fikcyjny",
                    "kamil.syntetyczny": "Kamil Syntetyczny",
                },
            },
            allow_unicode=True,
            sort_keys=False,
        ),
        encoding="utf-8",
    )
    return path


# Szerokości kolumn Wzorca rozliczenia pracownika: różne, z bardzo wąskimi C, E, L i M.
WORKER_TEMPLATE_COLUMN_WIDTHS: dict[int, float] = {
    1: 12.0,
    2: 30.5,
    3: 2.29,
    4: 14.0,
    5: 3.0,
    6: 20.33,
    7: 16.66,
    8: 18.0,
    9: 23.33,
    10: 11.0,
    11: 15.5,
    12: 2.0,
    13: 2.0,
    14: 6.5,
}


def write_template(path: Path, *, view: View = TEMPLATE_VIEW) -> None:
    workbook = Workbook()
    sheet = active_worksheet(workbook)
    sheet.title = "Sheet1"
    set_view(sheet, view)
    for column, width in WORKER_TEMPLATE_COLUMN_WIDTHS.items():
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.cell(17, 8).value = "WYKONAWCA"
    sheet["A17"].fill = PatternFill(fill_type="solid", fgColor="00AA55")
    sheet.cell(18, 47).value = "=N18"
    sheet.cell(19, 47).value = "=N19"
    sheet.cell(18, 14).value = 0
    sheet.cell(19, 14).value = 0
    rates = workbook.create_sheet("Rates")
    rates["A1"] = "synthetic rate"
    rates["B1"] = 17.5
    workbook.save(path)


def write_source(path: Path, *, include_unmapped: bool = True) -> None:
    workbook = Workbook()
    sheet = active_worksheet(workbook)
    sheet.title = "Sheet1"
    sheet.cell(17, 8).value = "WYKONAWCA"
    rows = [
        ("POZNAŃ", "syntetyczny adres 1", "adrian.testowy", "#1"),
        ("POZNAŃ", "syntetyczny adres 2", "dariusz.fikcyjny2", "#2"),
        ("POZNAŃ", "syntetyczny adres 3", "adrian.testowy", "#3"),
    ]
    if include_unmapped:
        rows.append(("POZNAŃ", "syntetyczny adres 4", "andrzej.nieznany2", "#4"))
    for row_number, (city, address, worker, order_number) in enumerate(rows, start=18):
        sheet.cell(row_number, 1).value = city
        sheet.cell(row_number, 2).value = address
        sheet.cell(row_number, 6).value = order_number
        sheet.cell(row_number, 8).value = worker
        sheet.cell(row_number, 14).value = row_number - 17
    workbook.save(path)


CellValue = str | int | float | datetime | date | None

# Nagłówki A17:K17 Wzorca pliku zbiorczego taty, pod którymi Przygotowanie zapisuje kolumny wejścia.
COLLECTIVE_HEADERS = (
    "MIASTO",
    "ADRES",
    "UMOWA",
    "URZĄDZENIA",
    "CZYNNOŚCI",
    "NUMER ZL.",
    "DATA",
    "WYKONAWCA",
    "NUMER UMOWY/ZLECENIA ORANGE",
    "SKUTECZNE",
    "CZYNNOŚCI ZREALIZOWANE",
)


# Wiersz danych ze znacznikiem w każdej kolumnie z treścią zlecenia (MIASTO, ADRES, UMOWA, URZĄDZENIA, CZYNNOŚCI,
# NUMER ZL., NUMER UMOWY, CZYNNOŚCI ZREALIZOWANE, AQ:AT i ilość w V), w środę 09.09.2026 o 13:45 i z Wykonawcą
# ``adrian.testowy``, który Konfiguracja wykonawców mapuje na „Adrian Testowy”.
MARKED: dict[str, CellValue] = {
    "A": "ZNACZNIK-MIASTO",
    "B": "ZNACZNIK-ADRES",
    "C": "u ZNACZNIK-UMOWA",
    "D": "nokia ZNACZNIK-URZADZENIA",
    "E": "ZNACZNIK-CZYNNOSCI",
    "F": "ZNACZNIK-NUMER-ZL",
    "G": datetime(2026, 9, 9, 13, 45),
    "I": "ZNACZNIK-NUMER-UMOWY",
    "J": "tak",
    "K": "ZNACZNIK-ZREALIZOWANE",
    "AQ": "ZNACZNIK-AQ",
    "AR": "ZNACZNIK-AR",
    "AS": "ZNACZNIK-AS",
    "AT": "ZNACZNIK-AT",
    "V": "ZNACZNIK-ILOSC",
}


def write_quantity_source(path: Path, rows: Sequence[Mapping[str, CellValue]]) -> None:
    """Plik zbiorczy z nagłówkami A17:K17 Wzorca i Wierszami danych od 18: litera kolumny → wartość.

    Bez podanej wartości wiersz ma syntetyczne MIASTO, ADRES, NUMER ZL. i SKUTECZNE oraz Wykonawcę
    ``adrian.testowy``, więc Przebieg z zapisem kończy się wynikiem.
    """

    workbook = Workbook()
    sheet = active_worksheet(workbook)
    sheet.title = "Sheet1"
    for column, header in enumerate(COLLECTIVE_HEADERS, start=1):
        sheet.cell(17, column).value = header
    for number, cells in enumerate(rows, start=18):
        defaults: dict[str, CellValue] = {"A": "POZNAŃ", "B": f"syntetyczny adres {number}", "F": f"#{number}", "H": "adrian.testowy", "J": "tak"}
        for letter, value in {**defaults, **cells}.items():
            sheet[f"{letter}{number}"] = value
    workbook.save(path)


def link_order_numbers(path: Path, address: str) -> dict[int, str]:
    """Hiperłącze w ``NUMER ZL.`` (F) każdego Wiersza danych Pliku zbiorczego, jak przy zleceniu z aplikacji firmowej.

    Adres to ``address`` z numerem wiersza na końcu; zwraca numer wiersza → adres.
    """

    workbook = load_workbook(path)
    try:
        sheet = active_worksheet(workbook)
        links = {row: f"{address}{row}" for row in range(18, sheet.max_row + 1) if sheet.cell(row, 6).value is not None}
        for row, link in links.items():
            sheet.cell(row, 6).hyperlink = link
        workbook.save(path)
    finally:
        workbook.close()
    return links


def make_fixture(tmp_path: Path, period: str = PERIOD) -> tuple[Path, Path, Path]:
    period_directory = tmp_path / period
    period_directory.mkdir(parents=True)
    target_directory = period_directory / f"Rozliczenie pracowników {period}"
    source_path = period_directory / f"Rozliczenie {period} - zbiorcze.xlsx"
    write_source(source_path, include_unmapped=False)
    config_path = write_mapping(tmp_path / "worker_mapping.yaml")
    write_template(tmp_path / "Wzorzec rozliczenia pracownika.xlsx")
    return source_path, target_directory, config_path


def projected(projection: ProgressProjection) -> ProgressObserver:
    def observer(event: ProgressEvent) -> None:
        projection.update(event)

    return observer


def make_preview_fixture(tmp_path: Path) -> tuple[Path, Path, Path]:
    source_path, target_directory, config_path = make_fixture(tmp_path)
    write_source(source_path, include_unmapped=True)
    return source_path, target_directory, config_path
