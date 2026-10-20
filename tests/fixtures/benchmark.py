"""Benchmark Przebiegu rozliczeń w testach: syntetyczne Wzorce i CLI narzędzia jako osobny proces.

Narzędzie leży w ``tools/``, poza pakietem, więc testy uruchamiają je tak jak Maks:
osobnym procesem Pythona z tego checkoutu.
"""

from __future__ import annotations

import os
import subprocess
import sys
import zipfile
from pathlib import Path

from openpyxl import load_workbook

from tests.fixtures.preparation import write_collective_template, write_internal_template
from tests.fixtures.settlement import write_template
from tests.fixtures.workbooks import active_worksheet

ROOT = Path(__file__).resolve().parents[2]
BENCHMARK = ROOT / "tools" / "benchmark_przebiegu.py"
SYNTHETIC_TEMPLATES = ROOT / "tools" / "syntetyczne_wzorce.py"

WORKER_TEMPLATE = "Wzorzec rozliczenia pracownika.xlsx"
COLLECTIVE_TEMPLATE = "Wzorzec pliku zbiorczego.xlsx"
INTERNAL_TEMPLATE = "Wzorzec rozliczenia wewnętrznego.xlsx"
# Mała odmiana syntetycznych Wzorców: kilka grup wspólnych formuł (co 64 wiersze), szybka w pytest.
SMALL_DATA_ROWS = 150

# Pliki repozytorium, których benchmark nie może zmienić: Ustawienia taty i historie metryk procesów.
PROTECTED_REPOSITORY_FILES = (
    ROOT / "config" / "centrum-ustawienia.yaml",
    ROOT / ".rozliczenia-metrics.jsonl",
    ROOT / ".przygotowanie-pliku-zbiorczego-metrics.jsonl",
)

# Kolumny Pliku zbiorczego, które identyfikują klienta albo zlecenie: ADRES, UMOWA, NUMER ZL., NUMER UMOWY.
_IDENTIFYING_COLUMNS = (2, 3, 6, 9)


def write_templates(directory: Path) -> Path:
    """Folder z trzema syntetycznymi Wzorcami pod nazwami, pod którymi leżą w ``config``."""

    directory.mkdir(parents=True)
    write_template(directory / "Wzorzec rozliczenia pracownika.xlsx")
    write_internal_template(directory / "Wzorzec rozliczenia wewnętrznego.xlsx")
    write_collective_template(directory / "Wzorzec pliku zbiorczego.xlsx")
    return directory


def write_synthetic_templates(directory: Path, data_rows: int | None = SMALL_DATA_ROWS) -> Path:
    """Folder z syntetycznymi Wzorcami kształtu Wzorców taty, z generatora w ``tools/``.

    ``data_rows=None`` to pełny rozmiar, jak u taty (domyślny w generatorze).
    """

    size = [] if data_rows is None else ["--data-rows", str(data_rows)]
    generated = _run_tool(SYNTHETIC_TEMPLATES, str(directory), *size)
    assert generated.returncode == 0, generated.stdout + generated.stderr
    return directory


def zip_part(path: Path, name: str) -> str:
    with zipfile.ZipFile(path) as archive:
        return archive.read(name).decode("utf-8")


def zip_part_names(path: Path) -> set[str]:
    with zipfile.ZipFile(path) as archive:
        return set(archive.namelist())


def run_benchmark(*arguments: str) -> subprocess.CompletedProcess[str]:
    return _run_tool(BENCHMARK, *arguments)


def _run_tool(tool: Path, *arguments: str) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [sys.executable, str(tool), *arguments],
        stdin=subprocess.DEVNULL,
        capture_output=True,
        text=True,
        encoding="utf-8",
        env={**os.environ, "PYTHONIOENCODING": "utf-8"},
        timeout=100,
        check=False,
    )


def protected_repository_files() -> dict[Path, bytes | None]:
    return {path: path.read_bytes() if path.is_file() else None for path in PROTECTED_REPOSITORY_FILES}


def identifying_values(collective_file: Path) -> set[str]:
    """Adresy, numery umów i numery zleceń z wierszy danych Pliku zbiorczego lustra."""

    workbook = load_workbook(collective_file)
    try:
        sheet = active_worksheet(workbook)
        return {
            str(row[column - 1])
            for row in sheet.iter_rows(min_row=18, max_col=11, values_only=True)
            for column in _IDENTIFYING_COLUMNS
            if row[column - 1] is not None
        }
    finally:
        workbook.close()


def data_row_entries(collective_file: Path) -> list[tuple[object, ...]]:
    """Wpisy taty w każdym Wierszu danych Pliku zbiorczego lustra: ilości ``N:AP`` i obciążenia z uwagami ``AQ:AT``."""

    workbook = load_workbook(collective_file)
    try:
        sheet = active_worksheet(workbook)
        return [
            row[13:46]
            for row in sheet.iter_rows(min_row=18, max_col=46, values_only=True)
            if any(value is not None for value in row[:11])
        ]
    finally:
        workbook.close()


def frozen_reference(reference_directory: Path) -> Path:
    """Jedyny zamrożony Folder rozliczeń pracowników pod ``reference_directory``."""

    (folder,) = (path for path in reference_directory.rglob("Rozliczenie pracowników *") if path.is_dir())
    return folder
