"""Syntetyczne skoroszyty: aktywny arkusz, odwołania zewnętrzne i celowo zepsute pliki."""

from __future__ import annotations

import os
import re
import tempfile
import xml.etree.ElementTree as ET
import zipfile
from collections.abc import Callable
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font
from openpyxl.worksheet.views import SheetView
from openpyxl.worksheet.worksheet import Worksheet

# Resztki formatowania w Wierszach danych Wzorców taty (wiersze 18–31 i od 32 w dół): inny krój, zawijanie,
# wyrównanie, brak obramowania, nietypowy format daty i własna wysokość wiersza.
LEFTOVER_FONT = Font(name="Arial", size=12)
LEFTOVER_ALIGNMENT = Alignment(horizontal="center", vertical="top", wrap_text=True)
LEFTOVER_DATE_FORMAT = "yyyy-mm-dd h:mm:ss"
LEFTOVER_ROW_HEIGHT = 30.0


def active_worksheet(workbook: Workbook) -> Worksheet:
    sheet = workbook.active
    assert isinstance(sheet, Worksheet)
    return sheet


@dataclass(frozen=True)
class View:
    """Widok arkusza po otwarciu w Excelu: górna lewa komórka, zaznaczona komórka i powiększenie w procentach."""

    top_left_cell: str
    selected_cell: str
    zoom: int


TEMPLATE_VIEW = View("A1", "A1", 75)
"""Widok Wzorców taty od #465 i syntetycznych Wzorców: lewy górny róg arkusza w powiększeniu 75%."""


def set_view(sheet: Worksheet, view: View) -> None:
    """Ustawia widok arkusza Wzorca; blokada wierszy (``freeze_panes``) zostaje."""

    sheet_view = sheet.sheet_view
    sheet_view.topLeftCell = view.top_left_cell
    sheet_view.zoomScale = sheet_view.zoomScaleNormal = view.zoom
    for selection in sheet_view.selection:
        selection.activeCell = selection.sqref = view.selected_cell


def opening_view(path: Path) -> SheetView:
    """Cały widok aktywnego arkusza, z którym Excel otwiera plik: górna lewa komórka, blokada, zaznaczenie i
    powiększenie."""

    workbook = load_workbook(path)
    try:
        return active_worksheet(workbook).sheet_view
    finally:
        workbook.close()


def view_of(sheet_view: SheetView) -> View:
    """Górna lewa komórka, zaznaczona komórka aktywnego okienka i powiększenie widoku (bez zapisanego: 100%, jak w
    Excelu)."""

    active_pane = None if sheet_view.pane is None else sheet_view.pane.activePane
    (selection,) = [selection for selection in sheet_view.selection if selection.pane == active_pane]
    return View(sheet_view.topLeftCell or "A1", selection.activeCell, sheet_view.zoomScale or 100)


def leave_data_row_leftovers(path: Path, *, first_row: int, rows: int, columns: int) -> None:
    """Wiersze danych Wzorca od ``first_row`` z resztkami formatowania w kolumnach A..``columns``."""

    workbook = load_workbook(path)
    try:
        sheet = active_worksheet(workbook)
        for row in range(first_row, first_row + rows):
            sheet.row_dimensions[row].height = LEFTOVER_ROW_HEIGHT
            for column in range(1, columns + 1):
                cell = sheet.cell(row, column)
                cell.font = LEFTOVER_FONT
                cell.alignment = LEFTOVER_ALIGNMENT
                cell.border = Border()
                cell.number_format = LEFTOVER_DATE_FORMAT
        workbook.save(path)
    finally:
        workbook.close()


def column_widths(sheet: Worksheet, max_col: int) -> list[float | None]:
    """Szerokość kolumn A..``max_col`` tak, jak widzi ją Excel; ``None`` dla kolumny bez własnej szerokości.

    Excel zapisuje sąsiednie kolumny o tej samej szerokości jednym zakresem, więc zakresy są rozwijane.
    """

    widths: list[float | None] = [None] * max_col
    for dimension in sheet.column_dimensions.values():
        if dimension.min is None or dimension.max is None or not dimension.customWidth:
            continue
        for column in range(dimension.min, min(dimension.max, max_col) + 1):
            widths[column - 1] = dimension.width
    return widths


def add_external_link_fixture(path: Path) -> None:
    """Dodaje minimalny, prawidłowy link zewnętrzny do syntetycznego XLSX."""

    ns_main = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    ns_rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    ns_pkg_rel = "http://schemas.openxmlformats.org/package/2006/relationships"
    ns_content = "http://schemas.openxmlformats.org/package/2006/content-types"
    ET.register_namespace("", ns_main)
    ET.register_namespace("r", ns_rel)
    ET.register_namespace("", ns_pkg_rel)

    descriptor = {
        "xl/externalLinks/externalLink1.xml": (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<externalLink xmlns="{ns_main}" xmlns:r="{ns_rel}">'
            '<externalBook><sheetNames><sheetName val="Słownik"/></sheetNames>'
            '<definedNames/><sheetDataSet><sheetData sheetId="0"/></sheetDataSet>'
            "</externalBook></externalLink>"
        ).encode("utf-8"),
        "xl/externalLinks/_rels/externalLink1.xml.rels": (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<Relationships xmlns="{ns_pkg_rel}">'
            '<Relationship Id="rId2" Type="http://schemas.microsoft.com/office/2019/04/relationships/externalLinkLongPath" '
            'Target="https://example.invalid/workbook.xlsm" TargetMode="External"/>'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/externalLinkPath" '
            'Target="file:///C:/workbook.xlsm" TargetMode="External"/>'
            "</Relationships>"
        ).encode("utf-8"),
    }
    file_descriptor, temp_name = tempfile.mkstemp(suffix=".xlsx", dir=path.parent)
    os.close(file_descriptor)
    temporary_path = Path(temp_name)
    try:
        with zipfile.ZipFile(path) as source, zipfile.ZipFile(temporary_path, "w") as target:
            for info in source.infolist():
                data = source.read(info.filename)
                if info.filename == "xl/workbook.xml":
                    root = ET.fromstring(data)
                    references = ET.SubElement(root, f"{{{ns_main}}}externalReferences")
                    ET.SubElement(references, f"{{{ns_main}}}externalReference", {f"{{{ns_rel}}}id": "rIdSynExternal1"})
                    data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                elif info.filename == "xl/_rels/workbook.xml.rels":
                    root = ET.fromstring(data)
                    ET.SubElement(
                        root,
                        f"{{{ns_pkg_rel}}}Relationship",
                        {
                            "Id": "rIdSynExternal1",
                            "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/externalLink",
                            "Target": "externalLinks/externalLink1.xml",
                        },
                    )
                    data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                elif info.filename == "[Content_Types].xml":
                    root = ET.fromstring(data)
                    ET.SubElement(
                        root,
                        f"{{{ns_content}}}Override",
                        {
                            "PartName": "/xl/externalLinks/externalLink1.xml",
                            "ContentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.externalLink+xml",
                        },
                    )
                    data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                target.writestr(info, data)
            for name, data in descriptor.items():
                target.writestr(name, data)
        os.replace(temporary_path, path)
    finally:
        temporary_path.unlink(missing_ok=True)


def external_link_parts(path: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(path) as archive:
        return {
            info.filename: archive.read(info.filename)
            for info in archive.infolist()
            if info.filename.startswith("xl/externalLinks/")
        }


def cut_sheet_after_data_rows(path: Path, sheet_part: str = "xl/worksheets/sheet1.xml") -> None:
    """Skoroszyt się otwiera, ale XML arkusza ``sheet_part`` urywa się za jego ostatnim wierszem.

    Początek arkusza, np. WYKONAWCA w H17, daje się odczytać; błąd wychodzi dopiero przy dalszym czytaniu.
    """

    rewrite_parts(path, lambda name, data: data[: data.rindex(b"</sheetData>")] if name == sheet_part else data)


def cut_sheet_after_sheet_data(path: Path) -> None:
    """Wszystkie Wiersze danych arkusza są całe, ale XML arkusza urywa się zaraz za ``</sheetData>``."""

    rewrite_parts(
        path,
        lambda name, data: data[: data.index(b"</sheetData>") + len(b"</sheetData>")]
        if name == "xl/worksheets/sheet1.xml"
        else data,
    )


def break_row_far_below_data_rows(path: Path) -> None:
    """XML arkusza jest zepsuty dopiero w wierszu 10000, daleko pod Wierszami danych i wierszami z zerami.

    Puste wiersze przed nim odsuwają uszkodzenie o ponad 16 KB XML: openpyxl tylko do odczytu parsuje
    arkusz porcjami tej wielkości, więc uszkodzenie w pierwszej porcji wyszłoby już przy wierszu nagłówka.
    """

    padding = b"".join(b'<row r="%d"/>' % number for number in range(5000, 7000))
    broken_row = b'<row r="10000"><c r="A10000"><v>1</v></wiersz></row>'

    def change(name: str, data: bytes) -> bytes:
        if name != "xl/worksheets/sheet1.xml":
            return data
        data = re.sub(rb'(<dimension ref="[A-Z]+[0-9]+:[A-Z]+)[0-9]+"', rb'\g<1>10000"', data)  # jak zapisuje Excel
        return data.replace(b"</sheetData>", padding + broken_row + b"</sheetData>")

    rewrite_parts(path, change)


def rewrite_parts(path: Path, change: Callable[[str, bytes], bytes | None]) -> None:
    """Przepisuje części archiwum .xlsx; ``None`` usuwa część."""

    with zipfile.ZipFile(path) as source:
        members = [(info, source.read(info.filename)) for info in source.infolist()]
    with zipfile.ZipFile(path, "w") as target:
        for info, data in members:
            changed = change(info.filename, data)
            if changed is not None:
                target.writestr(info, changed)


def break_empty_data_row(path: Path) -> None:
    """XML arkusza jest zepsuty w środku: wiersz 18, bez danych w A:K, kończy się złym znacznikiem."""

    def change(name: str, data: bytes) -> bytes:
        if name != "xl/worksheets/sheet1.xml":
            return data
        start = data.index(b'<row r="18"')
        end = data.index(b"</row>", start)
        return data[:end] + b"</wiersz>" + data[end + len(b"</row>") :]

    rewrite_parts(path, change)


def replace_cells(path: Path, cells: dict[str, str]) -> None:
    """Podmienia istniejące komórki arkusza ``xl/worksheets/sheet1.xml`` na podany XML, np. jak zapisałby je Excel.

    Klucz to adres komórki (``A18``), wartość to cały element ``<c>``; komórka musi już być w arkuszu.
    """

    def change(name: str, data: bytes) -> bytes:
        if name != "xl/worksheets/sheet1.xml":
            return data
        for reference, cell in cells.items():
            pattern = rb'<c r="' + reference.encode("ascii") + rb'"(?:[^>]*?/>|[^>]*?>.*?</c>)'
            data, count = re.subn(pattern, lambda _match: cell.encode("utf-8"), data, count=1, flags=re.S)
            assert count == 1, reference
        return data

    rewrite_parts(path, change)


def without_styles(path: Path) -> None:
    """Komórki zostają ze stylem, ale ``xl/styles.xml`` znika."""

    rewrite_parts(path, lambda name, data: None if name == "xl/styles.xml" else data)


def without_sheet(path: Path) -> None:
    rewrite_parts(
        path,
        lambda name, data: re.sub(rb"<sheets>.*?</sheets>", b"<sheets/>", data) if name == "xl/workbook.xml" else data,
    )


_REFUSED = b"<!-- openpyxl -->"
_COPIED = b"<!-- kopia -->"


def _rewrite_worksheets(path: Path, change: Callable[[bytes], bytes]) -> None:
    """Zmienia XML każdego arkusza; każdy musi się zmienić, więc znacznik trafia też do aktywnego."""

    unchanged: list[str] = []

    def each(name: str, data: bytes) -> bytes:
        if not (name.startswith("xl/worksheets/") and name.endswith(".xml")):
            return data
        changed = change(data)
        if changed == data:
            unchanged.append(name)
        return changed

    rewrite_parts(path, each)
    assert not unchanged, f"arkusze bez znacznika w {path.name}: {unchanged}"


def refused_by_copy(path: Path) -> Path:
    """Bliźniak Wzorca, którego kopia (ADR 0008) nie przyjmuje: komentarz XML w danych każdego arkusza.

    openpyxl pomija komentarz, więc z bliźniaka plik powstaje drogą openpyxl.
    """

    def change(data: bytes) -> bytes:
        empty = re.sub(rb"<sheetData\s*/>", b"<sheetData>" + _REFUSED + b"</sheetData>", data, count=1)
        return empty if empty != data else data.replace(b"<sheetData>", b"<sheetData>" + _REFUSED, 1)

    _rewrite_worksheets(path, change)
    return path


def marked_for_copy(path: Path) -> Path:
    """Wzorzec ze znacznikiem za danymi każdego arkusza: kopia przepisuje go bajt w bajt, openpyxl pomija."""

    _rewrite_worksheets(path, lambda data: data.replace(b"</worksheet>", _COPIED + b"</worksheet>", 1))
    return path


def written_by_copy(path: Path) -> bool:
    """Plik z Wzorca oznaczonego ``marked_for_copy`` powstał kopią, a nie drogą openpyxl."""

    with zipfile.ZipFile(path) as archive:
        return any(_COPIED in archive.read(name) for name in archive.namelist() if name.startswith("xl/worksheets/"))


def has_xml_comment(path: Path) -> bool:
    """Arkusz skoroszytu ma komentarz XML: kopia Wzorca go zachowuje, openpyxl nie zapisuje."""

    with zipfile.ZipFile(path) as archive:
        return any(b"<!--" in archive.read(name) for name in archive.namelist() if name.startswith("xl/worksheets/"))


def hyperlinks(path: Path) -> dict[str, str]:
    """Hiperłącza aktywnego arkusza tak, jak otwiera je Excel: adres komórki → adres (z ``#miejsce``)."""

    workbook = load_workbook(path)
    try:
        sheet = active_worksheet(workbook)
        return {
            cell.coordinate: cell.hyperlink.target + (f"#{cell.hyperlink.location}" if cell.hyperlink.location else "")
            for row in sheet.iter_rows()
            for cell in row
            if cell.hyperlink is not None and cell.hyperlink.target
        }
    finally:
        workbook.close()
