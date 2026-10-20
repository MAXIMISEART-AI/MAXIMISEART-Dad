"""Porównanie Folderów rozliczeń pracowników w testach: moduł z ``tools/`` i pary syntetycznych skoroszytów.

Moduł porównania służy tylko benchmarkowi, więc leży w ``tools/``, poza pakietem
``rozliczenia``. Testy importują go stąd, tak jak benchmark: z folderu ``tools``.
"""

from __future__ import annotations

import re
import shutil
import sys
import zipfile
from collections.abc import Callable
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.worksheet import Worksheet

from tests.fixtures.workbooks import active_worksheet, add_external_link_fixture, rewrite_parts

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "tools"))

from porownanie_rozliczen import CATEGORIES, compare_worker_settlements  # noqa: E402

__all__ = [
    "CATEGORIES",
    "SETTLEMENT_NAME",
    "add_calc_chain",
    "formulas_in_row_13",
    "renumber_cell_styles",
    "share_formulas_in_row_13",
    "SECRET_VALUES",
    "compare_worker_settlements",
    "settlement_folders",
    "write_settlement",
]

SETTLEMENT_NAME = "Rozliczenie 08_14_09_2026 - Adrian Testowy.xlsx"
# Treść wierszy danych, która nie może wyjść poza skoroszyt: adres i numer zlecenia.
SECRET_VALUES = ("ul. Tajna 7", "ZL-SEKRET-123")

Change = Callable[[Workbook, Worksheet], None]


def write_settlement(path: Path, change: Change | None = None) -> Path:
    """Małe Rozliczenie pracownika: nagłówek w 17, dwa Wiersze danych, formuły, style, scalenie, link zewnętrzny."""

    workbook = Workbook()
    sheet = active_worksheet(workbook)
    sheet.title = "Sheet1"
    sheet.merge_cells("A16:M16")
    sheet["A16"] = "ROZLICZENIE"
    sheet.cell(17, 8).value = "WYKONAWCA"
    sheet["A17"].fill = PatternFill(fill_type="solid", fgColor="00AA55")
    sheet["A17"].font = Font(bold=True)
    sheet.row_dimensions[17].height = 30
    sheet.column_dimensions["B"].width = 40
    for row, (address, order) in enumerate([SECRET_VALUES, ("ul. Syntetyczna 2", "ZL-2")], start=18):
        sheet.cell(row, 2).value = address
        sheet.cell(row, 6).value = order
        sheet.cell(row, 8).value = "adrian.testowy"
        sheet.cell(row, 14).value = 3
        sheet.cell(row, 15).value = 1.5
        sheet.cell(row, 14).number_format = "0"
        sheet.cell(row, 14).border = Border(left=Side(style="thin"))
        sheet.cell(row, 14).alignment = Alignment(horizontal="center")
        sheet.cell(row, 14).protection = Protection(locked=False)
        sheet.cell(row, 47).value = f"=SUM(N{row}:AP{row})"
    workbook.defined_names["Stawki"] = DefinedName("Stawki", attr_text="Sheet1!$N$1:$AP$1")
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    if change is not None:
        change(workbook, sheet)
    workbook.save(path)
    add_external_link_fixture(path)
    return path


def settlement_folders(tmp_path: Path) -> tuple[Path, Path]:
    """Wynik referencyjny i kandydat: dwa Foldery rozliczeń pracowników z tym samym Rozliczeniem pracownika."""

    reference = tmp_path / "referencja"
    candidate = tmp_path / "kandydat"
    reference.mkdir()
    candidate.mkdir()
    write_settlement(reference / SETTLEMENT_NAME)
    shutil.copy2(reference / SETTLEMENT_NAME, candidate / SETTLEMENT_NAME)
    return reference, candidate


def formulas_in_row_13(_workbook: Workbook, sheet: Worksheet) -> None:
    """Trzy formuły o tej samej treści względnej w N13:P13, jak sumy we Wzorcu."""

    for column in ("N", "O", "P"):
        sheet[f"{column}13"] = f"={column}14*{column}15"


def share_formulas_in_row_13(path: Path) -> None:
    """N13:P13 jako jedna wspólna formuła (``t="shared"``), tak jak zapisuje je Excel we Wzorcu."""

    def change(name: str, data: bytes) -> bytes:
        if name != "xl/worksheets/sheet1.xml":
            return data
        shared = data.replace(b"<f>N14*N15</f>", b'<f t="shared" ref="N13:P13" si="0">N14*N15</f>')
        for column in (b"O", b"P"):
            shared = shared.replace(b"<f>" + column + b"14*" + column + b"15</f>", b'<f t="shared" si="0" />')
        assert shared.count(b't="shared"') == 3
        return shared

    rewrite_parts(path, change)


def add_calc_chain(path: Path) -> None:
    """Dodaje ``xl/calcChain.xml`` z relacją i typem zawartości, jak w skoroszycie zapisanym przez Excel."""

    calc_chain = (
        b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        b'<calcChain xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        b'<c r="AU18" i="1"/><c r="AU19"/></calcChain>'
    )

    def change(name: str, data: bytes) -> bytes:
        if name == "[Content_Types].xml":
            changed, count = re.subn(
                rb"</((?:\w+:)?)Types>",
                rb'<\1Override PartName="/xl/calcChain.xml" ContentType='
                rb'"application/vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml"/></\1Types>',
                data,
            )
        elif name == "xl/_rels/workbook.xml.rels":
            changed, count = re.subn(
                rb"</((?:\w+:)?)Relationships>",
                rb'<\1Relationship Id="rIdCalcChain" Target="calcChain.xml" Type='
                rb'"http://schemas.openxmlformats.org/officeDocument/2006/relationships/calcChain"/></\1Relationships>',
                data,
            )
        else:
            return data
        assert count == 1
        return changed

    rewrite_parts(path, change)
    with zipfile.ZipFile(path, "a") as archive:
        archive.writestr("xl/calcChain.xml", calc_chain)


def renumber_cell_styles(path: Path) -> None:
    """Komórki z ostatnim stylem dostają inne numery stylu i czcionki o tej samej treści.

    Na końcu ``fonts`` powstaje kopia czcionki tego stylu, a na końcu ``cellXfs`` kopia stylu, która na nią wskazuje.
    """

    with zipfile.ZipFile(path) as archive:
        styles = archive.read("xl/styles.xml")
    fonts = re.search(rb'<fonts count="(\d+)">(.*?)</fonts>', styles, re.DOTALL)
    cell_styles = re.search(rb'<cellXfs count="(\d+)">(.*?)</cellXfs>', styles, re.DOTALL)
    assert fonts is not None and cell_styles is not None
    font_count, style_count = int(fonts[1]), int(cell_styles[1])
    last_style = b"<xf " + cell_styles[2].rsplit(b"<xf ", 1)[1]
    font_id = re.search(rb'fontId="(\d+)"', last_style)
    assert font_id is not None
    font = re.findall(rb"<font>.*?</font>", fonts[2], re.DOTALL)[int(font_id[1])]
    new_style = last_style.replace(font_id[0], b'fontId="%d"' % font_count)
    new_styles = (
        styles[: fonts.start()]
        + b'<fonts count="%d">%s%s</fonts>' % (font_count + 1, fonts[2], font)
        + styles[fonts.end() : cell_styles.start()]
        + b'<cellXfs count="%d">%s%s</cellXfs>' % (style_count + 1, cell_styles[2], new_style)
        + styles[cell_styles.end() :]
    )

    def change(name: str, data: bytes) -> bytes:
        if name == "xl/styles.xml":
            return new_styles
        if name == "xl/worksheets/sheet1.xml":
            renumbered = data.replace(b' s="%d"' % (style_count - 1), b' s="%d"' % style_count)
            assert renumbered != data
            return renumbered
        return data

    rewrite_parts(path, change)
