from __future__ import annotations

from dataclasses import dataclass, replace
from datetime import date, datetime, timedelta
from pathlib import Path
import time

from openpyxl import Workbook, load_workbook
import pytest

from rozliczenia.domain import (
    IssueCountReady,
    PhaseEnded,
    PhaseFailed,
    PhaseStarted,
    PlanReady,
    ProgressEvent,
    ProgressPhase,
)
from rozliczenia.progress import PhaseState, ProgressProjection
from rozliczenia.message import StopError
from rozliczenia.workbook_files import ExternalLinksError
import rozliczenia.workbook_files as workbook_files

from tests.fixtures.preparation import (
    COLLECTIVE_FILE_NAME,
    COLLECTIVE_TEMPLATE_ZERO_ROWS,
    INTERNAL_SETTLEMENT_NAME,
    INTERNAL_TEMPLATE_HEADERS,
    NOT_TRANSFERRED_PREFIXES,
    OLD_REGISTER_HEADERS,
    PERIOD,
    TEMPLATE_NARROW_COLUMN_WIDTHS,
    PreparationFixture,
    RegisterRow,
    default_register_rows,
    in_old_layout,
    make_dad_layout_fixture,
    make_preparation_fixture,
    register_order,
    template_column_width,
    write_collective_template,
    write_internal_template,
    write_register,
)
from tests.fixtures.benchmark import COLLECTIVE_TEMPLATE, INTERNAL_TEMPLATE, write_synthetic_templates
from tests.fixtures.workbooks import active_worksheet, column_widths, external_link_parts
from tests.fixtures.settlement import projected
from tests.fixtures.preparation import cell_values, collective_row, incomplete_period_fixture, run_preparation
from tests.fixtures.interlocutor import Tata


def sheet_rows(path: Path, *, min_row: int, max_col: int) -> list[tuple[object, ...]]:
    workbook = load_workbook(path)
    try:
        sheet = active_worksheet(workbook)
        return list(sheet.iter_rows(min_row=min_row, max_col=max_col, values_only=True))
    finally:
        workbook.close()


def internal_row(number: int, started: datetime, successful: str, worker: str) -> tuple[object, ...]:
    return (
        started,
        "POZNAŃ",
        f"syntetyczny adres SYN-{number:03d}",
        f"UM-SYN-{number:03d}",
        "dekoder",
        "instalacja",
        f"ZL-SYN-{number:03d}",
        worker,
        f"OR-SYN-{number:03d}",
        successful,
        "montaż",
        None,
        None,
    )


def test_rozliczenie_wewnetrzne_contains_exactly_period_orders_in_register_order(tmp_path: Path) -> None:
    fixture = make_preparation_fixture(tmp_path)

    run_preparation(fixture)

    assert sheet_rows(fixture.internal_settlement_path, min_row=2, max_col=len(INTERNAL_TEMPLATE_HEADERS)) == [
        internal_row(2, datetime(2026, 9, 8, 0, 0), "tak", "adrian.testowy"),
        internal_row(3, datetime(2026, 9, 9, 10, 30), "NIE", "dariusz.fikcyjny2"),
        internal_row(4, datetime(2026, 9, 12, 12, 0), " Tak ", "kamil.syntetyczny"),
        internal_row(6, datetime(2026, 9, 14, 18, 30), "TAK", "dariusz.fikcyjny2"),
    ]


def test_plik_zbiorczy_contains_exactly_zlecenia_skuteczne_from_row_18(tmp_path: Path) -> None:
    fixture = make_preparation_fixture(tmp_path)

    run_preparation(fixture)

    assert cell_values(fixture.collective_file_path, min_row=18, max_row=24, min_col=1, max_col=11) == [
        collective_row(2, datetime(2026, 9, 8, 0, 0), "tak", "adrian.testowy"),
        collective_row(4, datetime(2026, 9, 12, 12, 0), " Tak ", "kamil.syntetyczny"),
        collective_row(6, datetime(2026, 9, 14, 18, 30), "TAK", "dariusz.fikcyjny2"),
        *[(None,) * 11] * 4,
    ]


def test_akcje_zwrot_and_uwagi_from_rejestr_stay_out_of_both_period_files(tmp_path: Path) -> None:
    fixture = make_preparation_fixture(tmp_path)

    run_preparation(fixture)

    for path in (fixture.internal_settlement_path, fixture.collective_file_path):
        workbook = load_workbook(path)
        try:
            texts = [
                cell.value
                for sheet in workbook.worksheets
                for row in sheet.iter_rows()
                for cell in row
                if isinstance(cell.value, str)
            ]
        finally:
            workbook.close()
        assert [text for text in texts if any(prefix in text for prefix in NOT_TRANSFERRED_PREFIXES.values())] == []


def test_rejestr_in_old_layout_still_fills_both_period_files(tmp_path: Path) -> None:
    fixture = make_preparation_fixture(tmp_path)
    write_register(fixture.register_path, in_old_layout(default_register_rows()), headers=OLD_REGISTER_HEADERS)

    run_preparation(fixture)

    assert sheet_rows(fixture.internal_settlement_path, min_row=2, max_col=len(INTERNAL_TEMPLATE_HEADERS)) == [
        internal_row(2, datetime(2026, 9, 8, 0, 0), "tak", "adrian.testowy"),
        internal_row(3, datetime(2026, 9, 9, 10, 30), "NIE", "dariusz.fikcyjny2"),
        internal_row(4, datetime(2026, 9, 12, 12, 0), " Tak ", "kamil.syntetyczny"),
        internal_row(6, datetime(2026, 9, 14, 18, 30), "TAK", "dariusz.fikcyjny2"),
    ]
    assert cell_values(fixture.collective_file_path, min_row=18, max_row=20, min_col=1, max_col=11) == [
        collective_row(2, datetime(2026, 9, 8, 0, 0), "tak", "adrian.testowy"),
        collective_row(4, datetime(2026, 9, 12, 12, 0), " Tak ", "kamil.syntetyczny"),
        collective_row(6, datetime(2026, 9, 14, 18, 30), "TAK", "dariusz.fikcyjny2"),
    ]


def test_plik_zbiorczy_keeps_wzorzec_rows_1_to_17_quantities_and_formulas(tmp_path: Path) -> None:
    fixture = make_preparation_fixture(tmp_path)

    run_preparation(fixture)

    template, created = fixture.collective_template_path, fixture.collective_file_path
    assert cell_values(created, min_row=1, max_row=17, min_col=1, max_col=47) == cell_values(
        template, min_row=1, max_row=17, min_col=1, max_col=47
    )
    assert cell_values(created, min_row=18, max_row=30, min_col=12, max_col=47) == cell_values(
        template, min_row=18, max_row=30, min_col=12, max_col=47
    )
    assert cell_values(created, min_row=18, max_row=18, min_col=14, max_col=14) == [(0,)]
    assert cell_values(created, min_row=18, max_row=18, min_col=47, max_col=47) == [("=N18*$N$1",)]
    workbook = load_workbook(created)
    try:
        assert workbook["Rates"]["B1"].value == 17.5
    finally:
        workbook.close()


def test_przygotowanie_na_syntetycznych_wzorcach_ksztaltu_taty_wypelnia_pliki_okresu_jak_na_testowych(
    tmp_path: Path,
) -> None:
    (tmp_path / "testowe").mkdir()
    (tmp_path / "syntetyczne").mkdir()
    reference = make_preparation_fixture(tmp_path / "testowe")
    templates = write_synthetic_templates(tmp_path / "wzorce")
    fixture = replace(
        make_preparation_fixture(tmp_path / "syntetyczne"),
        internal_template_path=templates / INTERNAL_TEMPLATE,
        collective_template_path=templates / COLLECTIVE_TEMPLATE,
    )

    run_preparation(reference)
    summary = run_preparation(fixture)

    assert summary.warnings == []
    created, template = fixture.collective_file_path, fixture.collective_template_path
    assert cell_values(created, min_row=18, max_row=21, min_col=1, max_col=11) == cell_values(
        reference.collective_file_path, min_row=18, max_row=21, min_col=1, max_col=11
    )
    assert cell_values(created, min_row=1, max_row=17, min_col=1, max_col=51) == cell_values(
        template, min_row=1, max_row=17, min_col=1, max_col=51
    )
    assert cell_values(created, min_row=18, max_row=21, min_col=14, max_col=51) == cell_values(
        template, min_row=18, max_row=21, min_col=14, max_col=51
    )
    assert external_link_parts(created) == external_link_parts(template)
    def filled_rows(path: Path) -> int:
        return sum(any(value is not None for value in row) for row in sheet_rows(path, min_row=2, max_col=13))

    assert filled_rows(fixture.internal_settlement_path) == filled_rows(reference.internal_settlement_path) == 4


def test_przygotowanie_leaves_rejestr_and_wzorce_byte_identical(tmp_path: Path) -> None:
    fixture = make_preparation_fixture(tmp_path)
    inputs = (fixture.register_path, fixture.internal_template_path, fixture.collective_template_path)
    before = [path.read_bytes() for path in inputs]

    run_preparation(fixture)

    assert [path.read_bytes() for path in inputs] == before


def test_summary_reports_counts_and_file_names_without_client_data(tmp_path: Path) -> None:
    fixture = make_preparation_fixture(tmp_path)

    summary = run_preparation(fixture)

    assert summary.period == PERIOD
    assert summary.period_order_count == 4
    assert summary.successful_order_count == 3
    assert summary.skipped_empty_rows == 5
    assert summary.warnings == []
    assert summary.internal_settlement_file == fixture.internal_settlement_path
    assert summary.collective_file == fixture.collective_file_path
    assert sorted(path.name for path in fixture.period_directory.iterdir()) == [
        COLLECTIVE_FILE_NAME,
        INTERNAL_SETTLEMENT_NAME,
    ]
    assert "SYN-" not in repr(summary)


def test_existing_empty_folder_okresu_is_allowed(tmp_path: Path) -> None:
    fixture = make_preparation_fixture(tmp_path)
    fixture.period_directory.mkdir()

    run_preparation(fixture)

    assert fixture.internal_settlement_path.is_file()
    assert fixture.collective_file_path.is_file()


def test_in_dad_layout_period_files_go_to_folder_okresu_found_in_its_month(tmp_path: Path) -> None:
    fixture = make_dad_layout_fixture(tmp_path)

    summary = run_preparation(fixture)

    assert summary.period_directory == fixture.period_directory
    assert fixture.period_directory.parent.parent.name == "09_Wrzesień_2026"
    assert sorted(path.name for path in fixture.period_directory.iterdir()) == sorted(
        [INTERNAL_SETTLEMENT_NAME, COLLECTIVE_FILE_NAME]
    )
    assert not (fixture.periods_directory / PERIOD).exists()


@pytest.mark.parametrize("levels_up", [0, 1, 2], ids=["INEA", "miesiąc", "rok"])
def test_in_dad_layout_folder_okresow_chosen_inside_another_month_finds_folder_okresu_in_its_month(
    tmp_path: Path, levels_up: int
) -> None:
    fixture = make_dad_layout_fixture(tmp_path)
    august = fixture.periods_directory / "Rozliczenia 2026" / "08_Sierpień_2026" / "INEA"
    chosen = ([august, *august.parents])[levels_up]

    summary = run_preparation(replace(fixture, periods_directory=chosen))

    assert summary.period_directory == fixture.period_directory
    assert sorted(path.name for path in fixture.period_directory.iterdir()) == sorted(
        [INTERNAL_SETTLEMENT_NAME, COLLECTIVE_FILE_NAME]
    )
    assert not (chosen / PERIOD).exists()
    assert sorted(path.name for path in august.iterdir()) == ["25_31_08_2026"]


@pytest.mark.parametrize("existing_name", [INTERNAL_SETTLEMENT_NAME, COLLECTIVE_FILE_NAME])
def test_nienadpisywanie_stops_when_period_file_exists(tmp_path: Path, existing_name: str) -> None:
    fixture = make_preparation_fixture(tmp_path)
    fixture.period_directory.mkdir()
    existing = fixture.period_directory / existing_name
    existing.write_bytes(b"reczne poprawki taty")

    with pytest.raises(StopError) as raised:
        run_preparation(fixture)

    message = raised.value.message.text
    assert message == (
        "Co się stało: Plik tego okresu już istnieje, więc nie został nadpisany. "
        "Nie utworzono żadnego pliku.\n"
        f"Gdzie: folder INEA\\{PERIOD}, plik {existing_name}\n"
        "Co zrobić: Sprawdź, czy wybrano właściwy okres. Jeśli plik ma powstać od nowa, "
        "przenieś albo usuń istniejący plik i uruchom ponownie."
    )
    assert raised.value.message.code == "PLIK_OKRESU_ISTNIEJE"
    assert [path.name for path in fixture.period_directory.iterdir()] == [existing_name]
    assert existing.read_bytes() == b"reczne poprawki taty"
    assert [path.name for path in fixture.periods_directory.iterdir()] == [PERIOD]


@pytest.mark.parametrize(
    "token",
    [
        "8_14_09_2026",
        "09_15_09_2026",
        "08_15_09_2026",
        "okres",
        "../14_09_2026",
        # Na przełomie miesięcy i lat miesiąc i rok tokenu są z poniedziałku, nie z wtorku.
        "27_02_10_2026",
        "29_04_12_2026",
    ],
)
def test_invalid_period_token_stops_before_reading_rejestr(tmp_path: Path, token: str) -> None:
    fixture = make_preparation_fixture(tmp_path)

    with pytest.raises(StopError) as raised:
        run_preparation(fixture, token)

    assert raised.value.message.code == "ZLY_TOKEN_OKRESU"
    assert list(fixture.periods_directory.iterdir()) == []


@pytest.mark.parametrize(
    ("token", "tuesday"),
    [("27_02_11_2026", datetime(2026, 10, 27)), ("29_04_01_2027", datetime(2026, 12, 29))],
)
def test_period_crossing_month_or_year_takes_month_and_year_of_monday(
    tmp_path: Path, token: str, tuesday: datetime
) -> None:
    rows: list[RegisterRow] = [
        register_order(1, tuesday - timedelta(minutes=1)),
        register_order(2, tuesday + timedelta(hours=8)),
        register_order(3, tuesday + timedelta(days=7, minutes=-1)),
        register_order(4, tuesday + timedelta(days=7)),
    ]
    fixture = make_preparation_fixture(tmp_path, rows)

    summary = run_preparation(fixture, token)

    assert summary.period_order_count == 2
    assert list(summary.orders_per_day)[0] == tuesday.date()
    assert (fixture.periods_directory / token / f"Rozliczenie wew. {token}.xlsx").is_file()


EXPECTED_ORDERS_PER_DAY = {
    date(2026, 9, 8): 1,
    date(2026, 9, 9): 1,
    date(2026, 9, 10): 0,
    date(2026, 9, 11): 0,
    date(2026, 9, 12): 1,
    date(2026, 9, 13): 0,
    date(2026, 9, 14): 1,
}


def test_summary_reports_orders_per_day_for_every_day_of_the_period(tmp_path: Path) -> None:
    fixture = make_preparation_fixture(tmp_path)

    summary = run_preparation(fixture)

    assert summary.orders_per_day == EXPECTED_ORDERS_PER_DAY
    assert list(summary.orders_per_day) == sorted(EXPECTED_ORDERS_PER_DAY)


def test_dry_run_returns_plan_with_orders_per_day_and_writes_nothing(tmp_path: Path) -> None:
    fixture = make_preparation_fixture(tmp_path)
    inputs = (fixture.register_path, fixture.internal_template_path, fixture.collective_template_path)
    before = [path.read_bytes() for path in inputs]

    summary = run_preparation(fixture, dry_run=True)

    assert summary.period == PERIOD
    assert summary.period_order_count == 4
    assert summary.successful_order_count == 3
    assert summary.skipped_empty_rows == 5
    assert summary.orders_per_day == EXPECTED_ORDERS_PER_DAY
    assert summary.internal_settlement_file == fixture.internal_settlement_path
    assert summary.collective_file == fixture.collective_file_path
    assert list(fixture.periods_directory.iterdir()) == []
    assert [path.read_bytes() for path in inputs] == before


STAGES_IN_ORDER = [
    (ProgressPhase.CHECKING, PhaseStarted),
    (ProgressPhase.CHECKING, PhaseEnded),
    (ProgressPhase.READING, PhaseStarted),
    (ProgressPhase.READING, PhaseEnded),
    (ProgressPhase.READING, IssueCountReady),
    (ProgressPhase.PLANNING, PhaseStarted),
    (ProgressPhase.PLANNING, PhaseEnded),
    (ProgressPhase.PLANNING, PlanReady),
    (ProgressPhase.SAVING, PhaseStarted),
    (ProgressPhase.SAVING, PhaseEnded),
]


@pytest.mark.parametrize("dry_run", [False, True])
def test_zdarzenia_przebiegu_pass_four_stages_in_order_without_worker_events(
    tmp_path: Path, dry_run: bool
) -> None:
    fixture = make_preparation_fixture(tmp_path, [register_order(1, datetime(2026, 9, 13, 23, 59))])
    recorded: list[ProgressEvent] = []
    projection = ProgressProjection()

    def recording_observer(event: ProgressEvent) -> None:
        recorded.append(event)
        projection.update(event)

    summary = run_preparation(fixture, dry_run=dry_run, observer=recording_observer)

    assert [(event.phase, type(event)) for event in recorded] == STAGES_IN_ORDER
    assert set(projection.snapshot.phase_states.values()) == {PhaseState.COMPLETED}
    assert [warning.code for warning in summary.warnings] == ["NIEPELNY_OKRES"]
    assert projection.snapshot.issue_count == 1
    assert projection.snapshot.template_total == 0
    assert "SYN-" not in repr(recorded)


def test_progress_projection_accepts_przygotowanie_pliku_zbiorczego_interrupted_in_zapisywanie_with_its_time(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    fixture = make_preparation_fixture(tmp_path)
    projection = ProgressProjection()
    save_s = 0.05

    def fail_save(workbook: Workbook, target: Path, *, external_links_from: Path | None = None) -> None:
        time.sleep(save_s)
        raise ExternalLinksError("wymuszony błąd")

    monkeypatch.setattr(workbook_files, "save_workbook", fail_save)

    with pytest.raises(StopError):
        run_preparation(fixture, observer=projected(projection))

    assert projection.snapshot.phase_states == {
        ProgressPhase.CHECKING: PhaseState.COMPLETED,
        ProgressPhase.READING: PhaseState.COMPLETED,
        ProgressPhase.PLANNING: PhaseState.COMPLETED,
        ProgressPhase.SAVING: PhaseState.FAILED,
    }
    assert projection.snapshot.phase_durations_ms[ProgressPhase.SAVING] >= save_s * 1000


def test_interrupted_przygotowanie_emits_failed_without_exception_text(tmp_path: Path) -> None:
    fixture = make_preparation_fixture(tmp_path)
    fixture.period_directory.mkdir()
    fixture.collective_file_path.write_bytes(b"reczne poprawki taty")
    recorded: list[ProgressEvent] = []

    with pytest.raises(StopError):
        run_preparation(fixture, observer=recorded.append)

    assert [(event.phase, type(event)) for event in recorded] == [
        (ProgressPhase.CHECKING, PhaseStarted),
        (ProgressPhase.CHECKING, PhaseFailed),
    ]
    assert "istnieje" not in repr(recorded)
    assert COLLECTIVE_FILE_NAME not in repr(recorded)


def test_failing_observer_is_detached_and_does_not_change_result_or_files(tmp_path: Path) -> None:
    fixture = make_preparation_fixture(tmp_path)
    calls = 0

    def failing_observer(event: ProgressEvent) -> None:
        nonlocal calls
        calls += 1
        raise RuntimeError("awaria terminala")

    summary = run_preparation(fixture, observer=failing_observer)

    assert calls == 1
    assert summary.period_order_count == 4
    assert summary.successful_order_count == 3
    assert summary.orders_per_day == EXPECTED_ORDERS_PER_DAY
    assert fixture.internal_settlement_path.is_file()
    assert fixture.collective_file_path.is_file()


def cell_formats(path: Path, coordinates: list[str]) -> dict[str, tuple[object, ...]]:
    """Wypełnienie, kolor czcionki, pogrubienie i kursywa wskazanych komórek."""

    workbook = load_workbook(path)
    try:
        sheet = active_worksheet(workbook)
        formats: dict[str, tuple[object, ...]] = {}
        for coordinate in coordinates:
            cell = sheet[coordinate]
            fill = cell.fill.fgColor.rgb if cell.fill.fill_type == "solid" else None
            # Kolor motywu (domyślny tekst) nie jest wyróżnieniem.
            color = cell.font.color.rgb if cell.font.color is not None and cell.font.color.type == "rgb" else None
            formats[coordinate] = (fill, color, bool(cell.font.b), bool(cell.font.i))
        return formats
    finally:
        workbook.close()


def test_highlights_from_rejestr_reach_both_files_and_plain_cells_get_none(tmp_path: Path) -> None:
    fixture = make_preparation_fixture(tmp_path)

    run_preparation(fixture)

    # Zlecenie 2: ADRES wyróżniony; zlecenie 4: WYKONAWCA kursywą; reszta bez wyróżnień.
    assert cell_formats(fixture.internal_settlement_path, ["C2", "H4", "B2", "C3", "H2"]) == {
        "C2": ("FFFFFF00", "FFC00000", True, False),
        "H4": (None, None, False, True),
        "B2": (None, None, False, False),
        "C3": (None, None, False, False),
        "H2": (None, None, False, False),
    }
    assert cell_formats(fixture.collective_file_path, ["B18", "H19", "A18", "B19", "H18"]) == {
        "B18": ("FFFFFF00", "FFC00000", True, False),
        "H19": (None, None, False, True),
        "A18": (None, None, False, False),
        "B19": (None, None, False, False),
        "H18": (None, None, False, False),
    }


def data_fonts(path: Path, *, min_row: int, max_row: int, max_col: int) -> set[tuple[object, object]]:
    workbook = load_workbook(path)
    try:
        sheet = active_worksheet(workbook)
        return {
            (cell.font.name, cell.font.sz)
            for row in sheet.iter_rows(min_row=min_row, max_row=max_row, max_col=max_col)
            for cell in row
        }
    finally:
        workbook.close()


def test_data_rows_take_font_face_and_size_from_wzorzec_not_rejestr(tmp_path: Path) -> None:
    fixture = make_preparation_fixture(tmp_path)

    run_preparation(fixture)

    # Rejestr: Arial 8; wzorce: Calibri 10.
    assert data_fonts(fixture.internal_settlement_path, min_row=2, max_row=5, max_col=11) == {("Calibri", 10)}
    assert data_fonts(fixture.collective_file_path, min_row=18, max_row=20, max_col=11) == {("Calibri", 10)}


def test_cells_without_own_style_take_wzorzec_row_font_before_column_font(tmp_path: Path) -> None:
    fixture = make_preparation_fixture(tmp_path)
    write_internal_template(fixture.internal_template_path, row_style=True)
    write_collective_template(fixture.collective_template_path, row_style=True)

    run_preparation(fixture)

    # Wiersze danych wzorców: styl wiersza Calibri 10; kolumny: Calibri 11; komórki bez własnego stylu.
    assert data_fonts(fixture.internal_settlement_path, min_row=2, max_row=5, max_col=11) == {("Calibri", 10)}
    assert data_fonts(fixture.collective_file_path, min_row=18, max_row=20, max_col=11) == {("Calibri", 10)}


def test_data_column_keeps_date_format_from_rejestr(tmp_path: Path) -> None:
    fixture = make_preparation_fixture(tmp_path)

    run_preparation(fixture)

    internal = load_workbook(fixture.internal_settlement_path)
    collective = load_workbook(fixture.collective_file_path)
    try:
        internal_dates = [active_worksheet(internal).cell(row, 1) for row in range(2, 6)]
        collective_dates = [active_worksheet(collective).cell(row, 7) for row in range(18, 21)]
        assert {
            (type(cell.value), cell.number_format) for cell in internal_dates + collective_dates
        } == {(datetime, "dd.mm.yyyy hh:mm")}
    finally:
        internal.close()
        collective.close()


@dataclass(frozen=True)
class SheetLayout:
    """Szerokości kolumn, wysokości wierszy, nagłówek z formatowaniem i blokada wierszy."""

    widths: list[float | None]
    heights: list[float | None]
    header: list[tuple[object, ...]]
    freeze_panes: str | None


def sheet_layout(path: Path, *, header_row: int, max_col: int) -> SheetLayout:
    workbook = load_workbook(path)
    try:
        sheet = active_worksheet(workbook)
        return SheetLayout(
            widths=column_widths(sheet, max_col),
            heights=[sheet.row_dimensions[row].height for row in range(1, header_row + 6)],
            header=[
                (cell.value, cell.font.name, cell.font.sz, cell.font.b, cell.fill.fgColor.rgb)
                for cell in sheet[header_row][:max_col]
            ],
            freeze_panes=sheet.freeze_panes,
        )
    finally:
        workbook.close()


@pytest.mark.parametrize("kind", ["internal", "collective"])
def test_column_widths_header_and_frozen_row_equal_wzorzec(tmp_path: Path, kind: str) -> None:
    fixture = make_preparation_fixture(tmp_path)

    run_preparation(fixture)

    if kind == "internal":
        template, created, header_row = fixture.internal_template_path, fixture.internal_settlement_path, 1
    else:
        template, created, header_row = fixture.collective_template_path, fixture.collective_file_path, 17
    expected = sheet_layout(template, header_row=header_row, max_col=47)
    assert expected.freeze_panes == f"A{header_row + 1}"
    # Szerokości Wzorca, łącznie z jego bardzo wąskimi kolumnami.
    leading = range(1, max(TEMPLATE_NARROW_COLUMN_WIDTHS) + 1)
    assert expected.widths[: len(leading)] == [template_column_width(column) for column in leading]
    assert sheet_layout(created, header_row=header_row, max_col=47) == expected


def test_plik_zbiorczy_columns_outside_a_to_k_keep_wzorzec_formatting(tmp_path: Path) -> None:
    fixture = make_preparation_fixture(tmp_path)

    run_preparation(fixture)

    def formats(path: Path) -> list[tuple[object, ...]]:
        workbook = load_workbook(path)
        try:
            sheet = active_worksheet(workbook)
            return [
                (cell.coordinate, cell.fill.fill_type, cell.fill.fgColor.rgb, cell.number_format, cell.font.name)
                for row in sheet.iter_rows(min_row=18, max_row=24, min_col=12, max_col=47)
                for cell in row
            ]
        finally:
            workbook.close()

    expected = formats(fixture.collective_template_path)
    assert ("N18", "solid", "FFCCE5FF", "0", "Calibri") in expected
    assert formats(fixture.collective_file_path) == expected


def autofilter_range(path: Path) -> str | None:
    workbook = load_workbook(path)
    try:
        ref: str | None = active_worksheet(workbook).auto_filter.ref
        return ref
    finally:
        workbook.close()


def test_autofilter_covers_header_and_last_order_row_in_both_files(tmp_path: Path) -> None:
    fixture = make_preparation_fixture(tmp_path)

    run_preparation(fixture)

    # Wzorce mają zakres B:C do wiersza +40, który ma zostać zignorowany.
    # Rozliczenie wewnętrzne: nagłówki A:M i 4 zlecenia okresu.
    assert autofilter_range(fixture.internal_settlement_path) == "A1:M5"
    # Plik zbiorczy: nagłówki wiersza 17 od A do AP i 3 Zlecenia skuteczne.
    assert autofilter_range(fixture.collective_file_path) == "A17:AP20"


def test_autofilter_covers_only_header_when_period_has_no_zlecenia_skuteczne(tmp_path: Path) -> None:
    rows: list[RegisterRow] = [register_order(1, datetime(2026, 9, 9, 10, 0), successful="nie")]
    fixture = make_preparation_fixture(tmp_path, rows)

    run_preparation(fixture)

    assert autofilter_range(fixture.internal_settlement_path) == "A1:M2"
    assert autofilter_range(fixture.collective_file_path) == "A17:AP17"


REGISTER_NAME = "ZLECENIA INTERMAX OD 08_11_2021.xlsx"


def stop_message(fixture: PreparationFixture) -> StopError:
    """Uruchamia Przygotowanie, oczekuje zatrzymania i sprawdza, że nic nie powstało."""

    with pytest.raises(StopError) as raised:
        run_preparation(fixture)
    assert list(fixture.periods_directory.iterdir()) == []
    assert "SYN-" not in raised.value.message.text
    return raised.value


def test_row_without_date_between_period_orders_stops_with_register_row_numbers(tmp_path: Path) -> None:
    undated = register_order(90, datetime(2026, 9, 9, 8, 0))
    undated["DATA GODZ. POCZ."] = "środa rano"
    blank_date = register_order(91, datetime(2026, 9, 9, 9, 0))
    blank_date["DATA GODZ. POCZ."] = None
    rows: list[RegisterRow] = [
        register_order(2, datetime(2026, 9, 8, 8, 0)),
        undated,
        None,
        blank_date,
        register_order(3, datetime(2026, 9, 10, 8, 0)),
    ]
    fixture = make_preparation_fixture(tmp_path, rows)

    error = stop_message(fixture)

    assert error.message.code == "WIERSZ_BEZ_DATY"
    assert error.message.text == (
        "Co się stało: Między zleceniami okresu są wiersze z danymi bez poprawnej daty "
        "w kolumnie DATA GODZ. POCZ. Nie utworzono żadnego pliku.\n"
        f"Gdzie: plik {REGISTER_NAME}, arkusz ZLECENIA, wiersze 3, 5\n"
        "Co zrobić: Wpisz poprawną datę w kolumnie DATA GODZ. POCZ. we wskazanych wierszach Rejestru "
        "zleceń albo wyczyść wiersz, jeśli jest zbędny. Zapisz i zamknij Rejestr, i uruchom ponownie."
    )


def test_rows_without_date_outside_period_orders_do_not_stop(tmp_path: Path) -> None:
    note_before = register_order(90, datetime(2026, 9, 1))
    note_before["DATA GODZ. POCZ."] = "uwaga"
    note_after = register_order(91, datetime(2026, 9, 1))
    note_after["DATA GODZ. POCZ."] = "uwaga"
    rows: list[RegisterRow] = [
        note_before,
        register_order(2, datetime(2026, 9, 8, 8, 0)),
        None,
        register_order(3, datetime(2026, 9, 10, 8, 0)),
        note_after,
    ]
    fixture = make_preparation_fixture(tmp_path, rows)

    summary = run_preparation(fixture)

    assert summary.period_order_count == 2
    assert summary.skipped_empty_rows == 1


def test_period_without_orders_stops(tmp_path: Path) -> None:
    rows: list[RegisterRow] = [
        register_order(1, datetime(2026, 9, 7, 15, 0)),
        None,
        register_order(5, datetime(2026, 9, 15, 0, 0)),
    ]
    fixture = make_preparation_fixture(tmp_path, rows)

    error = stop_message(fixture)

    assert error.message.code == "BRAK_ZLECEN_OKRESU"
    assert error.message.text == (
        "Co się stało: W Rejestrze zleceń nie ma żadnego zlecenia z tego okresu. Nie utworzono żadnego pliku.\n"
        f"Gdzie: plik {REGISTER_NAME}, arkusz ZLECENIA, okres {PERIOD}\n"
        "Co zrobić: Sprawdź, czy wybrano właściwy okres i czy eksport zleceń do Rejestru został "
        "zrobiony i zapisany, i uruchom ponownie."
    )


SKUTECZNE_ACTION = (
    "Co zrobić: Wpisz tak albo nie we wskazanych komórkach Rejestru zleceń. "
    "Zapisz i zamknij Rejestr, i uruchom ponownie."
)


def with_skuteczne(order: dict[str, object], value: object) -> dict[str, object]:
    """Zlecenie z dowolną odczytaną wartością ``SKUTECZNE``, także nie-tekstem."""

    return order | {"SKUTECZNE": value}


def test_period_orders_without_tak_or_nie_in_skuteczne_stop_with_empty_and_other_cells(tmp_path: Path) -> None:
    rows: list[RegisterRow] = [
        register_order(1, datetime(2026, 9, 7, 15, 0), successful="???"),  # poprzedni okres, wiersz 2
        register_order(2, datetime(2026, 9, 8, 8, 0), successful=""),
        with_skuteczne(register_order(3, datetime(2026, 9, 8, 9, 0)), None),
        register_order(4, datetime(2026, 9, 9, 8, 0), successful=" NIE "),
        register_order(5, datetime(2026, 9, 9, 9, 0), successful="   "),
        register_order(6, datetime(2026, 9, 10, 8, 0), successful="Tak"),
        register_order(7, datetime(2026, 9, 11, 8, 0), successful="TAK/NIE\nsprawdzić"),
        register_order(8, datetime(2026, 9, 12, 8, 0), successful="nie"),
    ]
    fixture = make_preparation_fixture(tmp_path, rows)

    error = stop_message(fixture)

    assert error.message.code == "ZLE_SKUTECZNE"
    assert error.message.text == (
        "Co się stało: 4 zlecenia okresu nie mają tak ani nie w SKUTECZNE. Nie utworzono żadnego pliku.\n"
        f"Gdzie: plik {REGISTER_NAME}, arkusz ZLECENIA, puste: L3–L4, L6; inna wartość: L8 („TAK/NIE sprawdzić”)\n"
        f"{SKUTECZNE_ACTION}"
    )


def test_single_empty_skuteczne_points_to_cell_in_column_of_skuteczne_header(tmp_path: Path) -> None:
    fixture = make_preparation_fixture(tmp_path)
    rows: list[RegisterRow] = [
        register_order(1, datetime(2026, 9, 8, 8, 0), successful=""),
        register_order(2, datetime(2026, 9, 9, 8, 0)),
    ]
    write_register(fixture.register_path, in_old_layout(rows), headers=OLD_REGISTER_HEADERS)

    error = stop_message(fixture)

    assert error.message.text == (
        "Co się stało: 1 zlecenie okresu nie ma tak ani nie w SKUTECZNE. Nie utworzono żadnego pliku.\n"
        f"Gdzie: plik {REGISTER_NAME}, arkusz ZLECENIA, puste: J2\n"
        f"{SKUTECZNE_ACTION}"
    )


def test_every_other_skuteczne_value_is_listed_with_exact_cell_content(tmp_path: Path) -> None:
    values: list[object] = ["???", "tsk", "tsk", " nia ", "ne", "trak", "niie", "tak/nie", 1, "TAKK", "x", "-"]
    rows: list[RegisterRow] = [
        with_skuteczne(register_order(number, datetime(2026, 9, 10, 8, number)), value)
        for number, value in enumerate(values, start=1)
    ]
    rows.append(register_order(99, datetime(2026, 9, 11, 8, 0)))
    fixture = make_preparation_fixture(tmp_path, rows)

    error = stop_message(fixture)

    assert error.message.text == (
        "Co się stało: 12 zleceń okresu nie ma tak ani nie w SKUTECZNE. Nie utworzono żadnego pliku.\n"
        f"Gdzie: plik {REGISTER_NAME}, arkusz ZLECENIA, inna wartość: L2 („???”), L3–L4 („tsk”), "
        "L5 („ nia ”), L6 („ne”), L7 („trak”), L8 („niie”), L9 („tak/nie”), L10 („1”), L11 („TAKK”), "
        "L12 („x”), L13 („-”)\n"
        f"{SKUTECZNE_ACTION}"
    )


def test_zlecenie_skuteczne_with_worker_outside_worker_config_stops_with_identifier(tmp_path: Path) -> None:
    rows: list[RegisterRow] = [
        register_order(2, datetime(2026, 9, 8, 8, 0), worker="nowy.wykonawca"),
        register_order(3, datetime(2026, 9, 9, 8, 0), worker=" Kamil.Syntetyczny "),
        register_order(4, datetime(2026, 9, 10, 8, 0), successful="TAK", worker="NOWY.wykonawca"),
        register_order(5, datetime(2026, 9, 11, 8, 0), successful="nie", worker="ktos.nieskuteczny"),
        register_order(6, datetime(2026, 9, 12, 8, 0), worker=""),
    ]
    fixture = make_preparation_fixture(tmp_path, rows)

    error = stop_message(fixture)

    assert error.message.code == "NIEZNANY_WYKONAWCA"
    assert error.message.text == (
        "Co się stało: Zlecenia skuteczne mają wykonawców spoza konfiguracji wykonawców: "
        "nowy.wykonawca, (pusty). Nie utworzono żadnego pliku.\n"
        f"Gdzie: plik {REGISTER_NAME}, arkusz ZLECENIA, nowy.wykonawca: wiersze 2, 4; (pusty): wiersz 6\n"
        "Co zrobić: Dopisz identyfikator wykonawcy do konfiguracji wykonawców (worker_mapping.yaml), "
        "tej samej co w Przebiegu rozliczeń, albo popraw WYKONAWCA we wskazanych wierszach Rejestru zleceń. "
        "Zapisz i zamknij Rejestr, i uruchom ponownie."
    )
    assert "ktos.nieskuteczny" not in error.message.text


def test_worker_outside_worker_config_in_unsuccessful_order_does_not_stop(tmp_path: Path) -> None:
    rows: list[RegisterRow] = [
        register_order(2, datetime(2026, 9, 8, 8, 0)),
        register_order(3, datetime(2026, 9, 9, 8, 0), successful="nie", worker="ktos.nieskuteczny"),
    ]
    fixture = make_preparation_fixture(tmp_path, rows)

    summary = run_preparation(fixture)

    assert (summary.period_order_count, summary.successful_order_count) == (2, 1)
    assert fixture.internal_settlement_path.is_file()
    assert fixture.collective_file_path.is_file()


def test_more_zlecenia_skuteczne_than_wzorzec_zero_rows_warns_and_creates_both_files(tmp_path: Path) -> None:
    successful_count = COLLECTIVE_TEMPLATE_ZERO_ROWS + 2
    rows: list[RegisterRow] = [
        register_order(number, datetime(2026, 9, 14, 8, number)) for number in range(1, successful_count + 1)
    ]
    rows.insert(3, register_order(99, datetime(2026, 9, 14, 9, 0), successful="nie"))
    fixture = make_preparation_fixture(tmp_path, rows)

    summary = run_preparation(fixture)

    assert [warning.text for warning in summary.warnings] == [
        "Co się stało: Zleceń skutecznych: 8, wierszy przygotowanych zerami we Wzorcu pliku zbiorczego: 6. "
        "Dolne wiersze Pliku zbiorczego nie mają zer w kolumnach ilości.\n"
        f"Gdzie: plik {COLLECTIVE_FILE_NAME}, wiersze 24–25\n"
        "Co zrobić: Uzupełnij ilości w tych wierszach ręcznie. Jeśli zdarza się to często, przekaż to "
        "osobie utrzymującej narzędzie, żeby dodała wiersze z zerami we Wzorcu pliku zbiorczego."
    ]
    assert summary.warnings[0].code == "MALO_WIERSZY_Z_ZERAMI"
    assert summary.successful_order_count == successful_count
    assert len(sheet_rows(fixture.internal_settlement_path, min_row=2, max_col=1)) == successful_count + 1
    created = cell_values(fixture.collective_file_path, min_row=18, max_row=26, min_col=1, max_col=14)
    assert [row[5] for row in created] == [f"ZL-SYN-{number:03d}" for number in range(1, 9)] + [None]
    assert [row[13] for row in created] == [0] * COLLECTIVE_TEMPLATE_ZERO_ROWS + [None] * 3


def test_zlecenia_skuteczne_filling_exactly_wzorzec_zero_rows_do_not_warn(tmp_path: Path) -> None:
    rows: list[RegisterRow] = [
        register_order(number, datetime(2026, 9, 14, 8, number))
        for number in range(1, COLLECTIVE_TEMPLATE_ZERO_ROWS + 1)
    ]
    fixture = make_preparation_fixture(tmp_path, rows)

    summary = run_preparation(fixture)

    assert summary.warnings == []


def test_last_order_before_monday_warns_about_incomplete_period_with_orders_per_day(tmp_path: Path) -> None:
    rows: list[RegisterRow] = [
        register_order(1, datetime(2026, 9, 8, 8, 0)),
        None,
        register_order(2, datetime(2026, 9, 12, 23, 59), successful="nie"),
        register_order(3, datetime(2026, 9, 10, 9, 0)),
    ]
    fixture = make_preparation_fixture(tmp_path, rows)

    plan = run_preparation(fixture, dry_run=True)

    assert [warning.text for warning in plan.warnings] == [
        "Co się stało: Ostatnie zlecenie okresu jest z soboty 12.09.2026, a okres kończy się w poniedziałek "
        "14.09.2026. Rejestr zleceń może nie mieć jeszcze wszystkich zleceń tego okresu.\n"
        f"Gdzie: plik {REGISTER_NAME}, arkusz ZLECENIA, okres {PERIOD}\n"
        "Co zrobić: Sprawdź liczby zleceń na dzień. Jeśli brakuje zleceń, przerwij, uzupełnij eksport "
        "w Rejestrze zleceń, zapisz go i uruchom ponownie. Jeśli okres jest pełny, kontynuuj."
    ]
    assert plan.warnings[0].code == "NIEPELNY_OKRES"
    assert plan.orders_per_day == {
        date(2026, 9, 8): 1,
        date(2026, 9, 9): 0,
        date(2026, 9, 10): 1,
        date(2026, 9, 11): 0,
        date(2026, 9, 12): 1,
        date(2026, 9, 13): 0,
        date(2026, 9, 14): 0,
    }
    assert list(fixture.periods_directory.iterdir()) == []


def test_incomplete_period_warning_does_not_block_creating_both_files(tmp_path: Path) -> None:
    fixture = make_preparation_fixture(tmp_path, [register_order(1, datetime(2026, 9, 13, 23, 59))])

    summary = run_preparation(fixture)

    assert [warning.code for warning in summary.warnings] == ["NIEPELNY_OKRES"]
    assert fixture.internal_settlement_path.is_file()
    assert fixture.collective_file_path.is_file()


def test_order_on_monday_anywhere_in_rejestr_means_full_period(tmp_path: Path) -> None:
    rows: list[RegisterRow] = [
        register_order(1, datetime(2026, 9, 14, 0, 0)),
        register_order(2, datetime(2026, 9, 9, 8, 0)),
    ]
    fixture = make_preparation_fixture(tmp_path, rows)

    plan = run_preparation(fixture, dry_run=True)

    assert plan.warnings == []


def test_skipped_empty_rows_are_all_empty_rows_from_first_to_last_period_order(tmp_path: Path) -> None:
    rows: list[RegisterRow] = [
        None,
        register_order(1, datetime(2026, 9, 7, 15, 0)),  # poprzedni okres
        None,
        register_order(2, datetime(2026, 9, 8, 8, 0)),
        None,
        register_order(3, datetime(2026, 9, 15, 8, 0)),  # następny okres, między zleceniami okresu
        None,
        None,
        register_order(4, datetime(2026, 9, 14, 8, 0)),  # spóźnione zlecenie okresu
        None,
    ]
    fixture = make_preparation_fixture(tmp_path, rows)

    summary = run_preparation(fixture, dry_run=True)

    assert summary.period_order_count == 2
    assert summary.skipped_empty_rows == 3


def test_decyzja_taty_about_incomplete_period_comes_after_planowanie_and_before_zapisywanie(tmp_path: Path) -> None:
    fixture = incomplete_period_fixture(tmp_path)
    recorded: list[ProgressEvent] = []
    seen: list[tuple[list[tuple[ProgressPhase, type[ProgressEvent]]], bool]] = []
    tata = Tata(
        before_answer=lambda: seen.append(
            ([(event.phase, type(event)) for event in recorded], fixture.period_directory.exists())
        )
    )

    summary = run_preparation(fixture, observer=recorded.append, conversation=tata)

    [decision] = tata.decisions
    assert (decision.title, decision.message.code, decision.question) == (
        "Ostrzeżenie: niepełny okres",
        "NIEPELNY_OKRES",
        "Kontynuować mimo niepełnego okresu?",
    )
    assert decision.counts_title == "Zlecenia na dzień"
    assert list(decision.counts) == [
        ("wt 08.09.2026", 1),
        ("śr 09.09.2026", 1),
        ("czw 10.09.2026", 0),
        ("pt 11.09.2026", 0),
        ("sob 12.09.2026", 1),
        ("niedz 13.09.2026", 0),
        ("pon 14.09.2026", 0),
    ]
    assert seen == [(STAGES_IN_ORDER[:8], False)]
    assert [(event.phase, type(event)) for event in recorded] == STAGES_IN_ORDER
    assert summary.period_order_count == 3
    assert fixture.collective_file_path.is_file()


def test_full_period_does_not_ask_tata(tmp_path: Path) -> None:
    fixture = make_preparation_fixture(tmp_path)
    tata = Tata(consent=False)

    run_preparation(fixture, conversation=tata)

    assert tata.decisions == []
    assert fixture.collective_file_path.is_file()


def test_stage_times_ended_before_zapisywanie_stay_unchanged_after_measured_zapisywanie(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    fixture = make_preparation_fixture(tmp_path)
    save_s = 0.05
    real_save_workbook = workbook_files.save_workbook
    ended_before_saving: dict[ProgressPhase, int | None] = {}

    def slow_save(workbook: Workbook, target: Path, *, external_links_from: Path | None = None) -> None:
        time.sleep(save_s)
        real_save_workbook(workbook, target, external_links_from=external_links_from)

    def observer(event: ProgressEvent) -> None:
        if isinstance(event, PhaseEnded) and event.phase is not ProgressPhase.SAVING:
            ended_before_saving[event.phase] = event.phase_elapsed_ms

    monkeypatch.setattr(workbook_files, "save_workbook", slow_save)

    summary = run_preparation(fixture, observer=observer)

    assert set(ended_before_saving) == {ProgressPhase.CHECKING, ProgressPhase.READING, ProgressPhase.PLANNING}
    assert summary.phase_durations_ms[ProgressPhase.SAVING] >= save_s * 1000
    assert dict(summary.phase_durations_ms) == {
        **ended_before_saving,
        ProgressPhase.SAVING: summary.phase_durations_ms[ProgressPhase.SAVING],
    }
    with pytest.raises(TypeError):
        summary.phase_durations_ms[ProgressPhase.SAVING] = 0  # type: ignore[index]


def test_refused_incomplete_period_stops_with_hint_of_the_conversation_and_creates_nothing(tmp_path: Path) -> None:
    fixture = incomplete_period_fixture(tmp_path)
    recorded: list[ProgressEvent] = []
    tata = Tata(consent=False, continue_hint="przy decyzji wybierz „Kontynuuj”")

    with pytest.raises(StopError) as raised:
        run_preparation(fixture, observer=recorded.append, conversation=tata)

    assert raised.value.message.text == (
        "Co się stało: Przerwano z powodu niepełnego okresu, zgodnie z odpowiedzią. Nie utworzono żadnego pliku.\n"
        f"Gdzie: okres {PERIOD}\n"
        "Co zrobić: Uzupełnij eksport zleceń w Rejestrze zleceń, zapisz go, zamknij i uruchom ponownie. "
        "Jeśli okres jest pełny, przy decyzji wybierz „Kontynuuj”."
    )
    assert [(event.phase, type(event)) for event in recorded] == STAGES_IN_ORDER[:8]
    assert list(fixture.periods_directory.iterdir()) == []


def test_dry_run_of_incomplete_period_does_not_ask_tata(tmp_path: Path) -> None:
    fixture = incomplete_period_fixture(tmp_path)
    tata = Tata(consent=False)

    plan = run_preparation(fixture, dry_run=True, conversation=tata)

    assert tata.decisions == []
    assert [warning.code for warning in plan.warnings] == ["NIEPELNY_OKRES"]


def test_time_of_decyzja_taty_is_not_counted_in_przygotowanie_time(tmp_path: Path) -> None:
    fixture = incomplete_period_fixture(tmp_path)
    recorded: list[ProgressEvent] = []
    pause_s = 1.5

    summary = run_preparation(
        fixture, observer=recorded.append, conversation=Tata(before_answer=lambda: time.sleep(pause_s))
    )

    assert sum(summary.phase_durations_ms.values()) < pause_s * 1000
    assert max(event.elapsed_ms or 0 for event in recorded) < pause_s * 1000
