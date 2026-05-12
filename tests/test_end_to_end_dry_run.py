"""End-to-end test — full Phase 1 pipeline w test-mode.

Verifies fixtures/sample_emails/ → daily_workflow.main() → vault populated.
"""
from __future__ import annotations

import os
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
BUFFER_SHEET = "MAXIMISEART_DAILY_APPEND"


@pytest.fixture
def isolated_runtime(tmp_path: Path, monkeypatch) -> Path:
    """Spawn pipeline w tmp runtime — vault/state/logs w tmp_path."""
    runtime = tmp_path / "runtime"
    runtime.mkdir()
    monkeypatch.setenv("DAD_RUNTIME_PATH", str(runtime))
    monkeypatch.setenv("DAD_VAULT_PATH", str(runtime / "vault"))
    monkeypatch.setenv("DAD_REPORTS_DIR", str(runtime / "vault" / "reports"))
    monkeypatch.setenv("DAD_LOGS_DIR", str(runtime / "logs"))
    monkeypatch.setenv("DAD_USE_REAL_API", "false")
    return runtime


def test_test_mode_populates_vault_from_fixtures(isolated_runtime: Path) -> None:
    """Run `daily_workflow.py --test-mode` w subprocess z isolated DAD_RUNTIME_PATH."""
    excel_path = _create_test_workbook(isolated_runtime / "test.xlsx")
    env = os.environ.copy()
    env["DAD_RUNTIME_PATH"] = str(isolated_runtime)
    env["DAD_VAULT_PATH"] = str(isolated_runtime / "vault")
    env["DAD_REPORTS_DIR"] = str(isolated_runtime / "vault" / "reports")
    env["DAD_LOGS_DIR"] = str(isolated_runtime / "logs")
    env["DAD_EXCEL_PATH"] = str(excel_path)
    env["DAD_USE_REAL_API"] = "false"
    env["PYTHONIOENCODING"] = "utf-8"

    result = subprocess.run(
        [sys.executable, "scripts/daily_workflow.py",
         "--test-mode", "--date", "2026-04-24"],
        cwd=str(PROJECT_ROOT),
        env=env,
        capture_output=True,
        text=True,
        encoding="utf-8",
        timeout=30,
    )

    assert result.returncode == 0, f"Pipeline failed:\nSTDOUT:\n{result.stdout}\nSTDERR:\n{result.stderr}"

    vault = isolated_runtime / "vault"
    assert vault.exists()

    pracownicy = vault / "pracownicy"
    assert pracownicy.exists()

    # Fixture 01 + 02 both from jan.kowalski@example.com → przyklad-jan-kowalski
    jan_folder = pracownicy / "przyklad-jan-kowalski"
    assert jan_folder.exists(), f"Expected jan folder, got: {list(pracownicy.iterdir())}"

    jan_daily = jan_folder / "2026-04-24.md"
    assert jan_daily.exists()

    content = jan_daily.read_text(encoding="utf-8")
    assert "18f2a3b4c5d6e7f8" in content
    assert "18f2a3b4c5d6e7f9" in content
    assert "emails_count: 2" in content

    # Fixture 03 unknown sender — NIE tworzy folderu
    unknown_folder_candidates = [p.name for p in pracownicy.iterdir()]
    assert "nieznany-sender" not in unknown_folder_candidates

    wb = load_workbook(excel_path, data_only=False)
    assert BUFFER_SHEET in wb.sheetnames
    ws = wb[BUFFER_SHEET]
    assert ws.max_row == 4
    assert [ws.cell(row=row, column=8).value for row in range(2, 5)] == [
        "18f2a3b4c5d6e7f8",
        "18f2a3b4c5d6e7f8",
        "18f2a3b4c5d6e7f9",
    ]
    assert wb["Istniejacy"]["A1"].value == "Nie dotykac"
    assert list((isolated_runtime / "state" / "excel-snapshots").glob("test-*.xlsx"))


def test_dry_run_does_not_write_vault(isolated_runtime: Path) -> None:
    excel_path = _create_test_workbook(isolated_runtime / "test.xlsx")
    env = os.environ.copy()
    env["DAD_RUNTIME_PATH"] = str(isolated_runtime)
    env["DAD_VAULT_PATH"] = str(isolated_runtime / "vault")
    env["DAD_REPORTS_DIR"] = str(isolated_runtime / "vault" / "reports")
    env["DAD_LOGS_DIR"] = str(isolated_runtime / "logs")
    env["DAD_EXCEL_PATH"] = str(excel_path)
    env["PYTHONIOENCODING"] = "utf-8"

    result = subprocess.run(
        [sys.executable, "scripts/daily_workflow.py",
         "--test-mode", "--dry-run", "--date", "2026-04-24"],
        cwd=str(PROJECT_ROOT),
        env=env,
        capture_output=True,
        text=True,
        encoding="utf-8",
        timeout=30,
    )

    assert result.returncode == 0

    pracownicy = isolated_runtime / "vault" / "pracownicy"
    if pracownicy.exists():
        # Folder może istnieć (ensure_runtime_dirs), ale żadne pliki pracowników
        employee_folders = [p for p in pracownicy.iterdir() if p.is_dir()]
        assert employee_folders == [], f"Dry-run created folders: {employee_folders}"

    wb = load_workbook(excel_path, data_only=False)
    assert BUFFER_SHEET not in wb.sheetnames


def test_idempotent_double_run(isolated_runtime: Path) -> None:
    """Run twice, expect total dedup — no duplicate entries."""
    env = os.environ.copy()
    env["DAD_RUNTIME_PATH"] = str(isolated_runtime)
    env["DAD_VAULT_PATH"] = str(isolated_runtime / "vault")
    env["DAD_REPORTS_DIR"] = str(isolated_runtime / "vault" / "reports")
    env["DAD_LOGS_DIR"] = str(isolated_runtime / "logs")
    env["PYTHONIOENCODING"] = "utf-8"

    for _ in range(2):
        result = subprocess.run(
            [sys.executable, "scripts/daily_workflow.py",
             "--test-mode", "--date", "2026-04-24"],
            cwd=str(PROJECT_ROOT),
            env=env,
            capture_output=True,
            text=True,
            encoding="utf-8",
            timeout=30,
        )
        assert result.returncode == 0

    jan_daily = isolated_runtime / "vault" / "pracownicy" / "przyklad-jan-kowalski" / "2026-04-24.md"
    content = jan_daily.read_text(encoding="utf-8")
    # markers are unique per entry (deduped); raw ID appears in marker + meta = 2×
    assert content.count("<!-- email-id: 18f2a3b4c5d6e7f8 -->") == 1
    assert content.count("<!-- email-id: 18f2a3b4c5d6e7f9 -->") == 1
    assert "emails_count: 2" in content


def _create_test_workbook(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Istniejacy"
    ws["A1"] = "Nie dotykac"
    wb.save(path)
    return path
