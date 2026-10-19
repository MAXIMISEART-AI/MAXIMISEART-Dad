from __future__ import annotations

from datetime import datetime
import os
from pathlib import Path
import shutil
import subprocess

from openpyxl import load_workbook
import pytest

from tests.collective_preparation_fixtures import PERIOD as PREPARATION_PERIOD
from tests.collective_preparation_fixtures import PreparationFixture, make_preparation_fixture, register_order
from tests.test_settlement_engine import PERIOD, active_worksheet, make_fixture, write_template


REPOSITORY_ROOT = Path(__file__).resolve().parents[1]
SETTLEMENT_LAUNCHER = "Utwórz rozliczenia.cmd"
PREPARATION_LAUNCHER = "Przygotuj plik zbiorczy.cmd"


def _copy_launcher_runtime(destination: Path, launcher_name: str = SETTLEMENT_LAUNCHER) -> Path:
    """Disposable runtime with both launchers, both CLIs, the engine, and ``config``; returns ``launcher_name``."""

    destination.mkdir()
    for launcher in (SETTLEMENT_LAUNCHER, PREPARATION_LAUNCHER):
        shutil.copy2(REPOSITORY_ROOT / launcher, destination)
    shutil.copy2(REPOSITORY_ROOT / "run.py", destination)
    shutil.copy2(REPOSITORY_ROOT / "run_preparation.py", destination)
    shutil.copytree(
        REPOSITORY_ROOT / "src",
        destination / "src",
        ignore=shutil.ignore_patterns("__pycache__"),
    )
    shutil.copytree(REPOSITORY_ROOT / "config", destination / "config")
    return destination / launcher_name


def _configure_launcher(launcher_path: Path, config_path: Path) -> None:
    config_directory = launcher_path.parent / "config"
    shutil.copy2(config_path, config_directory / "worker_mapping.yaml")
    shutil.copy2(config_path.parent / "Wzorzec rozliczenia pracownika.xlsx", config_directory / "Wzorzec rozliczenia pracownika.xlsx")


def _input_value(path: Path) -> object:
    workbook = load_workbook(path, data_only=False)
    try:
        return active_worksheet(workbook)["A18"].value
    finally:
        workbook.close()


@pytest.mark.skipif(os.name != "nt", reason="The CMD launcher is Windows-only.")
def test_cmd_launcher_runs_przebieg_rozliczen(tmp_path: Path) -> None:
    source_path, target_directory, _ = make_fixture(tmp_path / "fixture")
    launcher_path = _copy_launcher_runtime(tmp_path / "launcher")
    _configure_launcher(launcher_path, tmp_path / "fixture" / "worker_mapping.yaml")

    result = _run_launcher(launcher_path, f'"{source_path}"')

    assert result.returncode == 0, result.stdout + result.stderr
    assert "Status końcowy: OK" in result.stdout
    assert {path.name for path in target_directory.glob("*.xlsx")} == {
        f"Rozliczenie {PERIOD} - {worker}.xlsx"
        for worker in ("Adrian Maciejewski", "Darek Nowak", "Kamil Frontczak")
    }
    assert _input_value(target_directory / f"Rozliczenie {PERIOD} - Adrian Maciejewski.xlsx") == "POZNAŃ"
    assert _input_value(target_directory / f"Rozliczenie {PERIOD} - Darek Nowak.xlsx") == "POZNAŃ"
    assert _input_value(target_directory / f"Rozliczenie {PERIOD} - Kamil Frontczak.xlsx") is None


def _run_preparation_launcher(tmp_path: Path, fixture: PreparationFixture, answers: str) -> subprocess.CompletedProcess[str]:
    launcher_path = _copy_launcher_runtime(tmp_path / "launcher", PREPARATION_LAUNCHER)
    config_directory = launcher_path.parent / "config"
    for path in (fixture.config_path, fixture.internal_template_path, fixture.collective_template_path):
        shutil.copy2(path, config_directory / path.name)
    arguments = (
        f'--register "{fixture.register_path}" '
        f'--periods-directory "{fixture.periods_directory}" --period {PREPARATION_PERIOD}'
    )
    return _run_launcher(launcher_path, arguments, answers)


def _run_launcher(launcher_path: Path, arguments: str, answers: str = "") -> subprocess.CompletedProcess[str]:
    environment = os.environ.copy()
    environment["PYTHONUTF8"] = "1"
    command_line = f'cmd.exe /d /c ""{launcher_path}" {arguments}"'
    return subprocess.run(
        command_line,
        cwd=launcher_path.parent,
        input=answers,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        env=environment,
        timeout=120,
    )


@pytest.mark.skipif(os.name != "nt", reason="The CMD launcher is Windows-only.")
def test_cmd_launcher_runs_przygotowanie_pliku_zbiorczego(tmp_path: Path) -> None:
    fixture = make_preparation_fixture(tmp_path)

    result = _run_preparation_launcher(tmp_path, fixture, "")

    assert result.returncode == 0, result.stdout + result.stderr
    assert "Status końcowy: OK" in result.stdout
    assert "Zlecenia skuteczne: 3" in result.stdout
    assert fixture.internal_settlement_path.is_file()
    workbook = load_workbook(fixture.collective_file_path, data_only=False)
    try:
        assert active_worksheet(workbook)["F18"].value == "ZL-SYN-002"
    finally:
        workbook.close()


@pytest.mark.skipif(os.name != "nt", reason="The CMD launcher is Windows-only.")
def test_cmd_launcher_passes_exit_code_1_after_refusing_incomplete_period(tmp_path: Path) -> None:
    fixture = make_preparation_fixture(tmp_path, [register_order(1, datetime(2026, 9, 12, 12, 0))])

    result = _run_preparation_launcher(tmp_path, fixture, "n\n")

    assert result.returncode == 1, result.stdout + result.stderr
    assert "Ostrzeżenie: niepełny okres" in result.stdout
    assert "Status końcowy: ZATRZYMANO" in result.stdout
    assert list(fixture.periods_directory.iterdir()) == []


@pytest.mark.skipif(os.name != "nt", reason="The CMD launcher is Windows-only.")
def test_cmd_launchers_run_przygotowanie_then_przebieg_rozliczen_on_its_plik_zbiorczy(tmp_path: Path) -> None:
    fixture = make_preparation_fixture(tmp_path)
    preparation = _run_preparation_launcher(tmp_path, fixture, "")
    assert preparation.returncode == 0, preparation.stdout + preparation.stderr
    runtime = tmp_path / "launcher"
    write_template(runtime / "config" / "Wzorzec rozliczenia pracownika.xlsx")

    settlement = _run_launcher(runtime / SETTLEMENT_LAUNCHER, f'"{fixture.collective_file_path}"')

    assert settlement.returncode == 0, settlement.stdout + settlement.stderr
    assert "Status końcowy: OK" in settlement.stdout
    target_directory = fixture.period_directory / f"Rozliczenie pracowników {PREPARATION_PERIOD}"
    assert {path.name for path in target_directory.iterdir()} == {
        f"Rozliczenie {PREPARATION_PERIOD} - {worker}.xlsx"
        for worker in ("Adrian Maciejewski", "Darek Nowak", "Kamil Frontczak")
    }
    workbook = load_workbook(target_directory / f"Rozliczenie {PREPARATION_PERIOD} - Adrian Maciejewski.xlsx")
    try:
        assert active_worksheet(workbook)["F18"].value == "ZL-SYN-002"
    finally:
        workbook.close()
