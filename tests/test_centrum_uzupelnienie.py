"""Uzupełnienie pliku zbiorczego w Centrum rozliczeń (#415, #417, ADR 0016), przez HTTP API jak przeglądarka.

Każdy test przygotowuje Plik zbiorczy z syntetycznego Rejestru zleceń w Centrum, potem zmienia go jak tata w
Excelu (openpyxl) i pyta Centrum o Ilości do uzupełnienia okresu.
"""

from __future__ import annotations

from collections.abc import Callable
from datetime import datetime
from pathlib import Path
import sys
from typing import Any

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from tests.fixtures.centrum import FakeOneDrive, view_timing
from tests.fixtures.centrum_processes import PreparingCentrum, Surroundings, preparing_centrum, run_to_end
from tests.fixtures.preparation import COLLECTIVE_FILE_NAME, PERIOD, register_order

TUE = datetime(2026, 9, 8, 10, 0)
MON = datetime(2026, 9, 14, 9, 0)
URL = f"/api/periods/{PERIOD}/quantities-to-fill"
LOG_EVENT = "uzupelnienie-odczyt-nieudany"

# Zlecenia skuteczne okresu po kolei, więc w Pliku zbiorczym od wiersza 18: wiersz → URZĄDZENIA.
DEVICES = {
    18: "nokia zte",  # Reguły ilości znają każdy składnik
    19: "przeniesienie",  # nieznane URZĄDZENIA: wiersz bez reguły
    20: "prewencja",  # wiersz z opisu
    21: "",  # puste URZĄDZENIA: wiersz z opisu
    22: "router x9",  # nieznane URZĄDZENIA: wiersz bez reguły
}
DESCRIPTION = ("Y", "AB", "AD", "AJ", "AK", "AL", "AM")


def _prepared(tmp_path: Path, devices: dict[int, str] = DEVICES, **options: Any) -> PreparingCentrum:
    # Ostatnie zlecenie w poniedziałek: okres jest pełny, więc Przygotowanie nie czeka na Decyzję taty.
    starts = [TUE.replace(hour=8 + index) for index in range(len(devices) - 1)] + [MON]
    rows = [
        {**register_order(index, started), "URZĄDZENIA": text, "UMOWA": f"u UM-SYN-{index:03d}"}
        for index, (text, started) in enumerate(zip(devices.values(), starts, strict=True), start=1)
    ]
    centrum = preparing_centrum(tmp_path, list(rows), **options)
    assert run_to_end(centrum, dry_run=False)["result"]["result"] == "WYMAGA_SPRAWDZENIA"
    return centrum


def _quantities(centrum: PreparingCentrum) -> dict[str, Any]:
    response = centrum.client.get(URL)
    assert response.status_code == 200, response.text
    found: dict[str, Any] = response.json()
    return found


def _kinds(view: dict[str, Any]) -> dict[str, list[tuple[int, bool]]]:
    return {kind["kind"]: [(row["row"], row["state"] == "done") for row in kind["rows"]] for kind in view["kinds"]}


def _kind_done(view: dict[str, Any]) -> dict[str, bool]:
    return {kind["kind"]: kind["done"] for kind in view["kinds"]}


def _description(view: dict[str, Any]) -> dict[str, int]:
    return {column["column"]: column["rows"] for column in view["description_columns"]}


def _edit(path: Path, change: Callable[[Worksheet, dict[str, str]], None]) -> None:
    """Zmiana Pliku zbiorczego jak tata w Excelu; ``change`` dostaje arkusz i nagłówki wiersza 17 → litera kolumny."""

    workbook = load_workbook(path)
    sheet = workbook.active
    assert isinstance(sheet, Worksheet)
    header = {str(cell.value): get_column_letter(cell.column) for cell in sheet[17] if cell.value is not None}
    change(sheet, header)
    workbook.save(path)


def _write(path: Path, cells: dict[str, int | float | str | None]) -> None:
    def change(sheet: Worksheet, _header: dict[str, str]) -> None:
        for cell, value in cells.items():
            sheet[cell] = value

    _edit(path, change)


def _log_lines(centrum: PreparingCentrum) -> list[str]:
    if not centrum.log_path.exists():
        return []
    return [line for line in centrum.log_path.read_text(encoding="utf-8").splitlines() if LOG_EVENT in line]


def _collective_file(centrum: PreparingCentrum) -> Path:
    return centrum.fixture.collective_file_path


def test_po_przygotowaniu_wiersze_bez_reguly_i_z_opisu_czekaja_na_tate(tmp_path: Path) -> None:
    view = _quantities(_prepared(tmp_path))

    assert view["state"] == "ready"
    assert _kinds(view) == {
        "without_rule": [(19, False), (22, False)],
        "from_description": [(20, False), (21, False)],
    }
    assert _kind_done(view) == {"without_rule": False, "from_description": False}
    assert _description(view) == dict.fromkeys(DESCRIPTION, 0)


def test_ilosc_taty_poza_v_t_i_n_robi_wiersz_bez_reguly_a_dowolna_w_n_ap_wiersz_z_opisu(tmp_path: Path) -> None:
    centrum = _prepared(tmp_path)
    _write(_collective_file(centrum), {"O19": 1, "AP21": 2})

    view = _quantities(centrum)

    assert _kinds(view) == {
        "without_rule": [(19, True), (22, False)],
        "from_description": [(20, False), (21, True)],
    }


def test_v_t_i_n_nie_robia_wiersza_bez_reguly_ale_robia_wiersz_z_opisu(tmp_path: Path) -> None:
    centrum = _prepared(tmp_path)
    _write(_collective_file(centrum), {"N19": 45, "T22": 1, "V22": 1, "N20": 45})

    view = _quantities(centrum)

    assert _kinds(view)["without_rule"] == [(19, False), (22, False)]
    assert _kinds(view)["from_description"] == [(20, True), (21, False)]


def test_zero_i_pusta_komorka_to_to_samo(tmp_path: Path) -> None:
    centrum = _prepared(tmp_path)
    _write(_collective_file(centrum), {"O19": 0, "Y20": "0", "AB21": None, "AE22": 0.0})

    view = _quantities(centrum)

    assert _kinds(view) == {
        "without_rule": [(19, False), (22, False)],
        "from_description": [(20, False), (21, False)],
    }
    assert _description(view) == dict.fromkeys(DESCRIPTION, 0)


def test_rodzaj_jest_zrobiony_gdy_kazdy_jego_wiersz_jest_zrobiony(tmp_path: Path) -> None:
    centrum = _prepared(tmp_path)
    _write(_collective_file(centrum), {"O19": 1, "Z22": "dwa", "Y20": 1})

    view = _quantities(centrum)

    assert _kind_done(view) == {"without_rule": True, "from_description": False}


def test_rodzaj_bez_wierszy_jest_zrobiony(tmp_path: Path) -> None:
    view = _quantities(_prepared(tmp_path, {18: "nokia zte", 19: "gniazdo"}))

    assert _kinds(view) == {"without_rule": [], "from_description": []}
    assert _kind_done(view) == {"without_rule": True, "from_description": True}


def test_kolumny_z_opisu_licza_wiersze_z_wpisem_taty_w_kazdej_kolumnie(tmp_path: Path) -> None:
    centrum = _prepared(tmp_path)
    _write(_collective_file(centrum), {"Y18": 1, "Y19": 2, "Y20": 0, "AB21": 1, "AM22": "dwa", "AK18": 1})

    view = _quantities(centrum)

    assert _description(view) == {"Y": 2, "AB": 1, "AD": 0, "AJ": 0, "AK": 1, "AL": 0, "AM": 1}


def test_ak_wyznaczone_przez_regule_nie_liczy_sie_jak_wpis_taty(tmp_path: Path) -> None:
    centrum = _prepared(tmp_path, {18: "technicolor", 19: "przeniesienie"})
    _write(_collective_file(centrum), {"AK19": 1})

    view = _quantities(centrum)

    assert _description(view)["AK"] == 1


def test_usuniecie_wiersza_w_excelu_przenumerowuje_wiersze(tmp_path: Path) -> None:
    centrum = _prepared(tmp_path)
    _write(_collective_file(centrum), {"O22": 1})
    _edit(_collective_file(centrum), lambda sheet, _header: sheet.delete_rows(19))

    view = _quantities(centrum)

    assert _kinds(view) == {
        "without_rule": [(21, True)],
        "from_description": [(19, False), (20, False)],
    }


def test_wyczyszczony_wiersz_nie_przesuwa_numerow_pozostalych(tmp_path: Path) -> None:
    centrum = _prepared(tmp_path)

    def change(sheet: Worksheet, _header: dict[str, str]) -> None:
        for cell in sheet[19]:
            cell.value = None

    _edit(_collective_file(centrum), change)

    view = _quantities(centrum)

    assert _kinds(view) == {
        "without_rule": [(22, False)],
        "from_description": [(20, False), (21, False)],
    }


def test_zmiana_urzadzen_przenosi_wiersz_miedzy_rodzajami(tmp_path: Path) -> None:
    centrum = _prepared(tmp_path)

    def change(sheet: Worksheet, header: dict[str, str]) -> None:
        sheet[f"{header['URZĄDZENIA']}19"] = "prewencja"
        sheet[f"{header['URZĄDZENIA']}20"] = "zyxel"
        sheet[f"{header['URZĄDZENIA']}21"] = "pralka"

    _edit(_collective_file(centrum), change)

    view = _quantities(centrum)

    # V i T, które Przygotowanie wpisało w wierszu bez reguły, zostają w nim, więc wiersz z opisu jest zrobiony.
    assert _kinds(view) == {
        "without_rule": [(21, False), (22, False)],
        "from_description": [(19, True)],
    }


def test_wiersz_bez_daty_w_data_nie_ma_rodzaju_a_jego_kolumny_z_opisu_sie_licza(tmp_path: Path) -> None:
    centrum = _prepared(tmp_path)

    def change(sheet: Worksheet, header: dict[str, str]) -> None:
        sheet[f"{header['DATA']}19"] = "wtorek"
        sheet["Y19"] = 1

    _edit(_collective_file(centrum), change)

    view = _quantities(centrum)

    assert _kinds(view)["without_rule"] == [(22, False)]
    assert _description(view)["Y"] == 1


def test_plik_otwarty_w_excelu_jest_czytany(tmp_path: Path) -> None:
    centrum = _prepared(tmp_path)
    path = _collective_file(centrum)
    _write(path, {"O19": 1})
    path.with_name(f"~${path.name}").write_bytes(b"blokada Excela")

    view = _quantities(centrum)

    assert _kinds(view)["without_rule"] == [(19, True), (22, False)]
    assert _log_lines(centrum) == []


def test_odpowiedz_i_dziennik_nie_niosa_tresci_wierszy(tmp_path: Path) -> None:
    centrum = _prepared(tmp_path)
    shown = centrum.client.get(URL).text
    path = _collective_file(centrum)
    path.write_bytes(b"")
    centrum.client.get(URL)
    logged = "\n".join(_log_lines(centrum))

    for text in (shown, logged):
        assert "SYN-" not in text
        assert "przeniesienie" not in text
        assert "prewencja" not in text
    assert logged != ""


def _problem(centrum: PreparingCentrum) -> dict[str, str]:
    view = _quantities(centrum)
    assert view["state"] == "problem"
    message: dict[str, str] = view["message"]
    return message


def test_uszkodzony_plik_zbiorczy_daje_komunikat_i_wpis_w_dzienniku(tmp_path: Path) -> None:
    centrum = _prepared(tmp_path)
    path = _collective_file(centrum)
    path.write_bytes(b"to nie jest skoroszyt")
    path.with_name(f"~${path.name}").write_bytes(b"blokada Excela")

    message = _problem(centrum)

    assert message["what"] == "Plik zbiorczy jest uszkodzony albo nie jest skoroszytem Excela."
    assert message["where"] == f"plik {COLLECTIVE_FILE_NAME}, folder {PERIOD}"
    [line] = _log_lines(centrum)
    assert f"okres={PERIOD} przyczyna=USZKODZONY_PLIK_ZBIORCZY excel=tak typ=BadZipFile" in line


@pytest.mark.skipif(sys.platform == "win32", reason="chmod w Windows nie odbiera prawa odczytu")
def test_niedostepny_plik_zbiorczy_nie_jest_nazwany_uszkodzonym(tmp_path: Path) -> None:
    centrum = _prepared(tmp_path)
    path = _collective_file(centrum)
    path.chmod(0)
    try:
        message = _problem(centrum)
    finally:
        path.chmod(0o644)

    assert message["what"].startswith("Nie można teraz odczytać Pliku zbiorczego.")
    assert message["where"] == f"plik {COLLECTIVE_FILE_NAME}, folder {PERIOD}"
    assert "uszkodzon" not in message["what"] + message["action"]
    [line] = _log_lines(centrum)
    assert "przyczyna=NIEDOSTEPNY_PLIK_ZBIORCZY excel=nie typ=PermissionError" in line


@pytest.mark.skipif(sys.platform == "win32", reason="chmod w Windows nie odbiera prawa odczytu")
def test_folder_okresu_bez_odczytu_daje_komunikat_bez_stanu_blokady(tmp_path: Path) -> None:
    centrum = _prepared(tmp_path)
    folder = _collective_file(centrum).parent
    folder.chmod(0)
    try:
        message = _problem(centrum)
    finally:
        folder.chmod(0o755)

    assert message["what"].startswith("Nie można teraz odczytać Pliku zbiorczego.")
    [line] = _log_lines(centrum)
    assert "excel=nieznany" in line


def test_plik_tylko_w_chmurze_daje_komunikat_onedrive(tmp_path: Path) -> None:
    onedrive = FakeOneDrive()
    centrum = _prepared(tmp_path, surroundings=Surroundings(onedrive=onedrive))
    path = _collective_file(centrum)
    onedrive.online_only.add(path)
    path.write_bytes(b"")

    message = _problem(centrum)

    assert message["what"] == "Plik zbiorczy jest tylko w OneDrive w chmurze, bez kopii na tym komputerze."
    assert message["where"] == f"plik {COLLECTIVE_FILE_NAME}, folder {PERIOD}"


def test_brak_pliku_zbiorczego_daje_komunikat(tmp_path: Path) -> None:
    centrum = _prepared(tmp_path)
    _collective_file(centrum).unlink()

    message = _problem(centrum)

    assert message["what"] == "Nie ma Pliku zbiorczego tego okresu."
    [line] = _log_lines(centrum)
    assert "przyczyna=BRAK_PLIKU_ZBIORCZEGO excel=nie typ=-" in line


def test_brak_naglowka_urzadzen_nazywa_naglowek(tmp_path: Path) -> None:
    centrum = _prepared(tmp_path)

    def change(sheet: Worksheet, header: dict[str, str]) -> None:
        sheet[f"{header['URZĄDZENIA']}17"] = None
        sheet[f"{header['UMOWA']}17"] = None

    _edit(_collective_file(centrum), change)

    message = _problem(centrum)

    assert message["what"] == "Brakuje nagłówka w Pliku zbiorczym: UMOWA i URZĄDZENIA w wierszu 17."
    [line] = _log_lines(centrum)
    assert "przyczyna=BRAK_NAGLOWKA_PLIKU_ZBIORCZEGO" in line


def test_brak_naglowka_wykonawca_nazywa_komorke(tmp_path: Path) -> None:
    centrum = _prepared(tmp_path)
    _write(_collective_file(centrum), {"H17": None})

    message = _problem(centrum)

    assert message["what"] == "Brakuje nagłówka w Pliku zbiorczym: WYKONAWCA w komórce H17."


def test_okres_bez_folderu_okresu_daje_komunikat_o_braku_pliku_zbiorczego(tmp_path: Path) -> None:
    centrum = _prepared(tmp_path)

    view = centrum.client.get("/api/periods/15_21_09_2026/quantities-to-fill").json()

    assert view["state"] == "problem"
    assert view["message"]["what"] == "Nie ma Pliku zbiorczego tego okresu."


def test_zly_token_okresu_i_brak_folderu_okresow_sa_odmowa(tmp_path: Path) -> None:
    centrum = _prepared(tmp_path)

    wrong = centrum.client.get("/api/periods/nie-okres/quantities-to-fill")
    centrum.fixture.periods_directory.rename(tmp_path / "przeniesiony")
    without_periods = centrum.client.get(URL)

    assert (wrong.status_code, wrong.json()["message"]) == (404, "Nie ma takiego okresu.")
    assert (without_periods.status_code, without_periods.json()["message"]) == (409, "Najpierw wskaż folder okresów.")


def test_pomiar_widoku_strony_uzupelnienia_trafia_do_dziennika(tmp_path: Path) -> None:
    centrum = _prepared(tmp_path)
    timing = view_timing(
        kind="navigation",
        page="uzupelnienie",
        fcp_ms=None,
        requests=[{"operation": "getQuantitiesToFill", "browser_ms": 310, "server_ms": 280}],
    )

    response = centrum.client.post("/api/view-timings", json=timing)

    assert response.status_code == 204
    assert "strona=uzupelnienie" in centrum.log_path.read_text(encoding="utf-8")
    assert "zapytania=getQuantitiesToFill:310/280" in centrum.log_path.read_text(encoding="utf-8")
