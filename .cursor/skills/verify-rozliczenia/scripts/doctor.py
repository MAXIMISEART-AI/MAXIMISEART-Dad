"""Read-only build and fixture check for the rozliczenia verification skill."""

from __future__ import annotations

import argparse
from pathlib import Path
import sys
import tomllib

from openpyxl import load_workbook


SCRIPT = Path(__file__).resolve()
REPO_ROOT = SCRIPT.parents[4]
sys.path.insert(0, str(REPO_ROOT / "src"))

from rozliczenia.engine import load_worker_mapping, period_from_source  # noqa: E402


def _check(root: Path, expected_version: str) -> tuple[str, int]:
    root = root.expanduser().resolve()
    if not root.is_dir():
        raise ValueError(f"missing fixture root: {root}")
    with (REPO_ROOT / "pyproject.toml").open("rb") as stream:
        version = tomllib.load(stream)["project"]["version"]
    if version != expected_version:
        raise ValueError(f"expected version {expected_version}, found {version}")
    if sys.version_info < (3, 11):
        raise ValueError("Python 3.11 or newer is required")

    source = root / "08_14_09_2026" / "Rozliczenie 08_14_09_2026 - zbiorcze.xlsx"
    period = period_from_source(source)
    target_directory = source.parent / f"Rozliczenie pracowników {period}"
    mapping = load_worker_mapping(root / "worker_mapping.yaml")
    if len(mapping) != 3:
        raise ValueError("synthetic mapping must contain three workers")

    workbook = load_workbook(source, read_only=True, data_only=False)
    try:
        if workbook.active["H17"].value != "WYKONAWCA":
            raise ValueError("source H17 is not WYKONAWCA")
    finally:
        workbook.close()

    templates = sorted(target_directory.glob("*.xlsx"))
    if len(templates) != 4:
        raise ValueError(f"expected four synthetic workbooks, found {len(templates)}")
    if list(target_directory.glob("~$*.xlsx")):
        raise ValueError("an Excel lock file is present")
    for target in templates:
        workbook = load_workbook(target, read_only=True, data_only=False)
        try:
            if workbook.active["H17"].value != "WYKONAWCA":
                raise ValueError(f"invalid worker header in {target.name}")
        finally:
            workbook.close()
    return period, len(templates)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", required=True, type=Path)
    parser.add_argument("--expected-version", default="0.1.0")
    args = parser.parse_args(argv)
    try:
        period, template_count = _check(args.root, args.expected_version)
    except (OSError, ValueError, KeyError, TypeError) as exc:
        print(f"DOCTOR FAIL: {exc}", file=sys.stderr)
        return 1
    print(
        f"DOCTOR OK | app=MAXIMISEART-Dad | version={args.expected_version} | "
        f"python={sys.version_info.major}.{sys.version_info.minor}.{sys.version_info.micro} | "
        f"period={period} | templates={template_count}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
