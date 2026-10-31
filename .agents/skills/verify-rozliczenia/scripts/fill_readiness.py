"""Gotowość: fill a Plik zbiorczy like Dad does by hand before the Przebieg rozliczeń.

In some Wiersze danych the harness types quantities into ``N:AP`` and the
customer charge and remarks into ``AQ:AT``. Every other cell, every formula, and
every ``xl/externalLinks/*`` part (the link to the Słownik) stays as it was.
Prints counts only, never a cell value.
"""

from __future__ import annotations

import hashlib
import json
from pathlib import Path
import sys
import zipfile

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

sys.path.insert(0, str(Path(__file__).resolve().parents[4] / "tools/lib"))
sys.path.insert(0, str(Path(__file__).resolve().parents[4] / "src"))

from rozliczenia.workbook_save import save_workbook
from skrypt_skilla import Parser, require_run_root


PERIOD = "08_14_09_2026"
HEADER_ROW = 17
FIRST_DATA_ROW = 18
QUANTITY_COLUMNS = range(14, 43)  # N:AP
CHARGE_DESCRIPTION, CHARGE_AMOUNT, REMARKS_AMOUNT, REMARKS = 43, 44, 45, 46  # AQ, AR, AS, AT
# Adrian twice and Darek once; Darek's gniazdo order and Kamil's two keep what Przygotowanie pliku zbiorczego wrote.
FILLED_ORDERS = ("PREP-SYN-ZL-2", "PREP-SYN-ZL-7", "PREP-SYN-ZL-6")


def _sheet(workbook: Workbook) -> Worksheet:
    sheet = workbook.active
    if not isinstance(sheet, Worksheet):
        raise ValueError("Plik zbiorczy does not contain an active worksheet")
    return sheet


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _external_link_parts(path: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(path) as archive:
        return {name: archive.read(name) for name in archive.namelist() if name.startswith("xl/externalLinks/")}


def _order_rows(sheet: Worksheet) -> dict[object, int]:
    """Row of every NUMER ZL. value in the Wiersze danych, found by its header like the processes do."""

    columns = [
        cell.column
        for cell in sheet[HEADER_ROW]
        if isinstance(cell.value, str) and cell.value.strip().casefold() == "numer zl."
    ]
    if len(columns) != 1:
        raise ValueError(f"Plik zbiorczy has {len(columns)} NUMER ZL. headers in row {HEADER_ROW}")
    column = columns[0]
    return {
        row[0]: number
        for number, row in enumerate(
            sheet.iter_rows(min_row=FIRST_DATA_ROW, min_col=column, max_col=column, values_only=True),
            start=FIRST_DATA_ROW,
        )
        if row[0] is not None
    }


def _readiness_values(index: int, order: str) -> dict[int, int | float | str]:
    """Dad's entries for one Wiersz danych: distinct non-zero quantities, a charge, and remarks."""

    suffix = order.removeprefix("PREP-SYN-ZL-")
    values: dict[int, int | float | str] = {column: 100 * index + column - 13 for column in QUANTITY_COLUMNS}
    values[CHARGE_DESCRIPTION] = f"PREP-SYN-OBC-{suffix}"
    values[CHARGE_AMOUNT] = 25.5 + index
    values[REMARKS_AMOUNT] = 10 + index
    values[REMARKS] = f"PREP-SYN-UWAGI-{suffix}"
    return values


def _cells(workbook: Workbook) -> dict[tuple[str, int, int], object]:
    return {
        (sheet.title, cell.row, cell.column): cell.value
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    }


def _fill(root: Path) -> int:
    period_directory = root / "INEA" / PERIOD
    collective_file = period_directory / f"Rozliczenie {PERIOD} - zbiorcze.xlsx"
    internal_settlement = period_directory / f"Rozliczenie wew. {PERIOD}.xlsx"
    links_before = _external_link_parts(collective_file)
    if not links_before:
        raise ValueError("Plik zbiorczy has no xl/externalLinks parts to preserve")

    workbook = load_workbook(collective_file, data_only=False, keep_links=True)
    try:
        sheet = _sheet(workbook)
        before = _cells(workbook)
        rows = _order_rows(sheet)
        missing = [order for order in FILLED_ORDERS if order not in rows]
        if missing:
            raise ValueError(f"Plik zbiorczy lacks {len(missing)} of the Wiersze danych to fill")
        expected = dict(before)
        for index, order in enumerate(FILLED_ORDERS):
            for column, value in _readiness_values(index, order).items():
                cell = sheet.cell(rows[order], column)
                if cell.data_type == "f":
                    raise ValueError(f"Plik zbiorczy has a formula in {cell.coordinate}, where Dad types values")
                cell.value = value
                expected[(sheet.title, rows[order], column)] = value
        save_workbook(workbook, collective_file)
    finally:
        workbook.close()

    errors: list[str] = []
    written = load_workbook(collective_file, data_only=False, keep_links=True)
    try:
        if _cells(written) != expected:
            errors.append("Plik zbiorczy differs from the Gotowość entries or lost other cells or formulas")
    finally:
        written.close()
    if _external_link_parts(collective_file) != links_before:
        errors.append("Plik zbiorczy xl/externalLinks parts changed")
    if errors:
        for error in errors:
            print(f"READINESS FAIL: {error}", file=sys.stderr)
        return 1

    manifest = {
        "collective_file_sha256": _sha256(collective_file),
        "internal_settlement_sha256": _sha256(internal_settlement),
        "filled_orders": list(FILLED_ORDERS),
    }
    (root / "readiness_manifest.json").write_text(json.dumps(manifest, sort_keys=True) + "\n", encoding="utf-8")
    print(
        f"READINESS OK | period={PERIOD} | rows_filled={len(FILLED_ORDERS)} | columns=N:AT | "
        "other_cells_unchanged=true | external_links_preserved=true"
    )
    return 0


def main(argv: list[str] | None = None) -> int:
    parser = Parser(description=__doc__)
    parser.add_root("run root of a chain fixture")
    args = parser.parse_args(argv)
    try:
        return _fill(require_run_root(args.root))
    except (OSError, ValueError, KeyError, zipfile.BadZipFile) as exc:
        print(f"READINESS FAIL: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
