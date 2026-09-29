"""Create disposable synthetic workbooks for the rozliczenia verification skill."""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
import sys

import yaml
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.worksheet.worksheet import Worksheet


PERIOD = "08_14_09_2026"
SOURCE_NAME = f"Rozliczenie {PERIOD} - zbiorcze.xlsx"
WORKERS = {
    "adrian.maciejewski": "Adrian Maciejewski",
    "dariusz.nowak2": "Darek Nowak",
    "kamil.frontczak": "Kamil Frontczak",
}


def _repo_root() -> Path:
    return Path(__file__).resolve().parents[4]


def _run_parent() -> Path:
    return (_repo_root() / ".verification" / "runs").resolve()


def _active_worksheet(workbook: Workbook) -> Worksheet:
    sheet = workbook.active
    if not isinstance(sheet, Worksheet):
        raise ValueError("Workbook does not contain an active worksheet")
    return sheet


def _safe_run_root(value: str) -> Path:
    root = Path(value).expanduser().resolve()
    parent = _run_parent()
    try:
        root.relative_to(parent)
    except ValueError as exc:
        raise ValueError(f"Fixture root must be below {parent}") from exc
    if root == parent:
        raise ValueError("Fixture root must be a child of the runs directory")
    return root


def _write_template(path: Path) -> None:
    workbook = Workbook()
    sheet = _active_worksheet(workbook)
    sheet.title = "Sheet1"
    sheet.cell(17, 8).value = "WYKONAWCA"
    sheet["A17"].fill = PatternFill(fill_type="solid", fgColor="00AA55")
    sheet.cell(18, 47).value = "=N18"
    sheet.cell(19, 47).value = "=N19"
    sheet.cell(18, 14).value = 0
    sheet.cell(19, 14).value = 0
    rates = workbook.create_sheet("Rates")
    rates["A1"] = "synthetic rate"
    rates["B1"] = 17.5
    workbook.save(path)
    workbook.close()


def _write_source(path: Path, *, include_unmapped: bool) -> None:
    workbook = Workbook()
    sheet = _active_worksheet(workbook)
    sheet.title = "Sheet1"
    sheet.cell(17, 8).value = "WYKONAWCA"
    rows = [
        ("TEST-CITY", "adrian.maciejewski", "SYNTHETIC-1"),
        ("TEST-CITY", "dariusz.nowak2", "SYNTHETIC-2"),
        ("TEST-CITY", "adrian.maciejewski", "SYNTHETIC-3"),
    ]
    if include_unmapped:
        rows.append(("TEST-CITY", "unknown.synthetic", "SYNTHETIC-4"))
    for row_number, (city, worker, marker) in enumerate(rows, start=18):
        sheet.cell(row_number, 1).value = city
        sheet.cell(row_number, 6).value = marker
        sheet.cell(row_number, 8).value = worker
        sheet.cell(row_number, 14).value = row_number - 17
    workbook.save(path)
    workbook.close()


def _write_mapping(path: Path) -> None:
    path.write_text(
        yaml.safe_dump(
            {"schema_version": 1, "workers": WORKERS},
            allow_unicode=True,
            sort_keys=False,
        ),
        encoding="utf-8",
    )


def _write_fixture(root: Path, *, include_unmapped: bool) -> None:
    period_directory = root / PERIOD
    period_directory.mkdir(parents=True, exist_ok=False)
    source = period_directory / SOURCE_NAME
    placeholder = root / "placeholder.xlsx"
    _write_source(source, include_unmapped=include_unmapped)
    _write_mapping(root / "worker_mapping.yaml")
    _write_template(placeholder)
    manifest = {
        "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        "placeholder_sha256": hashlib.sha256(placeholder.read_bytes()).hexdigest(),
        "include_unmapped": include_unmapped,
    }
    (root / "fixture_manifest.json").write_text(
        json.dumps(manifest, sort_keys=True) + "\n",
        encoding="utf-8",
    )


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=str, required=True)
    parser.add_argument("--include-unmapped", action="store_true")
    args = parser.parse_args(argv)
    try:
        root = _safe_run_root(args.root)
        if root.exists() and any(root.iterdir()):
            raise ValueError(f"Fixture root is not empty: {root}")
        root.mkdir(parents=True, exist_ok=False)
        _write_fixture(root, include_unmapped=args.include_unmapped)
    except (OSError, ValueError) as exc:
        print(f"FIXTURE FAILED: {exc}", file=sys.stderr)
        return 1
    print(f"FIXTURE_READY root={root} period={PERIOD} workers={len(WORKERS)} unmapped={args.include_unmapped}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
