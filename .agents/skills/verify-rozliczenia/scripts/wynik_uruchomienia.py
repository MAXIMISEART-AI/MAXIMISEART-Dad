"""Read what an Uruchomienie procesu left behind: its JSON result through the contract, metrics, and Hiperłącza.

The assertion scripts share this module, so they read the result of ``tools/uruchom_proces.py`` as the
contract model ``ProcessResult`` it is printed from, never as a raw dict: a field renamed or dropped in the
contract fails validation here and type checking in the scripts. Run alone, it checks one evidence file:

    py -3 .agents\\skills\\verify-rozliczenia\\scripts\\wynik_uruchomienia.py --result "$evidenceRoot\\run.json"

Messages name fields and files, never a value, because values can hold row content.
"""

from __future__ import annotations

import argparse
from collections.abc import Iterable, Sequence
import json
from pathlib import Path
import sys
from typing import Any
import zipfile

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import ValidationError

sys.path.insert(0, str(Path(__file__).resolve().parents[4] / "tools"))

from uruchom_proces import ProcessResult

from rozliczenia.centrum import contract

__all__ = [
    "Link",
    "ProcessResult",
    "check_metrics",
    "contract",
    "external_link_parts",
    "figures",
    "finish",
    "hyperlinks",
    "load_result",
]

Link = tuple[object, object]


def _contract_errors(error: ValidationError) -> str:
    return ", ".join(f"{'.'.join(map(str, detail['loc'])) or '<result>'} ({detail['type']})" for detail in error.errors())


def load_result(
    path: Path | None, row_content: Iterable[str], errors: list[str], *, naming: Iterable[str] = ()
) -> ProcessResult | None:
    """The result of a process step's evidence JSON as ``ProcessResult``; ``None`` after recording why not.

    ``naming``: exact Komunikat texts allowed to carry a row value (an unknown WYKONAWCA identifier), nowhere else.
    """

    if path is None or not path.is_file():
        errors.append("process step evidence is missing")
        return None
    text = path.read_text(encoding="utf-8")
    raw_text = text
    for allowed in naming:
        text = text.replace(allowed, "").replace(json.dumps(allowed)[1:-1], "")
    if any(value in text for value in row_content):
        errors.append("process result contains row content")
    raw = json.loads(raw_text)["result"]
    if not isinstance(raw, dict):
        errors.append("process step printed no JSON result")
        return None
    try:
        result = ProcessResult.model_validate(raw)
    except ValidationError as error:
        errors.append(f"process result does not match the contract ProcessResult: {_contract_errors(error)}")
        return None
    if not result.metrics_saved:
        errors.append("process result reports its metrics record as not saved")
    return result


def figures(items: Sequence[contract.Figure]) -> list[tuple[str, int]]:
    return [(figure.label, figure.value) for figure in items]


def check_metrics(
    metrics: Path,
    root: Path,
    *,
    period: str,
    expected: list[tuple[str, str]],
    row_content: Iterable[str],
    errors: list[str],
) -> list[dict[str, Any]]:
    """One completed record of ``period`` per expected (mode, result), with no row content or path.

    Returns the records for the process's own counter checks.
    """

    content = metrics.read_text(encoding="utf-8") if metrics.is_file() else ""
    records = [json.loads(line) for line in content.splitlines() if line.strip()]
    if [(record.get("mode"), record.get("result")) for record in records] != expected:
        errors.append(f"{metrics.name} does not contain exactly {expected}")
    for record in records:
        if record.get("period") != period or record.get("completed") is not True:
            errors.append(f"{metrics.name} record is not a completed run of the period")
    root_spellings = (str(root), json.dumps(str(root))[1:-1])
    if any(value in content for value in row_content) or any(spelling in content for spelling in root_spellings):
        errors.append(f"{metrics.name} contains row content or a filesystem path")
    return records


def hyperlinks(sheet: Worksheet) -> dict[tuple[int, int], Link]:
    """(row, column) -> (address, location) of every cell hyperlink; reads only cells stored in the file."""

    return {
        coordinate: (cell.hyperlink.target, cell.hyperlink.location)
        for coordinate, cell in sheet._cells.items()
        if cell.hyperlink is not None
    }


def look_mismatches(sheet: Worksheet, header_row: int, written_rows: range, look_columns: int) -> list[str]:
    """Wygląd Wiersza danych of every written cell in A..``look_columns``, named by column and property only.

    Font face and size of the header cell, left and centre alignment, wrapping only for text with a line break,
    a thin automatic border on four sides, and no stored height on a written row.
    """

    found: set[str] = set()
    for row in written_rows:
        dimension = sheet.row_dimensions.get(row)
        if dimension is not None and dimension.ht is not None:
            found.add("stored row height")
        for column in range(1, look_columns + 1):
            cell = sheet.cell(row, column)
            letter = get_column_letter(column)
            face, header_face = _face(sheet, row, column), _face(sheet, header_row, column)
            if face != header_face:
                found.add(f"{letter} font {face} vs header {header_face}")
            wraps = isinstance(cell.value, str) and "\n" in cell.value
            if (cell.alignment.horizontal, cell.alignment.vertical, bool(cell.alignment.wrap_text)) != ("left", "center", wraps):
                found.add(f"{letter} alignment")
            sides = (cell.border.left, cell.border.right, cell.border.top, cell.border.bottom)
            if any(side.style != "thin" or side.color is None or side.color.indexed != 64 for side in sides):
                found.add(f"{letter} border")
    return sorted(found)


def _face(sheet: Worksheet, row: int, column: int) -> str:
    font = sheet.cell(row, column).font
    return f"{font.name} {font.sz}"


def external_link_parts(path: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(path) as archive:
        return {name: archive.read(name) for name in archive.namelist() if name.startswith("xl/externalLinks/")}


def finish(errors: list[str], success_message: str) -> int:
    if errors:
        for error in errors:
            print(f"ASSERT FAIL: {error}", file=sys.stderr)
        return 1
    print(success_message)
    return 0


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--result", required=True, type=Path, help="evidence JSON of a process step")
    args = parser.parse_args(argv)
    errors: list[str] = []
    try:
        result = load_result(args.result.expanduser().resolve(), (), errors)
    except (OSError, ValueError, KeyError) as exc:
        errors.append(f"evidence cannot be read: {exc.__class__.__name__}")
        result = None
    if result is None:
        return finish(errors, "")
    return finish(errors, f"WYNIK OK | result={result.result} | stop={result.stop.message.code if result.stop else None}")


if __name__ == "__main__":
    raise SystemExit(main())
