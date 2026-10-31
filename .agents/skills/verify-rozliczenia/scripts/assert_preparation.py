"""Assert Przygotowanie pliku zbiorczego side effects without printing row contents."""

from __future__ import annotations

from datetime import datetime
import hashlib
import json
from pathlib import Path
import sys
import zipfile

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from create_preparation_fixture import REGISTER_NAME, completed_activities, order_contract, order_devices
sys.path.insert(0, str(Path(__file__).resolve().parents[4] / "tools/lib"))

from wynik_uruchomienia import (
    check_metrics,
    contract,
    external_link_parts,
    figures,
    finish,
    hyperlinks,
    load_result,
    look_mismatches,
)
from skrypt_skilla import Parser

from rozliczenia.workbook_files import QUANTITY_COLUMNS


PERIOD = "08_14_09_2026"
INTERNAL_NAME = f"Rozliczenie wew. {PERIOD}.xlsx"
COLLECTIVE_NAME = f"Rozliczenie {PERIOD} - zbiorcze.xlsx"
ROW_MARKER = "PREP-SYN-"
ORDER_LINK = "https://przyklad.invalid/zlecenie?caseId=PREP-SYN-LINK&tw=1&wiersz="
EXPECTED_COUNTERS = {"period_orders": 7, "successful_orders": 6, "skipped_empty_rows": 4, "warnings": 3}
# Ilości usług Dad completes: Plik zbiorczy rows 20 (order 7, prewencja) and 21 (order 6, unknown device), and
# the reminder every Przygotowanie carries. Only the written Plik zbiorczy needs review; the dry-run stays OK.
EXPECTED_WARNINGS = [
    ("NIEZNANE_URZADZENIA", f"plik {COLLECTIVE_NAME}, wiersz 21"),
    ("WIERSZE_Z_OPISU", f"plik {COLLECTIVE_NAME}, wiersz 20"),
    ("KOLUMNY_Z_OPISU", f"plik {COLLECTIVE_NAME}, kolumny Y, AB, AD, AJ, AK, AL i AM"),
]
EXPECTED_RESULT = {"DRY-RUN": "OK", "RUN": "WYMAGA_SPRAWDZENIA"}
EXPECTED_SUMMARY = contract.RunSummary(
    figures=[
        contract.Figure(label="Zlecenia okresu", value=7, unit="count"),
        contract.Figure(label="Zlecenia skuteczne", value=6, unit="count"),
        contract.Figure(label="Pominięte puste wiersze", value=4, unit="count"),
    ],
    counts=contract.CountGroup(
        title="Zlecenia na dzień",
        counts=[
            contract.Count(label=label, count=count)
            for label, count in (
                ("wt 08.09.2026", 2),
                ("śr 09.09.2026", 1),
                ("czw 10.09.2026", 0),
                ("pt 11.09.2026", 0),
                ("sob 12.09.2026", 1),
                ("niedz 13.09.2026", 1),
                ("pon 14.09.2026", 2),
            )
        ],
    ),
    files=[
        contract.PeriodFile(name=INTERNAL_NAME, kind="internal_settlement"),
        contract.PeriodFile(name=COLLECTIVE_NAME, kind="collective_file"),
    ],
    worker_settlements=None,
)
INTERNAL_TEMPLATE = Path("config") / "Wzorzec rozliczenia wewnętrznego.xlsx"
COLLECTIVE_TEMPLATE = Path("config") / "Wzorzec pliku zbiorczego.xlsx"
# Target headers the process fills; a real Wzorzec may order them differently.
TARGET_HEADERS = (
    "DATA",
    "MIASTO",
    "ADRES",
    "UMOWA",
    "URZĄDZENIA",
    "CZYNNOŚCI",
    "NUMER ZL.",
    "WYKONAWCA",
    "NUMER UMOWY/ZLECENIA ORANGE",
    "SKUTECZNE",
    "CZYNNOŚCI ZREALIZOWANE",
)


def _internal_row(number: int, started: datetime, worker: str, successful: str) -> tuple[object, ...]:
    """Expected A:M of the Rozliczenie wewnętrzne; zwrot and UWAGI stay empty."""

    return (
        started,
        "PREP-SYN-MIASTO",
        f"PREP-SYN-ADRES-{number}",
        order_contract(number),
        order_devices(number),
        "instalacja",
        f"PREP-SYN-ZL-{number}",
        worker,
        f"PREP-SYN-OR-{number}",
        successful,
        completed_activities(number),
        None,
        None,
    )


def _collective_row(number: int, started: datetime, worker: str, successful: str) -> tuple[object, ...]:
    """Expected A:K of the Plik zbiorczy in the Wzorzec header order."""

    return (
        "PREP-SYN-MIASTO",
        f"PREP-SYN-ADRES-{number}",
        order_contract(number),
        order_devices(number),
        "instalacja",
        f"PREP-SYN-ZL-{number}",
        started,
        worker,
        f"PREP-SYN-OR-{number}",
        successful,
        completed_activities(number),
    )


# Period orders in Rejestr order: 2, 3 (NIE), 4, 7 (after the next-period order), the late order 6, then 8 and 9.
EXPECTED_INTERNAL_ROWS = (
    _internal_row(2, datetime(2026, 9, 8, 0, 0), "adrian.testowy", "tak"),
    _internal_row(3, datetime(2026, 9, 9, 10, 30), "dariusz.fikcyjny2", "NIE"),
    _internal_row(4, datetime(2026, 9, 12, 12, 0), "kamil.syntetyczny", " Tak "),
    _internal_row(7, datetime(2026, 9, 13, 9, 0), "adrian.testowy", "tak"),
    _internal_row(6, datetime(2026, 9, 14, 18, 30), "dariusz.fikcyjny2", "TAK"),
    _internal_row(8, datetime(2026, 9, 8, 11, 0), "dariusz.fikcyjny2", "tak"),
    _internal_row(9, datetime(2026, 9, 14, 10, 0), "kamil.syntetyczny", "tak"),
)
EXPECTED_COLLECTIVE_ROWS = (
    _collective_row(2, datetime(2026, 9, 8, 0, 0), "adrian.testowy", "tak"),
    _collective_row(4, datetime(2026, 9, 12, 12, 0), "kamil.syntetyczny", " Tak "),
    _collective_row(7, datetime(2026, 9, 13, 9, 0), "adrian.testowy", "tak"),
    _collective_row(6, datetime(2026, 9, 14, 18, 30), "dariusz.fikcyjny2", "TAK"),
    _collective_row(8, datetime(2026, 9, 8, 11, 0), "dariusz.fikcyjny2", "tak"),
    _collective_row(9, datetime(2026, 9, 14, 10, 0), "kamil.syntetyczny", "tak"),
)
# Ilości usług of rows 18..23 (orders 2, 4, 7, 6, 8, 9) from the Reguły ilości; every other N:AP cell keeps the Wzorzec 0.
EXPECTED_QUANTITIES: tuple[dict[str, int], ...] = (
    {"O": 1, "P": 1, "Q": 1, "S": 3, "T": 1, "V": 1, "Z": 1},
    {"N": 45, "O": 1, "S": 2, "V": 1, "AG": 1, "AI": 1, "AO": 1},
    {},
    {"N": 45, "T": 1, "V": 1},
    {"V": 1, "AE": 1, "AP": 1},
    {"O": 1, "P": 1, "Q": 2, "S": 4, "V": 1, "Z": 1, "AK": 1},
)
EXPECTED_INTERNAL_ORDER_NUMBERS = tuple(row[6] for row in EXPECTED_INTERNAL_ROWS)
EXPECTED_COLLECTIVE_ORDER_NUMBERS = tuple(row[5] for row in EXPECTED_COLLECTIVE_ROWS)
# Rejestr rows of those orders; their NUMER ZL. Hiperłącze ends with the row number.
EXPECTED_INTERNAL_REGISTER_ROWS = (4, 6, 9, 12, 13, 14, 15)
EXPECTED_COLLECTIVE_REGISTER_ROWS = (4, 9, 12, 13, 14, 15)


def _sheet(workbook_sheet: object) -> Worksheet:
    if not isinstance(workbook_sheet, Worksheet):
        raise ValueError("Workbook does not contain an active worksheet")
    return workbook_sheet


def _rows(sheet: Worksheet, first: int, last: int, max_column: int) -> list[tuple[object, ...]]:
    return [tuple(row) for row in sheet.iter_rows(min_row=first, max_row=last, max_col=max_column, values_only=True)]


def _unchanged_inputs(root: Path) -> list[str]:
    manifest = json.loads((root / "fixture_manifest.json").read_text(encoding="utf-8"))
    errors: list[str] = []
    for label, path, key in (
        ("Rejestr zleceń", root / REGISTER_NAME, "register_sha256"),
        ("Wzorzec rozliczenia wewnętrznego", root / INTERNAL_TEMPLATE, "internal_template_sha256"),
        ("Wzorzec pliku zbiorczego", root / COLLECTIVE_TEMPLATE, "collective_template_sha256"),
    ):
        if hashlib.sha256(path.read_bytes()).hexdigest() != manifest.get(key):
            errors.append(f"run changed {label}")
    return errors


def _check_result(result_path: Path | None, mode: str, errors: list[str]) -> None:
    result = load_result(result_path, (ROW_MARKER,), errors)
    if result is None:
        return
    if result.result != EXPECTED_RESULT[mode] or result.stop is not None:
        errors.append(f"result is not {EXPECTED_RESULT[mode]} without a stop")
    if [(warning.code, warning.where) for warning in result.warnings] != EXPECTED_WARNINGS:
        errors.append("warnings are not the Ilości usług warnings of the synthetic Rejestr")
    if result.summary != EXPECTED_SUMMARY:
        errors.append("summary does not match the synthetic Rejestr")
    if dict(figures(result.progress.figures)).get("Ostrzeżenia") != len(EXPECTED_WARNINGS):
        errors.append(f"progress does not report {len(EXPECTED_WARNINGS)} Ostrzeżenia")


def _check_metrics(root: Path, metrics: Path, modes: list[str], errors: list[str]) -> None:
    """Every record is a completed run of the period with its mode's result and the synthetic Rejestr's counters."""

    expected = [(mode, EXPECTED_RESULT[mode]) for mode in modes]
    records = check_metrics(metrics, root, period=PERIOD, expected=expected, row_content=(ROW_MARKER,), errors=errors)
    for record in records:
        if record.get("counters") != EXPECTED_COUNTERS:
            errors.append("metrics counters do not match the synthetic Rejestr")


def _assert_dry_run(root: Path, metrics: Path, result_path: Path | None) -> int:
    errors = _unchanged_inputs(root)
    leftovers = sorted(path.name for path in (root / "INEA").iterdir())
    if leftovers:
        errors.append(f"dry-run created entries in the periods directory: {leftovers}")
    _check_result(result_path, "DRY-RUN", errors)
    _check_metrics(root, metrics, ["DRY-RUN"], errors)
    return finish(
        errors,
        "ASSERT OK | process=przygotowanie | mode=dry-run | period_folder_absent=true | "
        "inputs_unchanged=true | plan_reported=true | rows_hidden=true",
    )


def _check_internal_settlement(path: Path, errors: list[str]) -> None:
    workbook = load_workbook(path, data_only=False)
    try:
        sheet = _sheet(workbook.active)
        if _rows(sheet, 2, 2 + len(EXPECTED_INTERNAL_ROWS), 13) != [*EXPECTED_INTERNAL_ROWS, (None,) * 13]:
            errors.append("Rozliczenie wewnętrzne does not hold exactly the period orders in Rejestr order")
        highlighted = sheet["C2"]
        if (
            highlighted.fill.fgColor.rgb != "FFFFFF00"
            or highlighted.font.color is None
            or highlighted.font.color.rgb != "FFC00000"
            or not highlighted.font.b
            or (highlighted.font.name, highlighted.font.sz) != ("Calibri", 10)
        ):
            errors.append("Rozliczenie wewnętrzne lost a Rejestr highlight or the Wzorzec font")
        if not sheet["H4"].font.i or sheet["A2"].number_format != "dd.mm.yyyy hh:mm":
            errors.append("Rozliczenie wewnętrzne lost italic or the Rejestr date format")
        if sheet.auto_filter.ref != f"A1:M{1 + len(EXPECTED_INTERNAL_ROWS)}" or sheet.freeze_panes != "A2" or sheet.column_dimensions["A"].width != 10:
            errors.append("Rozliczenie wewnętrzne autofilter or Wzorzec layout is wrong")
        look = look_mismatches(sheet, 1, range(2, 2 + len(EXPECTED_INTERNAL_ROWS)), 13)
        if look:
            errors.append(f"Rozliczenie wewnętrzne: Wygląd Wiersza danych differs: {'; '.join(look)}")
    finally:
        workbook.close()


def _check_collective_file(path: Path, errors: list[str]) -> None:
    workbook = load_workbook(path, data_only=False)
    try:
        sheet = _sheet(workbook.active)
        if _rows(sheet, 18, 18 + len(EXPECTED_COLLECTIVE_ROWS), 11) != [*EXPECTED_COLLECTIVE_ROWS, (None,) * 11]:
            errors.append("Plik zbiorczy does not hold exactly the Zlecenia skuteczne from row 18")
        quantities = [
            {get_column_letter(column): sheet.cell(row, column).value for column in QUANTITY_COLUMNS} for row in range(18, 26)
        ]
        expected = [
            {get_column_letter(column): rules.get(get_column_letter(column), 0) for column in QUANTITY_COLUMNS}
            for rules in (*EXPECTED_QUANTITIES, {}, {})
        ]
        if quantities != expected:
            errors.append("Plik zbiorczy N:AP is not the Ilości usług from the Reguły ilości over the Wzorzec zeros")
        formulas = [sheet.cell(row, 47).value for row in range(18, 26)]
        if formulas != [f"=N{row}*$N$1" for row in range(18, 26)]:
            errors.append("Plik zbiorczy changed the Wzorzec formulas")
        if (
            sheet["N1"].value != 17.5
            or sheet["N16"].value != "=SUM(N18:N40)"
            or sheet["A17"].fill.fgColor.rgb != "0000AA55"
            or workbook.sheetnames != ["Sheet1", "Rates"]
            or workbook["Rates"]["B1"].value != 17.5
        ):
            errors.append("Plik zbiorczy lost Wzorzec rates, sums, sheets or header formatting")
        if sheet["B18"].fill.fgColor.rgb != "FFFFFF00" or sheet.auto_filter.ref != f"A17:AP{17 + len(EXPECTED_COLLECTIVE_ROWS)}":
            errors.append("Plik zbiorczy lost a Rejestr highlight or has a wrong autofilter")
        look = look_mismatches(sheet, 17, range(18, 18 + len(EXPECTED_COLLECTIVE_ROWS)), 12)
        if look:
            errors.append(f"Plik zbiorczy: Wygląd Wiersza danych differs: {'; '.join(look)}")
    finally:
        workbook.close()


def _target_positions(sheet: Worksheet, header_row: int, label: str, errors: list[str]) -> dict[str, int]:
    """Columns of the target headers in a real Wzorzec, found by header text like the process does."""

    by_text = {
        str(cell.value).strip().casefold(): cell.column
        for cell in sheet[header_row]
        if isinstance(cell.value, str) and cell.value.strip()
    }
    positions = {header: by_text[header.casefold()] for header in TARGET_HEADERS if header.casefold() in by_text}
    if len(positions) != len(TARGET_HEADERS):
        errors.append(f"{label}: Wzorzec lacks {len(TARGET_HEADERS) - len(positions)} target headers in row {header_row}")
    return positions


def _cell_values(workbook: Workbook, skipped: set[tuple[int, int]]) -> dict[tuple[str, int, int], object]:
    """Every stored cell value except the ``skipped`` (row, column) cells of the active sheet."""

    active = workbook.active
    values: dict[tuple[str, int, int], object] = {}
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None or (sheet is active and (cell.row, cell.column) in skipped):
                    continue
                values[(sheet.title, cell.row, cell.column)] = cell.value
    return values


def _check_real_template_file(
    template: Path,
    output: Path,
    *,
    label: str,
    header_row: int,
    look_columns: int,
    order_numbers: tuple[object, ...],
    quantities: tuple[dict[str, int], ...],
    errors: list[str],
) -> None:
    """Rows routed by header, the Wygląd Wiersza danych in A..``look_columns``, and the rest of the Wzorzec unchanged.

    ``quantities`` are the Ilości usług of the written rows by column letter; they replace the Wzorzec value.

    Messages name columns and fonts only, never a cell value of the Wzorzec or the Rejestr.
    """

    template_book = load_workbook(template, data_only=False, keep_links=True)
    output_book = load_workbook(output, data_only=False, keep_links=True)
    try:
        template_sheet = _sheet(template_book.active)
        sheet = _sheet(output_book.active)
        positions = _target_positions(template_sheet, header_row, label, errors)
        if not positions:
            return
        first_row = header_row + 1
        written_rows = range(first_row, first_row + len(order_numbers))
        written = {(row, column) for row in written_rows for column in positions.values()}

        order_column = positions.get("NUMER ZL.")
        if order_column is not None:
            routed = tuple(sheet.cell(row, order_column).value for row in written_rows)
            after = sheet.cell(first_row + len(order_numbers), order_column).value
            if routed != order_numbers or after != template_sheet.cell(first_row + len(order_numbers), order_column).value:
                errors.append(f"{label}: NUMER ZL. does not hold exactly the expected orders from row {first_row}")

        look = look_mismatches(sheet, header_row, written_rows, look_columns)
        if look:
            errors.append(f"{label}: Wygląd Wiersza danych differs: {'; '.join(look)}")

        expected_values = _cell_values(template_book, written)
        for row, row_quantities in zip(written_rows, quantities, strict=True):
            for letter, value in row_quantities.items():
                expected_values[(template_sheet.title, row, column_index_from_string(letter))] = value
        actual_values = _cell_values(output_book, written)
        changed = sorted(
            key for key in expected_values.keys() | actual_values.keys() if expected_values.get(key) != actual_values.get(key)
        )
        if output_book.sheetnames != template_book.sheetnames or changed:
            where = ", ".join(f"{title}!{get_column_letter(column)}{row}" for title, row, column in changed[:5])
            errors.append(f"{label}: {len(changed)} Wzorzec cells outside the written rows changed ({where})")
    finally:
        template_book.close()
        output_book.close()


def _fill(cell: Cell | MergedCell) -> tuple[object, object]:
    fill_type = cell.fill.fill_type
    return fill_type, cell.fill.fgColor.rgb if fill_type else None


def _check_address_fills(
    register: Path, output: Path, *, label: str, header_row: int, register_rows: tuple[int, ...], errors: list[str]
) -> None:
    """ADRES of every written row has the fill of that order's ADRES in the Rejestr, whatever the colour or none."""

    register_book = load_workbook(register, data_only=False)
    output_book = load_workbook(output, data_only=False)
    try:
        register_sheet = register_book["ZLECENIA"]
        [register_column] = [cell.column for cell in register_sheet[1] if cell.value == "ADRES"]
        expected = [_fill(register_sheet.cell(row, register_column)) for row in register_rows]
        if len(set(expected)) < 3:
            errors.append("Rejestr ADRES fills of the period orders lack yellow, another colour or none, so the check proves nothing")
        sheet = _sheet(output_book.active)
        column = _target_positions(sheet, header_row, label, errors).get("ADRES", 0)
        rows = range(header_row + 1, header_row + 1 + len(register_rows))
        differing = [
            f"{get_column_letter(column)}{row}"
            for row, fill in zip(rows, expected, strict=True)
            if _fill(sheet.cell(row, column)) != fill
        ]
        if differing:
            errors.append(f"{label}: ADRES fill differs from the Rejestr ADRES of the same order ({', '.join(differing)})")
    finally:
        register_book.close()
        output_book.close()


def _check_hyperlinks(
    template: Path, output: Path, *, label: str, header_row: int, register_rows: tuple[int, ...], errors: list[str]
) -> None:
    """Written NUMER ZL. cells carry the Rejestr Hiperłącze, other written cells none, the rest the Wzorzec's own.

    Messages name cells only, never an address.
    """

    template_book = load_workbook(template, data_only=False)
    output_book = load_workbook(output, data_only=False)
    try:
        template_sheet = _sheet(template_book.active)
        positions = _target_positions(template_sheet, header_row, label, [])
        first_row = header_row + 1
        written = {(first_row + index, column) for index in range(len(register_rows)) for column in positions.values()}
        expected = {cell: link for cell, link in hyperlinks(template_sheet).items() if cell not in written}
        order_column = positions.get("NUMER ZL.", 0)
        expected |= {(first_row + index, order_column): (f"{ORDER_LINK}{row}", None) for index, row in enumerate(register_rows)}
        actual = hyperlinks(_sheet(output_book.active))
        differing = sorted(cell for cell in expected.keys() | actual.keys() if expected.get(cell) != actual.get(cell))
        if differing:
            where = ", ".join(f"{get_column_letter(column)}{row}" for row, column in differing[:5])
            errors.append(f"{label}: {len(differing)} cells lack the Rejestr NUMER ZL. Hiperłącze or keep a wrong one ({where})")
    finally:
        template_book.close()
        output_book.close()


def _assert_run(root: Path, metrics: Path, result_path: Path | None, templates: str) -> int:
    errors = _unchanged_inputs(root)
    periods_directory = root / "INEA"
    period_directory = periods_directory / PERIOD
    if sorted(path.name for path in periods_directory.iterdir()) != [PERIOD]:
        errors.append("periods directory holds more than the Folder okresu (staging left behind?)")
    if not period_directory.is_dir() or sorted(path.name for path in period_directory.iterdir()) != sorted(
        [INTERNAL_NAME, COLLECTIVE_NAME]
    ):
        errors.append("Folder okresu does not contain exactly both period files")
    else:
        try:
            if templates == "synthetic":
                _check_internal_settlement(period_directory / INTERNAL_NAME, errors)
                _check_collective_file(period_directory / COLLECTIVE_NAME, errors)
            else:
                _check_real_template_file(
                    root / INTERNAL_TEMPLATE,
                    period_directory / INTERNAL_NAME,
                    label="Rozliczenie wewnętrzne",
                    header_row=1,
                    look_columns=13,
                    order_numbers=EXPECTED_INTERNAL_ORDER_NUMBERS,
                    quantities=({},) * len(EXPECTED_INTERNAL_ORDER_NUMBERS),
                    errors=errors,
                )
                _check_real_template_file(
                    root / COLLECTIVE_TEMPLATE,
                    period_directory / COLLECTIVE_NAME,
                    label="Plik zbiorczy",
                    header_row=17,
                    look_columns=12,
                    order_numbers=EXPECTED_COLLECTIVE_ORDER_NUMBERS,
                    quantities=EXPECTED_QUANTITIES,
                    errors=errors,
                )
            for template, output, label, header_row, register_rows in (
                (INTERNAL_TEMPLATE, INTERNAL_NAME, "Rozliczenie wewnętrzne", 1, EXPECTED_INTERNAL_REGISTER_ROWS),
                (COLLECTIVE_TEMPLATE, COLLECTIVE_NAME, "Plik zbiorczy", 17, EXPECTED_COLLECTIVE_REGISTER_ROWS),
            ):
                _check_hyperlinks(
                    root / template,
                    period_directory / output,
                    label=label,
                    header_row=header_row,
                    register_rows=register_rows,
                    errors=errors,
                )
                _check_address_fills(
                    root / REGISTER_NAME,
                    period_directory / output,
                    label=label,
                    header_row=header_row,
                    register_rows=register_rows,
                    errors=errors,
                )
                template_links = external_link_parts(root / template)
                if external_link_parts(period_directory / output) != template_links:
                    errors.append(f"{output} xl/externalLinks parts differ from the Wzorzec")
            if not external_link_parts(root / COLLECTIVE_TEMPLATE):
                errors.append("Wzorzec pliku zbiorczego has no xl/externalLinks parts to preserve")
        except (OSError, ValueError, KeyError, zipfile.BadZipFile):
            errors.append("one or both period files cannot be read")
    _check_result(result_path, "RUN", errors)
    _check_metrics(root, metrics, ["DRY-RUN", "RUN"], errors)
    return finish(
        errors,
        f"ASSERT OK | process=przygotowanie | mode=run | templates={templates} | workbooks_checked=2 | "
        "formatting_preserved=true | row_look=A:M,A:L | hyperlinks_from_rejestr=NUMER ZL. | address_fill_from_rejestr=ADRES | "
        "external_links_preserved=true | inputs_unchanged=true",
    )


def main(argv: list[str] | None = None) -> int:
    parser = Parser(description=__doc__)
    parser.add_root("run root of a preparation fixture")
    parser.add_argument("--expect", choices=("dry-run", "run"), required=True)
    parser.add_argument(
        "--templates",
        choices=("synthetic", "config"),
        default="synthetic",
        help="'config': the fixture copied the real Wzorce, so check them generically",
    )
    parser.add_argument("--metrics", type=Path, help="metrics history of the Przygotowanie (default: <root>/metrics.jsonl)")
    parser.add_argument("--result", type=Path, help="evidence JSON of the Przygotowanie step")
    args = parser.parse_args(argv)
    root = args.root
    metrics = args.metrics.expanduser().resolve() if args.metrics else root / "metrics.jsonl"
    result_path = args.result.expanduser().resolve() if args.result else None
    try:
        if args.expect == "dry-run":
            return _assert_dry_run(root, metrics, result_path)
        return _assert_run(root, metrics, result_path, args.templates)
    except (OSError, ValueError, KeyError) as exc:
        print(f"ASSERT FAIL: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
