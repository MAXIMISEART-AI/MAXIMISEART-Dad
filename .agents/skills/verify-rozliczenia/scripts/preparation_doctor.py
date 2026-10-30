"""Read-only build and fixture check before driving Przygotowanie pliku zbiorczego.

``--with-worker-template`` also checks the Wzorzec rozliczenia pracownika of the Przebieg rozliczeń that follows.
"""

from __future__ import annotations

from pathlib import Path
import sys
import tomllib
import zipfile

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[4] / "tools/lib"))
sys.path.insert(0, str(Path(__file__).resolve().parents[4] / "src"))

from rozliczenia.worker_config import load_worker_mapping
from rozliczenia.message import StopError
from rozliczenia.settlement_period import parse_period_token
from rozliczenia.template_file import TemplateFile
from rozliczenia.workbook_files import WorkbookError, is_open_in_excel
from skrypt_skilla import Parser

SCRIPT = Path(__file__).resolve()
REPO_ROOT = SCRIPT.parents[4]
PERIOD = "08_14_09_2026"
REQUIRED_REGISTER_HEADERS = frozenset(
    {
        "DATA GODZ. POCZ.",
        "MIASTO",
        "ADRES",
        "UMOWA",
        "URZĄDZENIA",
        "CZYNNOŚCI",
        "NUMER ZL.",
        "WYKONAWCA",
        "NUMER UMOWY",
        "SKUTECZNE",
        "CZYNNOŚCI ZREALIZOWANE",
    }
)
TARGET_HEADERS = frozenset(
    {
        "DATA",
        "MIASTO",
        "ADRES",
        "UMOWA",
        "URZĄDZENIA",
        "CZYNNOŚCI",
        "NUMER ZL.",
        "WYKONAWCA",
        "NUMER UMOWY/ZLECENIA ORANGE",
        "SKUTECZNE",
        "CZYNNOŚCI ZREALIZOWANE",
    }
)


def _available(path: Path) -> None:
    if not path.is_file():
        raise ValueError(f"missing workbook: {path.name}")
    if is_open_in_excel(path):
        raise ValueError(f"workbook has an Excel lock: {path.name}")


def _header_row(path: Path, row: int, max_column: int | None = None) -> set[object]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = workbook.active
        if sheet is None:
            raise ValueError(f"no active worksheet: {path.name}")
        values = next(sheet.iter_rows(min_row=row, max_row=row, max_col=max_column, values_only=True), ())
        return set(values)
    finally:
        workbook.close()


def _register_sheets(path: Path) -> int:
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        return sum(
            1
            for sheet in workbook.worksheets
            if REQUIRED_REGISTER_HEADERS <= set(next(sheet.iter_rows(max_row=1, values_only=True), ()))
        )
    finally:
        workbook.close()


def _check(root: Path, expected_version: str, *, worker_template: bool) -> tuple[str, int, int]:
    root = root.expanduser().resolve()
    if not root.is_dir():
        raise ValueError(f"missing fixture root: {root}")
    with (REPO_ROOT / "pyproject.toml").open("rb") as stream:
        version = tomllib.load(stream)["project"]["version"]
    if version != expected_version:
        raise ValueError(f"expected version {expected_version}, found {version}")
    if sys.version_info < (3, 11):
        raise ValueError("Python 3.11 or newer is required")

    period = parse_period_token(PERIOD).token
    register = root / "ZLECENIA INTERMAX OD 08_11_2021.xlsx"
    internal_template = root / "config" / "Wzorzec rozliczenia wewnętrznego.xlsx"
    collective_template = root / "config" / "Wzorzec pliku zbiorczego.xlsx"
    for path in (register, internal_template, collective_template):
        _available(path)

    order_sheets = _register_sheets(register)
    if order_sheets < 1:
        raise ValueError("Rejestr zleceń has no sheet with the required headers in row 1")
    if not TARGET_HEADERS <= _header_row(internal_template, 1):
        raise ValueError("Wzorzec rozliczenia wewnętrznego lacks expected headers in row 1")
    if not TARGET_HEADERS <= _header_row(collective_template, 17, 11):
        raise ValueError("Wzorzec pliku zbiorczego lacks expected headers in A17:K17")
    with zipfile.ZipFile(collective_template) as archive:
        external_parts = [name for name in archive.namelist() if name.startswith("xl/externalLinks/")]
    if not external_parts:
        raise ValueError("Wzorzec pliku zbiorczego has no xl/externalLinks parts")

    if worker_template:
        TemplateFile.worker_settlement(root / "config" / "Wzorzec rozliczenia pracownika.xlsx").check()

    mapping = load_worker_mapping(root / "config" / "worker_mapping.yaml")
    periods_directory = root / "INEA"
    if not periods_directory.is_dir():
        raise ValueError("periods directory is missing")
    if (periods_directory / period).exists():
        raise ValueError("Folder okresu must not exist before dry-run")
    return period, len(mapping), order_sheets


def main(argv: list[str] | None = None) -> int:
    parser = Parser(description=__doc__)
    parser.add_root("run root of a preparation fixture")
    parser.add_argument("--expected-version", default="0.1.0")
    parser.add_argument("--with-worker-template", action="store_true", help="also check the Wzorzec rozliczenia pracownika")
    args = parser.parse_args(argv)
    try:
        period, workers, order_sheets = _check(args.root, args.expected_version, worker_template=args.with_worker_template)
    except (
        OSError,
        ValueError,
        KeyError,
        TypeError,
        zipfile.BadZipFile,
        WorkbookError,
    ) as exc:
        print(f"DOCTOR FAIL: {exc}", file=sys.stderr)
        return 1
    except StopError as stop:
        print(f"DOCTOR FAIL: {stop.message.code} | {stop.message.what}", file=sys.stderr)
        return 1
    print(
        f"DOCTOR OK | app=MAXIMISEART-Dad | process=przygotowanie | version={args.expected_version} | "
        f"python={sys.version_info.major}.{sys.version_info.minor}.{sys.version_info.micro} | "
        f"period={period} | workers={workers} | order_sheets={order_sheets} | "
        f"templates=valid | worker_template={'valid' if args.with_worker_template else 'unchecked'} | "
        "external_links=present | period_folder_absent=true"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
