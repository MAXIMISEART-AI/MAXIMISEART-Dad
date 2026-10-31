"""Assert safe, user-visible workbook results without printing row contents."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path
import sys
from typing import NamedTuple

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

sys.path.insert(0, str(Path(__file__).resolve().parents[4] / "tools/lib"))

from wynik_uruchomienia import Link, ProcessResult, check_metrics, contract, figures, finish, hyperlinks, load_result
from skrypt_skilla import Parser


PERIOD = "08_14_09_2026"
ORDER_LINK = "https://przyklad.invalid/zlecenie?caseId=SYNTHETIC-LINK&tw=1&wiersz="
WORKERS = ("Adrian Testowy", "Darek Fikcyjny", "Kamil Syntetyczny")
PHASES = ("Sprawdzanie", "Odczyt danych", "Planowanie", "Zapisywanie")
WORKER_SETTLEMENTS = [
    contract.WorkerSettlement(worker=worker, file=f"Rozliczenie {PERIOD} - {worker}.xlsx", orders=orders, empty=orders == 0)
    for worker, orders in zip(WORKERS, (2, 1, 0), strict=True)
]
ORDER_COLUMN = 6  # F


class WorkbookSnapshot(NamedTuple):
    formula_au18: object
    formula_au19: object
    sheet_names: tuple[str, ...]
    header_fill_rgb: object
    rate_value: object


def _path(root: Path, worker: str) -> Path:
    return root / PERIOD / f"Rozliczenie pracowników {PERIOD}" / f"Rozliczenie {PERIOD} - {worker}.xlsx"


def _read_workbook_snapshot(path: Path) -> WorkbookSnapshot:
    workbook = load_workbook(path, data_only=False)
    try:
        sheet = workbook.active
        if not isinstance(sheet, Worksheet):
            raise ValueError("Workbook does not contain an active worksheet")
        return WorkbookSnapshot(
            sheet["AU18"].value,
            sheet["AU19"].value,
            tuple(workbook.sheetnames),
            sheet["A17"].fill.fgColor.rgb,
            workbook["Rates"]["B1"].value,
        )
    finally:
        workbook.close()


def _read_input_row(path: Path, row_number: int) -> tuple[object, ...]:
    workbook = load_workbook(path, data_only=False)
    try:
        sheet = workbook.active
        if not isinstance(sheet, Worksheet):
            raise ValueError("Workbook does not contain an active worksheet")
        return next(sheet.iter_rows(min_row=row_number, max_row=row_number, max_col=46, values_only=True))
    finally:
        workbook.close()


def _hyperlinks(path: Path) -> dict[tuple[int, int], Link]:
    workbook = load_workbook(path)
    try:
        sheet = workbook.active
        if not isinstance(sheet, Worksheet):
            raise ValueError("Workbook does not contain an active worksheet")
        return hyperlinks(sheet)
    finally:
        workbook.close()


def _order_links(rows: dict[int, int]) -> dict[tuple[int, int], Link]:
    """Expected Hiperłącza of a worker workbook: its row -> the Plik zbiorczy row the order came from."""

    return {(row, ORDER_COLUMN): (f"{ORDER_LINK}{source_row}", None) for row, source_row in rows.items()}


def _expected_source_row(worker_id: str, marker: str, quantity: int) -> tuple[object, ...]:
    return (
        "TEST-CITY",
        None,
        None,
        None,
        None,
        marker,
        None,
        worker_id,
        None,
        None,
        None,
        None,
        None,
        quantity,
    ) + (None,) * 32


def _unchanged_inputs(root: Path, manifest: dict[str, object]) -> list[str]:
    period_directory = root / PERIOD
    errors: list[str] = []
    for label, path, key in (
        ("source", period_directory / f"Rozliczenie {PERIOD} - zbiorcze.xlsx", "source_sha256"),
        ("worker template", root / "Wzorzec rozliczenia pracownika.xlsx", "worker_template_sha256"),
    ):
        actual = hashlib.sha256(path.read_bytes()).hexdigest()
        if actual != manifest.get(key):
            errors.append(f"run changed {label}")
    return errors


_UNMAPPED = "unknown.synthetic"
_ROW_CONTENT = ("TEST-CITY", "SYNTHETIC-1", "SYNTHETIC-2", "SYNTHETIC-3", "SYNTHETIC-4", _UNMAPPED)
_UNMAPPED_WARNING = f"Nieznany identyfikator WYKONAWCA: {_UNMAPPED}. Nie ma go w konfiguracji wykonawców."
_UNMAPPED_STOP = f"W Pliku zbiorczym jest WYKONAWCA spoza konfiguracji wykonawców: {_UNMAPPED}."
_NAMING_UNMAPPED = (_UNMAPPED_WARNING, _UNMAPPED_STOP)


def _check_metrics(root: Path, expected: list[tuple[str, str]], errors: list[str]) -> None:
    check_metrics(root / "metrics.jsonl", root, period=PERIOD, expected=expected, row_content=_ROW_CONTENT, errors=errors)


def _routes(copies: int) -> dict[str, dict[str, int]]:
    """``routes`` of a settlement record: ``copies`` Rozliczenia pracownika by copy, no other file (ADR 0008)."""

    none = {"kopia": 0, "openpyxl": 0}
    return {"worker_settlement": {"kopia": copies, "openpyxl": 0}, "collective_file": none, "internal_settlement": none}


def _check_stop_metrics(root: Path, expected_stop: str, errors: list[str]) -> None:
    """The stopped run's record carries the Rodzaj Zatrzymania and the Komunikat code, never its text.

    Every stopping variant stops „dla taty”: an existing output folder, a locked Wzorzec, an unknown WYKONAWCA.
    """

    records = [json.loads(line) for line in (root / "metrics.jsonl").read_text(encoding="utf-8").splitlines()]
    last = records[-1] if records else {}
    expected = (3, "BLAD_KRYTYCZNY", False, {"kind": "for_tata", "code": expected_stop})
    if (last.get("schema_version"), last.get("result"), last.get("completed"), last.get("stop")) != expected:
        errors.append(f"stopped run's metrics record is not {expected}")


def _check_progress(progress: contract.Progress, errors: list[str]) -> None:
    if [(phase.phase, phase.state) for phase in progress.phases] != [(phase, "completed") for phase in PHASES]:
        errors.append("run progress does not complete the four stages in order")
    if (progress.template_total, progress.templates_completed) != (3, 3) or figures(progress.figures) != [
        ("Wiersze danych", 3),
        ("Zapisane szablony", 2),
        ("Puste szablony", 1),
        ("Problemy", 0),
    ]:
        errors.append("run progress does not report 3/3 templates with two written and one empty")


def _check_stop(result: ProcessResult | None, expected_stop: str, expected_what: str | None, errors: list[str]) -> None:
    if result is None:
        return
    stop = result.stop
    if result.result != "BLAD_KRYTYCZNY" or result.summary is not None or stop is None or stop.message.code != expected_stop:
        errors.append(f"run did not stop with {expected_stop} (stopped with {stop.message.code if stop else None})")
    elif expected_what is not None and expected_what not in stop.message.what:
        errors.append(f"{expected_stop} does not say: {expected_what}")


def _assert_dry_run(root: Path, result_path: Path | None) -> int:
    period_directory = root / PERIOD
    target_directory = period_directory / f"Rozliczenie pracowników {PERIOD}"
    manifest = json.loads((root / "fixture_manifest.json").read_text(encoding="utf-8"))
    errors = _unchanged_inputs(root, manifest)
    if target_directory.exists():
        errors.append("dry-run created the target directory")
    unmapped = bool(manifest.get("include_unmapped"))
    result = load_result(result_path, _ROW_CONTENT, errors, naming=_NAMING_UNMAPPED if unmapped else ())
    expected_summary = contract.RunSummary(
        figures=[
            contract.Figure(label="Wiersze danych", value=3, unit="count"),
            contract.Figure(label="Planowane szablony", value=2, unit="count"),
            contract.Figure(label="Puste szablony", value=1, unit="count"),
        ],
        counts=None,
        files=[],
        worker_settlements=WORKER_SETTLEMENTS,
    )
    if result is not None:
        if result.result != ("WYMAGA_SPRAWDZENIA" if unmapped else "OK") or result.stop is not None:
            errors.append("dry-run result is not the expected finished result")
        if result.summary != expected_summary:
            errors.append("dry-run plan does not list every worker file with its rows")
        expected_warnings = [("BRAK_MAPOWANIA", _UNMAPPED_WARNING)] if unmapped else []
        if [(warning.code, warning.what) for warning in result.warnings] != expected_warnings:
            errors.append("dry-run warnings do not name each unmapped worker once")
        problems = dict(figures(result.progress.figures)).get("Problemy")
        if problems != (1 if unmapped else 0):
            errors.append(f"dry-run reports {problems} Problemy")
    _check_metrics(root, [("DRY-RUN", "WYMAGA_SPRAWDZENIA" if unmapped else "OK")], errors)

    return finish(
        errors,
        "ASSERT OK | mode=dry-run | output_folder_absent=true | "
        "inputs_unchanged=true | rows_hidden=true | metrics_safe=true",
    )


def _assert_run(root: Path, result_path: Path | None) -> int:
    target_directory = root / PERIOD / f"Rozliczenie pracowników {PERIOD}"
    manifest = json.loads((root / "fixture_manifest.json").read_text(encoding="utf-8"))
    errors = _unchanged_inputs(root, manifest)
    expected_files = {f"Rozliczenie {PERIOD} - {worker}.xlsx" for worker in WORKERS}
    if not target_directory.is_dir():
        errors.append("published folder is missing")
    elif {path.name for path in target_directory.glob("*.xlsx")} != expected_files:
        errors.append("published folder does not contain exactly the configured workbooks")

    try:
        adrian = _read_workbook_snapshot(_path(root, WORKERS[0]))
        darek = _read_workbook_snapshot(_path(root, WORKERS[1]))
        kamil = _read_workbook_snapshot(_path(root, WORKERS[2]))
        adrian_file = _path(root, WORKERS[0])
        darek_file = _path(root, WORKERS[1])
        kamil_file = _path(root, WORKERS[2])
        for name, snapshot in zip(WORKERS, (adrian, darek, kamil), strict=True):
            if (snapshot.formula_au18, snapshot.formula_au19) != ("=N18", "=N19"):
                errors.append(f"formula changed for {name}")
            if (
                snapshot.sheet_names != ("Sheet1", "Rates")
                or snapshot.header_fill_rgb != "0000AA55"
                or snapshot.rate_value != 17.5
            ):
                errors.append(f"worker template structure or formatting changed for {name}")
        if (
            _read_input_row(adrian_file, 18) != _expected_source_row("adrian.testowy", "SYNTHETIC-1", 1)
            or _read_input_row(adrian_file, 19) != _expected_source_row("adrian.testowy", "SYNTHETIC-3", 3)
            or _read_input_row(darek_file, 18) != _expected_source_row("dariusz.fikcyjny2", "SYNTHETIC-2", 2)
            or _read_input_row(kamil_file, 18) != (None,) * 13 + (0,) + (None,) * 32
        ):
            errors.append("worker rows were not routed or the empty workbook was changed")
        if (
            _hyperlinks(adrian_file) != _order_links({18: 18, 19: 20})
            or _hyperlinks(darek_file) != _order_links({18: 19})
            or _hyperlinks(kamil_file) != {}
        ):
            errors.append("order number hyperlinks of the Plik zbiorczy did not reach the worker workbooks")
    except (OSError, ValueError, KeyError):
        errors.append("one or more published workbooks cannot be read")

    result = load_result(result_path, _ROW_CONTENT, errors)
    expected_summary = contract.RunSummary(
        figures=[
            contract.Figure(label="Zapisane szablony", value=2, unit="count"),
            contract.Figure(label="Puste szablony", value=1, unit="count"),
        ],
        counts=None,
        files=[],
        worker_settlements=WORKER_SETTLEMENTS,
    )
    if result is not None:
        if result.result != "OK" or result.stop is not None or result.warnings != [] or result.summary != expected_summary:
            errors.append("run result does not report a successful complete run")
        _check_progress(result.progress, errors)
    records = check_metrics(
        root / "metrics.jsonl", root, period=PERIOD, expected=[("DRY-RUN", "OK"), ("RUN", "OK")],
        row_content=_ROW_CONTENT, errors=errors,
    )
    if [record.get("routes") for record in records] != [_routes(0), _routes(len(WORKERS))]:
        errors.append("metrics records do not count each written Rozliczenie pracownika as a copy of the Wzorzec")

    return finish(
        errors,
        "ASSERT OK | mode=run | workbooks_checked=3 | routes_counted=true | "
        "formulas_preserved=true | hyperlinks_equal_plik_zbiorczy=F | worker_template_unchanged=true | "
        "progress_reported=true | metrics_safe=true",
    )


def _assert_no_publish(root: Path, result_path: Path | None, expected_stop: str, expected_what: str | None) -> int:
    period_directory = root / PERIOD
    target_directory = period_directory / f"Rozliczenie pracowników {PERIOD}"
    manifest = json.loads((root / "fixture_manifest.json").read_text(encoding="utf-8"))
    errors = _unchanged_inputs(root, manifest)
    if target_directory.exists():
        errors.append("failed run published the target directory")
    if list(period_directory.glob(f".{target_directory.name}.staging-*")):
        errors.append("failed run left a staging directory")
    naming = _NAMING_UNMAPPED if manifest.get("include_unmapped") else ()
    _check_stop(load_result(result_path, _ROW_CONTENT, errors, naming=naming), expected_stop, expected_what, errors)
    _check_stop_metrics(root, expected_stop, errors)
    return finish(
        errors,
        "ASSERT OK | mode=no-publish | output_folder_absent=true | inputs_unchanged=true | stop_recorded=true",
    )


def _assert_existing_output(root: Path, result_path: Path | None, expected_stop: str, expected_what: str | None) -> int:
    target_directory = root / PERIOD / f"Rozliczenie pracowników {PERIOD}"
    sentinel = target_directory / "existing-synthetic.txt"
    manifest = json.loads((root / "fixture_manifest.json").read_text(encoding="utf-8"))
    errors = _unchanged_inputs(root, manifest)
    if not target_directory.is_dir() or sentinel.read_text(encoding="utf-8") != "preserve":
        errors.append("existing output folder was changed")
    if {path.name for path in target_directory.iterdir()} != {sentinel.name}:
        errors.append("existing output folder gained or lost files")
    _check_stop(load_result(result_path, _ROW_CONTENT, errors), expected_stop, expected_what, errors)
    _check_stop_metrics(root, expected_stop, errors)
    return finish(
        errors,
        "ASSERT OK | mode=existing-output | existing_folder_unchanged=true | stop_recorded=true",
    )


def main(argv: list[str] | None = None) -> int:
    parser = Parser(description=__doc__)
    parser.add_root("run root of a settlement fixture")
    parser.add_argument(
        "--expect",
        choices=("dry-run", "run", "no-publish", "existing-output"),
        required=True,
    )
    parser.add_argument("--expected-stop", help="code of the Komunikat the run stops with (no-publish, existing-output)")
    parser.add_argument("--expected-what", help="text the stop's Komunikat must contain in what happened")
    parser.add_argument("--result", type=Path, help="evidence JSON of the process step")
    args = parser.parse_args(argv)
    root = args.root
    result_path = args.result.expanduser().resolve() if args.result else None
    try:
        if args.expect == "dry-run":
            return _assert_dry_run(root, result_path)
        if args.expect == "run":
            return _assert_run(root, result_path)
        if not args.expected_stop:
            parser.error(f"--expected-stop is required for {args.expect}")
        if args.expect == "no-publish":
            return _assert_no_publish(root, result_path, args.expected_stop, args.expected_what)
        return _assert_existing_output(root, result_path, args.expected_stop, args.expected_what)
    except (OSError, ValueError, KeyError) as exc:
        print(f"ASSERT FAIL: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
