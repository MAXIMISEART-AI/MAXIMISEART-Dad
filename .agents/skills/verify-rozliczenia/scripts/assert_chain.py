"""Assert the Przebieg rozliczeń of a chain run on the Plik zbiorczy from Przygotowanie, without printing row contents.

The Plik zbiorczy was written by Przygotowanie pliku zbiorczego and then filled by
the Gotowość step, so its Wiersze danych carry values in all of ``A:AT``, and the
Wyróżnienia and NUMER ZL. Hiperłącza Przygotowanie carried over from the Rejestr. A Rozliczenie pracownika
holds those rows without the Ilości tylko dla INEA (ADR 0018). The checks are generic over the Wzorce, so the
same helper serves both the synthetic and the real Wzorce. Messages name workers, rows, and columns, never a cell value.
"""

from __future__ import annotations

import hashlib
import json
from pathlib import Path
import sys
import zipfile

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

sys.path.insert(0, str(Path(__file__).resolve().parents[4] / "tools/lib"))

from wynik_uruchomienia import (
    Link,
    check_metrics,
    contract,
    external_link_parts,
    figures,
    finish,
    hyperlinks,
    load_result,
    look_mismatches,
)
from skrypt_skilla import Parser


PERIOD = "08_14_09_2026"
HEADER_ROW = 17
FIRST_DATA_ROW = 18
INPUT_COLUMNS = 11  # A:K
COPIED_COLUMNS = 46  # A:AT
LOOK_COLUMNS = 12  # A:L: Wygląd Wiersza danych, with the number format of the Plik zbiorczy cell
READINESS_COLUMNS = range(13, 46)  # N:AT as 0-based tuple indices
# Ilości tylko dla INEA (ADR 0018): column -> word in URZĄDZENIA (any case, anywhere in the text) at which the cell stays
# empty in the Rozliczenie pracownika. Written out here, not imported, so a broken rule in the product fails the check.
INEA_ONLY = {"AP": "gniazdo", "AK": "technicolor"}
ROW_MARKER = "PREP-SYN-"
INTERNAL_NAME = f"Rozliczenie wew. {PERIOD}.xlsx"
COLLECTIVE_NAME = f"Rozliczenie {PERIOD} - zbiorcze.xlsx"
TARGET_DIRECTORY_NAME = f"Rozliczenie pracowników {PERIOD}"
WORKER_TEMPLATE = Path("config") / "Wzorzec rozliczenia pracownika.xlsx"
# Zlecenia skuteczne of every Wykonawca in Rejestr order; Adrian's two are split by Kamil's.
EXPECTED_WORKER_ORDERS = {
    "Adrian Testowy": ("PREP-SYN-ZL-2", "PREP-SYN-ZL-7"),
    "Darek Fikcyjny": ("PREP-SYN-ZL-6", "PREP-SYN-ZL-8"),
    "Kamil Syntetyczny": ("PREP-SYN-ZL-4", "PREP-SYN-ZL-9"),
}
# Rejestr row of every Zlecenie skuteczne; its NUMER ZL. Hiperłącze ends with that row number.
ORDER_LINK = "https://przyklad.invalid/zlecenie?caseId=PREP-SYN-LINK&tw=1&wiersz="
REGISTER_ROWS_OF_ORDERS = {
    "PREP-SYN-ZL-2": 4,
    "PREP-SYN-ZL-4": 9,
    "PREP-SYN-ZL-7": 12,
    "PREP-SYN-ZL-6": 13,
    "PREP-SYN-ZL-8": 14,
    "PREP-SYN-ZL-9": 15,
}
EXPECTED_ROWS = sum(len(orders) for orders in EXPECTED_WORKER_ORDERS.values())
TEMPLATES = len(EXPECTED_WORKER_ORDERS)
WORKER_SETTLEMENTS = [
    contract.WorkerSettlement(
        worker=name, file=f"Rozliczenie {PERIOD} - {name}.xlsx", orders=len(orders), empty=not orders
    )
    for name, orders in EXPECTED_WORKER_ORDERS.items()
]
PHASES = ("Sprawdzanie", "Odczyt danych", "Planowanie", "Zapisywanie")
UNSUCCESSFUL_ORDER = "PREP-SYN-ZL-3"  # Darek's NIE order, present only in the Rozliczenie wewnętrzne
INPUTS = (
    ("Rejestr zleceń", Path("ZLECENIA INTERMAX OD 08_11_2021.xlsx"), "fixture_manifest.json", "register_sha256"),
    (
        "Wzorzec rozliczenia wewnętrznego",
        Path("config") / "Wzorzec rozliczenia wewnętrznego.xlsx",
        "fixture_manifest.json",
        "internal_template_sha256",
    ),
    (
        "Wzorzec pliku zbiorczego",
        Path("config") / "Wzorzec pliku zbiorczego.xlsx",
        "fixture_manifest.json",
        "collective_template_sha256",
    ),
    ("Wzorzec rozliczenia pracownika", WORKER_TEMPLATE, "fixture_manifest.json", "worker_template_sha256"),
    (
        "Rozliczenie wewnętrzne",
        Path("INEA") / PERIOD / INTERNAL_NAME,
        "readiness_manifest.json",
        "internal_settlement_sha256",
    ),
    ("Plik zbiorczy", Path("INEA") / PERIOD / COLLECTIVE_NAME, "readiness_manifest.json", "collective_file_sha256"),
)

Row = tuple[object, ...]
Looks = tuple[tuple[object, ...], ...]
PLAIN_LOOK = (None, None, None, False, False)


def _sheet(workbook: Workbook) -> Worksheet:
    sheet = workbook.active
    if not isinstance(sheet, Worksheet):
        raise ValueError("Workbook does not contain an active worksheet")
    return sheet


def _manifest(root: Path, name: str) -> dict[str, object]:
    loaded = json.loads((root / name).read_text(encoding="utf-8"))
    if not isinstance(loaded, dict):
        raise ValueError(f"{name} is not a JSON object")
    return loaded


def _unchanged_inputs(root: Path, errors: list[str]) -> None:
    """Rejestr, all Wzorce, and both period files are byte-identical to before the Przebieg rozliczeń."""

    manifests = {name: _manifest(root, name) for name in ("fixture_manifest.json", "readiness_manifest.json")}
    for label, path, manifest, key in INPUTS:
        if hashlib.sha256((root / path).read_bytes()).hexdigest() != manifests[manifest].get(key):
            errors.append(f"{label} changed")


def _check_metrics(root: Path, expected: list[tuple[str, str]], errors: list[str]) -> None:
    records = check_metrics(
        root / "settlement-metrics.jsonl", root, period=PERIOD, expected=expected, row_content=(ROW_MARKER,), errors=errors
    )
    for record in records:
        counters = record.get("counters") or {}
        if (counters.get("templates_total"), counters.get("rows")) != (TEMPLATES, EXPECTED_ROWS):
            errors.append("settlement metrics counters do not match the Plik zbiorczy")


def _assert_dry_run(root: Path, result_path: Path | None) -> int:
    errors: list[str] = []
    _unchanged_inputs(root, errors)
    period_directory = root / "INEA" / PERIOD
    if sorted(path.name for path in period_directory.iterdir()) != sorted([INTERNAL_NAME, COLLECTIVE_NAME]):
        errors.append("dry-run added entries to the Folder okresu")
    result = load_result(result_path, (ROW_MARKER,), errors)
    if result is not None:
        if result.result != "OK" or result.stop is not None or result.warnings != []:
            errors.append("dry-run result is not OK without warnings")
        if result.summary != contract.RunSummary(
            figures=[
                contract.Figure(label="Wiersze danych", value=EXPECTED_ROWS, unit="count"),
                contract.Figure(label="Planowane szablony", value=TEMPLATES, unit="count"),
                contract.Figure(label="Puste szablony", value=0, unit="count"),
            ],
            counts=None,
            files=[],
            worker_settlements=WORKER_SETTLEMENTS,
        ):
            errors.append("dry-run plan does not list every Wykonawca's file with its Zlecenia skuteczne")
        if dict(figures(result.progress.figures)).get("Problemy") != 0:
            errors.append("dry-run reports Problemy")
    _check_metrics(root, [("DRY-RUN", "OK")], errors)
    return finish(
        errors,
        "ASSERT OK | process=przebieg | source=przygotowanie | mode=dry-run | output_folder_absent=true | "
        "inputs_unchanged=true | plan_reported=true | rows_hidden=true",
    )


def _header_column(sheet: Worksheet, header: str) -> int:
    columns = [
        cell.column
        for cell in sheet[HEADER_ROW]
        if isinstance(cell.value, str) and cell.value.strip().casefold() == header.casefold()
    ]
    if len(columns) != 1:
        raise ValueError(f"Plik zbiorczy has {len(columns)} {header} headers in row {HEADER_ROW}")
    return columns[0]


def _inea_only(values: Row, devices_index: int) -> frozenset[int]:
    """0-based indices of the Ilości tylko dla INEA in a Plik zbiorczy row, from its URZĄDZENIA."""

    devices = str(values[devices_index] or "").casefold()
    return frozenset(column_index_from_string(column) - 1 for column, word in INEA_ONLY.items() if word in devices)


def _data_rows(sheet: Worksheet) -> list[tuple[int, Row]]:
    """Wiersze danych with their row numbers: rows from 18 with a value in A:K, as A:AT."""

    return [
        (number, tuple(values))
        for number, values in enumerate(
            sheet.iter_rows(min_row=FIRST_DATA_ROW, max_col=COPIED_COLUMNS, values_only=True),
            start=FIRST_DATA_ROW,
        )
        if any(value is not None for value in values[:INPUT_COLUMNS])
    ]


def _looks(sheet: Worksheet, in_file: set[tuple[int, int]], number: int) -> Looks:
    """Wyróżnienia of A:AT in row ``number``: fill, font color, bold, italic, and in A:L the number format; a cell
    not ``in_file`` has none.

    ``in_file`` is taken right after loading: reading a row through openpyxl adds its missing cells to the sheet.
    """

    looks = []
    for column in range(1, COPIED_COLUMNS + 1):
        if (number, column) not in in_file:
            looks.append(PLAIN_LOOK)
            continue
        cell = sheet.cell(number, column)
        fill_type = getattr(cell.fill, "fill_type", None)
        fill = (fill_type, cell.fill.fgColor if fill_type else None)
        number_format = (cell.number_format,) if column <= LOOK_COLUMNS else ()
        looks.append((*fill, cell.font.color, bool(cell.font.b), bool(cell.font.i), *number_format))
    return tuple(looks)


def _row_links(links: dict[tuple[int, int], Link], number: int) -> dict[int, Link]:
    return {column: link for (row, column), link in links.items() if row == number and column <= COPIED_COLUMNS}


def _outside_rows(links: dict[tuple[int, int], Link], rows: range) -> dict[tuple[int, int], Link]:
    return {cell: link for cell, link in links.items() if not (cell[0] in rows and cell[1] <= COPIED_COLUMNS)}


def _cell_values(workbook: Workbook, skipped_rows: range) -> dict[tuple[str, int, int], object]:
    """Every stored value except the A:AT cells of ``skipped_rows`` on the active sheet."""

    active = workbook.active
    return {
        (sheet.title, cell.row, cell.column): cell.value
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
        and not (sheet is active and cell.row in skipped_rows and cell.column <= COPIED_COLUMNS)
    }


def _check_worker_file(
    name: str,
    path: Path,
    expected_rows: list[Row],
    expected_empty: list[frozenset[int]],
    expected_looks: list[Looks],
    expected_links: list[dict[int, Link]],
    template: Workbook,
    template_links: dict[str, bytes],
    readiness_orders: set[object],
    order_index: int,
    errors: list[str],
) -> None:
    workbook = load_workbook(path, data_only=False, keep_links=True)
    try:
        sheet = _sheet(workbook)
        in_file = set(sheet._cells)
        links = hyperlinks(sheet)
        data_rows = _data_rows(sheet)
        written = range(FIRST_DATA_ROW, FIRST_DATA_ROW + len(expected_rows))
        if [number for number, _ in data_rows] != list(written):
            errors.append(f"{name}: Wiersze danych are not exactly rows {written.start}-{written.stop - 1}")
        actual_rows = [values for _, values in data_rows]
        if any(values[order_index] == UNSUCCESSFUL_ORDER for values in actual_rows):
            errors.append(f"{name}: holds a Zlecenie nieskuteczne")
        if [values[order_index] for values in actual_rows] != [values[order_index] for values in expected_rows]:
            errors.append(f"{name}: orders differ from its Zlecenia skuteczne in Plik zbiorczy order")
        # A missing or extra row already failed the row-number check above; compare the rows both files have.
        for number, (actual, expected, empty, looks, row_links) in enumerate(
            zip(actual_rows, expected_rows, expected_empty, expected_looks, expected_links, strict=False), start=FIRST_DATA_ROW
        ):
            if _row_links(links, number) != {column: link for column, link in row_links.items() if column - 1 not in empty}:
                errors.append(f"{name}: row {number} lacks the Hiperłącza of the Plik zbiorczy")
            differing = [
                get_column_letter(index + 1)
                for index, value in enumerate(expected)
                if (actual[index] is not None if index in empty else actual[index] != value)
            ]
            if differing:
                errors.append(f"{name}: row {number} differs from the Plik zbiorczy in {', '.join(differing[:8])}")
            unlike = [get_column_letter(index + 1) for index, look in enumerate(_looks(sheet, in_file, number)) if look != looks[index]]
            if unlike:
                errors.append(
                    f"{name}: row {number} lacks the Wyróżnienia or number format of the Plik zbiorczy in {', '.join(unlike[:8])}"
                )
            elif expected[order_index] in readiness_orders and any(
                actual[index] in (None, 0) for index in READINESS_COLUMNS if index not in empty
            ):
                errors.append(f"{name}: row {number} lacks the Gotowość entries in N:AT")
        look = look_mismatches(sheet, HEADER_ROW, written, LOOK_COLUMNS)
        if look:
            errors.append(f"{name}: Wygląd Wiersza danych differs: {'; '.join(look)}")
        if workbook.sheetnames != template.sheetnames or _cell_values(workbook, written) != _cell_values(
            template, written
        ):
            errors.append(f"{name}: cells outside the Wiersze danych differ from the Wzorzec rozliczenia pracownika")
        if _outside_rows(links, written) != _outside_rows(hyperlinks(_sheet(template)), written):
            errors.append(f"{name}: hyperlinks outside the Wiersze danych differ from the Wzorzec rozliczenia pracownika")
    finally:
        workbook.close()
    if external_link_parts(path) != template_links:
        errors.append(f"{name}: xl/externalLinks parts differ from the Wzorzec rozliczenia pracownika")


def _check_worker_files(root: Path, period_directory: Path, target_directory: Path, errors: list[str]) -> int:
    """Every Rozliczenie pracownika against the Plik zbiorczy it was made from and the Wzorzec rozliczenia pracownika.

    Returns the number of Ilości tylko dla INEA the Rozliczenia pracowników must hold empty.
    """

    filled_orders = _manifest(root, "readiness_manifest.json").get("filled_orders")
    if not isinstance(filled_orders, list) or not filled_orders:
        raise ValueError("readiness_manifest.json names no filled Wiersze danych")
    readiness_orders = set(filled_orders)
    collective = load_workbook(period_directory / COLLECTIVE_NAME, data_only=False, keep_links=True)
    try:
        collective_sheet = _sheet(collective)
        collective_in_file = set(collective_sheet._cells)
        collective_hyperlinks = hyperlinks(collective_sheet)
        order_index = _header_column(collective_sheet, "NUMER ZL.") - 1
        devices_index = _header_column(collective_sheet, "URZĄDZENIA") - 1
        numbered_rows = _data_rows(collective_sheet)
        collective_rows = [values for _, values in numbered_rows]
        collective_looks = [_looks(collective_sheet, collective_in_file, number) for number, _ in numbered_rows]
        collective_links = [_row_links(collective_hyperlinks, number) for number, _ in numbered_rows]
        highlighted = sum(look != PLAIN_LOOK and look[0] is not None for looks in collective_looks for look in looks)
    finally:
        collective.close()
    if not highlighted:
        errors.append("Plik zbiorczy has no filled cell in its Wiersze danych, so the Wyróżnienia check proves nothing")
    if {
        values[order_index]: links.get(order_index + 1)
        for values, links in zip(collective_rows, collective_links, strict=True)
    } != {order: (f"{ORDER_LINK}{row}", None) for order, row in REGISTER_ROWS_OF_ORDERS.items()}:
        errors.append("Plik zbiorczy NUMER ZL. cells lack the Hiperłącze of their Rejestr cell")
    inea_only = [_inea_only(values, devices_index) for values in collective_rows]
    if {index for indices in inea_only for index in indices} != {column_index_from_string(column) - 1 for column in INEA_ONLY}:
        errors.append("Plik zbiorczy lacks a Wiersz danych for every Ilość tylko dla INEA, so their check proves nothing")
    valueless = [
        f"{get_column_letter(index + 1)}{number}"
        for (number, values), indices in zip(numbered_rows, inea_only, strict=True)
        for index in sorted(indices)
        if values[index] is None
    ]
    if valueless:
        errors.append(f"Plik zbiorczy has no Ilość tylko dla INEA in {', '.join(valueless)}, so its emptying proves nothing")
    if UNSUCCESSFUL_ORDER in {values[order_index] for values in collective_rows}:
        errors.append("Plik zbiorczy holds a Zlecenie nieskuteczne")
    filled = [values for values in collective_rows if values[order_index] in readiness_orders]
    if len(filled) != len(readiness_orders) or any(
        values[index] in (None, 0) for values in filled for index in READINESS_COLUMNS
    ):
        errors.append("Plik zbiorczy lacks the Gotowość entries in N:AT")

    template_path = root / WORKER_TEMPLATE
    template_links = external_link_parts(template_path)
    template = load_workbook(template_path, data_only=False, keep_links=True)
    try:
        for name, orders in EXPECTED_WORKER_ORDERS.items():
            expected_rows = [values for values in collective_rows if values[order_index] in orders]
            expected_empty = [
                empty for values, empty in zip(collective_rows, inea_only, strict=True) if values[order_index] in orders
            ]
            expected_looks = [
                looks for values, looks in zip(collective_rows, collective_looks, strict=True) if values[order_index] in orders
            ]
            expected_links = [
                links for values, links in zip(collective_rows, collective_links, strict=True) if values[order_index] in orders
            ]
            if tuple(values[order_index] for values in expected_rows) != orders:
                errors.append(f"Plik zbiorczy does not hold the Zlecenia skuteczne of {name} in Rejestr order")
            _check_worker_file(
                name,
                target_directory / f"Rozliczenie {PERIOD} - {name}.xlsx",
                expected_rows,
                expected_empty,
                expected_looks,
                expected_links,
                template,
                template_links,
                readiness_orders,
                order_index,
                errors,
            )
    finally:
        template.close()
    return sum(len(indices) for indices in inea_only)


def _assert_run(root: Path, result_path: Path | None) -> int:
    errors: list[str] = []
    _unchanged_inputs(root, errors)
    period_directory = root / "INEA" / PERIOD
    target_directory = period_directory / TARGET_DIRECTORY_NAME
    if sorted(path.name for path in period_directory.iterdir()) != sorted(
        [INTERNAL_NAME, COLLECTIVE_NAME, TARGET_DIRECTORY_NAME]
    ):
        errors.append("Folder okresu does not hold exactly both period files and the Folder rozliczeń pracowników")
    expected_files = {f"Rozliczenie {PERIOD} - {name}.xlsx" for name in EXPECTED_WORKER_ORDERS}
    inea_only = 0
    if not target_directory.is_dir() or {path.name for path in target_directory.iterdir()} != expected_files:
        errors.append("Folder rozliczeń pracowników does not hold exactly one workbook per Wykonawca")
    else:
        inea_only = _check_worker_files(root, period_directory, target_directory, errors)

    result = load_result(result_path, (ROW_MARKER,), errors)
    if result is not None:
        progress = result.progress
        if [(phase.phase, phase.state) for phase in progress.phases] != [(phase, "completed") for phase in PHASES]:
            errors.append("run progress does not complete the four stages in order")
        if (
            result.result != "OK"
            or result.stop is not None
            or result.warnings != []
            or result.summary
            != contract.RunSummary(
                figures=[
                    contract.Figure(label="Zapisane szablony", value=TEMPLATES, unit="count"),
                    contract.Figure(label="Puste szablony", value=0, unit="count"),
                ],
                counts=None,
                files=[],
                worker_settlements=WORKER_SETTLEMENTS,
            )
            or (progress.template_total, progress.templates_completed) != (TEMPLATES, TEMPLATES)
        ):
            errors.append("run result does not report a complete run with every template written")
    _check_metrics(root, [("DRY-RUN", "OK"), ("RUN", "OK")], errors)
    return finish(
        errors,
        f"ASSERT OK | process=przebieg | source=przygotowanie | mode=run | workbooks_checked={len(expected_files)} | "
        f"rows_equal_plik_zbiorczy=A:AT | inea_only_empty={inea_only} | highlights_equal_plik_zbiorczy=A:AT | number_formats_equal_plik_zbiorczy=A:L | "
        "row_look=A:L | hyperlinks_equal_plik_zbiorczy=A:AT | "
        "hyperlinks_from_rejestr=NUMER ZL. | "
        "readiness_carried=true | unsuccessful_absent=true | "
        "wzorzec_preserved=true | external_links_preserved=true | inputs_unchanged=true",
    )


def main(argv: list[str] | None = None) -> int:
    parser = Parser(description=__doc__)
    parser.add_root("run root of a chain fixture")
    parser.add_argument("--expect", choices=("dry-run", "run"), required=True)
    parser.add_argument("--result", type=Path, help="evidence JSON of the Przebieg rozliczeń step")
    args = parser.parse_args(argv)
    root = args.root
    result_path = args.result.expanduser().resolve() if args.result else None
    try:
        if args.expect == "dry-run":
            return _assert_dry_run(root, result_path)
        return _assert_run(root, result_path)
    except (OSError, ValueError, KeyError, zipfile.BadZipFile) as exc:
        print(f"ASSERT FAIL: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
