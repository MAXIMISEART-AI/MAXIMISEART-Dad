"""Drive a running Centrum rozliczeń in Chromium the way Dad does and record the proof.

Each feature step pairs a user action with the resulting screen, the server's
answer, and the side effect on disk (settings file, picker dialogs, opened
paths). Screenshots and ARIA snapshots go to ``<evidence>/screenshots`` and
``<evidence>/aria``; API answers to ``<evidence>/api``; the run's state files
to ``<evidence>/state``.
"""

from __future__ import annotations

import argparse
from collections.abc import Callable
from datetime import datetime
import json
import os
from pathlib import Path
import shutil
import socket
import subprocess
import sys
import time
import tomllib
import urllib.error
import urllib.parse
import urllib.request
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from playwright.sync_api import Error as PlaywrightError
from playwright.sync_api import Locator, Page, expect, sync_playwright
import yaml

REPO_ROOT = Path(__file__).resolve().parents[4]
# Dad's computer runs on Polish time; dates in Centrum are shown in the browser's time zone.
DAD_TIME_ZONE = ZoneInfo("Europe/Warsaw")
# The period the CLI step names: the fixture's last complete period, without a Folder okresu.
LAST_COMPLETE_FOR_CLI = "06_12_10_2026"

# Przygotowanie pliku zbiorczego appends to the metrics files that Metryki counts, so it runs after
# Metryki; Przebieg rozliczeń continues Dad's path from its Plik zbiorczy; start-i-zamykanie closes
# Centrum, so it stays last.
FEATURES = (
    "ustawienia",
    "tablica-okresow",
    "konfiguracja",
    "metryki",
    "przygotowanie-pliku-zbiorczego",
    "przebieg-rozliczen",
    "start-i-zamykanie",
)


class DriveFailure(Exception):
    pass


class Run:
    def __init__(self, root: Path, url: str, evidence: Path) -> None:
        self.root = root
        self.url = url
        self.evidence = evidence
        self.fixture = json.loads((root / "fixture.json").read_text(encoding="utf-8"))
        for name in ("screenshots", "aria", "api", "state"):
            (evidence / name).mkdir(parents=True, exist_ok=True)

    def step(self, feature: str, name: str) -> None:
        print(f"STEP {feature} | {name}", flush=True)

    def check(self, condition: bool, message: str) -> None:
        if not condition:
            raise DriveFailure(message)
        print(f"  ok: {message}", flush=True)

    def capture(self, page: Page, name: str, *, full_page: bool = True, live: bool = False) -> None:
        """Screenshot and ARIA snapshot after every running transition has finished.

        ``live`` captures while a stage runs: its indeterminate bar loops forever, so do not wait.
        """

        if not live:
            page.wait_for_function(
                "document.getAnimations().every((animation) => animation.playState !== 'running')"
            )
        page.screenshot(path=self.evidence / "screenshots" / f"{name}.png", full_page=full_page)
        snapshot = page.locator("body").aria_snapshot()
        (self.evidence / "aria" / f"{name}.txt").write_text(snapshot + "\n", encoding="utf-8")

    def api(self, path: str, name: str) -> dict[str, object]:
        with urllib.request.urlopen(self.url.rstrip("/") + path, timeout=10) as response:
            body = json.load(response)
        (self.evidence / "api" / f"{name}.json").write_text(
            json.dumps(body, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )
        assert isinstance(body, dict)
        return body

    def refused(self, request: urllib.request.Request, name: str) -> int:
        try:
            with urllib.request.urlopen(request, timeout=10) as response:
                status, body = int(response.status), response.read()
        except urllib.error.HTTPError as error:
            status, body = error.code, error.read()
        record = {"status": status, "body": json.loads(body or b"null")}
        (self.evidence / "api" / f"{name}.json").write_text(
            json.dumps(record, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )
        return status

    def answer_picker(self, path: str) -> None:
        (self.root / "picker" / "answer.txt").write_text(path, encoding="utf-8")

    def lines(self, name: str) -> list[dict[str, str]]:
        log = self.root / name
        if not log.exists():
            return []
        return [json.loads(line) for line in log.read_text(encoding="utf-8").splitlines() if line]

    def settings_text(self) -> str:
        settings = Path(self.fixture["settings"])
        return settings.read_text(encoding="utf-8") if settings.exists() else ""

    def remembered(self, key: str) -> object:
        """One value of the settings file as Centrum wrote it (``folder_okresow``, ``rejestr_zlecen``, ``motyw``)."""

        document = yaml.safe_load(self.settings_text()) or {}
        assert isinstance(document, dict)
        return document.get(key)

    def wait_for(self, condition: Callable[[], bool], message: str, timeout_s: float = 10) -> None:
        deadline = time.monotonic() + timeout_s
        while not condition():
            if time.monotonic() > deadline:
                raise DriveFailure(f"timed out: {message}")
            time.sleep(0.1)
        print(f"  ok: {message}", flush=True)


def _choose_periods_directory(run: Run, page: Page) -> None:
    run.answer_picker(run.fixture["periods_directory"])
    page.get_by_role("button", name="Wskaż folder okresów", exact=True).click()
    expect(page.get_by_role("table", name="Tablica okresów")).to_be_visible()


def drive_settings(run: Run, page: Page) -> None:
    feature = "ustawienia"
    register_name = Path(run.fixture["order_register"]).name

    run.step(feature, "ustawienia-pusty-stan")
    page.goto(run.url)
    expect(page.get_by_text("Wskaż folder okresów", exact=True).first).to_be_visible()
    settings = run.api("/api/settings", "ustawienia-przed")
    run.check(
        settings["order_register"] == {"state": "not_set"} and settings["periods_directory"] == {"state": "not_set"},
        "no settings file means nothing is set",
    )
    run.capture(page, "01-ustawienia-pusty-stan")

    run.step(feature, "ustawienia-folder-okresow")
    _choose_periods_directory(run, page)
    run.check(
        run.lines("dialogs.jsonl")[-1]["title"] == "Wybierz folder okresów (np. INEA)",
        "the folder picker opened with the CLI's title",
    )
    run.check(
        run.remembered("folder_okresow") == run.fixture["periods_directory"],
        "the settings file remembers the folder okresów",
    )
    run.capture(page, "02-ustawienia-folder-okresow")

    run.step(feature, "ustawienia-rejestr-zlecen")
    run.answer_picker(run.fixture["order_register"])
    page.get_by_role("button", name="Wskaż Rejestr zleceń", exact=True).click()
    register = page.get_by_role("group", name="Rejestr zleceń")
    expect(register.get_by_text(register_name)).to_be_visible()
    run.check(run.lines("dialogs.jsonl")[-1]["title"] == "Wybierz plik Rejestru zleceń", "the file picker opened")
    run.check(
        run.remembered("rejestr_zlecen") == run.fixture["order_register"],
        "the settings file remembers the Rejestr zleceń",
    )
    run.capture(page, "03-ustawienia-rejestr-zlecen")

    run.step(feature, "ustawienia-anulowanie")
    before = run.settings_text()
    dialogs = len(run.lines("dialogs.jsonl"))
    page.get_by_role("button", name="Zmień Rejestr zleceń", exact=True).click()
    run.wait_for(lambda: len(run.lines("dialogs.jsonl")) == dialogs + 1, "the cancelled picker was shown")
    expect(page.get_by_role("button", name="Zmień Rejestr zleceń", exact=True)).to_be_enabled()
    expect(register.get_by_text(register_name)).to_be_visible()
    run.check(run.settings_text() == before, "a cancelled picker leaves the settings file unchanged")

    run.step(feature, "ustawienia-nieistniejaca-sciezka")
    moved = Path(run.fixture["order_register"] + ".przeniesiony")
    Path(run.fixture["order_register"]).rename(moved)
    try:
        page.reload()
        expect(register.get_by_text(f"Nie ma już: {register_name}")).to_be_visible()
        run.check(run.api("/api/settings", "ustawienia-brak-rejestru")["order_register"] == {
            "state": "missing",
            "path": run.fixture["order_register"],
            "name": register_name,
        }, "a remembered Rejestr zleceń that is gone is reported as missing")
        run.capture(page, "04-ustawienia-brak-rejestru")
    finally:
        moved.rename(run.fixture["order_register"])
    page.reload()
    expect(register.get_by_text(register_name, exact=True)).to_be_visible()

    run.step(feature, "ustawienia-motyw")
    page.get_by_role("button", name="Włącz motyw jasny").click()
    expect(page.locator("html")).not_to_have_class("dark")
    run.wait_for(lambda: run.remembered("motyw") == "light", "the settings file remembers the light theme")
    page.reload()
    expect(page.get_by_role("button", name="Włącz motyw ciemny")).to_be_visible()
    expect(page.locator("html")).not_to_have_class("dark")
    run.capture(page, "05-ustawienia-motyw-jasny")
    page.get_by_role("button", name="Włącz motyw ciemny").click()
    expect(page.locator("html")).to_have_class("dark")
    run.wait_for(lambda: run.remembered("motyw") == "dark", "the dark theme is remembered again")


def _row(page: Page, token: str) -> Locator:
    return page.get_by_role("table", name="Tablica okresów").get_by_role("row").filter(has_text=token)


def _last_opened_is(run: Run, path: Path) -> Callable[[], bool]:
    return lambda: run.lines("opened.jsonl")[-1:] == [{"path": str(path)}]


def drive_period_board(run: Run, page: Page) -> None:
    feature = "tablica-okresow"
    periods = Path(run.fixture["periods_directory"])

    run.step(feature, "tablica-kolejnosc-i-status")
    page.goto(run.url)
    if run.api("/api/settings", "tablica-ustawienia")["periods_directory"] == {"state": "not_set"}:
        _choose_periods_directory(run, page)
    board = run.api("/api/periods", "tablica-okresow")
    rows = board["periods"]
    assert isinstance(rows, list)
    expected = run.fixture["board"]
    run.check(
        [[row["token"], row["status"], row["worker_settlement_count"], row["last_complete"]] for row in rows]
        == expected,
        "the server lists every Folder okresu newest first with its Status okresu",
    )
    table_rows = page.get_by_role("table", name="Tablica okresów").get_by_role("row")
    expect(table_rows).to_have_count(len(expected) + 1)
    labels = {"no_period_files": "Brak plików okresu", "collective_file_prepared": "Plik zbiorczy przygotowany",
              "worker_settlements_ready": "Rozliczenia pracowników gotowe"}
    for index, (token, status, _, last_complete) in enumerate(expected, start=1):
        row = table_rows.nth(index)
        expect(row).to_contain_text(token)
        expect(row).to_contain_text(labels[status])
        if last_complete:
            expect(row).to_have_attribute("data-state", "selected")
            expect(row).to_contain_text("ostatni pełny okres")
    run.check(True, "the screen shows the same order, Status okresu, and highlighted last complete period")
    run.capture(page, "06-tablica-okresow")

    run.step(feature, "tablica-nastepny-krok")
    next_step = page.get_by_role("region", name="Następny krok")
    expect(next_step).to_contain_text("Przygotuj plik zbiorczy")
    expect(next_step).to_contain_text(expected[0][0])
    expect(next_step).to_contain_text("Brak Folderu okresu")
    run.check(True, "Następny krok targets the last complete period, whose Folder okresu does not exist")

    run.step(feature, "tablica-otwieranie")
    opened_before = len(run.lines("opened.jsonl"))
    clicks = [
        ("Plik zbiorczy, okres 29_05_09_2026", periods / "29_05_09_2026" / "Rozliczenie 29_05_09_2026 - zbiorcze.xlsx"),
        ("Rozliczenie wewnętrzne, okres 22_28_09_2026", periods / "22_28_09_2026" / "Rozliczenie wew. 22_28_09_2026.xlsx"),
        ("Folder okresu, okres 15_21_09_2026", periods / "15_21_09_2026"),
        ("Rozliczenia pracowników: 3, okres 22_28_09_2026",
         periods / "22_28_09_2026" / "Rozliczenie pracowników 22_28_09_2026"),
    ]
    for name, path in clicks:
        page.get_by_role("button", name=name, exact=True).click()
        run.wait_for(_last_opened_is(run, path), f"{name} opened {path.name}")
        expect(page.get_by_text(path.name, exact=True).last).to_be_visible()
    run.check(len(run.lines("opened.jsonl")) == opened_before + len(clicks), "one open per click, nothing else")
    run.capture(page, "07-tablica-otwieranie", full_page=False)
    expect(page.get_by_role("button", name="Plik zbiorczy, okres 15_21_09_2026", exact=True)).to_be_disabled()
    run.check(True, "a period file that does not exist cannot be opened")

    run.step(feature, "tablica-zmiana-na-dysku")
    changed = periods / "15_21_09_2026"
    for name in ("Rozliczenie wew. 15_21_09_2026.xlsx", "Rozliczenie 15_21_09_2026 - zbiorcze.xlsx"):
        (changed / name).write_bytes(b"")
    page.reload()
    expect(_row(page, "15_21_09_2026")).to_contain_text("Plik zbiorczy przygotowany")
    run.check(True, "after a reload the board shows a Plik zbiorczy added outside Centrum")
    (changed / "Rozliczenie pracowników 15_21_09_2026").mkdir()
    # Headless Chromium keeps every tab visible, so returning to the tab is simulated with the
    # `focus` event the browser fires then; the page itself is not touched.
    page.evaluate("window.dispatchEvent(new FocusEvent('focus'))")
    expect(_row(page, "15_21_09_2026")).to_contain_text("Rozliczenia pracowników gotowe")
    run.check(True, "a (simulated) return to the tab shows a Folder rozliczeń pracowników added outside Centrum")
    run.capture(page, "08-tablica-po-zmianie-na-dysku")

    run.step(feature, "tablica-podpowiedz")
    _row(page, "29_05_09_2026").get_by_text("Utwórz rozliczenia pracowników").hover()
    expect(page.get_by_text("Najpierw uzupełnij Plik zbiorczy w Excelu, zapisz go i zamknij.")).to_be_visible()
    run.capture(page, "09-tablica-podpowiedz", full_page=False)

    run.step(feature, "tablica-ochrona")
    opened = len(run.lines("opened.jsonl"))
    open_url = run.url.rstrip("/") + "/api/periods/22_28_09_2026/open"
    payload = json.dumps({"target": "collective_file"}).encode()
    foreign_origin = urllib.request.Request(
        open_url, data=payload, method="POST",
        headers={"Content-Type": "application/json", "Origin": "http://zly.example"},
    )
    foreign_host = urllib.request.Request(run.url.rstrip("/") + "/api/health", headers={"Host": "zly.example"})
    run.check(run.refused(foreign_origin, "odmowa-obcy-origin") == 403, "an action from a foreign Origin is refused")
    run.check(run.refused(foreign_host, "odmowa-obcy-host") == 403, "a request with a foreign Host is refused")
    run.check(len(run.lines("opened.jsonl")) == opened, "a refused request opens nothing")


def _seconds(duration_ms: int) -> str:
    """The UI's duration text: tenths of a second rounded half up, Polish comma (``7250`` → ``7,3 s``)."""

    tenths = (duration_ms + 50) // 100
    return f"{tenths // 10},{tenths % 10} s"


def _local(started_at: str) -> str:
    """The UI's date text in Dad's time zone (``06.10.2026, 08:21``)."""

    return datetime.fromisoformat(started_at).astimezone(DAD_TIME_ZONE).strftime("%d.%m.%Y, %H:%M")


def _history_rows(page: Page) -> Locator:
    return page.get_by_role("table", name="Historia Uruchomień procesu").get_by_role("row")


def drive_metrics(run: Run, page: Page) -> None:
    feature = "metryki"
    fixture_metrics = run.fixture["metrics"]

    run.step(feature, "metryki-karty-okresow")
    page.goto(run.url + "#/okresy")
    if run.api("/api/settings", "metryki-ustawienia")["periods_directory"] == {"state": "not_set"}:
        _choose_periods_directory(run, page)
    metrics = run.api("/api/metrics", "metryki")
    runs = metrics["runs"]
    recent = metrics["recent"]
    assert isinstance(runs, list) and isinstance(recent, dict)
    run.check(len(runs) == fixture_metrics["runs"], "the server reads every valid record of both metrics files")
    run.check(
        {entry["process"] for entry in runs} == {"preparation", "settlement"},
        "the history holds both processes, and the records the CLI's reader skips are absent",
    )
    stats = page.get_by_role("group", name="Najważniejsze liczby")
    board = run.api("/api/periods", "metryki-tablica")["periods"]
    assert isinstance(board, list)
    prepared = next(period for period in board if period["collective_file_exists"])
    expect(stats.get_by_role("group", name="Ostatni przygotowany okres")).to_contain_text(prepared["token"])
    latest = runs[0]
    latest_card = stats.get_by_role("group", name="Ostatnie uruchomienie")
    expect(latest_card).to_contain_text(f"{latest['process_label']} · {_local(latest['started_at'])}")
    expect(latest_card).to_contain_text("OK")
    results = recent["results"]
    recent_card = stats.get_by_role("group", name=f"Ostatnie {recent['weeks']} tygodni")
    expect(recent_card).to_contain_text(f"{results['ok']} OK")
    expect(recent_card).to_contain_text(f"{results['needs_review']} do sprawdzenia · {results['critical']} zatrzymane")
    run.check(True, "Okresy shows the last prepared period, the last Uruchomienie procesu and the recent results")
    run.capture(page, "20-metryki-karty-okresow")

    run.step(feature, "metryki-p50-p95")
    page.get_by_role("link", name="Metryki", exact=True).click()
    expect(page.get_by_role("heading", name="Metryki", exact=True)).to_be_visible()
    entries = metrics["statistics"]
    assert isinstance(entries, list)
    statistics = {(entry["process"], entry["mode"]): entry for entry in entries}
    preparation_run = statistics[("preparation", "RUN")]
    percentiles = preparation_run["percentiles"]
    run.check(
        preparation_run["completed_runs"] >= 5 and percentiles is not None,
        "Przygotowanie RUN has five completed observations, so P50/P95 exist",
    )
    strip = page.get_by_role("group", name="Statystyki")
    expect(strip.get_by_role("group", name="P50")).to_contain_text(_seconds(percentiles["p50_ms"]))
    expect(strip.get_by_role("group", name="P95")).to_contain_text(_seconds(percentiles["p95_ms"]))
    counts = preparation_run["results"]
    expect(strip.get_by_role("group", name="Wyniki")).to_contain_text(f"{counts['ok']} OK")
    expect(strip.get_by_role("group", name="Wyniki")).to_contain_text(f"{counts['critical']} zatrzymane")
    run.check(True, "the strip shows the server's P50/P95 and result counts for Przygotowanie RUN")

    run.step(feature, "metryki-wykres")
    chart = page.get_by_role("figure", name="Wykres czasów etapów")
    expect(chart).to_be_visible()
    shown_runs = [entry for entry in runs if (entry["process"], entry["mode"]) == ("preparation", "RUN")][:12]
    bars = chart.locator(".recharts-bar-rectangle")
    expect(bars).to_have_count(len(shown_runs) * 4 - sum(
        1 for entry in shown_runs for phase in entry["phase_durations"] if phase["duration_ms"] == 0
    ))
    for phase in ("Sprawdzanie", "Odczyt danych", "Planowanie", "Zapisywanie"):
        expect(chart.get_by_text(phase, exact=True)).to_be_visible()
    run.check(True, "the stage-times chart stacks the four phases of each Przygotowanie RUN")
    run.capture(page, "21-metryki-wykres")

    run.step(feature, "metryki-tabela-etapow")
    page.get_by_role("tab", name="Tabela", exact=True).click()
    table = page.get_by_role("table", name="Czasy etapów")
    expect(table.get_by_role("row")).to_have_count(1 + len(shown_runs) + 2)
    p50_row = table.get_by_role("row").filter(has_text="P50")
    expect(p50_row).to_contain_text(_seconds(percentiles["p50_ms"]))
    for phase in percentiles["phases"]:
        if any(p["phase"] == phase["phase"] and p["duration_ms"] for e in shown_runs for p in e["phase_durations"]):
            expect(p50_row).to_contain_text(_seconds(phase["p50_ms"]))
    run.check(True, "the table view lists each run's phase times with per-phase P50/P95 from the server")
    run.capture(page, "22-metryki-tabela-etapow")
    page.get_by_role("tab", name="Wykres", exact=True).click()

    run.step(feature, "metryki-za-malo-obserwacji")
    page.get_by_role("group", name="Tryb", exact=True).get_by_role("button", name="DRY-RUN", exact=True).click()
    dry_run = statistics[("preparation", "DRY-RUN")]
    run.check(
        dry_run["percentiles"] is None and dry_run["completed_runs"] < 5,
        "Przygotowanie DRY-RUN has fewer than five completed observations",
    )
    expect(strip.get_by_role("group", name="P50")).to_contain_text(f"od 5 ukończonych (jest {dry_run['completed_runs']})")
    run.check(True, "below five completed observations P50/P95 show a dash and the count")
    run.capture(page, "23-metryki-za-malo-obserwacji")
    page.get_by_role("group", name="Tryb", exact=True).get_by_role("button", name="RUN", exact=True).click()

    run.step(feature, "metryki-przerwane")
    period, phase_times = fixture_metrics["interrupted"]
    interrupted = page.get_by_role("region", name="Przerwane")
    expect(interrupted).to_contain_text(period)
    expect(interrupted).to_contain_text(
        f"Sprawdzanie {_seconds(phase_times[0])} · Odczyt danych {_seconds(phase_times[1])}"
    )
    expect(interrupted).to_contain_text(f"do przerwania {_seconds(sum(phase_times))}")
    run.check(True, "an interrupted Uruchomienie procesu shows its phase times until the interruption")

    run.step(feature, "metryki-historia")
    rows = _history_rows(page)
    expect(rows).to_have_count(1 + min(10, len(runs)))
    expect(rows.nth(1)).to_contain_text(_local(runs[0]["started_at"]))
    dry_runs_on_page = sum(1 for entry in runs[:10] if entry["mode"] == "DRY-RUN")
    history = page.get_by_role("table", name="Historia Uruchomień procesu")
    expect(history.get_by_text("DRY-RUN", exact=True)).to_have_count(dry_runs_on_page)
    expect(history.get_by_text(".cmd", exact=True)).to_have_count(min(10, len(runs)))
    run.capture(page, "24-metryki-historia")
    history.get_by_role("button", name="Czas", exact=True).click()
    longest = max(runs, key=lambda entry: entry["total_duration_ms"])
    expect(rows.nth(1)).to_contain_text(_seconds(longest["total_duration_ms"]))
    expect(history.get_by_role("columnheader", name="Czas")).to_have_attribute("aria-sort", "descending")
    page.get_by_role("group", name="Proces w historii").get_by_role("button", name="Przebieg", exact=True).click()
    settlements = [entry for entry in runs if entry["process"] == "settlement"]
    expect(rows).to_have_count(1 + min(10, len(settlements)))
    for index in range(1, min(10, len(settlements)) + 1):
        expect(rows.nth(index)).to_contain_text("Przebieg rozliczeń")
    page.get_by_role("group", name="Proces w historii").get_by_role("button", name="Wszystkie", exact=True).click()
    page.get_by_role("button", name="Pokaż starsze", exact=True).click()
    expect(rows).to_have_count(1 + min(20, len(runs)))
    run.check(True, "the history sorts by time, filters by process and shows older runs on demand")
    run.capture(page, "25-metryki-historia-filtry")

    run.step(feature, "metryki-nowy-rekord-z-cli")
    completed = subprocess.run(
        [sys.executable, str(REPO_ROOT / "run_preparation.py"),
         "--register", str(run.root / "brak-rejestru.xlsx"),
         "--periods-directory", run.fixture["periods_directory"],
         "--period", LAST_COMPLETE_FOR_CLI, "--dry-run",
         "--metrics", fixture_metrics["preparation"]],
        cwd=REPO_ROOT, capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=120, check=False,
    )
    run.check(completed.returncode == 1, "the CLI stopped (no Rejestr zleceń) and wrote its metrics record")
    after = run.api("/api/metrics", "metryki-po-cli")["runs"]
    assert isinstance(after, list)
    run.check(len(after) == len(runs) + 1, "the next read includes the CLI's record without restarting Centrum")
    page.reload()
    newest = _history_rows(page).nth(1)
    expect(newest).to_contain_text(_local(after[0]["started_at"]))
    expect(newest).to_contain_text(LAST_COMPLETE_FOR_CLI)
    expect(newest).to_contain_text("Zatrzymano")
    expect(newest).to_contain_text(".cmd")
    run.check(True, "the reloaded page shows the CLI's DRY-RUN as the newest Uruchomienie procesu")
    run.capture(page, "26-metryki-nowy-rekord-z-cli")

    run.step(feature, "metryki-pusty-stan")
    moved = [Path(fixture_metrics[name]) for name in ("preparation", "settlement")]
    for path in moved:
        path.rename(path.with_name(path.name + ".odlozony"))
    try:
        run.check(run.api("/api/metrics", "metryki-brak-plikow")["runs"] == [], "no metrics files is an empty list")
        page.reload()
        expect(page.get_by_text("Jeszcze nie ma metryk", exact=True)).to_be_visible()
        run.capture(page, "27-metryki-pusty-stan")
    finally:
        for path in moved:
            path.with_name(path.name + ".odlozony").rename(path)
TEMPLATE_LABELS = ("Wzorzec rozliczenia wewnętrznego", "Wzorzec pliku zbiorczego", "Wzorzec rozliczenia pracownika")
REPORT_TO_MAINTAINER = "Przekaż ten komunikat osobie utrzymującej narzędzie."


def _project_version() -> str:
    with (REPO_ROOT / "pyproject.toml").open("rb") as file:
        version = tomllib.load(file)["project"]["version"]
    assert isinstance(version, str)
    return version


def _configuration_snapshot(files: dict[str, Path]) -> dict[str, tuple[bytes, int]]:
    """Bytes and modification time of the konfiguracja wykonawców and the three Wzorce, plus every lock file."""

    config = next(iter(files.values())).parent
    paths = [*files.values(), *config.glob("~$*")]
    return {path.name: (path.read_bytes(), path.stat().st_mtime_ns) for path in sorted(paths)}


def _template_group(page: Page, label: str) -> Locator:
    return page.get_by_role("group", name=label, exact=True)


def _shown_as_ready(page: Page, label: str) -> None:
    expect(_template_group(page, label).get_by_role("img", name="Na miejscu i czytelny")).to_be_visible()


def _template_answer(run: Run, name: str, label: str) -> dict[str, object]:
    templates = run.api("/api/configuration", name)["templates"]
    assert isinstance(templates, list)
    return next(template for template in templates if template["label"] == label)


def drive_configuration(run: Run, page: Page) -> None:
    """Every broken state is made by renaming or adding files and undone by renaming back, so a
    Centrum that wrote nothing leaves bytes and modification times exactly as the fixture made them."""

    feature = "konfiguracja"
    files = {name: Path(path) for name, path in run.fixture["configuration"].items()}
    before = _configuration_snapshot(files)

    run.step(feature, "konfiguracja-gotowa")
    page.goto(run.url)
    page.get_by_role("link", name="Konfiguracja", exact=True).click()
    expect(page.get_by_role("heading", name="Konfiguracja", level=1)).to_be_visible()
    configuration = run.api("/api/configuration", "konfiguracja-gotowa")
    templates = configuration["templates"]
    assert isinstance(templates, list)
    run.check(
        [(template["label"], template["state"]) for template in templates]
        == [(label, "ready") for label in TEMPLATE_LABELS],
        "the server reports all three Wzorce present and readable",
    )
    run.check(
        configuration["workers"] == {"state": "ready", "names": run.fixture["workers"]},
        "the server lists the Wykonawcy in the order of the konfiguracja wykonawców",
    )
    version = _project_version()
    run.check(configuration["version"] == version, "the server reports the pyproject.toml version")
    for label in TEMPLATE_LABELS:
        _shown_as_ready(page, label)
    expect(page.get_by_text("3 z 3", exact=True)).to_be_visible()
    expect(page.get_by_role("list", name="Wykonawcy").get_by_role("listitem")).to_have_text(run.fixture["workers"])
    expect(page.get_by_text(f"Wersja {version}", exact=True)).to_be_visible()
    expect(page.get_by_role("alert")).to_have_count(0)
    run.check(True, "the screen shows the same Wzorce, Wykonawcy, and version, with no problem")
    run.capture(page, "10-konfiguracja-gotowa")

    run.step(feature, "konfiguracja-tylko-do-odczytu")
    page.get_by_text("Tylko do odczytu", exact=True).hover()
    expect(page.get_by_text("Zmiany robi osoba utrzymująca narzędzie.", exact=True)).to_be_visible()
    run.check(
        page.url.endswith("#/konfiguracja") and page.locator("main").get_by_role("button").count() == 2,
        "the page offers no action besides the header's menu and theme buttons",
    )
    run.capture(page, "11-konfiguracja-tylko-do-odczytu", full_page=False)
    page.mouse.move(0, 0)

    run.step(feature, "konfiguracja-wzorzec-otwarty")
    collective = files["collective_template"]
    lock = collective.with_name(f"~${collective.name}")
    lock.write_bytes(b"")
    try:
        page.reload()
        expect(_template_group(page, "Wzorzec pliku zbiorczego")).to_contain_text("Otwarty w Excelu")
        message = page.get_by_role("alert").filter(has_text="Wzorzec jest otwarty w Excelu.")
        expect(message).to_contain_text(f"plik {collective.name}, folder config")
        expect(message).to_contain_text(f"usuń z tego samego folderu pozostawiony plik ~${collective.name}")
        expect(page.get_by_text("2 z 3", exact=True)).to_be_visible()
        answer = _template_answer(run, "konfiguracja-wzorzec-otwarty", "Wzorzec pliku zbiorczego")
        run.check(answer["problem"] == "open_in_excel", "the server reports the Wzorzec as open in Excel")
        run.capture(page, "12-konfiguracja-wzorzec-otwarty")
        run.check(lock.exists(), "Centrum leaves Excel's lock file alone")
    finally:
        lock.unlink()
    # Headless Chromium keeps every tab visible, so returning to the tab is simulated with `focus`.
    page.evaluate("window.dispatchEvent(new FocusEvent('focus'))")
    _shown_as_ready(page, "Wzorzec pliku zbiorczego")
    run.check(True, "a (simulated) return to the tab shows the Wzorzec ready once Excel has closed it")

    run.step(feature, "konfiguracja-brak-wzorca")
    worker_template = files["worker_template"]
    moved = worker_template.with_name(worker_template.name + ".przeniesiony")
    worker_template.rename(moved)
    try:
        page.reload()
        expect(_template_group(page, "Wzorzec rozliczenia pracownika")).to_contain_text("Brak pliku")
        message = page.get_by_role("alert").filter(has_text=f"Nie znaleziono Wzorca rozliczenia pracownika: {worker_template.name}.")
        expect(message).to_contain_text(REPORT_TO_MAINTAINER)
        answer = _template_answer(run, "konfiguracja-brak-wzorca", "Wzorzec rozliczenia pracownika")
        run.check(answer["problem"] == "missing", "the server reports the Wzorzec rozliczenia pracownika as missing")
        run.capture(page, "13-konfiguracja-brak-wzorca")
    finally:
        moved.rename(worker_template)

    run.step(feature, "konfiguracja-nieczytelna-konfiguracja-wykonawcow")
    worker_config = files["worker_config"]
    moved = worker_config.with_name(worker_config.name + ".przeniesiony")
    worker_config.rename(moved)
    try:
        worker_config.write_text("workers: [niedomknięta lista\n", encoding="utf-8")
        page.reload()
        message = page.get_by_role("alert").filter(has_text="Nie można odczytać konfiguracji wykonawców.")
        expect(message).to_contain_text(f"plik {worker_config.name}")
        expect(message).to_contain_text("Przekaż to osobie utrzymującej narzędzie.")
        expect(page.get_by_role("list", name="Wykonawcy")).to_have_count(0)
        workers = run.api("/api/configuration", "konfiguracja-nieczytelna-konfiguracja-wykonawcow")["workers"]
        assert isinstance(workers, dict)
        run.check(workers["state"] == "unreadable", "the server reports the konfiguracja wykonawców as unreadable")
        run.capture(page, "14-konfiguracja-nieczytelna-konfiguracja-wykonawcow")
    finally:
        worker_config.unlink(missing_ok=True)
        moved.rename(worker_config)

    run.step(feature, "konfiguracja-bez-zmian")
    page.reload()
    for label in TEMPLATE_LABELS:
        _shown_as_ready(page, label)
    run.check(
        _configuration_snapshot(files) == before,
        "the konfiguracja wykonawców and every Wzorzec keep their bytes and modification times; no lock file is left",
    )


PREPARATION = "przygotowanie-pliku-zbiorczego"
# The fixture puts this marker in every synthetic workbook row; it must never reach the Centrum log.
ROW_MARKER = "CENTRUM-SYN-ROW"
INCOMPLETE_PERIOD_WHAT = "Ostatnie zlecenie okresu jest z soboty 10.10.2026"


def _period_files(run: Run) -> list[Path]:
    folder = Path(run.fixture["periods_directory"]) / LAST_COMPLETE_FOR_CLI
    return [
        folder / f"Rozliczenie wew. {LAST_COMPLETE_FOR_CLI}.xlsx",
        folder / f"Rozliczenie {LAST_COMPLETE_FOR_CLI} - zbiorcze.xlsx",
    ]


def _preparation_records(run: Run, name: str) -> list[dict[str, object]]:
    """Valid records of the CLI's Przygotowanie metrics file as the server reads them, newest first."""

    runs = run.api("/api/metrics", name)["runs"]
    assert isinstance(runs, list)
    return [entry for entry in runs if entry["process"] == "preparation"]


def _current_run(run: Run, name: str) -> dict[str, object]:
    view = run.api("/api/runs/current", name)
    assert isinstance(view, dict)
    return view


def _finished_result(run: Run, name: str) -> dict[str, object]:
    run.wait_for(lambda: _current_run(run, name)["state"] == "finished", "the Uruchomienie procesu finished")
    result = _current_run(run, name)["result"]
    assert isinstance(result, dict)
    return result


def _ensure_files_set(run: Run, page: Page) -> None:
    settings = run.api("/api/settings", "przygotowanie-ustawienia")
    if settings["periods_directory"] == {"state": "not_set"}:
        _choose_periods_directory(run, page)
    if settings["order_register"] == {"state": "not_set"}:
        run.answer_picker(run.fixture["order_register"])
        page.get_by_role("button", name="Wskaż Rejestr zleceń", exact=True).click()
        run.wait_for(lambda: run.remembered("rejestr_zlecen") == run.fixture["order_register"], "Rejestr zleceń set")


def drive_preparation(run: Run, page: Page) -> None:
    feature = PREPARATION
    expected = run.fixture["preparation"]
    period_folder = _period_files(run)[0].parent
    register_name = Path(run.fixture["order_register"]).name
    records_before = len(_preparation_records(run, "przygotowanie-metryki-przed"))

    run.step(feature, "przygotowanie-wybor-okresu")
    page.goto(run.url)
    _ensure_files_set(run, page)
    page.reload()
    page.get_by_role("button", name=f"Przygotuj plik zbiorczy, okres {LAST_COMPLETE_FOR_CLI}", exact=True).click()
    expect(page.get_by_text("Nowe Uruchomienie procesu", exact=True)).to_be_visible()
    run.check(page.url.endswith("#/uruchomienia"), "the row's next step opens Uruchomienia")
    expect(page.get_by_role("combobox", name="Okres")).to_contain_text(f"{LAST_COMPLETE_FOR_CLI} · ostatni pełny")
    expect(page.get_by_role("group", name="Rejestr zleceń")).to_contain_text(register_name)
    expect(page.get_by_text("Zapisz i zamknij Rejestr zleceń w Excelu", exact=True)).to_be_visible()
    setup = run.api("/api/runs/setup", "przygotowanie-wybor-okresu")["processes"]
    assert isinstance(setup, list)
    run.check(
        setup[0]["preselected_period"] == LAST_COMPLETE_FOR_CLI and setup[0]["blocker"] is None,
        "the server preselects the last complete period and nothing is missing",
    )
    run.capture(page, "30-przygotowanie-nowe-uruchomienie")

    run.step(feature, "przygotowanie-podglad")
    page.goto(run.url)
    page.get_by_role("region", name="Następny krok").get_by_role("button", name="Podgląd bez zapisu").click()
    live = page.get_by_role("region", name="Podgląd na żywo")
    expect(live).to_contain_text("Podgląd gotowy · nic nie zapisano")
    result = _finished_result(run, "przygotowanie-podglad")
    summary = result["summary"]
    assert isinstance(summary, dict)
    run.check(
        result["result"] == "OK"
        and summary["period_order_count"] == expected["period_order_count"]
        and summary["successful_order_count"] == expected["successful_order_count"]
        and [day["count"] for day in summary["orders_per_day"]] == expected["orders_per_day"],
        "the DRY-RUN plan has the Rejestr zleceń's counts and Zlecenia na dzień",
    )
    expect(live.get_by_role("listitem", name="sob 10.10.2026: 1", exact=True)).to_be_visible()
    expect(live.get_by_label("Ostrzeżenie").filter(has_text=INCOMPLETE_PERIOD_WHAT)).to_be_visible()
    expect(live).to_contain_text(f"Rozliczenie {LAST_COMPLETE_FOR_CLI} - zbiorcze.xlsx")
    run.check(not period_folder.exists(), "the DRY-RUN created neither files nor the Folder okresu")
    records = _preparation_records(run, "przygotowanie-metryki-podglad")
    run.check(
        len(records) == records_before + 1 and records[0]["mode"] == "DRY-RUN",
        "the DRY-RUN appended one record to the CLI's metrics file",
    )
    run.capture(page, "31-przygotowanie-podglad")

    run.step(feature, "przygotowanie-potwierdzenie")
    live.get_by_role("button", name="Uruchom", exact=True).click()
    confirm = page.get_by_role("alertdialog", name="Uruchomić z zapisem?")
    expect(confirm).to_contain_text("Zapisz i zamknij Rejestr zleceń w Excelu")
    expect(confirm.get_by_text("Nie było podglądu bez zapisu")).to_have_count(0)
    run.check(True, "Uruchom asks to confirm the write and reminds to close the Rejestr zleceń")
    run.capture(page, "32-przygotowanie-potwierdzenie")
    confirm.get_by_role("button", name="Uruchom", exact=True).click()

    run.step(feature, "przygotowanie-decyzja")
    decision = page.get_by_role("alertdialog", name="Kontynuować mimo niepełnego okresu?")
    expect(decision).to_contain_text(INCOMPLETE_PERIOD_WHAT)
    expect(decision.get_by_role("listitem", name="pon 12.10.2026: 0", exact=True)).to_be_visible()
    expect(decision.get_by_role("button", name="Przerwij bez zapisu")).to_be_focused()
    view = _current_run(run, "przygotowanie-decyzja")
    run.check(view["state"] == "awaiting_decision", "the server waits for Dad's decision")
    second = urllib.request.Request(
        run.url.rstrip("/") + "/api/runs",
        data=json.dumps({"process": "preparation", "period": LAST_COMPLETE_FOR_CLI, "dry_run": True}).encode(),
        method="POST",
        headers={"Content-Type": "application/json", "Origin": run.url.rstrip("/")},
    )
    run.check(run.refused(second, "przygotowanie-trwa-juz") == 409, "a second Uruchomienie procesu is refused")
    close = urllib.request.Request(
        run.url.rstrip("/") + "/api/shutdown", data=b"", method="POST", headers={"Origin": run.url.rstrip("/")}
    )
    run.check(
        run.refused(close, "przygotowanie-zamkniecie-odmowa") == 409,
        "closing Centrum is refused while the Uruchomienie procesu waits for the decision",
    )
    page.reload()
    expect(decision).to_be_visible()
    run.check(True, "the decision survives a reload of the tab")
    run.capture(page, "33-przygotowanie-decyzja")

    run.step(feature, "przygotowanie-przerwij")
    decision.get_by_role("button", name="Przerwij bez zapisu").click()
    stop = page.get_by_label("Zatrzymanie")
    expect(stop).to_contain_text("Przerwano z powodu niepełnego okresu")
    expect(stop).to_contain_text("Jeśli okres jest pełny, przy decyzji wybierz „Kontynuuj”.")
    result = _finished_result(run, "przygotowanie-przerwij")
    run.check(
        result["ending"] == "stop" and result["status_line"] == "Status końcowy: ZATRZYMANO",
        "the refusal is a stop with the three-part message",
    )
    run.check(not period_folder.exists(), "Przerwij bez zapisu wrote nothing")
    run.capture(page, "34-przygotowanie-zatrzymanie")

    run.step(feature, "przygotowanie-kontynuuj")
    stop.get_by_role("button", name="Uruchom ponownie").click()
    decision.get_by_role("button", name="Kontynuuj").click()
    expect(page.get_by_label("Wynik")).to_contain_text("Plik zbiorczy przygotowany")
    result = _finished_result(run, "przygotowanie-kontynuuj")
    run.check(result["result"] == "OK", "the RUN finished OK")
    run.check(all(path.is_file() for path in _period_files(run)), "both period files are in the Folder okresu")
    records = _preparation_records(run, "przygotowanie-metryki-zapisane")
    run.check(
        len(records) == records_before + 3
        and [(record["mode"], record["result"]) for record in records[:3]]
        == [("RUN", "OK"), ("RUN", "BLAD_KRYTYCZNY"), ("DRY-RUN", "OK")],
        "every Uruchomienie procesu appended one record to the CLI's metrics file",
    )
    log = (run.root / "logs" / "centrum-rozliczen.log").read_text(encoding="utf-8")
    run.check(
        f"uruchomienie-koniec proces=preparation tryb=RUN okres={LAST_COMPLETE_FOR_CLI} wynik=OK koniec=finished" in log
        and "uruchomienie-start proces=preparation tryb=RUN" in log
        and ROW_MARKER not in log,
        "the Centrum log records the start and end of each Uruchomienie procesu with safe facts only",
    )
    run.capture(page, "35-przygotowanie-zapisane")

    run.step(feature, "przygotowanie-status-okresu")
    page.goto(run.url)
    expect(_row(page, LAST_COMPLETE_FOR_CLI)).to_contain_text("Plik zbiorczy przygotowany")
    expect(page.get_by_role("region", name="Następny krok")).to_contain_text("Utwórz rozliczenia pracowników")
    run.check(True, "the Status okresu on Tablica okresów is Plik zbiorczy przygotowany")
    run.capture(page, "36-przygotowanie-status-okresu")

    run.step(feature, "przygotowanie-nienadpisywanie")
    before = [path.read_bytes() for path in _period_files(run)]
    page.goto(run.url + "#/uruchomienia")
    page.get_by_role("button", name="Nowe Uruchomienie procesu").click()
    page.get_by_role("button", name="Uruchom", exact=True).click()
    page.get_by_role("alertdialog", name="Uruchomić z zapisem?").get_by_role("button", name="Uruchom", exact=True).click()
    expect(page.get_by_label("Zatrzymanie")).to_contain_text("Plik tego okresu już istnieje, więc nie został nadpisany.")
    _finished_result(run, "przygotowanie-nienadpisywanie")
    run.check([path.read_bytes() for path in _period_files(run)] == before, "Nienadpisywanie: both files are unchanged")
    run.capture(page, "37-przygotowanie-nienadpisywanie")

    run.step(feature, "przygotowanie-historia")
    page.get_by_role("tab", name="Historia").click()
    history = run.api("/api/metrics", "przygotowanie-historia")["runs"]
    assert isinstance(history, list)
    run.check(
        [entry["origin"] for entry in history[:4]] == ["centrum"] * 4,
        "the metrics history marks the four Uruchomienia from Centrum without a schema change",
    )
    newest = _history_rows(page).nth(1)
    expect(newest).to_contain_text(LAST_COMPLETE_FOR_CLI)
    expect(newest).not_to_contain_text(".cmd")
    run.check(True, "the Historia tab shows the newest Uruchomienie without the .cmd badge")
    run.capture(page, "38-przygotowanie-historia")


SETTLEMENT = "przebieg-rozliczen"
WORKER_SETTLEMENTS = f"Rozliczenie pracowników {LAST_COMPLETE_FOR_CLI}"
# serve_centrum.py waits this long before each Szablon pracownika while the pacing file exists.
TEMPLATE_PACING_MS = 1500
WORKER_COLUMN, QUANTITY_COLUMN, FIRST_DATA_ROW = 8, 14, 18  # H (WYKONAWCA), N (first quantity), row 18


def _settlement_records(run: Run, name: str) -> list[dict[str, object]]:
    """Valid records of the CLI's Przebieg rozliczeń metrics file as the server reads them, newest first."""

    runs = run.api("/api/metrics", name)["runs"]
    assert isinstance(runs, list)
    return [entry for entry in runs if entry["process"] == "settlement"]


def _fill_readiness(collective_file: Path) -> int:
    """Gotowość as Dad reaches it in Excel: a quantity in N of every Wiersz danych, saved and closed.

    Excel is not driven; the harness edits the workbook like Dad would and leaves no lock file.
    Returns the number of filled Wiersze danych, never a cell value.
    """

    workbook = load_workbook(collective_file)
    sheet = workbook.active
    assert isinstance(sheet, Worksheet)
    filled = 0
    for row in range(FIRST_DATA_ROW, sheet.max_row + 1):
        if sheet.cell(row, WORKER_COLUMN).value:
            sheet.cell(row, QUANTITY_COLUMN).value = 1
            filled += 1
    workbook.save(collective_file)
    return filled


def _post(run: Run, path: str, body: dict[str, object], name: str) -> int:
    request = urllib.request.Request(
        run.url.rstrip("/") + path,
        data=json.dumps(body).encode(),
        method="POST",
        headers={"Content-Type": "application/json", "Origin": run.url.rstrip("/")},
    )
    return run.refused(request, name)


def drive_settlement(run: Run, page: Page) -> None:
    feature = SETTLEMENT
    internal_settlement, collective_file = _period_files(run)
    period_folder = collective_file.parent
    worker_settlements = period_folder / WORKER_SETTLEMENTS
    workers = run.fixture["workers"]
    files = [f"Rozliczenie {LAST_COMPLETE_FOR_CLI} - {worker}.xlsx" for worker in workers]
    needs_check = run.fixture["settlement"]["needs_check"]
    if not collective_file.is_file():
        # Dad's path starts with Przygotowanie pliku zbiorczego; alone, this feature drives it first.
        drive_preparation(run, page)
    records_before = len(_settlement_records(run, "przebieg-metryki-przed"))

    run.step(feature, "przebieg-gotowosc")
    filled = _fill_readiness(collective_file)
    run.check(
        filled == len(workers) and not collective_file.with_name(f"~${collective_file.name}").exists(),
        "Gotowość (Excel simulated by the harness): quantities typed into every Wiersz danych, saved and closed",
    )

    run.step(feature, "przebieg-wybor-okresu")
    page.goto(run.url)
    page.get_by_role("button", name=f"Utwórz rozliczenia pracowników, okres {LAST_COMPLETE_FOR_CLI}", exact=True).click()
    expect(page.get_by_text("Nowe Uruchomienie procesu", exact=True)).to_be_visible()
    expect(page.get_by_role("radio", name="Przebieg rozliczeń")).to_be_checked()
    expect(page.get_by_role("combobox", name="Okres")).to_contain_text(f"{LAST_COMPLETE_FOR_CLI} · ostatni pełny")
    expect(page.get_by_role("group", name="Plik zbiorczy")).to_contain_text(collective_file.name)
    expect(page.get_by_text("Zapisz i zamknij Plik zbiorczy w Excelu", exact=True)).to_be_visible()
    processes = run.api("/api/runs/setup", "przebieg-wybor-okresu")["processes"]
    assert isinstance(processes, list)
    [setup] = [process for process in processes if process["process"] == "settlement"]
    run.check(
        setup["preselected_period"] == LAST_COMPLETE_FOR_CLI
        and [choice["token"] for choice in setup["periods"]] == [LAST_COMPLETE_FOR_CLI, needs_check]
        and setup["blocker"] is None,
        "only periods with Status Plik zbiorczy przygotowany are offered; the Plik zbiorczy follows from the period",
    )
    run.capture(page, "40-przebieg-nowe-uruchomienie")

    run.step(feature, "przebieg-podglad")
    page.goto(run.url)
    next_step = page.get_by_role("region", name="Następny krok")
    expect(next_step).to_contain_text("Utwórz rozliczenia pracowników")
    next_step.get_by_role("button", name="Podgląd bez zapisu").click()
    live = page.get_by_role("region", name="Podgląd na żywo")
    expect(live).to_contain_text("Podgląd gotowy · nic nie zapisano")
    result = _finished_result(run, "przebieg-podglad")
    summary = result["summary"]
    assert isinstance(summary, dict)
    run.check(
        result["result"] == "OK"
        and [file["file"] for file in summary["files"]] == files
        and [file["rows"] for file in summary["files"]] == run.fixture["settlement"]["rows_per_worker"],
        "the DRY-RUN plans one Rozliczenie pracownika per Wykonawca with its number of rows",
    )
    plan = live.get_by_role("table", name="Plan plików")
    expect(plan.get_by_role("row")).to_have_count(len(files))
    expect(plan.get_by_role("row").filter(has_text=files[0])).to_contain_text("1 wiersz")
    run.check(not worker_settlements.exists(), "the DRY-RUN created no Folder rozliczeń pracowników")
    records = _settlement_records(run, "przebieg-metryki-podglad")
    run.check(
        len(records) == records_before + 1 and records[0]["mode"] == "DRY-RUN",
        "the DRY-RUN appended one record to the CLI's Przebieg rozliczeń metrics file",
    )
    run.capture(page, "41-przebieg-podglad")

    run.step(feature, "przebieg-potwierdzenie")
    live.get_by_role("button", name="Uruchom", exact=True).click()
    confirm = page.get_by_role("alertdialog", name="Uruchomić z zapisem?")
    expect(confirm).to_contain_text("Zapisz i zamknij Plik zbiorczy w Excelu")
    expect(confirm).to_contain_text("„Uruchom” potwierdza Gotowość Pliku zbiorczego")
    expect(confirm.get_by_text("Nie było podglądu bez zapisu")).to_have_count(0)
    run.check(True, "Uruchom asks to confirm the write; the click is Dad's confirmation of Gotowość")
    run.capture(page, "42-przebieg-potwierdzenie")
    pacing = run.root / "pacing" / "szablony-ms.txt"
    pacing.parent.mkdir(exist_ok=True)
    pacing.write_text(str(TEMPLATE_PACING_MS), encoding="utf-8")
    confirm.get_by_role("button", name="Uruchom", exact=True).click()

    run.step(feature, "przebieg-na-zywo")
    try:
        expect(live.get_by_role("progressbar", name="Szablony pracowników")).to_be_visible()
        expect(live).to_contain_text(f"Szablony pracowników · {workers[0]}")
        view = _current_run(run, "przebieg-na-zywo")
        progress = view["progress"]
        assert isinstance(progress, dict)
        run.check(
            view["state"] == "running" and progress["template_total"] == len(workers),
            "the live view shows the Szablony pracowników progress from the projection (harness pacing on)",
        )
        close = urllib.request.Request(
            run.url.rstrip("/") + "/api/shutdown", data=b"", method="POST", headers={"Origin": run.url.rstrip("/")}
        )
        run.check(
            run.refused(close, "przebieg-zamkniecie-odmowa") == 409,
            "closing Centrum is refused while the Przebieg rozliczeń runs",
        )
        run.capture(page, "43-przebieg-na-zywo", live=True)
    finally:
        pacing.unlink(missing_ok=True)

    run.step(feature, "przebieg-zapisane")
    outcome = page.get_by_label("Wynik")
    expect(outcome).to_contain_text("Rozliczenia pracowników gotowe", timeout=20_000)
    result = _finished_result(run, "przebieg-zapisane")
    summary = result["summary"]
    assert isinstance(summary, dict)
    run.check(result["result"] == "OK" and summary["written_count"] == len(workers), "the RUN finished OK")
    run.check(
        sorted(path.name for path in worker_settlements.iterdir()) == sorted(files),
        "the Folder rozliczeń pracowników holds one Rozliczenie pracownika per Wykonawca",
    )
    run.check(
        sorted(path.name for path in period_folder.iterdir())
        == sorted([internal_settlement.name, collective_file.name, WORKER_SETTLEMENTS]),
        "no staging folder is left in the Folder okresu: the folder was published at once",
    )
    records = _settlement_records(run, "przebieg-metryki-zapisane")
    run.check(
        len(records) == records_before + 2
        and [(record["mode"], record["result"], record["origin"]) for record in records[:2]]
        == [("RUN", "OK", "centrum"), ("DRY-RUN", "OK", "centrum")],
        "both Uruchomienia appended one record each to the CLI's metrics file, marked as started from Centrum",
    )
    log = (run.root / "logs" / "centrum-rozliczen.log").read_text(encoding="utf-8")
    run.check(
        f"uruchomienie-koniec proces=settlement tryb=RUN okres={LAST_COMPLETE_FOR_CLI} wynik=OK koniec=finished" in log
        and ROW_MARKER not in log,
        "the Centrum log records the Przebieg rozliczeń with safe facts only",
    )
    outcome.get_by_role("button", name="Folder okresu").click()
    run.wait_for(_last_opened_is(run, period_folder), "Folder okresu opens in Explorer")
    run.capture(page, "44-przebieg-zapisane")

    run.step(feature, "przebieg-status-okresu")
    page.goto(run.url)
    expect(_row(page, LAST_COMPLETE_FOR_CLI)).to_contain_text("Rozliczenia pracowników gotowe")
    expect(page.get_by_role("region", name="Następny krok")).to_contain_text("Okres zamknięty")
    board = run.api("/api/periods", "przebieg-status-okresu")["periods"]
    assert isinstance(board, list)
    [row] = [row for row in board if row["token"] == LAST_COMPLETE_FOR_CLI]
    run.check(
        row["status"] == "worker_settlements_ready" and row["worker_settlement_count"] == len(workers),
        "the Status okresu is Rozliczenia pracowników gotowe with three Rozliczenia pracowników",
    )
    run.capture(page, "45-przebieg-status-okresu")

    run.step(feature, "przebieg-nienadpisywanie")
    before = {path.name: path.read_bytes() for path in worker_settlements.iterdir()}
    started = _post(
        run,
        "/api/runs",
        {"process": "settlement", "period": LAST_COMPLETE_FOR_CLI, "dry_run": False},
        "przebieg-nienadpisywanie-start",
    )
    run.check(started == 200, "a second RUN of the closed period starts (the Okres list no longer offers it)")
    page.goto(run.url + "#/uruchomienia")
    stop = page.get_by_label("Zatrzymanie")
    expect(stop).to_contain_text("Folder rozliczeń pracowników tego okresu już istnieje, więc nie został nadpisany.")
    expect(stop).to_contain_text(f"folder {LAST_COMPLETE_FOR_CLI}\\{WORKER_SETTLEMENTS}")
    result = _finished_result(run, "przebieg-nienadpisywanie")
    stops = result["stops"]
    assert isinstance(stops, list)
    [message] = stops
    assert isinstance(message, dict)
    run.check(
        result["ending"] == "stop"
        and all(message[part] for part in ("what", "where", "action"))
        and str(run.root) not in json.dumps(result, ensure_ascii=False),
        "the stop has three parts and no full path",
    )
    run.check(
        {path.name: path.read_bytes() for path in worker_settlements.iterdir()} == before,
        "Nienadpisywanie: the Rozliczenia pracowników are unchanged",
    )
    run.capture(page, "46-przebieg-nienadpisywanie")

    run.step(feature, "przebieg-wymaga-sprawdzenia")
    page.get_by_role("button", name="Nowe Uruchomienie procesu").click()
    page.get_by_role("radio", name="Przebieg rozliczeń").click()
    expect(page.get_by_role("combobox", name="Okres")).to_contain_text(needs_check)
    page.get_by_role("button", name="Podgląd bez zapisu", exact=True).click()
    outcome = page.get_by_label("Wynik")
    expect(outcome).to_contain_text("Podgląd wymaga sprawdzenia · nic nie zapisano")
    expect(outcome.get_by_role("button", name="Uruchom ponownie")).to_be_visible()
    expect(page.get_by_label("Ostrzeżenie")).to_contain_text("Nieznany identyfikator WYKONAWCA")
    result = _finished_result(run, "przebieg-wymaga-sprawdzenia")
    run.check(
        (result["result"], result["next_action"]) == ("WYMAGA_SPRAWDZENIA", "rerun"),
        "a Wykonawca missing from the konfiguracja wykonawców gives Wymaga sprawdzenia, not a success",
    )
    run.check(
        not (Path(run.fixture["periods_directory"]) / needs_check / f"Rozliczenie pracowników {needs_check}").exists(),
        "the preview wrote nothing",
    )
    run.capture(page, "47-przebieg-wymaga-sprawdzenia")
def _copy_tool(run: Run) -> Path:
    """A copy of the tool inside the run root, so the real launcher logs there and never in Dad's ``logs``."""

    tool = run.root / "narzedzie"
    if not tool.exists():
        tool.mkdir()
        for name in ("run_centrum.py", "pyproject.toml"):
            shutil.copy2(REPO_ROOT / name, tool)
        shutil.copytree(REPO_ROOT / "src", tool / "src", ignore=shutil.ignore_patterns("__pycache__"))
    return tool


def _browser_recorder(run: Run) -> tuple[Path, Path]:
    """Stands in for the browser (``BROWSER``) and appends each opened address to ``browser.txt``."""

    record = run.root / "browser.txt"
    if os.name == "nt":
        recorder = run.root / "browser.cmd"
        recorder.write_text(f'@echo %~1>>"{record}"\r\n', encoding="utf-8")
    else:
        recorder = run.root / "browser.sh"
        recorder.write_text(f'#!/bin/sh\nprintf "%s\\n" "$1" >> "{record}"\n', encoding="utf-8")
        recorder.chmod(0o755)
    return recorder, record


def _launch(run: Run, port: int, name: str) -> subprocess.CompletedProcess[str]:
    """The real launcher (``run_centrum.py``, the console path) from the run's copy of the tool."""

    tool = _copy_tool(run)
    recorder, _ = _browser_recorder(run)
    environment = os.environ | {"PYTHONUTF8": "1", "BROWSER": str(recorder)}
    result = subprocess.run(
        [sys.executable, str(tool / "run_centrum.py"), "--port", str(port)],
        cwd=tool, env=environment, capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=60,
    )
    (run.evidence / "api" / f"{name}.txt").write_text(
        f"exit_code={result.returncode}\nstdout:\n{result.stdout}\nstderr:\n{result.stderr}", encoding="utf-8"
    )
    return result


def _log_events(path: Path) -> list[str]:
    return [line.split(" ")[1] for line in path.read_text(encoding="utf-8").splitlines()] if path.exists() else []


def _answers(url: str) -> bool:
    try:
        with urllib.request.urlopen(url.rstrip("/") + "/api/health", timeout=2):
            return True
    except OSError:
        return False


def drive_start_and_close(run: Run, page: Page) -> None:
    feature = "start-i-zamykanie"
    port = int(urllib.parse.urlsplit(run.url).port or 0)
    served_log = run.root / "logs" / "centrum-rozliczen.log"
    tool_log = run.root / "narzedzie" / "logs" / "centrum-rozliczen.log"

    run.step(feature, "start-gotowosc")
    health = run.api("/api/health", "start-health")
    run.check(health["service"] == "centrum-rozliczen", "health names Centrum rozliczeń and its version")
    run.check(_log_events(served_log)[:1] == ["start"], "the Centrum log starts with a start entry")
    run.check(f"port={port}" in served_log.read_text(encoding="utf-8"), "the start entry names the port")

    run.step(feature, "start-druga-instancja")
    again = _launch(run, port, "start-druga-instancja")
    run.check(again.returncode == 0, "a second launch on the same port exits 0")
    run.check(f"Centrum rozliczeń już działa: {run.url}" in again.stdout, "it says Centrum is already running")
    run.check((run.root / "browser.txt").read_text(encoding="utf-8").split() == [run.url], "it only opens the tab")
    run.check(_log_events(tool_log) == ["juz-dziala"], "its log records the recognised running Centrum")
    run.check(_answers(run.url), "the first Centrum keeps serving")

    run.step(feature, "start-obcy-program")
    with socket.socket() as foreign_program:
        foreign_program.bind(("127.0.0.1", 0))
        foreign_program.listen()
        foreign_port = foreign_program.getsockname()[1]
        refused = _launch(run, foreign_port, "start-obcy-program")
    run.check(refused.returncode == 1, "a launch on a port held by another program exits 1")
    run.check(
        [line.split(":", 1)[0] for line in refused.stdout.splitlines()] == ["Co się stało", "Gdzie", "Co zrobić"],
        "the console shows the three-part start message",
    )
    run.check("zajmuje inny program" in refused.stdout and refused.stderr == "", "no exception text, only the message")
    run.check(_log_events(tool_log)[-1:] == ["blad-startu"], "the log records the start failure")

    run.step(feature, "zamykanie-karta")
    page.goto(run.url)
    other_tab = page.context.new_page()
    other_tab.goto(run.url)
    expect(other_tab.get_by_role("heading", name="Okresy", exact=True)).to_be_visible()
    other_tab.close()
    run.check(_answers(run.url), "closing a tab leaves Centrum running")

    run.step(feature, "zamykanie-potwierdzenie")
    close_button = page.get_by_role("button", name="Zamknij Centrum rozliczeń", exact=True)
    close_button.click()
    dialog = page.get_by_role("alertdialog", name="Zamknąć Centrum rozliczeń?")
    expect(dialog).to_be_visible()
    expect(dialog).to_contain_text("Wrócisz skrótem „Centrum rozliczeń” na pulpicie.")
    run.capture(page, "90-start-i-zamykanie-potwierdzenie", full_page=False)
    dialog.get_by_role("button", name="Anuluj", exact=True).click()
    expect(dialog).to_be_hidden()
    run.check(_answers(run.url) and "zamkniecie" not in _log_events(served_log), "Anuluj leaves Centrum running")

    run.step(feature, "zamykanie")
    close_button.click()
    page.get_by_role("button", name="Zamknij Centrum", exact=True).click()
    expect(page.get_by_role("heading", name="Centrum rozliczeń jest zamknięte", exact=True)).to_be_visible()
    run.wait_for(lambda: not _answers(run.url), "the Centrum process stops serving")
    run.wait_for(lambda: _log_events(served_log)[-2:] == ["zamkniecie", "stop"], "the log ends with zamkniecie, stop")
    run.capture(page, "91-start-i-zamykanie-zamkniete")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--root", required=True, help="Run root with fixture.json and a served Centrum.")
    parser.add_argument("--url", required=True, help="Centrum URL from <root>/server.json.")
    parser.add_argument("--evidence", required=True, help="Evidence directory that survives cleanup.")
    parser.add_argument("--feature", choices=("all", *FEATURES), default="all")
    args = parser.parse_args(argv)
    run = Run(Path(args.root).resolve(), args.url, Path(args.evidence).resolve())
    drives = {
        "ustawienia": drive_settings,
        "tablica-okresow": drive_period_board,
        "konfiguracja": drive_configuration,
        "metryki": drive_metrics,
        "przygotowanie-pliku-zbiorczego": drive_preparation,
        "przebieg-rozliczen": drive_settlement,
        "start-i-zamykanie": drive_start_and_close,
    }
    selected = FEATURES if args.feature == "all" else (args.feature,)
    try:
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch()
            try:
                context = browser.new_context(
                    viewport={"width": 1440, "height": 1000}, locale="pl-PL", timezone_id=DAD_TIME_ZONE.key
                )
                page = context.new_page()
                page.set_default_timeout(10_000)
                for feature in selected:
                    drives[feature](run, page)
            except Exception:
                if "page" in locals():
                    page.screenshot(path=run.evidence / "screenshots" / "zz-failure.png", full_page=True)
                raise
            finally:
                browser.close()
    except (DriveFailure, AssertionError, PlaywrightError) as exc:
        print(f"DRIVE FAIL | {exc}", file=sys.stderr)
        return 1
    finally:
        for name in ("opened.jsonl", "dialogs.jsonl", "config/centrum-ustawienia.yaml", "logs/centrum-rozliczen.log"):
            source = run.root / name
            if source.exists():
                shutil.copy2(source, run.evidence / "state" / source.name)
    print(f"DRIVE OK | features={','.join(selected)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
