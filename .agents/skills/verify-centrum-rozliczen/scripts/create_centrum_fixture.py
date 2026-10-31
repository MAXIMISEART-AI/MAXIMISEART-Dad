"""Create a synthetic folder okresów, Rejestr zleceń, configuration directory, and metrics files for one Centrum run.

The configuration directory holds what Konfiguracja reads: a synthetic konfiguracja
wykonawców and the three Wzorce, valid for the engine's Sprawdzanie. Centrum's
settings file is written there too, as in Dad's ``config/``. The metrics files
hold a synthetic history in the metrics schema for Metryki.

The Rejestr zleceń holds orders of the last complete period ``06_12_10_2026``
whose last order is on Saturday, so a RUN of Przygotowanie pliku zbiorczego
asks Dad's decision about the incomplete period. Its order cells carry
``ROW_MARKER``. The Plik zbiorczy of ``29_05_10_2026`` has a Wiersz danych of a
Wykonawca missing from the konfiguracja wykonawców, so a Podgląd bez zapisu of
Przebieg rozliczeń for that period needs checking.

The tool itself is a synthetic Git clone in ``<root>/aktualizacja/klon`` whose
"origin" is a local bare repository, ``<root>/aktualizacja/origin.git``, one commit
ahead of the clone: Aktualizacja narzędzia in this run moves that clone, never
the checkout the verification runs from, and never reaches GitHub. The clone
stands on the first version, the Wydanie ``v0.0.1`` (an annotated tag it got
from the origin, as an install from the README leaves it). The clone's
``config/`` holds Dad's files, among them his locally changed konfiguracja
wykonawców, which the newest commit stops tracking. Its
``run_centrum.py``, which the restart guardian starts after an update, is a shim
that serves this run's scripted Centrum (``serve_centrum.py`` with this run root)
on the port it is given; while ``<root>/pacing/ponowny-start-ms.txt`` exists, it
waits that many milliseconds first. One version on the origin's branch
``nie-startuje`` has a ``run_centrum.py`` that fails at once with an import error.

``<root>/INTERMAX/`` holds two synthetic INTERMAX files for ``uporzadkowanie-intermax``: one with the Surowe
zlecenia of Monday 26.10.2026 pasted under the ordered days (6 orders of 3 Wykonawcy, the first row shifted), and
one whose pasted rows hold a formula, which Uporządkowanie INTERMAX stops „do Maksa”. Kopie INTERMAX go to
``<root>/Kopie INTERMAX``. ``<root>/INTERMAX/przywracanie/`` holds a third INTERMAX for ``kopie-intermax`` and the
bytes of three Kopie of both names (one ordered order fewer in the oldest), which that feature puts into the folder
of Kopie only while it runs.

``--wariant rok-pracy`` (for ``wydajnosc``) adds a year of Dad's work on top:
``<root>/Rok pracy/Załączniki`` with 52 periods up to the last complete one in
Dad's layout, each before it with a Folder rozliczeń pracowników of 14
Wykonawcy, and 258 more synthetic records in both metrics files.
"""

from __future__ import annotations

from datetime import date, datetime, timedelta
import json
from pathlib import Path
import random
import os
import shutil
import subprocess
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet
import yaml

sys.path.insert(0, str(Path(__file__).resolve().parents[4] / "tools/lib"))
sys.path.insert(0, str(Path(__file__).resolve().parents[4]))

from skrypt_skilla import Parser, create_run_root
from tests.fixtures.intermax import ORDER_SHEET, Intermax

# Dad's "today" for the run: Monday 19.10.2026, so the last complete period is 06_12_10_2026,
# whose Folder okresu deliberately does not exist.
TODAY = "2026-10-19"
LAST_COMPLETE = "06_12_10_2026"
# A period whose Plik zbiorczy has a Wiersz danych of a Wykonawca the konfiguracja wykonawców does not know.
NEEDS_CHECK = "29_05_10_2026"
UNKNOWN_WORKER = "wykonawca.spoza.konfiguracji"
WORKERS = ("Wykonawca Testowy A", "Wykonawca Testowy B", "Wykonawca Testowy C")
# A Wykonawca of the konfiguracja wykonawców without orders: his Szablon pracownika stays empty.
EMPTY_WORKER = "Wykonawca Testowy D"
# A synthetic row value inside every workbook. Centrum must never read or show workbook content,
# so this marker must never appear in evidence (centrum_self_test.py checks it).
ROW_MARKER = "CENTRUM-SYN-ROW"
PHASES = ("Sprawdzanie", "Odczyt danych", "Planowanie", "Zapisywanie")
PREPARATION_COUNTERS = {"period_orders": 9, "successful_orders": 8, "skipped_empty_rows": 1, "warnings": 0}
SETTLEMENT_COUNTERS = {
    "templates_total": 3,
    "templates_completed": 3,
    "rows": 8,
    "written": 3,
    "empty": 0,
    "planned": 3,
    "skipped": 0,
    "issues": 0,
}
# Synthetic history as the core writes it (metrics schema_version 2):
# (started_at UTC, mode, period, result, phase times in ms; the total is their sum).
MetricRun = tuple[str, str, str, str, tuple[int, int, int, int]]
# The stops' Rodzaj Zatrzymania and Komunikat code by started_at. A stop missing here, the BLAD_KRYTYCZNY of
# 2026-08-25, is an old schema_version 1 record written before records carried them.
RECORDED_STOPS: dict[str, dict[str, str]] = {
    "2026-09-17T15:40:00+00:00": {"kind": "for_maks", "code": "BLAD_NIEOCZEKIWANY"},
    "2026-10-01T15:40:00+00:00": {"kind": "for_tata", "code": "PLIK_ZBIORCZY_OTWARTY"},
}
PREPARATION_RUNS: tuple[MetricRun, ...] = (
    ("2026-08-25T06:15:00+00:00", "RUN", "18_24_08_2026", "BLAD_KRYTYCZNY", (400, 1300, 0, 0)),
    ("2026-09-01T06:21:00+00:00", "RUN", "25_31_08_2026", "OK", (300, 7100, 900, 2500)),
    ("2026-09-08T06:21:00+00:00", "RUN", "01_07_09_2026", "OK", (300, 7400, 900, 2600)),
    ("2026-09-15T06:21:00+00:00", "RUN", "08_14_09_2026", "OK", (300, 6800, 800, 2500)),
    ("2026-09-22T06:12:00+00:00", "DRY-RUN", "15_21_09_2026", "OK", (300, 6700, 800, 0)),
    ("2026-09-22T06:21:00+00:00", "RUN", "15_21_09_2026", "OK", (400, 7900, 900, 2700)),
    ("2026-09-29T06:12:00+00:00", "DRY-RUN", "22_28_09_2026", "OK", (300, 6900, 900, 0)),
    ("2026-09-29T06:21:00+00:00", "RUN", "22_28_09_2026", "OK", (300, 7100, 900, 2600)),
    ("2026-10-06T06:12:00+00:00", "DRY-RUN", "29_05_10_2026", "OK", (300, 7100, 900, 0)),
    ("2026-10-06T06:21:00+00:00", "RUN", "29_05_10_2026", "OK", (400, 7600, 900, 2700)),
)
SETTLEMENT_RUNS: tuple[MetricRun, ...] = (
    ("2026-09-03T15:46:00+00:00", "RUN", "25_31_08_2026", "OK", (500, 3100, 700, 8000)),
    ("2026-09-10T15:46:00+00:00", "RUN", "01_07_09_2026", "OK", (500, 3000, 700, 7600)),
    ("2026-09-17T15:40:00+00:00", "RUN", "08_14_09_2026", "BLAD_KRYTYCZNY", (500, 3100, 700, 1200)),
    ("2026-09-17T15:46:00+00:00", "RUN", "08_14_09_2026", "WYMAGA_SPRAWDZENIA", (500, 3200, 800, 7600)),
    ("2026-09-24T15:40:00+00:00", "DRY-RUN", "15_21_09_2026", "OK", (500, 2100, 700, 0)),
    ("2026-09-24T15:46:00+00:00", "RUN", "15_21_09_2026", "OK", (400, 2900, 700, 7100)),
    ("2026-10-01T15:40:00+00:00", "RUN", "22_28_09_2026", "BLAD_KRYTYCZNY", (500, 0, 0, 0)),
    ("2026-10-01T15:46:00+00:00", "RUN", "22_28_09_2026", "OK", (500, 3400, 800, 8000)),
)
WORKER_IDS = ("wykonawca.a", "wykonawca.b", "wykonawca.c", "wykonawca.d")
INTERNAL_TEMPLATE = "Wzorzec rozliczenia wewnętrznego.xlsx"
COLLECTIVE_TEMPLATE = "Wzorzec pliku zbiorczego.xlsx"
WORKER_TEMPLATE = "Wzorzec rozliczenia pracownika.xlsx"
# The target headers the engine looks for; in the Wzorzec pliku zbiorczego WYKONAWCA must be in H17.
INTERNAL_HEADERS = (
    "DATA", "MIASTO", "ADRES", "UMOWA", "URZĄDZENIA", "CZYNNOŚCI", "NUMER ZL.", "WYKONAWCA",
    "NUMER UMOWY/ZLECENIA ORANGE", "SKUTECZNE", "CZYNNOŚCI ZREALIZOWANE",
)
COLLECTIVE_HEADERS = (
    "MIASTO", "ADRES", "UMOWA", "URZĄDZENIA", "CZYNNOŚCI", "NUMER ZL.", "DATA", "WYKONAWCA",
    "NUMER UMOWY/ZLECENIA ORANGE", "SKUTECZNE", "CZYNNOŚCI ZREALIZOWANE",
)
# Synthetic price-list names in N17:AP17 of the Wzorzec pliku zbiorczego, which Przygotowanie copies into the Plik
# zbiorczy; the kind page of Uzupełnienie shows them over the quantity columns.
PRICE_LIST = (
    "Dodatek za usterkę po godzinach", "Instalacja internetu", "Instalacja telewizji", "Multiroom",
    "Aktywacja telefonu", "Dostarczenie urządzenia", "Drukowanie umowy", "Prace niestandardowe (rbh)",
    "Dojazd do klienta", "Przyłącze w mikrokanalizacji", "Przyłącze napowietrzne",
    "Przyłącze w budynku wielorodzinnym", "Montaż ONT", "Przyłącze doziemne", "Spawanie w mufie",
    "Przedłużenie przyłącza", "Spawanie gniazdka", "Trasa do splittera", "Instalacja wewnętrzna",
    "Dodatkowe urządzenie", "Dostarczenie dodatkowego urządzenia", "Usterka", "Sieć domowa do 30 m",
    "Sieć domowa do 10 m bez przebicia", "Sieć domowa do 15 m z przebiciem", "Modyfikacja sieci domowej",
    "Wywiad techniczny", "Prace dodatkowe z premii", "Deinstalacja usługi z odbiorem urządzeń",
)
# CZYNNOŚCI ZREALIZOWANE: a technician's description long enough to be cut to one line, with the row marker.
COMPLETED = "Wymiana gniazdka abonenckiego i przedłużenie kabla do salonu, klient prosi o kontakt przed wizytą"


YEAR_OF_WORK = "rok-pracy"
YEAR_PERIODS = 52
YEAR_WORKERS = tuple(f"Wykonawca Testowy {chr(ord('A') + number)}" for number in range(14))
MONTHS = (
    "Styczeń", "Luty", "Marzec", "Kwiecień", "Maj", "Czerwiec",
    "Lipiec", "Sierpień", "Wrzesień", "Październik", "Listopad", "Grudzień",
)


REGISTER_HEADERS = (
    "nr zlecenia", "DATA GODZ. POCZ.", "WYKONAWCA", "MIASTO", "adres", "STATUS APLIKACJI", "UMOWA",
    "URZĄDZENIA", "CZYNNOŚCI", "SKUTECZNE", "CZYNNOŚCI ZREALIZOWANE", "NUMER UMOWY/ZLECENIA ORANGE",
)
# (start, worker, SKUTECZNE): one order of the previous period, then 06_12_10_2026 ending on Saturday.
ORDERS = (
    (datetime(2026, 10, 5, 9, 0), "wykonawca.a", "tak"),
    (datetime(2026, 10, 6, 8, 0), "wykonawca.a", "tak"),
    (datetime(2026, 10, 7, 10, 0), "wykonawca.b", "nie"),
    (datetime(2026, 10, 8, 12, 0), "wykonawca.b", "tak"),
    (datetime(2026, 10, 10, 9, 30), "wykonawca.c", "tak"),
)
# URZĄDZENIA by order number; every other order's URZĄDZENIA is its row marker, a device the Reguły ilości do not
# know. Plik zbiorczy rows 18 and 20 (orders 2 and 5) warn about the unknown device, row 19 (order 4) about prewencja.
DEVICES = {4: "prewencja"}
# What Przygotowanie pliku zbiorczego of LAST_COMPLETE reports for these orders (orders per day from Tuesday).
EXPECTED_PREPARATION = {
    "period_order_count": 4,
    "successful_order_count": 3,
    "skipped_empty_rows": 0,
    "orders_per_day": [1, 1, 1, 0, 1, 0, 0],
}


def _register(path: Path) -> None:
    """A Rejestr zleceń with one data sheet; every order cell with content carries the row marker."""

    workbook = Workbook()
    sheet = workbook.active
    assert isinstance(sheet, Worksheet)
    sheet.title = "ZLECENIA"
    sheet.append(REGISTER_HEADERS)
    for number, (started, worker, successful) in enumerate(ORDERS, start=1):
        content = f"{ROW_MARKER}-{number:03d}"
        values: dict[str, object] = dict.fromkeys(REGISTER_HEADERS, content)
        values.update({"DATA GODZ. POCZ.": started, "WYKONAWCA": worker, "SKUTECZNE": successful})
        values["URZĄDZENIA"] = DEVICES.get(number, content)
        values["CZYNNOŚCI ZREALIZOWANE"] = f"{COMPLETED} {content}"
        sheet.append([values[header] for header in REGISTER_HEADERS])
    workbook.save(path)


def _workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["A18"] = ROW_MARKER
    workbook.save(path)


def _template(path: Path, header_row: int, headers: tuple[str, ...], price_list: tuple[str, ...] = ()) -> None:
    """A Wzorzec without orders; the row marker sits on a second sheet the Sprawdzanie never reads."""

    workbook = Workbook()
    sheet = workbook.active
    assert isinstance(sheet, Worksheet)
    for column, header in enumerate(headers, start=1):
        sheet.cell(header_row, column).value = header
    for column, name in enumerate(price_list, start=14):  # N:AP
        sheet.cell(header_row, column).value = name
    workbook.create_sheet("Stawki")["A1"] = ROW_MARKER
    workbook.save(path)


def _configuration(config: Path) -> dict[str, str]:
    worker_config = config / "worker_mapping.yaml"
    worker_config.write_text(
        yaml.safe_dump(
            {"schema_version": 1, "workers": dict(zip(WORKER_IDS, (*WORKERS, EMPTY_WORKER), strict=True))},
            allow_unicode=True,
            sort_keys=False,
        ),
        encoding="utf-8",
    )
    _template(config / INTERNAL_TEMPLATE, 1, INTERNAL_HEADERS)
    _template(config / COLLECTIVE_TEMPLATE, 17, COLLECTIVE_HEADERS, PRICE_LIST)
    _template(config / WORKER_TEMPLATE, 17, COLLECTIVE_HEADERS)
    return {
        "worker_config": str(worker_config),
        "internal_template": str(config / INTERNAL_TEMPLATE),
        "collective_template": str(config / COLLECTIVE_TEMPLATE),
        "worker_template": str(config / WORKER_TEMPLATE),
    }


def _period_files(periods: Path, token: str) -> Path:
    period_folder = periods / token
    period_folder.mkdir()
    _workbook(period_folder / f"Rozliczenie wew. {token}.xlsx")
    _workbook(period_folder / f"Rozliczenie {token} - zbiorcze.xlsx")
    return period_folder


def _collective_file(path: Path, workers: tuple[str, ...]) -> None:
    """A Plik zbiorczy in the engine's layout: headers in row 17, one Wiersz danych per worker id from row 18."""

    workbook = Workbook()
    sheet = workbook.active
    assert isinstance(sheet, Worksheet)
    for column, header in enumerate(COLLECTIVE_HEADERS, start=1):
        sheet.cell(17, column).value = header
    worker_column = COLLECTIVE_HEADERS.index("WYKONAWCA") + 1
    for row, worker in enumerate(workers, start=18):
        for column in range(1, len(COLLECTIVE_HEADERS) + 1):
            sheet.cell(row, column).value = f"{ROW_MARKER}-{row:03d}"
        sheet.cell(row, worker_column).value = worker
    workbook.save(path)


def _metric_record(run: MetricRun, counters: dict[str, int]) -> dict[str, object]:
    started_at, mode, period, result, phases = run
    completed = result != "BLAD_KRYTYCZNY"
    record: dict[str, object] = {
        "schema_version": 2,
        "started_at": started_at,
        "mode": mode,
        "period": period,
        "completed": completed,
        "result": result,
        "total_duration_ms": sum(phases),
        "phase_durations_ms": dict(zip(PHASES, phases, strict=True)),
        "counters": counters if completed else dict.fromkeys(counters, 0),
    }
    if completed:
        return record | {"stop": None}
    if started_at not in RECORDED_STOPS:
        return record | {"schema_version": 1}
    return record | {"stop": RECORDED_STOPS[started_at]}


def _intermax(root: Path) -> dict[str, object]:
    """INTERMAX with Surowe zlecenia to order and one whose pasted rows Uporządkowanie INTERMAX stops „do Maksa”."""

    folder = root / "INTERMAX"
    folder.mkdir()
    register = folder / "ZLECENIA INTERMAX OD 08_11_2021-NOWA WERSJA.xlsx"
    Intermax().write(register)
    with_formula = folder / "ZLECENIA INTERMAX Z FORMUŁĄ.xlsx"
    layout = Intermax().write(with_formula)
    workbook = load_workbook(with_formula)
    workbook[ORDER_SHEET].cell(layout.first_raw + 1, 4).value = "=1+1"
    workbook.save(with_formula)
    restore = folder / "przywracanie"
    (restore / "kopie").mkdir(parents=True)
    restored = restore / register.name
    Intermax().write(restored)
    earlier = Intermax()
    earlier.ordered = earlier.ordered[:-1]
    # Kopie of both names; by name „przed przywróceniem” sorts before „przed uporządkowaniem”, in time it is between.
    copies = {
        "INTERMAX przed uporządkowaniem 2026-10-16 07-30-00.xlsx": earlier,
        "INTERMAX przed przywróceniem 2026-10-17 12-15-00.xlsx": earlier,
        "INTERMAX przed uporządkowaniem 2026-10-18 18-45-00.xlsx": Intermax(),
    }
    for name, intermax in copies.items():
        intermax.write(restore / "kopie" / name)
    return {
        "register": str(register),
        "with_formula": str(with_formula),
        "copies": str(root / "Kopie INTERMAX"),
        "orders": len(layout.raw),
        "workers": len({order.worker for order in layout.raw}),
        "day": "pon 26.10.2026",
        "restore": {
            "register": str(restored),
            "copies": [str(restore / "kopie" / name) for name in copies],
        },
    }


def _metrics(root: Path) -> dict[str, object]:
    """Both metrics files with a valid history and one record each that the metrics reader skips.

    Uporządkowanie INTERMAX has its own metrics file, absent until its first Uruchomienie.
    """

    metrics = root / "metrics"
    metrics.mkdir()
    preparation = metrics / "przygotowanie-pliku-zbiorczego-metrics.jsonl"
    settlement = metrics / "rozliczenia-metrics.jsonl"
    preparation_lines = [json.dumps(_metric_record(run, PREPARATION_COUNTERS)) for run in PREPARATION_RUNS]
    # Przebieg rozliczeń counters in the Przygotowanie file: invalid there, so skipped.
    preparation_lines.append(json.dumps(_metric_record(SETTLEMENT_RUNS[0], SETTLEMENT_COUNTERS)))
    settlement_lines = [json.dumps(_metric_record(run, SETTLEMENT_COUNTERS)) for run in SETTLEMENT_RUNS]
    settlement_lines.append('{"schema_version":1,"started_at":"2026-10-0')  # a torn last line
    preparation.write_text("".join(line + "\n" for line in preparation_lines), encoding="utf-8")
    settlement.write_text("".join(line + "\n" for line in settlement_lines), encoding="utf-8")
    interrupted = PREPARATION_RUNS[0]
    return {
        "preparation": str(preparation),
        "settlement": str(settlement),
        "intermax": str(metrics / "uporzadkowanie-intermax-metrics.jsonl"),
        "runs": len(PREPARATION_RUNS) + len(SETTLEMENT_RUNS),
        "oldest": min(run[0] for run in (*PREPARATION_RUNS, *SETTLEMENT_RUNS)),
        # period, then the measured phase times until the interruption
        "interrupted": [interrupted[2], [phase for phase in interrupted[4] if phase]],
    }


def _dad_layout(root: Path) -> dict[str, object]:
    """Dad's folder layout: ``Załączniki\\Rozliczenia {rok}\\{miesiąc}\\INEA\\{okres}``, months as billing cycles.

    ``06_12_10_2026`` (the last complete period) has both period files in the October folder, so the
    board shows it as prepared; ``29_05_10_2026`` crosses the month and sits in October by its Monday.
    """

    attachments = root / "Załączniki"
    september = attachments / "Rozliczenia 2026" / "09_Wrzesień_2026" / "INEA"
    october = attachments / "Rozliczenia 2026" / "10_Październik_2026" / "INEA"
    november = attachments / "Rozliczenia 2026" / "11_Listopad_2026" / "INEA"
    for folder in (september, october, november):
        folder.mkdir(parents=True)
    (attachments / "Rozliczenia 2026" / "Nowe rozliczenia wzór od 01_2026" / "Podpisane").mkdir(parents=True)
    (september / "22_28_09_2026").mkdir()
    (september / "Podpisane").mkdir()
    _period_files(october, "29_05_10_2026")
    prepared = _period_files(october, LAST_COMPLETE)
    (october / "13_19_10_2026").mkdir()
    (november / "27_02_11_2026").mkdir()
    return {
        "periods_directory": str(attachments),
        "opened": str(prepared / f"Rozliczenie {LAST_COMPLETE} - zbiorcze.xlsx"),
        # token, Status okresu, last complete period
        "board": [
            ["27_02_11_2026", "no_period_files", False],
            ["13_19_10_2026", "no_period_files", False],
            [LAST_COMPLETE, "collective_file_prepared", True],
            ["29_05_10_2026", "collective_file_prepared", False],
            ["22_28_09_2026", "no_period_files", False],
        ],
    }


def _year_metric_runs(rnd: random.Random, token: str, monday: date) -> tuple[list[MetricRun], list[MetricRun]]:
    """One week's history as Dad works: Przygotowanie on Tuesday (preview, RUN, a repeated RUN), Przebieg on Friday."""

    def started(day: date, hour: int) -> str:
        return f"{day.isoformat()}T{hour:02d}:{rnd.randint(10, 59)}:00+00:00"

    preparation: list[MetricRun] = []
    for mode in ("DRY-RUN", "RUN", "RUN"):
        writing = 0 if mode == "DRY-RUN" else rnd.randint(2000, 4000)
        phases = (rnd.randint(200, 600), rnd.randint(5000, 9000), rnd.randint(600, 1200), writing)
        preparation.append((started(monday + timedelta(days=1), 6), mode, token, "OK", phases))
    settlement: list[MetricRun] = []
    for mode in ("DRY-RUN", "RUN"):
        writing = 0 if mode == "DRY-RUN" else rnd.randint(6000, 12000)
        phases = (rnd.randint(300, 700), rnd.randint(2500, 4000), rnd.randint(500, 900), writing)
        result = rnd.choice(("OK", "OK", "OK", "WYMAGA_SPRAWDZENIA"))
        settlement.append((started(monday + timedelta(days=4), 15), mode, token, result, phases))
    return preparation, settlement


def _year_of_work(root: Path, metrics: dict[str, object]) -> dict[str, object]:
    """A year of Dad's work: 52 periods up to LAST_COMPLETE in Dad's layout and their metrics records.

    Every period before LAST_COMPLETE has its Rozliczenia pracowników of ``YEAR_WORKERS``; LAST_COMPLETE
    has only its two period files, as on Monday before Przebieg rozliczeń. The records go before the
    existing ones, so each file still ends with the record the metrics reader skips.
    """

    attachments = root / "Rok pracy" / "Załączniki"
    sample = root / "Rok pracy" / "wzor.xlsx"
    sample.parent.mkdir()
    _workbook(sample)
    rnd = random.Random(165)
    last_tuesday = date(2026, 10, 6)  # LAST_COMPLETE
    preparation: list[MetricRun] = []
    settlement: list[MetricRun] = []
    for weeks_back in range(YEAR_PERIODS - 1, -1, -1):
        tuesday = last_tuesday - timedelta(weeks=weeks_back)
        monday = tuesday + timedelta(days=6)
        token = f"{tuesday.day:02d}_{monday.day:02d}_{monday.month:02d}_{monday.year}"
        month = f"{monday.month:02d}_{MONTHS[monday.month - 1]}_{monday.year}"
        folder = attachments / f"Rozliczenia {monday.year}" / month / "INEA" / token
        folder.mkdir(parents=True)
        shutil.copyfile(sample, folder / f"Rozliczenie wew. {token}.xlsx")
        shutil.copyfile(sample, folder / f"Rozliczenie {token} - zbiorcze.xlsx")
        if weeks_back > 0:
            workers = folder / f"Rozliczenie pracowników {token}"
            workers.mkdir()
            for worker in YEAR_WORKERS:
                shutil.copyfile(sample, workers / f"Rozliczenie {token} - {worker}.xlsx")
        week_preparation, week_settlement = _year_metric_runs(rnd, token, monday)
        preparation += week_preparation
        settlement += week_settlement if weeks_back > 0 else []
    sample.unlink()
    year_settlement = {**SETTLEMENT_COUNTERS, "templates_total": len(YEAR_WORKERS), "templates_completed": len(YEAR_WORKERS)}
    for name, runs, counters in (
        ("preparation", preparation, PREPARATION_COUNTERS),
        ("settlement", settlement, year_settlement),
    ):
        path = Path(str(metrics[name]))
        lines = "".join(json.dumps(_metric_record(run, counters)) + "\n" for run in runs)
        path.write_text(lines + path.read_text(encoding="utf-8"), encoding="utf-8")
    return {
        "periods_directory": str(attachments),
        "periods": YEAR_PERIODS,
        "worker_settlements_ready": YEAR_PERIODS - 1,
        "workers": len(YEAR_WORKERS),
        "metrics_runs": len(preparation) + len(settlement),
        "oldest": min(run[0] for run in (*preparation, *settlement)),
    }


def _git(directory: Path, *args: str) -> str:
    """Git without the machine's own configuration, with a synthetic author; never the network."""

    environment = os.environ | {
        "GIT_CONFIG_GLOBAL": os.devnull,
        "GIT_CONFIG_NOSYSTEM": "1",
        "GIT_AUTHOR_NAME": "Maks Syntetyczny",
        "GIT_AUTHOR_EMAIL": "maks@example.invalid",
        "GIT_COMMITTER_NAME": "Maks Syntetyczny",
        "GIT_COMMITTER_EMAIL": "maks@example.invalid",
    }
    result = subprocess.run(
        ["git", "-C", str(directory), *args], env=environment, capture_output=True, text=True, check=True
    )
    return result.stdout.strip()


def _publish(publisher: Path, version: int, message: str, *, untrack: tuple[str, ...] = ()) -> str:
    (publisher / "pyproject.toml").write_text(
        f'[project]\nname = "narzedzie-syntetyczne"\nversion = "0.0.{version}"\n', encoding="utf-8"
    )
    if untrack:
        _git(publisher, "rm", "--cached", "--quiet", *untrack)
    _git(publisher, "add", "--all")
    _git(publisher, "commit", "--quiet", "--message", message)
    return _git(publisher, "rev-parse", "HEAD")


# The synthetic tool's Centrum: the run's scripted Centrum, so a restarted Centrum keeps the run-local boundaries.
_CENTRUM_SHIM = """\
\"\"\"Centrum of the run's synthetic tool: the run's scripted Centrum on the port the restart guardian gives.\"\"\"

from pathlib import Path
import runpy
import sys
import time

SERVE = {serve!r}
ROOT = Path({root!r})

pacing = ROOT / "pacing" / "ponowny-start-ms.txt"
if pacing.exists():
    time.sleep(int(pacing.read_text(encoding="utf-8")) / 1000)
port = sys.argv[sys.argv.index("--port") + 1]
sys.argv = [SERVE, "--root", str(ROOT), "--port", port]
runpy.run_path(SERVE, run_name="__main__")
"""


# A version whose Centrum does not start: an import error, as when a new version needs a library this Python lacks.
_NOT_STARTING_CENTRUM = """\
\"\"\"Centrum of a version that does not start: it needs a library no Python on this computer has.\"\"\"

import biblioteka_ktorej_nie_ma  # noqa: F401
"""


# The tool's konfiguracja wykonawców: tracked in the first version with an example, untracked by the newest (#161).
WORKER_MAPPING = "config/worker_mapping.yaml"
# A library no index has: with pip offline (PIP_NO_INDEX in serve_centrum.py) its installation always fails.
UNINSTALLABLE_REQUIREMENTS = "biblioteka-syntetyczna-nie-do-zainstalowania>=1.0\n"
# Dad's files in the clone's config/ after the first version: his konfiguracja wykonawców with synthetic
# "real" names (a local change) and files Git never tracks (Ustawienia taty, the key to Zgłoszenia, a Wzorzec).
DAD_FILES = {
    WORKER_MAPPING: "schema_version: 1\nworkers:\n  wykonawca.a: Wykonawca Syntetyczny A\n",
    "config/centrum-ustawienia.yaml": "folder_okresow: INEA\n",
    "config/zgloszenia-klucz.txt": "klucz-syntetyczny\n",
    "config/Wzorzec rozliczenia pracownika.xlsx": "wzorzec syntetyczny\n",
}


def _tool_clone(root: Path) -> dict[str, object]:
    """The run's tool: a clone on the first version (Wydanie v0.0.1) with Dad's files, and a local "origin" one version ahead.

    The newest version stops tracking the konfiguracja wykonawców, which Dad changed locally: the first
    update after #161 on Dad's computer. One version later (``libraries``, on the origin's branch
    ``biblioteki`` until the drive makes it the Wydanie) needs a library whose installation fails. Another
    one after the newest (``not_starting``, on the branch ``nie-startuje``) has a Centrum that does not start.
    """

    tool = root / "aktualizacja"
    publisher, origin, clone = tool / "wydawca", tool / "origin.git", tool / "klon"
    (publisher / "config").mkdir(parents=True)
    _git(publisher, "init", "--quiet", "--initial-branch", "main")
    shim = _CENTRUM_SHIM.format(serve=str(Path(__file__).resolve().with_name("serve_centrum.py")), root=str(root))
    (publisher / "run_centrum.py").write_text(shim, encoding="utf-8")
    (publisher / WORKER_MAPPING).write_text("schema_version: 1\nworkers:\n  przyklad: Przykład\n", encoding="utf-8")
    first = _publish(publisher, 1, "Pierwsza wersja narzędzia")
    _git(publisher, "tag", "--annotate", "--message", "Wydanie v0.0.1", "v0.0.1")
    _git(tool, "clone", "--quiet", "--bare", str(publisher), str(origin))
    _git(tool, "clone", "--quiet", str(origin), str(clone))
    for name, content in DAD_FILES.items():
        (clone / name).write_text(content, encoding="utf-8")
    (publisher / ".gitignore").write_text(f"{WORKER_MAPPING}\n", encoding="utf-8")
    newest = _publish(publisher, 2, "Nowa wersja narzędzia", untrack=(WORKER_MAPPING,))
    _git(publisher, "push", "--quiet", str(origin), "main")
    _git(origin, "tag", "v0.1.0", newest)
    (publisher / "requirements.txt").write_text(UNINSTALLABLE_REQUIREMENTS, encoding="utf-8")
    libraries = _publish(publisher, 3, "Wersja z nową biblioteką")
    _git(publisher, "push", "--quiet", str(origin), "main:biblioteki")
    _git(publisher, "checkout", "--quiet", "-b", "nie-startuje", newest)
    (publisher / "run_centrum.py").write_text(_NOT_STARTING_CENTRUM, encoding="utf-8")
    not_starting = _publish(publisher, 4, "Wersja, która nie wystartuje")
    _git(publisher, "push", "--quiet", str(origin), "nie-startuje")
    return {
        "clone": str(clone),
        "origin": str(origin),
        "first": first,
        "newest": newest,
        "libraries": libraries,
        "not_starting": not_starting,
        "dad_files": DAD_FILES,
    }


def create(root: Path, variant: str | None = None) -> dict[str, object]:
    create_run_root(root)
    (root / "config").mkdir()
    configuration = _configuration(root / "config")
    (root / "picker").mkdir()
    periods = root / "INEA"
    periods.mkdir()

    ready = _period_files(periods, "22_28_09_2026")
    workers = ready / "Rozliczenie pracowników 22_28_09_2026"
    workers.mkdir()
    for worker in WORKERS:
        _workbook(workers / f"Rozliczenie 22_28_09_2026 - {worker}.xlsx")
    needs_check = _period_files(periods, NEEDS_CHECK)
    _collective_file(needs_check / f"Rozliczenie {NEEDS_CHECK} - zbiorcze.xlsx", ("wykonawca.a", UNKNOWN_WORKER))
    (periods / "15_21_09_2026").mkdir()
    (periods / "Archiwum").mkdir()

    register = root / "ZLECENIA SYNTETYCZNE OD 01_01_2026.xlsx"
    _register(register)

    fixture: dict[str, object] = {
        "today": TODAY,
        "periods_directory": str(periods),
        "order_register": str(register),
        "settings": str(root / "config" / "centrum-ustawienia.yaml"),
        "metrics": _metrics(root),
        # Konfiguracja: the files Centrum checks, and the Wykonawcy in the order of the konfiguracja wykonawców
        "configuration": configuration,
        "workers": [*WORKERS, EMPTY_WORKER],
        # Aktualizacja narzędzia: the run's tool clone, its local "origin", both commits, and Dad's files in the clone
        "tool": _tool_clone(root),
        # Uporządkowanie INTERMAX: two INTERMAX files, the folder of Kopie INTERMAX and the day it orders
        "intermax": _intermax(root),
        # Przygotowanie pliku zbiorczego of the last complete period from the Rejestr zleceń above
        "preparation": EXPECTED_PREPARATION,
        # Przebieg rozliczeń: one Wiersz danych per Wykonawca with orders after Przygotowanie of the last complete
        # period, none for the last one, and the period whose Podgląd bez zapisu needs checking
        "settlement": {"rows_per_worker": [1, 1, 1, 0], "needs_check": NEEDS_CHECK},
        # Dad's own layout with months as billing cycles, for uklad-folderow-taty
        "dad_layout": _dad_layout(root),
        # token, Status okresu, number of Rozliczenia pracowników, last complete period
        "board": [
            [LAST_COMPLETE, "no_period_files", 0, True],
            ["29_05_10_2026", "collective_file_prepared", 0, False],
            ["22_28_09_2026", "worker_settlements_ready", len(WORKERS), False],
            ["15_21_09_2026", "no_period_files", 0, False],
        ],
    }
    if variant == YEAR_OF_WORK:
        metrics = fixture["metrics"]
        assert isinstance(metrics, dict)
        year = _year_of_work(root, metrics)
        metrics["runs"] = len(PREPARATION_RUNS) + len(SETTLEMENT_RUNS) + int(str(year["metrics_runs"]))
        metrics["oldest"] = year["oldest"]
        # Dad's year for wydajnosc: its folder okresów replaces the flat INEA in the measured Centrum
        fixture["year_of_work"] = year
    (root / "fixture.json").write_text(json.dumps(fixture, ensure_ascii=False, indent=2), encoding="utf-8")
    return fixture


def main(argv: list[str] | None = None) -> int:
    parser = Parser(description=__doc__)
    parser.add_root("run root to create; absent or empty")
    parser.add_argument(
        "--wariant", choices=(YEAR_OF_WORK,), help="rok-pracy: add a year of Dad's work for wydajnosc (see above)."
    )
    args = parser.parse_args(argv)
    try:
        fixture = create(args.root, args.wariant)
    except (OSError, ValueError, subprocess.CalledProcessError) as exc:
        print(f"FIXTURE FAILED: {exc}", file=sys.stderr)
        return 1
    print(f"FIXTURE OK | root={args.root} | today={fixture['today']} | variant={args.wariant or 'podstawowy'}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
