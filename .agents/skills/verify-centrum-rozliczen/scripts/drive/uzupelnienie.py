"""Drive features/uzupelnienie.md."""

from __future__ import annotations

from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string
from openpyxl.worksheet.worksheet import Worksheet
from playwright.sync_api import Page, expect

from drive.common import LAST_COMPLETE, Run, last_opened_is, period_files
from drive.przygotowanie_pliku_zbiorczego import drive_preparation

FEATURE = "uzupelnienie"
SUMMARY = f"#/uzupelnienie/{LAST_COMPLETE}"
API = f"/api/periods/{LAST_COMPLETE}/quantities-to-fill"
VERSION_API = f"/api/periods/{LAST_COMPLETE}/collective-file-version"
# The page asks about the saved version every 3 s (frontend/src/lib/use-live-file.ts); a full read follows.
LIVE_TIMEOUT_MS = 10_000
KIND_TITLES = {"without_rule": "Wiersze bez reguły", "from_description": "Wiersze z opisu"}
# The fixture's Plik zbiorczy after Przygotowanie: rows 18 and 20 carry a device the Reguły ilości do not know,
# row 19 is prewencja (create_centrum_fixture.py, DEVICES).
EXPECTED_ROWS = {"without_rule": [18, 20], "from_description": [19]}
# Dad's value goes into a column from the technician's description: it is outside V, T and N, so it also makes
# its row without rule done, and it moves that column's counter.
DAD_COLUMN = "Y"
# A period whose Folder okresu holds no Plik zbiorczy.
EMPTY_PERIOD = "15_21_09_2026"


def _in_rows(rows: int) -> str:
    return "nic nie wpisano" if rows == 0 else f"wpisane w {rows} {'wierszu' if rows == 1 else 'wierszach'}"


def _ready(run: Run, name: str) -> dict[str, object]:
    answer = run.api(API, name)
    run.check(answer["state"] == "ready", f"{name}: the server read the Plik zbiorczy")
    return answer


def _without_read(answer: dict[str, object]) -> dict[str, object]:
    """The answer without its read stamp: the time of every read differs."""

    return {key: value for key, value in answer.items() if key != "read"}


def _rows(answer: dict[str, object]) -> dict[str, list[tuple[int, bool]]]:
    """Each kind's rows as (row number, done)."""

    kinds = answer["kinds"]
    assert isinstance(kinds, list)
    return {kind["kind"]: [(row["row"], row["done"]) for row in kind["rows"]] for kind in kinds}


def _columns(answer: dict[str, object]) -> dict[str, int]:
    columns = answer["description_columns"]
    assert isinstance(columns, list)
    return {column["column"]: column["rows"] for column in columns}


def _check_screen(run: Run, page: Page, answer: dict[str, object], where: str) -> None:
    """The summary shows every kind's counter and row numbers and every column counter of the server's answer."""

    kinds = answer["kinds"]
    assert isinstance(kinds, list)
    for kind in kinds:
        rows = [(row["row"], row["done"]) for row in kind["rows"]]
        section = page.get_by_role("region", name=KIND_TITLES[kind["kind"]], exact=True)
        if rows:
            expect(section).to_contain_text(f"{sum(done for _, done in rows)} z {len(rows)}")
        else:
            expect(section).to_contain_text("Nie ma takich wierszy.")
        expect(section.get_by_role("img", name="Zrobione" if kind["done"] else "Do uzupełnienia")).to_be_visible()
        expect(section.get_by_role("listitem")).to_have_text([f"{row}{' zrobiony' if done else ''}" for row, done in rows])
    columns = _columns(answer)
    reminder = page.get_by_role("region", name="Kolumny z opisu technika", exact=True)
    expect(reminder.locator("dt")).to_have_text(list(columns))
    expect(reminder.locator("dd")).to_have_text([_in_rows(rows) for rows in columns.values()])
    run.check(True, f"{where}: counters, row numbers with their done state, and column counters match the server")


def _write_dad_value(collective_file: Path, row: int) -> None:
    """Dad types one quantity in Excel and saves (harness: openpyxl edits the run's Plik zbiorczy, no lock file)."""

    workbook = load_workbook(collective_file)
    try:
        sheet = workbook.active
        assert isinstance(sheet, Worksheet)
        sheet.cell(row, column_index_from_string(DAD_COLUMN)).value = 1
        workbook.save(collective_file)
    finally:
        workbook.close()


def _light_theme(page: Page) -> bool:
    """Switches to the light theme for the screenshot; True when it was dark and must be switched back."""

    light = page.locator("header").get_by_role("button", name="Włącz motyw jasny", exact=True)
    if light.count() == 0:
        return False
    light.click()
    expect(page.locator("header").get_by_role("button", name="Włącz motyw ciemny", exact=True)).to_be_visible()
    return True


def drive_filling(run: Run, page: Page) -> None:
    _, collective_file = period_files(run)
    if not collective_file.is_file():
        # Dad's path starts with Przygotowanie pliku zbiorczego; alone, this feature drives it first.
        drive_preparation(run, page)
    original = collective_file.read_bytes()
    lock = collective_file.with_name(f"~${collective_file.name}")
    try:
        _drive(run, page, collective_file, lock)
    finally:
        # The harness's value must not reach Przebieg rozliczeń, which continues from this Plik zbiorczy.
        collective_file.write_bytes(original)
        lock.unlink(missing_ok=True)

    run.step(FEATURE, "uzupelnienie-bez-pliku")
    page.goto(f"{run.url.rstrip('/')}/#/uzupelnienie/{EMPTY_PERIOD}")
    answer = run.api(f"/api/periods/{EMPTY_PERIOD}/quantities-to-fill", "uzupelnienie-bez-pliku")
    message = answer["message"]
    assert isinstance(message, dict)
    notice = page.get_by_label("Komunikat", exact=True)
    for part in ("what", "where", "action"):
        expect(notice).to_contain_text(message[part])
    run.check(
        answer["state"] == "problem" and page.get_by_role("alert").filter(has_text="Nie udało się").count() == 0,
        "a period without a Plik zbiorczy shows the server's three-part Komunikat, calm and not as a failure",
    )
    log = (run.root / "logs" / "centrum-rozliczen.log").read_text(encoding="utf-8").splitlines()
    run.check(
        any(" uzupelnienie-odczyt-nieudany " in line and f"okres={EMPTY_PERIOD}" in line for line in log),
        "the failed read is in the Centrum log with the period",
    )
    run.capture(page, "224-uzupelnienie-bez-pliku")


def _drive(run: Run, page: Page, collective_file: Path, lock: Path) -> None:
    run.step(FEATURE, "uzupelnienie-karta")
    page.goto(run.url)
    next_step = page.get_by_role("region", name="Następny krok")
    expect(next_step).to_contain_text(LAST_COMPLETE)
    expect(next_step).to_contain_text("Utwórz rozliczenia pracowników")
    fill = next_step.get_by_role("button", name="Uzupełnij", exact=True)
    expect(fill).to_be_visible()
    expect(next_step.get_by_role("button", name="Podgląd bez zapisu", exact=True)).to_be_visible()
    board = run.api("/api/periods", "uzupelnienie-karta")
    periods = board["periods"]
    assert isinstance(periods, list)
    [last] = [period for period in periods if period["last_complete"]]
    run.check(
        last["token"] == LAST_COMPLETE and last["status"] == "collective_file_prepared",
        "Następny krok offers Uzupełnij for the last complete period with Plik zbiorczy przygotowany, "
        "with Podgląd bez zapisu next to it",
    )
    run.capture(page, "220-uzupelnienie-karta")

    run.step(FEATURE, "uzupelnienie-podsumowanie")
    log_path = run.root / "logs" / "centrum-rozliczen.log"
    since = len(log_path.read_text(encoding="utf-8").splitlines())
    fill.click()
    expect(page).to_have_url(f"{run.url.rstrip('/')}/{SUMMARY}")
    expect(page.get_by_role("heading", level=1)).to_have_text("Uzupełnienie pliku zbiorczego")
    expect(page.get_by_role("main").get_by_text(LAST_COMPLETE, exact=True)).to_be_visible()
    before = _ready(run, "uzupelnienie-podsumowanie")
    run.check(
        _rows(before) == {kind: [(row, False) for row in rows] for kind, rows in EXPECTED_ROWS.items()},
        "the server lists rows 18 and 20 without rule and row 19 from the description, none done yet",
    )
    _check_screen(run, page, before, "after Uzupełnij")
    menu = page.get_by_role("list", name="Sekcje Centrum rozliczeń")
    expect(menu.get_by_role("link", name="Okresy", exact=True)).to_have_attribute("data-active", "")
    run.wait_for(
        lambda: any(
            " pomiar " in line and "strona=uzupelnienie" in line and "getQuantitiesToFill:" in line
            for line in log_path.read_text(encoding="utf-8").splitlines()[since:]
        ),
        "the view is measured as page uzupelnienie with its getQuantitiesToFill request",
    )
    run.capture(page, "221-uzupelnienie-podsumowanie")

    run.step(FEATURE, "uzupelnienie-wpis-taty")
    row = EXPECTED_ROWS["without_rule"][0]
    asked: list[str] = []
    page.on("request", lambda request: asked.append(urlsplit(request.url).path))
    page.evaluate("() => { window.__bez_przeladowania = true }")
    _write_dad_value(collective_file, row)
    run.check(not lock.exists(), f"harness: one value typed into column {DAD_COLUMN} of row {row}, saved and closed")
    without_rule = page.get_by_role("region", name=KIND_TITLES["without_rule"], exact=True)
    expect(without_rule).to_contain_text("1 z 2", timeout=LIVE_TIMEOUT_MS)
    run.check(
        page.evaluate("() => window.__bez_przeladowania === true") and asked.count(VERSION_API) >= 1,
        f"the counter changed to 1 z 2 without a click or reload, after {asked.count(VERSION_API)} "
        "check(s) of the saved version",
    )
    after = _ready(run, "uzupelnienie-po-wpisie")
    columns_before, columns_after = _columns(before), _columns(after)
    run.check(
        _rows(after) == {**_rows(before), "without_rule": [(18, True), (20, False)]}
        and columns_after == {**columns_before, DAD_COLUMN: columns_before[DAD_COLUMN] + 1},
        f"row {row} is done and only column {DAD_COLUMN} counts one more row",
    )
    views = asked.count(API)
    for _ in range(2):
        with page.expect_request(lambda request: urlsplit(request.url).path == VERSION_API, timeout=LIVE_TIMEOUT_MS):
            pass
    run.check(asked.count(API) == views, "two more checks of the unchanged saved version read nothing again")
    _check_screen(run, page, after, "after Dad's value")
    switched = _light_theme(page)
    run.capture(page, "222-uzupelnienie-po-wpisie")
    if switched:
        page.locator("header").get_by_role("button", name="Włącz motyw ciemny", exact=True).click()

    run.step(FEATURE, "uzupelnienie-blokada-excela")
    lock.write_bytes(b"\x00" * 165)
    page.reload()
    held = _ready(run, "uzupelnienie-blokada-excela")
    held_read = held["read"]
    assert isinstance(held_read, dict)
    expect(page.get_by_role("status").filter(has_text="Stan z")).to_have_count(0)
    run.check(
        held_read["retrying"] is False and _without_read(held) == _without_read(after),
        "with Excel's lock file next to the Plik zbiorczy the server reads its last saved version unchanged",
    )
    _check_screen(run, page, held, "with Excel's lock file")
    run.capture(page, "223-uzupelnienie-blokada-excela")

    run.step(FEATURE, "uzupelnienie-odczyt-nieudany")
    lock.unlink()
    saved = collective_file.read_bytes()
    collective_file.write_bytes(b"Excel zapisuje")
    status = page.get_by_role("status").filter(has_text="Stan z")
    expect(status).to_be_visible(timeout=LIVE_TIMEOUT_MS)
    kept = _ready(run, "uzupelnienie-odczyt-nieudany")
    read = kept["read"]
    assert isinstance(read, dict)
    run.check(
        read["retrying"] is True and _without_read(kept) == _without_read(held),
        "a Plik zbiorczy Centrum cannot read keeps the last good state, marked as retried",
    )
    _check_screen(run, page, kept, "while the read fails")
    expect(page.get_by_label("Komunikat", exact=True)).to_have_count(0)
    run.capture(page, "229-uzupelnienie-odczyt-nieudany")
    collective_file.write_bytes(saved)
    expect(status).to_have_count(0, timeout=LIVE_TIMEOUT_MS)
    run.check(True, "once the file reads again the time disappears, without a click")

    run.step(FEATURE, "uzupelnienie-otworz")
    page.get_by_role("button", name="Otwórz Plik zbiorczy", exact=True).click()
    run.wait_for(last_opened_is(run, collective_file), "Otwórz Plik zbiorczy opened the period's Plik zbiorczy")
    expect(page.get_by_text(collective_file.name, exact=True).last).to_be_visible()

    run.step(FEATURE, "uzupelnienie-powrot")
    page.get_by_role("main").get_by_role("link", name="Okresy", exact=True).click()
    expect(page.get_by_role("heading", level=1)).to_have_text("Okresy")
    page.go_back()
    expect(page.get_by_role("heading", level=1)).to_have_text("Uzupełnienie pliku zbiorczego")
    run.check(True, "the link Okresy leads back, and Wstecz returns to the summary of the same period")
