"""Drive features/przebieg-rozliczen.md."""

from __future__ import annotations

import json
from pathlib import Path
import urllib.request

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from playwright.sync_api import Page, expect

from drive.common import (
    LAST_COMPLETE,
    ROW_MARKER,
    Run,
    active_run_signs,
    current_run,
    figure,
    finished_result,
    history_rows,
    last_opened_is,
    period_choice,
    period_files,
    period_row,
)
from drive.przygotowanie_pliku_zbiorczego import drive_preparation

SETTLEMENT = "przebieg-rozliczen"
WORKER_SETTLEMENTS = f"Rozliczenie pracowników {LAST_COMPLETE}"
# serve_centrum.py waits this long before each Szablon pracownika while the pacing file exists.
TEMPLATE_PACING_MS = 2500
WORKER_COLUMN, QUANTITY_COLUMN, FIRST_DATA_ROW = 8, 14, 18  # H (WYKONAWCA), N (first quantity), row 18


def _settlement_records(run: Run, name: str) -> list[dict[str, object]]:
    """Valid records of the Przebieg rozliczeń metrics file as the server reads them, newest first."""

    runs = run.api("/api/metrics", name)["runs"]
    assert isinstance(runs, list)
    return [entry for entry in runs if entry["process"] == "settlement"]


def _fill_readiness(collective_file: Path) -> int:
    """Gotowość as Dad reaches it in Excel: a quantity in N of every Wiersz danych, saved and closed.

    Excel is not driven; the harness edits the workbook like Dad would and leaves no lock file.
    Returns the number of filled Wiersze danych, never a cell value.
    """

    workbook = load_workbook(collective_file)
    sheet = workbook.active
    assert isinstance(sheet, Worksheet)
    filled = 0
    for row in range(FIRST_DATA_ROW, sheet.max_row + 1):
        if sheet.cell(row, WORKER_COLUMN).value:
            sheet.cell(row, QUANTITY_COLUMN).value = 1
            filled += 1
    workbook.save(collective_file)
    return filled


def _post(run: Run, path: str, body: dict[str, object], name: str) -> int:
    request = urllib.request.Request(
        run.url.rstrip("/") + path,
        data=json.dumps(body).encode(),
        method="POST",
        headers={"Content-Type": "application/json", "Origin": run.url.rstrip("/")},
    )
    return run.refused(request, name)


def drive_settlement(run: Run, page: Page) -> None:
    feature = SETTLEMENT
    internal_settlement, collective_file = period_files(run)
    period_folder = collective_file.parent
    worker_settlements = period_folder / WORKER_SETTLEMENTS
    workers = run.fixture["workers"]
    files = [f"Rozliczenie {LAST_COMPLETE} - {worker}.xlsx" for worker in workers]
    needs_check = run.fixture["settlement"]["needs_check"]
    if not collective_file.is_file():
        # Dad's path starts with Przygotowanie pliku zbiorczego; alone, this feature drives it first.
        drive_preparation(run, page)
    records_before = len(_settlement_records(run, "przebieg-metryki-przed"))

    run.step(feature, "przebieg-gotowosc")
    filled = _fill_readiness(collective_file)
    run.check(
        filled == len(workers) and not collective_file.with_name(f"~${collective_file.name}").exists(),
        "Gotowość (Excel simulated by the harness): quantities typed into every Wiersz danych, saved and closed",
    )

    run.step(feature, "przebieg-wybor-okresu")
    page.goto(run.url)
    page.get_by_role("button", name=f"Utwórz rozliczenia pracowników, okres {LAST_COMPLETE}", exact=True).click()
    expect(page.get_by_text("Nowe Uruchomienie procesu", exact=True)).to_be_visible()
    expect(page.get_by_role("radio", name="Przebieg rozliczeń")).to_be_checked()
    expect(page.get_by_role("combobox", name="Okres")).to_contain_text(period_choice(LAST_COMPLETE, "ostatni pełny"))
    expect(page.get_by_role("group", name="Plik zbiorczy")).to_contain_text(collective_file.name)
    expect(page.get_by_text("Zapisz i zamknij Plik zbiorczy w Excelu", exact=True)).to_be_visible()
    processes = run.api("/api/runs/setup", "przebieg-wybor-okresu")["processes"]
    assert isinstance(processes, list)
    [setup] = [process for process in processes if process["process"] == "settlement"]
    run.check(
        setup["preselected_period"] == LAST_COMPLETE
        and [choice["token"] for choice in setup["periods"]] == [LAST_COMPLETE, needs_check]
        and setup["blocker"] is None,
        "only periods with Status Plik zbiorczy przygotowany are offered; the Plik zbiorczy follows from the period",
    )
    run.capture(page, "40-przebieg-nowe-uruchomienie")

    run.step(feature, "przebieg-podglad")
    page.goto(run.url)
    next_step = page.get_by_role("region", name="Następny krok")
    expect(next_step).to_contain_text("Utwórz rozliczenia pracowników")
    next_step.get_by_role("button", name="Podgląd bez zapisu").click()
    live = page.get_by_role("region", name="Podgląd na żywo")
    expect(live).to_contain_text("Podgląd gotowy · nic nie zapisano")
    result = finished_result(run, "przebieg-podglad")
    summary = result["summary"]
    assert isinstance(summary, dict)
    run.check(
        result["result"] == "OK"
        and [file["file"] for file in summary["file_plan"]] == files
        and [file["rows"] for file in summary["file_plan"]] == run.fixture["settlement"]["rows_per_worker"],
        "the DRY-RUN plans one Rozliczenie pracownika per Wykonawca with its number of rows",
    )
    plan = live.get_by_role("table", name="Plan plików")
    expect(plan.get_by_role("row")).to_have_count(len(files))
    expect(plan.get_by_role("row").filter(has_text=files[0])).to_contain_text("1 wiersz")
    run.check(not worker_settlements.exists(), "the DRY-RUN created no Folder rozliczeń pracowników")
    records = _settlement_records(run, "przebieg-metryki-podglad")
    run.check(
        len(records) == records_before + 1 and records[0]["mode"] == "DRY-RUN",
        "the DRY-RUN appended one record to the Przebieg rozliczeń metrics file",
    )
    run.capture(page, "41-przebieg-podglad")

    run.step(feature, "przebieg-potwierdzenie")
    live.get_by_role("button", name="Uruchom", exact=True).click()
    confirm = page.get_by_role("alertdialog", name="Uruchomić z zapisem?")
    expect(confirm).to_contain_text("Zapisz i zamknij Plik zbiorczy w Excelu")
    expect(confirm).to_contain_text("„Uruchom” potwierdza Gotowość Pliku zbiorczego")
    expect(confirm.get_by_text("Nie było podglądu bez zapisu")).to_have_count(0)
    run.check(True, "Uruchom asks to confirm the write; the click is Dad's confirmation of Gotowość")
    run.capture(page, "42-przebieg-potwierdzenie")
    pacing = run.root / "pacing" / "szablony-ms.txt"
    pacing.parent.mkdir(exist_ok=True)
    pacing.write_text(str(TEMPLATE_PACING_MS), encoding="utf-8")
    confirm.get_by_role("button", name="Uruchom", exact=True).click()

    run.step(feature, "przebieg-na-zywo")
    try:
        expect(live.get_by_role("progressbar", name="Szablony pracowników")).to_be_visible()
        expect(live).to_contain_text(f"Szablony pracowników · {workers[0]}")
        counters = live.locator('dl[aria-label="Liczniki"]')
        expect(counters).to_contain_text("Zapisane szablony0")
        # The second Szablon pracownika is paced, so the counters stay at one written while it runs.
        expect(live).to_contain_text(f"Szablony pracowników · {workers[1]}", timeout=10_000)
        expect(counters).to_contain_text("Zapisane szablony1")
        view = current_run(run, "przebieg-na-zywo")
        progress = view["progress"]
        assert isinstance(progress, dict)
        run.check(
            view["state"] == "running" and progress["template_total"] == len(workers),
            "the live view shows the Szablony pracowników progress from the projection (harness pacing on)",
        )
        rows = figure(progress, "Wiersze danych")
        expect(counters).to_contain_text(f"Wiersze danych{rows}")
        run.check(
            figure(progress, "Zapisane szablony") == 1 and rows == run.fixture["settlement"]["rows_per_worker"][0],
            "the live counters (Wiersze danych, Zapisane szablony) update from the projection during the run",
        )
        close = urllib.request.Request(
            run.url.rstrip("/") + "/api/shutdown", data=b"", method="POST", headers={"Origin": run.url.rstrip("/")}
        )
        run.check(
            run.refused(close, "przebieg-zamkniecie-odmowa") == 409,
            "closing Centrum is refused while the Przebieg rozliczeń runs",
        )
        run.capture(page, "43-przebieg-na-zywo", live=True)

        run.step(feature, "przebieg-znak-uruchomienia")
        menu_sign, header_link = active_run_signs(page, "Trwa Uruchomienie")
        page.get_by_role("link", name="Okresy", exact=True).click()
        expect(page.get_by_role("heading", name="Okresy", level=1)).to_be_visible()
        expect(menu_sign).to_be_visible()
        expect(header_link).to_be_visible()
        run.check(
            current_run(run, "przebieg-znak-uruchomienia")["state"] == "running",
            "on Okresy during the RUN the menu and the header show „Trwa Uruchomienie” (harness pacing on)",
        )
        run.capture(page, "150-przebieg-znak-uruchomienia", live=True)  # the sign's spinner turns while it runs
        header_link.click()
        expect(live).to_be_visible()
        expect(header_link).to_have_count(0)
        run.check(page.url.endswith("#/uruchomienia"), "the header link leads back to the live view; there it is gone")
    finally:
        pacing.unlink(missing_ok=True)

    run.step(feature, "przebieg-zapisane")
    outcome = page.get_by_label("Wynik")
    expect(outcome).to_contain_text("Rozliczenia pracowników gotowe", timeout=20_000)
    result = finished_result(run, "przebieg-zapisane")
    summary = result["summary"]
    assert isinstance(summary, dict)
    run.check(result["result"] == "OK" and figure(summary, "Zapisane szablony") == len(workers), "the RUN finished OK")
    run.check(
        sorted(path.name for path in worker_settlements.iterdir()) == sorted(files),
        "the Folder rozliczeń pracowników holds one Rozliczenie pracownika per Wykonawca",
    )
    run.check(
        sorted(path.name for path in period_folder.iterdir())
        == sorted([internal_settlement.name, collective_file.name, WORKER_SETTLEMENTS]),
        "no staging folder is left in the Folder okresu: the folder was published at once",
    )
    records = _settlement_records(run, "przebieg-metryki-zapisane")
    run.check(
        len(records) == records_before + 2
        and [(record["mode"], record["result"], record["origin"]) for record in records[:2]]
        == [("RUN", "OK", "centrum"), ("DRY-RUN", "OK", "centrum")],
        "both Uruchomienia appended one record each to the process's metrics file, marked as started from Centrum",
    )
    log = (run.root / "logs" / "centrum-rozliczen.log").read_text(encoding="utf-8")
    run.check(
        f"uruchomienie-koniec proces=settlement tryb=RUN okres={LAST_COMPLETE} wynik=OK koniec=finished" in log
        and ROW_MARKER not in log,
        "the Centrum log records the Przebieg rozliczeń with safe facts only",
    )
    outcome.get_by_role("button", name="Folder okresu").click()
    run.wait_for(last_opened_is(run, period_folder), "Folder okresu opens in Explorer")
    run.capture(page, "44-przebieg-zapisane")

    run.step(feature, "przebieg-odczyt-po-uruchomieniu")
    # Same tab, no reload: the history and the period choice Centrum showed before the RUN are read again.
    page.get_by_role("tab", name="Historia", exact=True).click()
    newest, previous = history_rows(page).nth(1), history_rows(page).nth(2)
    expect(newest).to_contain_text("Przebieg rozliczeń")
    expect(newest).not_to_contain_text("DRY-RUN")
    expect(previous).to_contain_text("Przebieg rozliczeń")
    expect(previous).to_contain_text("DRY-RUN")
    run.check(True, "the Historia tab shows the finished RUN and its preview at the top")
    page.get_by_role("tab", name="Uruchomienie", exact=True).click()
    page.get_by_role("button", name="Nowe Uruchomienie procesu", exact=True).click()
    page.get_by_role("radio", name="Przebieg rozliczeń").click()
    expect(page.get_by_role("combobox", name="Okres")).to_contain_text(needs_check)
    processes = run.api("/api/runs/setup", "przebieg-odczyt-po-uruchomieniu")["processes"]
    assert isinstance(processes, list)
    [setup] = [process for process in processes if process["process"] == "settlement"]
    run.check(
        [choice["token"] for choice in setup["periods"]] == [needs_check],
        "the period choice is read again: the closed period is no longer offered",
    )
    run.capture(page, "48-przebieg-odczyt-po-uruchomieniu")

    run.step(feature, "przebieg-status-okresu")
    page.goto(run.url)
    expect(period_row(page, LAST_COMPLETE)).to_contain_text("Rozliczenia pracowników gotowe")
    expect(page.get_by_role("region", name="Następny krok")).to_contain_text("Okres zamknięty")
    board = run.api("/api/periods", "przebieg-status-okresu")["periods"]
    assert isinstance(board, list)
    [row] = [row for row in board if row["token"] == LAST_COMPLETE]
    run.check(
        row["status"] == "worker_settlements_ready" and row["worker_settlement_count"] == len(workers),
        "the Status okresu is Rozliczenia pracowników gotowe with three Rozliczenia pracowników",
    )
    run.capture(page, "45-przebieg-status-okresu")

    run.step(feature, "przebieg-nienadpisywanie")
    before = {path.name: path.read_bytes() for path in worker_settlements.iterdir()}
    started = _post(
        run,
        "/api/runs",
        {"process": "settlement", "period": LAST_COMPLETE, "dry_run": False},
        "przebieg-nienadpisywanie-start",
    )
    run.check(started == 200, "a second RUN of the closed period starts (the Okres list no longer offers it)")
    page.goto(run.url + "#/uruchomienia")
    stop = page.get_by_label("Zatrzymanie")
    expect(stop).to_contain_text("Folder rozliczeń pracowników tego okresu już istnieje, więc nie został nadpisany.")
    expect(stop).to_contain_text(f"folder {LAST_COMPLETE}\\{WORKER_SETTLEMENTS}")
    result = finished_result(run, "przebieg-nienadpisywanie")
    stops = result["stops"]
    assert isinstance(stops, list)
    [message] = stops
    assert isinstance(message, dict)
    run.check(
        result["ending"] == "stop"
        and all(message[part] for part in ("what", "where", "action"))
        and str(run.root) not in json.dumps(result, ensure_ascii=False),
        "the stop has three parts and no full path",
    )
    run.check(
        {path.name: path.read_bytes() for path in worker_settlements.iterdir()} == before,
        "Nienadpisywanie: the Rozliczenia pracowników are unchanged",
    )
    run.capture(page, "46-przebieg-nienadpisywanie")

    run.step(feature, "przebieg-wymaga-sprawdzenia")
    page.get_by_role("button", name="Nowe Uruchomienie procesu").click()
    page.get_by_role("radio", name="Przebieg rozliczeń").click()
    expect(page.get_by_role("combobox", name="Okres")).to_contain_text(needs_check)
    page.get_by_role("button", name="Podgląd bez zapisu", exact=True).click()
    outcome = page.get_by_label("Wynik")
    expect(outcome).to_contain_text("Podgląd wymaga sprawdzenia · nic nie zapisano")
    expect(outcome.get_by_role("button", name="Uruchom ponownie")).to_be_visible()
    expect(page.get_by_label("Ostrzeżenie")).to_contain_text("Nieznany identyfikator WYKONAWCA")
    result = finished_result(run, "przebieg-wymaga-sprawdzenia")
    run.check(
        (result["result"], result["next_action"]) == ("WYMAGA_SPRAWDZENIA", "rerun"),
        "a Wykonawca missing from the konfiguracja wykonawców gives Wymaga sprawdzenia, not a success",
    )
    run.check(
        not (Path(run.fixture["periods_directory"]) / needs_check / f"Rozliczenie pracowników {needs_check}").exists(),
        "the preview wrote nothing",
    )
    run.capture(page, "47-przebieg-wymaga-sprawdzenia")
