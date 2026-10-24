"""Check that the Rozliczenie pracownika made by the template copy (ADR 0008) is the one openpyxl would make.

Przebieg rozliczeń writes every Rozliczenie pracownika with ``template_copy.py``: it splices the
Wiersze danych into the bytes of the Wzorzec rozliczenia pracownika instead of loading and saving the
workbook with openpyxl. This checker generates Wiersze danych with Hypothesis (every value type a cell
can hold, edge texts and numbers, Wyróżnienia, Hiperłącza, row counts on both sides of each compression chunk boundary), writes
them through the real ``TemplateFile.create()`` and, as the reference, through the old openpyxl path: the
same ``TemplateFile.create()`` on a twin of the Wzorzec that the copy refuses (an XML comment in its sheet
data, which openpyxl skips). It checks five rules:

- WIERNA_KOPIA: the two files are the same for what Dad sees in Excel, as judged by the benchmark's
  correctness gate (``tools/porownanie_rozliczen.py``): values, formulas, resolved styles with Wyróżnienia,
  Hiperłącza, merges, dimensions, defined names, external links and recalculation flags.
- TA_SAMA_ODMOWA: the copy refuses a value exactly when the openpyxl path refuses it, and a refusal
  leaves no file.
- WYGLAD_WIERSZA: every cell A:L of a written Wiersz danych has the Wygląd Wiersza danych, derived here from the
  rule and not from the code: font face and size of the Wzorzec's header cell in row 17, left and centre alignment,
  wrapping exactly for text with a line break, a thin automatic border on four sides and the number format of the
  Plik zbiorczy cell; a written row stores no height.
- NIETKNIETY_WZORZEC: what ADR 0008 says stays byte for byte stays: every zip part except the sheet,
  ``xl/workbook.xml`` (only ``calcPr``), ``xl/styles.xml`` (only appended number formats, fonts, fills, borders and
  styles) and the sheet relationships (only appended relationships of Hiperłącza, or a new part at the end when the
  Wzorzec has none), every sheet row outside the written Wiersze danych apart from one inserted ``<hyperlinks>``,
  the row attributes of written rows apart from the removed height (``ht``, ``customHeight``), and the cells right
  of AT in the written rows.
- KSZTALT_EXCELA: the file has the shape Excel needs and openpyxl does not check when reading it:
  valid zip CRCs, rows and cells in ascending order with matching references, ``dimension`` covering
  the written rows, ``xml:space="preserve"`` on text with spaces at its edges, style lists whose counts and
  references match, sheet elements in schema order, a content type for every part, unique relationship ids and
  every hyperlink pointing at an external hyperlink relationship.

One Hypothesis example is one Przebieg rozliczeń: one ``TemplateFile`` loaded once and
1 to 3 Rozliczenia pracowników made from it in a row, as ``run_settlements()`` does for its workers.
"""

from __future__ import annotations

import argparse
import re
import shutil
import sys
import tempfile
import time as clock
import warnings
import zipfile
from collections import Counter
from collections.abc import Callable, Iterator, Sequence
from dataclasses import dataclass, field, replace
from datetime import date, datetime, time, timedelta, timezone
from decimal import Decimal
from pathlib import Path
from typing import Any
from unittest import mock
from xml.etree import ElementTree

REPO_ROOT = Path(__file__).resolve().parents[4]
SOURCE_ROOT = REPO_ROOT / "src"
sys.path[:0] = [str(Path(__file__).resolve().parents[4] / folder) for folder in ("src", "tools", ".")]

from hypothesis import HealthCheck, Phase, example, given, seed, settings
from hypothesis import strategies as st
from hypothesis.internal.conjecture import engine as hypothesis_engine
from hypothesis.internal.cache import LRUCache
from hypothesis.internal.conjecture import providers as hypothesis_providers
from hypothesis.internal.constants_ast import Constants
from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Color, Font, GradientFill, PatternFill
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.formula import ArrayFormula

import porownanie_rozliczen as gate
from rozliczenia import workbook_files
from rozliczenia.highlights import PLAIN, Highlights
from rozliczenia.hyperlinks import Hyperlink
from rozliczenia.workbook_files import DataRow, ExcelRow, Layout, TemplateFile
from syntetyczne_wzorce import write_settlement_template
from tests.fixtures.settlement import write_template
from tests.fixtures.workbooks import add_external_link_fixture, has_xml_comment, refused_by_copy

VERIFIED_FILE = "template_copy.py"
SHEET = "xl/worksheets/sheet1.xml"
WORKBOOK = "xl/workbook.xml"
STYLES = "xl/styles.xml"
SHEET_RELS = "xl/worksheets/_rels/sheet1.xml.rels"
CONTENT_TYPES = "[Content_Types].xml"
ORDER_NUMBER = 6  # F, NUMER ZL.: the column whose Hiperłącza Dad clicks
FIRST_DATA_ROW = workbook_files.FIRST_DATA_ROW
HEADER_ROW = workbook_files.HEADER_ROW
DATA_COLUMNS = workbook_files.DATA_ROW_COLUMNS[-1]  # AT
LOOK_COLUMNS = 12  # A:L: the columns with the Wygląd Wiersza danych
ROWS_PER_CHUNK = 64  # Excel's shared-formula block and the copy's compression chunk (``_ROWS_PER_CHUNK``)
DEFAULT_EXAMPLES = 60
# Hypothesis shrinks a counterexample for up to 300 s; a minute gives a short story here and keeps a failing run cheap.
SHRINK_SECONDS = 60
CELL_DIFFERENCES_SHOWN = 5
MAX_TEXT = 32767

_ROW = re.compile(rb'<row r="(\d+)"[^>]*?(?:/>|>.*?</row>)', re.S)
_CELL = re.compile(rb'<c r="([A-Z]{1,3})(\d+)"([^>]*?)(/>|>.*?</c>)', re.S)
_DIMENSION = re.compile(rb'<dimension ref="([^"]*)"\s*/>')
_CALC_PR = re.compile(rb"<(?:[\w.-]+:)?calcPr\b[^>]*?/?>")
_COUNT = re.compile(rb'\scount="(\d+)"')
_FIRST_CUSTOM_FORMAT = 164  # numFmtId below it are Excel's built-in formats
_XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
_MAIN = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_OFFICE_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PACKAGE = "{http://schemas.openxmlformats.org/package/2006/relationships}"
_TYPES = "{http://schemas.openxmlformats.org/package/2006/content-types}"
_HYPERLINKS_BLOCK = re.compile(rb"<hyperlinks\b[^>]*>.*?</hyperlinks>", re.S)
_RELATIONSHIP = re.compile(rb"<Relationship\b[^>]*/>")
# Children of <worksheet> in the order of the schema (CT_Worksheet), written down here independently of the copy.
WORKSHEET_ORDER = (
    "sheetPr", "dimension", "sheetViews", "sheetFormatPr", "cols", "sheetData", "sheetCalcPr", "sheetProtection",
    "protectedRanges", "scenarios", "autoFilter", "sortState", "dataConsolidate", "customSheetViews", "mergeCells",
    "phoneticPr", "conditionalFormatting", "dataValidations", "hyperlinks", "printOptions", "pageMargins", "pageSetup",
    "headerFooter", "rowBreaks", "colBreaks", "customProperties", "cellWatches", "ignoredErrors", "smartTags", "drawing",
    "legacyDrawing", "legacyDrawingHF", "drawingHF", "picture", "oleObjects", "controls", "webPublishItems", "tableParts",
    "extLst",
)


@dataclass(frozen=True)
class StyleList:
    """A list in ``xl/styles.xml`` the copy may only append to; a cell style refers to its entries by position."""

    tag: str
    block: re.Pattern[bytes]
    entry: re.Pattern[bytes]

    @classmethod
    def of(cls, tag: str, item: str) -> StyleList:
        block = re.compile(rb"<%b\b([^>]*?)(?:/>|>(.*?)</%b>)" % (tag.encode(), tag.encode()), re.S)
        return cls(tag, block, re.compile(rb"<%b\b[^>]*?(?:/>|>.*?</%b>)" % (item.encode(), item.encode()), re.S))

    def entries(self, styles: bytes) -> list[bytes]:
        found = self.block.search(styles)
        return self.entry.findall(found.group(2) or b"") if found else []

    def count(self, styles: bytes) -> int | None:
        found = self.block.search(styles)
        count = _COUNT.search(found.group(1)) if found else None
        return int(count.group(1)) if count else None

    def without(self, styles: bytes) -> bytes:
        return self.block.sub(b"", styles, count=1)


NUM_FMTS, FONTS, FILLS, BORDERS, CELL_XFS = (
    StyleList.of("numFmts", "numFmt"),
    StyleList.of("fonts", "font"),
    StyleList.of("fills", "fill"),
    StyleList.of("borders", "border"),
    StyleList.of("cellXfs", "xf"),
)
STYLE_LISTS = (NUM_FMTS, FONTS, FILLS, BORDERS, CELL_XFS)
_XF_ID = {name: re.compile(rb'\s%b="(\d+)"' % name.encode()) for name in ("numFmtId", "fontId", "fillId", "borderId")}
_ROW_HEIGHT = re.compile(rb'\s(?:ht|customHeight)="[^"]*"')


# --- Wzorce ----------------------------------------------------------------------------------------


@dataclass(frozen=True)
class Template:
    """A Wzorzec rozliczenia pracownika under test, with what the checker needs to know about its sheet."""

    name: str
    description: str
    path: Path
    content: bytes
    cells: frozenset[tuple[int, int]]  # (row, column) of every cell element in the sheet XML
    chunk_rows: tuple[int, ...]  # sheet rows within reach where a compressed chunk of the copy starts
    last_row: int  # last row of the sheet XML
    max_rows: int  # bound on Wiersze danych per Rozliczenie pracownika
    header_faces: tuple[tuple[str | None, float | None], ...]  # font face and size of A17:L17, the Wygląd's source

    @classmethod
    def of(cls, name: str, description: str, path: Path, max_rows: int | None = None) -> Template:
        content = path.read_bytes()
        xml = zipfile.ZipFile(path).read(SHEET)
        rows = [int(match.group(1)) for match in _ROW.finditer(xml)]
        data_rows = [number for number in rows if number >= FIRST_DATA_ROW]
        cells = frozenset((int(match.group(2)), column_index_from_string(match.group(1).decode())) for match in _CELL.finditer(xml))
        last_row = rows[-1] if rows else 0
        bound = max_rows if max_rows is not None else max(last_row - FIRST_DATA_ROW + 1, 0) + 10
        # Chunk starts a Rozliczenie pracownika within the bound can reach; later ones only hold Wzorzec rows.
        chunk_rows = tuple(row for row in data_rows[::ROWS_PER_CHUNK] if row <= FIRST_DATA_ROW + bound)
        workbook = load_workbook(path)
        header = workbook.active
        faces = tuple(
            (header.cell(HEADER_ROW, column).font.name, header.cell(HEADER_ROW, column).font.sz)  # type: ignore[union-attr]
            for column in range(1, LOOK_COLUMNS + 1)
        )
        workbook.close()
        return cls(name, description, path, content, cells, chunk_rows, last_row, bound, faces)

    def interesting_counts(self) -> list[int]:
        """Row counts whose last written row is right before, on, or right after a chunk start or the sheet end."""

        edges = [*self.chunk_rows, self.last_row, self.last_row + 1]
        counts = {0, 1, self.max_rows}
        for row in edges:
            counts.update(row - FIRST_DATA_ROW + offset for offset in (0, 1, 2))
        return sorted(count for count in counts if 0 <= count <= self.max_rows)


SYNTHETIC_TEMPLATES = ("syntetyczny", "maly")


def build_templates(directory: Path, real_template: Path | None, names: Sequence[str] = SYNTHETIC_TEMPLATES) -> list[Template]:
    """The synthetic Wzorce in ``names``, or only Dad's real Wzorzec when it is given."""

    if real_template is not None:
        copied = directory / "wzorzec-taty.xlsx"
        shutil.copyfile(real_template, copied)
        return [Template.of("wzorzec-taty", f"prawdziwy Wzorzec: {real_template.name}", copied, max_rows=700)]
    templates = []
    if "syntetyczny" in names:
        synthetic = directory / "syntetyczny.xlsx"
        write_settlement_template(synthetic, "GODZINA", data_rows=200, comments=True)
        templates.append(
            Template.of(
                "syntetyczny",
                "kształt Wzorca taty z tools/syntetyczne_wzorce.py, 200 Wierszy danych: wspólne formuły, calcChain, "
                "link zewnętrzny, komentarz z relacjami arkusza rId1 i rId2, wiersze zer w N:V i X, A i B bez komórek od wiersza 139",
                synthetic,
                max_rows=260,
            )
        )
    if "maly" in names:
        small = directory / "maly.xlsx"
        write_template(small)
        add_external_link_fixture(small)
        templates.append(
            Template.of(
                "maly",
                "skoroszyt zapisany przez openpyxl, dwa sformatowane Wiersze danych, drugi arkusz, bez relacji arkusza",
                small,
                max_rows=80,
            )
        )
    return templates


# --- Wiersze danych ---------------------------------------------------------------------------------

EDGE_TEXTS = (
    "",
    " ",
    "=",
    "==",
    "=N18*2",
    '="a<b&c"',
    "#N/A",
    "#DIV/0!",
    "#ARG!",
    " spacja na początku",
    "spacja na końcu ",
    "\tz tabulatorem",
    "linia 1\nlinia 2",
    "linia 1\r\nlinia 2\r",
    "_x000D_",
    '<b>&amp; "cudzysłów" \'apostrof\' ]]></b>',
    "zażółć gęślą jaźń ŻÓŁĆ",
    "😀 poza BMP",
    " twarda spacja ",
    "a" * 32767,
    "b" * 32767 + "ucięte",
)
EDGE_NUMBERS: tuple[int | float | Decimal, ...] = (
    0,
    -0.0,
    1,
    -1,
    2**53 + 1,
    10**16,
    -(10**20),
    0.1,
    1 / 3,
    1e-300,
    5e-324,
    1.7976931348623157e308,
    123456789.12345679,
    float("nan"),
    float("inf"),
    float("-inf"),
    Decimal("2.5"),
    Decimal("0.1000000000000000055511151231257827"),
    Decimal("-Infinity"),
)
EDGE_DATES: tuple[date | datetime | time | timedelta, ...] = (
    date(1900, 1, 1),
    date(1900, 2, 28),
    date(1900, 3, 1),
    date(1899, 12, 30),
    date(2026, 9, 8),
    datetime(2026, 9, 8, 7, 30, 15),
    datetime(2026, 9, 8, 23, 59, 59, 999999),
    datetime(1900, 2, 28, 12),
    time(0, 0),
    time(13, 45),
    time(23, 59, 59, 999999),
    timedelta(0),
    timedelta(hours=36, minutes=5),
    timedelta(days=-1, seconds=1),
)


def cell_values() -> st.SearchStrategy[object]:
    """Every value a Wiersz danych cell can hold: what Plik zbiorczy reading gives, and what openpyxl accepts."""

    printable = st.characters(exclude_categories=("Cs", "Cc"), include_characters="\t\n\r")
    texts = st.one_of(
        st.sampled_from(EDGE_TEXTS), st.text(printable, max_size=40), st.text(alphabet="=#<>&\"' \n\r\t_x0D", max_size=12)
    )
    numbers = st.one_of(
        st.sampled_from(EDGE_NUMBERS),
        st.integers(min_value=-(2**70), max_value=2**70),
        st.floats(),
        st.decimals(allow_nan=False, allow_infinity=True),
    )
    dates = st.one_of(st.sampled_from(EDGE_DATES), st.datetimes(), st.dates(), st.times(), st.timedeltas())
    common = st.one_of(st.none(), texts, numbers, st.booleans(), dates)
    # Values that make a path refuse the whole Rozliczenie pracownika or send it the old openpyxl way. They are
    # rare, so that most Rozliczenia pracowników of a run still go through the copy.
    rare = st.one_of(
        st.text(st.characters(max_codepoint=0x1F, exclude_characters="\t\n\r"), min_size=1, max_size=3),
        st.sampled_from((Decimal("NaN"), Decimal("sNaN"))),
        st.datetimes(timezones=st.just(timezone.utc)),
        st.times(timezones=st.just(timezone.utc)),
        st.sampled_from(
            (
                CellRichText("zwykły ", TextBlock(InlineFont(b=True), "pogrubiony")),
                ArrayFormula("Z1", "=SUM(N18:N19)"),
                b"bajty",
            )
        ),
    )
    return st.integers(min_value=0, max_value=59).flatmap(lambda pick: rare if pick == 0 else common)


def is_rare(value: object) -> bool:
    """A value of the rare kind in ``cell_values()``: the only kind either path may refuse."""

    if isinstance(value, str):
        return any(ord(character) < 0x20 and character not in "\t\n\r" for character in value)
    if isinstance(value, Decimal):
        return value.is_nan()
    if isinstance(value, (datetime, time)):
        return value.tzinfo is not None
    return isinstance(value, (CellRichText, ArrayFormula, bytes))


def colors() -> st.SearchStrategy[Color]:
    return st.one_of(
        st.integers(min_value=0, max_value=0xFFFFFF).map(lambda rgb: Color(f"FF{rgb:06X}")),
        st.builds(
            lambda theme, tint: Color(theme=theme, tint=tint),
            st.integers(min_value=0, max_value=9),
            st.sampled_from((0.0, 0.3999755851924192, -0.249977111117893)),
        ),
        st.integers(min_value=0, max_value=63).map(lambda index: Color(indexed=index)),
    )


def highlights() -> st.SearchStrategy[Highlights]:
    """Wyróżnienia of one cell as reading a Plik zbiorczy gives them: ``Highlights.of`` over its font and fill."""

    fills = st.one_of(
        st.none(),
        st.builds(lambda color: PatternFill(fill_type="solid", fgColor=color), colors()),
        st.builds(
            lambda pattern, foreground, background: PatternFill(fill_type=pattern, fgColor=foreground, bgColor=background),
            st.sampled_from(("gray125", "darkGrid", "lightUp")),
            colors(),
            colors(),
        ),
        st.just(GradientFill(degree=90, stop=("FFFFFF00", "FF0070C0"))),
    )
    fonts = st.builds(lambda color, bold, italic: Font(color=color, b=bold, i=italic), st.none() | colors(), st.booleans(), st.booleans())
    return st.builds(Highlights.of, fonts, fills)


EDGE_ADDRESSES = (
    "https://przyklad.invalid/zlecenie?caseId=SYN-1&tw=1",
    "https://przyklad.invalid/zlecenie?a=1&b=2#sekcja-2",
    "https://przyklad.invalid/\"cudzysłów\"/'apostrof'/<b>&amp;</b>",
    " https://przyklad.invalid/spacje na brzegach ",
    "https://przyklad.invalid/zażółć/gęślą?jaźń=ŻÓŁĆ",
    "https://przyklad.invalid/😀",
    "https://przyklad.invalid/tab\tnowa\nlinia\r",
    "mailto:biuro@przyklad.invalid",
    "file:///C:/Syntetyczny%20folder/plik.xlsx",
    "https://przyklad.invalid/" + "a" * 2000,
)


def addresses() -> st.SearchStrategy[str]:
    """Addresses of Hiperłącza as reading an XML attribute can give them: any character XML allows."""

    characters = st.characters(exclude_categories=("Cs", "Cc", "Cn"), include_characters="\t\n\r")
    return st.one_of(st.sampled_from(EDGE_ADDRESSES), st.text(characters, max_size=30).map("https://przyklad.invalid/".__add__))


def _data_row(
    values: ExcelRow,
    marked: dict[int, Highlights] | None = None,
    links: dict[int, str] | None = None,
    number_formats: tuple[str, ...] | None = None,
) -> DataRow:
    """A Wiersz danych of 46 values (A:AT); cells missing from ``marked`` have no Wyróżnienia, ``links`` maps index → address."""

    cells = marked or {}
    hyperlinks = {index + 1: Hyperlink(address) for index, address in (links or {}).items()}
    highlights = tuple(cells.get(index, PLAIN) for index in range(DATA_COLUMNS))
    return DataRow(FIRST_DATA_ROW, values, highlights, number_formats or _source_formats(values), hyperlinks)


# Number formats a Plik zbiorczy cell can carry: built-in ones, Dad's date format, and custom codes with quotes,
# colours and XML specials that the copy has to escape in ``xl/styles.xml``.
NUMBER_FORMATS = (
    "General",
    "@",
    "0.00",
    "m/d/yy h:mm",
    "yyyy-mm-dd h:mm:ss",
    'dd.mm.yyyy "godz." hh:mm',
    '#,##0.00 "zł";[Red]-#,##0.00 "zł"',
    '"<&>" 0',
)


def number_formats() -> st.SearchStrategy[tuple[str, ...] | None]:
    """Number formats of A:L of one row; ``None`` (most rows) as reading a Plik zbiorczy written by Przygotowanie gives."""

    return st.one_of(st.none(), st.none(), st.lists(st.sampled_from(NUMBER_FORMATS), min_size=LOOK_COLUMNS, max_size=LOOK_COLUMNS).map(tuple))


def _source_formats(values: ExcelRow) -> tuple[str, ...]:
    """Number formats of the look columns A:L as reading a Plik zbiorczy gives them: a date format on dates."""

    return tuple(
        "m/d/yy h:mm" if isinstance(value, (date, datetime, time, timedelta)) else "General"
        for value in values[:LOOK_COLUMNS]
    )


@st.composite
def data_rows(draw: st.DrawFn) -> DataRow:
    """One Wiersz danych as from Plik zbiorczy: 46 values, their Wyróżnienia and Hiperłącza, often none at all.

    A Hiperłącze sits mostly in F (``NUMER ZL.``) and only on a cell with a value, as ``read_data_rows`` gives it.
    """

    columns = st.integers(min_value=0, max_value=DATA_COLUMNS - 1)
    values = draw(st.dictionaries(columns, cell_values(), max_size=8))
    marked = draw(st.one_of(st.just({}), st.dictionaries(columns, highlights(), max_size=10)))
    linked_columns = st.one_of(st.just(ORDER_NUMBER - 1), columns)
    links = draw(st.one_of(st.just({}), st.dictionaries(linked_columns, addresses(), min_size=1, max_size=3)))
    for index in links:
        if values.get(index) in (None, ""):
            values[index] = f"ZL-SYN-{index}"
    formats = draw(number_formats())
    return _data_row(tuple(values.get(index) for index in range(DATA_COLUMNS)), marked, links, formats)  # type: ignore[misc]


def _highlighted(rows: Sequence[DataRow]) -> bool:
    return any(cell != PLAIN for row in rows for cell in row.highlights)


def _linked(rows: Sequence[DataRow]) -> bool:
    return any(row.hyperlinks for row in rows)


@st.composite
def settlements(draw: st.DrawFn, template: Template) -> list[DataRow]:
    """Wiersze danych of one Rozliczenie pracownika: a few distinct rows repeated, with some rows overridden."""

    interesting = template.interesting_counts()
    count = draw(st.one_of(st.sampled_from(interesting), st.integers(min_value=0, max_value=template.max_rows)))
    palette = draw(st.lists(data_rows(), min_size=1, max_size=4))
    overrides = draw(st.dictionaries(st.integers(min_value=0, max_value=max(count - 1, 0)), data_rows(), max_size=3)) if count else {}
    return [overrides.get(index, palette[index % len(palette)]) for index in range(count)]


@dataclass(frozen=True)
class Case:
    """One Przebieg rozliczeń: one Wzorzec loaded once, then each Rozliczenie pracownika in order."""

    template: Template
    workers: tuple[tuple[DataRow, ...], ...]

    def __repr__(self) -> str:
        return f"Case({self.template.name}, {[len(rows) for rows in self.workers]} Wierszy danych)"


def cases(template: Template) -> st.SearchStrategy[Case]:
    return st.lists(settlements(template), min_size=1, max_size=3).map(
        lambda workers: Case(template, tuple(tuple(rows) for rows in workers))
    )


# --- one Rozliczenie pracownika --------------------------------------------------------------------


@dataclass(frozen=True)
class Divergence:
    rule: str
    detail: str


@dataclass(frozen=True)
class Attempt:
    path: Path
    refusal: str | None  # exception type name, or None when the file was written


def _attempt(target: Path, write: Callable[[], None]) -> Attempt:
    try:
        write()
    except Exception as exc:  # noqa: BLE001 - a path may refuse a rare value in any way; a refusal without one is a LUKA
        return Attempt(target, type(exc).__name__)
    return Attempt(target, None)


def _write_with_openpyxl(twin: Path, target: Path, rows: Sequence[DataRow], via_copy: bool) -> None:
    """The reference; it gets the text an XML reader reads (``_xml_line_ends``) only when the copy wrote the file."""

    if via_copy:
        rows = [replace(row, values=tuple(_xml_line_ends(value) for value in row.values)) for row in rows]
    TemplateFile(twin, Layout.WORKER_SETTLEMENT).create(target, rows)
    if has_xml_comment(target):
        raise AssertionError("bliźniak Wzorca poszedł kopią: wzorzec porównania nie jest drogą openpyxl")


def _xml_line_ends(value: Any) -> Any:
    """Text as every XML reader sees it: ``\\r\\n`` and a lone ``\\r`` become ``\\n``.

    openpyxl writes the sheet through et_xmlfile, which opens its temporary file in text mode, so on
    Windows every ``\\n`` is written as ``\\r\\n`` and a raw ``\\r\\n`` in a value comes back as two line
    breaks. The copy writes bytes and is the same on every system. Giving the reference the text an XML
    reader would read anyway keeps the reference independent of the system; the copy still gets the raw
    ``\\r``, so dropping or mangling it would still be a divergence. When the copy itself takes the old path,
    the reference gets the raw text as well: both are openpyxl, and its Windows quirk is not checked.
    """

    if isinstance(value, str):
        # openpyxl cuts text to 32 767 characters first (``Cell.check_string``); normalizing before that would
        # move the cut.
        return value[:MAX_TEXT].replace("\r\n", "\n").replace("\r", "\n")
    return value


class Routes:
    """Which path ``TemplateFile.create`` took: the copy, or the old openpyxl path (``_create_with_openpyxl``)."""

    def __init__(self) -> None:
        self.fallbacks: list[Path] = []
        original = TemplateFile._create_with_openpyxl

        def recording(worker: TemplateFile, target: Path, rows: Sequence[DataRow]) -> None:
            self.fallbacks.append(target)
            original(worker, target, rows)

        setattr(TemplateFile, "_create_with_openpyxl", recording)


@dataclass
class Statistics:
    examples: int = 0
    settlements: Counter[str] = field(default_factory=Counter)  # kopia / droga openpyxl / odmowa
    refusals: Counter[str] = field(default_factory=Counter)
    highlighted: int = 0  # Rozliczenia pracowników with Wyróżnienia written through the copy
    linked: int = 0  # Rozliczenia pracowników with Hiperłącza written through the copy
    unexplained: Counter[str] = field(default_factory=Counter)  # refusals of Wiersze danych without a rare value
    max_rows: int = 0
    chunk_rows_reached: set[int] = field(default_factory=set)
    violated: bool = False  # from the first counterexample on, every part checks every example Hypothesis shrinks


def check_settlement(
    template: Template, copy: Attempt, reference: Attempt, rows: Sequence[DataRow], via_copy: bool
) -> list[Divergence]:
    if copy.refusal is not None or reference.refusal is not None:
        found = []
        if (copy.refusal is None) != (reference.refusal is None):
            found.append(
                Divergence(
                    "TA_SAMA_ODMOWA",
                    f"kopia: {copy.refusal or 'zapisała plik'}, droga openpyxl: {reference.refusal or 'zapisała plik'}",
                )
            )
        leftovers = sorted(path.name for path in copy.path.parent.iterdir()) if copy.refusal is not None else []
        if leftovers:
            found.append(Divergence("TA_SAMA_ODMOWA", f"po odmowie zostały pliki: {leftovers}"))
        return found
    if via_copy:
        try:
            shape = _excel_shape_divergences(copy.path, len(rows))
        except Exception as exc:  # noqa: BLE001 - a file that cannot be read is the divergence itself
            return [Divergence("KSZTALT_EXCELA", f"pliku nie da się przeczytać: {type(exc).__name__}: {exc}")]
        if shape:
            return shape
        return (
            _gate_divergences(copy.path, reference.path)
            + _look_divergences(template, copy.path, rows)
            + _untouched_divergences(template, copy.path, len(rows))
        )
    return _gate_divergences(copy.path, reference.path) + _look_divergences(template, copy.path, rows)


def _look_divergences(template: Template, path: Path, rows: Sequence[DataRow]) -> list[Divergence]:
    """WYGLAD_WIERSZA: A:L of every written row look as the rule says, and the row stores no height."""

    workbook = load_workbook(path)
    sheet = workbook.active
    assert sheet is not None
    thin = [("thin", 64)] * 4
    found: list[str] = []
    for index, row in enumerate(rows):
        number = FIRST_DATA_ROW + index
        dimension = sheet.row_dimensions.get(number)
        if dimension is not None and dimension.ht is not None:
            found.append(f"wiersz {number} ma zapisaną wysokość {dimension.ht}")
        for column in range(1, LOOK_COLUMNS + 1):
            cell = sheet.cell(number, column)
            value = row.values[column - 1]
            wraps = isinstance(value, str) and ("\n" in value or "\r" in value)  # Excel czyta \r w tekście jako koniec linii
            sides = [
                (side.style, side.color.indexed if side.color is not None else None)
                for side in (cell.border.left, cell.border.right, cell.border.top, cell.border.bottom)
            ]
            actual = (
                (cell.font.name, cell.font.sz),
                (cell.alignment.horizontal, cell.alignment.vertical, bool(cell.alignment.wrap_text)),
                sides,
                cell.number_format,
            )
            expected = (template.header_faces[column - 1], ("left", "center", wraps), thin, row.number_formats[column - 1])
            if actual != expected:
                found.append(f"{cell.coordinate} {_value_text(actual)} zamiast {_value_text(expected)}")
    workbook.close()
    if not found:
        return []
    return [Divergence("WYGLAD_WIERSZA", f"{len(found)} komórek albo wierszy: {'; '.join(found[:CELL_DIFFERENCES_SHOWN])}")]


def _gate_divergences(copy: Path, reference: Path) -> list[Divergence]:
    differences = gate.compare_worker_settlements(reference.parent, copy.parent)
    if not any(differences.values()):
        return []
    found: list[Divergence] = []
    cell_counts = Counter({category: differences[category] for category in ("wartosci", "formuly", "style", "hiperlacza")})
    other = {category: count for category, count in differences.items() if count and category not in cell_counts}
    if other:
        found.append(Divergence("WIERNA_KOPIA", f"bramka poprawności: {other}"))
    if sum(cell_counts.values()):
        cells = _cell_differences(copy, reference)
        shown = "; ".join(cells[:CELL_DIFFERENCES_SHOWN]) or f"liczniki bramki {dict(cell_counts)}"
        found.append(Divergence("WIERNA_KOPIA", f"{len(cells)} komórek: {shown}"))
    return found


def _cell_differences(copy: Path, reference: Path) -> list[str]:
    """Cells that differ, with the gate's own semantics."""

    reference_book, copy_book = load_workbook(reference), load_workbook(copy)
    styles = gate._Styles()
    cells: list[str] = []
    for title in dict.fromkeys([*reference_book.sheetnames, *copy_book.sheetnames]):
        expected, actual = gate._sheet(reference_book, title), gate._sheet(copy_book, title)
        expected_columns, actual_columns = gate._column_dimensions(expected), gate._column_dimensions(actual)
        for coordinate in sorted(expected._cells.keys() | actual._cells.keys()):
            row, column = coordinate
            expected_cell, actual_cell = expected._cells.get(coordinate), actual._cells.get(coordinate)
            label = f"{title}!{get_column_letter(column)}{row}"
            expected_formula, actual_formula = gate._formula(expected_cell), gate._formula(actual_cell)
            if expected_formula is not None or actual_formula is not None:
                if expected_formula != actual_formula:
                    cells.append(f"{label} formuła kopia={_value_text(actual_formula)} openpyxl={_value_text(expected_formula)}")
                continue
            expected_value, actual_value = gate._value(expected_cell), gate._value(actual_cell)
            if expected_value != actual_value:
                cells.append(f"{label} wartość kopia={_value_text(actual_value)} openpyxl={_value_text(expected_value)}")
            expected_link, actual_link = gate._hyperlink(expected_cell), gate._hyperlink(actual_cell)
            if expected_link != actual_link:
                cells.append(f"{label} hiperłącze kopia={_value_text(actual_link)} openpyxl={_value_text(expected_link)}")
            expected_style = gate._style(expected, expected_cell, coordinate, expected_columns, styles)
            actual_style = gate._style(actual, actual_cell, coordinate, actual_columns, styles)
            if expected_style != actual_style:
                cells.append(f"{label} styl (czcionka, wypełnienie, obramowanie, wyrównanie albo format liczb)")
    reference_book.close()
    copy_book.close()
    return cells


def _rows(xml: bytes) -> dict[int, tuple[int, int]]:
    return {int(match.group(1)): match.span() for match in _ROW.finditer(xml)}


def _start_tag(row_xml: bytes) -> bytes:
    tag = row_xml[: row_xml.index(b">")]
    return tag.rstrip(b"/ ")


def _cells_right_of_data(row_xml: bytes) -> dict[int, bytes]:
    return {
        column_index_from_string(match.group(1).decode()): match.group(0)
        for match in _CELL.finditer(row_xml)
        if column_index_from_string(match.group(1).decode()) > DATA_COLUMNS
    }


def _without_calc_pr(xml: bytes) -> bytes:
    return _CALC_PR.sub(b"", xml, count=1)


def _without_style_lists(xml: bytes) -> bytes:
    for style_list in STYLE_LISTS:
        xml = style_list.without(xml)
    return xml


def _styles_untouched(before: bytes, after: bytes) -> list[Divergence]:
    found = []
    if _without_style_lists(before) != _without_style_lists(after):
        found.append(Divergence("NIETKNIETY_WZORZEC", f"{STYLES} zmienił się poza {', '.join(item.tag for item in STYLE_LISTS)}"))
    for style_list in STYLE_LISTS:
        original = style_list.entries(before)
        if style_list.entries(after)[: len(original)] != original:
            found.append(Divergence("NIETKNIETY_WZORZEC", f"{STYLES}: {style_list.tag} Wzorca zmienione, a nie tylko dopisane"))
    return found


def _untouched_divergences(template: Template, copy: Path, written: int) -> list[Divergence]:
    found: list[Divergence] = []
    original, result = zipfile.ZipFile(template.path), zipfile.ZipFile(copy)
    names_before, names_after = original.namelist(), result.namelist()
    linked = _HYPERLINKS_BLOCK.search(result.read(SHEET)) is not None
    added_rels = linked and SHEET_RELS not in names_before and names_after == [*names_before, SHEET_RELS]
    if names_before != names_after and not added_rels:
        return [Divergence("NIETKNIETY_WZORZEC", "inny zestaw albo kolejność części zip niż we Wzorcu")]
    for before in original.infolist():
        after = result.getinfo(before.filename)
        name = before.filename
        if name == SHEET:
            continue
        if name == SHEET_RELS and linked:
            found += _rels_untouched(original.read(name), result.read(name))
            continue
        if name == WORKBOOK:
            if _without_calc_pr(original.read(name)) != _without_calc_pr(result.read(name)):
                found.append(Divergence("NIETKNIETY_WZORZEC", f"{name} zmienił się poza calcPr"))
            continue
        if name == STYLES:
            found += _styles_untouched(original.read(name), result.read(name))
            continue
        same = (before.CRC, before.compress_size, before.compress_type) == (after.CRC, after.compress_size, after.compress_type)
        if not same or original.read(name) != result.read(name):
            found.append(Divergence("NIETKNIETY_WZORZEC", f"{name} nie jest przepisana bajt w bajt"))
    found += _sheet_untouched(original.read(SHEET), result.read(SHEET), written)
    return found


def _rels_untouched(before: bytes, after: bytes) -> list[Divergence]:
    """The Wzorzec's sheet relationships with only whole ``<Relationship/>`` entries inserted before the end."""

    end = before.rindex(b"</Relationships>")
    middle = after[end : len(after) - (len(before) - end)]
    if not (after.startswith(before[:end]) and after.endswith(before[end:]) and not _RELATIONSHIP.sub(b"", middle)):
        return [Divergence("NIETKNIETY_WZORZEC", f"{SHEET_RELS}: relacje Wzorca zmienione, a nie tylko dopisane")]
    return []


def _region(xml: bytes, rows: dict[int, tuple[int, int]], last: int) -> tuple[int, int]:
    """Span of the rows from the first Wiersz danych to ``last``; empty span where the first one would start."""

    inside = [span for number, span in rows.items() if FIRST_DATA_ROW <= number <= last]
    if inside:
        return inside[0][0], inside[-1][1]
    later = [span[0] for number, span in rows.items() if number > last]
    position = later[0] if later else xml.index(b"</sheetData>")
    return position, position


def _sheet_untouched(before: bytes, after: bytes, written: int) -> list[Divergence]:
    last = FIRST_DATA_ROW + written - 1
    rows_before, rows_after = _rows(before), _rows(after)
    start_before, end_before = _region(before, rows_before, last)
    start_after, end_after = _region(after, rows_after, last)
    found: list[Divergence] = []
    if _DIMENSION.sub(b"", before[:start_before]) != _DIMENSION.sub(b"", after[:start_after]):
        found.append(Divergence("NIETKNIETY_WZORZEC", "arkusz zmienił się przed pierwszym Wierszem danych"))
    if before[end_before:] != _HYPERLINKS_BLOCK.sub(b"", after[end_after:], count=1):
        found.append(Divergence("NIETKNIETY_WZORZEC", f"arkusz zmienił się za ostatnim wpisanym wierszem {last} poza <hyperlinks>"))
    for number in range(FIRST_DATA_ROW, last + 1):
        if number not in rows_after:
            found.append(Divergence("NIETKNIETY_WZORZEC", f"brak wiersza {number} w arkuszu"))
            break
        if number not in rows_before:
            continue
        row_before = before[slice(*rows_before[number])]
        row_after = after[slice(*rows_after[number])]
        if _ROW_HEIGHT.sub(b"", _start_tag(row_before)) != _start_tag(row_after):
            found.append(Divergence("NIETKNIETY_WZORZEC", f"atrybuty wiersza {number} zmienione"))
            break
        if _cells_right_of_data(row_before) != _cells_right_of_data(row_after):
            found.append(Divergence("NIETKNIETY_WZORZEC", f"komórki za AT w wierszu {number} nie są bajt w bajt"))
            break
    return found


def _excel_shape_divergences(copy: Path, written: int) -> list[Divergence]:
    found: list[Divergence] = []
    archive = zipfile.ZipFile(copy)
    broken = archive.testzip()
    if broken is not None:
        return [Divergence("KSZTALT_EXCELA", f"zła suma CRC części {broken}")]
    for name in archive.namelist():
        if name.endswith((".xml", ".rels")):
            ElementTree.fromstring(archive.read(name))
    names = archive.namelist()
    if len(names) != len(set(names)):
        found.append(Divergence("KSZTALT_EXCELA", "powtórzona część zip"))
    xml = archive.read(SHEET)
    root = ElementTree.fromstring(xml)
    previous_row = 0
    max_row, max_column = 0, 0
    for row in root.iter(f"{_MAIN}row"):
        number = int(row.get("r", "0"))
        if number <= previous_row:
            found.append(Divergence("KSZTALT_EXCELA", f"wiersz {number} po wierszu {previous_row}"))
        previous_row = number
        previous_column = 0
        for cell in row.iter(f"{_MAIN}c"):
            match = re.fullmatch(r"([A-Z]{1,3})(\d+)", cell.get("r", ""))
            if match is None or int(match.group(2)) != number:
                found.append(Divergence("KSZTALT_EXCELA", f"komórka {cell.get('r')} w wierszu {number}"))
                continue
            column = column_index_from_string(match.group(1))
            if column <= previous_column:
                found.append(Divergence("KSZTALT_EXCELA", f"komórka {cell.get('r')} nie rośnie w wierszu {number}"))
            previous_column = column
            if FIRST_DATA_ROW <= number < FIRST_DATA_ROW + written:
                max_row, max_column = max(max_row, number), max(max_column, column)
        if len(found) > CELL_DIFFERENCES_SHOWN:
            break
    for text in root.iter(f"{_MAIN}t"):
        value = text.text or ""
        if value != value.strip() and text.get(_XML_SPACE) != "preserve":
            found.append(Divergence("KSZTALT_EXCELA", f"tekst ze spacją na brzegu bez xml:space=\"preserve\": {value[:20]!r}"))
            break
    dimension = _DIMENSION.search(xml)
    if dimension is not None and written:
        reference = dimension.group(1).decode()
        last_cell = reference.split(":")[-1]
        bound = re.fullmatch(r"([A-Z]+)(\d+)", last_cell)
        if bound is None or int(bound.group(2)) < FIRST_DATA_ROW + written - 1 or column_index_from_string(bound.group(1)) < max_column:
            found.append(Divergence("KSZTALT_EXCELA", f"dimension {reference} nie obejmuje wpisanych wierszy do {FIRST_DATA_ROW + written - 1}"))
    return found + _styles_shape(archive.read(STYLES), xml) + _package_shape(archive, root)


def _package_shape(archive: zipfile.ZipFile, root: ElementTree.Element) -> list[Divergence]:
    """Sheet elements in schema order, a content type for every part, hyperlinks pointing at external hyperlink relationships."""

    found = []
    names = [child.tag.removeprefix(_MAIN) for child in root]
    if names.count("hyperlinks") > 1:
        found.append(Divergence("KSZTALT_EXCELA", "więcej niż jeden <hyperlinks>"))
    positions = [WORKSHEET_ORDER.index(name) for name in names if name in WORKSHEET_ORDER]
    if positions != sorted(positions):
        found.append(Divergence("KSZTALT_EXCELA", f"elementy arkusza poza kolejnością schematu: {names[names.index('sheetData') + 1:]}"))
    types = ElementTree.fromstring(archive.read(CONTENT_TYPES))
    extensions = {default.get("Extension", "").lower() for default in types.iter(f"{_TYPES}Default")}
    overrides = {override.get("PartName", "") for override in types.iter(f"{_TYPES}Override")}
    untyped = [
        name
        for name in archive.namelist()
        if name != CONTENT_TYPES and f"/{name}" not in overrides and name.rsplit(".", 1)[-1].lower() not in extensions
    ]
    if untyped:
        found.append(Divergence("KSZTALT_EXCELA", f"części bez typu zawartości: {untyped}"))
    relationships = (
        list(ElementTree.fromstring(archive.read(SHEET_RELS)).iter(f"{_PACKAGE}Relationship"))
        if SHEET_RELS in archive.namelist()
        else []
    )
    ids = [relationship.get("Id") for relationship in relationships]
    if len(ids) != len(set(ids)):
        found.append(Divergence("KSZTALT_EXCELA", f"powtórzony Id w {SHEET_RELS}: {ids}"))
    by_id = {relationship.get("Id"): relationship for relationship in relationships}
    references = []
    for link in root.iter(f"{_MAIN}hyperlink"):
        references.append(link.get("ref"))
        relationship = by_id.get(link.get(f"{{{_OFFICE_REL}}}id"))
        if (
            relationship is None
            or not relationship.get("Type", "").endswith("/hyperlink")
            or relationship.get("TargetMode") != "External"
        ):
            found.append(Divergence("KSZTALT_EXCELA", f"hiperłącze {link.get('ref')} nie wskazuje zewnętrznej relacji hiperłącza"))
            break
    if len(references) != len(set(references)):
        found.append(Divergence("KSZTALT_EXCELA", "powtórzona komórka w <hyperlinks>"))
    return found


def _styles_shape(styles: bytes, sheet: bytes) -> list[Divergence]:
    """Every list's ``count`` matches its entries and every reference by position points at an existing entry."""

    found = []
    for style_list in STYLE_LISTS:
        count, entries = style_list.count(styles), len(style_list.entries(styles))
        if count is not None and count != entries:
            found.append(Divergence("KSZTALT_EXCELA", f"{style_list.tag} count={count}, a pozycji jest {entries}"))
    formats = [int(_XF_ID["numFmtId"].search(entry).group(1)) for entry in NUM_FMTS.entries(styles)]  # type: ignore[union-attr]
    if len(formats) != len(set(formats)):
        found.append(Divergence("KSZTALT_EXCELA", "powtórzony numFmtId"))
    xfs = CELL_XFS.entries(styles)
    fonts, fills, borders = len(FONTS.entries(styles)), len(FILLS.entries(styles)), len(BORDERS.entries(styles))
    for position, xf in enumerate(xfs):
        tag = xf[: xf.index(b">")]
        ids = {name: int(match.group(1)) for name, pattern in _XF_ID.items() if (match := pattern.search(tag))}
        if (
            ids.get("fontId", 0) >= fonts
            or ids.get("fillId", 0) >= fills
            or ids.get("borderId", 0) >= borders
            or ids.get("numFmtId", 0) >= _FIRST_CUSTOM_FORMAT and ids["numFmtId"] not in formats
        ):
            found.append(Divergence("KSZTALT_EXCELA", f"styl komórki {position} wskazuje brakującą pozycję: {ids}"))
            break
    used = [int(style) for style in re.findall(rb'<c r="[A-Z]+\d+" s="(\d+)"', sheet)]
    if used and max(used) >= len(xfs):
        found.append(Divergence("KSZTALT_EXCELA", f"komórka ze stylem {max(used)} spoza {len(xfs)} stylów"))
    return found


# --- one Przebieg rozliczeń ------------------------------------------------------------------------


class Violation(AssertionError):
    def __init__(self, case: Case, divergences: list[tuple[int, Divergence]]) -> None:
        super().__init__(f"{case!r}: {[divergence.rule for _, divergence in divergences]}")
        self.case = case
        self.divergences = divergences


@dataclass
class Run:
    """Checks examples; part ``part`` of ``parts`` compares only every ``parts``-th example of each Wzorzec.

    Every part writes every Rozliczenie pracownika through the copy, so the coverage it reports is that of all
    examples; the rules (the openpyxl reference and the comparison, most of the time) run in one part per example.
    Hypothesis generates the same examples in every part, as no check fails until a counterexample.
    """

    routes: Routes
    statistics: dict[str, Statistics]
    part: int = 0
    parts: int = 1

    def check(self, case: Case) -> None:
        template = case.template
        statistics = self.statistics[template.name]
        compared = statistics.violated or statistics.examples % self.parts == self.part
        statistics.examples += 1
        divergences: list[tuple[int, Divergence]] = []
        with tempfile.TemporaryDirectory(prefix="kopia-wzorca-") as scratch:
            twin = refused_by_copy(Path(shutil.copy(template.path, Path(scratch) / "bliźniak Wzorca.xlsx"))) if compared else None
            worker = TemplateFile(template.path, Layout.WORKER_SETTLEMENT)
            worker.check()
            for index, rows in enumerate(case.workers, start=1):
                copy_folder, reference_folder = Path(scratch) / f"kopia-{index}", Path(scratch) / f"openpyxl-{index}"
                copy_folder.mkdir()
                reference_folder.mkdir()
                name = f"Rozliczenie 08_14_09_2026 - Wykonawca {index}.xlsx"
                fallbacks = len(self.routes.fallbacks)
                copy = _attempt(copy_folder / name, lambda: worker.create(copy_folder / name, rows))
                via_copy = len(self.routes.fallbacks) == fallbacks
                reference = (
                    _attempt(reference_folder / name, lambda: _write_with_openpyxl(twin, reference_folder / name, rows, via_copy))
                    if twin is not None
                    else None
                )
                route = "odmowa" if copy.refusal else "kopia" if via_copy else "droga openpyxl"
                statistics.settlements[route] += 1
                if copy.refusal:
                    statistics.refusals[copy.refusal] += 1
                refusal = copy.refusal or (reference.refusal if reference is not None else None)
                if refusal and not any(is_rare(value) for row in rows for value in row.values):
                    statistics.unexplained[refusal] += 1
                statistics.max_rows = max(statistics.max_rows, len(rows))
                if via_copy and copy.refusal is None:
                    statistics.highlighted += _highlighted(rows)
                    statistics.linked += _linked(rows)
                    last = FIRST_DATA_ROW + len(rows) - 1
                    statistics.chunk_rows_reached.update(row for row in template.chunk_rows[1:] if row <= last + 1)
                if reference is not None:
                    divergences += [(index, divergence) for divergence in check_settlement(template, copy, reference, rows, via_copy)]
        if template.path.read_bytes() != template.content:
            divergences.append((0, Divergence("NIETKNIETY_WZORZEC", "plik Wzorca na dysku się zmienił")))
        if divergences:
            statistics.violated = True
            raise Violation(case, divergences)


# --- report -----------------------------------------------------------------------------------------


def _first(items: Sequence[str]) -> str:
    return ", ".join(items[:8]) + (f" i {len(items) - 8} innych" if len(items) > 8 else "")


def _highlights_text(highlights: Highlights) -> str:
    fill = highlights.fill
    parts = [
        f"wypełnienie {getattr(fill, 'patternType', None) or type(fill).__name__}" if fill is not None else "",
        f"kolor {_color_text(highlights.color)}" if highlights.color is not None else "",
        "pogrubienie" if highlights.bold else "",
        "kursywa" if highlights.italic else "",
    ]
    return "(" + ", ".join(part for part in parts if part) + ")"


def _color_text(color: Color) -> str:
    if color.type == "theme":
        return f"motyw {color.theme} tint {color.tint}"
    return f"{color.type} {color.value}"


def _value_text(value: object) -> str:
    text = repr(value)
    return text if len(text) <= 60 else f"{text[:40]}… ({len(text)} znaków)"


def story(case: Case, divergences: Sequence[tuple[int, Divergence]]) -> list[str]:
    template = case.template
    lines = [
        f"  Wzorzec {template.name} ({template.description}); kawałki kompresji kopii od wierszy "
        f"{', '.join(map(str, template.chunk_rows))}.",
        "  1. Przebieg rozliczeń wczytuje Wzorzec raz (TemplateFile.check).",
    ]
    for index, rows in enumerate(case.workers, start=1):
        if not rows:
            lines.append(f"  {index + 1}. Rozliczenie pracownika {index}: bez Wierszy danych.")
            continue
        last = FIRST_DATA_ROW + len(rows) - 1
        values = [
            f"{get_column_letter(column)}{FIRST_DATA_ROW + offset}={_value_text(value)}"
            for offset, row in enumerate(rows)
            for column, value in enumerate(row.values, start=1)
            if value is not None
        ]
        marked = [
            f"{get_column_letter(column)}{FIRST_DATA_ROW + offset} {_highlights_text(cell)}"
            for offset, row in enumerate(rows)
            for column, cell in enumerate(row.highlights, start=1)
            if cell != PLAIN
        ]
        linked = [
            f"{get_column_letter(column)}{FIRST_DATA_ROW + offset}→{_value_text(link.address)}"
            for offset, row in enumerate(rows)
            for column, link in sorted(row.hyperlinks.items())
        ]
        lines.append(
            f"  {index + 1}. Rozliczenie pracownika {index}: {len(rows)} Wierszy danych (wiersze {FIRST_DATA_ROW}–{last}); "
            f"niepuste: {_first(values) or 'brak'}; Wyróżnienia: {_first(marked) or 'brak'}; Hiperłącza: {_first(linked) or 'brak'}."
        )
    for index, divergence in divergences:
        where = f"Rozliczenie pracownika {index}" if index else "Wzorzec"
        lines.append(f"  → {divergence.rule} ({where}): {divergence.detail}")
    return lines


@dataclass(frozen=True)
class Outcome:
    exit_code: int
    lines: list[str]


def run_check(
    examples: int = DEFAULT_EXAMPLES,
    random_seed: int | None = None,
    shrink: bool = True,
    real_template: Path | None = None,
    template_names: Sequence[str] = SYNTHETIC_TEMPLATES,
    part: tuple[int, int] = (0, 1),
) -> Outcome:
    if not 0 <= part[0] < part[1]:
        raise ValueError(f"część {part[0]} spoza 0..{part[1] - 1}")
    started = clock.perf_counter()
    lines: list[str] = []
    if hasattr(hypothesis_engine, "MAX_SHRINKING_SECONDS"):
        hypothesis_engine.MAX_SHRINKING_SECONDS = SHRINK_SECONDS
    with tempfile.TemporaryDirectory(prefix="kopia-wzorca-wzorce-") as scratch:
        templates = build_templates(Path(scratch), real_template, template_names)
        run = Run(Routes(), {template.name: Statistics() for template in templates}, *part)
        phases = [Phase.explicit, Phase.generate] + ([Phase.shrink] if shrink else [])
        failures: list[Violation] = []
        errors: list[str] = []
        for template in templates:
            # Each Wzorzec gets its own Hypothesis run with the full budget of examples.
            explicit = [Case(template, (_value_rows(),)), _longest_example(template), *_carriage_return_examples(template)]
            missing_cells = _missing_cells_example(template)
            if missing_cells is not None:
                explicit.append(missing_cells)
            test = _property(run, template, explicit, examples, random_seed, phases)
            try:
                # Hypothesis draws about one value in forty from the literals of the non-test modules loaded in the
                # process and keeps the ones it may draw in a cache shared by every Hypothesis run in the process.
                # Without both the test alone and the full pytest check the same examples (#288).
                with (
                    mock.patch.object(hypothesis_providers, "_get_local_constants", Constants),
                    mock.patch.object(hypothesis_providers, "CONSTANTS_CACHE", LRUCache(1024)),
                ):
                    test()
            except Violation as violation:
                failures.append(violation)
            except BaseExceptionGroup as group:
                leaves = list(_leaves(group))
                failures += [error for error in leaves if isinstance(error, Violation)]
                errors += [_checker_error(error) for error in leaves if not isinstance(error, Violation)]
            except Exception as error:  # noqa: BLE001 - a bug in the checker itself ends as LUKA SPRAWDZENIA, not a traceback
                errors.append(_checker_error(error))
        for template in templates:
            statistics = run.statistics[template.name]
            routes = ", ".join(f"{route} {count}" for route, count in sorted(statistics.settlements.items()))
            refusals = f" (odmowy: {dict(statistics.refusals)})" if statistics.refusals else ""
            reached = len(statistics.chunk_rows_reached)
            lines.append(
                f"Wzorzec {template.name}: {statistics.examples} przykładów, Rozliczenia pracowników: {routes}{refusals}; "
                f"z Wyróżnieniami przez kopię {statistics.highlighted}; z Hiperłączami przez kopię {statistics.linked}; do {statistics.max_rows} Wierszy danych; granice kawałków przekroczone: {reached}/{len(template.chunk_rows) - 1}"
            )
    gaps = []
    for template in templates:
        statistics = run.statistics[template.name]
        if failures or not statistics.examples:
            continue
        if statistics.settlements["kopia"] < statistics.settlements["droga openpyxl"]:
            gaps.append(f"Wzorzec {template.name}: kopia użyta rzadziej niż droga openpyxl, sprawdzenie byłoby puste")
        if statistics.settlements["kopia"] < statistics.settlements["odmowa"]:
            gaps.append(f"Wzorzec {template.name}: odmów więcej niż kopii, sprawdzenie byłoby puste")
        if statistics.unexplained:
            gaps.append(
                f"Wzorzec {template.name}: odmowy Wierszy danych bez rzadkiej wartości {dict(statistics.unexplained)}; "
                "to zmiana interfejsu albo nowa odmowa, nie wartość, której obie drogi nie zapiszą"
            )
        if len(statistics.chunk_rows_reached) < len(template.chunk_rows) - 1:
            gaps.append(f"Wzorzec {template.name}: nie każda granica kawałka kompresji została przekroczona")
        if not statistics.highlighted:
            gaps.append(f"Wzorzec {template.name}: żadne Rozliczenie pracownika z Wyróżnieniami nie przeszło przez kopię")
        if not statistics.linked:
            gaps.append(f"Wzorzec {template.name}: żadne Rozliczenie pracownika z Hiperłączami nie przeszło przez kopię")
    for failure in failures:
        lines.append("NOWE NARUSZENIE:")
        lines += story(failure.case, failure.divergences)
    gaps += errors
    lines += [f"LUKA SPRAWDZENIA: {gap}" for gap in gaps]
    share = f", porównana część {part[0] + 1} z {part[1]}" if part[1] > 1 else ""
    lines.append(f"Czas: {clock.perf_counter() - started:.0f} s, {examples} przykładów na Wzorzec{share}.")
    if errors:
        return Outcome(2, [*lines, f"KOPIA WZORCA: LUKA SPRAWDZENIA ({len(gaps)})"])
    if failures:
        return Outcome(1, [*lines, f"KOPIA WZORCA: NOWE NARUSZENIE ({len(failures)})"])
    if gaps:
        return Outcome(2, [*lines, f"KOPIA WZORCA: LUKA SPRAWDZENIA ({len(gaps)})"])
    return Outcome(0, [*lines, "KOPIA WZORCA OK"])


def _property(
    run: Run, template: Template, explicit: Sequence[Case], examples: int, random_seed: int | None, phases: list[Phase]
) -> Callable[[], None]:
    @settings(
        max_examples=examples,
        deadline=None,
        database=None,
        derandomize=random_seed is None,
        phases=phases,
        print_blob=False,
        suppress_health_check=list(HealthCheck),
    )
    @given(cases(template))
    def property_holds(case: Case) -> None:
        run.check(case)

    test = property_holds
    for case in explicit:
        test = example(case)(test)
    return seed(random_seed)(test) if random_seed is not None else test


def _checker_error(error: BaseException) -> str:
    notes = " | ".join(getattr(error, "__notes__", []))
    return f"błąd samego sprawdzenia {type(error).__name__}: {error}" + (f" ({notes})" if notes else "")


def _leaves(group: BaseExceptionGroup[BaseException] | None) -> Iterator[BaseException]:
    for error in group.exceptions if group is not None else ():
        if isinstance(error, BaseExceptionGroup):
            yield from _leaves(error)
        else:
            yield error


def _longest_example(template: Template) -> Case:
    """One Rozliczenie pracownika with the most Wiersze danych in scope, each with a Hiperłącze in F: it crosses every reachable chunk start."""

    values = ("syntetyczne miasto",) + (None,) * 4 + ("ZL-SYN-1",) + (None,) * 7 + (1,) + (None,) * (DATA_COLUMNS - 14)
    row = _data_row(values, links={ORDER_NUMBER - 1: EDGE_ADDRESSES[0]})
    return Case(template, ((row,) * template.max_rows,))


def _carriage_return_examples(template: Template) -> list[Case]:
    """Text with a lone ``\\r`` in a look column after an empty Rozliczenie pracownika: Excel reads it as a line break.

    Counterexamples from CI: A18='\\r' on ``maly`` and B18='x\\rxD&' on ``syntetyczny``. The third, from CI on
    Windows: AC='linia 1\\r\\nlinia 2\\r' beside b'bajty', which sends the file down the old path with the raw text.
    """

    old_path = (None,) * 26 + (b"bajty", None, "linia 1\r\nlinia 2\r") + (None,) * (DATA_COLUMNS - 29)
    return [
        Case(template, ((), (_data_row(("\r",) + (None,) * (DATA_COLUMNS - 1)),))),
        Case(template, ((), (_data_row((None, "x\rxD&") + (None,) * (DATA_COLUMNS - 2)),))),
        Case(template, ((_data_row(old_path),),)),
    ]


def _missing_cells_example(template: Template) -> Case | None:
    """Empty A and B up to the first row of the Wzorzec without a cell in A, which the copy has to add to the XML."""

    last = FIRST_DATA_ROW + template.max_rows - 1
    gaps = [row for row in range(FIRST_DATA_ROW, min(template.last_row, last) + 1) if (row, 1) not in template.cells]
    if not gaps:
        return None
    rows = [_data_row((None, None, "x") + (None,) * (DATA_COLUMNS - 3))] * (gaps[0] - FIRST_DATA_ROW + 1)
    return Case(template, (tuple(rows),))


def _value_rows() -> tuple[DataRow, ...]:
    """The value-type and Wyróżnienia rows of the owner tests, plus Windows line ends and Hiperłącza, as a fixed first example."""

    values: list[Any] = [None] * DATA_COLUMNS
    for letter, value in {
        "A": '<b>Miasto & "Wieś"</b> \'1\' ]]>',
        "B": " spacje na brzegach ",
        "C": "=N18*2",
        "D": "=",
        "E": "#N/A",
        "G": date(2026, 9, 8),
        "I": 12345678901234567,
        "J": True,
        "K": datetime(2026, 9, 8, 7, 30, 15),
        "O": 1.25,
        "P": Decimal("2.5"),
        "Q": time(13, 45),
        "AT": "zażółć gęślą jaźń\nnowa linia",
        "R": "koniec linii Windows\r\ni sam\r",
    }.items():
        values[column_index_from_string(letter) - 1] = value
    yellow = Highlights(PatternFill(fill_type="solid", fgColor="FFFFFF00"), Color("FFFF0000"), bold=True, italic=True)
    theme_fill = Highlights(PatternFill(fill_type="solid", fgColor=Color(theme=5, tint=0.4)), Color(theme=1))
    marked = {column_index_from_string(letter) - 1: cell for letter, cell in {"A": yellow, "G": theme_fill, "N": yellow, "AT": yellow}.items()}
    links = {0: EDGE_ADDRESSES[2], 8: EDGE_ADDRESSES[1]}
    return (
        _data_row(tuple(values), marked, links),
        _data_row(("tylko miasto",) + (None,) * (DATA_COLUMNS - 1), links={0: EDGE_ADDRESSES[3]}),
    )


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--examples", type=int, default=DEFAULT_EXAMPLES, help="Przykłady Hypothesis (domyślnie %(default)s).")
    parser.add_argument("--seed", type=int, help="Losowe przykłady z tego ziarna zamiast stałych (głębsze sprawdzenie).")
    parser.add_argument("--no-shrink", action="store_true", help="Bez zmniejszania kontrprzykładu (szybciej, dla mutacji).")
    parser.add_argument("--wzorzec", type=Path, help="Prawdziwy Wzorzec rozliczenia pracownika zamiast syntetycznych (tylko odczyt).")
    args = parser.parse_args(argv)
    # openpyxl warns about generated numbers too large for a date in a date-formatted column; both paths
    # write them the same, and the gate compares what openpyxl reads.
    warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")
    module = Path(workbook_files.__file__).resolve()
    if not module.is_relative_to(SOURCE_ROOT.resolve()):
        print(f"KOPIA WZORCA: LUKA SPRAWDZENIA (1): rozliczenia importuje się z {module}, nie z {SOURCE_ROOT}")
        return 2
    outcome = run_check(args.examples, args.seed, not args.no_shrink, args.wzorzec)
    print("\n".join(outcome.lines))
    return outcome.exit_code


if __name__ == "__main__":
    raise SystemExit(main())
