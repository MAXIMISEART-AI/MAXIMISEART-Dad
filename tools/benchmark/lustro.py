"""Lustro folderów taty dla scenariusza benchmarku, bez danych klientów.

Lustro to ``Załączniki/Rozliczenia 2026/10_Październik_2026/INEA/06_12_10_2026`` z
Plikiem zbiorczym gotowym do Przebiegu rozliczeń, zbudowane atomowo pod nazwą
``<scenariusz>-<odcisk wejść>``: inne Wzorce Przygotowania pliku zbiorczego,
Konfiguracja wykonawców albo scenariusz dają inne lustro.
"""

from __future__ import annotations

import hashlib
import random
import shutil
import subprocess
import sys
import tempfile
import time
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass
from datetime import UTC, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from benchmark import ROOT, SCRIPT, BenchmarkError
import potomne
from rozliczenia import settlement_messages
from rozliczenia.period_folder import PeriodFolder, collective_file_name
from rozliczenia.settlement_period import parse_period_token
from rozliczenia.workbook_files import DATA_ROW_COLUMNS, QUANTITY_COLUMNS, read_data_rows
from rozliczenia.workbook_save import save_workbook
from rozliczenia.worker_config import WorkerConfigError, load_worker_mapping
from syntetyczne_wzorce import COLLECTIVE_TEMPLATE, INTERNAL_TEMPLATE, WORKER_CONFIG, WORKER_TEMPLATE


@dataclass(frozen=True)
class Scenario:
    """Tydzień lustra: Zlecenia skuteczne, nieskuteczne i Wykonawcy z Konfiguracji bez wierszy."""

    successful_orders: int
    unsuccessful_orders: int
    workers_without_rows: int


SCENARIOS = {
    "typowy-tydzien": Scenario(successful_orders=500, unsuccessful_orders=50, workers_without_rows=2),
    # 2× najcięższego tygodnia w Rejestrze zleceń (592 zlecenia).
    "ciezki-tydzien": Scenario(successful_orders=1200, unsuccessful_orders=120, workers_without_rows=0),
    # Test dymny benchmarku: kilka zleceń, sekunda na uruchomienie.
    "maly-tydzien": Scenario(successful_orders=6, unsuccessful_orders=1, workers_without_rows=1),
}


TEMPLATE_NAMES = (WORKER_TEMPLATE, INTERNAL_TEMPLATE, COLLECTIVE_TEMPLATE)


PERIOD = parse_period_token("06_12_10_2026")


MONTH_FOLDER = "10_Październik_2026"


REGISTER_NAME = "ZLECENIA INTERMAX OD 08_11_2021.xlsx"
# Kody wyjścia Skryptu dla maszyn, po których Plik zbiorczy lustra powstał: OK i wynik wymagający sprawdzenia, którym
# Przygotowanie kończy zapis tylko z powodu ostrzeżeń ilości (preparation_process._finished).
_PREPARED_EXITS = (0, 2)


# Zmiana sposobu generowania lustra zmienia jego odcisk, więc stare lustro nie zostanie użyte.
GENERATOR_VERSION = 2


SEED = 141


REGISTER_HEADERS = (
    "DATA GODZ. POCZ.",
    "MIASTO",
    "adres",
    "UMOWA",
    "URZĄDZENIA",
    "CZYNNOŚCI",
    "NUMER ZL.",
    "WYKONAWCA",
    "NUMER UMOWY",
    "SKUTECZNE",
    "CZYNNOŚCI ZREALIZOWANE",
)


CITIES = ("Miasto Syntetyczne", "Wieś Syntetyczna", "Osada Syntetyczna")


DEVICES = ("dekoder", "router", "terminal ONT")


ACTIVITIES = ("instalacja", "serwis", "demontaż")


# Wpisy taty w Wierszach danych Pliku zbiorczego: ilości N:AP, obciążenie AQ:AR, uwagi AS:AT (za ilościami
# do końca Wiersza danych).
CHARGE_DESCRIPTION, CHARGE_AMOUNT, REMARKS_AMOUNT, REMARKS = range(QUANTITY_COLUMNS[-1] + 1, DATA_ROW_COLUMNS[-1] + 1)


def _main_checkout_config() -> Path | None:
    """``config/`` głównego checkoutu: worktree dzieli z nim katalog gita, ale nie Wzorce spoza gita."""

    try:
        common = subprocess.run(
            ["git", "rev-parse", "--git-common-dir"], cwd=ROOT, capture_output=True, text=True, check=True, timeout=10
        ).stdout.strip()
    except (OSError, subprocess.SubprocessError):
        return None
    return (ROOT / common).resolve().parent / "config"


def _config_with(name: str) -> Path:
    """``config/`` tego checkoutu, gdy ma plik ``name``, inaczej ``config/`` głównego checkoutu (jeśli go ma)."""

    local = ROOT / "config"
    if (local / name).is_file():
        return local
    main_config = _main_checkout_config()
    return main_config if main_config is not None and (main_config / name).is_file() else local


def default_templates_directory() -> Path:
    """``config/`` tego checkoutu, gdy ma Wzorce, inaczej ``config/`` głównego checkoutu."""

    return _config_with(WORKER_TEMPLATE)


def default_worker_config() -> Path:
    """Prawdziwa Konfiguracja wykonawców spoza gita: z ``config/`` tego checkoutu albo głównego checkoutu."""

    return _config_with(WORKER_CONFIG) / WORKER_CONFIG


def template_paths(directory: Path) -> dict[str, Path]:
    paths = {name: directory / name for name in TEMPLATE_NAMES}
    missing = [name for name, path in paths.items() if not path.is_file()]
    if missing:
        raise BenchmarkError(f"Brak Wzorców w {directory}: {', '.join(missing)}. Wskaż folder przez --templates-dir.")
    return paths


def _worker_ids(config: Path) -> list[str]:
    if not config.is_file():
        raise BenchmarkError(
            f"Brak Konfiguracji wykonawców {config}. Wskaż plik przez --config "
            "(np. syntetyczny z tools/syntetyczne_wzorce.py)."
        )
    try:
        return list(load_worker_mapping(config))
    except WorkerConfigError as error:
        reason = settlement_messages.worker_config_problem(config, error).what
        raise BenchmarkError(f"{reason} Plik: {config}") from None


def _fingerprint(scenario_name: str, config: Path, templates: Mapping[str, Path]) -> str:
    """Odcisk wejść lustra: inne Wzorce, Konfiguracja albo scenariusz dają inne lustro."""

    digest = hashlib.sha256(f"{GENERATOR_VERSION}|{SEED}|{scenario_name}|{SCENARIOS[scenario_name]}".encode())
    for path in (config, templates[INTERNAL_TEMPLATE], templates[COLLECTIVE_TEMPLATE]):
        digest.update(path.read_bytes())
    return digest.hexdigest()[:12]


def period_directory(mirror: Path) -> Path:
    return mirror / "Załączniki" / f"Rozliczenia {PERIOD.monday.year}" / MONTH_FOLDER / "INEA" / PERIOD.token


def _register_rows(scenario: Scenario, worker_ids: Sequence[str]) -> list[dict[str, object]]:
    """Deterministyczne, fikcyjne zlecenia okresu; jest wśród nich zlecenie z poniedziałku."""

    active = list(worker_ids[: len(worker_ids) - scenario.workers_without_rows])
    if not active:
        raise BenchmarkError("Scenariusz nie zostawia żadnego Wykonawcy z wierszami.")
    rng = random.Random(SEED)
    total = scenario.successful_orders + scenario.unsuccessful_orders
    successful = [True] * scenario.successful_orders + [False] * scenario.unsuccessful_orders
    rng.shuffle(successful)
    workers = active + [rng.choice(active) for _ in range(scenario.successful_orders - len(active))]
    rng.shuffle(workers)
    successful_workers = iter(workers)
    rows: list[dict[str, object]] = []
    for number in range(1, total + 1):
        started = PERIOD.start + timedelta(days=number % 7, hours=rng.randint(7, 19), minutes=rng.randrange(60))
        activity = rng.choice(ACTIVITIES)
        is_successful = successful[number - 1]
        rows.append(
            {
                "DATA GODZ. POCZ.": started,
                "MIASTO": rng.choice(CITIES),
                "adres": f"ul. Syntetyczna {number:05d}",
                "UMOWA": f"BENCH-UM-{number:05d}",
                "URZĄDZENIA": rng.choice(DEVICES),
                "CZYNNOŚCI": activity,
                "NUMER ZL.": f"BENCH-ZL-{number:05d}",
                "WYKONAWCA": next(successful_workers) if is_successful else rng.choice(active),
                "NUMER UMOWY": f"BENCH-OR-{number:05d}",
                "SKUTECZNE": "tak" if is_successful else "nie",
                "CZYNNOŚCI ZREALIZOWANE": activity if is_successful else "",
            }
        )
    rows.sort(key=lambda row: str(row["DATA GODZ. POCZ."]))
    return rows


def _write_register(path: Path, rows: Iterable[Mapping[str, object]]) -> None:
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("ZLECENIA")
    sheet.append(list(REGISTER_HEADERS))
    for row in rows:
        sheet.append([row[header] for header in REGISTER_HEADERS])
    workbook.save(path)


def _entries(rng: random.Random, number: int) -> dict[int, int | float | str]:
    """Wpisy taty w jednym Wierszu danych: 1–3 ilości w ``N:AP``, czasem obciążenie i uwagi w ``AQ:AT``."""

    entries: dict[int, int | float | str] = {
        column: rng.randint(1, 4) for column in rng.sample(QUANTITY_COLUMNS, rng.randint(1, 3))
    }
    if rng.random() < 0.1:
        entries[CHARGE_DESCRIPTION] = f"BENCH-OBC-{number:05d}"
        entries[CHARGE_AMOUNT] = round(rng.uniform(10, 200), 2)
    if rng.random() < 0.05:
        entries[REMARKS_AMOUNT] = rng.randint(5, 50)
        entries[REMARKS] = f"BENCH-UWAGI-{number:05d}"
    return entries


def _fill_entries(collective_file: Path) -> None:
    """Gotowość lustra: tata wpisuje ilości i obciążenia w ``N:AT`` każdego Wiersza danych Pliku zbiorczego."""

    rng = random.Random(SEED)
    data_rows = [row.number for row in read_data_rows(collective_file).rows]
    workbook = load_workbook(collective_file)
    try:
        sheet = workbook.active
        if not isinstance(sheet, Worksheet):
            raise BenchmarkError("Plik zbiorczy lustra nie ma aktywnego arkusza.")
        for number, row in enumerate(data_rows, start=1):
            for column, value in _entries(rng, number).items():
                cell = sheet.cell(row, column)
                if cell.data_type == "f":
                    raise BenchmarkError("Plik zbiorczy lustra ma formułę w kolumnach N:AT Wiersza danych.")
                cell.value = value
        save_workbook(workbook, collective_file)
    finally:
        workbook.close()


@dataclass(frozen=True)
class _ChildRun:
    exit_code: int
    spawned_at: datetime
    wall_ms: int
    stdout: str


def run_quietly(arguments: Sequence[str | Path], failure_log: Path) -> _ChildRun:
    """Uruchamia skrypt Pythona bez stdin i mierzy czas ściany procesu.

    Wyjście procesu zawiera nazwy Wykonawców, więc nie trafia na ekran, tylko do lokalnego logu przy błędzie.
    """

    spawned_at = datetime.now(UTC)
    started = time.perf_counter()
    completed = potomne.run([sys.executable, *map(str, arguments)], cwd=Path.cwd())
    wall_ms = round((time.perf_counter() - started) * 1000)
    if completed.returncode != 0:
        failure_log.parent.mkdir(parents=True, exist_ok=True)
        failure_log.write_text(completed.stdout + completed.stderr, encoding="utf-8")
    return _ChildRun(completed.returncode, spawned_at, wall_ms, completed.stdout)


def ensure_mirror(
    mirrors: Path, scenario_name: str, config: Path, templates: Mapping[str, Path]
) -> tuple[Path, bool]:
    """Lustro scenariusza i to, czy zostało właśnie zbudowane; budowa jest atomowa (folder roboczy + rename)."""

    mirror = mirrors / f"{scenario_name}-{_fingerprint(scenario_name, config, templates)}"
    if mirror.is_dir():
        return mirror, False
    mirrors.mkdir(parents=True, exist_ok=True)
    staging = Path(tempfile.mkdtemp(prefix=f".{mirror.name}.staging-", dir=mirrors))
    try:
        period_directory(staging).mkdir(parents=True)
        register = staging / REGISTER_NAME
        _write_register(register, _register_rows(SCENARIOS[scenario_name], _worker_ids(config)))
        with tempfile.TemporaryDirectory(prefix="benchmark-przygotowanie-") as scratch:
            preparation = run_quietly(
                [
                    ROOT / SCRIPT,
                    "preparation",
                    "--register",
                    register,
                    "--periods-directory",
                    staging / "Załączniki",
                    "--period",
                    PERIOD.token,
                    "--config",
                    config,
                    "--internal-template",
                    templates[INTERNAL_TEMPLATE],
                    "--collective-template",
                    templates[COLLECTIVE_TEMPLATE],
                    "--metrics",
                    Path(scratch) / "metrics.jsonl",
                ],
                mirrors / "ostatni-blad-przygotowania.txt",
            )
        if preparation.exit_code not in _PREPARED_EXITS:
            raise BenchmarkError(
                f"Przygotowanie pliku zbiorczego lustra zakończyło się kodem {preparation.exit_code}; "
                f"wyjście w {mirrors / 'ostatni-blad-przygotowania.txt'}."
            )
        _fill_entries(period_directory(staging) / collective_file_name(PERIOD))
        staging.rename(mirror)
    finally:
        shutil.rmtree(staging, ignore_errors=True)
    return mirror, True


def worker_settlements_folder(mirror: Path) -> Path:
    return PeriodFolder.from_collective_file(period_directory(mirror) / collective_file_name(PERIOD)).worker_settlements
