"""Pomiar Reguł ilości na Plikach zbiorczych wypełnionych przez tatę (spec #268). Tylko na Macu Maksa, poza CI.

Dla każdego Pliku zbiorczego z argumentów porównuje Ilości usług z Reguł ilości (``evaluate``) z wpisami taty w
kolumnach Reguł ilości (``QuantityColumn``). Wypisuje same liczby i teksty ``URZĄDZENIA``, bez adresów, numerów,
umów i opisów technika, więc wynik nie niesie danych klientów. Pliki są tylko czytane.

Porównanie jak w Kontroli ilości: puste i 0 to to samo, liczba zapisana tekstem to ta liczba, inna treść komórki
to „nieliczbowa”. AK i AM liczą się tylko tam, gdzie reguła daje w nich wartość różną od zera, bo poza tym tata
wpisuje je z opisu technika. Wiersze bez reguły porównują tylko V, T i N. Wiersze rozliczane z opisu nie mają
porównania: wynik podaje wpisy taty w ``N:AP``.

Klasa rozbieżności to zbiór różniących się kolumn (kolumna, reguła, tata) razem z tym, od czego zależy ich
wartość: dzień i godzina ``DATA`` dla N, liczba w ``CZYNNOŚCI`` dla V i AN, pierwsze słowo ``UMOWA`` dla T, a dla
AE, AP i AK wpis taty w Y (Last drop MDU z opisu technika).

    .venv/bin/python tools/pomiar_regul_ilosci.py "…/Rozliczenie 08_14_09_2026 - zbiorcze.xlsx" […]
"""

from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
import re
import sys

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

from rozliczenia.quantity_rules import FromDescription, OrderContent, QuantityColumn, RuleResult, ServiceQuantities, WeekTime, WithoutRule, evaluate
from rozliczenia.workbook_files import HEADER_ROW, QUANTITY_COLUMNS, DataRow, read_data_rows

NONNUMERIC = "nieliczbowa"
Value = int | str
"""Ilość z komórki taty: liczba albo ``NONNUMERIC``."""
Difference = tuple[QuantityColumn, int, Value]
"""Kolumna, wartość z reguły, wartość taty."""

_Q = QuantityColumn
_INPUTS: Mapping[str, str] = {"C": "UMOWA", "D": "URZĄDZENIA", "E": "CZYNNOŚCI", "G": "DATA"}
_ALL = tuple(get_column_letter(column) for column in QUANTITY_COLUMNS)
_ONLY_WHEN_RULED = frozenset({_Q.AK, _Q.AM})
_WITHOUT_RULE = (_Q.V, _Q.T, _Q.N)
_DAYS = ("pn", "wt", "śr", "cz", "pt", "sb", "nd")
_DIGIT = re.compile(r"[0-9]")
_SHORT_WORD = re.compile(r"[a-ząćęłńóśźż]{1,4}")


@dataclass(frozen=True)
class Row:
    """Wiersz danych Pliku zbiorczego: tekst ``URZĄDZENIA``, wynik reguł, wpisy taty w ``N:AP`` i kontekst reguł."""

    devices: str
    result: RuleResult
    tata: Mapping[str, Value]
    context: Mapping[QuantityColumn, str]

    def compares(self, column: QuantityColumn) -> bool:
        if isinstance(self.result, FromDescription) or column in _ONLY_WHEN_RULED and not self.ruled(column):
            return False
        return isinstance(self.result, ServiceQuantities) or column in _WITHOUT_RULE

    def ruled(self, column: QuantityColumn) -> int:
        return self.result.quantities.get(column, 0)

    def differences(self) -> tuple[Difference, ...]:
        return tuple(
            (column, self.ruled(column), self.tata[column])
            for column in QuantityColumn
            if self.compares(column) and self.ruled(column) != self.tata[column]
        )


def value(cell: object) -> Value:
    """Ilość z komórki taty jak w Kontroli ilości."""

    if cell is None or isinstance(cell, bool):
        return 0 if cell is None else NONNUMERIC
    if isinstance(cell, int):
        return cell
    if isinstance(cell, float):
        return int(cell) if cell.is_integer() else NONNUMERIC
    text = str(cell).strip().replace(",", ".")
    if not text:
        return 0
    try:
        number = float(text)
    except ValueError:
        return NONNUMERIC
    return int(number) if number.is_integer() else NONNUMERIC


def rows(path: Path) -> tuple[list[Row], int]:
    """Wiersze danych z wynikiem reguł i liczba wierszy, których ``DATA`` nie jest datą."""

    _check_header(path)
    measured: list[Row] = []
    undated = 0
    for row in read_data_rows(path).rows:
        started = _cell(row, "G")
        if not isinstance(started, datetime):
            undated += 1
            continue
        content = OrderContent(contract=_cell(row, "C"), devices=_cell(row, "D"), activities=_cell(row, "E"), started=WeekTime.of(started))
        tata = {letter: value(_cell(row, letter)) for letter in _ALL}
        devices = " ".join(str(_cell(row, "D") or "").split())
        measured.append(Row(devices, evaluate(content), tata, _context(row, started, tata)))
    return measured, undated


def _check_header(path: Path) -> None:
    workbook = load_workbook(path, read_only=True)
    try:
        sheet = workbook.active
        header = {letter: sheet[f"{letter}{HEADER_ROW}"].value for letter in _INPUTS} if sheet is not None else {}
    finally:
        workbook.close()
    if header != _INPUTS:
        raise SystemExit(f"{path.name}: nagłówki w wierszu {HEADER_ROW} to {header}, a nie {dict(_INPUTS)}")


def _cell(row: DataRow, letter: str) -> object:
    return row.values[column_index_from_string(letter) - 1]


def _context(row: DataRow, started: datetime, tata: Mapping[str, Value]) -> dict[QuantityColumn, str]:
    day = _DAYS[started.weekday()]
    slot = day if started.weekday() >= 5 else f"{day} {'od' if started.hour >= 17 else 'przed'} 17:00"
    activities = f"CZYNNOŚCI {'z liczbą' if _DIGIT.search(str(_cell(row, 'E') or '')) else 'bez liczby'}"
    words = str(_cell(row, "C") or "").casefold().split()
    contract = "puste" if not words else words[0] if _SHORT_WORD.fullmatch(words[0]) else "inne"
    last_drop = f"Y taty {'≠' if tata['Y'] != 0 else '='} 0"
    return {
        _Q.N: f"DATA {slot}",
        _Q.V: activities,
        _Q.AN: activities,
        _Q.T: f"UMOWA {contract}",
        _Q.AE: last_drop,
        _Q.AP: last_drop,
        _Q.AK: last_drop,
    }


def report(name: str, measured: list[Row], undated: int) -> list[str]:
    service = [row for row in measured if isinstance(row.result, ServiceQuantities)]
    without_rule = [row for row in measured if isinstance(row.result, WithoutRule)]
    from_description = [row for row in measured if isinstance(row.result, FromDescription)]
    lines = [
        f"== {name}",
        f"wiersze danych: {len(measured) + undated} | z regułą: {len(service)} | bez reguły: {len(without_rule)}"
        f" | z opisu: {len(from_description)} | DATA nie jest datą: {undated}",
        f"zgodne całe wiersze z regułą (AK, AM tylko gdy reguła ≠ 0): {_share([not row.differences() for row in service])}",
        f"zgodne całe wiersze z regułą, ściśle we wszystkich {len(QuantityColumn)} kolumnach: {_share([all(row.ruled(column) == row.tata[column] for column in QuantityColumn) for row in service])}",
        f"wiersze bez reguły zgodne w V, T, N: {_share([not row.differences() for row in without_rule])}",
        "-- kolumna: zgodne / porównane | reguła ≠ 0 | tata ≠ 0 (wiersze z regułą; V, T, N także bez reguły)",
    ]
    for column in QuantityColumn:
        compared = [row for row in measured if row.compares(column)]
        lines.append(
            f"{column}: {_share([row.ruled(column) == row.tata[column] for row in compared])}"
            f" | {sum(row.ruled(column) != 0 for row in compared)} | {sum(row.tata[column] != 0 for row in compared)}"
        )
    classes: dict[tuple[str, tuple[Difference, ...], tuple[str, ...]], Counter[str]] = defaultdict(Counter)
    for row in service + without_rule:
        found = row.differences()
        if found:
            kind = "z regułą" if isinstance(row.result, ServiceQuantities) else "bez reguły"
            context = tuple(sorted({row.context[column] for column, _, _ in found if column in row.context}))
            classes[(kind, found, context)][row.devices] += 1
    lines.append(f"-- klasy rozbieżności: {len(classes)}")
    for (kind, found, context), texts in sorted(classes.items(), key=lambda item: -item[1].total()):
        columns = "; ".join(f"{column}: reguła {ruled}, tata {tata}" for column, ruled, tata in found)
        lines.append(f"{texts.total()} wierszy {kind} | {columns}{''.join(f' | {part}' for part in context)} | {_texts(texts)}")
    lines.append("-- wiersze bez reguły według URZĄDZENIA: wiersze | wpisy taty poza V, T, N → wiersze")
    lines += _groups(without_rule, skip=_WITHOUT_RULE)
    lines.append("-- wiersze z opisu według URZĄDZENIA: wiersze | wpisy taty w N:AP → wiersze")
    lines += _groups(from_description, skip=())
    beyond = Counter(
        letter
        for row in service
        for letter in _ALL
        if row.tata[letter] != 0 and not (letter in QuantityColumn and row.compares(QuantityColumn(letter)))
    )
    lines.append(f"-- wiersze z regułą z wpisem taty w kolumnie spoza porównania: {_counts(beyond)}")
    return lines


def _share(matches: Sequence[bool]) -> str:
    part, whole = sum(matches), len(matches)
    return f"{part}/{whole} ({100 * part / whole:.1f}%)" if whole else "0/0"


def _texts(texts: Counter[str], limit: int = 12) -> str:
    shown = ", ".join(f"{count}× {text!r}" for text, count in texts.most_common(limit))
    return shown + (f" i {len(texts) - limit} innych tekstów" if len(texts) > limit else "")


def _counts(values: Mapping[str, object]) -> str:
    return " ".join(f"{letter}={values[letter]}" for letter in _ALL if letter in values and values[letter] != 0) or "nic"


def _groups(group: list[Row], skip: Sequence[str]) -> list[str]:
    by_text: dict[str, Counter[str]] = defaultdict(Counter)
    for row in group:
        by_text[row.devices][_counts({letter: entry for letter, entry in row.tata.items() if letter not in skip})] += 1
    return [
        f"{entries.total()}× {text!r} | " + ", ".join(f"{filled} → {count}" for filled, count in entries.most_common())
        for text, entries in sorted(by_text.items(), key=lambda item: -item[1].total())
    ]


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Pomiar Reguł ilości na Plikach zbiorczych wypełnionych przez tatę.")
    parser.add_argument("pliki", nargs="+", type=Path, help="Pliki zbiorcze wypełnione przez tatę")
    for path in parser.parse_args(argv).pliki:
        measured, undated = rows(path)
        print("\n".join(report(path.name, measured, undated)))
    return 0


if __name__ == "__main__":
    sys.exit(main())
