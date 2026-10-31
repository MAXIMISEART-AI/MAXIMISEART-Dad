"""Porównuje znaczeniowo dwa Foldery rozliczeń pracowników (spec #141); użyj jako bramki, gdy zmiana silnika ma dać te same rozliczenia.

Porównanie znaczeniowe dwóch Folderów rozliczeń pracowników: bramka poprawności wspinaczki (spec #141).

Użycie (z katalogu repozytorium):

    python tools/porownanie_rozliczen.py FOLDER_REFERENCYJNY FOLDER_KANDYDATA

Jedna operacja, ``compare_worker_settlements``: dla dwóch Folderów rozliczeń
pracowników zwraca liczbę różnic w każdej kategorii (``CATEGORIES``). Wynik to
same liczby: bez treści komórek, adresów komórek, nazw plików i numerów zleceń.

„Ten sam wynik” znaczy to, co widzi tata w Excelu:

- ``pliki``: plik tylko w jednym folderze albo plik, który nie otwiera się jako skoroszyt;
- ``wartosci`` i ``formuly``: każda komórka każdego arkusza; wspólne formuły są
  rozwinięte do postaci jawnej (robi to openpyxl przy wczytaniu), wartości
  zapamiętane przy formułach się nie liczą, bo Excel przelicza skoroszyt przy otwarciu;
- ``style``: rozwiązany styl komórki (czcionka, wypełnienie, obramowanie, wyrównanie,
  format liczb, ochrona), nie numer stylu; komórka nieobecna w XML ma styl wiersza,
  a bez niego styl kolumny, tak jak w Excelu;
- ``hiperlacza``: adres (``Target``) i miejsce (``location``) hiperłącza każdej komórki,
  nie numer relacji;
- ``scalenia``, ``wymiary`` (wysokość, ukrycie i poziom konspektu wierszy i kolumn),
  ``nazwy_zdefiniowane`` (nazwa arkusza w apostrofach i bez nich to to samo), ``external_links`` (bajty każdej części ``xl/externalLinks/*``)
  i ``flagi_przeliczania`` (``calcMode``, ``fullCalcOnLoad``, ``forceFullCalc``).

``xl/calcChain.xml`` nie jest porównywany.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import zipfile
from collections import Counter
from collections.abc import Hashable, Iterator, Mapping, Sequence
from concurrent.futures import ProcessPoolExecutor
from copy import copy
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any, Final
from xml.etree import ElementTree

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.styles.cell_style import StyleArray
from openpyxl.utils.cell import column_index_from_string
from openpyxl.worksheet.dimensions import Dimension
from openpyxl.worksheet.worksheet import Worksheet

CATEGORIES: Final = (
    "pliki",
    "wartosci",
    "formuly",
    "style",
    "hiperlacza",
    "scalenia",
    "wymiary",
    "nazwy_zdefiniowane",
    "external_links",
    "flagi_przeliczania",
)

_MAIN_NS: Final = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_EXTERNAL_LINKS: Final = "xl/externalLinks/"
_CALC_FLAGS: Final[dict[str, str | bool]] = {"calcMode": "auto", "fullCalcOnLoad": False, "forceFullCalc": False}

_StyleSource = Cell | MergedCell | Dimension
# Nazwa arkusza w apostrofach w odwołaniu (``'Sheet1'!$A$1``); openpyxl dodaje je przy zapisie, Excel nie.
_QUOTED_SHEET: Final = re.compile(r"'((?:[^']|'')+)'!")


def compare_worker_settlements(reference: Path, candidate: Path) -> dict[str, int]:
    """Liczba różnic w każdej kategorii między Folderem referencyjnym a Folderem kandydata."""

    reference_files = _file_names(reference)
    candidate_files = _file_names(candidate)
    differences = Counter({"pliki": len(reference_files ^ candidate_files)})
    pairs = [(reference / name, candidate / name) for name in sorted(reference_files & candidate_files)]
    for result in _compare_all(pairs):
        differences.update(result)
    return {category: differences[category] for category in CATEGORIES}


def _file_names(folder: Path) -> set[str]:
    return {path.name for path in folder.iterdir() if path.is_file()} if folder.is_dir() else set()


def _compare_all(pairs: Sequence[tuple[Path, Path]]) -> Iterator[dict[str, int]]:
    """Pary skoroszytów porównywane równolegle: pełne wczytanie skoroszytu trwa sekundy."""

    if len(pairs) < 2:
        yield from (_compare_workbooks(*pair) for pair in pairs)
        return
    with ProcessPoolExecutor(max_workers=min(len(pairs), os.cpu_count() or 1)) as pool:
        yield from pool.map(_compare_workbooks, *zip(*pairs, strict=True))


def _compare_workbooks(reference: Path, candidate: Path) -> dict[str, int]:
    try:
        reference_parts = _parts(reference)
        candidate_parts = _parts(candidate)
        reference_book = load_workbook(reference)
        candidate_book = load_workbook(candidate)
    except Exception:  # noqa: BLE001 - skoroszyt, którego nie da się otworzyć, liczy się jako różnica pliku
        return {"pliki": 1}
    differences = Counter(
        {
            "nazwy_zdefiniowane": _mapping_difference(_defined_names(reference_parts), _defined_names(candidate_parts)),
            "external_links": _mapping_difference(_external_links(reference_parts), _external_links(candidate_parts)),
            "flagi_przeliczania": _mapping_difference(_calc_flags(reference_parts), _calc_flags(candidate_parts)),
        }
    )
    styles = _Styles()
    titles = dict.fromkeys([*reference_book.sheetnames, *candidate_book.sheetnames])
    for title in titles:
        differences.update(
            _compare_sheets(_sheet(reference_book, title), _sheet(candidate_book, title), styles)
        )
    reference_book.close()
    candidate_book.close()
    return dict(differences)


def _parts(path: Path) -> dict[str, bytes]:
    """Części archiwum potrzebne poza openpyxl: ``xl/workbook.xml`` i ``xl/externalLinks/*``."""

    with zipfile.ZipFile(path) as archive:
        return {
            name: archive.read(name)
            for name in archive.namelist()
            if name == "xl/workbook.xml" or name.startswith(_EXTERNAL_LINKS)
        }


def _workbook_element(parts: dict[str, bytes], tag: str) -> list[ElementTree.Element]:
    data = parts.get("xl/workbook.xml")
    return [] if data is None else list(ElementTree.fromstring(data).iter(f"{_MAIN_NS}{tag}"))


def _flag(value: str | None, default: str | bool) -> str | bool:
    if value is None:
        return default
    if isinstance(default, bool):
        return value.strip().lower() in ("1", "true")
    return value


def _defined_names(parts: dict[str, bytes]) -> dict[tuple[str | None, str | None], tuple[str | bool, str]]:
    """Nazwa zdefiniowana w swoim zakresie (skoroszyt albo arkusz) z ukryciem i definicją."""

    return {
        (element.get("name"), element.get("localSheetId")): (_flag(element.get("hidden"), False), _definition(element.text))
        for element in _workbook_element(parts, "definedName")
    }


def _definition(text: str | None) -> str:
    """Definicja nazwy znaczeniowo: ``'Sheet1'!$A$1`` i ``Sheet1!$A$1`` to to samo odwołanie."""

    return _QUOTED_SHEET.sub(lambda match: match.group(1).replace("''", "'") + "!", text or "")


def _calc_flags(parts: dict[str, bytes]) -> dict[str, str | bool]:
    elements = _workbook_element(parts, "calcPr")
    attributes = elements[0].attrib if elements else {}
    return {name: _flag(attributes.get(name), default) for name, default in _CALC_FLAGS.items()}




def _external_links(parts: dict[str, bytes]) -> dict[str, bytes]:
    return {name: data for name, data in parts.items() if name.startswith(_EXTERNAL_LINKS)}


def _mapping_difference(reference: Mapping[Any, object], candidate: Mapping[Any, object]) -> int:
    """Klucze, których wartość jest inna albo jest tylko po jednej stronie."""

    return sum(reference.get(key) != candidate.get(key) for key in reference.keys() | candidate.keys())


def _sheet(workbook: Workbook, title: str) -> Worksheet:
    """Arkusz o tej nazwie; brakujący arkusz porównuje się jak pusty arkusz bez wymiarów i scaleń."""

    if title in workbook.sheetnames:
        sheet = workbook[title]
        if isinstance(sheet, Worksheet):
            return sheet
    stand_in = Workbook().active
    assert isinstance(stand_in, Worksheet)
    return stand_in


class _Styles:
    """Rozwiązany styl jako mała liczba: ten sam numer dla równych stylów z obu skoroszytów pary."""

    def __init__(self) -> None:
        self._canonical: dict[tuple[Hashable, ...], int] = {}
        self._by_style_array: dict[tuple[int, tuple[int, ...]], int] = {}

    def of(self, source: _StyleSource, sheet: Worksheet) -> int:
        """Numer rozwiązanego stylu komórki albo wymiaru; openpyxl trzyma styl jako indeksy list skoroszytu."""

        # Prywatne pole openpyxl, bez typu w stubach; komórka utworzona przy wczytaniu hiperłącza dostaje je dopiero
        # przy pierwszym odczycie stylu, a do tego czasu ma styl domyślny.
        style_array: Sequence[int] = getattr(source, "_style") or StyleArray()
        key = (id(sheet), tuple(style_array))
        found = self._by_style_array.get(key)
        if found is None:
            resolved = (
                copy(source.font),
                copy(source.fill),
                copy(source.border),
                copy(source.alignment),
                source.number_format,
                copy(source.protection),
            )
            found = self._canonical.setdefault(resolved, len(self._canonical))
            self._by_style_array[key] = found
        return found


def _compare_sheets(reference: Worksheet, candidate: Worksheet, styles: _Styles) -> Counter[str]:
    differences: Counter[str] = Counter()
    reference_columns = _column_dimensions(reference)
    candidate_columns = _column_dimensions(candidate)
    for coordinate in reference._cells.keys() | candidate._cells.keys():
        reference_cell = reference._cells.get(coordinate)
        candidate_cell = candidate._cells.get(coordinate)
        reference_formula, candidate_formula = _formula(reference_cell), _formula(candidate_cell)
        if reference_formula is not None or candidate_formula is not None:
            differences["formuly"] += reference_formula != candidate_formula
        else:
            differences["wartosci"] += _value(reference_cell) != _value(candidate_cell)
        differences["style"] += _style(reference, reference_cell, coordinate, reference_columns, styles) != _style(
            candidate, candidate_cell, coordinate, candidate_columns, styles
        )
        differences["hiperlacza"] += _hyperlink(reference_cell) != _hyperlink(candidate_cell)
    differences["scalenia"] = len(
        {str(merged) for merged in reference.merged_cells.ranges}
        ^ {str(merged) for merged in candidate.merged_cells.ranges}
    )
    differences["wymiary"] = _dimension_difference(reference, candidate, reference_columns, candidate_columns)
    return differences


def _formula(cell: Cell | MergedCell | None) -> tuple[str, ...] | None:
    if cell is None or cell.data_type != "f":
        return None
    value: Any = cell.value
    if isinstance(value, str):
        return ("formuła", value)
    text = getattr(value, "text", None)
    if text is not None:  # formuła tablicowa
        return ("tablicowa", str(getattr(value, "ref", "")), str(text))
    return ("tabela danych", *(f"{key}={item}" for key, item in sorted(vars(value).items())))


def _hyperlink(cell: Cell | MergedCell | None) -> tuple[str | None, str | None] | None:
    link = getattr(cell, "hyperlink", None)
    return None if link is None else (link.target, link.location)


def _value(cell: Cell | MergedCell | None) -> tuple[str, object] | None:
    """Wartość znaczeniowo: 3 i 3.0 to ta sama liczba, ale prawda to nie 1."""

    value: object = None if cell is None else cell.value
    if value is None:
        return None
    if isinstance(value, bool):
        return ("logiczna", value)
    if isinstance(value, int | float):
        return ("liczba", float(value))
    if isinstance(value, datetime | date | time | timedelta):
        return ("data", value)
    return (type(value).__name__, value)


def _column_dimensions(sheet: Worksheet) -> dict[int, Dimension]:
    """Wymiar każdej kolumny, także tych z zakresu ``min``–``max`` jednego wpisu ``<col>``."""

    columns: dict[int, Dimension] = {}
    for letter, dimension in sheet.column_dimensions.items():
        first = dimension.min or column_index_from_string(letter)
        last = dimension.max or first
        for index in range(first, last + 1):
            columns[index] = dimension
    return columns


def _style(
    sheet: Worksheet,
    cell: Cell | MergedCell | None,
    coordinate: tuple[int, int],
    columns: dict[int, Dimension],
    styles: _Styles,
) -> int:
    if cell is not None:
        return styles.of(cell, sheet)
    row_dimension = sheet.row_dimensions.get(coordinate[0])
    if row_dimension is not None and row_dimension.has_style:
        return styles.of(row_dimension, sheet)
    column_dimension = columns.get(coordinate[1])
    if column_dimension is not None and column_dimension.has_style:
        return styles.of(column_dimension, sheet)
    return styles.of(Cell(sheet, row=1, column=1, style_array=StyleArray()), sheet)


def _dimension_difference(
    reference: Worksheet,
    candidate: Worksheet,
    reference_columns: dict[int, Dimension],
    candidate_columns: dict[int, Dimension],
) -> int:
    def row(sheet: Worksheet, index: int) -> tuple[object, ...]:
        dimension = sheet.row_dimensions.get(index)
        return (None, False, 0) if dimension is None else (dimension.ht, bool(dimension.hidden), dimension.outlineLevel)

    def column(columns: dict[int, Dimension], index: int) -> tuple[object, ...]:
        dimension = columns.get(index)
        if dimension is None:
            return (None, False, 0)
        return (getattr(dimension, "width", None), bool(dimension.hidden), dimension.outlineLevel)

    rows = reference.row_dimensions.keys() | candidate.row_dimensions.keys()
    differing_rows = sum(row(reference, index) != row(candidate, index) for index in rows)
    indices = reference_columns.keys() | candidate_columns.keys()
    differing_columns = sum(column(reference_columns, index) != column(candidate_columns, index) for index in indices)
    return differing_rows + differing_columns


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("reference", type=Path, help="Folder rozliczeń pracowników z wyniku referencyjnego.")
    parser.add_argument("candidate", type=Path, help="Folder rozliczeń pracowników kandydata.")
    args = parser.parse_args(argv)
    differences = compare_worker_settlements(args.reference, args.candidate)
    print(json.dumps(differences, ensure_ascii=False))
    return 0 if not any(differences.values()) else 1


if __name__ == "__main__":
    raise SystemExit(main())

