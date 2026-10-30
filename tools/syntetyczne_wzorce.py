"""Zapisuje syntetyczne Wzorce w kształcie Wzorców taty; użyj, gdy benchmark, weryfikacja w Excelu albo CI nie ma prawdziwych Wzorców.

Użycie (z katalogu repozytorium):

    python tools/syntetyczne_wzorce.py FOLDER [--data-rows N]

Do ``FOLDER`` trafiają trzy Wzorce pod nazwami, pod którymi leżą w ``config`` (folder nadaje się
dla ``tools/benchmark_przebiegu.py --templates-dir``), oraz ``worker_mapping.yaml``: Konfiguracja
20 syntetycznych Wykonawców. Wzorców taty nie ma w repozytorium, a CI potrzebuje skoroszytów,
których wczytanie i zapis kosztują tyle samo.

Wzorzec rozliczenia pracownika i Wzorzec pliku zbiorczego mają kształt Wzorców taty (spec #141):
nagłówki w wierszu 17, ~10 tys. sformatowanych Wierszy danych (18..10093), ~500 tys. komórek,
~50 tys. formuł zapisanych jako wspólne (wzór co 64 wiersze) z zapisanymi wartościami,
``calcChain`` z każdą komórką z formułą, ``sharedStrings``, jedno odwołanie zewnętrzne
(``xl/externalLinks/*``), scalenia, pasy stylów i 483 wiersze zer w N:V i X. Nie wszystkie
wiersze mają komórki w A i B, jak u taty. Wzorzec rozliczenia pracownika ma, jak u taty,
komentarz komórki, więc relacje arkusza (``rId1`` rysunek VML, ``rId2`` komentarze). openpyxl nie zapisuje wspólnych formuł ani
``calcChain``, więc generator pisze XML części zip sam. Wszystkie trzy Wzorce otwierają się, jak
Wzorce taty od #465, w lewym górnym rogu arkusza (górna lewa komórka i zaznaczenie A1) w powiększeniu 75%.

Etykiety A:K, AQ:AY i opisy wierszy 13–16 są takie jak u taty; usługi N:AP, stawki, formuły
i odwołanie zewnętrzne są syntetyczne. Generator jest deterministyczny: te same argumenty dają
te same bajty, więc lustro benchmarku zbudowane z nich ma stały odcisk.
"""

from __future__ import annotations

import argparse
import sys
import zipfile
from collections.abc import Iterable, Sequence
from dataclasses import dataclass
from pathlib import Path
from xml.sax.saxutils import escape

import yaml
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from rozliczenia.collective_preparation import (
    default_collective_template_path,
    default_internal_template_path,
)
from rozliczenia.engine import default_worker_template_path
from rozliczenia.workbook_files import DATA_ROW_COLUMNS, FIRST_DATA_ROW, HEADER_ROW, QUANTITY_COLUMNS
from rozliczenia.worker_config import default_config_path

# Nazwy, pod którymi Wzorce i Konfiguracja wykonawców leżą w ``config``: te same, których szukają procesy.
WORKER_TEMPLATE = default_worker_template_path().name
COLLECTIVE_TEMPLATE = default_collective_template_path().name
INTERNAL_TEMPLATE = default_internal_template_path().name
WORKER_CONFIG = default_config_path().name

# Wiersze danych Wzorca taty: 18..10093.
FULL_DATA_ROWS = 10_076
WORKERS = 20

_HEADER_ROW = HEADER_ROW
_FIRST_DATA_ROW = FIRST_DATA_ROW
# Wzorzec taty: stałe 0 w N:V i X w wierszach 18..500.
_ZERO_ROWS = 483
# Excel zapisuje wzór wspólnej formuły kolumny co 64 wiersze.
_SHARED_FORMULA_ROWS = 64
# Pusty wiersz z wysokością daleko pod danymi, jak w Wzorcu taty (10145 przy 10093).
_TRAILING_ROW_OFFSET = 52

_FIXED_TIME = (1980, 1, 1, 0, 0, 0)
_MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_XML = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'

_COLUMN = {index: get_column_letter(index) for index in range(1, 54)}
# Ilości N:AP i koniec Wiersza danych AT z układu procesu; pozostałe kolumny to szczegóły Wzorca taty.
_N, _AP, _AT = QUANTITY_COLUMNS[0], QUANTITY_COLUMNS[-1], DATA_ROW_COLUMNS[-1]
_V, _W, _X, _AU, _AY, _BA = 22, 23, 24, 47, 51, 53

_INPUT_HEADERS = (
    "ADRES",
    "UMOWA",
    "URZĄDZENIA",
    "CZYNNOŚCI",
    "NUMER ZL.",
    "DATA",
    "WYKONAWCA",
    "NUMER UMOWY/ZLECENIA ORANGE",
    "SKUTECZNE",
    "CZYNNOŚCI ZREALIZOWANE",
)
_SUMMARY_HEADERS = (
    "Obciążenie klienta - opis",
    "Obciążenie klienta - kwota",
    "Uwagi - kwota",
    "Uwagi",
    "Kwota",
    "Montaż urządzeń",
    "Wykonanie instalacji",
    "Budowa Przyłącza",
    "Obszar",
)
_SERVICE_HEADERS = tuple(f"Usługa syntetyczna {number:02d}" for number in range(1, _AP - _N + 2))
_INTERNAL_HEADERS = ("MIASTO", *_INPUT_HEADERS, "zwrot", "UWAGI")


# --- style ------------------------------------------------------------------------------------


class _Styles:
    """``xl/styles.xml`` z kilkudziesięcioma stylami komórek, jak Wzorzec taty (74 style)."""

    def __init__(self) -> None:
        self._fonts = ['<font><sz val="11"/><color rgb="FF000000"/><name val="Calibri"/><family val="2"/></font>']
        self._fills = ['<fill><patternFill patternType="none"/></fill>', '<fill><patternFill patternType="gray125"/></fill>']
        self._borders = ["<border><left/><right/><top/><bottom/><diagonal/></border>"]
        self._xfs: list[str] = ['<xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>']

    @staticmethod
    def _index(items: list[str], item: str) -> int:
        if item not in items:
            items.append(item)
        return items.index(item)

    def add(
        self,
        *,
        size: int = 11,
        bold: bool = False,
        color: str = "FF000000",
        fill: str | None = None,
        border: str | None = None,
        number_format: int = 0,
        horizontal: str | None = None,
        wrap: bool = False,
    ) -> int:
        font = (
            f'<font>{"<b/>" if bold else ""}<sz val="{size}"/><color rgb="{color}"/>'
            '<name val="Calibri"/><family val="2"/><charset val="238"/></font>'
        )
        fill_xml = (
            '<fill><patternFill patternType="none"/></fill>'
            if fill is None
            else f'<fill><patternFill patternType="solid"><fgColor rgb="{fill}"/><bgColor indexed="64"/></patternFill></fill>'
        )
        side = "" if border is None else f' style="{border}"><color indexed="64"/'
        border_xml = (
            "<border><left/><right/><top/><bottom/><diagonal/></border>"
            if border is None
            else f"<border><left{side}></left><right{side}></right><top{side}></top><bottom{side}></bottom><diagonal/></border>"
        )
        font_id = self._index(self._fonts, font)
        fill_id = self._index(self._fills, fill_xml)
        border_id = self._index(self._borders, border_xml)
        alignment = ""
        if horizontal is not None or wrap:
            attributes = (f' horizontal="{horizontal}"' if horizontal else "") + (' wrapText="1"' if wrap else "")
            alignment = f"<alignment{attributes}/>"
        applied = ' applyNumberFormat="1"' if number_format else ""
        xf = (
            f'<xf numFmtId="{number_format}" fontId="{font_id}" fillId="{fill_id}" borderId="{border_id}" xfId="0"'
            f'{applied} applyFont="1" applyFill="1" applyBorder="1"'
            + (f' applyAlignment="1">{alignment}</xf>' if alignment else "/>")
        )
        return self._index(self._xfs, xf)

    def xml(self) -> str:
        def block(tag: str, items: Sequence[str]) -> str:
            return f'<{tag} count="{len(items)}">{"".join(items)}</{tag}>'

        return (
            f'{_XML}<styleSheet xmlns="{_MAIN_NS}">'
            '<numFmts count="1"><numFmt numFmtId="164" formatCode="#,##0.00\\ &quot;zł&quot;"/></numFmts>'
            f"{block('fonts', self._fonts)}{block('fills', self._fills)}{block('borders', self._borders)}"
            '<cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>'
            f"{block('cellXfs', self._xfs)}"
            '<cellStyles count="1"><cellStyle name="Normalny" xfId="0" builtinId="0"/></cellStyles>'
            '<dxfs count="0"/><tableStyles count="0"/></styleSheet>'
        )


class _SharedStrings:
    def __init__(self) -> None:
        self._strings: list[str] = []
        self._uses = 0

    def cell(self, reference: str, text: str, style: int) -> str:
        if text not in self._strings:
            self._strings.append(text)
        self._uses += 1
        return f'<c r="{reference}" s="{style}" t="s"><v>{self._strings.index(text)}</v></c>'

    def xml(self) -> str:
        items = "".join(f"<si><t>{escape(text)}</t></si>" for text in self._strings)
        return f'{_XML}<sst xmlns="{_MAIN_NS}" count="{self._uses}" uniqueCount="{len(self._strings)}">{items}</sst>'


# --- arkusz Wzorca z wierszem 17 --------------------------------------------------------------


@dataclass(frozen=True)
class _Bands:
    """Pasy stylów Wierszy danych w proporcjach Wzorca taty (u taty 18–31, 32–608, 609, 610+)."""

    last_row: int
    zero_end: int
    middle_end: int
    side_rows: range

    @classmethod
    def for_rows(cls, data_rows: int) -> _Bands:
        last_row = _HEADER_ROW + data_rows
        side_start = _HEADER_ROW + data_rows * 28 // 100
        side_end = _HEADER_ROW + data_rows * 33 // 100
        return cls(
            last_row=last_row,
            zero_end=_HEADER_ROW + min(_ZERO_ROWS, data_rows // 2),
            middle_end=_HEADER_ROW + min(591, data_rows * 6 // 10),
            side_rows=range(side_start, side_end + 1),
        )


class _TemplateSheet:
    """XML arkusza ``Sheet1`` i ``calcChain`` z każdą komórką z formułą tego arkusza."""

    def __init__(self, first_header: str, data_rows: int, styles: _Styles, strings: _SharedStrings) -> None:
        self._bands = _Bands.for_rows(data_rows)
        self._first_header = first_header
        self._styles = styles
        self._strings = strings
        self._formula_cells: list[str] = []
        self._shared_index = 0
        self._style = self._define_styles()

    def _define_styles(self) -> dict[str, int]:
        add = self._styles.add
        return {
            "hidden": add(size=10),
            "label": add(size=11, bold=True),
            "rate": add(size=11, number_format=2, fill="FFFFF2CC", border="thin"),
            "rate-sum": add(size=11, number_format=164, fill="FFFFF2CC", border="thin"),
            "count": add(size=11, number_format=1, fill="FFE2EFDA", border="thin"),
            "subtotal": add(size=12, bold=True, number_format=164, border="medium", horizontal="center"),
            "block": add(size=14, bold=True, fill="FFD9E1F2", border="medium", horizontal="center"),
            "block-service": add(size=11, bold=True, fill="FFFCE4D6", border="thin", horizontal="center"),
            "header": add(size=10, bold=True, fill="FF00AA55", border="thin", horizontal="center", wrap=True),
            "header-input": add(size=10, bold=True, fill="FF00AA55", border="medium", wrap=True),
            "header-service": add(size=9, bold=True, fill="FFFCE4D6", border="thin", wrap=True),
            "header-summary": add(size=10, bold=True, fill="FFD9E1F2", border="thin", wrap=True),
            "top": add(size=10, fill="FFDDDDDD", border="thin"),
            "top-address": add(size=10, fill="FFDDDDDD", border="thin", wrap=True),
            "top-order": add(size=10, fill="FFDDDDDD", border="thin", horizontal="center"),
            "top-date": add(size=10, fill="FFDDDDDD", border="thin", number_format=14),
            "top-worker": add(size=10, bold=True, fill="FFDDDDDD", border="thin"),
            "top-flag": add(size=10, fill="FFDDDDDD", border="thin", horizontal="center"),
            "middle": add(size=10, border="thin"),
            "middle-marked": add(size=10, fill="FFFFFF00", border="thin"),
            "middle-tail": add(size=10, border="hair"),
            "middle-date": add(size=10, border="thin", number_format=14),
            "seam": add(size=10, color="FFC00000", border="medium"),
            "seam-date": add(size=10, color="FFC00000", border="medium", number_format=14),
            "bottom": add(size=10, border="hair"),
            "bottom-text": add(size=10, border="hair", wrap=True),
            "bottom-date": add(size=10, border="hair", number_format=14),
            "zero": add(size=11, number_format=1, fill="FFCCE5FF", border="thin", horizontal="center"),
            "zero-last": add(size=11, number_format=1, fill="FFCCE5FF", border="medium", horizontal="center"),
            "quantity": add(size=11, number_format=1, border="hair", horizontal="center"),
            "quantity-last": add(size=11, number_format=1, border="medium", horizontal="center"),
            "separator": add(size=11, fill="FF808080"),
            "extra": add(size=11, number_format=2, border="hair"),
            "amount": add(size=11, bold=True, number_format=164, border="thin"),
            "amount-part": add(size=11, number_format=164, border="thin"),
            "side": add(size=11, color="FF7F7F7F", number_format=2),
        }

    def _cell(self, column: int, row: int, style: str) -> str:
        return f'<c r="{_COLUMN[column]}{row}" s="{self._style[style]}"/>'

    def _number(self, column: int, row: int, style: str, value: float) -> str:
        return f'<c r="{_COLUMN[column]}{row}" s="{self._style[style]}"><v>{value:g}</v></c>'

    def _text(self, column: int, row: int, style: str, text: str) -> str:
        return self._strings.cell(f"{_COLUMN[column]}{row}", text, self._style[style])

    def _formula(self, column: int, row: int, style: str | None, formula: str, *, error: bool = False) -> str:
        reference = f"{_COLUMN[column]}{row}"
        self._formula_cells.append(reference)
        attributes = (f' s="{self._style[style]}"' if style else "") + (' t="e"' if error else "")
        value = "#N/A" if error else "0"
        return f'<c r="{reference}"{attributes}><f>{escape(formula)}</f><v>{value}</v></c>'

    def _shared(self, column: int, row: int, style: str, formula: str | None, span: str, si: int) -> str:
        """Komórka wspólnej formuły: wzór z zakresem albo samo odwołanie do wzoru ``si``."""

        reference = f"{_COLUMN[column]}{row}"
        self._formula_cells.append(reference)
        body = f'<f t="shared" ref="{span}" si="{si}">{escape(formula)}</f>' if formula else f'<f t="shared" si="{si}"/>'
        return f'<c r="{reference}" s="{self._style[style]}">{body}<v>0</v></c>'

    def _next_si(self) -> int:
        self._shared_index += 1
        return self._shared_index - 1

    @staticmethod
    def _row(number: int, cells: Iterable[str], attributes: str = "") -> str:
        return f'<row r="{number}"{attributes} x14ac:dyDescent="0.25">{"".join(cells)}</row>'

    # wiersze 1–17

    def _top_rows(self) -> list[str]:
        last = self._bands.last_row
        rows = [
            self._row(number, (self._cell(column, number, "hidden") for column in range(1, _X + 1)), ' spans="1:50" hidden="1"')
            for number in range(1, 12)
        ]
        rows.append(
            self._row(
                12,
                [self._cell(13, 12, "label"), self._number(_N, 12, "rate", 100)]
                + [self._cell(column, 12, "rate") for column in range(_N + 1, _AP)],
                ' spans="1:50" ht="34.5" customHeight="1"',
            )
        )
        services = range(_N, _AP + 1)
        sum_si = self._next_si()
        row_13 = [self._text(13, 13, "label", "Suma")]
        row_13 += [
            self._shared(column, 13, "rate-sum", "N14*N15" if column == _N else None, "N13:AP13", sum_si)
            for column in services
        ]
        row_13 += [self._cell(column, 13, "rate") for column in range(_AP + 1, _AT + 1)]
        row_13 += [
            self._formula(column, 13, "subtotal", f"SUBTOTAL(109,{_COLUMN[column]}18:{_COLUMN[column]}{last})")
            for column in range(_AU, _AU + 4)
        ]
        rows.append(self._row(13, row_13, ' spans="1:50"'))
        rates = [self._number(column, 14, "rate", 5 + (column * 7) % 40 + (column % 3) / 2) for column in services]
        rows.append(
            self._row(
                14,
                [self._text(13, 14, "label", "Koszt"), *rates]
                + [self._cell(column, 14, "rate") for column in range(_AP + 1, _AY)],
                ' spans="1:50"',
            )
        )
        first_si, second_si = self._next_si(), self._next_si()
        row_15 = [self._text(13, 15, "label", "Ilość"), self._formula(_N, 15, "count", f"SUM(N18:N{last})")]
        for column in range(_N + 1, _AP + 1):
            si, master, span = (first_si, _N + 1, f"O15:{_COLUMN[_X - 1]}15") if column < _X else (second_si, _X, "X15:AP15")
            letter = _COLUMN[column]
            row_15.append(self._shared(column, 15, "count", f"SUM({letter}18:{letter}{last})" if column == master else None, span, si))
        row_15 += [self._cell(column, 15, "count") for column in range(_AP + 1, _AY)]
        rows.append(self._row(15, row_15, ' spans="1:50"'))
        row_16 = [self._text(1, 16, "block", "Dane wejściowe")] + [self._cell(column, 16, "block") for column in range(2, 14)]
        row_16 += [self._text(_N, 16, "block-service", "Instalacje")]
        row_16 += [self._cell(column, 16, "block-service") for column in range(_N + 1, _AU + 1)]
        rows.append(self._row(16, row_16, ' spans="1:50" ht="19.5" thickBot="1"'))
        header = [self._text(1, 17, "header-input", self._first_header)]
        header += [self._text(column, 17, "header-input", label) for column, label in enumerate(_INPUT_HEADERS, start=2)]
        header += [self._cell(12, 17, "header"), self._cell(13, 17, "header")]
        header += [self._text(column, 17, "header-service", label) for column, label in enumerate(_SERVICE_HEADERS, start=_N)]
        header += [self._text(column, 17, "header-summary", label) for column, label in enumerate(_SUMMARY_HEADERS, start=_AP + 1)]
        rows.append(self._row(17, header, ' spans="1:51" s="1" customFormat="1" ht="41.25" customHeight="1"'))
        return rows

    # Wiersze danych

    def _band(self, row: int) -> str:
        if row <= 31:
            return "top"
        if row <= self._bands.middle_end:
            return "middle"
        return "seam" if row == self._bands.middle_end + 1 else "bottom"

    def _input_cells(self, row: int) -> list[str]:
        band = self._band(row)
        tall = row <= self._bands.middle_end + 1
        styles = {
            "top": ("top", "top-address", "top", "top-order", "top-date", "top-worker", "top", "top-flag", "top-flag", "top"),
            "middle": ("middle", "middle", "middle", "middle", "middle-date", "middle", "middle", "middle", "middle", "middle-tail"),
            "seam": ("seam", "seam", "seam", "seam", "seam-date", "seam", "seam", "seam", "seam", "seam"),
            "bottom": ("bottom", "bottom", "bottom-text", "bottom", "bottom-date", "bottom", "bottom", "bottom", "bottom-text", "bottom"),
        }[band]
        cells = []
        if tall or row % 97 == 0:
            cells.append(self._cell(1, row, styles[0]))
        if tall or row % 33 == 0:
            marked = band == "middle" and row % 29 == 0
            cells.append(self._cell(2, row, "middle-marked" if marked else styles[0]))
        cells += [self._cell(column, row, styles[column - 3]) for column in range(3, 12)]
        cells += [self._cell(12, row, styles[9]), self._cell(13, row, styles[9])]
        return cells

    def _quantity_cells(self, row: int) -> list[str]:
        zero = row <= self._bands.zero_end
        cells = []
        for column in range(_N, _AT + 1):
            if column == _W:
                cells.append(self._cell(column, row, "separator"))
            elif zero and (column <= _V or column == _X):
                cells.append(self._number(column, row, "zero-last" if column == _V and row > 19 else "zero", 0))
            elif column <= _X:
                cells.append(self._cell(column, row, "quantity-last" if column == _V else "quantity"))
            else:
                cells.append(self._cell(column, row, "extra"))
        return cells

    def _formula_cells_of(self, row: int, masters: dict[int, int]) -> list[str]:
        """AU:AY: AU18 zwykła, potem wspólne grupy po 64 wiersze w AU:AX, AY zwykła z odwołaniem zewnętrznym."""

        terms = "+".join(f"{_COLUMN[column]}{{row}}*${_COLUMN[column]}$14" for column in range(_N, _AP + 1))
        formulas = {
            _AU: terms + "+AS{row}*0.5+AR{row}*0.5",
            _AU + 1: "IFERROR(SUMPRODUCT(N{row}:V{row},$N$14:$V$14)+Y{row}*$Y$14+AA{row}*$AA$14,0)",
            _AU + 2: "IFERROR(SUMPRODUCT(Z{row}:AH{row},$Z$14:$AH$14)+AS{row}*0.5+AB{row},0)",
            _AU + 3: "IFERROR(SUMPRODUCT(AI{row}:AP{row},$AI$14:$AP$14)+X{row}+W{row},0)",
        }
        cells = []
        for column, formula in formulas.items():
            style = "amount" if column == _AU else "amount-part"
            if column == _AU and row == _FIRST_DATA_ROW:
                cells.append(self._formula(column, row, style, formula.format(row=row)))
                continue
            start = _FIRST_DATA_ROW + (column == _AU)
            if (row - start) % _SHARED_FORMULA_ROWS == 0:
                masters[column] = self._next_si()
                end = min(row + _SHARED_FORMULA_ROWS - 1, self._bands.last_row)
                letter = _COLUMN[column]
                span = f"{letter}{row}:{letter}{end}"
                cells.append(self._shared(column, row, style, formula.format(row=row), span, masters[column]))
            else:
                cells.append(self._shared(column, row, style, None, "", masters[column]))
        cells.append(
            self._formula(
                _AY, row, None, f"IFERROR(VLOOKUP(E{row},[1]Słownik!E:F,2,0),VLOOKUP(A{row},[1]Słownik!A:B,2,0))", error=True
            )
        )
        return cells

    def _data_rows(self) -> list[str]:
        masters: dict[int, int] = {}
        rows = []
        for row in range(_FIRST_DATA_ROW, self._bands.last_row + 1):
            cells = self._input_cells(row) + self._quantity_cells(row) + self._formula_cells_of(row, masters)
            side = row in self._bands.side_rows
            if side:
                cells += [self._cell(column, row, "side") for column in (_AY + 1, _BA)]
            first = 1 if f'r="A{row}"' in cells[0] else 2 if f'r="B{row}"' in cells[0] else 3
            rows.append(self._row(row, cells, f' spans="{first}:{_BA if side else _AY}"'))
        trailing = self._bands.last_row + _TRAILING_ROW_OFFSET
        rows.append(f'<row r="{trailing}" ht="13.5" customHeight="1" x14ac:dyDescent="0.25"/>')
        return rows

    def parts(self, *, comments: bool) -> tuple[str, str]:
        """XML arkusza i ``calcChain``: łańcuch powstaje z formuł zapisanych w arkuszu."""

        last = self._bands.last_row
        widths = [6.7, 36.6, 5.6, 17.4, 18.6, 12.1, 10.4, 18.7, 9, 8.4, 56.1, 3.7, 6.3]
        columns = "".join(
            f'<col min="{index}" max="{index}" width="{width}" style="{self._style["hidden"]}" customWidth="1"/>'
            for index, width in enumerate(widths, start=1)
        )
        columns += f'<col min="14" max="22" width="8.7" style="{self._style["hidden"]}" customWidth="1"/>'
        columns += '<col min="23" max="23" width="2.7" customWidth="1"/><col min="24" max="46" width="8.7" customWidth="1"/>'
        columns += '<col min="47" max="50" width="15.7" customWidth="1"/><col min="51" max="53" width="10.9" customWidth="1"/>'
        rows = self._top_rows() + self._data_rows()
        merges = ["AU13:AU15", "AV13:AV15", "AW13:AW15", "AX13:AX15", "A16:M16"]
        sheet_xml = (
            f'{_XML}<worksheet xmlns="{_MAIN_NS}" xmlns:r="{_REL_NS}" '
            'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" mc:Ignorable="x14ac" '
            'xmlns:x14ac="http://schemas.microsoft.com/office/spreadsheetml/2009/9/ac">'
            f'<dimension ref="A1:BA{last + _TRAILING_ROW_OFFSET}"/>'
            '<sheetViews><sheetView tabSelected="1" topLeftCell="A1" zoomScale="75" zoomScaleNormal="75" '
            'workbookViewId="0"><selection activeCell="A1" sqref="A1"/></sheetView></sheetViews>'
            '<sheetFormatPr defaultRowHeight="15" x14ac:dyDescent="0.25"/>'
            f"<cols>{columns}</cols><sheetData>{''.join(rows)}</sheetData>"
            '<sheetProtection selectLockedCells="1" selectUnlockedCells="1"/>'
            f'<protectedRanges><protectedRange sqref="B18:D{last}" name="Zakres1"/></protectedRanges>'
            f'<autoFilter ref="A17:AY17"/><mergeCells count="{len(merges)}">'
            + "".join(f'<mergeCell ref="{merge}"/>' for merge in merges)
            + '</mergeCells><pageMargins left="0.7" right="0.7" top="0.75" bottom="0.75" header="0.3" footer="0.3"/>'
            + ('<legacyDrawing r:id="rId1"/>' if comments else "")
            + "</worksheet>"
        )
        # Jak w ``calcChain`` Excela: ``i`` (arkusz) podany raz, kolejne komórki go dziedziczą.
        chain = [f'<c r="{self._formula_cells[0]}" i="1"/>'] + [f'<c r="{cell}"/>' for cell in self._formula_cells[1:]]
        return sheet_xml, f'{_XML}<calcChain xmlns="{_MAIN_NS}">{"".join(chain)}</calcChain>'


# --- części pakietu ---------------------------------------------------------------------------


def _content_types(*, linked: bool, comments: bool = False) -> str:
    overrides = {
        "/xl/workbook.xml": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
        "/xl/worksheets/sheet1.xml": "application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml",
        "/xl/styles.xml": "application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml",
        "/xl/sharedStrings.xml": "application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml",
        "/docProps/core.xml": "application/vnd.openxmlformats-package.core-properties+xml",
        "/docProps/app.xml": "application/vnd.openxmlformats-officedocument.extended-properties+xml",
    }
    if linked:
        overrides["/xl/externalLinks/externalLink1.xml"] = (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.externalLink+xml"
        )
        overrides["/xl/calcChain.xml"] = "application/vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml"
    if comments:
        overrides["/xl/comments1.xml"] = "application/vnd.openxmlformats-officedocument.spreadsheetml.comments+xml"
    vml = '<Default Extension="vml" ContentType="application/vnd.openxmlformats-officedocument.vmlDrawing"/>'
    return (
        f'{_XML}<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        + (vml if comments else "")
        + '<Default Extension="xml" ContentType="application/xml"/>'
        + "".join(f'<Override PartName="{name}" ContentType="{kind}"/>' for name, kind in overrides.items())
        + "</Types>"
    )


def _relationships(targets: Sequence[tuple[str, str]], *, external: bool = False) -> str:
    mode = ' TargetMode="External"' if external else ""
    items = "".join(
        f'<Relationship Id="rId{index}" Type="{kind}" Target="{target}"{mode}/>'
        for index, (kind, target) in enumerate(targets, start=1)
    )
    return f'{_XML}<Relationships xmlns="{_PACKAGE_REL_NS}">{items}</Relationships>'


_OFFICE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PACKAGE_RELS = _relationships(
    [
        (f"{_OFFICE}/officeDocument", "xl/workbook.xml"),
        ("http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties", "docProps/core.xml"),
        (f"{_OFFICE}/extended-properties", "docProps/app.xml"),
    ]
)
_CORE = (
    f'{_XML}<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
    'xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/" '
    'xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
    "<dc:creator>Generator syntetycznych Wzorców</dc:creator>"
    '<dcterms:created xsi:type="dcterms:W3CDTF">2026-01-01T00:00:00Z</dcterms:created>'
    '<dcterms:modified xsi:type="dcterms:W3CDTF">2026-01-01T00:00:00Z</dcterms:modified>'
    "</cp:coreProperties>"
)


def _app(sheet_name: str) -> str:
    return (
        f'{_XML}<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
        'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
        "<Application>Microsoft Excel</Application><DocSecurity>0</DocSecurity><ScaleCrop>false</ScaleCrop>"
        f'<TitlesOfParts><vt:vector size="1" baseType="lpstr"><vt:lpstr>{sheet_name}</vt:lpstr></vt:vector></TitlesOfParts>'
        "<AppVersion>16.0300</AppVersion></Properties>"
    )


def _workbook(sheet_name: str, filter_range: str, *, linked: bool) -> str:
    external = '<externalReferences><externalReference r:id="rId2"/></externalReferences>' if linked else ""
    return (
        f'{_XML}<workbook xmlns="{_MAIN_NS}" xmlns:r="{_REL_NS}">'
        '<fileVersion appName="xl" lastEdited="7" lowestEdited="6" rupBuild="27928"/><workbookPr/>'
        '<bookViews><workbookView xWindow="0" yWindow="0" windowWidth="29040" windowHeight="15720"/></bookViews>'
        f'<sheets><sheet name="{sheet_name}" sheetId="1" r:id="rId1"/></sheets>{external}'
        f'<definedNames><definedName name="_xlnm._FilterDatabase" localSheetId="0" hidden="1">{filter_range}'
        '</definedName></definedNames><calcPr calcId="191029"/></workbook>'
    )


_EXTERNAL_LINK = (
    f'{_XML}<externalLink xmlns="{_MAIN_NS}"><externalBook xmlns:r="{_REL_NS}" r:id="rId1">'
    '<sheetNames><sheetName val="Słownik"/></sheetNames><sheetDataSet><sheetData sheetId="0">'
    + "".join(
        f'<row r="{row}"><cell r="A{row}" t="str"><v>Firma syntetyczna {row}</v></cell>'
        f'<cell r="B{row}" t="str"><v>Obszar {row % 3 + 1}</v></cell>'
        f'<cell r="E{row}" t="str"><v>Powiat syntetyczny {row}</v></cell>'
        f'<cell r="F{row}" t="str"><v>Obszar {row % 3 + 1}</v></cell></row>'
        for row in range(1, 41)
    )
    + "</sheetData></sheetDataSet></externalBook></externalLink>"
)
_EXTERNAL_LINK_RELS = _relationships(
    [(f"{_OFFICE}/externalLinkPath", "Syntetyczny%20skoroszyt%20zewnetrzny.xlsm")], external=True
)


# Komentarz komórki nagłówka, jak we Wzorcu taty: część komentarzy i rysunek VML, który Excel pokazuje jako notatkę.
_COMMENTS = (
    f'{_XML}<comments xmlns="{_MAIN_NS}"><authors><author>Autor syntetyczny</author></authors><commentList>'
    '<comment ref="H17" authorId="0"><text><r><t>Identyfikator z Konfiguracji wykonawców</t></r></text></comment>'
    "</commentList></comments>"
)
_VML_DRAWING = (
    '<xml xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office" '
    'xmlns:x="urn:schemas-microsoft-com:office:excel"><o:shapelayout v:ext="edit"><o:idmap v:ext="edit" data="1"/>'
    '</o:shapelayout><v:shapetype id="_x0000_t202" coordsize="21600,21600" o:spt="202" path="m,l,21600r21600,l21600,xe">'
    '<v:stroke joinstyle="miter"/><v:path gradientshapeok="t" o:connecttype="rect"/></v:shapetype>'
    '<v:shape id="_x0000_s1025" type="#_x0000_t202" style="position:absolute;margin-left:59.25pt;margin-top:1.5pt;'
    'width:108pt;height:59.25pt;z-index:1;visibility:hidden" fillcolor="#ffffe1" o:insetmode="auto">'
    '<v:fill color2="#ffffe1"/><v:shadow on="t" color="black" obscured="t"/><v:path o:connecttype="none"/>'
    '<v:textbox style="mso-direction-alt:auto"><div style="text-align:left"></div></v:textbox>'
    '<x:ClientData ObjectType="Note"><x:MoveWithCells/><x:SizeWithCells/><x:Anchor>8, 15, 15, 2, 10, 15, 19, 1</x:Anchor>'
    "<x:AutoFill>False</x:AutoFill><x:Row>16</x:Row><x:Column>7</x:Column></x:ClientData></v:shape></xml>"
)
_SHEET_RELS = _relationships(
    [(f"{_OFFICE}/vmlDrawing", "../drawings/vmlDrawing1.vml"), (f"{_OFFICE}/comments", "../comments1.xml")]
)


def _write_package(path: Path, parts: Sequence[tuple[str, str]]) -> None:
    """Zip z częściami w podanej kolejności i stałym czasem, więc te same części dają te same bajty."""

    path.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(path, "w") as archive:
        for name, content in parts:
            info = zipfile.ZipInfo(name, date_time=_FIXED_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            archive.writestr(info, content.encode("utf-8"))


def write_settlement_template(
    path: Path, first_header: str, data_rows: int = FULL_DATA_ROWS, *, comments: bool = False
) -> None:
    """Wzorzec z nagłówkami w wierszu 17: rozliczenia pracownika (``GODZINA``) albo pliku zbiorczego (``MIASTO``).

    ``comments`` dodaje komentarz komórki H17 z relacjami arkusza, jak we Wzorcu rozliczenia pracownika taty.
    """

    styles, strings = _Styles(), _SharedStrings()
    sheet_xml, calc_chain_xml = _TemplateSheet(first_header, data_rows, styles, strings).parts(comments=comments)
    workbook_rels = _relationships(
        [
            (f"{_OFFICE}/worksheet", "worksheets/sheet1.xml"),
            (f"{_OFFICE}/externalLink", "externalLinks/externalLink1.xml"),
            (f"{_OFFICE}/styles", "styles.xml"),
            (f"{_OFFICE}/sharedStrings", "sharedStrings.xml"),
            (f"{_OFFICE}/calcChain", "calcChain.xml"),
        ]
    )
    commented = (
        [
            ("xl/comments1.xml", _COMMENTS),
            ("xl/drawings/vmlDrawing1.vml", _VML_DRAWING),
            ("xl/worksheets/_rels/sheet1.xml.rels", _SHEET_RELS),
        ]
        if comments
        else []
    )
    _write_package(
        path,
        [
            ("[Content_Types].xml", _content_types(linked=True, comments=comments)),
            ("_rels/.rels", _PACKAGE_RELS),
            ("xl/workbook.xml", _workbook("Sheet1", "Sheet1!$A$17:$AY$17", linked=True)),
            ("xl/_rels/workbook.xml.rels", workbook_rels),
            ("xl/worksheets/sheet1.xml", sheet_xml),
            ("xl/styles.xml", styles.xml()),
            ("xl/sharedStrings.xml", strings.xml()),
            ("xl/externalLinks/externalLink1.xml", _EXTERNAL_LINK),
            ("xl/calcChain.xml", calc_chain_xml),
            ("docProps/core.xml", _CORE),
            ("docProps/app.xml", _app("Sheet1")),
            ("xl/externalLinks/_rels/externalLink1.xml.rels", _EXTERNAL_LINK_RELS),
            *commented,
        ],
    )


def write_internal_template(path: Path) -> None:
    """Wzorzec rozliczenia wewnętrznego: arkusz ``Arkusz1``, nagłówki w wierszu 1, ~260 sformatowanych wierszy."""

    styles, strings = _Styles(), _SharedStrings()
    header = styles.add(size=10, bold=True, fill="FFD9E1F2", border="thin", wrap=True)
    plain = styles.add(size=10, border="thin")
    date = styles.add(size=10, border="thin", number_format=22)
    wide = styles.add(size=10, border="thin", wrap=True)
    last = 263
    rows = [
        '<row r="1" spans="1:13" ht="30" customHeight="1">'
        + "".join(strings.cell(f"{_COLUMN[column]}1", label, header) for column, label in enumerate(_INTERNAL_HEADERS, start=1))
        + "</row>"
    ]
    column_styles = [plain, wide, plain, plain, plain, plain, date, plain, plain, plain, wide, plain, wide]
    for row in range(2, last + 1):
        cells = "".join(f'<c r="{_COLUMN[column]}{row}" s="{style}"/>' for column, style in enumerate(column_styles, start=1))
        rows.append(f'<row r="{row}" spans="1:13">{cells}</row>')
    widths = [12.2, 38, 7.8, 18.2, 11.5, 9.5, 11.5, 22, 27.8, 9.7, 52.3, 5.5, 46.2]
    columns = "".join(
        f'<col min="{index}" max="{index}" width="{width}" customWidth="1"/>' for index, width in enumerate(widths, start=1)
    )
    sheet_xml = (
        f'{_XML}<worksheet xmlns="{_MAIN_NS}" xmlns:r="{_REL_NS}"><dimension ref="A1:M{last}"/>'
        '<sheetViews><sheetView tabSelected="1" topLeftCell="A1" zoomScale="75" zoomScaleNormal="75" workbookViewId="0">'
        '<pane ySplit="1" topLeftCell="A2" activePane="bottomLeft" state="frozen"/>'
        '<selection pane="bottomLeft" activeCell="A1" sqref="A1"/>'
        '</sheetView></sheetViews><sheetFormatPr defaultRowHeight="15"/>'
        f"<cols>{columns}</cols><sheetData>{''.join(rows)}</sheetData>"
        f'<autoFilter ref="A1:M{last}"/>'
        '<pageMargins left="0.7" right="0.7" top="0.75" bottom="0.75" header="0.3" footer="0.3"/></worksheet>'
    )
    workbook_rels = _relationships(
        [
            (f"{_OFFICE}/worksheet", "worksheets/sheet1.xml"),
            (f"{_OFFICE}/styles", "styles.xml"),
            (f"{_OFFICE}/sharedStrings", "sharedStrings.xml"),
        ]
    )
    _write_package(
        path,
        [
            ("[Content_Types].xml", _content_types(linked=False)),
            ("_rels/.rels", _PACKAGE_RELS),
            ("xl/workbook.xml", _workbook("Arkusz1", f"Arkusz1!$A$1:$M${last}", linked=False)),
            ("xl/_rels/workbook.xml.rels", workbook_rels),
            ("xl/worksheets/sheet1.xml", sheet_xml),
            ("xl/styles.xml", styles.xml()),
            ("xl/sharedStrings.xml", strings.xml()),
            ("docProps/core.xml", _CORE),
            ("docProps/app.xml", _app("Arkusz1")),
        ],
    )


def write_worker_config(path: Path) -> None:
    """Konfiguracja wykonawców z 20 syntetycznymi Wykonawcami (tyle Szablonów pracowników ma tata)."""

    workers = {f"wykonawca.syntetyczny{number:02d}": f"Wykonawca Syntetyczny {number:02d}" for number in range(1, WORKERS + 1)}
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        yaml.safe_dump({"schema_version": 1, "workers": workers}, allow_unicode=True, sort_keys=False), encoding="utf-8"
    )


def write_synthetic_templates(directory: Path, data_rows: int = FULL_DATA_ROWS) -> None:
    """Trzy Wzorce pod nazwami z ``config`` i Konfiguracja wykonawców w ``directory``."""

    write_settlement_template(directory / WORKER_TEMPLATE, "GODZINA", data_rows, comments=True)
    write_settlement_template(directory / COLLECTIVE_TEMPLATE, "MIASTO", data_rows)
    write_internal_template(directory / INTERNAL_TEMPLATE)
    write_worker_config(directory / WORKER_CONFIG)


def _parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("directory", type=Path, help="Folder na Wzorce i Konfigurację wykonawców.")
    parser.add_argument(
        "--data-rows",
        type=int,
        default=FULL_DATA_ROWS,
        help=f"Liczba sformatowanych Wierszy danych (domyślnie {FULL_DATA_ROWS}, jak u taty).",
    )
    return parser


def main(argv: list[str] | None = None) -> int:
    args = _parser().parse_args(argv)
    if args.data_rows < 2 * _SHARED_FORMULA_ROWS:
        print(f"--data-rows musi wynosić co najmniej {2 * _SHARED_FORMULA_ROWS}.")
        return 2
    write_synthetic_templates(args.directory.expanduser().resolve(), args.data_rows)
    print(f"Syntetyczne Wzorce w {args.directory}.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
